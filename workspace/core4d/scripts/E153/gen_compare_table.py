"""
生成 gate 方法对比表格 (E152+E153 合并视角)
输出: workspace/core4d/report/gate_method_compare.xlsx

5 行 × 10 列 (方法 + 9 个指标)
  - baseline / gateA / b1            数据来源: E152 (reuse E148/E151)
  - gateA+b1  (-0.010 / 0.050)       数据来源: E152 fresh run
  - gateA+b1  (-0.010 / 0.100)       数据来源: E153 fresh run
"""
import pathlib
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = pathlib.Path(__file__).parents[2] / "report" / "gate_method_compare.xlsx"

# ── Style helpers ──────────────────────────────────────────────────────────────
def fill(hex6):
    return PatternFill("solid", fgColor="FF" + hex6)

def font(bold=False, color="000000", size=10, italic=False):
    return Font(bold=bold, color="FF" + color, size=size, italic=italic)

THIN   = Side(style="thin",   color="FF000000")
MEDIUM = Side(style="medium", color="FF000000")

def border(left="thin", right="thin", top="thin", bottom="thin"):
    s = {"thin": THIN, "medium": MEDIUM}
    return Border(left=s[left], right=s[right], top=s[top], bottom=s[bottom])

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

def set_cell(ws, row, col, value="", bold=False, fg="000000", bg=None,
             align=CENTER, size=10, italic=False, bord="thin"):
    c = ws.cell(row=row, column=col, value=value)
    c.font      = Font(bold=bold, color="FF" + fg, size=size, italic=italic)
    c.alignment = align
    if bg:
        c.fill = fill(bg)
    b = THIN if bord == "thin" else MEDIUM
    c.border = Border(left=b, right=b, top=b, bottom=b)
    return c

def merge_cell(ws, r1, c1, r2, c2, value="", bold=False, fg="000000", bg=None, size=11):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    c = ws.cell(row=r1, column=c1, value=value)
    c.font      = Font(bold=bold, color="FF" + fg, size=size)
    c.alignment = CENTER
    if bg:
        c.fill = fill(bg)
    c.border = Border(
        left=MEDIUM, right=MEDIUM, top=MEDIUM, bottom=MEDIUM
    )
    return c

def pct(v):
    return f"{v:.1%}" if v is not None else "—"

# ── Table definition ───────────────────────────────────────────────────────────
METRICS = [
    "5cm几何接触\n(↑更好)",
    "2mm几何穿透\n(↓更好)",
    "5mm几何穿透\n(↓更好)",
    "物体接触\n(↑更好)",
    "物理穿透>5mm\n(↓更好)",
    "腿穿透\n(↓更好)",
    "success\n(0mm | 2mm)",
    "fall",
    "fallback",
]

# +1 = higher is better, -1 = lower is better, None = no colour coding
DIRECTION = [+1, -1, -1, +1, -1, -1, None, None, None]

# Baseline reference values for colour thresholds
#            [5cm,   2mm,   5mm,   obj,   phy,   leg]
BASELINE_REF = [0.6150, 0.2940, 0.2053, 0.3491, 0.5531, 0.0450]

# ── Row data ──────────────────────────────────────────────────────────────────
# [method_label, 5cm, 2mm, 5mm, obj_contact, phy_pen, leg_pen,
#  success_str, fall_str, fallback_str, source_note]
ROWS = [
    # ── reference configurations ──────────────────────────────────────────────
    ("baseline",
     0.6150, 0.2940, 0.2053, 0.3491, 0.5531, 0.0450,
     "ref", "0/3", "0.000",
     "reuse E148"),

    ("gateA\n(min_sdf=-0.010\nmax_viol=0.050)",
     0.6062, 0.2688, 0.1451, 0.3109, 0.3467, 0.0311,
     "0/3 | 0/3", "0/3", "0.017",
     "E152 新跑\n(gate only, no b1 reward)"),

    ("b1",
     0.6939, 0.4711, 0.2727, 0.5649, 0.5492, 0.0375,
     "ref", "0/3", "0.000",
     "reuse E151"),

    # ── gate + b1 variants ────────────────────────────────────────────────────
    ("gateA + b1\n(min_sdf=-0.010\nmax_viol=0.050)",
     0.7130, 0.3998, 0.2180, 0.5334, 0.3817, 0.0222,
     "2/3 | 3/3", "0/3", "0.032",
     "E152 新跑"),

    ("gateA + b1\n(min_sdf=-0.010\nmax_viol=0.100)★",
     0.6939, 0.3638, 0.2183, 0.5480, 0.3767, 0.0222,
     "3/3 | 3/3", "0/3", "0.001",
     "E153 新跑\n(阈值扫描最优点)"),
]

# ── Build workbook ─────────────────────────────────────────────────────────────
wb  = openpyxl.Workbook()
ws  = wb.active
ws.title = "gate方法对比"

NCOLS     = 1 + len(METRICS)   # 10
HDR_BG    = "1F4E79"
SUB_BG    = "BDD7EE"
GOOD_BG   = "C6EFCE"; GOOD_FG = "276221"
BAD_BG    = "FFC7CE"; BAD_FG  = "9C0006"
NA_BG     = "D9D9D9"; NA_FG   = "595959"
REF_BG    = "FFF2CC"; REF_FG  = "7F6000"
METHOD_BG = ["E6E6E6", "DDEEFF"]   # alternating rows

# row 1: title
merge_cell(ws, 1, 1, 1, NCOLS,
           value="Hand-Object Gate 方法对比  (3 cases 均值: box021 / box004 / box023)",
           bold=True, fg="FFFFFF", bg="2F2F2F", size=13)
ws.row_dimensions[1].height = 30

# row 2: column headers
set_cell(ws, 2, 1, "方法", bold=True, bg=HDR_BG, fg="FFFFFF", size=11)
for i, label in enumerate(METRICS):
    set_cell(ws, 2, 2 + i, label, bold=True, bg=SUB_BG, size=9)
ws.row_dimensions[2].height = 50

# rows 3-7: data
for ri, row_data in enumerate(ROWS):
    r = 3 + ri
    (method, v5cm, v2mm, v5mm, vobj, vphy, vleg,
     success, fall, fallback, source) = row_data

    mbg = METHOD_BG[ri % 2]
    set_cell(ws, r, 1, method, bold=True, bg=mbg, align=LEFT, size=10)

    numeric_vals = [v5cm, v2mm, v5mm, vobj, vphy, vleg]
    str_vals     = [success, fall, fallback]

    # numeric metric columns (2-7)
    for ci, (val, direction) in enumerate(zip(numeric_vals, DIRECTION[:6])):
        col   = 2 + ci
        ref   = BASELINE_REF[ci]
        text  = pct(val)
        is_ref_row = (success == "ref")

        if is_ref_row:
            bg, fg = REF_BG, REF_FG
        elif direction == +1:
            bg, fg = (GOOD_BG, GOOD_FG) if val >= ref else (BAD_BG, BAD_FG)
        else:
            bg, fg = (GOOD_BG, GOOD_FG) if val <= ref else (BAD_BG, BAD_FG)

        set_cell(ws, r, col, text, bold=(bg == GOOD_BG), fg=fg, bg=bg, size=10)

    # string columns (8-10): success / fall / fallback
    for ci, val in enumerate(str_vals):
        col = 8 + ci
        if val == "ref":
            set_cell(ws, r, col, val, bg=REF_BG, fg=REF_FG, size=10)
        else:
            set_cell(ws, r, col, val, bg=mbg, size=10)

    ws.row_dimensions[r].height = 48

# footnote row
merge_cell(ws, 8, 1, 8, NCOLS,
           value="★ gateA+b1 (-0.010/0.100) 为 E153 阈值扫描最优点，"
                 "max_viol 放宽至 0.10 后 gate 有效率 ↑ 89.5%，3/3 cases 穿透改善",
           bold=False, fg="595959", bg="F7F7F7", size=9)
ws.row_dimensions[8].height = 22

# ── Column widths ──────────────────────────────────────────────────────────────
ws.column_dimensions["A"].width = 24
for c in range(2, NCOLS + 1):
    ws.column_dimensions[get_column_letter(c)].width = 13

ws.freeze_panes = "B3"

# ── 说明 sheet ─────────────────────────────────────────────────────────────────
ws2 = wb.create_sheet("说明")
notes = [
    ("指标",             "来源字段",                              "说明"),
    ("5cm几何接触",      "hand_geom_near_5cm_frac",              "手部几何体距物体表面≤5cm的帧比例 (↑更好)"),
    ("2mm几何穿透",      "hand_geom_penetration_2mm_frac",       "手部几何体与物体重叠≥2mm的帧比例 (↓更好)"),
    ("5mm几何穿透",      "hand_geom_penetration_5mm_frac",       "手部几何体与物体重叠≥5mm的帧比例 (↓更好)"),
    ("物体接触",         "hand_object_physics_contact_frac",     "手与物体存在MuJoCo物理接触的帧比例 (↑更好)"),
    ("物理穿透>5mm",     "hand_object_con_dist_frac_lt_neg5mm",  "物理接触距离<-5mm的帧比例 (↓更好)"),
    ("腿穿透",           "leg_penetration_frac",                 "腿部与物体穿透的帧比例 (↓更好)"),
    ("success (0mm|2mm)","—",
     "pen改善 且 contact不下降, 分别按0mm和2mm穿透阈值统计, X/3 cases"),
    ("fall",             "fall_flag",                            "摔倒的case数 / 3"),
    ("fallback",         "cem_gate_fallback_used_mean",
     "本轮CEM所有sample均违反gate时退化到无gate采样的帧比例均值"),
    ("",                 "",                                     ""),
    ("颜色",             "",                                     ""),
    ("绿色",             "",                                     "vs baseline改善"),
    ("红色",             "",                                     "vs baseline劣化"),
    ("黄色(ref)",        "",                                     "参考行 (baseline / b1), 不做颜色比较"),
    ("",                 "",                                     ""),
    ("数据来源",         "",                                     ""),
    ("baseline",         "reuse E148",                          "3 cases: box021_029_p2, box004_083_p2, box023_person2"),
    ("gateA",            "E152 新跑",                            "gate only (min_sdf=-0.010, max_viol=0.050), no b1 reward"),
    ("b1",               "reuse E151",                          "b1 hand surface contact reward only"),
    ("gateA+b1 /0.050",  "E152 新跑",                           "gate + b1, (min_sdf=-0.010, max_viol=0.050)"),
    ("gateA+b1 /0.100",  "E153 新跑",                           "gate + b1, (min_sdf=-0.010, max_viol=0.100), 阈值扫描最优"),
]
for i, (a, b, c) in enumerate(notes, 1):
    ws2.cell(row=i, column=1, value=a).font = Font(bold=(i == 1 or a in ("颜色", "数据来源")), size=10)
    ws2.cell(row=i, column=2, value=b)
    cl = ws2.cell(row=i, column=3, value=c)
    cl.alignment = Alignment(wrap_text=True)
ws2.column_dimensions["A"].width = 22
ws2.column_dimensions["B"].width = 38
ws2.column_dimensions["C"].width = 55

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"saved → {OUT}")
