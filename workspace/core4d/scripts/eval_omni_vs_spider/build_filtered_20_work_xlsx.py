#!/usr/bin/env python3
"""Build an xlsx-only filtered Spider work table.

Starting point is the separate 24-case work table definition:
11 ref_fk strict cases + 13 ref_fk upper-WORK/non-strict cases.
This script excludes requested case_ids and writes only one xlsx file.
"""

from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import REPO
from build_expanded_24_work_eval import (
    BETTER_FILL,
    COLOR_RULE_FIELDS,
    COLOR_RULE_ROWS,
    COMPARISON_COLOR_RULES,
    EXPANDED_FIELDS,
    METRIC_FIELDS,
    TIER_FIELDS,
    VALIDATION_FIELDS,
    build_comparison,
    comparison_status,
    current_main_items,
    evaluate_items,
    extra_upper_work_items,
    tier_summary,
)
from unified_replay_eval import method_summary, ppt_method_summary


DEFAULT_EXCLUDE_CASE_IDS = [
    "bucket004_20231003_1_012_p1",
    "e091_box026_20231020_134_p1",
    "e091_box026_20231020_134_p1",
    "e091_box026_20231020_141_p2",
    "e091_box026_20231023_139_p2",
]

EXCLUDE_FIELDS = ["字段", "值"]


def add_sheet(wb: Workbook, name: str, rows: list[dict[str, Any]], fields: list[str]):
    ws = wb.create_sheet(name)
    ws.append(fields)
    for row in rows:
        ws.append([row.get(field, "") for field in fields])
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, field in enumerate(fields, start=1):
        width = min(max(len(field) + 2, 10), 34)
        ws.column_dimensions[get_column_letter(idx)].width = width
    return ws


def apply_case_colors(ws) -> None:
    header_to_col = {cell.value: cell.column for cell in ws[1]}
    for row_idx in range(2, ws.max_row + 1):
        for spider_field, omni_field, direction, threshold in COMPARISON_COLOR_RULES:
            spider_col = header_to_col.get(spider_field)
            omni_col = header_to_col.get(omni_field)
            if not spider_col or not omni_col:
                continue
            spider_cell = ws.cell(row=row_idx, column=spider_col)
            omni_cell = ws.cell(row=row_idx, column=omni_col)
            status = comparison_status(spider_cell.value, omni_cell.value, direction, threshold)
            if status == "better":
                spider_cell.fill = BETTER_FILL
            elif status == "worse":
                spider_cell.fill = PatternFill("solid", fgColor="FFC7CE")


def write_xlsx_only(path: Path, comparison: list[dict[str, Any]], metrics: list[dict[str, Any]], summary: list[dict[str, Any]], tiers: list[dict[str, str]], validation: list[dict[str, Any]], exclude_rows: list[dict[str, str]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    case_ws = add_sheet(wb, "20case逐case对比", comparison, EXPANDED_FIELDS)
    apply_case_colors(case_ws)
    ppt = add_sheet(wb, "PPT方法汇总20", ppt_method_summary(summary), list(ppt_method_summary(summary)[0].keys()) if summary else [])
    for col, width in {
        "A": 12,
        "B": 4,
        "C": 7,
        "D": 9,
        "E": 8,
        "F": 8,
        "G": 8,
        "H": 8,
        "I": 8,
        "J": 10,
        "K": 8,
        "L": 9,
        "M": 8,
    }.items():
        ppt.column_dimensions[col].width = width
    add_sheet(wb, "方法汇总20", summary, list(summary[0].keys()) if summary else [])
    add_sheet(wb, "分层统计", tiers, TIER_FIELDS)
    legend = add_sheet(wb, "颜色规则", COLOR_RULE_ROWS, COLOR_RULE_FIELDS)
    legend_header = {cell.value: cell.column for cell in legend[1]}
    for row in range(2, legend.max_row + 1):
        legend.cell(row=row, column=legend_header["绿色含义"]).fill = BETTER_FILL
        legend.cell(row=row, column=legend_header["红色含义"]).fill = PatternFill("solid", fgColor="FFC7CE")
    add_sheet(wb, "排除说明", exclude_rows, EXCLUDE_FIELDS)
    add_sheet(wb, "完整method_metrics", metrics, METRIC_FIELDS)
    add_sheet(wb, "历史对齐校验", validation, VALIDATION_FIELDS)
    wb.save(path)


def build_exclude_rows(requested: list[str], matched: list[str], unmatched: list[str], num_before: int, num_after: int) -> list[dict[str, str]]:
    return [
        {"字段": "原始请求排除case_id", "值": ", ".join(requested)},
        {"字段": "去重后请求排除case_id", "值": ", ".join(dict.fromkeys(requested))},
        {"字段": "实际匹配并删除case_id", "值": ", ".join(matched)},
        {"字段": "未匹配case_id", "值": ", ".join(unmatched)},
        {"字段": "过滤前case数", "值": str(num_before)},
        {"字段": "过滤后case数", "值": str(num_after)},
        {"字段": "备注", "值": "`e091_box026_20231020_134_p1` 在请求中重复出现，实际只删除一条匹配行。"},
    ]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--existing-cases", type=Path, default=REPO / "workspace/core4d/data_construction_v3/existing_cases.tsv")
    parser.add_argument("--audit-tsv", type=Path, default=REPO / "workspace/core4d/results/E109/spider_cem_work_audit/spider_cem_work_candidates.tsv")
    parser.add_argument("--out-xlsx", type=Path, default=REPO / "workspace/core4d/results/E109/filtered_20_work_cases_xlsx/filtered_20_omni_vs_spider_work_eval.xlsx")
    args = parser.parse_args()

    main_items, main_meta = current_main_items(args.existing_cases)
    extra_items, extra_meta = extra_upper_work_items(args.audit_tsv)
    all_items = main_items + extra_items
    meta = {**main_meta, **extra_meta}
    requested_unique = list(dict.fromkeys(DEFAULT_EXCLUDE_CASE_IDS))
    requested_set = set(requested_unique)
    all_case_ids = [item["case"]["case_id"] for item in all_items]
    matched = [case_id for case_id in requested_unique if case_id in set(all_case_ids)]
    unmatched = [case_id for case_id in requested_unique if case_id not in set(all_case_ids)]

    filtered_items = [item for item in all_items if item["case"]["case_id"] not in requested_set]
    metric_rows, validation_rows = evaluate_items(filtered_items)
    comparison = build_comparison(metric_rows, validation_rows, meta)
    summary = method_summary(metric_rows)
    tiers = tier_summary(comparison)
    exclude_rows = build_exclude_rows(DEFAULT_EXCLUDE_CASE_IDS, matched, unmatched, len(all_items), len(filtered_items))

    args.out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    write_xlsx_only(args.out_xlsx, comparison, metric_rows, summary, tiers, validation_rows, exclude_rows)
    counts = Counter(row["分层"] for row in comparison)
    print(f"[build_filtered_20_work_xlsx] wrote {len(comparison)} cases to {args.out_xlsx}")
    print(f"[build_filtered_20_work_xlsx] tiers={dict(counts)} validation_mismatch={sum(row['status'] != 'PASS' for row in validation_rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
