#!/usr/bin/env python3
"""Unified replay evaluator for existing OmniRetarget vs Spider cases.

The script recomputes geometry/contact metrics from qpos + MuJoCo scenes, then
validates the recomputed Spider metrics against historical summaries when those
historical columns exist.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from collections import defaultdict
from pathlib import Path
from typing import Any

import mujoco
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import REPO, as_float, read_csv, resolve_path, write_json, write_tsv
from build_existing_cases_comparison import (
    repo_path_text,
    load_summary_row,
    resolve_omnirt_trajectory,
)


FPS = 30.0
FALL_PELVIS_Z_M = 0.45
EEF_OFFSET = np.asarray([0.05, 0.0, 0.0], dtype=np.float64)
NEAR_THRESHOLDS_M = (0.03, 0.05, 0.08, 0.10, 0.12, 0.15)
DEEP_PENETRATION_M = -0.02

LOWERBODY_GEOMS = [
    "left_hip_collision",
    "right_hip_collision",
    "left_thigh_collision",
    "right_thigh_collision",
    "left_shin_collision",
    "right_shin_collision",
    "left_linkage_brace_collision",
    "right_linkage_brace_collision",
    "lf0",
    "lf1",
    "lf2",
    "lf3",
    "rf0",
    "rf1",
    "rf2",
    "rf3",
]
HAND_GEOMS = ["lh", "rh"]
UPPER_BODY_NAMES = [
    "head_link",
    "torso_link",
    "pelvis_contour_link",
    "left_shoulder_yaw_link",
    "right_shoulder_yaw_link",
    "left_shoulder_pitch_link",
    "right_shoulder_pitch_link",
    "left_elbow_link",
    "right_elbow_link",
]

METRIC_FIELDS = [
    "case_id",
    "method",
    "object_key",
    "target_variant_id",
    "run_id",
    "qpos_path",
    "scene_xml",
    "qpos_frames",
    "duration_s",
    "pelvis_min_m",
    "pelvis_end_m",
    "fall_flag",
    "root_xy_displacement_m",
    "root_xy_path_m",
    "object_xy_displacement_m",
    "object_z_range_m",
    "object_floor_contact_frac",
    "eef_near_3cm_frac",
    "eef_near_5cm_frac",
    "eef_near_8cm_frac",
    "eef_near_10cm_frac",
    "eef_near_12cm_frac",
    "eef_near_15cm_frac",
    "hand_geom_near_3cm_frac",
    "hand_geom_near_5cm_frac",
    "hand_geom_near_8cm_frac",
    "hand_geom_near_10cm_frac",
    "hand_geom_near_12cm_frac",
    "hand_geom_near_15cm_frac",
    "hand_geom_penetration_frac",
    "hand_geom_deep_penetration_2cm_frac",
    "hand_object_physics_contact_frac",
    "leg_near_2cm_frac",
    "leg_penetration_frac",
    "leg_object_physics_contact_frac",
    "body_penetration_frac",
    "head_penetration_frac",
    "upper_body_penetration_frac",
    "obj_err_mean_m",
    "obj_err_max_m",
    "notes",
]

VALIDATION_FIELDS = [
    "case_id",
    "run_id",
    "historical_field",
    "unified_field",
    "historical_value",
    "unified_value",
    "abs_diff",
    "tolerance",
    "status",
    "explanation",
]

COMPARISON_FIELDS = [
    "case_id",
    "object_key",
    "target_variant_id",
    "spider_run_id",
    "omni_qpos_path",
    "spider_qpos_path",
    "Omni pelvis_min_m",
    "Spider pelvis_min_m",
    "Omni fall",
    "Spider fall",
    "Omni eef_near_3cm",
    "Spider eef_near_3cm",
    "Omni eef_near_5cm",
    "Spider eef_near_5cm",
    "Omni eef_near_8cm",
    "Spider eef_near_8cm",
    "Omni eef_near_10cm",
    "Spider eef_near_10cm",
    "Omni hand_penetration",
    "Spider hand_penetration",
    "Omni hand_deep_penetration_2cm",
    "Spider hand_deep_penetration_2cm",
    "Omni leg_penetration",
    "Spider leg_penetration",
    "Omni leg_near_2cm",
    "Spider leg_near_2cm",
    "Omni body_penetration",
    "Spider body_penetration",
    "Omni object_floor_contact",
    "Spider object_floor_contact",
    "Omni object_xy_displacement_m",
    "Spider object_xy_displacement_m",
    "Omni object_z_range_m",
    "Spider object_z_range_m",
    "Spider obj_err_mean_m",
    "Spider obj_err_max_m",
    "validation_status",
]

PPT_SUMMARY_FIELDS = [
    "方法",
    "N",
    "跌倒↓",
    "站立高度↑",
    "手15cm↑",
    "手12cm↑",
    "手10cm↑",
    "手8cm↑",
    "手5cm↑",
    "手深穿透↓",
    "手geom12cm↑",
    "手物理接触↑",
    "腿穿透↓",
    "身体穿透↓",
    "物体XY↑",
]


def finite_fmt(value: Any, digits: int = 6) -> str:
    v = as_float(value)
    return f"{v:.{digits}f}" if math.isfinite(v) else ""


def read_existing_cem_pass(path: Path) -> list[dict[str, str]]:
    return [
        row
        for row in read_csv(path, delimiter="\t")
        if row.get("cem_status") == "pass" and row.get("target_variant_id") != "adaptive"
    ]


def npz_qpos(path: Path, *, channel: str = "sim") -> tuple[np.ndarray | None, np.ndarray | None]:
    if not path.exists():
        return None, None
    data = np.load(path, allow_pickle=True)
    if "qpos" not in data:
        return None, None
    arr = np.asarray(data["qpos"], dtype=np.float64)
    if arr.ndim == 3 and arr.shape[1] >= 2:
        primary = arr[:, 0, :]
        ref = arr[:, 1, :]
        return primary, ref
    if arr.ndim == 2:
        return arr, None
    return None, None


def scene_for_qpos(qpos_path: Path, summary: dict[str, str], method: str) -> Path | None:
    if method == "omniretarget":
        parent_scene = qpos_path.parent.parent / "scene.xml"
        if parent_scene.exists():
            return parent_scene
        for key in ("scene_xml", "scene_used", "legobj_scene_used"):
            p = resolve_path(summary.get(key))
            if p and p.exists():
                scene = p.parent / "scene.xml"
                if scene.exists():
                    return scene
                return p
    else:
        for key in ("scene_xml", "scene_used", "legobj_scene_used"):
            p = resolve_path(summary.get(key))
            if p and p.exists():
                return p
    return None


def mj_id(model: mujoco.MjModel, obj_type: int, name: str) -> int:
    return int(mujoco.mj_name2id(model, obj_type, name))


def mj_name(model: mujoco.MjModel, obj_type: int, idx: int) -> str:
    return mujoco.mj_id2name(model, obj_type, idx) or f"{obj_type}:{idx}"


def object_collision_geoms(model: mujoco.MjModel) -> list[int]:
    gids = []
    for gid in range(model.ngeom):
        name = mj_name(model, mujoco.mjtObj.mjOBJ_GEOM, gid)
        if name.startswith("object_collision"):
            gids.append(gid)
    if not gids:
        raise ValueError("no object_collision* geoms in scene")
    return gids


def signed_point_box(point: np.ndarray, box_pos: np.ndarray, box_mat: np.ndarray, half: np.ndarray) -> float:
    local = box_mat.T @ (point - box_pos)
    q = np.abs(local) - half
    outside = np.linalg.norm(np.maximum(q, 0.0))
    inside = min(max(q[0], q[1], q[2]), 0.0)
    return float(outside + inside)


def point_object_sdf(model: mujoco.MjModel, data: mujoco.MjData, point: np.ndarray, object_gids: list[int]) -> float:
    vals = []
    for gid in object_gids:
        vals.append(
            signed_point_box(
                point,
                data.geom_xpos[gid].copy(),
                data.geom_xmat[gid].reshape(3, 3).copy(),
                model.geom_size[gid, :3].copy(),
            )
        )
    return float(min(vals))


def geom_object_sdf(model: mujoco.MjModel, data: mujoco.MjData, geom_id: int, object_gids: list[int]) -> float:
    geom_type = int(model.geom_type[geom_id])
    radius = float(model.geom_size[geom_id, 0])
    center = data.geom_xpos[geom_id].copy()
    mat = data.geom_xmat[geom_id].reshape(3, 3).copy()
    points = [center]
    if geom_type == int(mujoco.mjtGeom.mjGEOM_CAPSULE):
        half_len = float(model.geom_size[geom_id, 1])
        axis = mat[:, 2]
        points = [center + axis * s for s in np.linspace(-half_len, half_len, 9)]
    vals = []
    for obj_gid in object_gids:
        obj_pos = data.geom_xpos[obj_gid].copy()
        obj_mat = data.geom_xmat[obj_gid].reshape(3, 3).copy()
        half = model.geom_size[obj_gid, :3].copy()
        vals.append(min(signed_point_box(p, obj_pos, obj_mat, half) for p in points) - radius)
    return float(min(vals))


def quat_rotate(q_wxyz: np.ndarray, vec: np.ndarray) -> np.ndarray:
    out = np.zeros(3, dtype=np.float64)
    mujoco.mju_rotVecQuat(out, vec, q_wxyz)
    return out


def evaluate_sequence(
    *,
    case_id: str,
    method: str,
    object_key: str,
    target_variant_id: str,
    run_id: str,
    qpos_path: Path,
    scene_xml: Path,
) -> dict[str, Any]:
    qpos, ref_qpos = npz_qpos(qpos_path)
    if qpos is None:
        raise ValueError(f"missing qpos in {qpos_path}")
    model = mujoco.MjModel.from_xml_path(str(scene_xml))
    if qpos.shape[1] != model.nq:
        raise ValueError(f"qpos nq={qpos.shape[1]} != scene nq={model.nq}: {qpos_path} vs {scene_xml}")
    data = mujoco.MjData(model)
    object_gids = object_collision_geoms(model)
    floor_gid = mj_id(model, mujoco.mjtObj.mjOBJ_GEOM, "floor")
    object_body = mj_id(model, mujoco.mjtObj.mjOBJ_BODY, "object")
    left_wrist = mj_id(model, mujoco.mjtObj.mjOBJ_BODY, "left_wrist_yaw_link")
    right_wrist = mj_id(model, mujoco.mjtObj.mjOBJ_BODY, "right_wrist_yaw_link")
    lower_gids = [mj_id(model, mujoco.mjtObj.mjOBJ_GEOM, name) for name in LOWERBODY_GEOMS]
    lower_gids = [gid for gid in lower_gids if gid >= 0]
    hand_gids = [mj_id(model, mujoco.mjtObj.mjOBJ_GEOM, name) for name in HAND_GEOMS]
    hand_gids = [gid for gid in hand_gids if gid >= 0]
    upper_bodies = [mj_id(model, mujoco.mjtObj.mjOBJ_BODY, name) for name in UPPER_BODY_NAMES]
    upper_bodies = [bid for bid in upper_bodies if bid >= 0]
    head_body = mj_id(model, mujoco.mjtObj.mjOBJ_BODY, "head_link")

    pelvis_z = []
    root_xy = []
    obj_xyz = []
    eef_sdf_l = []
    eef_sdf_r = []
    hand_geom_sdf = []
    leg_sdf = []
    body_sdf = []
    head_sdf = []
    upper_sdf = []
    hand_physics = []
    leg_physics = []
    object_floor = []
    obj_err = []

    for i, q in enumerate(qpos):
        data.qpos[:] = q
        data.qvel[:] = 0.0
        mujoco.mj_forward(model, data)
        pelvis_z.append(float(data.xpos[mj_id(model, mujoco.mjtObj.mjOBJ_BODY, "pelvis"), 2]))
        root_xy.append(q[:2].copy())
        obj_xyz.append(data.xpos[object_body].copy())

        frame_eef = []
        for bid in (left_wrist, right_wrist):
            if bid < 0:
                frame_eef.append(math.nan)
                continue
            eef = data.xpos[bid].copy() + quat_rotate(data.xquat[bid].copy(), EEF_OFFSET)
            frame_eef.append(point_object_sdf(model, data, eef, object_gids))
        eef_sdf_l.append(float(frame_eef[0]) if len(frame_eef) > 0 else math.nan)
        eef_sdf_r.append(float(frame_eef[1]) if len(frame_eef) > 1 else math.nan)

        hand_vals = [geom_object_sdf(model, data, gid, object_gids) for gid in hand_gids]
        hand_geom_sdf.append(float(min(hand_vals)) if hand_vals else math.nan)
        leg_vals = [geom_object_sdf(model, data, gid, object_gids) for gid in lower_gids]
        leg_sdf.append(float(min(leg_vals)) if leg_vals else math.nan)
        body_vals = [point_object_sdf(model, data, data.xpos[bid].copy(), object_gids) for bid in upper_bodies]
        body_sdf.append(float(min(body_vals)) if body_vals else math.nan)
        if head_body >= 0:
            head_sdf.append(point_object_sdf(model, data, data.xpos[head_body].copy(), object_gids))
        else:
            head_sdf.append(math.nan)
        upper_vals = [point_object_sdf(model, data, data.xpos[bid].copy(), object_gids) for bid in upper_bodies if bid != head_body]
        upper_sdf.append(float(min(upper_vals)) if upper_vals else math.nan)

        hand_contact = False
        leg_contact = False
        floor_contact = False
        object_set = set(object_gids)
        for ci in range(data.ncon):
            con = data.contact[ci]
            pair = {int(con.geom1), int(con.geom2)}
            if not (object_set & pair):
                continue
            other = list(pair - object_set)
            if floor_gid >= 0 and floor_gid in pair:
                floor_contact = True
            if other:
                if other[0] in hand_gids:
                    hand_contact = True
                if other[0] in lower_gids:
                    leg_contact = True
        hand_physics.append(hand_contact)
        leg_physics.append(leg_contact)
        object_floor.append(floor_contact)

        if ref_qpos is not None and i < len(ref_qpos):
            data.qpos[:] = ref_qpos[i]
            data.qvel[:] = 0.0
            mujoco.mj_forward(model, data)
            obj_err.append(float(np.linalg.norm(obj_xyz[-1] - data.xpos[object_body].copy())))

    root_xy_arr = np.asarray(root_xy, dtype=np.float64)
    obj_xyz_arr = np.asarray(obj_xyz, dtype=np.float64)
    pelvis = np.asarray(pelvis_z, dtype=np.float64)
    eef_l = np.asarray(eef_sdf_l, dtype=np.float64)
    eef_r = np.asarray(eef_sdf_r, dtype=np.float64)
    hand = np.asarray(hand_geom_sdf, dtype=np.float64)
    leg = np.asarray(leg_sdf, dtype=np.float64)
    body = np.asarray(body_sdf, dtype=np.float64)
    head = np.asarray(head_sdf, dtype=np.float64)
    upper = np.asarray(upper_sdf, dtype=np.float64)
    obj_err_arr = np.asarray(obj_err, dtype=np.float64)

    def frac(mask: np.ndarray) -> float:
        if mask.size == 0:
            return math.nan
        return float(np.nanmean(mask))

    out: dict[str, Any] = {
        "case_id": case_id,
        "method": method,
        "object_key": object_key,
        "target_variant_id": target_variant_id,
        "run_id": run_id,
        "qpos_path": repo_path_text(qpos_path),
        "scene_xml": repo_path_text(scene_xml),
        "qpos_frames": int(qpos.shape[0]),
        "duration_s": float(qpos.shape[0] / FPS),
        "pelvis_min_m": float(np.nanmin(pelvis)),
        "pelvis_end_m": float(pelvis[-1]),
        "fall_flag": bool(np.nanmin(pelvis) < FALL_PELVIS_Z_M),
        "root_xy_displacement_m": float(np.linalg.norm(root_xy_arr[-1] - root_xy_arr[0])) if len(root_xy_arr) else math.nan,
        "root_xy_path_m": float(np.linalg.norm(np.diff(root_xy_arr, axis=0), axis=1).sum()) if len(root_xy_arr) > 1 else 0.0,
        "object_xy_displacement_m": float(np.linalg.norm(obj_xyz_arr[-1, :2] - obj_xyz_arr[0, :2])) if len(obj_xyz_arr) else math.nan,
        "object_z_range_m": float(np.nanmax(obj_xyz_arr[:, 2]) - np.nanmin(obj_xyz_arr[:, 2])) if len(obj_xyz_arr) else math.nan,
        "object_floor_contact_frac": frac(np.asarray(object_floor, dtype=bool)),
        "hand_geom_penetration_frac": frac(hand < 0.0),
        "hand_geom_deep_penetration_2cm_frac": frac(hand < DEEP_PENETRATION_M),
        "hand_object_physics_contact_frac": frac(np.asarray(hand_physics, dtype=bool)),
        "leg_near_2cm_frac": frac(leg < 0.02),
        "leg_penetration_frac": frac(leg < 0.0),
        "leg_object_physics_contact_frac": frac(np.asarray(leg_physics, dtype=bool)),
        "body_penetration_frac": frac(body < 0.0),
        "head_penetration_frac": frac(head < 0.0),
        "upper_body_penetration_frac": frac(upper < 0.0),
        "obj_err_mean_m": float(np.nanmean(obj_err_arr)) if obj_err_arr.size else math.nan,
        "obj_err_max_m": float(np.nanmax(obj_err_arr)) if obj_err_arr.size else math.nan,
        "notes": "unified_replay_sdf_v1",
    }
    for threshold in NEAR_THRESHOLDS_M:
        tag = int(round(threshold * 100))
        # Historical `contact_frac_either` is actually max(left_frac, right_frac),
        # not the union of both hands. Keep this field history-compatible.
        out[f"eef_near_{tag}cm_frac"] = max(frac(eef_l < threshold), frac(eef_r < threshold))
        out[f"hand_geom_near_{tag}cm_frac"] = frac(hand < threshold)
    return out


def read_case_inputs(existing_cases: Path) -> list[dict[str, Any]]:
    out = []
    for case in read_existing_cem_pass(existing_cases):
        summary, metrics_ref = load_summary_row(case)
        spider_qpos = resolve_path(case.get("cem_result_npz"))
        if not spider_qpos or not spider_qpos.exists():
            spider_qpos = resolve_path(summary.get("npz_path") or summary.get("root_npz_path"))
        omni_qpos = resolve_omnirt_trajectory(case, summary)
        if not spider_qpos or not omni_qpos:
            raise FileNotFoundError(f"missing qpos for {case.get('case_id')}")
        spider_scene = scene_for_qpos(spider_qpos, summary, "spider")
        omni_scene = scene_for_qpos(omni_qpos, summary, "omniretarget")
        if not spider_scene or not omni_scene:
            raise FileNotFoundError(f"missing scene for {case.get('case_id')}")
        out.append(
            {
                "case": case,
                "summary": summary,
                "metrics_ref": metrics_ref,
                "spider_qpos": spider_qpos,
                "spider_scene": spider_scene,
                "omni_qpos": omni_qpos,
                "omni_scene": omni_scene,
            }
        )
    return out


def compare_history(case: dict[str, str], summary: dict[str, str], unified: dict[str, Any]) -> list[dict[str, Any]]:
    pairs = [
        ("pelvis_min_m", "pelvis_min_m", 1e-6, "same qpos replay should match"),
        ("obj_err_mean_m", "obj_err_mean_m", 1e-6, "same qpos/ref object replay should match"),
        ("obj_err_max_m", "obj_err_max_m", 1e-6, "same qpos/ref object replay should match"),
        ("contact_frac_either", "eef_near_8cm_frac", 2e-3, "historical contact is wrist+5cm EEF near object proxy"),
        ("leg_box_interference_frac", "leg_penetration_frac", 1e-6, "E081/E105 lower-body SDF proxy"),
        ("leg_box_near_2cm_frac", "leg_near_2cm_frac", 1e-6, "E081/E105 lower-body 2cm proxy"),
        ("hand_object_contact_physics_frac", "hand_object_physics_contact_frac", 1e-6, "MuJoCo contact pair fraction"),
        ("head_pen_frac", "head_penetration_frac", 2e-3, "body point object proxy"),
        ("upper_pen_frac", "upper_body_penetration_frac", 2e-3, "upper body point object proxy"),
    ]
    rows = []
    for hist_field, new_field, tol, explanation in pairs:
        if hist_field not in summary or summary.get(hist_field, "") == "":
            continue
        h = as_float(summary.get(hist_field))
        u = as_float(unified.get(new_field))
        diff = abs(h - u) if math.isfinite(h) and math.isfinite(u) else math.nan
        status = "PASS" if math.isfinite(diff) and diff <= tol else "MISMATCH"
        rows.append(
            {
                "case_id": case.get("case_id", ""),
                "run_id": case.get("cem_run_id", ""),
                "historical_field": hist_field,
                "unified_field": new_field,
                "historical_value": finite_fmt(h),
                "unified_value": finite_fmt(u),
                "abs_diff": finite_fmt(diff),
                "tolerance": tol,
                "status": status,
                "explanation": explanation,
            }
        )
    return rows


def method_summary(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        groups[row["method"]].append(row)
    out = []
    for method, rs in sorted(groups.items()):
        def mean(key: str) -> str:
            vals = [as_float(r.get(key)) for r in rs]
            vals = [v for v in vals if math.isfinite(v)]
            return f"{sum(vals) / len(vals):.6f}" if vals else ""
        out.append(
            {
                "method": method,
                "num_cases": len(rs),
                "mean_pelvis_min_m": mean("pelvis_min_m"),
                "num_fall": sum(bool(r.get("fall_flag")) for r in rs),
                "mean_eef_near_3cm_frac": mean("eef_near_3cm_frac"),
                "mean_eef_near_5cm_frac": mean("eef_near_5cm_frac"),
                "mean_eef_near_8cm_frac": mean("eef_near_8cm_frac"),
                "mean_eef_near_10cm_frac": mean("eef_near_10cm_frac"),
                "mean_eef_near_12cm_frac": mean("eef_near_12cm_frac"),
                "mean_eef_near_15cm_frac": mean("eef_near_15cm_frac"),
                "mean_hand_penetration_frac": mean("hand_geom_penetration_frac"),
                "mean_hand_deep_penetration_2cm_frac": mean("hand_geom_deep_penetration_2cm_frac"),
                "mean_hand_geom_near_3cm_frac": mean("hand_geom_near_3cm_frac"),
                "mean_hand_geom_near_5cm_frac": mean("hand_geom_near_5cm_frac"),
                "mean_hand_geom_near_8cm_frac": mean("hand_geom_near_8cm_frac"),
                "mean_hand_geom_near_10cm_frac": mean("hand_geom_near_10cm_frac"),
                "mean_hand_geom_near_12cm_frac": mean("hand_geom_near_12cm_frac"),
                "mean_hand_geom_near_15cm_frac": mean("hand_geom_near_15cm_frac"),
                "mean_hand_object_physics_contact_frac": mean("hand_object_physics_contact_frac"),
                "mean_leg_penetration_frac": mean("leg_penetration_frac"),
                "mean_leg_near_2cm_frac": mean("leg_near_2cm_frac"),
                "mean_body_penetration_frac": mean("body_penetration_frac"),
                "mean_object_xy_displacement_m": mean("object_xy_displacement_m"),
                "mean_object_z_range_m": mean("object_z_range_m"),
                "mean_obj_err_mean_m": mean("obj_err_mean_m"),
            }
        )
    return out


def fmt_pct(value: Any) -> str:
    v = as_float(value)
    return f"{v * 100:.1f}%" if math.isfinite(v) else ""


def fmt_m(value: Any, digits: int = 3) -> str:
    v = as_float(value)
    return f"{v:.{digits}f}" if math.isfinite(v) else ""


def ppt_method_summary(summary: list[dict[str, Any]]) -> list[dict[str, str]]:
    """Compact method-summary rows for half-slide reporting.

    Keep only metrics shared by both OmniRetarget and Spider CEM. Spider-only
    CEM object error stays in the full sheets.
    """
    rows = []
    method_label = {
        "OmniRetarget": "OmniRetarget",
        "Spider CEM": "Spider",
    }
    for row in summary:
        rows.append(
            {
                "方法": method_label.get(str(row.get("method", "")), str(row.get("method", ""))),
                "N": str(row.get("num_cases", "")),
                "跌倒↓": str(row.get("num_fall", "")),
                "站立高度↑": fmt_m(row.get("mean_pelvis_min_m")),
                "手15cm↑": fmt_pct(row.get("mean_eef_near_15cm_frac")),
                "手12cm↑": fmt_pct(row.get("mean_eef_near_12cm_frac")),
                "手10cm↑": fmt_pct(row.get("mean_eef_near_10cm_frac")),
                "手8cm↑": fmt_pct(row.get("mean_eef_near_8cm_frac")),
                "手5cm↑": fmt_pct(row.get("mean_eef_near_5cm_frac")),
                "手深穿透↓": fmt_pct(row.get("mean_hand_deep_penetration_2cm_frac")),
                "手geom12cm↑": fmt_pct(row.get("mean_hand_geom_near_12cm_frac")),
                "手物理接触↑": fmt_pct(row.get("mean_hand_object_physics_contact_frac")),
                "腿穿透↓": fmt_pct(row.get("mean_leg_penetration_frac")),
                "身体穿透↓": fmt_pct(row.get("mean_body_penetration_frac")),
                "物体XY↑": fmt_m(row.get("mean_object_xy_displacement_m")),
            }
        )
    return rows


def build_wide_comparison(metric_rows: list[dict[str, Any]], validation_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_case: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    for row in metric_rows:
        key = f"{row['case_id']}|{row['target_variant_id']}|{row['run_id']}"
        by_case[key][row["method"]] = row
    val_status: dict[str, str] = {}
    for row in validation_rows:
        if row["status"] != "PASS":
            key = f"{row['case_id']}||{row['run_id']}"
            val_status[key] = "MISMATCH"
    out = []
    for key, methods in sorted(by_case.items()):
        omni = methods.get("OmniRetarget", {})
        spider = methods.get("Spider CEM", {})
        case_id = spider.get("case_id", omni.get("case_id", ""))
        target_variant_id = spider.get("target_variant_id", omni.get("target_variant_id", ""))
        run_id = spider.get("run_id", omni.get("run_id", ""))
        out.append(
            {
                "case_id": case_id,
                "object_key": spider.get("object_key", omni.get("object_key", "")),
                "target_variant_id": target_variant_id,
                "spider_run_id": run_id,
                "omni_qpos_path": omni.get("qpos_path", ""),
                "spider_qpos_path": spider.get("qpos_path", ""),
                "Omni pelvis_min_m": finite_fmt(omni.get("pelvis_min_m")),
                "Spider pelvis_min_m": finite_fmt(spider.get("pelvis_min_m")),
                "Omni fall": omni.get("fall_flag", ""),
                "Spider fall": spider.get("fall_flag", ""),
                "Omni eef_near_3cm": finite_fmt(omni.get("eef_near_3cm_frac")),
                "Spider eef_near_3cm": finite_fmt(spider.get("eef_near_3cm_frac")),
                "Omni eef_near_5cm": finite_fmt(omni.get("eef_near_5cm_frac")),
                "Spider eef_near_5cm": finite_fmt(spider.get("eef_near_5cm_frac")),
                "Omni eef_near_8cm": finite_fmt(omni.get("eef_near_8cm_frac")),
                "Spider eef_near_8cm": finite_fmt(spider.get("eef_near_8cm_frac")),
                "Omni eef_near_10cm": finite_fmt(omni.get("eef_near_10cm_frac")),
                "Spider eef_near_10cm": finite_fmt(spider.get("eef_near_10cm_frac")),
                "Omni hand_penetration": finite_fmt(omni.get("hand_geom_penetration_frac")),
                "Spider hand_penetration": finite_fmt(spider.get("hand_geom_penetration_frac")),
                "Omni hand_deep_penetration_2cm": finite_fmt(omni.get("hand_geom_deep_penetration_2cm_frac")),
                "Spider hand_deep_penetration_2cm": finite_fmt(spider.get("hand_geom_deep_penetration_2cm_frac")),
                "Omni leg_penetration": finite_fmt(omni.get("leg_penetration_frac")),
                "Spider leg_penetration": finite_fmt(spider.get("leg_penetration_frac")),
                "Omni leg_near_2cm": finite_fmt(omni.get("leg_near_2cm_frac")),
                "Spider leg_near_2cm": finite_fmt(spider.get("leg_near_2cm_frac")),
                "Omni body_penetration": finite_fmt(omni.get("body_penetration_frac")),
                "Spider body_penetration": finite_fmt(spider.get("body_penetration_frac")),
                "Omni object_floor_contact": finite_fmt(omni.get("object_floor_contact_frac")),
                "Spider object_floor_contact": finite_fmt(spider.get("object_floor_contact_frac")),
                "Omni object_xy_displacement_m": finite_fmt(omni.get("object_xy_displacement_m")),
                "Spider object_xy_displacement_m": finite_fmt(spider.get("object_xy_displacement_m")),
                "Omni object_z_range_m": finite_fmt(omni.get("object_z_range_m")),
                "Spider object_z_range_m": finite_fmt(spider.get("object_z_range_m")),
                "Spider obj_err_mean_m": finite_fmt(spider.get("obj_err_mean_m")),
                "Spider obj_err_max_m": finite_fmt(spider.get("obj_err_max_m")),
                "validation_status": val_status.get(f"{case_id}||{run_id}", "PASS"),
            }
        )
    return out


def write_markdown(path: Path, comparison: list[dict[str, Any]], summary: list[dict[str, Any]], validation: list[dict[str, Any]]) -> None:
    mismatches = [row for row in validation if row["status"] != "PASS"]
    lines = [
        "# 统一 replay 评测：OmniRetarget vs Spider",
        "",
        f"case 集合：`existing_cases.tsv` 中 `cem_status=pass` 且排除 `target_variant_id=adaptive` 的 {len(comparison)} 条。",
        "",
        "## 校验结论",
        "",
        f"- 历史对齐检查项：{len(validation)}",
        f"- mismatch：{len(mismatches)}",
        "- 只有同一脚本重算出的统一指标用于最终对比；历史指标只用于一致性校验。",
        "",
        "## 方法汇总",
        "",
        "| method | cases | pelvis | fall | eef 3cm | eef 5cm | eef 8cm | eef 10cm | hand pen | hand deep | leg pen | leg 2cm | body pen | obj xy | obj z | obj err |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for row in summary:
        lines.append(
            f"| {row['method']} | {row['num_cases']} | {row['mean_pelvis_min_m']} | {row['num_fall']} | "
            f"{row['mean_eef_near_3cm_frac']} | {row['mean_eef_near_5cm_frac']} | {row['mean_eef_near_8cm_frac']} | {row['mean_eef_near_10cm_frac']} | "
            f"{row['mean_hand_penetration_frac']} | {row['mean_hand_deep_penetration_2cm_frac']} | "
            f"{row['mean_leg_penetration_frac']} | {row['mean_leg_near_2cm_frac']} | {row['mean_body_penetration_frac']} | "
            f"{row['mean_object_xy_displacement_m']} | {row['mean_object_z_range_m']} | {row['mean_obj_err_mean_m']} |"
        )
    lines.extend(
        [
            "",
            "## 逐 case 对比",
            "",
            "| case | object | Omni eef 3/5/8/10 | Spider eef 3/5/8/10 | Omni pen(hand/leg/body) | Spider pen(hand/leg/body) | Spider obj err | validation |",
            "|---|---|---|---|---|---|---:|---|",
        ]
    )
    for row in comparison:
        lines.append(
            f"| `{row['case_id']}` | {row['object_key']} | "
            f"{row['Omni eef_near_3cm']}/{row['Omni eef_near_5cm']}/{row['Omni eef_near_8cm']}/{row['Omni eef_near_10cm']} | "
            f"{row['Spider eef_near_3cm']}/{row['Spider eef_near_5cm']}/{row['Spider eef_near_8cm']}/{row['Spider eef_near_10cm']} | "
            f"{row['Omni hand_penetration']}/{row['Omni leg_penetration']}/{row['Omni body_penetration']} | "
            f"{row['Spider hand_penetration']}/{row['Spider leg_penetration']}/{row['Spider body_penetration']} | "
            f"{row['Spider obj_err_mean_m']} | {row['validation_status']} |"
        )
    if mismatches:
        lines.extend(["", "## 历史对齐 mismatch", "", "| case | field | historical | unified | diff | explanation |", "|---|---|---:|---:|---:|---|"])
        for row in mismatches:
            lines.append(
                f"| `{row['case_id']}` | {row['historical_field']}->{row['unified_field']} | "
                f"{row['historical_value']} | {row['unified_value']} | {row['abs_diff']} | {row['explanation']} |"
            )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def md_cell(value: Any) -> str:
    return str(value if value is not None else "").replace("|", "\\|").replace("\n", " ")


def write_full_markdown(path: Path, title: str, rows: list[dict[str, Any]], fields: list[str]) -> None:
    lines = [
        f"# {title}",
        "",
        "说明：本表为全字段 Markdown，字段与 xlsx 对应 sheet 一致。",
        "",
        "|" + "|".join(fields) + "|",
        "|" + "|".join(["---"] * len(fields)) + "|",
    ]
    for row in rows:
        lines.append("|" + "|".join(md_cell(row.get(field, "")) for field in fields) + "|")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_xlsx(path: Path, comparison: list[dict[str, Any]], metrics: list[dict[str, Any]], summary: list[dict[str, Any]], validation: list[dict[str, Any]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    def add(name: str, rows: list[dict[str, Any]], fields: list[str]):
        ws = wb.create_sheet(name)
        ws.append(fields)
        for row in rows:
            ws.append([row.get(field, "") for field in fields])
        fill = PatternFill("solid", fgColor="1F4E78")
        font = Font(name="Arial", bold=True, color="FFFFFF")
        for c in ws[1]:
            c.fill = fill
            c.font = font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, field in enumerate(fields, start=1):
            width = min(max(len(field) + 2, 12), 36)
            ws.column_dimensions[get_column_letter(idx)].width = width
        return ws

    add("统一逐case对比", comparison, COMPARISON_FIELDS)
    add("统一method_metrics", metrics, METRIC_FIELDS)
    add("方法汇总", summary, list(summary[0].keys()) if summary else [])
    ppt_ws = add("PPT方法汇总", ppt_method_summary(summary), PPT_SUMMARY_FIELDS)
    ppt_widths = {
        "A": 12,
        "B": 4,
        "C": 7,
        "D": 9,
        "E": 8,
        "F": 8,
        "G": 8,
        "H": 10,
        "I": 8,
        "J": 9,
        "K": 8,
    }
    for col, width in ppt_widths.items():
        ppt_ws.column_dimensions[col].width = width
    for row in ppt_ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", bold=cell.row == 1, color="FFFFFF" if cell.row == 1 else "000000", size=9)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_idx in range(2, ppt_ws.max_row + 1):
        ppt_ws.row_dimensions[row_idx].height = 18
    ppt_ws.row_dimensions[1].height = 28
    add("历史对齐校验", validation, VALIDATION_FIELDS)
    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--existing-cases", type=Path, default=REPO / "workspace/core4d/data_construction_v3/existing_cases.tsv")
    parser.add_argument("--out-dir", type=Path, default=REPO / "workspace/core4d/results/E109/unified_replay_eval")
    args = parser.parse_args()

    inputs = read_case_inputs(args.existing_cases)
    metric_rows: list[dict[str, Any]] = []
    validation_rows: list[dict[str, Any]] = []
    for item in inputs:
        case = item["case"]
        summary = item["summary"]
        for method, qpos_key, scene_key in [
            ("OmniRetarget", "omni_qpos", "omni_scene"),
            ("Spider CEM", "spider_qpos", "spider_scene"),
        ]:
            metric_rows.append(
                evaluate_sequence(
                    case_id=case.get("case_id", ""),
                    method=method,
                    object_key=case.get("object_key", ""),
                    target_variant_id=case.get("target_variant_id", ""),
                    run_id=case.get("cem_run_id", ""),
                    qpos_path=item[qpos_key],
                    scene_xml=item[scene_key],
                )
            )
        validation_rows.extend(compare_history(case, summary, metric_rows[-1]))

    comparison = build_wide_comparison(metric_rows, validation_rows)
    summary = method_summary(metric_rows)

    args.out_dir.mkdir(parents=True, exist_ok=True)
    write_tsv(args.out_dir / "unified_method_metrics.tsv", metric_rows, METRIC_FIELDS)
    write_tsv(args.out_dir / "unified_case_comparison.tsv", comparison, COMPARISON_FIELDS)
    write_tsv(args.out_dir / "history_validation.tsv", validation_rows, VALIDATION_FIELDS)
    write_tsv(args.out_dir / "unified_method_summary.tsv", summary, list(summary[0].keys()) if summary else [])
    write_markdown(args.out_dir / "unified_omni_vs_spider_comparison.md", comparison, summary, validation_rows)
    write_full_markdown(args.out_dir / "unified_case_comparison_full.md", "统一 replay 完整逐 case 对比表", comparison, COMPARISON_FIELDS)
    write_full_markdown(args.out_dir / "unified_method_metrics_full.md", "统一 replay 完整 method metrics 表", metric_rows, METRIC_FIELDS)
    write_full_markdown(args.out_dir / "history_validation_full.md", "历史指标对齐校验全表", validation_rows, VALIDATION_FIELDS)
    write_xlsx(args.out_dir / "unified_omni_vs_spider_comparison.xlsx", comparison, metric_rows, summary, validation_rows)
    write_json(
        args.out_dir / "run_summary.json",
        {
            "num_cases": len(inputs),
            "num_metric_rows": len(metric_rows),
            "num_validation_rows": len(validation_rows),
            "num_validation_mismatch": sum(row["status"] != "PASS" for row in validation_rows),
            "near_thresholds_m": NEAR_THRESHOLDS_M,
            "deep_penetration_m": DEEP_PENETRATION_M,
            "fall_pelvis_z_m": FALL_PELVIS_Z_M,
        },
    )
    print(f"[unified_replay_eval] wrote {len(inputs)} cases to {args.out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
