#!/usr/bin/env python3
"""Build an OmniRetarget vs Spider comparison table for existing CEM-pass cases."""

from __future__ import annotations

import argparse
import csv
import json
import math
from collections import defaultdict
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import REPO, as_float, read_csv, resolve_path, write_json, write_tsv


FPS = 30.0
FALL_PELVIS_Z_M = 0.45

CASE_FIELDS = [
    "比较ID",
    "case_id",
    "物体",
    "retarget版本",
    "target路线",
    "Spider CEM run",
    "CEM状态",
    "RL状态",
    "OmniRetarget轨迹",
    "Spider结果npz",
    "指标覆盖",
    "主要结论",
    "Omni帧数",
    "Spider帧数",
    "Omni时长s",
    "Spider时长s",
    "Omni pelvis最小高度m",
    "Spider pelvis最小高度m",
    "pelvis高度差Spider-Omni m",
    "Omni跌倒",
    "Spider跌倒",
    "Omni root位移m",
    "Spider root位移m",
    "Spider相对Omni root均值误差m",
    "Spider相对Omni root最大误差m",
    "Omni物体XY位移m",
    "Omni物体高度变化m",
    "Spider物体均值误差m",
    "Spider物体最大误差m",
    "Spider手-物体8cm近接触比例",
    "Spider腿-物体干涉比例",
    "Spider腿-物体2cm邻近比例",
    "Spider手-物体穿透/物理接触比例",
    "Omni contact mask激活比例",
    "Omni contact_pos存在",
    "Spider replay gate",
    "Spider lower-body strict",
    "备注",
]

METHOD_FIELDS = [
    "方法",
    "case数",
    "有轨迹/指标case数",
    "平均pelvis最小高度m",
    "跌倒case数",
    "平均root位移m",
    "平均物体XY位移m",
    "平均物体高度变化m",
    "平均物体误差m",
    "平均8cm接触比例",
    "平均腿部干涉比例",
    "说明",
]

DEFINITION_FIELDS = ["指标", "方向", "单位", "来源", "说明"]


def repo_path_text(path: str | Path | None) -> str:
    if path is None:
        return ""
    text = str(path)
    if not text:
        return ""
    p = Path(text)
    try:
        if p.is_absolute():
            return str(p.relative_to(REPO))
    except ValueError:
        return text
    return text


def finite_mean(values: list[float]) -> str:
    vals = [v for v in values if math.isfinite(v)]
    return f"{sum(vals) / len(vals):.6f}" if vals else ""


def fmt(value: Any, digits: int = 6) -> str:
    v = as_float(value)
    return f"{v:.{digits}f}" if math.isfinite(v) else ""


def read_existing_cases(path: Path) -> list[dict[str, str]]:
    return [
        row
        for row in read_csv(path, delimiter="\t")
        if row.get("cem_status") == "pass" and row.get("target_variant_id") != "adaptive"
    ]


def load_summary_row(case_row: dict[str, str]) -> tuple[dict[str, str], Path | None]:
    metrics_ref = resolve_path(case_row.get("cem_metrics_ref"))
    if not metrics_ref or not metrics_ref.exists():
        return {}, metrics_ref
    if metrics_ref.suffix.lower() == ".md":
        csv_path = metrics_ref.with_suffix(".csv")
        if csv_path.exists():
            metrics_ref = csv_path
        else:
            sibling = metrics_ref.parent / "E108_cem_eval_summary.csv"
            if sibling.exists():
                metrics_ref = sibling
    if metrics_ref.suffix.lower() != ".csv":
        return {}, metrics_ref
    rows = read_csv(metrics_ref)
    run_id = case_row.get("cem_run_id", "")
    case_id = case_row.get("case_id", "")
    for row in rows:
        if run_id and row.get("variant") == run_id:
            return row, metrics_ref
    for row in rows:
        if case_id in {row.get("source_task", ""), row.get("case", ""), row.get("case_id", "")}:
            return row, metrics_ref
    return (rows[0] if rows else {}), metrics_ref


def candidate_scene_act_from_summary(summary: dict[str, str]) -> Path | None:
    for key in ("scene_xml", "scene_used", "legobj_scene_used"):
        path = resolve_path(summary.get(key))
        if path and path.exists():
            return path
    return None


def trajectory_from_scene(scene_act: Path | None) -> Path | None:
    if not scene_act:
        return None
    task_dir = scene_act.parent
    candidate = task_dir / "0" / "trajectory_kinematic.npz"
    if candidate.exists():
        return candidate
    info_path = task_dir / "task_info.json"
    if info_path.exists():
        try:
            info = json.loads(info_path.read_text(encoding="utf-8"))
        except Exception:
            info = {}
        source_qpos = info.get("source_qpos", "")
        if source_qpos:
            p = Path(source_qpos)
            candidates = []
            if p.is_absolute():
                candidates.append(p)
            candidates.append((task_dir / p).resolve())
            candidates.append((REPO / p).resolve())
            for item in candidates:
                if item.exists():
                    return item
    return None


def resolve_omnirt_trajectory(case_row: dict[str, str], summary: dict[str, str]) -> Path | None:
    scene_act = candidate_scene_act_from_summary(summary)
    traj = trajectory_from_scene(scene_act)
    if traj:
        return traj

    case_id = case_row.get("case_id", "")
    run_id = case_row.get("cem_run_id", "")
    task_root = REPO / "example_datasets/processed/core4d/unitree_g1/humanoid_object"
    derived = summary.get("derived_task", "")
    for name in [derived, case_id]:
        if name:
            candidate = task_root / name / "0" / "trajectory_kinematic.npz"
            if candidate.exists():
                return candidate

    if run_id.startswith("E108"):
        manifest = REPO / "workspace/core4d/results/E108/s5_handoff/handoff_manifest.tsv"
        for row in read_csv(manifest, delimiter="\t"):
            if row.get("case_id") == case_id and row.get("retarget_variant_id") == case_row.get("retarget_variant_id"):
                for key in ("trajectory", "spider_trajectory"):
                    p = resolve_path(row.get(key))
                    if p and p.exists():
                        return p
    return None


def load_npz(path: Path | None) -> dict[str, np.ndarray]:
    if not path or not path.exists():
        return {}
    try:
        z = np.load(path, allow_pickle=True)
        return {key: z[key] for key in z.files}
    except Exception:
        return {}


def qpos_main(npz: dict[str, np.ndarray]) -> np.ndarray | None:
    qpos = npz.get("qpos")
    if qpos is None:
        return None
    arr = np.asarray(qpos)
    if arr.ndim == 3:
        return arr[:, 0, :]
    if arr.ndim == 2:
        return arr
    return None


def trajectory_metrics(path: Path | None, *, method: str) -> dict[str, Any]:
    npz = load_npz(path)
    qpos = qpos_main(npz)
    out: dict[str, Any] = {
        "trajectory_exists": bool(path and path.exists()),
        "frames": "",
        "duration_s": "",
        "pelvis_min_z_m": "",
        "fall_flag": "",
        "root_xy_displacement_m": "",
        "root_xy_path_m": "",
        "object_xy_displacement_m": "",
        "object_z_range_m": "",
        "contact_mask_active_frac": "",
        "contact_pos_present": False,
    }
    if qpos is not None and qpos.shape[0] > 0 and qpos.shape[1] >= 3:
        root = qpos[:, :3].astype(float)
        out["frames"] = int(qpos.shape[0])
        out["duration_s"] = qpos.shape[0] / FPS
        out["pelvis_min_z_m"] = float(np.nanmin(root[:, 2]))
        out["fall_flag"] = bool(out["pelvis_min_z_m"] < FALL_PELVIS_Z_M)
        out["root_xy_displacement_m"] = float(np.linalg.norm(root[-1, :2] - root[0, :2]))
        if qpos.shape[0] > 1:
            out["root_xy_path_m"] = float(np.linalg.norm(np.diff(root[:, :2], axis=0), axis=1).sum())
        else:
            out["root_xy_path_m"] = 0.0
        if method == "omniretarget" and qpos.shape[1] >= 43:
            obj = qpos[:, -7:-4].astype(float)
            out["object_xy_displacement_m"] = float(np.linalg.norm(obj[-1, :2] - obj[0, :2]))
            out["object_z_range_m"] = float(np.nanmax(obj[:, 2]) - np.nanmin(obj[:, 2]))
    contact = npz.get("contact")
    if contact is not None:
        carr = np.asarray(contact)
        if carr.size:
            out["contact_mask_active_frac"] = float((carr > 0).any(axis=-1).mean()) if carr.ndim >= 2 else float((carr > 0).mean())
    out["contact_pos_present"] = "contact_pos" in npz
    return out


def paired_root_error(spider_path: Path | None, omni_path: Path | None) -> dict[str, Any]:
    s = qpos_main(load_npz(spider_path))
    o = qpos_main(load_npz(omni_path))
    if s is None or o is None or s.size == 0 or o.size == 0:
        return {"mean": "", "max": ""}
    n = min(s.shape[0], o.shape[0])
    if n == 0:
        return {"mean": "", "max": ""}
    diff = np.linalg.norm(s[:n, :3].astype(float) - o[:n, :3].astype(float), axis=1)
    return {"mean": float(diff.mean()), "max": float(diff.max())}


def get_float(row: dict[str, str], keys: list[str]) -> float:
    for key in keys:
        val = as_float(row.get(key))
        if math.isfinite(val):
            return val
    return math.nan


def build_rows(existing_cases: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    case_rows = []
    method_long = []
    warnings = []
    for idx, case in enumerate(read_existing_cases(existing_cases), start=1):
        summary, metrics_ref = load_summary_row(case)
        spider_npz = resolve_path(case.get("cem_result_npz"))
        if not spider_npz or not spider_npz.exists():
            spider_npz = resolve_path(summary.get("npz_path") or summary.get("root_npz_path"))
        omni_traj = resolve_omnirt_trajectory(case, summary)

        omni_m = trajectory_metrics(omni_traj, method="omniretarget")
        spider_m = trajectory_metrics(spider_npz, method="spider")
        root_err = paired_root_error(spider_npz, omni_traj)

        spider_contact = get_float(summary, ["contact_frac_either", "case_window_sim_contact_frames_pct", "post2_sim_contact_frames_pct"])
        if math.isfinite(spider_contact) and spider_contact > 1.0:
            spider_contact /= 100.0
        spider_leg = get_float(summary, ["leg_box_interference_frac", "full_sim_leg_box_interference_frames_pct", "case_window_sim_leg_box_interference_frames_pct"])
        if math.isfinite(spider_leg) and spider_leg > 1.0:
            spider_leg /= 100.0
        spider_leg_near = get_float(summary, ["leg_box_near_2cm_frac", "full_sim_leg_box_near_2cm_frames_pct", "case_window_sim_leg_box_near_2cm_frames_pct"])
        if math.isfinite(spider_leg_near) and spider_leg_near > 1.0:
            spider_leg_near /= 100.0
        spider_pen = get_float(summary, ["hand_object_penetration_frac", "hand_object_contact_physics_frac", "full_sim_hand_object_contact_frames_pct"])
        if math.isfinite(spider_pen) and spider_pen > 1.0:
            spider_pen /= 100.0

        spider_obj_mean = get_float(summary, ["obj_err_mean_m", "post2_obj_err_mean_m", "case_window_obj_err_mean_m"])
        spider_obj_max = get_float(summary, ["obj_err_max_m", "post2_obj_err_max_m", "case_window_obj_err_max_m"])
        spider_pelvis_summary = get_float(summary, ["pelvis_min_m", "post2_pelvis_z_min_m", "full_pelvis_z_min_m", "case_window_pelvis_z_min_m"])
        if math.isfinite(spider_pelvis_summary):
            spider_m["pelvis_min_z_m"] = spider_pelvis_summary
            spider_m["fall_flag"] = bool(spider_pelvis_summary < FALL_PELVIS_Z_M)

        comparison_id = f"{idx:02d}_{case.get('case_id','')}_{case.get('target_variant_id','')}_{case.get('cem_run_id','')}"
        coverage = []
        if omni_m["trajectory_exists"]:
            coverage.append("Omni轨迹")
        else:
            warnings.append({"比较ID": comparison_id, "问题": "缺少OmniRetarget trajectory_kinematic/source_qpos"})
        if spider_m["trajectory_exists"]:
            coverage.append("Spider轨迹")
        else:
            warnings.append({"比较ID": comparison_id, "问题": "缺少Spider CEM npz"})
        if summary:
            coverage.append("Spider summary")
        else:
            warnings.append({"比较ID": comparison_id, "问题": "缺少Spider CEM summary row"})

        spider_good = (
            math.isfinite(spider_obj_mean)
            and spider_obj_mean <= 0.02
            and math.isfinite(as_float(spider_m["pelvis_min_z_m"]))
            and as_float(spider_m["pelvis_min_z_m"]) >= FALL_PELVIS_Z_M
        )
        conclusion = "Spider CEM稳定且物体跟踪误差低" if spider_good else "需要复查"
        if math.isfinite(spider_leg) and spider_leg > 0.05:
            conclusion = "Spider上半身可用但腿部干涉偏高"

        row = {
            "比较ID": comparison_id,
            "case_id": case.get("case_id", ""),
            "物体": case.get("object_key", ""),
            "retarget版本": case.get("retarget_variant_id", ""),
            "target路线": case.get("target_variant_id", ""),
            "Spider CEM run": case.get("cem_run_id", ""),
            "CEM状态": case.get("cem_status", ""),
            "RL状态": case.get("rl_status", ""),
            "OmniRetarget轨迹": repo_path_text(omni_traj),
            "Spider结果npz": repo_path_text(spider_npz),
            "指标覆盖": "+".join(coverage),
            "主要结论": conclusion,
            "Omni帧数": omni_m["frames"],
            "Spider帧数": spider_m["frames"],
            "Omni时长s": fmt(omni_m["duration_s"], 3),
            "Spider时长s": fmt(spider_m["duration_s"], 3),
            "Omni pelvis最小高度m": fmt(omni_m["pelvis_min_z_m"], 4),
            "Spider pelvis最小高度m": fmt(spider_m["pelvis_min_z_m"], 4),
            "pelvis高度差Spider-Omni m": fmt(as_float(spider_m["pelvis_min_z_m"]) - as_float(omni_m["pelvis_min_z_m"]), 4),
            "Omni跌倒": omni_m["fall_flag"],
            "Spider跌倒": spider_m["fall_flag"],
            "Omni root位移m": fmt(omni_m["root_xy_displacement_m"], 4),
            "Spider root位移m": fmt(spider_m["root_xy_displacement_m"], 4),
            "Spider相对Omni root均值误差m": fmt(root_err["mean"], 4),
            "Spider相对Omni root最大误差m": fmt(root_err["max"], 4),
            "Omni物体XY位移m": fmt(omni_m["object_xy_displacement_m"], 4),
            "Omni物体高度变化m": fmt(omni_m["object_z_range_m"], 4),
            "Spider物体均值误差m": fmt(spider_obj_mean, 4),
            "Spider物体最大误差m": fmt(spider_obj_max, 4),
            "Spider手-物体8cm近接触比例": fmt(spider_contact, 4),
            "Spider腿-物体干涉比例": fmt(spider_leg, 4),
            "Spider腿-物体2cm邻近比例": fmt(spider_leg_near, 4),
            "Spider手-物体穿透/物理接触比例": fmt(spider_pen, 4),
            "Omni contact mask激活比例": fmt(omni_m["contact_mask_active_frac"], 4),
            "Omni contact_pos存在": omni_m["contact_pos_present"],
            "Spider replay gate": summary.get("replay_gate_pass", summary.get("stage_pass", "")),
            "Spider lower-body strict": summary.get("lowerbody_strict_pass", ""),
            "备注": f"metrics_ref={repo_path_text(metrics_ref)}; Omni接触SDF未全量重算，contact mask不是实际物理接触。",
        }
        case_rows.append(row)

        method_long.append(
            {
                "方法": "OmniRetarget",
                "case_id": row["case_id"],
                "pelvis_min": as_float(row["Omni pelvis最小高度m"]),
                "fall": row["Omni跌倒"] is True,
                "root_disp": as_float(row["Omni root位移m"]),
                "obj_xy": as_float(row["Omni物体XY位移m"]),
                "obj_z": as_float(row["Omni物体高度变化m"]),
                "obj_err": math.nan,
                "contact": math.nan,
                "leg": math.nan,
                "ready": bool(omni_m["trajectory_exists"]),
            }
        )
        method_long.append(
            {
                "方法": "Spider CEM",
                "case_id": row["case_id"],
                "pelvis_min": as_float(row["Spider pelvis最小高度m"]),
                "fall": row["Spider跌倒"] is True,
                "root_disp": as_float(row["Spider root位移m"]),
                "obj_xy": math.nan,
                "obj_z": math.nan,
                "obj_err": spider_obj_mean,
                "contact": spider_contact,
                "leg": spider_leg,
                "ready": bool(spider_m["trajectory_exists"] and summary),
            }
        )
    return case_rows, method_long, warnings


def summarize_methods(method_long: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in method_long:
        groups[row["方法"]].append(row)
    out = []
    for method, rows in groups.items():
        out.append(
            {
                "方法": method,
                "case数": len(rows),
                "有轨迹/指标case数": sum(bool(r["ready"]) for r in rows),
                "平均pelvis最小高度m": finite_mean([r["pelvis_min"] for r in rows]),
                "跌倒case数": sum(bool(r["fall"]) for r in rows),
                "平均root位移m": finite_mean([r["root_disp"] for r in rows]),
                "平均物体XY位移m": finite_mean([r["obj_xy"] for r in rows]),
                "平均物体高度变化m": finite_mean([r["obj_z"] for r in rows]),
                "平均物体误差m": finite_mean([r["obj_err"] for r in rows]),
                "平均8cm接触比例": finite_mean([r["contact"] for r in rows]),
                "平均腿部干涉比例": finite_mean([r["leg"] for r in rows]),
                "说明": "OmniRetarget为运动学轨迹指标；Spider CEM为物理 rollout + CEM summary 指标。",
            }
        )
    return out


def metric_definitions() -> list[dict[str, str]]:
    return [
        {"指标": "pelvis最小高度m", "方向": "越高越好", "单位": "m", "来源": "npz qpos root z / summary", "说明": "低于0.45m记为跌倒。"},
        {"指标": "root位移m", "方向": "任务相关", "单位": "m", "来源": "npz qpos root xy", "说明": "机器人根节点首末帧XY位移。"},
        {"指标": "Spider相对Omni root误差", "方向": "越低越接近Omni参考", "单位": "m", "来源": "Spider CEM qpos vs Omni trajectory", "说明": "诊断指标，使用Omni作为reference，不作为公平胜负核心。"},
        {"指标": "Omni物体XY位移m", "方向": "任务相关", "单位": "m", "来源": "Omni qpos最后7维object pose", "说明": "OmniRetarget运动学参考里的物体首末帧XY位移。"},
        {"指标": "Omni物体高度变化m", "方向": "任务相关", "单位": "m", "来源": "Omni qpos最后7维object pose", "说明": "物体z最大值减最小值。"},
        {"指标": "Spider物体均值/最大误差m", "方向": "越低越好", "单位": "m", "来源": "CEM eval summary", "说明": "相对Spider CEM参考物体轨迹的误差，是method-reference指标。"},
        {"指标": "Spider手-物体8cm近接触比例", "方向": "通常越高越好", "单位": "比例", "来源": "CEM eval summary", "说明": "8cm阈值近接触；高接触需同时检查穿透。"},
        {"指标": "Spider腿-物体干涉比例", "方向": "越低越好", "单位": "比例", "来源": "CEM eval summary", "说明": "腿部与物体穿透/干涉proxy。"},
        {"指标": "Omni contact mask激活比例", "方向": "仅诊断", "单位": "比例", "来源": "Omni trajectory contact", "说明": "这是目标mask，不等价于实际物理接触。"},
    ]


def write_markdown(path: Path, case_rows: list[dict[str, Any]], method_rows: list[dict[str, Any]], warnings: list[dict[str, Any]]) -> None:
    lines = [
        "# OmniRetarget vs Spider 指标对比表",
        "",
        "case 集合：`workspace/core4d/data_construction_v3/existing_cases.tsv` 中 `cem_status=pass` 的 12 条 Spider case。",
        "",
        "## 读表结论",
        "",
        "- Spider CEM 的 12 条成功 case 都纳入了对比表。",
        "- OmniRetarget 侧对齐的是进入 Spider CEM 的 `trajectory_kinematic.npz` 或其 `source_qpos`。",
        "- 表中 Spider 的接触、物体误差、腿部干涉来自 CEM eval summary；OmniRetarget 目前只做轨迹级指标，没有伪造同口径 SDF 接触。",
        "- `Spider相对Omni root误差` 使用 Omni 作为 reference，只是诊断项，不作为公平胜负核心。",
        "",
        "## 方法汇总",
        "",
        "| 方法 | case数 | ready | pelvis均值 | 跌倒case | root位移 | 物体XY位移 | 物体高度变化 | 物体误差 | 8cm接触 | 腿部干涉 |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for row in method_rows:
        lines.append(
            f"| {row['方法']} | {row['case数']} | {row['有轨迹/指标case数']} | {row['平均pelvis最小高度m']} | "
            f"{row['跌倒case数']} | {row['平均root位移m']} | {row['平均物体XY位移m']} | "
            f"{row['平均物体高度变化m']} | {row['平均物体误差m']} | {row['平均8cm接触比例']} | {row['平均腿部干涉比例']} |"
        )
    lines.extend(
        [
            "",
            "## 逐case对比",
            "",
            "| case | 物体 | target | Omni T | Spider T | Omni pelvis | Spider pelvis | Spider obj mean | Spider contact 8cm | Spider leg intf | 结论 |",
            "|---|---|---|---:|---:|---:|---:|---:|---:|---:|---|",
        ]
    )
    for row in case_rows:
        lines.append(
            f"| `{row['case_id']}` | {row['物体']} | {row['target路线']} | {row['Omni帧数']} | {row['Spider帧数']} | "
            f"{row['Omni pelvis最小高度m']} | {row['Spider pelvis最小高度m']} | {row['Spider物体均值误差m']} | "
            f"{row['Spider手-物体8cm近接触比例']} | {row['Spider腿-物体干涉比例']} | {row['主要结论']} |"
        )
    if warnings:
        lines.extend(["", "## 覆盖问题", "", "| 比较ID | 问题 |", "|---|---|"])
        for row in warnings:
            lines.append(f"| `{row['比较ID']}` | {row['问题']} |")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def md_cell(value: Any) -> str:
    text = str(value if value is not None else "")
    return text.replace("|", "\\|").replace("\n", " ")


def write_full_case_markdown(path: Path, case_rows: list[dict[str, Any]]) -> None:
    lines = [
        "# OmniRetarget vs Spider 完整逐case指标表",
        "",
        "说明：本表与 xlsx 的 `逐case对比` sheet 字段一致。空值表示该 case 的现有评测证据中没有同口径指标，不做填充或猜测。",
        "",
        "|" + "|".join(CASE_FIELDS) + "|",
        "|" + "|".join(["---"] * len(CASE_FIELDS)) + "|",
    ]
    for row in case_rows:
        lines.append("|" + "|".join(md_cell(row.get(field, "")) for field in CASE_FIELDS) + "|")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_xlsx(path: Path, case_rows: list[dict[str, Any]], method_rows: list[dict[str, Any]], definitions: list[dict[str, str]], warnings: list[dict[str, Any]]) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    def add_sheet(name: str, rows: list[dict[str, Any]], fields: list[str]) -> None:
        ws = wb.create_sheet(name)
        ws.append(fields)
        for row in rows:
            ws.append([row.get(field, "") for field in fields])
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(name="Arial", bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col_idx, field in enumerate(fields, start=1):
            width = min(max(len(str(field)) + 2, 12), 42)
            for cell in ws[get_column_letter(col_idx)][1:]:
                width = min(max(width, min(len(str(cell.value or "")) + 2, 42)), 42)
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    add_sheet("逐case对比", case_rows, CASE_FIELDS)
    add_sheet("方法汇总", method_rows, METHOD_FIELDS)
    add_sheet("指标定义", definitions, DEFINITION_FIELDS)
    add_sheet("覆盖问题", warnings, ["比较ID", "问题"])
    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--existing-cases", type=Path, default=REPO / "workspace/core4d/data_construction_v3/existing_cases.tsv")
    parser.add_argument("--out-dir", type=Path, default=REPO / "workspace/core4d/results/E109/omni_vs_spider_existing_cases")
    args = parser.parse_args()

    args.out_dir.mkdir(parents=True, exist_ok=True)
    case_rows, method_long, warnings = build_rows(args.existing_cases)
    method_rows = summarize_methods(method_long)
    definitions = metric_definitions()

    write_tsv(args.out_dir / "omni_vs_spider_case_metrics.tsv", case_rows, CASE_FIELDS)
    write_tsv(args.out_dir / "omni_vs_spider_method_summary.tsv", method_rows, METHOD_FIELDS)
    write_tsv(args.out_dir / "metric_definitions.tsv", definitions, DEFINITION_FIELDS)
    write_tsv(args.out_dir / "coverage_warnings.tsv", warnings, ["比较ID", "问题"])
    write_markdown(args.out_dir / "omni_vs_spider_comparison.md", case_rows, method_rows, warnings)
    write_full_case_markdown(args.out_dir / "omni_vs_spider_case_metrics_full.md", case_rows)
    write_xlsx(args.out_dir / "omni_vs_spider_comparison.xlsx", case_rows, method_rows, definitions, warnings)
    write_json(
        args.out_dir / "run_summary.json",
        {
            "existing_cases": repo_path_text(args.existing_cases),
            "num_cases": len(case_rows),
            "num_warnings": len(warnings),
            "fps_assumption": FPS,
            "fall_pelvis_z_m": FALL_PELVIS_Z_M,
        },
    )
    print(f"[build_existing_cases_comparison] wrote {len(case_rows)} cases to {args.out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
