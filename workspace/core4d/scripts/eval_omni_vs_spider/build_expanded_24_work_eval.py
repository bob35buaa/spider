#!/usr/bin/env python3
"""Build a separate 24-case Spider work table.

This keeps the canonical 11-case strict table unchanged, then adds 13 deduped
ref_fk upper-WORK/non-strict cases for reporting and failure-mode context.
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import REPO, read_csv, resolve_path, write_json, write_tsv
from build_existing_cases_comparison import trajectory_from_scene
from unified_replay_eval import (
    COMPARISON_FIELDS,
    METRIC_FIELDS,
    PPT_SUMMARY_FIELDS,
    VALIDATION_FIELDS,
    compare_history,
    evaluate_sequence,
    finite_fmt,
    method_summary,
    ppt_method_summary,
    read_case_inputs,
    write_full_markdown,
)


EXPANDED_FIELDS = [
    "分层",
    "纳入建议",
    "是否strict主表",
    "case_id",
    "object_key",
    "target_variant_id",
    "spider_run_id",
    "omni_qpos_path",
    "spider_qpos_path",
    "Omni pelvis_min_m",
    "Spider pelvis_min_m",
    "Omni fall",
    "Spider fall",
    "Omni eef_near_5cm",
    "Spider eef_near_5cm",
    "Omni eef_near_8cm",
    "Spider eef_near_8cm",
    "Omni eef_near_10cm",
    "Spider eef_near_10cm",
    "Omni eef_near_12cm",
    "Spider eef_near_12cm",
    "Omni eef_near_15cm",
    "Spider eef_near_15cm",
    "Omni hand_deep_penetration_2cm",
    "Spider hand_deep_penetration_2cm",
    "Omni hand_geom_near_5cm",
    "Spider hand_geom_near_5cm",
    "Omni hand_geom_near_8cm",
    "Spider hand_geom_near_8cm",
    "Omni hand_geom_near_10cm",
    "Spider hand_geom_near_10cm",
    "Omni hand_geom_near_12cm",
    "Spider hand_geom_near_12cm",
    "Omni hand_geom_near_15cm",
    "Spider hand_geom_near_15cm",
    "Omni hand_geom_penetration",
    "Spider hand_geom_penetration",
    "Omni hand_object_physics_contact",
    "Spider hand_object_physics_contact",
    "Omni leg_penetration",
    "Spider leg_penetration",
    "Omni body_penetration",
    "Spider body_penetration",
    "Omni object_xy_displacement_m",
    "Spider object_xy_displacement_m",
    "Spider obj_err_mean_m",
    "validation_status",
    "说明",
]

TIER_FIELDS = ["分层", "纳入建议", "case数", "说明"]
COLOR_RULE_FIELDS = ["Spider指标", "Omni指标", "方向", "明显阈值", "绿色含义", "红色含义"]

BETTER_FILL = PatternFill("solid", fgColor="C6EFCE")
WORSE_FILL = PatternFill("solid", fgColor="FFC7CE")
NEUTRAL_FILL = PatternFill("solid", fgColor="D9EAF7")

COMPARISON_COLOR_RULES = [
    ("Spider pelvis_min_m", "Omni pelvis_min_m", "higher", 0.03),
    ("Spider fall", "Omni fall", "lower_bool", 0.0),
    ("Spider eef_near_5cm", "Omni eef_near_5cm", "higher", 0.05),
    ("Spider eef_near_8cm", "Omni eef_near_8cm", "higher", 0.05),
    ("Spider eef_near_10cm", "Omni eef_near_10cm", "higher", 0.05),
    ("Spider eef_near_12cm", "Omni eef_near_12cm", "higher", 0.05),
    ("Spider eef_near_15cm", "Omni eef_near_15cm", "higher", 0.05),
    ("Spider hand_deep_penetration_2cm", "Omni hand_deep_penetration_2cm", "lower", 0.02),
    ("Spider hand_geom_near_5cm", "Omni hand_geom_near_5cm", "higher", 0.05),
    ("Spider hand_geom_near_8cm", "Omni hand_geom_near_8cm", "higher", 0.05),
    ("Spider hand_geom_near_10cm", "Omni hand_geom_near_10cm", "higher", 0.05),
    ("Spider hand_geom_near_12cm", "Omni hand_geom_near_12cm", "higher", 0.05),
    ("Spider hand_geom_near_15cm", "Omni hand_geom_near_15cm", "higher", 0.05),
    ("Spider hand_geom_penetration", "Omni hand_geom_penetration", "lower", 0.02),
    ("Spider hand_object_physics_contact", "Omni hand_object_physics_contact", "higher", 0.05),
    ("Spider leg_penetration", "Omni leg_penetration", "lower", 0.02),
    ("Spider body_penetration", "Omni body_penetration", "lower", 0.02),
    ("Spider object_xy_displacement_m", "Omni object_xy_displacement_m", "higher", 0.05),
]

COLOR_RULE_ROWS = [
    {
        "Spider指标": spider,
        "Omni指标": omni,
        "方向": {"higher": "越大越好", "lower": "越小越好", "lower_bool": "False 优于 True"}[direction],
        "明显阈值": str(threshold),
        "绿色含义": "Spider 明显好于 OmniRetarget",
        "红色含义": "Spider 明显差于 OmniRetarget",
    }
    for spider, omni, direction, threshold in COMPARISON_COLOR_RULES
]


def repo_text(path: Path | None) -> str:
    if not path:
        return ""
    try:
        return str(path.relative_to(REPO))
    except ValueError:
        return str(path)


def to_float(value: Any) -> float | None:
    try:
        if value in ("", None):
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def to_bool(value: Any) -> bool | None:
    text = str(value).strip().lower()
    if text in {"true", "1", "yes"}:
        return True
    if text in {"false", "0", "no"}:
        return False
    return None


def comparison_status(spider: Any, omni: Any, direction: str, threshold: float) -> str:
    if direction == "lower_bool":
        s = to_bool(spider)
        o = to_bool(omni)
        if s is None or o is None or s == o:
            return "neutral"
        return "better" if (not s and o) else "worse"

    s_val = to_float(spider)
    o_val = to_float(omni)
    if s_val is None or o_val is None:
        return "neutral"
    diff = s_val - o_val
    if direction == "higher":
        if diff >= threshold:
            return "better"
        if diff <= -threshold:
            return "worse"
    elif direction == "lower":
        if diff <= -threshold:
            return "better"
        if diff >= threshold:
            return "worse"
    return "neutral"


def load_csv_row(path: Path, run_id: str) -> dict[str, str]:
    for row in read_csv(path):
        if row.get("variant") == run_id:
            return row
    raise KeyError(f"{run_id} not found in {path}")


def current_main_items(existing_cases: Path) -> tuple[list[dict[str, Any]], dict[str, dict[str, str]]]:
    items = read_case_inputs(existing_cases)
    meta = {}
    for item in items:
        case = item["case"]
        key = row_key(case.get("case_id", ""), case.get("target_variant_id", ""), case.get("cem_run_id", ""))
        meta[key] = {
            "分层": "A_strict_main",
            "纳入建议": "主表纳入",
            "是否strict主表": "是",
            "说明": "当前 11-case ref_fk strict 主表 case",
        }
    return items, meta


def row_key(case_id: str, target_variant_id: str, run_id: str) -> str:
    return f"{case_id}|{target_variant_id}|{run_id}"


def extra_upper_work_items(audit_tsv: Path) -> tuple[list[dict[str, Any]], dict[str, dict[str, str]]]:
    """Return the 13 deduped E106/E107 ref_fk upper-WORK rows."""
    rows = [
        row
        for row in read_csv(audit_tsv, delimiter="\t")
        if row.get("tier") == "B_candidate_upper_work_not_strict"
        and row.get("target_variant_id") == "ref_fk"
        and row.get("source") in {"E106", "E107"}
    ]
    items: list[dict[str, Any]] = []
    meta: dict[str, dict[str, str]] = {}
    for row in rows:
        run_id = row["cem_run_id"]
        source = row["source"]
        summary_path = REPO / f"workspace/core4d/results/{source}/cem/full/full_eval_summary.csv"
        summary = load_csv_row(summary_path, run_id)
        spider_qpos = resolve_path(summary.get("npz_path") or row.get("npz_path"))
        spider_scene = resolve_path(summary.get("scene_xml") or row.get("scene_xml"))
        if not spider_qpos or not spider_scene:
            raise FileNotFoundError(f"missing Spider qpos/scene for {run_id}")
        omni_qpos = trajectory_from_scene(spider_scene)
        if not omni_qpos:
            raise FileNotFoundError(f"missing OmniRetarget trajectory for {run_id} scene={spider_scene}")
        omni_scene = spider_scene.parent / "scene.xml"
        if not omni_scene.exists():
            raise FileNotFoundError(f"missing OmniRetarget scene.xml for {run_id}: {omni_scene}")
        case = {
            "case_id": row["source_task"],
            "object_key": row["object_key"],
            "target_variant_id": "ref_fk",
            "cem_run_id": run_id,
        }
        items.append(
            {
                "case": case,
                "summary": summary,
                "spider_qpos": spider_qpos,
                "spider_scene": spider_scene,
                "omni_qpos": omni_qpos,
                "omni_scene": omni_scene,
            }
        )
        meta[row_key(case["case_id"], "ref_fk", run_id)] = {
            "分层": "B_upper_WORK_non_strict",
            "纳入建议": "补充表纳入；不要当 RL-ready positive",
            "是否strict主表": "否",
            "说明": "Spider CEM upper/object WORK，但 lower-body strict 未过",
        }
    return items, meta


def evaluate_items(items: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    metric_rows: list[dict[str, Any]] = []
    validation_rows: list[dict[str, Any]] = []
    for item in items:
        case = item["case"]
        summary = item["summary"]
        for method, qpos_key, scene_key in [
            ("OmniRetarget", "omni_qpos", "omni_scene"),
            ("Spider CEM", "spider_qpos", "spider_scene"),
        ]:
            metric_rows.append(
                evaluate_sequence(
                    case_id=case.get("case_id", ""),
                    method=method,
                    object_key=case.get("object_key", ""),
                    target_variant_id=case.get("target_variant_id", ""),
                    run_id=case.get("cem_run_id", ""),
                    qpos_path=item[qpos_key],
                    scene_xml=item[scene_key],
                )
            )
        validation_rows.extend(compare_history(case, summary, metric_rows[-1]))
    return metric_rows, validation_rows


def build_comparison(metric_rows: list[dict[str, Any]], validation_rows: list[dict[str, Any]], meta: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, dict[str, Any]]] = {}
    for row in metric_rows:
        key = row_key(row["case_id"], row["target_variant_id"], row["run_id"])
        grouped.setdefault(key, {})[row["method"]] = row
    bad_keys = {
        row_key(row["case_id"], "ref_fk", row["run_id"])
        for row in validation_rows
        if row["status"] != "PASS"
    }
    out = []
    for key in sorted(grouped):
        methods = grouped[key]
        omni = methods.get("OmniRetarget", {})
        spider = methods.get("Spider CEM", {})
        case_id = spider.get("case_id", omni.get("case_id", ""))
        target_variant_id = spider.get("target_variant_id", omni.get("target_variant_id", ""))
        run_id = spider.get("run_id", omni.get("run_id", ""))
        m = meta.get(key, {})
        out.append(
            {
                "分层": m.get("分层", ""),
                "纳入建议": m.get("纳入建议", ""),
                "是否strict主表": m.get("是否strict主表", ""),
                "case_id": case_id,
                "object_key": spider.get("object_key", omni.get("object_key", "")),
                "target_variant_id": target_variant_id,
                "spider_run_id": run_id,
                "omni_qpos_path": omni.get("qpos_path", ""),
                "spider_qpos_path": spider.get("qpos_path", ""),
                "Omni pelvis_min_m": finite_fmt(omni.get("pelvis_min_m")),
                "Spider pelvis_min_m": finite_fmt(spider.get("pelvis_min_m")),
                "Omni fall": omni.get("fall_flag", ""),
                "Spider fall": spider.get("fall_flag", ""),
                "Omni eef_near_5cm": finite_fmt(omni.get("eef_near_5cm_frac")),
                "Spider eef_near_5cm": finite_fmt(spider.get("eef_near_5cm_frac")),
                "Omni eef_near_8cm": finite_fmt(omni.get("eef_near_8cm_frac")),
                "Spider eef_near_8cm": finite_fmt(spider.get("eef_near_8cm_frac")),
                "Omni eef_near_10cm": finite_fmt(omni.get("eef_near_10cm_frac")),
                "Spider eef_near_10cm": finite_fmt(spider.get("eef_near_10cm_frac")),
                "Omni eef_near_12cm": finite_fmt(omni.get("eef_near_12cm_frac")),
                "Spider eef_near_12cm": finite_fmt(spider.get("eef_near_12cm_frac")),
                "Omni eef_near_15cm": finite_fmt(omni.get("eef_near_15cm_frac")),
                "Spider eef_near_15cm": finite_fmt(spider.get("eef_near_15cm_frac")),
                "Omni hand_deep_penetration_2cm": finite_fmt(omni.get("hand_geom_deep_penetration_2cm_frac")),
                "Spider hand_deep_penetration_2cm": finite_fmt(spider.get("hand_geom_deep_penetration_2cm_frac")),
                "Omni hand_geom_near_5cm": finite_fmt(omni.get("hand_geom_near_5cm_frac")),
                "Spider hand_geom_near_5cm": finite_fmt(spider.get("hand_geom_near_5cm_frac")),
                "Omni hand_geom_near_8cm": finite_fmt(omni.get("hand_geom_near_8cm_frac")),
                "Spider hand_geom_near_8cm": finite_fmt(spider.get("hand_geom_near_8cm_frac")),
                "Omni hand_geom_near_10cm": finite_fmt(omni.get("hand_geom_near_10cm_frac")),
                "Spider hand_geom_near_10cm": finite_fmt(spider.get("hand_geom_near_10cm_frac")),
                "Omni hand_geom_near_12cm": finite_fmt(omni.get("hand_geom_near_12cm_frac")),
                "Spider hand_geom_near_12cm": finite_fmt(spider.get("hand_geom_near_12cm_frac")),
                "Omni hand_geom_near_15cm": finite_fmt(omni.get("hand_geom_near_15cm_frac")),
                "Spider hand_geom_near_15cm": finite_fmt(spider.get("hand_geom_near_15cm_frac")),
                "Omni hand_geom_penetration": finite_fmt(omni.get("hand_geom_penetration_frac")),
                "Spider hand_geom_penetration": finite_fmt(spider.get("hand_geom_penetration_frac")),
                "Omni hand_object_physics_contact": finite_fmt(omni.get("hand_object_physics_contact_frac")),
                "Spider hand_object_physics_contact": finite_fmt(spider.get("hand_object_physics_contact_frac")),
                "Omni leg_penetration": finite_fmt(omni.get("leg_penetration_frac")),
                "Spider leg_penetration": finite_fmt(spider.get("leg_penetration_frac")),
                "Omni body_penetration": finite_fmt(omni.get("body_penetration_frac")),
                "Spider body_penetration": finite_fmt(spider.get("body_penetration_frac")),
                "Omni object_xy_displacement_m": finite_fmt(omni.get("object_xy_displacement_m")),
                "Spider object_xy_displacement_m": finite_fmt(spider.get("object_xy_displacement_m")),
                "Spider obj_err_mean_m": finite_fmt(spider.get("obj_err_mean_m")),
                "validation_status": "MISMATCH" if key in bad_keys else "PASS",
                "说明": m.get("说明", ""),
            }
        )
    return out


def tier_summary(comparison: list[dict[str, Any]]) -> list[dict[str, str]]:
    descriptions = {
        "A_strict_main": "原 11 条 ref_fk strict 主表 case",
        "B_upper_WORK_non_strict": "新增 13 条 ref_fk upper/object WORK，但 lower-body strict 未过",
    }
    counts = Counter((row["分层"], row["纳入建议"]) for row in comparison)
    return [
        {"分层": tier, "纳入建议": rec, "case数": str(count), "说明": descriptions.get(tier, "")}
        for (tier, rec), count in sorted(counts.items())
    ]


def write_xlsx(path: Path, comparison: list[dict[str, Any]], metrics: list[dict[str, Any]], summary: list[dict[str, Any]], tiers: list[dict[str, str]], validation: list[dict[str, Any]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    def add(name: str, rows: list[dict[str, Any]], fields: list[str]):
        ws = wb.create_sheet(name)
        ws.append(fields)
        for row in rows:
            ws.append([row.get(field, "") for field in fields])
        fill = PatternFill("solid", fgColor="1F4E78")
        font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, field in enumerate(fields, start=1):
            width = min(max(len(field) + 2, 10), 34)
            ws.column_dimensions[get_column_letter(idx)].width = width
        return ws

    case_ws = add("24case逐case对比", comparison, EXPANDED_FIELDS)
    header_to_col = {cell.value: cell.column for cell in case_ws[1]}
    for row_idx in range(2, case_ws.max_row + 1):
        for spider_field, omni_field, direction, threshold in COMPARISON_COLOR_RULES:
            spider_col = header_to_col.get(spider_field)
            omni_col = header_to_col.get(omni_field)
            if not spider_col or not omni_col:
                continue
            spider_cell = case_ws.cell(row=row_idx, column=spider_col)
            omni_cell = case_ws.cell(row=row_idx, column=omni_col)
            status = comparison_status(spider_cell.value, omni_cell.value, direction, threshold)
            if status == "better":
                spider_cell.fill = BETTER_FILL
            elif status == "worse":
                spider_cell.fill = WORSE_FILL
    ppt = add("PPT方法汇总24", ppt_method_summary(summary), PPT_SUMMARY_FIELDS)
    for col, width in {"A": 12, "B": 4, "C": 7, "D": 9, "E": 8, "F": 8, "G": 8, "H": 10, "I": 8, "J": 9, "K": 8}.items():
        ppt.column_dimensions[col].width = width
    add("方法汇总24", summary, list(summary[0].keys()) if summary else [])
    add("分层统计", tiers, TIER_FIELDS)
    legend = add("颜色规则", COLOR_RULE_ROWS, COLOR_RULE_FIELDS)
    legend_header = {cell.value: cell.column for cell in legend[1]}
    for row in range(2, legend.max_row + 1):
        legend.row_dimensions[row].height = 26
        legend.cell(row=row, column=legend_header["绿色含义"]).fill = BETTER_FILL
        legend.cell(row=row, column=legend_header["红色含义"]).fill = WORSE_FILL
    add("完整method_metrics", metrics, METRIC_FIELDS)
    add("历史对齐校验", validation, VALIDATION_FIELDS)
    wb.save(path)


def write_markdown(path: Path, comparison: list[dict[str, Any]], summary: list[dict[str, Any]], tiers: list[dict[str, str]], validation: list[dict[str, Any]]) -> None:
    mismatches = [row for row in validation if row["status"] != "PASS"]
    lines = [
        "# Expanded 24-case Spider work eval",
        "",
        "本表不覆盖原 11-case strict 主表；它额外加入 13 条 `ref_fk upper-WORK / non-strict` case。",
        "",
        f"- case 数：{len(comparison)}",
        f"- 历史对齐检查：{len(validation)}",
        f"- mismatch：{len(mismatches)}",
        "",
        "## 分层",
        "",
        "| 分层 | 纳入建议 | case数 | 说明 |",
        "|---|---|---:|---|",
    ]
    for row in tiers:
        lines.append(f"| `{row['分层']}` | {row['纳入建议']} | {row['case数']} | {row['说明']} |")
    lines.extend(["", "## 方法汇总", "", "| method | cases | pelvis | fall | hand10 | hand8 | hand5 | hand deep | leg pen | body pen | obj xy | obj err |", "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|"])
    for row in summary:
        lines.append(
            f"| {row['method']} | {row['num_cases']} | {row['mean_pelvis_min_m']} | {row['num_fall']} | "
            f"{row['mean_eef_near_10cm_frac']} | {row['mean_eef_near_8cm_frac']} | {row['mean_eef_near_5cm_frac']} | "
            f"{row['mean_hand_deep_penetration_2cm_frac']} | {row['mean_leg_penetration_frac']} | {row['mean_body_penetration_frac']} | "
            f"{row['mean_object_xy_displacement_m']} | {row['mean_obj_err_mean_m']} |"
        )
    lines.extend(["", "## 逐 case", "", "| tier | object | case | run | Spider hand10/8/5 | Spider deep/leg/body | validation |", "|---|---|---|---|---|---|---|"])
    for row in comparison:
        lines.append(
            f"| `{row['分层']}` | {row['object_key']} | `{row['case_id']}` | `{row['spider_run_id']}` | "
            f"{row['Spider eef_near_10cm']}/{row['Spider eef_near_8cm']}/{row['Spider eef_near_5cm']} | "
            f"{row['Spider hand_deep_penetration_2cm']}/{row['Spider leg_penetration']}/{row['Spider body_penetration']} | {row['validation_status']} |"
        )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--existing-cases", type=Path, default=REPO / "workspace/core4d/data_construction_v3/existing_cases.tsv")
    parser.add_argument("--audit-tsv", type=Path, default=REPO / "workspace/core4d/results/E109/spider_cem_work_audit/spider_cem_work_candidates.tsv")
    parser.add_argument("--out-dir", type=Path, default=REPO / "workspace/core4d/results/E109/expanded_24_work_cases")
    args = parser.parse_args()

    main_items, main_meta = current_main_items(args.existing_cases)
    extra_items, extra_meta = extra_upper_work_items(args.audit_tsv)
    all_items = main_items + extra_items
    meta = {**main_meta, **extra_meta}
    metric_rows, validation_rows = evaluate_items(all_items)
    comparison = build_comparison(metric_rows, validation_rows, meta)
    summary = method_summary(metric_rows)
    tiers = tier_summary(comparison)

    args.out_dir.mkdir(parents=True, exist_ok=True)
    write_tsv(args.out_dir / "expanded_24_case_comparison.tsv", comparison, EXPANDED_FIELDS)
    write_tsv(args.out_dir / "expanded_24_method_metrics.tsv", metric_rows, METRIC_FIELDS)
    write_tsv(args.out_dir / "expanded_24_method_summary.tsv", summary, list(summary[0].keys()) if summary else [])
    write_tsv(args.out_dir / "expanded_24_tier_summary.tsv", tiers, TIER_FIELDS)
    write_tsv(args.out_dir / "expanded_24_history_validation.tsv", validation_rows, VALIDATION_FIELDS)
    write_markdown(args.out_dir / "expanded_24_omni_vs_spider_work_eval.md", comparison, summary, tiers, validation_rows)
    write_full_markdown(args.out_dir / "expanded_24_case_comparison_full.md", "Expanded 24-case 完整逐 case 表", comparison, EXPANDED_FIELDS)
    write_xlsx(args.out_dir / "expanded_24_omni_vs_spider_work_eval.xlsx", comparison, metric_rows, summary, tiers, validation_rows)
    write_json(
        args.out_dir / "run_summary.json",
        {
            "num_cases": len(comparison),
            "num_strict_main_cases": sum(row["分层"] == "A_strict_main" for row in comparison),
            "num_upper_work_non_strict_cases": sum(row["分层"] == "B_upper_WORK_non_strict" for row in comparison),
            "num_metric_rows": len(metric_rows),
            "num_validation_rows": len(validation_rows),
            "num_validation_mismatch": sum(row["status"] != "PASS" for row in validation_rows),
            "note": "旧 11-case strict 表不覆盖；本输出是单独 expanded 24-case work 表。",
        },
    )
    print(f"[build_expanded_24_work_eval] wrote {len(comparison)} cases to {args.out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
