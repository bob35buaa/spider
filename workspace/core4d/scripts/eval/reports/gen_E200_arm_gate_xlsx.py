#!/usr/bin/env python3
"""E200 arm-comparison gate workbook: noPRG (E167A) vs PRG (E199 aug), same 249
translation-aug rollouts scored under one contract.

Reuses the E199 fullscale workbook machinery verbatim (gen_E199_fullscale_gate_xlsx:
the frozen E178 12-gate recompute via eval_E187_e178_compat, body_z p95 CPU FK,
release applicability, the 4 NEW acceptance gates, and the detail cell writer) so
this workbook's gate columns are byte-identical to E199 / E194 / E198.

Data sources (both scored by eval_E199_augmentation.score, same EvalConfig):
  * PRG   = E199 fullscale case_metrics, group==aug  (already on disk)
  * noPRG = E200 e200_noprg_case_metrics.tsv          (eval_E200_arm_augmentation)

Pairing key = (case_id, aug_variant); the two arms share the SAME aug trajectory
so every case/variant pairs 1:1. detail lists both arms per rollout (PRG then
noPRG, adjacent); summary reports per-gate pass rate noPRG vs PRG (overall +
per object) and key-metric means.

Usage:
    .venv/bin/python3 workspace/core4d/scripts/eval/reports/gen_E200_arm_gate_xlsx.py
"""

from __future__ import annotations

import argparse
import math
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[5]
sys.path.insert(0, str(REPO / "workspace/core4d/scripts"))
sys.path.insert(0, str(REPO / "workspace/core4d/scripts/eval/runners"))
sys.path.insert(0, str(REPO / "workspace/core4d/scripts/eval/reports"))
sys.path.insert(0, str(REPO / "workspace/core4d/scripts/experiments/E200"))

# reuse the E199 workbook machinery (12-gate recompute, cell writers, gate defs)
import gen_E199_fullscale_gate_xlsx as E199X  # noqa: E402
from gen_E199_fullscale_gate_xlsx import (  # noqa: E402
    GATE12, NEW_GATES, NAVY, HEAD_FONT, _cell, bool_fill, compute_row, finite, read_tsv,
)
import e200_common as C  # noqa: E402

PRG_CASE_METRICS = REPO / "workspace/core4d/results/E199/s6_downstream/eval/fullscale_augmentation/e199_fullscale_case_metrics.tsv"
PRG_MANIFEST = C.E199_FULLSCALE_MANIFEST
NOPRG_CASE_METRICS = C.RESULTS / "s6_downstream/eval/noprg/e200_noprg_case_metrics.tsv"
NOPRG_MANIFEST = C.manifest_path("noprg")
DEFAULT_OUT = C.RESULTS / "s6_downstream/eval/E200_noprg_vs_prg_gate_metrics.xlsx"

ARM_ORDER = {"PRG": 0, "noPRG": 1}
VARIANT_ORDER = {"trans0": 0, "trans1": 1, "trans2": 2}
KEY_METRICS_SUMMARY = [
    ("obj_pos_cm", "track_obj_pos_err_cm_mean"),
    ("obj_ori_deg", "track_obj_ori_err_deg_mean"),
    ("contact_in_mask", "hand_object_physics_contact_in_mask_frac"),
    ("hand_pen3mm", "hand_object_physics_penetration_3mm_frame_frac"),
    ("leg_pen", "leg_penetration_frac"),
    ("root_pos_cm", "track_root_pos_err_cm_mean"),
    ("eef_pos_cm", "track_eef_pos_err_cm_mean"),
    ("body_z_p95_m", "body_z_err_p95_m"),
]


def _norm(cid: str) -> str:
    return cid.replace("_person", "_p")


def load_arm(case_metrics: Path, manifest_path: Path, arm: str, group_filter: str | None) -> list[dict[str, Any]]:
    """Recompute the 12-gate + NEW gates for one arm's rollouts (reuses compute_row)."""
    man = {(m["case_id"], m["aug_variant"]): m for m in read_tsv(manifest_path)}
    out: list[dict[str, Any]] = []
    src = read_tsv(case_metrics)
    if group_filter:
        src = [r for r in src if r.get("group") == group_filter]
    for i, row in enumerate(src, start=1):
        r = compute_row({**row, "group": "aug"}, man)  # group=aug -> compute_row pulls traj from manifest
        r["arm"] = arm
        out.append(r)
        if i % 50 == 0 or i == len(src):
            print(f"[{arm}] recompute {i}/{len(src)}", file=sys.stderr, flush=True)
    return out


def write_detail(ws, rows: list[dict[str, Any]]) -> None:
    header: list[str] = ["object_key", "case_id", "arm", "aug_variant"]
    for _lbl, _key, _fld, _op, val_hdr in GATE12:
        header += [val_hdr, "→pass"]
    header += ["gate12_all_pass", "gate12_failure_modes", "qpos_jerk_l2_p95", "foot_slip_max_m"]
    new_val_hdr = {"NG_contact3mm_in_mask>=0.40": "contact3mm_in_mask",
                   "NG_hand_pen_3mm<=0.32": "hand_pen3mm",
                   "NG_ankle_jerk_p95<1000": "ankle_jerk_p95",
                   "NG_obj_speed_max<3": "obj_speed_max"}
    for label, *_ in NEW_GATES:
        header += [new_val_hdr[label], label]
    header += ["new_gates_all_pass", "note"]

    for c, name in enumerate(header, start=1):
        _cell(ws, 1, c, name); ws.cell(row=1, column=c).fill = NAVY
        ws.cell(row=1, column=c).font = HEAD_FONT
        ws.cell(row=1, column=c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "E2"

    def num(r, c, v):
        v = finite(v); _cell(ws, r, c, round(v, 4) if math.isfinite(v) else "")

    for r, row in enumerate(rows, start=2):
        c = 1
        for key in ("object_key", "case_id", "arm", "aug_variant"):
            _cell(ws, r, c, row.get(key, "")); c += 1
        for _lbl, key, field, _op, _vh in GATE12:
            if field == "fall_flag":
                _cell(ws, r, c, str(row.get("fall_flag", "")))
            else:
                num(r, c, row.get(field))
            c += 1
            v = row[f"{key}_gate_pass"]; _cell(ws, r, c, "PASS" if v else "FAIL", bool_fill(v)); c += 1
        _cell(ws, r, c, "PASS" if row["gate12_all_pass"] else "FAIL", bool_fill(row["gate12_all_pass"])); c += 1
        _cell(ws, r, c, row["gate12_failure_modes"]); c += 1
        num(r, c, row.get("qpos_jerk_l2_p95")); c += 1
        num(r, c, row.get("foot_slip_max_m")); c += 1
        for label, field, _comp, _thr in NEW_GATES:
            num(r, c, row.get(field)); c += 1
            v = row[label]; _cell(ws, r, c, "PASS" if v else "FAIL", bool_fill(v)); c += 1
        _cell(ws, r, c, "PASS" if row["new_gates_all_pass"] else "FAIL", bool_fill(row["new_gates_all_pass"])); c += 1
        _cell(ws, r, c, row.get("_note", ""))

    widths = [10, 34, 8, 12] + [14, 7] * len(GATE12) + [12, 22, 15, 14] + [16, 20] * len(NEW_GATES) + [16, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _rate(rows, key):
    return sum(1 for r in rows if r.get(key) is True), len(rows)


def write_summary(ws, noprg: list[dict[str, Any]], prg: list[dict[str, Any]]) -> None:
    objects = sorted({r["object_key"] for r in noprg + prg})
    gate_keys = ([("gate12_all", "gate12_all_pass")]
                 + [(lbl, f"{key}_gate_pass") for lbl, key, *_ in GATE12]
                 + [(lbl, lbl) for lbl, *_ in NEW_GATES]
                 + [("new_gates_all", "new_gates_all_pass")])

    def block(title, np_rows, prg_rows, r0):
        _cell(ws, r0, 1, title); ws.cell(row=r0, column=1).font = Font(bold=True)
        heads = ["gate", "noPRG pass", "noPRG rate", "PRG pass", "PRG rate", "Δrate(noPRG−PRG) pp"]
        for c, h in enumerate(heads, start=1):
            _cell(ws, r0 + 1, c, h, NAVY); ws.cell(row=r0 + 1, column=c).font = HEAD_FONT
        r = r0 + 2
        for lbl, key in gate_keys:
            nn, nt = _rate(np_rows, key); pn, pt = _rate(prg_rows, key)
            nr = nn / nt if nt else float("nan"); pr = pn / pt if pt else float("nan")
            _cell(ws, r, 1, lbl); _cell(ws, r, 2, f"{nn}/{nt}"); _cell(ws, r, 3, round(nr, 3) if nt else "")
            _cell(ws, r, 4, f"{pn}/{pt}"); _cell(ws, r, 5, round(pr, 3) if pt else "")
            _cell(ws, r, 6, round((nr - pr) * 100, 1) if nt and pt else "")
            r += 1
        # key-metric means (mean over rollouts)
        _cell(ws, r, 1, "— key-metric mean —"); ws.cell(row=r, column=1).font = Font(italic=True); r += 1
        import numpy as np
        for lbl, field in KEY_METRICS_SUMMARY:
            def mean(rows):
                v = np.array([finite(x.get(field)) for x in rows]); v = v[np.isfinite(v)]
                return float(v.mean()) if v.size else float("nan")
            nm, pm = mean(np_rows), mean(prg_rows)
            _cell(ws, r, 1, lbl); _cell(ws, r, 2, round(nm, 4)); _cell(ws, r, 4, round(pm, 4))
            _cell(ws, r, 6, round(nm - pm, 4) if math.isfinite(nm) and math.isfinite(pm) else "")
            r += 1
        return r + 1

    row_ptr = block(f"OVERALL (noPRG {len(noprg)} vs PRG {len(prg)})", noprg, prg, 1)
    for obj in objects:
        row_ptr = block(obj, [r for r in noprg if r["object_key"] == obj],
                        [r for r in prg if r["object_key"] == obj], row_ptr)
    ws.column_dimensions["A"].width = 26
    for col in "BCDEF":
        ws.column_dimensions[col].width = 15


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()

    prg = load_arm(PRG_CASE_METRICS, PRG_MANIFEST, "PRG", group_filter="aug")
    noprg = load_arm(NOPRG_CASE_METRICS, NOPRG_MANIFEST, "noPRG", group_filter=None)

    combined = prg + noprg
    combined.sort(key=lambda r: (r["object_key"], _norm(r["case_id"]),
                                 VARIANT_ORDER.get(r["aug_variant"], 9), ARM_ORDER.get(r["arm"], 9)))

    wb = Workbook()
    write_detail(wb.active, combined); wb.active.title = "detail"
    write_summary(wb.create_sheet("summary"), noprg, prg)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"[done] wrote {args.out} (noPRG {len(noprg)} + PRG {len(prg)} = {len(combined)} rows)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
