#!/usr/bin/env python3
"""Create E167 vs E166 vs E163 comparison workbook."""
from __future__ import annotations

import csv
import math
from pathlib import Path
from typing import Any

import mujoco
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


REPO = Path(__file__).resolve().parents[5]
OUT = REPO / "workspace/core4d/results/E167/holosoma_zonly/eval/comparison/E167_vs_E166_vs_E163_cem_offline_eval.xlsx"

E167_METRICS = REPO / "workspace/core4d/results/E167/holosoma_zonly/eval/cem_metrics/full/e167_arm_metrics.tsv"
E167_ZGATE = REPO / "workspace/core4d/results/E167/holosoma_zonly/eval/zgate/full/holosoma_zgate.tsv"
E167_AXIS = REPO / "workspace/core4d/results/E167/holosoma_zonly/eval/axis_audit/full/axis_audit.tsv"
E167_ARTIFACT = REPO / "workspace/core4d/results/E167/holosoma_zonly/eval/cem_metrics/full/e167_artifact_check.tsv"
E166_METRICS = [
    REPO / "workspace/core4d/results/E166/foot_smooth_retarget/eval/full/e166_arm_metrics.tsv",
    REPO / "workspace/core4d/results/E166/foot_smooth_retarget/eval/remaining4/e166_arm_metrics.tsv",
]
E166_ARTIFACTS = [
    REPO / "workspace/core4d/results/E166/foot_smooth_retarget/eval/full/e166_artifact_check.tsv",
    REPO / "workspace/core4d/results/E166/foot_smooth_retarget/eval/remaining4/e166_artifact_check.tsv",
]

CASE_ORDER = [
    "box023_person2",
    "box021_029_p2",
    "box021_035_p1",
    "box021_035_p2",
    "box004_082_p1",
    "box004_083_p1",
    "box004_083_p2",
]

METHODS = [
    {
        "key": "E163",
        "display": "E163 baseline",
        "source_exp": "E163",
        "arm": "baseline",
        "method_note": "E163 narrowSurfaceBand clean8；作为 7case 共同 baseline。",
    },
    {
        "key": "E166_A",
        "display": "E166 A",
        "source_exp": "E166",
        "arm": "A",
        "method_note": "E166 足/踝约束：包含 foot-slip XY 与 3D ankle extra-weight。",
    },
    {
        "key": "E166_A_B2",
        "display": "E166 A_B2_postSmooth",
        "source_exp": "E166",
        "arm": "A_B2_postSmooth",
        "method_note": "E166 A 后 qpos-wide postSmooth；不是 z-only。",
    },
    {
        "key": "E167A",
        "display": "E167A",
        "source_exp": "E167",
        "arm": "E167A",
        "method_note": "Holosoma 对齐：只做 body/ground z-only 约束，不限制 XY。",
    },
    {
        "key": "E167A_B1",
        "display": "E167A_B1",
        "source_exp": "E167",
        "arm": "E167A_B1",
        "method_note": "E167A + CEM-side body-z accel/jerk smooth。",
    },
    {
        "key": "E167A_B2",
        "display": "E167A_B2",
        "source_exp": "E167",
        "arm": "E167A_B2",
        "method_note": "E167A + root-z handoff postprocess；qpos/qvel non-z delta=0。",
    },
]

BODY_NAMES = [
    "left_ankle_roll_link",
    "right_ankle_roll_link",
    "left_wrist_yaw_link",
    "right_wrist_yaw_link",
]
Z_THRESHOLD_M = 0.25
SUGAR_3D_THRESHOLD_M = 0.30

SUMMARY_COLUMNS = [
    ("display", "方法"),
    ("source_exp", "来源"),
    ("n_cases", "case数"),
    ("pass_cases", "通过case"),
    ("failed_cases", "失败case"),
    ("tracked_cases", "tracked"),
    ("fall_cases", "fall"),
    ("raw_contact", "物理接触(raw)"),
    ("raw_delta", "raw Δ vs E163"),
    ("rl_contact", "RL mask接触"),
    ("clean3_contact", "clean3接触"),
    ("clean3_delta", "clean3 Δ"),
    ("clean5_contact", "clean5接触"),
    ("pen3", "物理穿透3mm"),
    ("pen3_delta", "物理穿透3mm Δ"),
    ("geom_pen2", "几何穿透2mm"),
    ("release_false3", "release误接触3mm"),
    ("joint_err", "关节误差(°)"),
    ("eef_err", "EEF误差(cm)"),
    ("obj_err", "物体位置误差(cm)"),
    ("root_err", "root误差(cm)"),
    ("qpos_jerk", "qpos jerk p95"),
    ("qpos_jerk_delta", "qpos jerk Δ"),
    ("trackbody_jerk", "trackbody jerk p95"),
    ("trackbody_jerk_delta", "trackbody jerk Δ"),
    ("ankle_acc", "ankle acc max"),
    ("ankle_acc_delta", "ankle acc Δ"),
    ("obj_speed", "obj speed max"),
    ("foot_slip", "foot slip max(m)"),
    ("foot_slip_delta", "foot slip Δ"),
    ("grounded_frac", "grounded frame frac"),
    ("z_pass", "Holosoma z pass"),
    ("sugar3d_pass", "SUGAR 3D offline pass"),
    ("z_peak_cm", "body-z peak(cm)"),
    ("xy_peak_cm", "xy peak(cm)"),
    ("note", "说明"),
]

CASE_COLUMNS = [
    ("display", "方法"),
    ("case", "case"),
    ("source_exp", "来源"),
    ("status", "状态"),
    ("spider_gate", "SPIDER gate"),
    ("contact_no_reg", "contact no-reg"),
    ("penetration_no_reg", "penetration no-reg"),
    ("raw_contact", "raw接触"),
    ("raw_delta", "raw Δ"),
    ("clean3_contact", "clean3接触"),
    ("clean3_delta", "clean3 Δ"),
    ("pen3", "物理穿透3mm"),
    ("pen3_delta", "物理穿透3mm Δ"),
    ("geom_pen2", "几何穿透2mm"),
    ("release_false3", "release误接触3mm"),
    ("tracked", "tracked"),
    ("fall", "fall"),
    ("joint_err", "关节误差(°)"),
    ("eef_err", "EEF误差(cm)"),
    ("obj_err", "物体位置误差(cm)"),
    ("root_err", "root误差(cm)"),
    ("qpos_jerk", "qpos jerk p95"),
    ("qpos_jerk_delta", "qpos jerk Δ"),
    ("trackbody_jerk", "trackbody jerk p95"),
    ("trackbody_jerk_delta", "trackbody jerk Δ"),
    ("foot_slip", "foot slip max(m)"),
    ("foot_slip_delta", "foot slip Δ"),
    ("z_pass", "Holosoma z pass"),
    ("z_peak_cm", "body-z peak(cm)"),
    ("xy_peak_cm", "xy peak(cm)"),
    ("qpos_path", "qpos_path"),
    ("video", "video"),
]

METRIC_MAP = {
    "raw_contact": "hand_object_physics_contact_in_mask_frac",
    "raw_delta": "hand_object_physics_contact_in_mask_frac_delta_vs_baseline",
    "rl_contact": "hand_object_physics_contact_in_rl_mask_frac",
    "clean3_contact": "hand_object_physics_contact_3mm_in_mask_frac",
    "clean3_delta": "hand_object_physics_contact_3mm_in_mask_frac_delta_vs_baseline",
    "clean5_contact": "hand_object_physics_contact_5mm_in_mask_frac",
    "pen3": "hand_object_physics_penetration_3mm_frame_frac",
    "pen3_delta": "hand_object_physics_penetration_3mm_frame_frac_delta_vs_baseline",
    "geom_pen2": "hand_geom_penetration_2mm_frac",
    "release_false3": "hand_object_release_false_contact_3mm_frac",
    "joint_err": "track_joint_err_deg_mean",
    "eef_err": "track_eef_pos_err_cm_mean",
    "obj_err": "track_obj_pos_err_cm_mean",
    "root_err": "track_root_pos_err_cm_mean",
    "qpos_jerk": "qpos_jerk_l2_p95",
    "qpos_jerk_delta": "qpos_jerk_l2_p95_delta_vs_baseline",
    "trackbody_jerk": "trackbody_jerk_p95",
    "trackbody_jerk_delta": "trackbody_jerk_p95_delta_vs_baseline",
    "ankle_acc": "ankle_acc_max",
    "ankle_acc_delta": "ankle_acc_max_delta_vs_baseline",
    "obj_speed": "obj_speed_max",
    "foot_slip": "foot_slip_max_m",
    "foot_slip_delta": "foot_slip_max_m_delta_vs_baseline",
    "grounded_frac": "foot_grounded_frame_frac",
}


def rel(path: str | Path) -> str:
    p = Path(path)
    try:
        return str(p.resolve().relative_to(REPO.resolve()))
    except Exception:
        return str(path)


def repo_path(path: str | Path) -> Path:
    p = Path(path)
    return p if p.is_absolute() else REPO / p


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def num(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if text == "":
        return None
    try:
        value_f = float(text)
    except Exception:
        return None
    return value_f if math.isfinite(value_f) else None


def mean(values: list[Any]) -> float | None:
    vals = [num(v) for v in values]
    vals = [v for v in vals if v is not None]
    return float(np.mean(vals)) if vals else None


def boolish(value: Any) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "on", "是", "pass", "通过"}


def cn_bool(value: Any) -> str:
    if value in ("", None):
        return ""
    return "是" if boolish(value) else "否"


def metric(row: dict[str, Any], key: str) -> Any:
    return row.get(METRIC_MAP[key], "")


def load_metric_rows() -> list[dict[str, Any]]:
    e167_rows = read_tsv(E167_METRICS)
    e166_rows: list[dict[str, Any]] = []
    for path in E166_METRICS:
        e166_rows.extend(read_tsv(path))

    out: list[dict[str, Any]] = []
    for spec in METHODS:
        source_rows = e167_rows if spec["source_exp"] in {"E163", "E167"} else e166_rows
        for row in source_rows:
            if row.get("arm") != spec["arm"]:
                continue
            if row.get("short_case_id") not in CASE_ORDER:
                continue
            item = dict(row)
            item.update(
                {
                    "compare_key": spec["key"],
                    "display": spec["display"],
                    "source_exp_compare": spec["source_exp"],
                    "method_note": spec["method_note"],
                    "case_order": CASE_ORDER.index(row["short_case_id"]),
                    "method_order": METHODS.index(spec),
                }
            )
            if spec["key"] == "E167A_B2":
                qpos = repo_path(item.get("qpos_path", ""))
                video = qpos.with_name(f"{qpos.stem}_full.mp4")
                if video.is_file():
                    item["video"] = rel(video)
            out.append(item)
    return sorted(out, key=lambda r: (int(r["method_order"]), int(r["case_order"])))


def p95(values: np.ndarray) -> float | None:
    vals = np.asarray(values, dtype=np.float64)
    vals = vals[np.isfinite(vals)]
    return float(np.percentile(vals, 95)) if vals.size else None


def load_qpos_pair(path: Path) -> tuple[np.ndarray, np.ndarray]:
    with np.load(path, allow_pickle=True) as data:
        qpos = np.asarray(data["qpos"], dtype=np.float64)
    if qpos.ndim == 3 and qpos.shape[1] >= 2:
        return qpos[:, 0, :], qpos[:, 1, :]
    raise ValueError(f"expected qpos (T,2,nq), got {qpos.shape} in {path}")


def body_positions(scene_xml: Path, qpos: np.ndarray) -> np.ndarray:
    model = mujoco.MjModel.from_xml_path(str(scene_xml))
    data = mujoco.MjData(model)
    if qpos.shape[1] != model.nq:
        raise ValueError(f"{scene_xml} model.nq={model.nq}, qpos width={qpos.shape[1]}")
    body_ids = []
    for name in BODY_NAMES:
        bid = mujoco.mj_name2id(model, mujoco.mjtObj.mjOBJ_BODY, name)
        if bid >= 0:
            body_ids.append(bid)
    out = np.zeros((qpos.shape[0], len(body_ids), 3), dtype=np.float64)
    for i, q in enumerate(qpos):
        data.qpos[:] = q
        data.qvel[:] = 0.0
        mujoco.mj_forward(model, data)
        out[i] = data.xpos[body_ids]
    return out


def compute_zgate(row: dict[str, Any]) -> dict[str, Any]:
    path = repo_path(row.get("qpos_path", row.get("result_npz", "")))
    scene_text = row.get("scene_xml") or row.get("scene") or ""
    scene = repo_path(scene_text)
    sim, ref = load_qpos_pair(path)
    n = min(sim.shape[0], ref.shape[0])
    sim_pos = body_positions(scene, sim[:n])
    ref_pos = body_positions(scene, ref[:n])
    diff = sim_pos - ref_pos
    z_err = np.abs(diff[..., 2])
    xy_err = np.linalg.norm(diff[..., :2], axis=-1)
    err3d = np.linalg.norm(diff, axis=-1)
    z_peak = float(np.max(z_err))
    err3d_peak = float(np.max(err3d))
    return {
        "compare_key": row["compare_key"],
        "display": row["display"],
        "source_exp": row["source_exp_compare"],
        "case": row["short_case_id"],
        "arm": row["arm"],
        "qpos_path": rel(path),
        "scene": rel(scene),
        "frames": n,
        "body_z_err_peak_m": z_peak,
        "body_z_err_p95_m": p95(z_err),
        "body_z_over_frac": float(np.mean(z_err > Z_THRESHOLD_M)),
        "holosoma_z_gate_pass": z_peak <= Z_THRESHOLD_M,
        "sugar_3d_err_peak_m": err3d_peak,
        "sugar_3d_err_p95_m": p95(err3d),
        "sugar_3d_over_frac": float(np.mean(err3d > SUGAR_3D_THRESHOLD_M)),
        "sugar_3d_gate_pass": err3d_peak <= SUGAR_3D_THRESHOLD_M,
        "xy_err_peak_m": float(np.max(xy_err)),
        "xy_err_p95_m": p95(xy_err),
        "xy_only_failure": err3d_peak > SUGAR_3D_THRESHOLD_M and z_peak <= Z_THRESHOLD_M,
    }


def build_zgate(metric_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    scene_by_qpos = {}
    scene_by_case = {}
    for row in read_tsv(E167_ZGATE):
        scene_by_qpos[row["qpos_path"]] = row.get("scene", "")
        scene_by_case.setdefault(row["short_case_id"], row.get("scene", ""))

    rows = []
    for row in metric_rows:
        item = dict(row)
        if not item.get("scene_xml"):
            item["scene"] = scene_by_qpos.get(item.get("qpos_path", ""), "") or scene_by_case.get(
                item.get("short_case_id", ""), ""
            )
        else:
            item["scene"] = item["scene_xml"]
        rows.append(compute_zgate(item))
    return rows


def summarize(metric_rows: list[dict[str, Any]], zgate_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    z_by_method = {}
    for zg in zgate_rows:
        z_by_method.setdefault(zg["compare_key"], []).append(zg)
    out = []
    for spec in METHODS:
        rows = [r for r in metric_rows if r["compare_key"] == spec["key"]]
        zrows = z_by_method.get(spec["key"], [])
        failed = [r["short_case_id"] for r in rows if not boolish(r.get("spider_gate_pass"))]
        item = {
            "display": spec["display"],
            "source_exp": spec["source_exp"],
            "n_cases": len(rows),
            "pass_cases": sum(1 for r in rows if boolish(r.get("spider_gate_pass"))),
            "failed_cases": ",".join(failed),
            "tracked_cases": sum(1 for r in rows if boolish(r.get("success_tracked"))),
            "fall_cases": sum(1 for r in rows if boolish(r.get("fall_flag"))),
            "z_pass": f"{sum(1 for r in zrows if r['holosoma_z_gate_pass'])}/{len(zrows)}" if zrows else "",
            "sugar3d_pass": f"{sum(1 for r in zrows if r['sugar_3d_gate_pass'])}/{len(zrows)}" if zrows else "",
            "z_peak_cm": mean([r["body_z_err_peak_m"] * 100 for r in zrows]),
            "xy_peak_cm": mean([r["xy_err_peak_m"] * 100 for r in zrows]),
            "note": spec["method_note"],
        }
        for out_key in METRIC_MAP:
            item[out_key] = mean([metric(r, out_key) for r in rows])
        out.append(item)
    return out


def per_case(metric_rows: list[dict[str, Any]], zgate_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    zmap = {(z["compare_key"], z["case"]): z for z in zgate_rows}
    rows = []
    for r in metric_rows:
        zg = zmap.get((r["compare_key"], r["short_case_id"]), {})
        rows.append(
            {
                "display": r["display"],
                "case": r["short_case_id"],
                "source_exp": r["source_exp_compare"],
                "status": "baseline" if r["arm"] == "baseline" else ("通过" if boolish(r.get("spider_gate_pass")) else "失败"),
                "spider_gate": cn_bool(r.get("spider_gate_pass")),
                "contact_no_reg": cn_bool(r.get("contact_no_regression")),
                "penetration_no_reg": cn_bool(r.get("penetration_no_regression")),
                "tracked": cn_bool(r.get("success_tracked")),
                "fall": cn_bool(r.get("fall_flag")),
                "z_pass": cn_bool(zg.get("holosoma_z_gate_pass")),
                "z_peak_cm": num(zg.get("body_z_err_peak_m")) * 100 if num(zg.get("body_z_err_peak_m")) is not None else "",
                "xy_peak_cm": num(zg.get("xy_err_peak_m")) * 100 if num(zg.get("xy_err_peak_m")) is not None else "",
                "qpos_path": r.get("qpos_path", ""),
                "video": r.get("video", ""),
                **{key: num(metric(r, key)) for key in METRIC_MAP},
            }
        )
    return rows


def collect_artifacts(metric_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    e167_art = read_tsv(E167_ARTIFACT)
    e166_art = []
    for p in E166_ARTIFACTS:
        e166_art.extend(read_tsv(p))
    selected = {(r["source_exp_compare"], r["short_case_id"], r["arm"]) for r in metric_rows}
    rows = []
    for src, source_rows in [("E167", e167_art), ("E166", e166_art)]:
        for row in source_rows:
            key = (src if row["arm"] != "baseline" else "E163", row["short_case_id"], row["arm"])
            if key not in selected:
                continue
            item = dict(row)
            item["source_exp"] = key[0]
            if key[0] == "E167" and item.get("arm") == "E167A_B2":
                qpos = repo_path(item.get("qpos_path", ""))
                video = qpos.with_name(f"{qpos.stem}_full.mp4")
                item["video"] = rel(video) if video.is_file() else item.get("video", "")
                item["video_exists"] = str(video.is_file()).lower()
            rows.append(item)
    return rows


def write_table(ws, rows: list[dict[str, Any]], columns: list[tuple[str, str]]) -> None:
    ws.append([label for _, label in columns])
    for row in rows:
        ws.append([row.get(key, "") for key, _ in columns])
    style_sheet(ws)


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(bottom=thin)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(cell.value, float):
                cell.number_format = "0.000"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = 8
        for cell in col:
            max_len = max(max_len, min(len(str(cell.value)) if cell.value is not None else 0, 60))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 42)


def add_hyperlinks(ws, header_names: set[str]) -> None:
    headers = {cell.value: cell.column for cell in ws[1]}
    for header in header_names:
        col = headers.get(header)
        if not col:
            continue
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            if isinstance(cell.value, str) and cell.value and cell.value != ".":
                target = repo_path(cell.value)
                if target.exists():
                    cell.hyperlink = str(target)
                    cell.style = "Hyperlink"


def write_raw(ws, rows: list[dict[str, Any]]) -> None:
    fields = []
    for row in rows:
        for key in row:
            if key not in fields:
                fields.append(key)
    write_table(ws, rows, [(f, f) for f in fields])
    add_hyperlinks(ws, {"qpos_path", "video", "config_act", "result_npz"})


def write_notes(ws) -> None:
    sugar_root = REPO.parent / "Loco-Manipulation/SUGAR/outputs/core4d/e167_zonly_refiner_rl"
    final_ckpt = len(list(sugar_root.glob("*/**/model_5999.pt"))) if sugar_root.exists() else 0
    success_csv = (
        len(list(sugar_root.glob("*/eval_staggered_phase_mw30/analysis/success_vs_phase.csv")))
        if sugar_root.exists()
        else 0
    )
    rows = [
        ("范围", "7 case: clean8 去除 box026_139_p1；CEM/offline metrics 为当前 workbook 主口径。"),
        ("主表方法", "E163 baseline、E166 A、E166 A_B2_postSmooth、E167A、E167A_B1、E167A_B2。"),
        ("E166来源", "E166 A/A_B2_postSmooth 由 full 三 case + remaining4 四 case 合并为 7case；E166 B1/B2/AplusB 不是 7case 完整集合，未放入主表。"),
        ("Holosoma z-only", f"离线重算踝+腕 body-z gate，阈值 {Z_THRESHOLD_M:.2f}m；SUGAR 3D gate 仅作诊断，阈值 {SUGAR_3D_THRESHOLD_M:.2f}m。"),
        ("E167 axis caveat", "E167A/B1/B2 轴审计通过：foot-slip XY 关闭，ankle 3D extra-weight 关闭；B2 qpos/qvel non-z delta 为 0。"),
        ("SUGAR downstream", f"E167 SUGAR downstream 仍未完成，当前本地可见 model_5999.pt={final_ckpt}/21、success_vs_phase.csv={success_csv}/21；本 xlsx 不含最终 SUGAR 成功率。"),
        ("生成脚本", "workspace/core4d/scripts/eval/reports/gen_E167_vs_E166_E163_xlsx.py"),
    ]
    write_table(ws, [{"项": k, "说明": v} for k, v in rows], [("项", "项"), ("说明", "说明")])


def validate_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=False)
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    bad.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    if bad:
        raise RuntimeError("formula-like errors found: " + "; ".join(bad[:20]))


def main() -> int:
    metric_rows = load_metric_rows()
    zgate_rows = build_zgate(metric_rows)
    summary_rows = summarize(metric_rows, zgate_rows)
    case_rows = per_case(metric_rows, zgate_rows)
    artifact_rows = collect_artifacts(metric_rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "主表"
    write_table(ws, summary_rows, SUMMARY_COLUMNS)

    ws = wb.create_sheet("逐case")
    write_table(ws, case_rows, CASE_COLUMNS)
    add_hyperlinks(ws, {"qpos_path", "video"})

    ws = wb.create_sheet("HolosomaZ门")
    z_cols = [
        ("display", "方法"),
        ("case", "case"),
        ("source_exp", "来源"),
        ("holosoma_z_gate_pass", "Holosoma z pass"),
        ("body_z_err_peak_m", "body-z peak(m)"),
        ("body_z_err_p95_m", "body-z p95(m)"),
        ("body_z_over_frac", "body-z over frac"),
        ("sugar_3d_gate_pass", "SUGAR 3D pass"),
        ("sugar_3d_err_peak_m", "SUGAR 3D peak(m)"),
        ("xy_err_peak_m", "xy peak(m)"),
        ("xy_only_failure", "xy-only failure"),
        ("qpos_path", "qpos_path"),
        ("scene", "scene"),
    ]
    write_table(ws, zgate_rows, z_cols)
    add_hyperlinks(ws, {"qpos_path", "scene"})

    ws = wb.create_sheet("产物与轴审计")
    write_raw(ws, artifact_rows + read_tsv(E167_AXIS))

    ws = wb.create_sheet("说明")
    write_notes(ws)

    ws = wb.create_sheet("原始metrics")
    raw_rows = []
    zmap = {(z["compare_key"], z["case"]): z for z in zgate_rows}
    for row in metric_rows:
        item = dict(row)
        item.update({f"zgate_{k}": v for k, v in zmap[(row["compare_key"], row["short_case_id"])].items()})
        raw_rows.append(item)
    write_raw(ws, raw_rows)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    validate_workbook(OUT)
    print(OUT)
    print(f"summary_rows={len(summary_rows)} case_rows={len(case_rows)} zgate_rows={len(zgate_rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
