#!/usr/bin/env python3
"""Export the E199 full-scale augmentation metrics workbook.

Emits an xlsx that carries, per rollout (249 aug + 83 orig):
  * the frozen E178 canonical **12 gates** (fall / body_z / contact / release /
    hand_penetration / lower_body / root_pos / root_ori / hand_pos / hand_ori /
    object_pos / object_ori) — body_z + release are recomputed here because the
    E199 fullscale case_metrics only stored the reduced 6-gate set;
  * four smoothness/health metrics: qpos_jerk_l2_p95, ankle_jerk_p95,
    obj_speed_max, foot_slip_max_m;
  * four NEW acceptance gates requested for this export:
      - contact_3mm_in_mask >= 0.40
      - hand-object 3mm penetration frame frac <= 0.32
      - ankle jerk p95 < 1000
      - object max speed < 3.0

Gate definitions/thresholds reuse the tested E178-compat helpers
(`eval_E187_e178_compat`) so the 12-gate here matches the viser review player
and E194/E198 arm sweeps exactly.

Usage:
    .venv/bin/python3 workspace/core4d/scripts/eval/reports/gen_E199_fullscale_gate_xlsx.py
"""

from __future__ import annotations

import argparse
import csv
import math
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[5]
sys.path.insert(0, str(REPO / "workspace/core4d/scripts"))
sys.path.insert(0, str(REPO / "workspace/core4d/scripts/eval/runners"))

# Reuse the frozen E178 gate machinery (body_z recompute, release applicability,
# thresholds) so this workbook's 12-gate is byte-identical to the review player.
import eval_E187_e178_compat as G  # noqa: E402

EVAL_DIR = REPO / "workspace/core4d/results/E199/s6_downstream/eval/fullscale_augmentation"
CASE_METRICS = EVAL_DIR / "e199_fullscale_case_metrics.tsv"
MANIFEST = (
    REPO
    / "workspace/core4d/results/E199/s6_downstream/manifests/e199_fullscale_priority_manifest.tsv"
)
CONTACT_MASK_ROOT = REPO / "workspace/core4d/results/E199/data_preprocess/contact_masks"
DEFAULT_OUT = EVAL_DIR.parent / "E199_fullscale_gate_metrics.xlsx"

# ---- the canonical E178 12 gates, in review-player order ---------------------
# (label, gate_key from G.apply_gates, underlying numeric field, op, value-col header)
GATE12 = [
    ("fall", "fall", "fall_flag", "fall", "fall_flag"),
    ("body_z", "body_z", "body_z_err_p95_m", "<=", "body_z_p95_m(≤0.20)"),
    ("contact", "contact", "hand_object_physics_contact_in_mask_frac", ">=", "contact_in_mask(≥0.50)"),
    ("release", "release", "hand_object_release_false_contact_3mm_frac", "<=", "release_fc3mm(≤0.30)"),
    ("hand_penetration", "hand_penetration", "hand_object_physics_penetration_3mm_frame_frac", "<=", "hand_pen3mm(≤0.30)"),
    ("lower_body", "lower_body", "leg_penetration_frac", "<=", "leg_pen(≤0.10)"),
    ("root_pos", "root_pos", "track_root_pos_err_cm_mean", "<=", "root_pos_cm(≤20)"),
    ("root_ori", "root_ori", "track_root_ori_err_deg_mean", "<=", "root_ori_deg(≤20)"),
    ("hand_pos", "hand_pos", "track_eef_pos_err_cm_mean", "<=", "eef_pos_cm(≤20)"),
    ("hand_ori", "hand_ori", "track_eef_ori_err_deg_mean", "<=", "eef_ori_deg(≤20)"),
    ("object_pos", "object_pos", "track_obj_pos_err_cm_mean", "<=", "obj_pos_cm(≤20)"),
    ("object_ori", "object_ori", "track_obj_ori_err_deg_mean", "<=", "obj_ori_deg(≤10)"),
]

METRICS4 = [
    ("qpos_jerk_l2_p95", "qpos_jerk_l2_p95"),
    ("ankle_jerk_p95", "ankle_jerk_p95"),
    ("obj_speed_max", "obj_speed_max"),
    ("foot_slip_max_m", "foot_slip_max_m"),
]

# NEW gates for this export: (label, field, comparator, threshold)
NEW_GATES = [
    ("NG_contact3mm_in_mask>=0.40", "hand_object_physics_contact_3mm_in_mask_frac", ">=", 0.40),
    ("NG_hand_pen_3mm<=0.32", "hand_object_physics_penetration_3mm_frame_frac", "<=", 0.32),
    ("NG_ankle_jerk_p95<1000", "ankle_jerk_p95", "<", 1000.0),
    ("NG_obj_speed_max<3", "obj_speed_max", "<", 3.0),
]

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
GREY = PatternFill("solid", fgColor="D9D9D9")
NAVY = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def finite(value: Any, default: float = math.nan) -> float:
    try:
        f = float(value)
        return f if math.isfinite(f) else default
    except (TypeError, ValueError):
        return default


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh, delimiter="\t"))


def actor_index(case_id: str) -> int:
    """CORE4D actor: person1/_p1 -> 0, person2/_p2 -> 1."""
    cid = case_id.lower()
    return 1 if cid.endswith(("_p2", "_person2")) else 0


def orig_paths(scene_xml: Path) -> tuple[Path, Path]:
    """Reference trajectory + shared 3cm mask for a reused orig (v1) row."""
    scene_dir = scene_xml.parent
    traj = scene_dir / "0" / "trajectory_kinematic.npz"
    mask = CONTACT_MASK_ROOT / scene_dir.name / "raw_contact_mask_3cm.npz"
    return traj, mask


def compute_row(row: dict[str, str], manifest: dict[tuple[str, str], dict[str, str]]) -> dict[str, Any]:
    """Attach recomputed body_z + release applicability, then the 12 gates."""
    out = dict(row)
    scene_xml = G.repo_path(row["scene_xml"])
    qpos_path = Path(row["qpos_path"])
    group = row["group"]

    if group == "aug":
        m = manifest.get((row["case_id"], row["aug_variant"]))
        trajectory = G.repo_path(m["trajectory"]) if m else None
        contact_mask = G.repo_path(m["contact_mask"]) if m else None
    else:  # orig: derive from the v1 scene dir
        traj, mask = orig_paths(scene_xml)
        trajectory, contact_mask = traj, mask

    note = ""
    # body_z p95 (recompute via CPU FK — not stored in the fullscale tsv)
    try:
        if trajectory is None or not trajectory.is_file():
            raise FileNotFoundError(f"missing reference trajectory: {trajectory}")
        out["body_z_err_p95_m"] = G.body_z_p95(qpos_path, scene_xml, trajectory)
    except Exception as exc:  # noqa: BLE001
        out["body_z_err_p95_m"] = math.nan
        note = f"body_z:{type(exc).__name__}"

    # release applicability from the shared 3cm mask
    frame_count = int(finite(row.get("qpos_frames"), 0))
    applicable = False
    if contact_mask is not None and contact_mask.is_file() and frame_count:
        try:
            applicable = G.release_applicable(contact_mask, actor_index(row["case_id"]), frame_count)
        except Exception as exc:  # noqa: BLE001
            note = (note + f";release:{type(exc).__name__}").strip(";")
    out["release_gate_applicable"] = applicable

    # coerce the gate source fields to the numeric types G.apply_gates expects
    item = {
        "fall_flag": str(row.get("fall_flag", "")).lower() in ("true", "1"),
        "body_z_err_p95_m": out["body_z_err_p95_m"],
        "hand_object_physics_contact_in_mask_frac": finite(row.get("hand_object_physics_contact_in_mask_frac")),
        "hand_object_release_false_contact_3mm_frac": finite(row.get("hand_object_release_false_contact_3mm_frac")),
        "hand_object_physics_penetration_3mm_frame_frac": finite(row.get("hand_object_physics_penetration_3mm_frame_frac")),
        "leg_penetration_frac": finite(row.get("leg_penetration_frac")),
        "track_root_pos_err_cm_mean": finite(row.get("track_root_pos_err_cm_mean")),
        "track_root_ori_err_deg_mean": finite(row.get("track_root_ori_err_deg_mean")),
        "track_eef_pos_err_cm_mean": finite(row.get("track_eef_pos_err_cm_mean")),
        "track_eef_ori_err_deg_mean": finite(row.get("track_eef_ori_err_deg_mean")),
        "track_obj_pos_err_cm_mean": finite(row.get("track_obj_pos_err_cm_mean")),
        "track_obj_ori_err_deg_mean": finite(row.get("track_obj_ori_err_deg_mean")),
        "release_gate_applicable": applicable,
    }
    G.apply_gates(item)  # writes fall_gate_pass ... object_ori_gate_pass + numeric_release_pass
    for _lbl, key, *_ in GATE12:
        out[f"{key}_gate_pass"] = bool(item[f"{key}_gate_pass"])
    out["gate12_all_pass"] = bool(item["numeric_release_pass"])
    out["gate12_failure_modes"] = item["numeric_failure_modes"]

    # NEW gates
    new_pass = []
    for label, field, comp, thr in NEW_GATES:
        val = finite(row.get(field))
        if comp == ">=":
            ok = math.isfinite(val) and val >= thr
        elif comp == "<=":
            ok = math.isfinite(val) and val <= thr
        else:  # "<"
            ok = math.isfinite(val) and val < thr
        out[label] = ok
        new_pass.append(ok)
    out["new_gates_all_pass"] = all(new_pass)
    out["_note"] = note
    return out


def _cell(ws, r: int, c: int, value: Any, fill: PatternFill | None = None) -> None:
    cell = ws.cell(row=r, column=c, value=value)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    if fill is not None:
        cell.fill = fill


def bool_fill(value: Any) -> PatternFill:
    if value is True:
        return GREEN
    if value is False:
        return RED
    return GREY


def _num_cell(ws, r: int, c: int, value: Any) -> None:
    v = finite(value)
    _cell(ws, r, c, round(v, 4) if math.isfinite(v) else "")


def write_detail(ws, rows: list[dict[str, Any]]) -> None:
    """One merged sheet: each gate shows its numeric value AND its PASS/FAIL."""
    # column plan: id block -> 12 gates (value|pass each) -> gate12 summary ->
    # ungated metrics -> 4 new gates (value|pass each) -> new_all + note
    header: list[str] = ["object_key", "case_id", "group", "aug_variant"]
    for _lbl, _key, _fld, _op, val_hdr in GATE12:
        header += [val_hdr, "→pass"]
    header += ["gate12_all_pass", "gate12_failure_modes",
               "qpos_jerk_l2_p95", "foot_slip_max_m"]
    new_val_hdr = {
        "NG_contact3mm_in_mask>=0.40": "contact3mm_in_mask",
        "NG_hand_pen_3mm<=0.32": "hand_pen3mm",
        "NG_ankle_jerk_p95<1000": "ankle_jerk_p95",
        "NG_obj_speed_max<3": "obj_speed_max",
    }
    for label, *_ in NEW_GATES:
        header += [new_val_hdr[label], label]
    header += ["new_gates_all_pass", "note"]

    for c, name in enumerate(header, start=1):
        _cell(ws, 1, c, name)
        ws.cell(row=1, column=c).fill = NAVY
        ws.cell(row=1, column=c).font = HEAD_FONT
        ws.cell(row=1, column=c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "E2"

    for r, row in enumerate(rows, start=2):
        c = 1
        for key in ("object_key", "case_id", "group", "aug_variant"):
            _cell(ws, r, c, row.get(key, ""))
            c += 1
        for _lbl, key, field, _op, _vh in GATE12:
            if field == "fall_flag":
                _cell(ws, r, c, str(row.get("fall_flag", "")))
            else:
                _num_cell(ws, r, c, row.get(field))
            c += 1
            v = row[f"{key}_gate_pass"]
            _cell(ws, r, c, "PASS" if v else "FAIL", bool_fill(v))
            c += 1
        _cell(ws, r, c, "PASS" if row["gate12_all_pass"] else "FAIL", bool_fill(row["gate12_all_pass"]))
        c += 1
        _cell(ws, r, c, row["gate12_failure_modes"])
        c += 1
        _num_cell(ws, r, c, row.get("qpos_jerk_l2_p95")); c += 1
        _num_cell(ws, r, c, row.get("foot_slip_max_m")); c += 1
        for label, field, _comp, _thr in NEW_GATES:
            _num_cell(ws, r, c, row.get(field))
            c += 1
            v = row[label]
            _cell(ws, r, c, "PASS" if v else "FAIL", bool_fill(v))
            c += 1
        _cell(ws, r, c, "PASS" if row["new_gates_all_pass"] else "FAIL", bool_fill(row["new_gates_all_pass"]))
        c += 1
        _cell(ws, r, c, row.get("_note", ""))

    widths = [10, 34, 8, 12] + [14, 7] * len(GATE12) + [12, 22, 15, 14] + [16, 20] * len(NEW_GATES) + [16, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _rate(rows: list[dict[str, Any]], key: str) -> tuple[int, int]:
    n = sum(1 for r in rows if r.get(key) is True)
    return n, len(rows)


def write_summary(ws, rows: list[dict[str, Any]]) -> None:
    aug = [r for r in rows if r["group"] == "aug"]
    orig = [r for r in rows if r["group"] == "orig"]
    objects = sorted({r["object_key"] for r in rows})

    gate_keys = (
        [("gate12_all", "gate12_all_pass")]
        + [(lbl, f"{key}_gate_pass") for lbl, key, *_ in GATE12]
        + [(lbl, lbl) for lbl, *_ in NEW_GATES]
        + [("new_gates_all", "new_gates_all_pass")]
    )

    def block(title: str, subset_aug, subset_orig, r0: int) -> int:
        _cell(ws, r0, 1, title)
        ws.cell(row=r0, column=1).font = Font(bold=True)
        _cell(ws, r0 + 1, 1, "gate", NAVY)
        _cell(ws, r0 + 1, 2, "aug pass", NAVY)
        _cell(ws, r0 + 1, 3, "aug rate", NAVY)
        _cell(ws, r0 + 1, 4, "orig pass", NAVY)
        _cell(ws, r0 + 1, 5, "orig rate", NAVY)
        for col in range(1, 6):
            ws.cell(row=r0 + 1, column=col).font = HEAD_FONT
        r = r0 + 2
        for lbl, key in gate_keys:
            an, at = _rate(subset_aug, key)
            on, ot = _rate(subset_orig, key)
            _cell(ws, r, 1, lbl)
            _cell(ws, r, 2, f"{an}/{at}")
            _cell(ws, r, 3, round(an / at, 3) if at else "")
            _cell(ws, r, 4, f"{on}/{ot}")
            _cell(ws, r, 5, round(on / ot, 3) if ot else "")
            r += 1
        return r + 1

    row_ptr = block("OVERALL (aug 249 vs orig 83)", aug, orig, 1)
    for obj in objects:
        row_ptr = block(
            f"{obj}",
            [r for r in aug if r["object_key"] == obj],
            [r for r in orig if r["object_key"] == obj],
            row_ptr,
        )
    ws.column_dimensions["A"].width = 26
    for col in "BCDE":
        ws.column_dimensions[col].width = 12


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()

    case_rows = read_tsv(CASE_METRICS)
    man_rows = read_tsv(MANIFEST)
    manifest = {(m["case_id"], m["aug_variant"]): m for m in man_rows}

    order = {"orig": 0, "trans0": 1, "trans1": 2, "trans2": 3}
    case_rows.sort(key=lambda r: (r["object_key"], r["case_id"].replace("_person", "_p"), order.get(r["aug_variant"], 9)))

    computed: list[dict[str, Any]] = []
    for i, row in enumerate(case_rows, start=1):
        computed.append(compute_row(row, manifest))
        if i % 40 == 0 or i == len(case_rows):
            print(f"[compute] {i}/{len(case_rows)}", file=sys.stderr)

    wb = Workbook()
    write_detail(wb.active, computed)
    wb.active.title = "detail"
    write_summary(wb.create_sheet("summary"), computed)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)

    n_note = sum(1 for r in computed if r.get("_note"))
    print(f"[done] wrote {args.out} ({len(computed)} rows, {n_note} with recompute notes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
