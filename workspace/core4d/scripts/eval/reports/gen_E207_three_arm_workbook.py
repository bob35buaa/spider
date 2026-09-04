#!/usr/bin/env python3
"""E207 three-arm comparison workbook (xlsx): PRG vs PRG+G1 vs PRG+G1A2.

The three arms form an incremental ladder on the same 9 bucket cases, same
reference trajectory, same 3cm mask, same CEM budget (1024x32 seed0):

    PRG        E178   leg-PRG on, A0 hand-gate, no gravcomp   (baseline)
    PRG+G1     E207   + object gravcomp                       (this experiment)
    PRG+G1A2   E205   + object gravcomp + A2 hand-gate

so each step is a single variable: PRG->PRG+G1 isolates gravcomp, and
PRG+G1->PRG+G1A2 isolates the A2 hand-gate.

Joins two already-computed sources (no re-scoring):
  * four_arm_rollout.tsv          E201 14-gate metrics (eval_E207_g1only.py)
  * e207_object_z_diff_by_case.tsv object z bias/MAE (gen_E207_object_z_diff.py)

Sheets: README / Summary / Per-Case / Paired

Usage:
    .venv/bin/python workspace/core4d/scripts/eval/reports/gen_E207_three_arm_workbook.py
"""

from __future__ import annotations

import argparse
import csv
import math
import statistics
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[5]
sys.path.insert(0, str(REPO / "workspace/core4d/scripts/experiments/E201"))
sys.path.insert(0, str(REPO / "workspace/core4d/scripts/experiments/E207"))
import funnel_config as FC  # noqa: E402
import e207_common as C207  # noqa: E402

EVAL_DIR = C207.RESULTS / "s6_downstream/eval/four_arm"
ROLLOUT_TSV = EVAL_DIR / "four_arm_rollout.tsv"
ZDIFF_TSV = EVAL_DIR / "e207_object_z_diff_by_case.tsv"
OUT_XLSX = EVAL_DIR / "E207_three_arm_comparison.xlsx"

#: display order = the incremental ladder. rollout arm label -> display label.
ARMS = [("PRG", "PRG"), ("G1only", "PRG+G1"), ("G1A2", "PRG+G1A2")]
ARM_DESC = {
    "PRG": "E178 · leg-PRG on, A0 hand-gate, gravcomp 0 —— 基线",
    "PRG+G1": "E207 · 基线 + object gravcomp=1（相对 PRG 唯一变量）",
    "PRG+G1A2": "E205 · 基线 + object gravcomp=1 + A2 hand-gate（相对 PRG+G1 唯一变量）",
}
#: z metrics come from the z-diff TSV, keyed by the rollout arm label.
Z_ARM = {"PRG": "PRG", "G1only": "G1only", "G1A2": "G1A2"}

NAVY, HEAD2, GREEN, RED, AMBER, GRAY = "1F3864", "2E5496", "C6EFCE", "F8CBAD", "FFEB9C", "F2F2F2"

#: (field, lower_is_better, label). z_* are joined from the z-diff TSV.
METRICS: list[tuple[str, bool, str]] = [
    ("z_bias_cm", None, "物体 z bias (cm)"),          # None = |value| lower better
    ("z_mae_cm", True, "物体 z MAE (cm)"),
    ("track_obj_pos_err_cm_mean", True, "obj pos (cm)"),
    ("track_obj_ori_err_deg_mean", True, "obj ori (°)"),
    ("hand_object_physics_contact_in_mask_frac", False, "承重接触 in-mask"),
    ("hand_object_release_false_contact_3mm_frac", True, "release 假接触"),
    ("hand_object_physics_penetration_3mm_frame_frac", True, "手穿透 3mm"),
    ("track_eef_pos_err_cm_mean", True, "eef pos (cm)"),
    ("track_eef_ori_err_deg_mean", True, "eef ori (°)"),
    ("track_root_pos_err_cm_mean", True, "root pos (cm)"),
    ("track_root_ori_err_deg_mean", True, "root ori (°)"),
    ("leg_penetration_frac", True, "腿穿透"),
    ("body_z_err_p95_m", True, "body z p95 (m)"),
    ("ankle_jerk_p95", True, "ankle jerk p95"),
]
GATE_ORDER = [("fall", "fall_flag"), ("body_z", "body_z_err_p95_m"),
              ("ankle_jerk", "ankle_jerk_p95"), ("obj_speed", "obj_speed_max")] + \
    [(g[0], g[1]) for g in FC.BANDED_GATES]


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream, delimiter="\t"))


def num(value: Any) -> float:
    try:
        out = float(value)
        return out if math.isfinite(out) else math.nan
    except (TypeError, ValueError):
        return math.nan


def head(sheet, row: int, values: list[str], fill: str = NAVY) -> None:
    for col, text in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=col, value=text)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 32


def widths(sheet, spec: dict[str, int]) -> None:
    for col, width in spec.items():
        sheet.column_dimensions[col].width = width


def build(rollout: list[dict[str, str]], zrows: list[dict[str, str]]) -> Workbook:
    by = {(r["arm"], r["case_id"]): r for r in rollout}
    zby = {(r["arm"], r["case_id"]): r for r in zrows}
    cases = [c for c in C207.CASES if ("PRG", c) in by]
    obj_of = {r["case_id"]: r["object_key"] for r in rollout}

    def value(arm: str, case_id: str, field: str) -> float:
        if field.startswith("z_"):
            return num(zby.get((Z_ARM[arm], case_id), {}).get(field))
        return num(by.get((arm, case_id), {}).get(field))

    def mean(arm: str, field: str) -> float:
        vals = [value(arm, c, field) for c in cases]
        vals = [v for v in vals if math.isfinite(v)]
        return statistics.fmean(vals) if vals else math.nan

    workbook = Workbook()

    # ---------- README ----------
    sheet = workbook.active
    sheet.title = "README"
    widths(sheet, {"A": 26, "B": 104})
    rows = [
        ("实验", "E207 (R293) · bucket G1-only · plan237 / log296"),
        ("对比", "PRG → PRG+G1 → PRG+G1A2，三臂为递进阶梯，每一步只改一个变量"),
        ("PRG", ARM_DESC["PRG"]),
        ("PRG+G1", ARM_DESC["PRG+G1"]),
        ("PRG+G1A2", ARM_DESC["PRG+G1A2"]),
        ("共同不变量", "同 9 个 bucket case、同 omnirt_v1 ref_fk 参考轨迹、同 3cm 掩码、"
                       "同 contact-aligned 五段代理、CEM 1024×32 seed0"),
        ("case 范围", f"{len(cases)} 例（bucket003=3 / bucket007=6，**无 bucket004**，结论不外推到它）"),
        ("主判定", "E201 14-gate 漏斗（4 硬门 + 10 带门 wide/narrow），与 E204/E205 同尺子"),
        ("z 指标", "z_diff(t)=z_sim(t)−z_ref(t) cm，正=高于参考；bias 保留符号，MAE 取绝对值"),
        ("单变量证据", "Hydra compose 逐 key diff：PRG→PRG+G1 差 1 个 key(scene_name)；"
                       "PRG+G1→PRG+G1A2 差 2 个 key(hand-gate)"),
        ("总判定", "PARTIAL SUCCESS(偏强)：C0/C1/C2/C3/C5/C7 过；C4 仅输「≥7/9 改善」子句(6/9)，"
                   "幅度子句超额(|bias| 1.379→0.566，门 0.8)"),
        ("关键机制", "gravcomp 是常量偏移 +1.945±0.711 cm，与抬升幅度无关(r=+0.11)、"
                     "与原始 bias 强负相关(r=−0.85)；故对原本几乎不下沉的 case 必然过冲"),
        ("A2 的作用", "PRG+G1 与 PRG+G1A2 的 z bias 仅差 0.169 cm → 物体跟踪改善全部来自 gravcomp"),
        ("口径提醒", "gravcomp 使物体在优化中失重，下游 RL/真机在真实重力下工作；"
                     "该建模落差已在 plan237 声明接受"),
        ("单元格配色", "绿=该行三臂中最优，红=最差（z bias 按绝对值比较）"),
        ("数据来源", "four_arm_rollout.tsv + e207_object_z_diff_by_case.tsv（均为已算结果，无重新评分）"),
    ]
    for index, (key, val) in enumerate(rows, start=1):
        cell = sheet.cell(row=index, column=1, value=key)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=GRAY)
        body = sheet.cell(row=index, column=2, value=val)
        body.alignment = Alignment(wrap_text=True, vertical="top")

    # ---------- Summary ----------
    sheet = workbook.create_sheet("Summary")
    widths(sheet, {"A": 34, "B": 14, "C": 14, "D": 14, "E": 15, "F": 15})
    sheet["A1"] = "14-gate 通过数 (n=%d)" % len(cases)
    sheet["A1"].font = Font(bold=True, size=12)
    head(sheet, 2, ["", "PRG", "PRG+G1", "PRG+G1A2", "Δ(G1−PRG)", "Δ(G1A2−G1)"])
    row = 3
    for label, key in (("hard", "hard_pass"), ("wide_all", "wide_pass"), ("narrow_all", "narrow_pass")):
        counts = [sum(str(by[(a, c)][key]).strip().lower() == "true" for c in cases) for a, _ in ARMS]
        sheet.cell(row=row, column=1, value=label).font = Font(bold=(label == "narrow_all"))
        for i, n in enumerate(counts):
            sheet.cell(row=row, column=2 + i, value=f"{n}/{len(cases)}").alignment = Alignment(horizontal="center")
        sheet.cell(row=row, column=5, value=counts[1] - counts[0]).alignment = Alignment(horizontal="center")
        sheet.cell(row=row, column=6, value=counts[2] - counts[1]).alignment = Alignment(horizontal="center")
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="逐门 narrow 通过数").font = Font(bold=True, size=12)
    row += 1
    head(sheet, row, ["gate", "PRG", "PRG+G1", "PRG+G1A2", "Δ(G1−PRG)", "Δ(G1A2−G1)"], HEAD2)
    row += 1
    hard = {h[1]: h for h in FC.HARD_GATES}
    banded = {g[1]: g for g in FC.BANDED_GATES}
    for gname, gfield in GATE_ORDER:
        counts = []
        for arm, _ in ARMS:
            group = [by[(arm, c)] for c in cases]
            if gname == "fall":
                ok = sum(str(r["fall_flag"]).strip().lower() not in ("true", "1") for r in group)
            elif gfield in hard:
                _n, _f, op, thr = hard[gfield]
                ok = sum(FC.passes(op, num(r[gfield]), thr) for r in group)
            else:
                bg = banded[gfield]
                ok = sum(FC.passes(bg[2], num(r[gfield]), bg[3]) for r in group)
            counts.append(ok)
        sheet.cell(row=row, column=1, value=gname)
        for i, n in enumerate(counts):
            cell = sheet.cell(row=row, column=2 + i, value=n)
            cell.alignment = Alignment(horizontal="center")
            if n < len(cases):
                cell.fill = PatternFill("solid", fgColor=AMBER if n >= len(cases) - 2 else RED)
        for off, d in ((3, counts[1] - counts[0]), (4, counts[2] - counts[1])):
            cell = sheet.cell(row=row, column=2 + off, value=d)
            cell.alignment = Alignment(horizontal="center")
            if d:
                cell.fill = PatternFill("solid", fgColor=GREEN if d > 0 else RED)
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="指标均值（宏平均）").font = Font(bold=True, size=12)
    row += 1
    head(sheet, row, ["指标", "PRG", "PRG+G1", "PRG+G1A2", "Δ(G1−PRG)", "Δ(G1A2−G1)"], HEAD2)
    row += 1
    for field, lower_better, label in METRICS:
        means = [mean(arm, field) for arm, _ in ARMS]
        sheet.cell(row=row, column=1, value=label)
        key = (lambda v: abs(v)) if lower_better is None else (lambda v: v)
        finite = [m for m in means if math.isfinite(m)]
        best = min(finite, key=key) if lower_better is not False else max(finite, key=key)
        worst = max(finite, key=key) if lower_better is not False else min(finite, key=key)
        for i, m in enumerate(means):
            cell = sheet.cell(row=row, column=2 + i, value=round(m, 4))
            cell.alignment = Alignment(horizontal="center")
            if math.isfinite(m) and len(set(finite)) > 1:
                if m == best:
                    cell.fill = PatternFill("solid", fgColor=GREEN)
                elif m == worst:
                    cell.fill = PatternFill("solid", fgColor=RED)
        for off, d in ((3, means[1] - means[0]), (4, means[2] - means[1])):
            cell = sheet.cell(row=row, column=2 + off, value=round(d, 4))
            cell.alignment = Alignment(horizontal="center")
        row += 1
    sheet.freeze_panes = "B3"

    # ---------- Per-Case ----------
    sheet = workbook.create_sheet("Per-Case")
    cols = ["case_id", "object", "arm", "narrow", "layer", "narrow_failed"] + [m[2] for m in METRICS]
    head(sheet, 1, cols)
    widths(sheet, {"A": 30, "B": 11, "C": 11, "D": 8, "E": 11, "F": 22})
    for i in range(len(METRICS)):
        sheet.column_dimensions[get_column_letter(7 + i)].width = 14
    row = 2
    thin = Side(style="thin", color="BFBFBF")
    for case_id in cases:
        block_start = row
        for arm, disp in ARMS:
            record = by[(arm, case_id)]
            npass = str(record["narrow_pass"]).strip().lower() == "true"
            values = [case_id, obj_of[case_id], disp, "PASS" if npass else "FAIL",
                      record["layer"], record["narrow_failed"]]
            for i, v in enumerate(values):
                sheet.cell(row=row, column=1 + i, value=v)
            sheet.cell(row=row, column=4).fill = PatternFill("solid", fgColor=GREEN if npass else RED)
            for i, (field, _lb, _label) in enumerate(METRICS):
                cell = sheet.cell(row=row, column=7 + i, value=round(value(arm, case_id, field), 4))
                cell.alignment = Alignment(horizontal="center")
            row += 1
        # highlight the best/worst arm per metric within this case block
        for i, (field, lower_better, _label) in enumerate(METRICS):
            col = 7 + i
            vals = [(r, sheet.cell(row=r, column=col).value) for r in range(block_start, row)]
            vals = [(r, v) for r, v in vals if isinstance(v, (int, float)) and math.isfinite(v)]
            if len(vals) < 2 or len({v for _r, v in vals}) == 1:
                continue
            key = (lambda t: abs(t[1])) if lower_better is None else (lambda t: t[1])
            best = min(vals, key=key) if lower_better is not False else max(vals, key=key)
            worst = max(vals, key=key) if lower_better is not False else min(vals, key=key)
            sheet.cell(row=best[0], column=col).fill = PatternFill("solid", fgColor=GREEN)
            sheet.cell(row=worst[0], column=col).fill = PatternFill("solid", fgColor=RED)
        for c in range(1, len(cols) + 1):
            sheet.cell(row=row - 1, column=c).border = Border(bottom=thin)
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{row - 1}"

    # ---------- Paired ----------
    sheet = workbook.create_sheet("Paired")
    sheet["A1"] = "逐 case paired delta —— 两条单变量对比"
    sheet["A1"].font = Font(bold=True, size=12)
    sheet["A2"] = "PRG → PRG+G1 隔离 gravcomp；PRG+G1 → PRG+G1A2 隔离 A2 hand-gate。负值=改善（除承重接触外）。"
    sheet["A2"].alignment = Alignment(wrap_text=True)
    widths(sheet, {"A": 30, "B": 11})
    row = 4
    for a, b, title in (("PRG", "G1only", "PRG → PRG+G1（隔离 gravcomp）"),
                        ("G1only", "G1A2", "PRG+G1 → PRG+G1A2（隔离 A2）")):
        sheet.cell(row=row, column=1, value=title).font = Font(bold=True, size=11)
        row += 1
        head(sheet, row, ["case_id", "object"] + [m[2] for m in METRICS], HEAD2)
        for i in range(len(METRICS)):
            sheet.column_dimensions[get_column_letter(3 + i)].width = 14
        row += 1
        first = row
        for case_id in cases:
            sheet.cell(row=row, column=1, value=case_id)
            sheet.cell(row=row, column=2, value=obj_of[case_id])
            for i, (field, lower_better, _label) in enumerate(METRICS):
                delta = value(b, case_id, field) - value(a, case_id, field)
                cell = sheet.cell(row=row, column=3 + i, value=round(delta, 4))
                cell.alignment = Alignment(horizontal="center")
                if math.isfinite(delta) and abs(delta) > 1e-12:
                    if lower_better is None:
                        good = abs(value(b, case_id, field)) < abs(value(a, case_id, field))
                    else:
                        good = (delta < 0) if lower_better else (delta > 0)
                    cell.fill = PatternFill("solid", fgColor=GREEN if good else RED)
            row += 1
        sheet.cell(row=row, column=1, value="宏平均").font = Font(bold=True)
        for i, (field, _lb, _label) in enumerate(METRICS):
            deltas = [value(b, c, field) - value(a, c, field) for c in cases]
            deltas = [d for d in deltas if math.isfinite(d)]
            cell = sheet.cell(row=row, column=3 + i, value=round(statistics.fmean(deltas), 4) if deltas else None)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=GRAY)
            cell.alignment = Alignment(horizontal="center")
        for i in range(len(METRICS)):
            col = get_column_letter(3 + i)
            sheet.conditional_formatting.add(
                f"{col}{first}:{col}{row - 1}",
                ColorScaleRule(start_type="min", start_color="FFFFFF",
                               end_type="max", end_color="FFFFFF"))
        row += 3
    return workbook


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=OUT_XLSX)
    args = parser.parse_args()
    for label, path in (("rollout", ROLLOUT_TSV), ("z-diff", ZDIFF_TSV)):
        if not path.is_file():
            raise SystemExit(f"missing {label}: {path}")
    workbook = build(read_tsv(ROLLOUT_TSV), read_tsv(ZDIFF_TSV))
    args.out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.out)
    print(f"[done] {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
