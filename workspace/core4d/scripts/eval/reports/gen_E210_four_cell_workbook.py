#!/usr/bin/env python3
"""E210: direction-aware contrasts + xlsx workbook over the 2x2 rollout table.

Reads `four_cell_rollout.tsv` (written by eval_E210_aug_g1only.py) and recomputes
everything the verdicts depend on. Kept separate from the runner because scoring
40 rollouts takes ~15 min and the arithmetic below needed two corrections that
would otherwise have cost a rescore each:

  1. DIRECTION. `hand_object_physics_contact_in_mask_frac` is better when it goes
     UP; every other metric here is an error and is better when it goes DOWN. A
     shared "n_worse = count(delta > 0)" silently reports contact backwards --
     which flips the sign of the headline finding, not just a label.
  2. FIELD. The funnel's `contact` gate reads
     `hand_object_physics_contact_in_mask_frac`, NOT the `..._3mm_...` variant
     that appears in the RL-export schema. Both are in the TSV; mixing them up
     compares a gate against a metric it does not gate on.

Usage:
    .venv/bin/python workspace/core4d/scripts/eval/reports/gen_E210_four_cell_workbook.py
    ... --rollout <tsv> --out <xlsx>
"""

from __future__ import annotations

import argparse
import json
import math
import statistics as st
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

REPO = Path(__file__).resolve().parents[5]
for _p in ("workspace/core4d/scripts",
           "workspace/core4d/scripts/experiments/E201",
           "workspace/core4d/scripts/experiments/E210"):
    sys.path.insert(0, str(REPO / _p))

import funnel_config as FC  # noqa: E402
import e210_common as C210  # noqa: E402

CELLS = ("A_orig_PRG", "B_orig_G1", "C_aug_PRG", "D_aug_G1")
CELL_LABEL = {"A_orig_PRG": "A orig x PRG", "B_orig_G1": "B orig x PRG+G1",
              "C_aug_PRG": "C aug x PRG", "D_aug_G1": "D aug x PRG+G1"}

CONTACT_GATE = "hand_object_physics_contact_in_mask_frac"   # what the funnel gates on
CONTACT_3MM = "hand_object_physics_contact_3mm_in_mask_frac"  # RL-export schema metric
EEF_ORI = "track_eef_ori_err_deg_mean"
HAND_PEN = "hand_object_physics_penetration_3mm_frame_frac"

#: True == lower is better. Anything absent is treated as an error metric.
LOWER_IS_BETTER = {CONTACT_GATE: False, CONTACT_3MM: False}
TRACKED = [EEF_ORI, "track_root_ori_err_deg_mean", HAND_PEN, CONTACT_GATE, CONTACT_3MM,
           "hand_object_release_false_contact_3mm_frac", "track_obj_pos_err_cm_mean",
           "track_obj_ori_err_deg_mean", "track_eef_pos_err_cm_mean", "leg_penetration_frac"]

BANDED = {g[1]: g for g in FC.BANDED_GATES}


def num(v: Any) -> float:
    try:
        f = float(v)
        return f if math.isfinite(f) else math.nan
    except (TypeError, ValueError):
        return math.nan


def truth(v: Any) -> bool:
    """write_tsv serialises booleans lowercase; `== "True"` silently counts zero."""
    return str(v).strip().lower() in ("true", "1")


def q(sorted_xs: list[float], p: float) -> float:
    return sorted_xs[min(len(sorted_xs) - 1, int(round(p * (len(sorted_xs) - 1))))]


def delta_stats(field: str, deltas: list[float]) -> dict[str, Any]:
    d = sorted(x for x in deltas if math.isfinite(x))
    if not d:
        return {"n": 0}
    lower_better = LOWER_IS_BETTER.get(field, True)
    worse = sum(x > 0 for x in d) if lower_better else sum(x < 0 for x in d)
    better = sum(x < 0 for x in d) if lower_better else sum(x > 0 for x in d)
    # "worst" = the most-degraded end, whichever direction that is
    worst = d[-1] if lower_better else d[0]
    return {"n": len(d), "mean": round(st.fmean(d), 4),
            "median": round(st.median(d), 4),
            "std": round(st.pstdev(d), 4) if len(d) > 1 else 0.0,
            "n_worse": worse, "n_better": better, "n_equal": len(d) - worse - better,
            "p10": round(q(d, 0.10), 4), "p90": round(q(d, 0.90), 4),
            "worst": round(worst, 4), "lower_is_better": lower_better}


def gate_rate(rows: list[dict], field: str, caliber: str = "narrow") -> float:
    if not rows:
        return math.nan
    g = BANDED[field]
    thr = g[3] if caliber == "narrow" else g[4]
    return sum(FC.passes(g[2], num(r[field]), thr) for r in rows) / len(rows)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--rollout", type=Path,
                    default=C210.RESULTS / "s6_downstream/eval/four_cell/four_cell_rollout.tsv")
    ap.add_argument("--out", type=Path, default=None)
    args = ap.parse_args()

    rows = C210.read_tsv(args.rollout)
    if not rows:
        raise SystemExit(f"empty rollout table: {args.rollout}")
    for r in rows:
        for k in ("hard_pass", "wide_pass", "narrow_pass"):
            r[k] = truth(r[k])
    by_cell = {c: [r for r in rows if r["cell"] == c] for c in CELLS}
    idx = {(r["cell"], r["case_id"], r["aug_variant"]): r for r in rows}
    cases = sorted({r["case_id"] for r in rows})
    variants = list(C210.TRANS_VARIANTS)

    def pair_up(a_cell: str, b_cell: str) -> list[tuple[dict, dict]]:
        out = []
        a_aug, b_aug = a_cell[0] in "CD", b_cell[0] in "CD"
        for case_id in cases:
            for v in (variants if (a_aug or b_aug) else ["orig"]):
                ra = idx.get((a_cell, case_id, v if a_aug else "orig"))
                rb = idx.get((b_cell, case_id, v if b_aug else "orig"))
                if ra and rb:
                    out.append((ra, rb))
        return out

    contrasts: dict[str, Any] = {}
    for a, b, label in (
        ("C_aug_PRG", "D_aug_G1", "C->D  gravcomp on aug (STRICT single-variable)"),
        ("A_orig_PRG", "B_orig_G1", "A->B  gravcomp on orig (E207's effect, recomputed)"),
        ("B_orig_G1", "D_aug_G1", "B->D  delivery (CONFOUNDED: aug + v1->v2 retarget)"),
        ("A_orig_PRG", "C_aug_PRG", "A->C  aug on PRG (CONFOUNDED: aug + v1->v2 retarget)"),
    ):
        pairs = pair_up(a, b)
        if not pairs:
            continue
        entry: dict[str, Any] = {
            "n": len(pairs),
            "narrow_gain": sum(rb["narrow_pass"] and not ra["narrow_pass"] for ra, rb in pairs),
            "narrow_loss": sum(ra["narrow_pass"] and not rb["narrow_pass"] for ra, rb in pairs),
            "metrics": {f: delta_stats(f, [num(rb[f]) - num(ra[f]) for ra, rb in pairs])
                        for f in TRACKED},
        }
        entry["gate_flips"] = Counter()
        for ra, rb in pairs:
            fa = set(filter(None, str(ra["narrow_failed"]).split(",")))
            fb = set(filter(None, str(rb["narrow_failed"]).split(",")))
            for g in fb - fa:
                entry["gate_flips"][f"+{g}"] += 1
            for g in fa - fb:
                entry["gate_flips"][f"-{g}"] += 1
        entry["gate_flips"] = dict(entry["gate_flips"])
        contrasts[label] = entry

    cd = contrasts["C->D  gravcomp on aug (STRICT single-variable)"]
    ab = contrasts["A->B  gravcomp on orig (E207's effect, recomputed)"]

    verdicts: dict[str, Any] = {}
    d_rate = st.fmean([float(r["narrow_pass"]) for r in by_cell["D_aug_G1"]])
    b_rate = st.fmean([float(r["narrow_pass"]) for r in by_cell["B_orig_G1"]])
    verdicts["C3_delivery"] = {
        "gate": "narrow rate(D) >= rate(B) - 0.15  AND  fall == 0",
        "rate_B_orig_G1": round(b_rate, 3), "rate_D_aug_G1": round(d_rate, 3),
        "gap": round(b_rate - d_rate, 3),
        "fall_count_D": sum(truth(r["fall_flag"]) for r in by_cell["D_aug_G1"]),
        "pass": bool(d_rate >= b_rate - 0.15
                     and not any(truth(r["fall_flag"]) for r in by_cell["D_aug_G1"])),
    }
    ex = cd["metrics"][EEF_ORI]["mean"] - ab["metrics"][EEF_ORI]["mean"]
    verdicts["C5a_phantom_support"] = {
        "gate": "eef_ori excess on aug <= +1.0 deg",
        "eef_ori_delta_C->D": cd["metrics"][EEF_ORI]["mean"],
        "eef_ori_delta_A->B": ab["metrics"][EEF_ORI]["mean"],
        "excess_on_aug": round(ex, 4), "pass": bool(ex <= 1.0),
        "eef_ori_narrow_C": round(gate_rate(by_cell["C_aug_PRG"], EEF_ORI), 3),
        "eef_ori_narrow_D": round(gate_rate(by_cell["D_aug_G1"], EEF_ORI), 3),
        "contact_gate_delta_C->D": cd["metrics"][CONTACT_GATE],
        "contact_gate_narrow_C": round(gate_rate(by_cell["C_aug_PRG"], CONTACT_GATE), 3),
        "contact_gate_narrow_D": round(gate_rate(by_cell["D_aug_G1"], CONTACT_GATE), 3),
    }
    verdicts["C5b_hand_penetration"] = {
        "gate": "hand_pen narrow rate(D) >= rate(B) - 0.20",
        "rate_B_orig_G1": round(gate_rate(by_cell["B_orig_G1"], HAND_PEN), 3),
        "rate_C_aug_PRG": round(gate_rate(by_cell["C_aug_PRG"], HAND_PEN), 3),
        "rate_D_aug_G1": round(gate_rate(by_cell["D_aug_G1"], HAND_PEN), 3),
        "delta_C->D": cd["metrics"][HAND_PEN], "delta_B->D": contrasts[
            "B->D  delivery (CONFOUNDED: aug + v1->v2 retarget)"]["metrics"][HAND_PEN],
        "pass": bool(gate_rate(by_cell["D_aug_G1"], HAND_PEN)
                     >= gate_rate(by_cell["B_orig_G1"], HAND_PEN) - 0.20),
    }
    d_rows = by_cell["D_aug_G1"]
    bands = Counter(r["offset_band"] for r in d_rows)
    offs = [num(r["approach_offset_m"]) for r in d_rows]
    verdicts["C6_effective_offset"] = {
        "bands": dict(bands), "below_floor_0.05m": sum(o < 0.05 for o in offs),
        "min": round(min(offs), 4), "median": round(st.median(offs), 4),
        "narrow_rate_by_band": {
            b: round(st.fmean([float(r["narrow_pass"]) for r in d_rows
                               if r["offset_band"] == b]), 3) for b in bands},
        "contact_gate_by_band": {
            b: round(gate_rate([r for r in d_rows if r["offset_band"] == b], CONTACT_GATE), 3)
            for b in bands},
    }
    # interaction: is gravcomp's cost different on aug than on orig?
    verdicts["C4_interaction"] = {
        f: {"delta_on_aug(C->D)": cd["metrics"][f]["mean"],
            "delta_on_orig(A->B)": ab["metrics"][f]["mean"],
            "excess_on_aug": round(cd["metrics"][f]["mean"] - ab["metrics"][f]["mean"], 4)}
        for f in (EEF_ORI, HAND_PEN, CONTACT_GATE, "track_obj_pos_err_cm_mean",
                  "track_obj_ori_err_deg_mean")
    }

    # per-case narrow table (which case carries the damage)
    per_case = []
    for case_id in cases:
        row: dict[str, Any] = {"case_id": case_id,
                               "offset_m": num(idx[("D_aug_G1", case_id, "trans0")]["approach_offset_m"])}
        for c in CELLS:
            sel = [r for r in by_cell[c] if r["case_id"] == case_id]
            row[c] = f"{sum(r['narrow_pass'] for r in sel)}/{len(sel)}"
            if c == "D_aug_G1":
                row["D_failed_gates"] = ",".join(sorted({
                    g for r in sel for g in filter(None, str(r["narrow_failed"]).split(","))}))
                row["D_contact_mean"] = round(st.fmean([num(r[CONTACT_GATE]) for r in sel]), 4)
                row["C_contact_mean"] = round(st.fmean(
                    [num(r[CONTACT_GATE]) for r in by_cell["C_aug_PRG"] if r["case_id"] == case_id]), 4)
        per_case.append(row)

    out_dir = args.rollout.parent
    C210.write_json(out_dir / "e210_four_cell_verdicts.json", {
        "created_at": C210.now(), "rollout": C210.rel(args.rollout),
        "n_rows": len(rows), "cases": cases,
        "cell_counts": {c: len(by_cell[c]) for c in CELLS},
        "contrasts": contrasts, "verdicts": verdicts, "per_case": per_case,
        "corrections_vs_runner": [
            "direction-aware: contact_in_mask is higher-is-better; the runner's "
            "n_worse=count(delta>0) reported it inverted",
            f"the funnel's contact gate reads {CONTACT_GATE}, not {CONTACT_3MM}",
        ],
    })

    # --- console -------------------------------------------------------------
    print("=== per-cell 14-gate funnel ===")
    print(f"  {'cell':18s} {'n':>3s} {'hard':>6s} {'wide':>6s} {'narrow':>12s}")
    for c in CELLS:
        g = by_cell[c]
        npass = sum(r["narrow_pass"] for r in g)
        print(f"  {CELL_LABEL[c]:18s} {len(g):3d} {sum(r['hard_pass'] for r in g):6d} "
              f"{sum(r['wide_pass'] for r in g):6d} {npass:6d} ({npass/len(g):3.0%})")

    print("\n=== contrasts (direction-aware; 'worse' means degraded, per metric) ===")
    for label, e in contrasts.items():
        print(f"\n  {label}   n={e['n']}   narrow +{e['narrow_gain']} / -{e['narrow_loss']}")
        if e["gate_flips"]:
            print(f"    gate flips: {e['gate_flips']}")
        for f in (EEF_ORI, CONTACT_GATE, HAND_PEN, "track_obj_pos_err_cm_mean",
                  "track_root_ori_err_deg_mean"):
            m = e["metrics"][f]
            if not m.get("n"):
                continue
            arrow = "v better" if m["lower_is_better"] else "^ better"
            print(f"    {f:46s} [{arrow}] mean={m['mean']:+8.4f} "
                  f"worse={m['n_worse']:2d} better={m['n_better']:2d} worst={m['worst']:+8.4f}")

    print("\n=== per-case narrow (n_pass / n) ===")
    print(f"  {'case':30s} {'off':>6s} {'A':>5s} {'B':>5s} {'C':>5s} {'D':>5s}  "
          f"{'C_cont':>7s} {'D_cont':>7s}  D_failed")
    for r in per_case:
        print(f"  {r['case_id']:30s} {r['offset_m']:6.3f} {r['A_orig_PRG']:>5s} "
              f"{r['B_orig_G1']:>5s} {r['C_aug_PRG']:>5s} {r['D_aug_G1']:>5s}  "
              f"{r['C_contact_mean']:7.3f} {r['D_contact_mean']:7.3f}  {r['D_failed_gates']}")

    print("\n=== verdicts ===")
    print(json.dumps(verdicts, ensure_ascii=False, indent=2))

    # --- xlsx ---------------------------------------------------------------
    out_xlsx = args.out or (out_dir / "E210_four_cell_comparison.xlsx")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        print(f"\n[skip xlsx] openpyxl unavailable; JSON+TSV are the authority")
        return 0
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    for i, line in enumerate([
        "E210 - bucket007 augmented variants x {PRG, PRG+G1}",
        "",
        "A orig x PRG      E178 (frozen e178_case_metrics.tsv)",
        "B orig x PRG+G1   E207",
        "C aug  x PRG      E202",
        "D aug  x PRG+G1   E210 (this experiment)",
        "",
        "C->D is the only strictly single-variable contrast (gravcomp only).",
        "A->C and B->D also change the retarget variant (omnirt_v1 -> v2), which",
        "E202 had to adopt because v1 is often IK-infeasible once the object moves.",
        "Never quote B->D as a causal effect of augmentation alone.",
        "",
        f"contact gate field = {CONTACT_GATE} (higher is better)",
        f"({CONTACT_3MM} is the RL-export schema metric and is NOT what the gate reads)",
    ], 1):
        ws.cell(row=i, column=1, value=line)
    ws.cell(row=1, column=1).font = Font(bold=True, size=13)

    ws = wb.create_sheet("PerCell")
    ws.append(["cell", "n", "hard", "wide", "narrow", "narrow_rate"])
    for c in CELLS:
        g = by_cell[c]
        ws.append([CELL_LABEL[c], len(g), sum(r["hard_pass"] for r in g),
                   sum(r["wide_pass"] for r in g), sum(r["narrow_pass"] for r in g),
                   round(sum(r["narrow_pass"] for r in g) / len(g), 3)])

    ws = wb.create_sheet("Contrasts")
    ws.append(["contrast", "metric", "lower_is_better", "n", "mean", "median", "std",
               "n_worse", "n_better", "p10", "p90", "worst"])
    for label, e in contrasts.items():
        for f, m in e["metrics"].items():
            if m.get("n"):
                ws.append([label, f, m["lower_is_better"], m["n"], m["mean"], m["median"],
                           m["std"], m["n_worse"], m["n_better"], m["p10"], m["p90"], m["worst"]])

    ws = wb.create_sheet("PerCase")
    hdr = ["case_id", "offset_m", *CELLS, "C_contact_mean", "D_contact_mean", "D_failed_gates"]
    ws.append(hdr)
    for r in per_case:
        ws.append([r.get(h, "") for h in hdr])

    ws = wb.create_sheet("PerRollout")
    cols = list(rows[0].keys())
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])

    ws = wb.create_sheet("Verdicts")
    ws.append(["claim", "field", "value"])
    for claim, body in verdicts.items():
        for k, v in body.items():
            ws.append([claim, k, json.dumps(v, ensure_ascii=False)
                       if isinstance(v, (dict, list)) else v])
    for sheet in wb:
        sheet.freeze_panes = "A2" if sheet.title != "README" else None
    wb.save(out_xlsx)
    print(f"\n[done] {C210.rel(out_xlsx)}")
    print(f"[done] {C210.rel(out_dir / 'e210_four_cell_verdicts.json')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
