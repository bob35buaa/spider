#!/usr/bin/env python3
"""Generate a box023 (E179 vs E173) + box021 (E170 vs E168) PRG vs no-PRG report.

Reuses already-computed metrics; does not re-run CEM or re-evaluate NPZ:

- box023: E179's own ``e179_vs_e173_paired.tsv`` / ``..._gate_matrix.tsv``
  (E173=PRG, E179=no-PRG) are read verbatim and relabeled prg/noprg.
- box021: E168 (no-PRG baseline) and E170 (frozen PRG) never had a twelve-gate
  scoring pass; their raw metric columns already exist in
  ``e168_case_metrics.tsv`` / ``e170_case_metrics.tsv``, so the same
  ``e189_common.apply_12gate_scoring`` adapter used by E179/E189 is applied
  here to both sides before pairing.

Output: one workbook with a Summary sheet plus one sheet per object, in the
same layout as ``E189_vs_PRG_boxes_report.xlsx`` (case_id/pass/failure-modes
+ 12 gates x [prg, noprg, delta], delta = prg - noprg, orange delta header).
"""

from __future__ import annotations

import math
import sys
from pathlib import Path
from typing import Any

SCRIPT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(SCRIPT_ROOT / "experiments/E189"))

import e189_common as C  # noqa: E402

REPO = C.REPO
E179_PAIRED = REPO / "workspace/core4d/results/E179/s6_downstream/eval/full/e179_vs_e173_paired.tsv"
E179_GATE_MATRIX = REPO / "workspace/core4d/results/E179/s6_downstream/eval/full/e179_vs_e173_gate_matrix.tsv"
E168_CASE_METRICS = REPO / "workspace/core4d/results/E168/s6_downstream/cem/eval/box021_all28_reviewed/e168_case_metrics.tsv"
E170_CASE_METRICS = REPO / "workspace/core4d/results/E170/s6_downstream/eval/full/e170_case_metrics.tsv"

OUTPUT_DIR = REPO / "workspace/core4d/results/E189/s6_downstream/eval/full"
OUTPUT_XLSX = OUTPUT_DIR / "box023_box021_PRG_vs_noPRG_report.xlsx"
OUTPUT_MD = OUTPUT_DIR / "box023_box021_PRG_vs_noPRG_report.md"

LABELS = {
    "fall": "fall",
    "body_z": "body-z",
    "contact": "contact",
    "release": "release",
    "hand_penetration": "hand penetration",
    "lower_body": "lower-body",
    "root_pos": "root pos",
    "root_ori": "root ori",
    "hand_pos": "EEF pos",
    "hand_ori": "EEF ori",
    "object_pos": "object pos",
    "object_ori": "object ori",
}
GATE_METRIC_MAP = {
    "body_z": "body_z_err_p95_m",
    "contact": "hand_object_physics_contact_in_mask_frac",
    "release": "hand_object_release_false_contact_3mm_frac",
    "hand_penetration": "hand_object_physics_penetration_3mm_frame_frac",
    "lower_body": "leg_penetration_frac",
    "root_pos": "track_root_pos_err_cm_mean",
    "root_ori": "track_root_ori_err_deg_mean",
    "hand_pos": "track_eef_pos_err_cm_mean",
    "hand_ori": "track_eef_ori_err_deg_mean",
    "object_pos": "track_obj_pos_err_cm_mean",
    "object_ori": "track_obj_ori_err_deg_mean",
}

OBJECTS = ("box023", "box021")
SOURCE_LABELS = {
    "box023": {"prg": "E173", "noprg": "E179"},
    "box021": {"prg": "E170", "noprg": "E168"},
}


def migration(prg_pass: bool, noprg_pass: bool) -> str:
    if prg_pass and noprg_pass:
        return "PASS_TO_PASS"
    if prg_pass:
        return "PASS_TO_FAIL"
    if noprg_pass:
        return "FAIL_TO_PASS"
    return "FAIL_TO_FAIL"


def load_box023() -> tuple[list[dict[str, Any]], dict[str, dict[str, bool]]]:
    paired_raw = C.read_tsv(E179_PAIRED)
    gate_raw = C.read_tsv(E179_GATE_MATRIX)
    rows: list[dict[str, Any]] = []
    for r in paired_raw:
        row: dict[str, Any] = {
            "object_key": "box023",
            "case_id": r["case_id"],
            "retarget_variant_id": r["retarget_variant_id"],
            "prg_pass": C.boolish(r["e173_12gate_pass"]),
            "noprg_pass": C.boolish(r["e179_12gate_pass"]),
            "prg_failure_modes": r["e173_failure_modes"] or "-",
            "noprg_failure_modes": r["e179_failure_modes"] or "-",
        }
        row["pass_migration"] = migration(row["prg_pass"], row["noprg_pass"])
        for gate, metric in GATE_METRIC_MAP.items():
            prg_v = C.finite(r.get(f"e173_{metric}"), math.nan)
            noprg_v = C.finite(r.get(f"e179_{metric}"), math.nan)
            row[f"{gate}_prg"] = prg_v
            row[f"{gate}_noprg"] = noprg_v
            row[f"{gate}_delta"] = (
                prg_v - noprg_v if math.isfinite(prg_v) and math.isfinite(noprg_v) else math.nan
            )
        rows.append(row)
    gates_by_case: dict[str, dict[str, bool]] = {}
    for g in gate_raw:
        gates_by_case.setdefault(g["case_id"], {})
        gates_by_case[g["case_id"]][f"{g['gate']}_prg"] = C.boolish(g["e173_pass"])
        gates_by_case[g["case_id"]][f"{g['gate']}_noprg"] = C.boolish(g["e179_pass"])
    for row in rows:
        cell = gates_by_case[row["case_id"]]
        row["fall_prg"] = int(cell["fall_prg"])
        row["fall_noprg"] = int(cell["fall_noprg"])
        row["fall_delta"] = row["fall_prg"] - row["fall_noprg"]
    return rows, gates_by_case


def score_box021_side(path: Path) -> dict[str, dict[str, Any]]:
    rows = C.read_tsv(path)
    by_case = {}
    for r in rows:
        if r.get("object_key") != "box021":
            continue
        item = dict(r)
        C.apply_12gate_scoring(item)
        by_case[r["case_id"]] = item
    if len(by_case) != 28:
        raise ValueError(f"{path}: expected 28 box021 rows, got {len(by_case)}")
    return by_case


def load_box021() -> tuple[list[dict[str, Any]], dict[str, dict[str, bool]]]:
    prg_by_case = score_box021_side(E170_CASE_METRICS)
    noprg_by_case = score_box021_side(E168_CASE_METRICS)
    if set(prg_by_case) != set(noprg_by_case):
        raise ValueError("box021 E170/E168 case_id sets differ")
    rows: list[dict[str, Any]] = []
    gates_by_case: dict[str, dict[str, bool]] = {}
    for case_id in sorted(prg_by_case):
        prg_item = prg_by_case[case_id]
        noprg_item = noprg_by_case[case_id]
        row: dict[str, Any] = {
            "object_key": "box021",
            "case_id": case_id,
            "retarget_variant_id": prg_item.get("retarget_variant_id", ""),
            "prg_pass": bool(prg_item["numeric_release_pass_12gate"]),
            "noprg_pass": bool(noprg_item["numeric_release_pass_12gate"]),
            "prg_failure_modes": prg_item["numeric_failure_modes"] or "-",
            "noprg_failure_modes": noprg_item["numeric_failure_modes"] or "-",
        }
        row["pass_migration"] = migration(row["prg_pass"], row["noprg_pass"])
        for gate, metric in GATE_METRIC_MAP.items():
            prg_v = C.finite(prg_item.get(metric), math.nan)
            noprg_v = C.finite(noprg_item.get(metric), math.nan)
            row[f"{gate}_prg"] = prg_v
            row[f"{gate}_noprg"] = noprg_v
            row[f"{gate}_delta"] = (
                prg_v - noprg_v if math.isfinite(prg_v) and math.isfinite(noprg_v) else math.nan
            )
        row["fall_prg"] = int(prg_item["fall_gate_pass"])
        row["fall_noprg"] = int(noprg_item["fall_gate_pass"])
        row["fall_delta"] = row["fall_prg"] - row["fall_noprg"]
        rows.append(row)
        gates_by_case[case_id] = {
            f"{gate}_prg": bool(prg_item[f"{gate}_gate_pass"])
            for gate in C.ALL_GATES
        } | {
            f"{gate}_noprg": bool(noprg_item[f"{gate}_gate_pass"])
            for gate in C.ALL_GATES
        }
    return rows, gates_by_case


def gate_pass_counts(rows: list[dict[str, Any]], gates_by_case: dict[str, dict[str, bool]]) -> dict[str, dict[str, int]]:
    counts: dict[str, dict[str, int]] = {
        gate: {"prg": 0, "noprg": 0} for gate in C.ALL_GATES
    }
    for row in rows:
        cell = gates_by_case[row["case_id"]]
        for gate in C.ALL_GATES:
            counts[gate]["prg"] += int(cell[f"{gate}_prg"])
            counts[gate]["noprg"] += int(cell[f"{gate}_noprg"])
    return counts


def write_xlsx(by_object: dict[str, list[dict[str, Any]]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    delta_header_fill = PatternFill("solid", fgColor="EA580C")
    fail_fill = PatternFill("solid", fgColor="FEE2E2")
    pass_fill = PatternFill("solid", fgColor="DCFCE7")

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(
        ["object_key", "n", "prg_source", "noprg_source", "prg_12gate_pass", "noprg_12gate_pass"]
    )
    for col_idx in range(1, 7):
        cell = ws_summary.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
    for obj in OBJECTS:
        rows = by_object[obj]
        n = len(rows)
        prg_pass = sum(row["prg_pass"] for row in rows)
        noprg_pass = sum(row["noprg_pass"] for row in rows)
        ws_summary.append(
            [
                obj,
                n,
                SOURCE_LABELS[obj]["prg"],
                SOURCE_LABELS[obj]["noprg"],
                f"{prg_pass}/{n}",
                f"{noprg_pass}/{n}",
            ]
        )
    for col_idx in range(1, 7):
        ws_summary.column_dimensions[get_column_letter(col_idx)].width = 20

    fixed_cols = ("case_id", "prg_pass", "noprg_pass", "prg_failure_modes", "noprg_failure_modes")
    for obj in OBJECTS:
        ws = wb.create_sheet(obj)
        rows = by_object[obj]
        fields = list(fixed_cols)
        for gate in C.ALL_GATES:
            fields.extend([f"{gate}_prg", f"{gate}_noprg", f"{gate}_delta"])

        ws.append(
            [""] * len(fixed_cols)
            + [g for gate in C.ALL_GATES for g in (LABELS[gate], "", "")]
        )
        ws.append(list(fixed_cols) + [c.rsplit("_", 1)[-1] for c in fields[len(fixed_cols):]])
        for col_idx in range(1, len(fixed_cols) + 1):
            for r in (1, 2):
                cell = ws.cell(row=r, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
        for gate_idx, gate in enumerate(C.ALL_GATES):
            base_col = len(fixed_cols) + gate_idx * 3 + 1
            ws.merge_cells(start_row=1, start_column=base_col, end_row=1, end_column=base_col + 2)
            top_cell = ws.cell(row=1, column=base_col)
            top_cell.font = header_font
            top_cell.fill = header_fill
            top_cell.alignment = Alignment(horizontal="center")
            for offset, is_delta in enumerate((False, False, True)):
                cell = ws.cell(row=2, column=base_col + offset)
                cell.font = header_font
                cell.fill = delta_header_fill if is_delta else header_fill

        for row in rows:
            values = []
            for f in fields:
                v = row.get(f, "")
                if isinstance(v, float):
                    v = round(v, 4) if math.isfinite(v) else "NA"
                values.append(v)
            ws.append(values)
        for r_idx, row in enumerate(rows, start=3):
            for bool_col in ("prg_pass", "noprg_pass"):
                col_idx = fields.index(bool_col) + 1
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.fill = pass_fill if row[bool_col] else fail_fill

        ws.column_dimensions["A"].width = 30
        for col_idx in range(2, len(fixed_cols) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 20
        for col_idx in range(len(fixed_cols) + 1, len(fields) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 11
        ws.freeze_panes = "F3"

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_XLSX)


def write_md(by_object: dict[str, list[dict[str, Any]]], gate_counts: dict[str, dict[str, dict[str, int]]]) -> None:
    lines = [
        "# box023 / box021 PRG vs no-PRG 十二门对比",
        "",
        "_box023: E173(PRG) vs E179(no-PRG, 16/16配对消融) · "
        "box021: E170(frozen PRG) vs E168(no-PRG baseline, 28 case)_",
        "",
        "两个物体的十二门 scoring 均复用 `core4d-e179-12gate-tracking-v1` 口径"
        "（六门physics + root/EEF/object tracking），box021 是本次新用同一套"
        "`apply_12gate_scoring` 只读重算，box023 直接复用 E179 已发布结果，均未"
        "重跑 CEM。Delta = PRG − no-PRG。",
        "",
        "## 总览",
        "",
        "| 物体 | n | PRG 来源 | no-PRG 来源 | PRG 12门 | no-PRG 12门 |",
        "|---|---:|---|---|---:|---:|",
    ]
    for obj in OBJECTS:
        rows = by_object[obj]
        n = len(rows)
        prg_pass = sum(row["prg_pass"] for row in rows)
        noprg_pass = sum(row["noprg_pass"] for row in rows)
        lines.append(
            f"| {obj} | {n} | {SOURCE_LABELS[obj]['prg']} | "
            f"{SOURCE_LABELS[obj]['noprg']} | {prg_pass}/{n} | {noprg_pass}/{n} |"
        )
    for obj in OBJECTS:
        rows = by_object[obj]
        n = len(rows)
        lines.extend([
            "",
            f"## {obj}（n={n}）逐门通过数",
            "",
            "| Gate | PRG | no-PRG | Delta |",
            "|---|---:|---:|---:|",
        ])
        for gate in C.ALL_GATES:
            c = gate_counts[obj][gate]
            lines.append(f"| {LABELS[gate]} | {c['prg']}/{n} | {c['noprg']}/{n} | {c['prg']-c['noprg']:+d} |")
        migrations = {"PASS_TO_PASS": 0, "PASS_TO_FAIL": 0, "FAIL_TO_PASS": 0, "FAIL_TO_FAIL": 0}
        for row in rows:
            migrations[row["pass_migration"]] += 1
        lines.extend([
            "",
            "| Migration | Count |",
            "|---|---:|",
        ])
        for key, count in migrations.items():
            lines.append(f"| {key} | {count} |")
    OUTPUT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    box023_rows, box023_gates = load_box023()
    box021_rows, box021_gates = load_box021()
    by_object = {"box023": box023_rows, "box021": box021_rows}
    gates_by_object = {"box023": box023_gates, "box021": box021_gates}
    gate_counts = {
        obj: gate_pass_counts(by_object[obj], gates_by_object[obj]) for obj in OBJECTS
    }
    write_xlsx(by_object)
    write_md(by_object, gate_counts)
    print(OUTPUT_XLSX)
    print(OUTPUT_MD)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
