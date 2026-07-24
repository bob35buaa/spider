#!/usr/bin/env python3
"""Generate the E171 box022/box026 full-validation workbook."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_EVAL = Path("workspace/core4d/results/E171/s6_downstream/eval/full")
DEFAULT_KEYFRAMES = Path(
    "workspace/core4d/results/E171/s6_downstream/evidence/visual_qc/keyframe_manifest.tsv"
)
SHEETS = (
    ("Case Metrics", "e171_case_metrics.tsv"),
    ("Paired Deltas", "e171_paired_deltas.tsv"),
    ("Group Summary", "e171_group_summary.tsv"),
    ("Worst Cases", "e171_worst_cases.tsv"),
    ("Manual Review", "user_manual_review_template.tsv"),
    ("Codex Verification", "codex_verification.tsv"),
    ("Not Ready", "e171_not_ready.tsv"),
    ("Eval Errors", "e171_evaluation_errors.tsv"),
)


def read_tsv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if not path.is_file() or path.stat().st_size <= 1:
        return [], []
    with path.open("r", encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream, delimiter="\t")
        return list(reader.fieldnames or []), list(reader)


def value(raw: str) -> Any:
    text = str(raw)
    if text.lower() == "true":
        return True
    if text.lower() == "false":
        return False
    if text == "":
        return None
    try:
        return float(text) if any(char in text.lower() for char in (".", "e")) else int(text)
    except ValueError:
        return text


def style_sheet(sheet: Any) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        width = min(45, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[letter].width = width
        for cell in column:
            cell.font = Font(
                name="Arial",
                size=9,
                bold=cell.row == 1,
                color="FFFFFF" if cell.row == 1 else "000000",
            )
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")


def add_tsv_sheet(
    workbook: Workbook, title: str, path: Path
) -> tuple[dict[str, int], int]:
    fields, rows = read_tsv(path)
    sheet = workbook.create_sheet(title)
    if not fields:
        sheet.append(["status"])
        sheet.append(["not_generated_or_empty"])
        style_sheet(sheet)
        return {"status": 1}, 0
    sheet.append(fields)
    for row in rows:
        sheet.append([value(row.get(field, "")) for field in fields])
    style_sheet(sheet)
    return {field: index + 1 for index, field in enumerate(fields)}, len(rows)


def col_range(sheet: str, column: int, last: int) -> str:
    letter = get_column_letter(column)
    return f"'{sheet}'!{letter}2:{letter}{last}"


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--eval-dir", type=Path, default=DEFAULT_EVAL)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--keyframe-manifest", type=Path, default=DEFAULT_KEYFRAMES)
    args = parser.parse_args()

    eval_dir = args.eval_dir
    output = args.output or eval_dir / "E171_box022_box026_prg_full_validation.xlsx"
    summary = json.loads((eval_dir / "summary.json").read_text(encoding="utf-8"))

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"

    mappings: dict[str, dict[str, int]] = {}
    row_counts: dict[str, int] = {}
    for title, filename in SHEETS:
        mappings[title], row_counts[title] = add_tsv_sheet(
            workbook, title, eval_dir / filename
        )
    mappings["Keyframe Evidence"], row_counts["Keyframe Evidence"] = add_tsv_sheet(
        workbook, "Keyframe Evidence", args.keyframe_manifest
    )

    evaluated = summary["counts"]["evaluated"]
    case_last = row_counts["Case Metrics"] + 1
    review_last = row_counts["Manual Review"] + 1
    case = mappings["Case Metrics"]
    review = mappings["Manual Review"]
    required_case = (
        "case_id",
        "numeric_release_pass",
        "strict_release_usable",
        "leg_gate_health_pass",
    )
    required_review = (
        "case_id",
        "user_manual_review_status",
        "manual_use_decision",
    )
    missing = [field for field in required_case if field not in case]
    missing += [field for field in required_review if field not in review]
    if missing:
        raise RuntimeError(f"Missing workbook fields: {missing}")
    if row_counts["Case Metrics"] != evaluated or row_counts["Manual Review"] != evaluated:
        raise RuntimeError(
            "Case/manual row count mismatch: "
            f"summary={evaluated}, case={row_counts['Case Metrics']}, "
            f"manual={row_counts['Manual Review']}"
        )

    case_ids = col_range("Case Metrics", case["case_id"], case_last)
    numeric = col_range(
        "Case Metrics", case["numeric_release_pass"], case_last
    )
    strict = col_range(
        "Case Metrics", case["strict_release_usable"], case_last
    )
    gate = col_range(
        "Case Metrics", case["leg_gate_health_pass"], case_last
    )
    review_status = col_range(
        "Manual Review", review["user_manual_review_status"], review_last
    )
    review_decision = col_range(
        "Manual Review", review["manual_use_decision"], review_last
    )

    overview.append(
        ["E171 Box022/Box026 PRG Full Validation", "Value", "Source / rule"]
    )
    rows = (
        ("Metric standard", summary["metric_standard_id"], "summary.json"),
        ("Evaluated rows", f"=COUNTA({case_ids})", "COUNTA Case Metrics case_id"),
        ("Numeric release pass", f"=COUNTIF({numeric},TRUE)", "COUNTIF numeric_release_pass"),
        (
            "User reviewed",
            f'=COUNTIF({review_status},"reviewed")',
            "Manual Review authority TSV",
        ),
        (
            "Manual operational USE",
            f'=COUNTIF({review_decision},"USE")',
            "Manual Review authority TSV",
        ),
        (
            "Manual DO_NOT_USE",
            f'=COUNTIF({review_decision},"DO_NOT_USE")',
            "Manual Review authority TSV",
        ),
        ("Strict release usable", f"=COUNTIF({strict},TRUE)", "COUNTIF strict_release_usable"),
        ("Gate-health pass", f"=COUNTIF({gate},TRUE)", "COUNTIF leg_gate_health_pass"),
        (
            "Machine recommendation",
            f'=IF(B5<{evaluated},"PENDING_USER_REVIEW","USER_REVIEW_COMPLETE")',
            f"requires {evaluated} manual reviews",
        ),
        ("Final promotion decision", "PENDING_USER_DECISION", "User authority"),
        ("Manifest SHA256", summary["manifest_sha256"], summary["manifest"]),
        ("Baseline SHA256", summary.get("baseline_sha256", ""), summary.get("baseline", "")),
    )
    for row in rows:
        overview.append(row)

    overview.freeze_panes = "A2"
    overview.column_dimensions["A"].width = 42
    overview.column_dimensions["B"].width = 34
    overview.column_dimensions["C"].width = 65
    overview.row_dimensions[12].height = 30
    overview.sheet_properties.pageSetUpPr.fitToPage = True
    overview.page_setup.orientation = "landscape"
    overview.page_setup.fitToWidth = 1
    overview.page_setup.fitToHeight = 1
    overview.print_area = "A1:C13"
    for cell in overview[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=12)
        cell.fill = PatternFill("solid", fgColor="17365D")
    for row in overview.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in (overview["B5"], overview["B10"], overview["B11"]):
        cell.fill = PatternFill("solid", fgColor="FFF2CC")

    workbook.move_sheet("Overview", -(len(workbook.sheetnames) - 1))
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(
        json.dumps(
            {
                "output": str(output),
                "sheets": workbook.sheetnames,
                "evaluated": evaluated,
                "manual_review_rows": row_counts["Manual Review"],
                "keyframe_rows": row_counts["Keyframe Evidence"],
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
