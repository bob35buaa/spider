#!/usr/bin/env python3
"""E209 P7 report: the PRG-vs-G1 workbook, formatted for reading rather than parsing.

Sheets:
  说明            what each sheet is + the verdict up front
  claims          C0-C9 with the gate, the measured value and the verdict
  funnel_summary  per-arm L1/L2/L3 + hard/wide/narrow + 12-gate counts
  per_gate        per gate: mean/std/min/max per arm + paired delta (G1 - PRG)
  paired_delta    per case: G1 - PRG on the decision-relevant metrics
  flips           the 9 cases whose narrow verdict changed, with the driver
  z_diff          per case z bias / MAE / 3D pos, with the S+/S- stratum
  model           C4d: the three pre-registered transfer functions, scored
  per_object      object x arm, with n and whether n permits a recommendation
  rollout         the full 44-row scored data (what every other sheet points at)

Style helpers are imported from gen_E206_two_arm_workbook (rule 13) so E206 and
E209 workbooks are visually one family and the palette lives in one place.

Reporting choices that are deliberate, not cosmetic:
  * `worst` sits next to every mean -- rule 5 forbids the mean alone.
  * the `model` sheet reports the *observed* OLS fit next to the three frozen
    candidates, and separates "which model family won" (settled: r=-0.915,
    slope=0.408) from "did E207's coefficients transfer" (they did not).
    Collapsing those two into one PASS/FAIL would misreport the finding.
  * per_object rows with n < 3 are greyed -- chair005 (n=1) must not be read as
    a per-object recommendation.
  * gates where G1 *improved* (object tracking) are shown in the same table as
    the ones where it regressed (eef_ori). Reporting only the regression would
    be as misleading as reporting only the win.

Usage:
    .venv/bin/python .../gen_E209_two_arm_workbook.py
"""

from __future__ import annotations

import csv
import json
import math
import statistics
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[5]
for _p in (
    "workspace/core4d/scripts/eval/reports",
    "workspace/core4d/scripts/experiments/E201",
    "workspace/core4d/scripts/experiments/E209",
):
    _s = str(REPO / _p)
    if _s not in sys.path:
        sys.path.insert(0, _s)

import funnel_config as FC  # noqa: E402
import e209_common as C  # noqa: E402
# style + numeric helpers, shared with the E206 workbook
from gen_E206_two_arm_workbook import (  # noqa: E402
    ARIAL, BLUE, GREEN_TXT, LAYER_FILL, LIGHT_AMBER, LIGHT_GRAY, LIGHT_GREEN,
    LIGHT_RED, LOWER_BETTER, NAVY, RED_TXT, WHITE,
    body_font, delta_scale, f, finish_table, header, set_widths, title,
)

EVAL = C.S6_DIR / "eval/two_arm"
OUT = EVAL / "e209_two_arm.xlsx"
ROLLOUT = EVAL / "e209_two_arm_rollout.tsv"
PER_GATE = EVAL / "e209_per_gate.tsv"
PER_OBJECT = EVAL / "e209_per_object.tsv"
SUMMARY = EVAL / "e209_two_arm_summary.json"
ZDIFF = EVAL / "e209_object_z_diff_by_case.tsv"
ZSUM = EVAL / "e209_object_z_diff_summary.json"

ARMS = ("prg", "g1")
ARM_TITLE = {"prg": "PRG (E206 基线)", "g1": "G1 (PRG+gravcomp)"}

#: All 14 funnel gates in E201 order: 4 hard then 10 banded. `paired_delta` and
#: `per_gate` both report the full set -- an earlier version showed only a
#: 7-metric hand-picked subset, which let a reader assume the untabulated gates
#: were unchanged when they had simply not been computed.
GATES_14: list[tuple[str, str, str]] = (
    [(g[0], g[1], "hard") for g in FC.HARD_GATES]
    + [(g[0], g[1], "banded") for g in FC.BANDED_GATES]
)
GATE_FIELDS_14 = [f for _n, f, _k in GATES_14]

#: `fall_flag` is boolean; a numeric delta on it is meaningless, so the
#: paired_delta sheet renders it as a state transition instead.
BOOL_FIELDS = {"fall_flag"}


def gate_label(name: str, kind: str) -> str:
    return f"{name}*" if kind == "hard" else name


def numv(v: Any) -> float:
    s = str(v).strip().lower()
    if s in ("true", "false"):
        return 1.0 if s == "true" else 0.0
    return f(v)


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh, delimiter="\t"))


def verdict_cell(ws, row: int, col: int, ok: bool | None) -> None:
    cell = ws.cell(row, col)
    if ok is None:
        cell.value, fill, txt = "待做", LIGHT_AMBER, "7F6000"
    else:
        cell.value, fill, txt = ("PASS", LIGHT_GREEN, GREEN_TXT) if ok else ("FAIL", LIGHT_RED, RED_TXT)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name=ARIAL, size=9, bold=True, color=txt)
    cell.alignment = Alignment(horizontal="center")


def sheet_readme(wb, zc: dict, sc: dict) -> None:
    ws = wb.create_sheet("说明")
    title(ws, "E209 · desk/chair PRG + G1（object gravcomp）单变量",
          "R295 · Phase 68 · plan239 / log298 · 22 例 E206 交付 case，基线复用 E206 PRG rollout 零重跑。"
          "唯一变量 = object body gravcomp 0→1。", 3)
    c4d = zc["C4d_model"]
    emp = zc["empirical"]
    lines = [
        ("总判定", "FAIL（主门 C3 破），但取得决定性机制结论", LIGHT_RED),
        ("", "", None),
        ("主门 C3 · 14-gate", f"narrow {sc['arms']['prg']['narrow_pass']}/22 → "
         f"{sc['arms']['g1']['narrow_pass']}/22（门 ≥12）；L1_reject 3 → 11", LIGHT_RED),
        ("退化机理", "几乎全在 eef_ori（narrow 失败 2→9，mean +2.78°）+ root_ori +1.81°；"
         "而 obj_pos/obj_ori/leg_pen 全改善、contact 仅 −0.016 —— 物体失重后 CEM 用姿态"
         "跑偏的机器人「托」着物体走（虚扶 / ghost-carry）", LIGHT_AMBER),
        ("", "", None),
        ("z 干预有效", f"bias {zc['C4c_all']['baseline_macro_bias_cm']:+.3f} → "
         f"{zc['C4c_all']['macro_bias_cm']:+.3f} cm；mean|bias| "
         f"{zc['C4c_all']['baseline_mean_abs_bias_cm']:.3f} → {zc['C4c_all']['mean_abs_bias_cm']:.3f}；"
         f"{zc['C4c_all']['improved']}/22 改善；Wilcoxon p="
         f"{zc['C4c_all']['wilcoxon_p_abs_bias_decreased']:.2e}", LIGHT_GREEN),
        ("核心科学结论",
         f"证伪 E207「gravcomp = +1.945cm 常量偏移」：实测 r(pre,delta)="
         f"{emp['r_pre_vs_delta']:+.3f}、OLS slope={emp['observed_fit_slope']:.3f}"
         f"（加性型预测 0 与 1.0）→ 收缩型。胜出模型 {c4d['winner']}"
         f"（RMSE {c4d['winner_rmse_cm']:.3f} vs {c4d['runner_up']} "
         f"{c4d['runner_up_rmse_cm']:.3f}）", LIGHT_GREEN),
        ("C4d 形式 FAIL 的读法",
         "「模型族已判定」与「E207 系数可否外推」是两件事。前者由 r 与 slope 决定，与门无关；"
         "后者是 C4d 两条门在问的，答案是不够精确（M_shrink 残差偏置 −0.910cm，"
         "desk/chair 截距比 bucket 高）。不要合并成一个 PASS/FAIL 读。", LIGHT_AMBER),
        ("", "", None),
        ("方法学产出", "contact_in_mask 不是虚扶的检测器（−0.016 毫无反应），eef_ori 才是。"
         "E194 C4 当年担心的风险是对的但盯错了指标；E207 在 bucket 上未触发故未暴露。", LIGHT_GREEN),
        ("显式不作判据", "z MAE 降幅（gravcomp 修 bias 不修 MAE）；2×2 交互项（未跑 noPRG+G1）；"
         "真实重力下可执行性（沿用 plan237 的声明式接受）", LIGHT_GRAY),
    ]
    r = 4
    for k, v, fill in lines:
        ws.cell(r, 1, k).font = Font(name=ARIAL, size=9, bold=True)
        ws.cell(r, 2, v).font = Font(name=ARIAL, size=9)
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        if fill:
            for c in (1, 2):
                ws.cell(r, c).fill = PatternFill("solid", fgColor=fill)
        ws.row_dimensions[r].height = 30 if v and len(v) > 70 else 16
        r += 1

    r += 1
    ws.cell(r, 1, "工作表").font = Font(name=ARIAL, size=10, bold=True, color=WHITE)
    ws.cell(r, 1).fill = PatternFill("solid", fgColor=BLUE)
    ws.cell(r, 2, "内容").font = Font(name=ARIAL, size=10, bold=True, color=WHITE)
    ws.cell(r, 2).fill = PatternFill("solid", fgColor=BLUE)
    r += 1
    for name, desc in [
        ("claims", "C0–C9 逐条：门、实测、判定"),
        ("funnel_summary", "两臂 L1/L2/L3 + hard/wide/narrow + 12-gate 计数"),
        ("per_gate", "逐门 mean/std/min/max + 配对 delta（G1 − PRG）"),
        ("paired_delta", "逐 case 的决策指标 G1 − PRG"),
        ("flips", "9 例 narrow 判定翻转的 case 与驱动门"),
        ("z_diff", "逐 case z bias / MAE / 3D pos，带 S+/S− 分层"),
        ("model", "C4d：三个预注册传递函数的评分 + 实测拟合"),
        ("per_object", "物体 × arm，带 n（n<3 灰显，不可作 per-object 推荐）"),
        ("rollout", "44 行完整打分数据（其余所有表的来源）"),
    ]:
        ws.cell(r, 1, name).font = Font(name=ARIAL, size=9, bold=True)
        ws.cell(r, 2, desc).font = Font(name=ARIAL, size=9)
        r += 1
    set_widths(ws, {1: 22, 2: 118, 3: 10})


def sheet_claims(wb, zc: dict, sc: dict) -> None:
    ws = wb.create_sheet("claims")
    title(ws, "Claims 结算（plan239）", "门在跑之前就写进 plan239 与 e209_common.GATES；"
          "此表只填实测与判定。", 4)
    ws.append(["Claim", "门", "实测", "判定"])
    header(ws, 3, 4)
    c4, c4b, c4c, c4d = zc["C4_s_minus"], zc["C4b_s_plus"], zc["C4c_all"], zc["C4d_model"]
    cl = sc["claims"]
    rows = [
        ("C0 输入零漂移", "22/22 traj+mask+baseline scene sha256 == E206 交付表", "22/22", True),
        ("C1 单变量·配置层", "compose 对称差 == {scene_name}", "22/22", True),
        ("C1b 单变量·场景层", "gravcomp 断言 + 编译层 ngeom/npair/nq/nv/nu/nbody 相等", "22/22", True),
        ("C1c 单变量·运行时", "config_act 全键 diff 仅 THE_VARIABLE；A0 hand-gate", "22/22", True),
        ("C2 执行闭合", "22/22，0 fail", "22/22，problem rows: []", cl["C2_execution"]["pass"]),
        ("C3 主门 · 14-gate", f"narrow ≥ {C.GATES['C3_narrow_pass_min']}/22 ∧ hard == 22/22",
         f"narrow {cl['C3_14gate']['g1_narrow_pass']}/22（基线 "
         f"{cl['C3_14gate']['prg_narrow_pass']}）；hard {cl['C3_14gate']['g1_hard_pass']}/22",
         cl["C3_14gate"]["pass"]),
        ("C4 z 修正 · S⁻(n=18)",
         f"①|宏|≤{C.GATES['C4_s_minus_abs_macro_bias_max_cm']} ②≥{C.GATES['C4_s_minus_improved_min']}/18 改善 "
         f"③|post|>{C.GATES['C4_s_minus_overshoot_abs_cm']}cm 者≤{C.GATES['C4_s_minus_overshoot_max_cases']}",
         f"①{abs(c4['macro_bias_cm']):.3f} ✓ ②{c4['improved']}/18 ✓ ③{len(c4['overshoot_cases'])} ✗",
         c4["pass"]),
        ("C4b z 危害上限 · S⁺(n=4)",
         f"|宏|≤{C.GATES['C4b_s_plus_abs_macro_bias_max_cm']} ∧ 无单例>{C.GATES['C4b_s_plus_abs_case_max_cm']}",
         f"{c4b['macro_bias_cm']:+.3f}；最差 {c4b['worst_abs_cm']:.3f}", c4b["pass"]),
        ("C4c 全 22 例反挑拣",
         f"|宏|<{C.GATES['C4c_all_abs_macro_bias_max_cm']} ∧ mean|b|<{C.GATES['C4c_all_mean_abs_bias_max_cm']} ∧ p<0.05",
         f"{c4c['macro_bias_cm']:+.3f} / {c4c['mean_abs_bias_cm']:.3f} / "
         f"p={c4c['wilcoxon_p_abs_bias_decreased']:.2e}", c4c["pass"]),
        ("C4d 传递函数判别",
         f"胜者 RMSE≤{C.GATES['C4d_rmse_max_cm']} ∧ 比次者低≥{C.GATES['C4d_rmse_win_ratio_min']}×",
         f"胜者 {c4d['winner']} RMSE {c4d['winner_rmse_cm']:.3f}；比值 {c4d['rmse_ratio']:.2f}×",
         c4d["pass"]),
        ("C5 承重接触不塌",
         f"contact≥{C.GATES['C5_contact_in_mask_min']} ∧ leg_pen≤{C.GATES['C5_leg_pen_frac_max']}",
         f"{cl['C5_load_bearing']['contact_in_mask_mean']:.4f} / "
         f"{cl['C5_load_bearing']['leg_pen_frac_mean']:.4f}", cl["C5_load_bearing"]["pass"]),
        ("C6 副作用天花板",
         f"release≤{C.GATES['C6_release_false_contact_max']} ∧ hand_pen≤{C.GATES['C6_hand_pen_3mm_max']}",
         f"{cl['C6_side_effects']['release_false_contact_mean']:.4f} / "
         f"{cl['C6_side_effects']['hand_pen_3mm_mean']:.4f}", cl["C6_side_effects"]["pass"]),
        ("C7 人审不降级", f"22/22 已审 ∧ USE≥{C.GATES['C7_use_min']}", "待 P8", None),
        ("C8 吞吐无回归",
         f"median ∈ [{C.GATES['C8_wall_median_min_min']}, {C.GATES['C8_wall_median_max_min']}] min",
         "42.2 min（基线 44.0）", True),
        ("C9 复现性双保障", "git add -f + 快照含双臂 sha256", "132 文件 + 22 目录", True),
    ]
    for i, (cid, gate, got, ok) in enumerate(rows, start=4):
        ws.cell(i, 1, cid)
        ws.cell(i, 2, gate)
        ws.cell(i, 3, got)
        verdict_cell(ws, i, 4, ok)
    body_font(ws, 4, 3 + len(rows))
    for i in range(4, 4 + len(rows)):
        ws.cell(i, 4).alignment = Alignment(horizontal="center")
        for c in (2, 3):
            ws.cell(i, c).alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        ws.row_dimensions[i].height = 26
    finish_table(ws, 3, 3 + len(rows), 4, "E209Claims")
    set_widths(ws, {1: 24, 2: 56, 3: 44, 4: 9})


def sheet_funnel(wb, sc: dict) -> None:
    ws = wb.create_sheet("funnel_summary")
    title(ws, "14-gate 漏斗汇总（E201 口径）",
          "两臂都从 rollout 重打分、走同一 code path；PRG 侧与 P0 冻结基线逐位一致，"
          "故差异确来自 gravcomp。", 11)
    cols = ["arm", "n", "L3_auto", "L2_review", "L1_reject", "hard_pass", "wide_pass",
            "narrow_pass", "gate12_pass", "physics6_pass", "tracking6_pass"]
    ws.append(cols)
    header(ws, 3, len(cols))
    for arm in ARMS:
        s = sc["arms"][arm]
        ws.append([ARM_TITLE[arm]] + [s[k] for k in cols[1:]])
    d = {k: sc["arms"]["g1"][k] - sc["arms"]["prg"][k] for k in cols[1:]}
    ws.append(["Δ (G1 − PRG)"] + [d[k] for k in cols[1:]])
    last = ws.max_row
    body_font(ws, 4, last)
    for c in range(2, len(cols) + 1):
        cell = ws.cell(last, c)
        v = cell.value
        cell.font = Font(name=ARIAL, size=9, bold=True,
                         color=(RED_TXT if isinstance(v, int) and v < 0 else GREEN_TXT))
    finish_table(ws, 3, last, len(cols), "E209Funnel")
    set_widths(ws, {1: 22}, default=13)


def sheet_per_gate(wb) -> None:
    rows = read_tsv(PER_GATE)
    ws = wb.create_sheet("per_gate")
    title(ws, f"逐门统计 + 配对 delta（G1 − PRG）— 全 {len(rows)} 门",
          "4 硬门 + 10 带门全列（kind 列标注）。每个 mean 旁都给 min/max"
          "（rule 5：禁止只报均值）。delta 列为零中心梯度：绿 = G1 改善，红 = G1 退化。"
          "注意 obj_pos/obj_ori/leg_pen 是绿的 —— gravcomp 在物体侧确实是净收益，"
          "代价出在 eef_ori / root_ori。fall 已按 0/1 计入。", 13)
    cols = ["gate", "kind", "field",
            "prg_mean", "prg_std", "prg_min", "prg_max",
            "g1_mean", "g1_std", "g1_min", "g1_max",
            "paired_delta_mean", "paired_delta_max"]
    ws.append(["门", "类型", "字段", "PRG mean", "PRG std", "PRG min", "PRG max",
               "G1 mean", "G1 std", "G1 min", "G1 max", "Δ mean", "Δ max"])
    header(ws, 3, len(cols))
    for r in rows:
        ws.append([r.get(c, "") if c in ("gate", "kind", "field") else f(r.get(c))
                   for c in cols])
        if r.get("kind") == "hard":
            for c in (1, 2):
                ws.cell(ws.max_row, c).fill = PatternFill("solid", fgColor=LIGHT_AMBER)
    last = ws.max_row
    body_font(ws, 4, last)
    for i, r in enumerate(rows, start=4):
        delta_scale(ws, f"L{i}:L{i}", LOWER_BETTER.get(r["field"], True))
    finish_table(ws, 3, last, len(cols), "E209PerGate")
    set_widths(ws, {1: 14, 2: 8, 3: 46}, default=11)


def sheet_paired(wb, rollout: list[dict]) -> None:
    by = {(r["arm"], r["case_id"]): r for r in rollout}
    cases = [c for c in C.CASES if ("g1", c) in by and ("prg", c) in by]
    ncol = 2 + len(GATES_14)
    ws = wb.create_sheet("paired_delta")
    title(ws, "逐 case · 全 14 门：G1 − PRG",
          "全部 14 门（4 硬门带 * + 10 带门），不是挑出来的子集。"
          "正 = G1 更差（contact 相反，其为 >= 门）。fall 是布尔，列出状态迁移而非差值。"
          "eef_ori 列是本次退化主因；obj_pos / obj_ori 列是干预收益 —— 必须并列看。", ncol)
    ws.append(["case_id", "层"] + [gate_label(n, k) for n, _fld, k in GATES_14])
    header(ws, 3, ncol)
    for cid in cases:
        p, g = by[("prg", cid)], by[("g1", cid)]
        row: list[Any] = [cid, p["stratum"]]
        for _n, fld, _k in GATES_14:
            if fld in BOOL_FIELDS:
                pv, gv = str(p[fld]).strip(), str(g[fld]).strip()
                row.append("—" if pv == gv else f"{pv}→{gv}")
            else:
                row.append(numv(g[fld]) - numv(p[fld]))
        ws.append(row)
    last = ws.max_row
    body_font(ws, 4, last)
    for j, (_n, fld, _k) in enumerate(GATES_14):
        col = get_column_letter(3 + j)
        if fld in BOOL_FIELDS:
            for r in range(4, last + 1):
                ws.cell(r, 3 + j).alignment = Alignment(horizontal="center")
            continue
        delta_scale(ws, f"{col}4:{col}{last}", LOWER_BETTER.get(fld, True))
    finish_table(ws, 3, last, ncol, "E209Paired")
    set_widths(ws, {1: 34, 2: 6}, default=12)


def sheet_flips(wb, rollout: list[dict]) -> None:
    by = {(r["arm"], r["case_id"]): r for r in rollout}
    ws = wb.create_sheet("flips")
    title(ws, "narrow 判定翻转的 case",
          "6 例转坏 / 3 例转好。转坏里 4 例是 eef_ori 越过 20° 窄门；"
          "desk007_20231030_034_p1 是彻底崩坏（不是擦边）。", 7)
    ws.append(["case_id", "层", "PRG narrow", "G1 narrow", "方向",
               "eef_ori PRG→G1 (°)", "G1 失败门"])
    header(ws, 3, 7)
    n_bad = n_good = 0
    for cid in C.CASES:
        p, g = by[("prg", cid)], by[("g1", cid)]
        if p["narrow_pass"] == g["narrow_pass"]:
            continue
        good = g["narrow_pass"] == "True"
        n_good, n_bad = (n_good + 1, n_bad) if good else (n_good, n_bad + 1)
        ws.append([cid, p["stratum"], p["narrow_pass"], g["narrow_pass"],
                   "转好" if good else "转坏",
                   f"{f(p['track_eef_ori_err_deg_mean']):.2f} → {f(g['track_eef_ori_err_deg_mean']):.2f}",
                   g["narrow_failed"] or ""])
        ws.cell(ws.max_row, 5).fill = PatternFill(
            "solid", fgColor=LIGHT_GREEN if good else LIGHT_RED)
    last = ws.max_row
    body_font(ws, 4, last)
    finish_table(ws, 3, last, 7, "E209Flips")
    set_widths(ws, {1: 34, 2: 6, 3: 12, 4: 12, 5: 8, 6: 20, 7: 40})


def sheet_zdiff(wb, zc: dict) -> None:
    rows = read_tsv(ZDIFF)
    by = {(r["arm"], r["case_id"]): r for r in rows}
    ws = wb.create_sheet("z_diff")
    title(ws, "物体 z 诊断（gravcomp 的直接目标）",
          "S⁺/S⁻ 分层由 e209_common.BASELINE_Z_BIAS_CM 在 P0 冻结并随代码提交，"
          "由基线数据机械决定，不按物体名切 —— chair006_20231003_2_015_p1(−0.726) 属 S⁻。", 9)
    ws.append(["case_id", "层", "物体", "参考抬升(m)", "PRG bias(cm)", "G1 bias(cm)",
               "Δ", "|bias| 改善", "PRG→G1 3D pos(cm)"])
    header(ws, 3, 9)
    for cid in C.CASES:
        p, g = by.get(("PRG", cid)), by.get(("G1", cid))
        if not p or not g:
            continue
        pb, gb = f(p["z_bias_cm"]), f(g["z_bias_cm"])
        better = abs(gb) < abs(pb)
        ws.append([cid, p["stratum"], p["object_key"], f(p["ref_z_range_m"]),
                   pb, gb, gb - pb, "✓" if better else "✗",
                   f"{f(p['obj_pos_err_3d_cm']):.2f} → {f(g['obj_pos_err_3d_cm']):.2f}"])
        ws.cell(ws.max_row, 8).fill = PatternFill(
            "solid", fgColor=LIGHT_GREEN if better else LIGHT_RED)
        ws.cell(ws.max_row, 8).alignment = Alignment(horizontal="center")
    last = ws.max_row
    body_font(ws, 4, last)
    finish_table(ws, 3, last, 9, "E209Z")
    set_widths(ws, {1: 34, 2: 6, 3: 11, 4: 13, 5: 14, 6: 14, 7: 10, 8: 11, 9: 20})

    r = last + 2
    c4, c4b, c4c = zc["C4_s_minus"], zc["C4b_s_plus"], zc["C4c_all"]
    for label, d in (("S⁻ (n=18)", c4), ("S⁺ (n=4)", c4b), ("全 22 例", c4c)):
        ws.cell(r, 1, label).font = Font(name=ARIAL, size=9, bold=True)
        ws.cell(r, 2, f"宏 bias {d['baseline_macro_bias_cm']:+.3f} → {d['macro_bias_cm']:+.3f} cm")
        ws.cell(r, 2).font = Font(name=ARIAL, size=9)
        r += 1


def sheet_model(wb, zc: dict) -> None:
    c4d, emp = zc["C4d_model"], zc["empirical"]
    ws = wb.create_sheet("model")
    title(ws, "C4d · gravcomp 传递函数判别（E209 的核心科学产出）",
          "三个候选模型的系数在跑之前冻结进 e209_common.PREREG_MODELS，此处只算 RMSE。"
          "E207 的 9 例里物体质量与 pre-bias 共线，两族不可分；E209 全 22 例 mass=5.000kg "
          "而 pre-bias 跨 −5.65…+2.67，可分。", 6)
    forms = {
        "M_shrink": "post = 1.0939 + 0.3827·pre（E207 OLS，收缩型）",
        "M_pooled": "post = pre + 1.945（E207 池化常量，加性型）",
        "M_mass": "post = pre + 2.836（E207 mass=5 单元，加性型）",
    }
    ws.append(["模型", "形式", "RMSE (cm)", "MAE (cm)", "残差偏置 (cm)", "最差 (cm)"])
    header(ws, 3, 6)
    for name, s in sorted(c4d["scores"].items(), key=lambda kv: kv[1]["rmse_cm"]):
        ws.append([name + ("  ⬅ 胜" if name == c4d["winner"] else ""), forms[name],
                   s["rmse_cm"], s["mae_cm"], s["bias_cm"], s["max_abs_cm"]])
        if name == c4d["winner"]:
            for c in range(1, 7):
                ws.cell(ws.max_row, c).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    last = ws.max_row
    body_font(ws, 4, last)
    finish_table(ws, 3, last, 6, "E209Model")
    set_widths(ws, {1: 18, 2: 46}, default=14)

    r = last + 2
    ws.cell(r, 1, "模型族判别（与门无关，由数据直接决定）").font = Font(
        name=ARIAL, size=10, bold=True, color=WHITE)
    ws.cell(r, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    r += 1
    for k, add, shr, got in [
        ("r(pre_bias, delta)", "≈ 0", "显著负", f"{emp['r_pre_vs_delta']:+.3f}"),
        ("OLS slope", "≈ 1.0", "< 1", f"{emp['observed_fit_slope']:.3f}"),
        ("delta 是否恒定", "是", "否",
         f"{emp['delta_mean_cm']:+.3f} ± {emp['delta_sd_cm']:.3f}（E207 池化 +1.945±0.711）"),
    ]:
        for c, v in enumerate([k, "加性型预测：" + add, "收缩型预测：" + shr, got], start=1):
            cell = ws.cell(r, c if c < 4 else 4, v)
            cell.font = Font(name=ARIAL, size=9, bold=(c == 4))
        ws.cell(r, 4).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
        r += 1
    r += 1
    for txt in [
        f"实测 OLS：post = {emp['observed_fit_intercept']:+.4f} + "
        f"{emp['observed_fit_slope']:.4f}·pre，与 M_shrink 的斜率 0.3827 几乎一致。",
        f"→ 收缩型，不动点 = {emp['observed_fit_intercept']:.4f}/(1−{emp['observed_fit_slope']:.4f}) "
        f"= {emp['observed_fit_intercept']/(1-emp['observed_fit_slope']):+.2f} cm。",
        "理论满悬挂沉降 m·g/kp = 5×9.81/500 = 9.81 cm；实测修正 "
        f"{emp['delta_mean_cm']:.3f} cm 只占 {100*emp['delta_mean_cm']/9.81:.0f}% "
        "—— 与 E207「手分担了大部分载荷」一致。",
        "",
        "C4d 形式 FAIL 的正确读法：「模型族已判定」与「E207 系数可否外推」是两件事。"
        "前者由上表的 r 与 slope 决定；后者才是 C4d 两条门在问的，答案是不够精确"
        f"（M_shrink 残差偏置 {c4d['scores']['M_shrink']['bias_cm']:+.3f} cm，desk/chair 截距比 bucket 高）。",
        "",
        "可执行推论：gravcomp 是连续量，gravcomp=1 是全额补偿。既然是收缩型且不动点在 "
        f"{emp['observed_fit_intercept']/(1-emp['observed_fit_slope']):+.2f} cm，"
        "部分补偿（如 gravcomp≈0.5）理论上能把 bias 停在 0 附近而少付 eef_ori 的代价。",
    ]:
        ws.cell(r, 1, txt).font = Font(name=ARIAL, size=9)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
        if len(txt) > 80:
            ws.row_dimensions[r].height = 28
        r += 1


def sheet_per_object(wb) -> None:
    rows = read_tsv(PER_OBJECT)
    ws = wb.create_sheet("per_object")
    title(ws, "逐物体 × arm",
          "n < 3 的行灰显：chair005 (n=1) 不可作为 per-object 推荐（rule 5）。", 12)
    cols = ["object_key", "n", "prg_L3", "prg_L2", "prg_L1", "prg_narrow",
            "g1_L3", "g1_L2", "g1_L1", "g1_narrow", "prg_contact_mean", "g1_contact_mean"]
    ws.append(cols)
    header(ws, 3, len(cols))
    for r in rows:
        ws.append([r.get(c, "") if c == "object_key" else f(r.get(c)) for c in cols])
        if int(f(r.get("n", 0))) < 3:
            for c in range(1, len(cols) + 1):
                ws.cell(ws.max_row, c).fill = PatternFill("solid", fgColor=LIGHT_GRAY)
    last = ws.max_row
    body_font(ws, 4, last)
    finish_table(ws, 3, last, len(cols), "E209PerObject")
    set_widths(ws, {1: 14}, default=11)


def sheet_rollout(wb, rollout: list[dict]) -> None:
    ws = wb.create_sheet("rollout")
    title(ws, "完整打分数据（44 行 = 22 case × 2 arm）",
          "其余所有表的唯一来源。两臂都从 rollout 重打分、走同一 code path。", 12)
    cols = list(rollout[0].keys())
    ws.append(cols)
    header(ws, 3, len(cols))
    numeric = {c for c in cols if c not in
               ("object_key", "case_id", "arm", "stratum", "layer", "hard_failed",
                "wide_failed", "narrow_failed", "gate12_failed")}
    for r in rollout:
        ws.append([f(r[c]) if c in numeric and r[c] not in ("", "True", "False")
                   else r[c] for c in cols])
        lay = r.get("layer", "")
        if lay in LAYER_FILL:
            ws.cell(ws.max_row, cols.index("layer") + 1).fill = PatternFill(
                "solid", fgColor=LAYER_FILL[lay])
    last = ws.max_row
    body_font(ws, 4, last)
    finish_table(ws, 3, last, len(cols), "E209Rollout")
    set_widths(ws, {1: 12, 2: 34, 3: 6, 4: 6}, default=13)


def main() -> int:
    for p in (ROLLOUT, PER_GATE, PER_OBJECT, SUMMARY, ZDIFF, ZSUM):
        if not p.is_file():
            raise SystemExit(f"missing input: {p} (run eval_E209_g1_gravcomp.py / "
                             f"gen_E209_object_z_diff.py first)")
    rollout = read_tsv(ROLLOUT)
    sc = json.loads(SUMMARY.read_text(encoding="utf-8"))
    zc = json.loads(ZSUM.read_text(encoding="utf-8"))["claims"]

    wb = Workbook()
    wb.remove(wb.active)
    sheet_readme(wb, zc, sc)
    sheet_claims(wb, zc, sc)
    sheet_funnel(wb, sc)
    sheet_per_gate(wb)
    sheet_paired(wb, rollout)
    sheet_flips(wb, rollout)
    sheet_zdiff(wb, zc)
    sheet_model(wb, zc)
    sheet_per_object(wb)
    sheet_rollout(wb, rollout)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"[done] {OUT.relative_to(REPO)}")
    print(f"  sheets: {wb.sheetnames}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
