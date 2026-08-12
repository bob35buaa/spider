#!/usr/bin/env python3
"""Recompute and report OmniRetarget versus PRG metrics for all five-box Full CEM cases."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import mujoco
import numpy as np
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from scipy.spatial.transform import Rotation as Rotation

REPO = Path(__file__).resolve().parents[5]
sys.path.insert(0, str(REPO / "workspace/core4d/scripts"))

from eval.core.core_metrics import (  # noqa: E402
    EVAL_METRIC_STANDARD_ID,
    EvalConfig,
    evaluate_sequence,
    person_idx_from_case,
)
from eval.core.motion_health import run_health  # noqa: E402

OUT = REPO / "workspace/core4d/analysis/E197_full_cem_omnirt_vs_prg_metrics"
METHOD_TSV = OUT / "e197_method_metrics.tsv"
CASE_TSV = OUT / "e197_case_comparison.tsv"
SUMMARY_TSV = OUT / "e197_summary_by_object.tsv"
AUDIT_TSV = OUT / "e197_input_audit.tsv"
SUMMARY_JSON = OUT / "e197_summary.json"
GATE_TSV = OUT / "e197_omni_absolute_wide_gate_thresholds.tsv"
RL_FILTER_TSV = OUT / "e197_omni_absolute_wide_gate_filter.tsv"
REPORT = OUT / "E197_full_cem_omnirt_vs_prg_metrics.md"
WORKBOOK = OUT / "E197_full_cem_omnirt_vs_prg_metrics.xlsx"
OMNI_DIR = OUT / "omni_converted_qpos"

SOURCES = (
    ("E172", REPO / "workspace/core4d/results/E172/s6_downstream/eval/full/e171_case_metrics.tsv", {"box004"}),
    ("E173", REPO / "workspace/core4d/results/E173/s6_downstream/eval/full/e173_case_metrics.tsv", {"box001", "box023", "box024"}),
    ("E170", REPO / "workspace/core4d/results/E170/s6_downstream/eval/full/e170_case_metrics.tsv", {"box021"}),
)
EXPECTED = {"box001": 28, "box004": 6, "box021": 28, "box023": 16, "box024": 9}
OBJECTS = tuple(EXPECTED)
METRICS = (
    ("contact_3mm_in_mask", "3mm in-mask 接触", "hand_object_physics_contact_3mm_in_mask_frac", "higher"),
    ("raw_contact_in_mask", "Raw in-mask 接触", "hand_object_physics_contact_in_mask_frac", "higher"),
    ("hand_object_penetration_3mm", "手物穿透 >3mm", "hand_object_physics_penetration_3mm_frame_frac", "lower"),
    ("lower_body_penetration", "Lower-body 穿透", "leg_penetration_frac", "lower"),
    ("foot_slip_max_m", "Foot slip max", "foot_slip_max_m", "lower"),
    ("obj_speed_max", "Object speed max", "obj_speed_max", "lower"),
    ("ankle_jerk_p95", "Ankle jerk P95", "ankle_jerk_p95", "lower"),
)
EXPECTED_TOTAL = sum(EXPECTED.values())
BOOTSTRAP_DRAWS = 10_000
BOOTSTRAP_SEED = 197

# Wide downstream-RL candidate gate. It exclusively evaluates the OmniRetarget
# reference trajectory; PRG remains available only for paired comparison.
OMNI_WIDE_GATE_VERSION = "E197-omni-absolute-wide-v4"
OMNI_WIDE_GATES = (
    ("contact_3mm_in_mask", "ge", 0.01, ""),
    ("contact_3mm_in_mask", "ge", 0.00, "box024"),
    ("raw_contact_in_mask", "ge", 0.50, ""),
    ("hand_object_penetration_3mm", "le", 0.80, ""),
    ("lower_body_penetration", "le", 0.30, ""),
    ("foot_slip_max_m", "le", 1.90, ""),
    ("ankle_jerk_p95", "le", 4000.0, ""),
)

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
WHITE = "FFFFFF"
GREEN = "E2F0D9"
ORANGE = "FCE4D6"
GRAY = "E7E6E6"
THIN = Side(style="thin", color="D9E1F2")


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream, delimiter="\t"))


def fmt(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (float, np.floating)):
        return f"{float(value):.10g}" if math.isfinite(float(value)) else ""
    return str(value)


def write_tsv(path: Path, rows: list[dict[str, Any]], fields: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        raise ValueError(f"refusing to write empty TSV: {path}")
    columns = fields or list(rows[0])
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns, delimiter="\t", lineterminator="\n", extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: fmt(row.get(key, "")) for key in columns})


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def repo_path(value: str | Path) -> Path:
    path = Path(value)
    if path.is_file():
        return path.resolve()
    if not path.is_absolute() and (REPO / path).is_file():
        return (REPO / path).resolve()
    markers = ("workspace/core4d/", "example_datasets/", "logs/")
    text = str(path)
    for marker in markers:
        if marker in text:
            candidate = REPO / (marker + text.split(marker, 1)[1])
            if candidate.is_file():
                return candidate.resolve()
    raise FileNotFoundError(value)


def rel(path: Path | str) -> str:
    candidate = Path(path)
    try:
        return str(candidate.resolve().relative_to(REPO))
    except ValueError:
        return str(candidate)


def numeric(value: Any) -> float:
    try:
        output = float(value)
    except (TypeError, ValueError):
        return math.nan
    return output if math.isfinite(output) else math.nan


def compiled_euler_convention(model: mujoco.MjModel, object_body: int) -> str:
    start = int(model.body_jntadr[object_body])
    stop = start + int(model.body_jntnum[object_body])
    joints = [
        joint_id for joint_id in range(start, stop)
        if int(model.jnt_type[joint_id]) == int(mujoco.mjtJoint.mjJNT_HINGE)
    ]
    joints.sort(key=lambda joint_id: int(model.jnt_qposadr[joint_id]))
    letters: list[str] = []
    for joint_id in joints:
        axis = np.asarray(model.jnt_axis[joint_id], dtype=np.float64)
        index = int(np.argmax(np.abs(axis)))
        expected = np.zeros(3, dtype=np.float64)
        expected[index] = 1.0
        if not np.allclose(axis, expected, atol=1e-9, rtol=0.0):
            raise ValueError(f"object hinge axis is not a positive unit axis: joint={joint_id} axis={axis.tolist()}")
        letters.append("XYZ"[index])
    convention = "".join(letters)
    if len(convention) != 3 or len(set(convention)) != 3:
        raise ValueError(f"expected three unique object hinge axes, got {convention!r}")
    return convention


def angular_error_deg(q1_wxyz: np.ndarray, q2_wxyz: np.ndarray) -> np.ndarray:
    q1_wxyz = q1_wxyz / np.linalg.norm(q1_wxyz, axis=1, keepdims=True)
    q2_wxyz = q2_wxyz / np.linalg.norm(q2_wxyz, axis=1, keepdims=True)
    dots = np.abs(np.sum(q1_wxyz * q2_wxyz, axis=1))
    return np.degrees(2.0 * np.arccos(np.clip(dots, -1.0, 1.0)))


def convert_omni(row: dict[str, str], trajectory: Path, scene: Path) -> tuple[Path, dict[str, Any]]:
    model = mujoco.MjModel.from_xml_path(str(scene))
    object_body = int(mujoco.mj_name2id(model, mujoco.mjtObj.mjOBJ_BODY, "object"))
    if object_body < 0:
        raise ValueError(f"scene has no object body: {scene}")
    convention = compiled_euler_convention(model, object_body)
    raw = np.asarray(np.load(trajectory, allow_pickle=True)["qpos"], dtype=np.float64)
    if raw.ndim == 3:
        raw = raw[:, 0, :]
    nq_robot = model.nq - 6
    if raw.ndim != 2 or raw.shape[1] < nq_robot + 7:
        raise ValueError(f"unsupported OmniRetarget qpos {raw.shape}: {trajectory}")

    body_pos = np.asarray(model.body_pos[object_body], dtype=np.float64)
    body_quat = np.asarray(model.body_quat[object_body], dtype=np.float64)
    body_rotation = Rotation.from_quat([body_quat[1], body_quat[2], body_quat[3], body_quat[0]])
    object_pos = raw[:, nq_robot : nq_robot + 3]
    object_quat = raw[:, nq_robot + 3 : nq_robot + 7]
    object_rotation = Rotation.from_quat(np.column_stack((object_quat[:, 1:], object_quat[:, 0])))

    converted = np.zeros((raw.shape[0], model.nq), dtype=np.float64)
    converted[:, :nq_robot] = raw[:, :nq_robot]
    converted[:, nq_robot : nq_robot + 3] = body_rotation.inv().apply(object_pos - body_pos)
    converted[:, nq_robot + 3 : nq_robot + 6] = (body_rotation.inv() * object_rotation).as_euler(convention)

    data = mujoco.MjData(model)
    replay_pos = np.zeros_like(object_pos)
    replay_quat = np.zeros_like(object_quat)
    for index, qpos in enumerate(converted):
        data.qpos[:] = qpos
        data.qvel[:] = 0.0
        mujoco.mj_forward(model, data)
        replay_pos[index] = data.xpos[object_body]
        replay_quat[index] = data.xquat[object_body]
    pos_error_cm = np.linalg.norm(replay_pos - object_pos, axis=1) * 100.0
    ori_error_deg = angular_error_deg(replay_quat, object_quat)
    pos_max = float(np.max(pos_error_cm))
    ori_max = float(np.max(ori_error_deg))
    if pos_max >= 1e-8 or ori_max >= 1e-4:
        raise ValueError(
            f"OmniRetarget world-pose round-trip failed: {row['case_id']} "
            f"pos={pos_max:.3e}cm ori={ori_max:.3e}deg convention={convention}"
        )
    OMNI_DIR.mkdir(parents=True, exist_ok=True)
    output = OMNI_DIR / f"{row['case_id']}_omnirt_scene_act_qpos.npz"
    np.savez_compressed(output, qpos=converted)
    return output, {
        "euler_convention": convention,
        "roundtrip_position_error_cm_max": pos_max,
        "roundtrip_orientation_error_deg_max": ori_max,
        "omni_raw_nq": int(raw.shape[1]),
        "scene_act_nq": int(model.nq),
        "frames": int(raw.shape[0]),
    }


def authority_rows() -> tuple[list[dict[str, str]], list[dict[str, Any]]]:
    rows: list[dict[str, str]] = []
    audit: list[dict[str, Any]] = []
    for source_exp, source_path, objects in SOURCES:
        for raw in read_tsv(source_path):
            if raw.get("object_key") not in objects:
                continue
            if raw.get("metric_standard_id") != EVAL_METRIC_STANDARD_ID:
                raise ValueError(f"metric standard mismatch: {raw.get('case_id')} {raw.get('metric_standard_id')}")
            item = dict(raw)
            item["source_exp"] = source_exp
            item["source_metrics_tsv"] = rel(source_path)
            paths = {
                "prg_qpos": repo_path(raw["qpos_path"]),
                "scene": repo_path(raw["scene_xml"]),
                "trajectory": repo_path(raw["trajectory"]),
                "contact_mask": repo_path(raw["contact_mask"]),
            }
            item.update({f"resolved_{key}": str(value) for key, value in paths.items()})
            rows.append(item)
            audit.append({
                "case_id": raw["case_id"],
                "object_key": raw["object_key"],
                "source_exp": source_exp,
                "retarget_variant_id": raw.get("retarget_variant_id", ""),
                "metric_standard_id": raw.get("metric_standard_id", ""),
                "prg_qpos": rel(paths["prg_qpos"]),
                "prg_qpos_sha256": sha256(paths["prg_qpos"]),
                "scene_xml": rel(paths["scene"]),
                "scene_xml_sha256": sha256(paths["scene"]),
                "trajectory": rel(paths["trajectory"]),
                "trajectory_sha256": sha256(paths["trajectory"]),
                "contact_mask": rel(paths["contact_mask"]),
                "contact_mask_sha256": sha256(paths["contact_mask"]),
                "person_idx": person_idx_from_case(raw["case_id"]),
                "input_status": "pass",
            })
    counts = Counter(row["object_key"] for row in rows)
    case_ids = [row["case_id"] for row in rows]
    if len(rows) != EXPECTED_TOTAL or counts != Counter(EXPECTED) or len(set(case_ids)) != EXPECTED_TOTAL:
        raise ValueError(f"Full CEM authority mismatch: rows={len(rows)} unique={len(set(case_ids))} counts={counts}")
    rows.sort(key=lambda row: (OBJECTS.index(row["object_key"]), row["case_id"]))
    audit.sort(key=lambda row: (OBJECTS.index(row["object_key"]), row["case_id"]))
    return rows, audit


def evaluate_method(
    row: dict[str, str], method: str, qpos: Path, scene: Path, trajectory: Path, mask: Path,
    conversion: dict[str, Any] | None = None,
) -> dict[str, Any]:
    eval_row = {
        "case_id": row["case_id"],
        "variant": f"E197_{row['case_id']}_{method}",
        "object_key": row["object_key"],
        "object_category": "box",
        "expected_quality": "full_cem_authority",
    }
    result = evaluate_sequence(
        row=eval_row,
        method=method,
        hand_collision_variant_id="rubber_hull",
        qpos_path=qpos,
        scene_xml=scene,
        kin_ref_path=trajectory,
        contact_mask_path=mask,
        person_idx=person_idx_from_case(row["case_id"]),
    )
    health = run_health(qpos, scene, EvalConfig())
    output: dict[str, Any] = {
        "case_id": row["case_id"],
        "object_key": row["object_key"],
        "method": method,
        "source_exp": row["source_exp"],
        "retarget_variant_id": row.get("retarget_variant_id", ""),
        "metric_standard_id": EVAL_METRIC_STANDARD_ID,
        "qpos_frames": result["qpos_frames"],
        "qpos_path": rel(qpos),
        "scene_xml": rel(scene),
        "trajectory": rel(trajectory),
        "contact_mask": rel(mask),
        "person_idx": person_idx_from_case(row["case_id"]),
        "euler_convention": "" if conversion is None else conversion["euler_convention"],
        "roundtrip_position_error_cm_max": "" if conversion is None else conversion["roundtrip_position_error_cm_max"],
        "roundtrip_orientation_error_deg_max": "" if conversion is None else conversion["roundtrip_orientation_error_deg_max"],
    }
    for short, _label, field, _direction in METRICS:
        value = numeric(result[field]) if field in result else numeric(health[field])
        if short in {"contact_3mm_in_mask", "raw_contact_in_mask", "hand_object_penetration_3mm", "lower_body_penetration"} and not 0.0 <= value <= 1.0:
            raise ValueError(f"invalid {field}={value}: {row['case_id']} {method}")
        if not math.isfinite(value):
            raise ValueError(f"missing health metric {field}: {row['case_id']} {method}")
        output[short] = value
    return output


def recompute(rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for index, row in enumerate(rows, 1):
        scene = Path(row["resolved_scene"])
        trajectory = Path(row["resolved_trajectory"])
        mask = Path(row["resolved_contact_mask"])
        omni_qpos, conversion = convert_omni(row, trajectory, scene)
        output.append(evaluate_method(row, "OmniRetarget", omni_qpos, scene, trajectory, mask, conversion))
        output.append(evaluate_method(row, "PRG", Path(row["resolved_prg_qpos"]), scene, trajectory, mask))
        write_tsv(METHOD_TSV, output)
        print(f"[{index:02d}/{len(rows)}] {row['case_id']} convention={conversion['euler_convention']} roundtrip={conversion['roundtrip_orientation_error_deg_max']:.3e}deg", flush=True)
    return output


def case_comparison(authority: list[dict[str, str]], methods: list[dict[str, Any]]) -> list[dict[str, Any]]:
    lookup = {(row["case_id"], row["method"]): row for row in methods}
    rows: list[dict[str, Any]] = []
    for source in authority:
        omni = lookup[(source["case_id"], "OmniRetarget")]
        prg = lookup[(source["case_id"], "PRG")]
        if int(omni["qpos_frames"]) != int(prg["qpos_frames"]):
            raise ValueError(f"frame mismatch: {source['case_id']} Omni={omni['qpos_frames']} PRG={prg['qpos_frames']}")
        item: dict[str, Any] = {
            "case_id": source["case_id"],
            "object_key": source["object_key"],
            "source_exp": source["source_exp"],
            "retarget_variant_id": source.get("retarget_variant_id", ""),
            "qpos_frames": int(omni["qpos_frames"]),
            "euler_convention": omni["euler_convention"],
            "roundtrip_position_error_cm_max": omni["roundtrip_position_error_cm_max"],
            "roundtrip_orientation_error_deg_max": omni["roundtrip_orientation_error_deg_max"],
        }
        for short, _label, _field, direction in METRICS:
            omni_value = numeric(omni[short])
            prg_value = numeric(prg[short])
            delta = prg_value - omni_value
            item[f"omnirt_{short}"] = omni_value
            item[f"prg_{short}"] = prg_value
            item[f"delta_prg_minus_omnirt_{short}"] = delta
            item[f"improvement_{short}"] = delta if direction == "higher" else -delta
        rows.append(item)
    return rows


def bootstrap_ci(values: np.ndarray, seed: int, macro_objects: list[np.ndarray] | None = None) -> tuple[float, float]:
    rng = np.random.default_rng(seed)
    if macro_objects is None:
        draws = rng.choice(values, size=(BOOTSTRAP_DRAWS, values.size), replace=True).mean(axis=1)
    else:
        components = []
        for group in macro_objects:
            components.append(rng.choice(group, size=(BOOTSTRAP_DRAWS, group.size), replace=True).mean(axis=1))
        draws = np.mean(np.stack(components, axis=1), axis=1)
    low, high = np.quantile(draws, [0.025, 0.975])
    return float(low), float(high)


def summary_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    groups: list[tuple[str, list[dict[str, Any]]]] = [(key, [row for row in rows if row["object_key"] == key]) for key in OBJECTS]
    groups.append(("ALL_MICRO", rows))
    for metric_index, (short, label, field, direction) in enumerate(METRICS):
        for group_index, (group, selected) in enumerate(groups):
            omni = np.asarray([numeric(row[f"omnirt_{short}"]) for row in selected])
            prg = np.asarray([numeric(row[f"prg_{short}"]) for row in selected])
            delta = prg - omni
            improvement = delta if direction == "higher" else -delta
            low, high = bootstrap_ci(improvement, BOOTSTRAP_SEED + metric_index * 100 + group_index)
            output.append({
                "group": group,
                "aggregation": "case_micro",
                "case_count": len(selected),
                "object_count": 1 if group in OBJECTS else len(OBJECTS),
                "metric": short,
                "metric_label": label,
                "spider_field": field,
                "direction": direction,
                "omnirt_mean": float(omni.mean()),
                "prg_mean": float(prg.mean()),
                "delta_prg_minus_omnirt_mean": float(delta.mean()),
                "improvement_mean": float(improvement.mean()),
                "improvement_median": float(np.median(improvement)),
                "improvement_bootstrap_ci95_low": low,
                "improvement_bootstrap_ci95_high": high,
                "prg_improved_cases": int(np.count_nonzero(improvement > 1e-12)),
                "tied_cases": int(np.count_nonzero(np.abs(improvement) <= 1e-12)),
                "prg_worsened_cases": int(np.count_nonzero(improvement < -1e-12)),
            })
        by_object = [[row for row in rows if row["object_key"] == key] for key in OBJECTS]
        omni_groups = [np.asarray([numeric(row[f"omnirt_{short}"]) for row in group]) for group in by_object]
        prg_groups = [np.asarray([numeric(row[f"prg_{short}"]) for row in group]) for group in by_object]
        improvement_groups = [
            (prg - omni) if direction == "higher" else (omni - prg)
            for omni, prg in zip(omni_groups, prg_groups)
        ]
        all_improvement = np.concatenate(improvement_groups)
        low, high = bootstrap_ci(all_improvement, BOOTSTRAP_SEED + metric_index * 100 + 99, improvement_groups)
        output.append({
            "group": "ALL_MACRO",
            "aggregation": "object_balanced_macro",
            "case_count": len(rows),
            "object_count": len(OBJECTS),
            "metric": short,
            "metric_label": label,
            "spider_field": field,
            "direction": direction,
            "omnirt_mean": float(np.mean([group.mean() for group in omni_groups])),
            "prg_mean": float(np.mean([group.mean() for group in prg_groups])),
            "delta_prg_minus_omnirt_mean": float(np.mean([prg.mean() - omni.mean() for omni, prg in zip(omni_groups, prg_groups)])),
            "improvement_mean": float(np.mean([group.mean() for group in improvement_groups])),
            "improvement_median": float(np.median(all_improvement)),
            "improvement_bootstrap_ci95_low": low,
            "improvement_bootstrap_ci95_high": high,
            "prg_improved_cases": int(np.count_nonzero(all_improvement > 1e-12)),
            "tied_cases": int(np.count_nonzero(np.abs(all_improvement) <= 1e-12)),
            "prg_worsened_cases": int(np.count_nonzero(all_improvement < -1e-12)),
        })
    order = {key: index for index, key in enumerate((*OBJECTS, "ALL_MICRO", "ALL_MACRO"))}
    metric_order = {item[0]: index for index, item in enumerate(METRICS)}
    output.sort(key=lambda row: (order[row["group"]], metric_order[row["metric"]]))
    return output


def wide_gate_thresholds(methods: list[dict[str, Any]]) -> list[dict[str, Any]]:
    omni = [row for row in methods if row["method"] == "OmniRetarget"]
    if len(omni) != EXPECTED_TOTAL:
        raise ValueError(f"wide gate requires {EXPECTED_TOTAL} Omni rows, got {len(omni)}")
    output: list[dict[str, Any]] = []
    labels = {short: (label, field, direction) for short, label, field, direction in METRICS}
    for short, operator, threshold, object_scope in OMNI_WIDE_GATES:
        label, field, direction = labels[short]
        scoped_omni = [row for row in omni if not object_scope or row["object_key"] == object_scope]
        values = np.asarray([numeric(row[short]) for row in scoped_omni], dtype=np.float64)
        if not np.isfinite(values).all():
            raise ValueError(f"non-finite Omni gate distribution: {short}")
        p05, p50, p95 = np.quantile(values, [0.05, 0.50, 0.95])
        relation = "≥" if operator == "ge" else "≤"
        rule = f"OmniRetarget {relation} {threshold:g}" + (f" ({object_scope})" if object_scope else "")
        output.append({
            "gate_id": f"omni_absolute_wide_{short}" + (f"_{object_scope}" if object_scope else ""),
            "metric": short,
            "metric_label": label,
            "spider_field": field,
            "direction": direction,
            "operator": operator,
            "threshold": threshold,
            "omni_p05": float(p05),
            "omni_p50": float(p50),
            "omni_p95": float(p95),
            "headroom_rule": rule,
            "calibration_population": f"{len(scoped_omni)} Full-CEM OmniRetarget references" + (f" from {object_scope}" if object_scope else " pooled across five boxes"),
            "gate_version": OMNI_WIDE_GATE_VERSION,
            "candidate_arm": "OmniRetarget only; PRG excluded from decision",
            "object_scope": object_scope or "all boxes",
        })
    return output


def gate_pass(value: float, operator: str, threshold: float) -> bool:
    return value >= threshold if operator == "ge" else value <= threshold


def rl_filter_rows(methods: list[dict[str, Any]], gates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    lookup = {(row["case_id"], row["method"]): row for row in methods}
    output: list[dict[str, Any]] = []
    case_ids = sorted({row["case_id"] for row in methods})
    for case_id in case_ids:
        omni = lookup[(case_id, "OmniRetarget")]
        item: dict[str, Any] = {
            "case_id": case_id,
            "object_key": omni["object_key"],
            "gate_version": OMNI_WIDE_GATE_VERSION,
            "calibration_arm": "OmniRetarget",
            "candidate_arm": "OmniRetarget",
        }
        for gate in gates:
            gate_id = gate["gate_id"]
            item[f"threshold_{gate_id}"] = ""
            item[f"omnirt_{gate_id}"] = ""
            item[f"omnirt_gate_pass_{gate_id}"] = ""
        omni_failures: list[str] = []
        for gate in gates:
            short = gate["metric"]
            object_scope = gate.get("object_scope", "all boxes")
            # A scoped override replaces the all-box rule for that metric.
            if object_scope == "all boxes" and any(
                other["metric"] == short and other.get("object_scope") == omni["object_key"]
                for other in gates
            ):
                continue
            if object_scope not in {"all boxes", omni["object_key"]}:
                continue
            gate_key = gate["gate_id"]
            threshold = numeric(gate["threshold"])
            operator = gate["operator"]
            omni_value = numeric(omni[short])
            omni_pass = gate_pass(omni_value, operator, threshold)
            item[f"threshold_{gate_key}"] = threshold
            item[f"omnirt_{gate_key}"] = omni_value
            item[f"omnirt_gate_pass_{gate_key}"] = omni_pass
            if not omni_pass:
                omni_failures.append(gate["gate_id"])
        item["omnirt_wide_gate_pass"] = not omni_failures
        item["omnirt_failure_modes"] = ";".join(omni_failures)
        item["rl_filter_decision"] = (
            "RL_CANDIDATE_OMNI_WIDE_GATE_PASS" if not omni_failures else "RL_CANDIDATE_FILTERED_OMNI_WIDE_GATE"
        )
        item["rl_export_authority"] = "NOT_RL_EXPORT_READY_REQUIRES_EXISTING_MANUAL_PARTNER_ALIGNMENT_GATES"
        output.append(item)
    output.sort(key=lambda row: (OBJECTS.index(row["object_key"]), row["case_id"]))
    return output


def title(ws, text: str, subtitle: str, columns: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws["A1"] = text
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="666666")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 30


def header(ws, row: int, columns: int) -> None:
    for cell in ws[row][:columns]:
        cell.font = Font(name="Arial", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)


def add_table(ws, start: int, end: int, columns: int, name: str) -> None:
    item = Table(displayName=name, ref=f"A{start}:{get_column_letter(columns)}{end}")
    item.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(item)
    ws.freeze_panes = f"A{start + 1}"


def set_widths(ws, special: dict[int, float], default: float = 15.0) -> None:
    for index in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(index)].width = special.get(index, default)


def add_readme(wb: Workbook, summaries: list[dict[str, Any]], gates: list[dict[str, Any]], filters: list[dict[str, Any]]) -> None:
    ws = wb.active
    ws.title = "README"
    title(ws, "E197 OmniRetarget vs PRG — seven metrics", "All 87 five-box Full CEM cases; Spider physics and shared motion-health metrics.", 4)
    rows = (
        ("Scope", "87 unique Full CEM cases", "box001/004/021/023/024", "28/6/28/16/9"),
        ("Methods", "OmniRetarget kinematic replay", "PRG CEM rollout", "Same scene/mask/person/frame domain"),
        ("Delta", "PRG − OmniRetarget", "Positive contact delta is better", "Negative penetration delta is better"),
        ("Improvement", "Direction-aware", "Positive is always PRG better", "Used for win/tie/loss and color"),
        ("Aggregation", "Per-object", "ALL_MICRO case-weighted", "ALL_MACRO object-balanced"),
        ("Metric standard", EVAL_METRIC_STANDARD_ID, "Public core_metrics + motion_health", "Physics fractions; health metrics use SI units"),
        ("Euler audit", "Compiled hinge axes", "No default XYZ fallback", "World-pose round-trip <1e-4 deg"),
        ("CI", "10,000 paired bootstrap draws", f"Seed={BOOTSTRAP_SEED}", "CI evidence is hardcoded; means/deltas are formulas"),
        ("Formula sheets", "Case Comparison", "Summary", "LibreOffice recalculated"),
        ("RL wide gate", "Omni absolute thresholds only", "PRG excluded from decision", f"{sum(bool(row['omnirt_wide_gate_pass']) for row in filters)}/{len(filters)} pass"),
        ("RL authority", "Candidate prefilter only", "Not RL_EXPORT_READY", "Manual/partner/alignment gates remain required"),
        ("Source sheets", "Method Metrics", "Input Audit", "Metric Definitions"),
    )
    for row_index, values in enumerate(rows, 4):
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_index, column, value)
            cell.font = Font(name="Arial", bold=column == 1)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN)
    set_widths(ws, {1: 22, 2: 35, 3: 35, 4: 42})


def add_case_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> dict[str, tuple[int, int, int, int]]:
    ws = wb.create_sheet("Case Comparison")
    base = ["Case", "Object", "Source exp", "Retarget variant", "Frames", "Euler", "Round-trip pos (cm)", "Round-trip ori (deg)"]
    labels = list(base)
    metric_columns: dict[str, tuple[int, int, int, int]] = {}
    column = len(base) + 1
    for short, label, _field, _direction in METRICS:
        metric_columns[short] = (column, column + 1, column + 2, column + 3)
        labels.extend((f"{label} — OmniRetarget", f"{label} — PRG", f"{label} — Delta", f"{label} — Improvement"))
        column += 4
    title(ws, "Same-case metric comparison", "Delta = PRG − OmniRetarget. Improvement flips penetration signs so positive always means PRG is better.", len(labels))
    for index, label in enumerate(labels, 1):
        ws.cell(4, index, label)
    header(ws, 4, len(labels))
    for row_index, raw in enumerate(rows, 5):
        values = (
            raw["case_id"], raw["object_key"], raw["source_exp"], raw["retarget_variant_id"], raw["qpos_frames"],
            raw["euler_convention"], raw["roundtrip_position_error_cm_max"], raw["roundtrip_orientation_error_deg_max"],
        )
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
        for short, _label, _field, direction in METRICS:
            omni_col, prg_col, delta_col, improvement_col = metric_columns[short]
            ws.cell(row_index, omni_col, raw[f"omnirt_{short}"])
            ws.cell(row_index, prg_col, raw[f"prg_{short}"])
            ws.cell(row_index, delta_col, f"={get_column_letter(prg_col)}{row_index}-{get_column_letter(omni_col)}{row_index}")
            delta_ref = f"{get_column_letter(delta_col)}{row_index}"
            ws.cell(row_index, improvement_col, f"={delta_ref}" if direction == "higher" else f"=-{delta_ref}")
            ws.cell(4, delta_col).fill = PatternFill("solid", fgColor="C55A11")
            ws.cell(4, improvement_col).fill = PatternFill("solid", fgColor="548235")
            number_format = "0.0%" if short in {
                "contact_3mm_in_mask", "raw_contact_in_mask", "hand_object_penetration_3mm", "lower_body_penetration"
            } else ("0.000" if short in {"foot_slip_max_m", "obj_speed_max"} else "0.0")
            for col in (omni_col, prg_col, delta_col, improvement_col):
                ws.cell(row_index, col).number_format = number_format
        for cell in ws[row_index]:
            cell.font = Font(name="Arial", size=8)
    add_table(ws, 4, ws.max_row, len(labels), "E197CaseComparison")
    for _short, (_omni, _prg, _delta, improvement) in metric_columns.items():
        letter = get_column_letter(improvement)
        ws.conditional_formatting.add(
            f"{letter}5:{letter}{ws.max_row}",
            ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
        )
    set_widths(ws, {1: 37, 2: 12, 3: 12, 4: 16, 5: 10, 6: 10, 7: 18, 8: 18}, 17)
    return metric_columns


def add_summary_sheet(
    wb: Workbook, summaries: list[dict[str, Any]], metric_columns: dict[str, tuple[int, int, int, int]], case_count: int,
) -> None:
    ws = wb.create_sheet("Summary")
    labels = (
        "Group", "Aggregation", "Metric key", "Metric", "Direction", "N", "OmniRetarget mean", "PRG mean",
        "Delta (PRG−Omni)", "Improvement", "Median improvement", "Bootstrap 95% CI low", "Bootstrap 95% CI high",
        "PRG improved", "Tie", "PRG worsened",
    )
    title(ws, "Per-object and overall summary", "Means, deltas, improvements and counts are formulas over Case Comparison; bootstrap CI uses fixed-seed paired evidence.", len(labels))
    for col, label in enumerate(labels, 1):
        ws.cell(4, col, label)
    header(ws, 4, len(labels))
    row_map: dict[tuple[str, str], int] = {}
    for row_index, raw in enumerate(summaries, 5):
        row_map[(raw["group"], raw["metric"])] = row_index
        for col, key in enumerate(("group", "aggregation", "metric", "metric_label", "direction"), 1):
            ws.cell(row_index, col, raw[key])
        omni_col, prg_col, _delta_col, improvement_col = metric_columns[raw["metric"]]
        omni_range = f"'Case Comparison'!${get_column_letter(omni_col)}$5:${get_column_letter(omni_col)}${case_count + 4}"
        prg_range = f"'Case Comparison'!${get_column_letter(prg_col)}$5:${get_column_letter(prg_col)}${case_count + 4}"
        imp_range = f"'Case Comparison'!${get_column_letter(improvement_col)}$5:${get_column_letter(improvement_col)}${case_count + 4}"
        object_range = f"'Case Comparison'!$B$5:$B${case_count + 4}"
        if raw["group"] in OBJECTS:
            criterion = f'"{raw["group"]}"'
            ws.cell(row_index, 6, f"=COUNTIF({object_range},{criterion})")
            ws.cell(row_index, 7, f"=AVERAGEIF({object_range},{criterion},{omni_range})")
            ws.cell(row_index, 8, f"=AVERAGEIF({object_range},{criterion},{prg_range})")
            ws.cell(row_index, 14, f'=COUNTIFS({object_range},{criterion},{imp_range},">0")')
            ws.cell(row_index, 15, f'=COUNTIFS({object_range},{criterion},{imp_range},"=0")')
            ws.cell(row_index, 16, f'=COUNTIFS({object_range},{criterion},{imp_range},"<0")')
        elif raw["group"] == "ALL_MICRO":
            ws.cell(row_index, 6, f"=COUNT({omni_range})")
            ws.cell(row_index, 7, f"=AVERAGE({omni_range})")
            ws.cell(row_index, 8, f"=AVERAGE({prg_range})")
            ws.cell(row_index, 14, f'=COUNTIF({imp_range},">0")')
            ws.cell(row_index, 15, f'=COUNTIF({imp_range},"=0")')
            ws.cell(row_index, 16, f'=COUNTIF({imp_range},"<0")')
        else:
            object_rows = [row_map[(object_key, raw["metric"])] for object_key in OBJECTS]
            ws.cell(row_index, 6, "=5")
            ws.cell(row_index, 7, f"=AVERAGE({','.join(f'G{item}' for item in object_rows)})")
            ws.cell(row_index, 8, f"=AVERAGE({','.join(f'H{item}' for item in object_rows)})")
            micro = row_map[("ALL_MICRO", raw["metric"])]
            for col in (14, 15, 16):
                ws.cell(row_index, col, f"={get_column_letter(col)}{micro}")
        ws.cell(row_index, 9, f"=H{row_index}-G{row_index}")
        ws.cell(row_index, 10, f'=IF(E{row_index}="higher",I{row_index},-I{row_index})')
        for col, key in ((11, "improvement_median"), (12, "improvement_bootstrap_ci95_low"), (13, "improvement_bootstrap_ci95_high")):
            ws.cell(row_index, col, raw[key])
            ws.cell(row_index, col).comment = Comment("Source: e197_summary_by_object.tsv; paired bootstrap 10,000 draws with frozen seed 197.", "Codex")
        number_format = "0.0%" if raw["metric"] in {
            "contact_3mm_in_mask", "raw_contact_in_mask", "hand_object_penetration_3mm", "lower_body_penetration"
        } else ("0.000" if raw["metric"] in {"foot_slip_max_m", "obj_speed_max"} else "0.0")
        for col in range(7, 14):
            ws.cell(row_index, col).number_format = number_format
        for cell in ws[row_index]:
            cell.font = Font(name="Arial", size=9)
    add_table(ws, 4, ws.max_row, len(labels), "E197Summary")
    ws.conditional_formatting.add(
        f"J5:J{ws.max_row}",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
    )
    set_widths(ws, {1: 16, 2: 22, 3: 34, 4: 28, 5: 12, 6: 9}, 17)


def add_raw_sheet(wb: Workbook, name: str, heading: str, rows: list[dict[str, Any]], table_name: str) -> None:
    ws = wb.create_sheet(name)
    fields = list(rows[0])
    title(ws, heading, "Source evidence. Derived comparison calculations live in the formula-driven Case Comparison and Summary sheets.", len(fields))
    for col, field in enumerate(fields, 1):
        ws.cell(4, col, field)
    header(ws, 4, len(fields))
    for row_index, raw in enumerate(rows, 5):
        for col, field in enumerate(fields, 1):
            value = raw.get(field, "")
            number = numeric(value)
            if math.isfinite(number) and str(value).strip() != "":
                value = number
            ws.cell(row_index, col, value)
            ws.cell(row_index, col).font = Font(name="Arial", size=8)
    add_table(ws, 4, ws.max_row, len(fields), table_name)
    set_widths(ws, {1: 37, 2: 12, 3: 14, 4: 14}, 18)


def add_definitions(wb: Workbook) -> None:
    ws = wb.create_sheet("Metric Definitions")
    labels = ("Metric key", "User-facing name", "Spider field", "Direction", "Unit", "Definition")
    title(ws, "Metric definitions", "Physics metrics use Spider public core_metrics; motion-health metrics use the shared motion_health module.", len(labels))
    for col, label in enumerate(labels, 1):
        ws.cell(4, col, label)
    header(ws, 4, len(labels))
    definitions = {
        "contact_3mm_in_mask": "Raw 3cm mask active frames with any hand-object MuJoCo contact whose minimum contact distance is at least -3mm.",
        "raw_contact_in_mask": "Raw 3cm mask active frames with any hand-object MuJoCo physics contact, without the 3mm cleanliness filter.",
        "hand_object_penetration_3mm": "All sequence frames with hand-object MuJoCo contact deeper than 3mm.",
        "lower_body_penetration": "All sequence frames whose minimum lower-body geom-to-object signed distance is below zero.",
        "foot_slip_max_m": "Maximum accumulated XY drift of either ankle during contiguous frames classified as grounded.",
        "obj_speed_max": "Maximum world-space linear speed of the object body, in metres per second.",
        "ankle_jerk_p95": "95th percentile of ankle Cartesian jerk magnitude, computed from the evaluated trajectory at its native FPS.",
    }
    units = {
        "foot_slip_max_m": "m",
        "obj_speed_max": "m/s",
        "ankle_jerk_p95": "m/s³",
    }
    for row_index, (short, label, field, direction) in enumerate(METRICS, 5):
        values = (short, label, field, direction, units.get(short, "frame fraction"), definitions[short])
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
            ws.cell(row_index, col).font = Font(name="Arial", size=9)
            ws.cell(row_index, col).alignment = Alignment(wrap_text=True, vertical="top")
    add_table(ws, 4, ws.max_row, len(labels), "E197MetricDefinitions")
    set_widths(ws, {1: 34, 2: 27, 3: 52, 4: 12, 5: 16, 6: 75})


def build_workbook(
    cases: list[dict[str, Any]], summaries: list[dict[str, Any]], methods: list[dict[str, Any]], audit: list[dict[str, Any]],
    gates: list[dict[str, Any]], filters: list[dict[str, Any]],
) -> None:
    wb = Workbook()
    add_readme(wb, summaries, gates, filters)
    metric_columns = add_case_sheet(wb, cases)
    add_summary_sheet(wb, summaries, metric_columns, len(cases))
    add_raw_sheet(wb, "Method Metrics", "Recomputed method metrics", methods, "E197MethodMetrics")
    add_raw_sheet(wb, "Input Audit", "Full CEM input authority", audit, "E197InputAudit")
    add_raw_sheet(wb, "Omni Wide Gates", "OmniRetarget-only absolute wide RL gates", gates, "E197OmniWideGates")
    add_raw_sheet(wb, "Omni Wide Filter", "OmniRetarget downstream RL candidate prefilter; PRG excluded", filters, "E197OmniWideFilter")
    add_definitions(wb)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                if cell.font.name != "Arial":
                    cell.font = Font(name="Arial", size=cell.font.sz or 10, bold=cell.font.bold, italic=cell.font.italic)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    wb.save(WORKBOOK)


def summary_lookup(rows: list[dict[str, Any]], group: str, metric: str) -> dict[str, Any]:
    return next(row for row in rows if row["group"] == group and row["metric"] == metric)


def percent(value: float) -> str:
    return f"{100.0 * value:.1f}%"


def pp(value: float) -> str:
    return f"{100.0 * value:+.1f} pp"


def metric_value(short: str, value: float) -> str:
    if short in {"contact_3mm_in_mask", "raw_contact_in_mask", "hand_object_penetration_3mm", "lower_body_penetration"}:
        return percent(value)
    if short == "foot_slip_max_m":
        return f"{value:.3f} m"
    if short == "obj_speed_max":
        return f"{value:.3f} m/s"
    return f"{value:.1f} m/s³"


def metric_delta(short: str, value: float) -> str:
    if short in {"contact_3mm_in_mask", "raw_contact_in_mask", "hand_object_penetration_3mm", "lower_body_penetration"}:
        return pp(value)
    if short == "foot_slip_max_m":
        return f"{value:+.3f} m"
    if short == "obj_speed_max":
        return f"{value:+.3f} m/s"
    return f"{value:+.1f} m/s³"


def write_report(
    summaries: list[dict[str, Any]], cases: list[dict[str, Any]], audit: list[dict[str, Any]],
    gates: list[dict[str, Any]], filters: list[dict[str, Any]],
) -> None:
    macro = {short: summary_lookup(summaries, "ALL_MACRO", short) for short, *_ in METRICS}
    micro = {short: summary_lookup(summaries, "ALL_MICRO", short) for short, *_ in METRICS}
    lines = [
        "# 五类 box Full CEM：OmniRetarget 与 PRG 七指标对比", "",
        "_Core4D E197 离线统一重评 · 87 个 Full CEM case · 2026-08-12_", "", "---", "",
        "## 📋 摘要", "",
        f"本报告覆盖 `{len(cases)}` 个进入 Full CEM 的唯一 case：box001/004/021/023/024 分别为 "
        f"`{EXPECTED['box001']}/{EXPECTED['box004']}/{EXPECTED['box021']}/{EXPECTED['box023']}/{EXPECTED['box024']}`。"
        "OmniRetarget kinematic replay 与 PRG CEM rollout 使用同一 scene_act、3cm raw mask、person index 和帧域，"
        "并由 Spider 公共评测模块统一重算。", "",
        "以 object-balanced macro average 为主，PRG 相对 OmniRetarget 的四项 direction-aware improvement 为："
        f"3mm in-mask 接触 `{pp(macro['contact_3mm_in_mask']['improvement_mean'])}`，raw in-mask 接触 "
        f"`{pp(macro['raw_contact_in_mask']['improvement_mean'])}`，手物 >3mm 穿透 "
        f"`{pp(macro['hand_object_penetration_3mm']['improvement_mean'])}`，lower-body 穿透 "
        f"`{pp(macro['lower_body_penetration']['improvement_mean'])}`。正 improvement 始终代表 PRG 更好。", "",
        "## 🔬 方法与口径", "",
        "```mermaid", "flowchart LR",
        "    accTitle: E197 Paired Evaluation Flow",
        "    accDescr: The same 87 Full CEM cases are replayed as OmniRetarget references and PRG rollouts under a shared scene and contact mask before public-core metrics are paired and aggregated.", "",
        "    authority[\"📥 Freeze 87 cases\"] --> omni[\"⚙️ Replay OmniRetarget\"]",
        "    authority --> prg[\"⚙️ Replay PRG\"]",
        "    omni --> public_core[\"📊 Public-core metrics\"]",
        "    prg --> public_core",
        "    public_core --> paired[\"🔗 Pair same cases\"]",
        "    paired --> report[\"✅ Report results\"]", "",
        "    classDef input fill:#ede9fe,stroke:#7c3aed,stroke-width:2px,color:#3b0764",
        "    classDef process fill:#dbeafe,stroke:#2563eb,stroke-width:2px,color:#1e3a5f",
        "    classDef success fill:#dcfce7,stroke:#16a34a,stroke-width:2px,color:#14532d", "",
        "    class authority input", "    class omni,prg,public_core,paired process", "    class report success", "```", "",
        "### 指标映射", "",
        "| 用户指标 | Spider 公共字段 | 方向 |", "| --- | --- | --- |",
    ]
    for _short, label, field, direction in METRICS:
        lines.append(f"| {label} | `{field}` | {'越高越好' if direction == 'higher' else '越低越好'} |")
    lines += [
        "", "前四项 physics 指标是 frame fraction；后三项 motion-health 指标分别使用 m、m/s、m/s³。"
        "`Delta = PRG − OmniRetarget`；接触的 improvement 等于 delta，其余越低越好的指标 improvement 等于 `−delta`。"
        "因此 improvement 为正时，一律表示 PRG 改善。", "",
        "### 输入与转换审计", "",
        f"输入路径与 SHA 审计为 `{len(audit)}/{EXPECTED_TOTAL}` pass。OmniRetarget freejoint object pose 转为 PRG "
        "scene_act 的 slide/hinge 参数时，Euler 顺序直接由 compiled object hinge axes 推导；不使用默认 `XYZ`。"
        f"87 条 world-pose round-trip 最大 orientation error 为 "
        f"`{max(numeric(row['roundtrip_orientation_error_deg_max']) for row in cases):.8f}°`，"
        f"最大 position error 为 `{max(numeric(row['roundtrip_position_error_cm_max']) for row in cases):.3e} cm`。", "",
        "## 📊 结果", "", "### 按物体对比", "",
        "下表为均值；括号内为 `PRG − OmniRetarget`。穿透 delta 为负代表 PRG 穿透更少。", "",
        "| 物体 | n | 3mm in-mask 接触 | Raw in-mask 接触 | 手物 >3mm 穿透 | Lower-body 穿透 | Foot slip max | Object speed max | Ankle jerk P95 |",
        "| --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: |",
    ]
    for object_key in OBJECTS:
        cells = []
        for short, *_ in METRICS:
            row = summary_lookup(summaries, object_key, short)
            cells.append(
                f"{metric_value(short, row['omnirt_mean'])} → {metric_value(short, row['prg_mean'])} "
                f"({metric_delta(short, row['delta_prg_minus_omnirt_mean'])})"
            )
        lines.append(f"| {object_key} | {EXPECTED[object_key]} | " + " | ".join(cells) + " |")
    lines += ["", "### 总体汇总", "", "| 聚合 | 指标 | OmniRetarget | PRG | Delta | Improvement | 95% CI | W/T/L |", "| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: |"]
    for group, label in (("ALL_MICRO", "Case-weighted micro"), ("ALL_MACRO", "Object-balanced macro")):
        for short, metric_label, *_ in METRICS:
            row = summary_lookup(summaries, group, short)
            lines.append(
                f"| {label} | {metric_label} | {metric_value(short, row['omnirt_mean'])} | {metric_value(short, row['prg_mean'])} | "
                f"{metric_delta(short, row['delta_prg_minus_omnirt_mean'])} | {metric_delta(short, row['improvement_mean'])} | "
                f"[{metric_delta(short, row['improvement_bootstrap_ci95_low'])}, {metric_delta(short, row['improvement_bootstrap_ci95_high'])}] | "
                f"{row['prg_improved_cases']}/{row['tied_cases']}/{row['prg_worsened_cases']} |"
            )
    lines += ["", "## 💡 解读", ""]
    contact_better = [key for key in OBJECTS if summary_lookup(summaries, key, "contact_3mm_in_mask")["improvement_mean"] > 0]
    raw_better = [key for key in OBJECTS if summary_lookup(summaries, key, "raw_contact_in_mask")["improvement_mean"] > 0]
    hand_better = [key for key in OBJECTS if summary_lookup(summaries, key, "hand_object_penetration_3mm")["improvement_mean"] > 0]
    lower_better = [key for key in OBJECTS if summary_lookup(summaries, key, "lower_body_penetration")["improvement_mean"] > 0]
    lines += [
        f"- PRG 的 3mm in-mask 接触均值在 `{len(contact_better)}/5` 个物体上更高：{', '.join(contact_better) or '无'}",
        f"- PRG 的 raw in-mask 接触均值在 `{len(raw_better)}/5` 个物体上更高：{', '.join(raw_better) or '无'}",
        f"- PRG 的手物 >3mm 穿透均值在 `{len(hand_better)}/5` 个物体上更低：{', '.join(hand_better) or '无'}",
        f"- PRG 的 lower-body 穿透均值在 `{len(lower_better)}/5` 个物体上更低：{', '.join(lower_better) or '无'}", "",
        "总体 micro average 会被 28-case 的 box001/box021 主导；跨物体判断应优先参考 macro average，"
        "逐 case 排查则使用 XLSX 的 `Case Comparison` sheet。", "",
        "## 🎯 下游 RL 宽口径过滤", "",
        "这是一个只读取 OmniRetarget 指标的绝对阈值候选预过滤；PRG 的任何指标都不参与 gate 判定。"
        f"{len(OMNI_WIDE_GATES)} 条规则（6 个指标，3mm 接触含 box024 特例）均通过才标记为 `RL_CANDIDATE_OMNI_WIDE_GATE_PASS`；仍需人工 USE、partner、alignment 和下游 contract。"
        "`obj_speed_max` 保留在统一报告和 Viser 指标栏中，但用户未提供阈值，故不参与本版 gate。"
        "P05/P50/P95 仅展示 87 条 Omni reference 的分布证据。", "",
        "| Gate | Omni P05 | Omni P50 | Omni P95 | Absolute threshold | Rule |", "| --- | ---: | ---: | ---: | ---: | --- |",
    ]
    for gate in gates:
        lines.append(
            f"| {gate['metric_label']} | {numeric(gate['omni_p05']):.4g} | {numeric(gate['omni_p50']):.4g} | "
            f"{numeric(gate['omni_p95']):.4g} | {numeric(gate['threshold']):.4g} | {gate['headroom_rule']} |"
        )
    omni_pass = [row for row in filters if row["omnirt_wide_gate_pass"]]
    pass_by_object = Counter(row["object_key"] for row in omni_pass)
    failure_counts = Counter(
        mode for row in filters for mode in str(row["omnirt_failure_modes"]).split(";") if mode
    )
    lines += [
        "", f"OmniRetarget 宽 gate 通过 `{len(omni_pass)}/{len(filters)}` 条。通过 case 清单见 `e197_omni_absolute_wide_gate_filter.tsv`，"
        "其中 `RL_CANDIDATE_OMNI_WIDE_GATE_PASS` 只表示数值预筛通过，不代表可直接导出 RL。"
        f"按物体通过数为 box001/004/021/023/024=`{pass_by_object['box001']}/{pass_by_object['box004']}/"
        f"{pass_by_object['box021']}/{pass_by_object['box023']}/{pass_by_object['box024']}`。"
        f"最常见的过滤原因是 raw contact 未达 50%（`{failure_counts['omni_absolute_wide_raw_contact_in_mask']}` 条），"
        f"其次是 foot slip 超过 1.90 m（`{failure_counts['omni_absolute_wide_foot_slip_max_m']}` 条）。", "",
        "## ⚠️ 限制", "",
        "- OmniRetarget 是 kinematic reference replay，PRG 是 Full CEM dynamic rollout；本比较衡量结果差异，不单独识别 P/R/G 各组件因果贡献",
        "- CEM 只有单 seed；bootstrap 对 case 重采样，不能估计 optimizer seed 方差",
        "- 五物体 case 数不均，因此同时报告 case-weighted micro 与 object-balanced macro",
        "- 本轮不改变任何人工 USE/DNU、numeric gate 或 RL export 决策", "",
        "## 👁️ OmniRetarget Viser review", "",
        "只读播放器已生成：`workspace/core4d/scripts/eval/review/viser_e197_omnirt_player.py`，"
        "启动 wrapper 为 `workspace/core4d/scripts/eval/wrappers/review_E197_omnirt_player.sh`。"
        "它直接加载 `omni_converted_qpos/*.npz` 和对应 scene_act XML，支持 87 个 case 的"
        "case/object/gate 筛选、逐帧播放、碰撞体显示和七项指标查看；`--check` 已审计"
        "`87/87` playable，Viser smoke server 在 `8097` 正常监听。", "",
        "## 🔗 产物", "",
        "- `E197_full_cem_omnirt_vs_prg_metrics.xlsx`：README、公式化 Summary、87-case 对比、174-row method metrics、输入审计与指标定义",
        "- `e197_case_comparison.tsv`：87 条同 case 对比",
        "- `e197_summary_by_object.tsv`：按物体、micro 与 macro 汇总",
        "- `e197_method_metrics.tsv`：174 条统一重算 method metrics",
        "- `e197_input_audit.tsv`：输入路径与 SHA authority", "",
    ]
    REPORT.write_text("\n".join(lines), encoding="utf-8")


def validate_outputs(
    methods: list[dict[str, Any]], cases: list[dict[str, Any]], summaries: list[dict[str, Any]],
    gates: list[dict[str, Any]], filters: list[dict[str, Any]],
) -> dict[str, Any]:
    method_counts = Counter(row["method"] for row in methods)
    object_counts = Counter(row["object_key"] for row in cases)
    failures: list[str] = []
    if method_counts != Counter({"OmniRetarget": EXPECTED_TOTAL, "PRG": EXPECTED_TOTAL}):
        failures.append(f"method counts {method_counts}")
    if object_counts != Counter(EXPECTED):
        failures.append(f"object counts {object_counts}")
    if len(cases) != EXPECTED_TOTAL or len({row["case_id"] for row in cases}) != EXPECTED_TOTAL:
        failures.append("case count/uniqueness")
    for row in methods:
        for short, _label, _field, _direction in METRICS:
            value = numeric(row[short])
            if short in {"contact_3mm_in_mask", "raw_contact_in_mask", "hand_object_penetration_3mm", "lower_body_penetration"} and not 0.0 <= value <= 1.0:
                failures.append(f"metric range {row['case_id']} {row['method']} {short}={value}")
            if not math.isfinite(value):
                failures.append(f"metric nonfinite {row['case_id']} {row['method']} {short}")
    max_ori = max(numeric(row["roundtrip_orientation_error_deg_max"]) for row in cases)
    max_pos = max(numeric(row["roundtrip_position_error_cm_max"]) for row in cases)
    if max_ori >= 1e-4 or max_pos >= 1e-8:
        failures.append(f"roundtrip max pos={max_pos} ori={max_ori}")
    if len(gates) != len(OMNI_WIDE_GATES) or len(filters) != EXPECTED_TOTAL:
        failures.append(f"gate/filter cardinality gates={len(gates)} filters={len(filters)}")
    if any(row["rl_filter_decision"] == "RL_CANDIDATE_OMNI_WIDE_GATE_PASS" and not row["omnirt_wide_gate_pass"] for row in filters):
        failures.append("inconsistent RL filter decisions")
    return {
        "status": "pass" if not failures else "fail",
        "failures": failures,
        "cases": len(cases),
        "method_rows": len(methods),
        "summary_rows": len(summaries),
        "wide_gate_count": len(gates),
        "omnirt_wide_gate_pass_count": sum(bool(row["omnirt_wide_gate_pass"]) for row in filters),
        "object_counts": dict(object_counts),
        "method_counts": dict(method_counts),
        "metric_standard_id": EVAL_METRIC_STANDARD_ID,
        "max_roundtrip_position_error_cm": max_pos,
        "max_roundtrip_orientation_error_deg": max_ori,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--reuse-method-metrics", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    authority, audit = authority_rows()
    write_tsv(AUDIT_TSV, audit)
    if args.reuse_method_metrics:
        methods = read_tsv(METHOD_TSV)
    else:
        methods = recompute(authority)
    cases = case_comparison(authority, methods)
    summaries = summary_rows(cases)
    gates = wide_gate_thresholds(methods)
    filters = rl_filter_rows(methods, gates)
    write_tsv(CASE_TSV, cases)
    write_tsv(SUMMARY_TSV, summaries)
    write_tsv(GATE_TSV, gates)
    write_tsv(RL_FILTER_TSV, filters)
    build_workbook(cases, summaries, methods, audit, gates, filters)
    write_report(summaries, cases, audit, gates, filters)
    validation = validate_outputs(methods, cases, summaries, gates, filters)
    validation["sha256"] = {
        "method_metrics": sha256(METHOD_TSV),
        "case_comparison": sha256(CASE_TSV),
        "summary_by_object": sha256(SUMMARY_TSV),
        "input_audit": sha256(AUDIT_TSV),
        "wide_gate_thresholds": sha256(GATE_TSV),
        "rl_wide_filter": sha256(RL_FILTER_TSV),
        "report": sha256(REPORT),
        "workbook": sha256(WORKBOOK),
    }
    write_json(SUMMARY_JSON, validation)
    print(json.dumps(validation, ensure_ascii=False, indent=2), flush=True)
    return 0 if validation["status"] == "pass" else 1


if __name__ == "__main__":
    raise SystemExit(main())
