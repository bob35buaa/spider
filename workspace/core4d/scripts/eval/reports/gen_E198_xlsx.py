#!/usr/bin/env python3
"""Build the E198 G1xA2 factorial results workbook (xlsx).

Reads the frozen eval TSVs under results/E198/.../eval/full_factorial/ and emits a
multi-sheet .xlsx: per-object 4-arm comparison (G1+A2 / G1 / A2 / PRG), interaction
term with 95% CI, 12-gate migration table, and the raw per-case dump.

Arm naming: PRG == A0 baseline (gravcomp off, PRG on), G1 == object gravcomp,
A2 == hand-gate 3-field override, G1+A2 == both.
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve()
sys.path.insert(0, str(HERE.parents[2] / "experiments" / "E198"))
import e198_common as C  # noqa: E402

EVAL = C.RESULTS_E198 / "s6_downstream/eval/full_factorial"
BY_OBJECT = EVAL / "e198_factorial_by_object.tsv"
BY_CASE = EVAL / "e198_factorial_by_case.tsv"
GATES = EVAL / "e198_gate_migrations.tsv"
ARM_CACHE = EVAL / "e198_arm_cache.tsv"
OUT = EVAL / "E198_G1xA2_factorial.xlsx"

# the 12 scoring gates (arm_cache col = {gate}_gate_pass), display order
GATES_12 = ["fall", "body_z", "contact", "release", "hand_penetration", "lower_body",
            "root_pos", "root_ori", "hand_pos", "hand_ori", "object_pos", "object_ori"]
OBJ_ORDER = {"box024": 0, "box021": 1, "box023": 2, "box004": 3, "box001": 4}

# object display order = experiment priority P0->P3, then box001 (plan227 supplement)
OBJECTS = [("box024", "box024 (长箱, P0)"), ("box021", "box021 (中箱, P1/P2)"),
           ("box023", "box023 (小箱, P1/P2)"), ("box004", "box004 (P3)"),
           ("box001", "box001 (plan227补充, E196修正参考)")]

# (raw metric key, 显示名, 方向 ↓越低越好 / ↑越高越好)
METRICS = [
    ("track_obj_z_abs_err_cm_mean", "物体 Z |误差| (cm)", "↓"),
    ("track_obj_pos_err_cm_mean", "物体 3D 位置误差 (cm)", "↓"),
    ("track_obj_ori_err_deg_mean", "物体朝向误差 (°)", "↓"),
    ("hand_object_physics_penetration_3mm_frame_frac", "手-物 3mm 穿透占比", "↓"),
    ("hand_object_physics_contact_3mm_in_mask_frac", "承重接触 3mm in-mask", "↑"),
    ("leg_penetration_frac", "腿穿透占比", "↓"),
    ("track_root_pos_err_cm_mean", "根位置误差 (cm)", "↓"),
    ("track_root_ori_err_deg_mean", "根朝向误差 (°)", "↓"),
    ("track_eef_pos_err_cm_mean", "末端位置误差 (cm)", "↓"),
    ("track_eef_ori_err_deg_mean", "末端朝向误差 (°)", "↓"),
    ("qpos_jerk_l2_p95", "关节 jerk p95", "↓"),
]

# column arms in the order the user asked for: G1+A2 vs G1 vs A2 vs PRG
ARM_COLS = [("G1A2", "G1+A2"), ("G1", "G1"), ("A2", "A2"), ("A0", "PRG (baseline)")]

# ---- styling ----
HDR = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="305496")
OBJ_FILL = PatternFill("solid", fgColor="D9E1F2")
BEST_FILL = PatternFill("solid", fgColor="C6EFCE")
SIG_FILL = PatternFill("solid", fgColor="FFEB9C")
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def read_tsv(p: Path) -> list[dict]:
    with open(p) as f:
        return list(csv.DictReader(f, delimiter="\t"))


def fnum(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def style_header(ws, row: int, ncol: int):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR
        cell.fill = HDR_FILL
        cell.alignment = CENTER
        cell.border = THIN


def sheet_by_object(wb, rows: list[dict]):
    ws = wb.create_sheet("四臂×物体对比")
    by_obj = {r["object_key"]: r for r in rows}
    headers = ["指标", "方向", "n", "G1+A2", "G1", "A2", "PRG",
               "INT (=G1A2−G1−A2+A0)", "INT 95%CI", "读数"]
    r = 1
    for okey, oname in OBJECTS:
        row = by_obj.get(okey)
        if not row:
            continue
        n = row.get("n", "")
        ws.cell(row=r, column=1, value=f"■ {oname}   n={n}").font = Font(bold=True, size=12)
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).fill = OBJ_FILL
        r += 1
        for c, h in enumerate(headers, 1):
            ws.cell(row=r, column=c, value=h)
        style_header(ws, r, len(headers))
        r += 1
        for key, disp, direction in METRICS:
            vals = {a: fnum(row.get(f"{key}__{a}_mean")) for a, _ in ARM_COLS}
            intv = fnum(row.get(f"{key}__INT_mean"))
            lo = fnum(row.get(f"{key}__INT_ci_lo"))
            hi = fnum(row.get(f"{key}__INT_ci_hi"))
            present = [v for v in vals.values() if v is not None]
            best = None
            if present:
                best = (min if direction == "↓" else max)(present)
            ws.cell(row=r, column=1, value=disp)
            ws.cell(row=r, column=2, value=direction).alignment = CENTER
            ws.cell(row=r, column=3, value=int(float(n)) if n else "").alignment = CENTER
            for c, (a, _) in enumerate(ARM_COLS, 4):
                v = vals[a]
                cell = ws.cell(row=r, column=c, value=None if v is None else round(v, 3))
                cell.alignment = CENTER
                if v is not None and best is not None and abs(v - best) < 1e-9:
                    cell.fill = BEST_FILL
            ic = ws.cell(row=r, column=8, value=None if intv is None else round(intv, 3))
            ic.alignment = CENTER
            ci_txt = "" if lo is None else f"[{lo:.3f}, {hi:.3f}]"
            cc = ws.cell(row=r, column=9, value=ci_txt)
            cc.alignment = CENTER
            sig = lo is not None and (lo > 0 or hi < 0)  # CI excludes 0
            if sig:
                ic.fill = SIG_FILL
                cc.fill = SIG_FILL
            # 读数: 简明方向判读
            note = ""
            if sig:
                syn = (intv < 0) == (direction == "↓")
                note = "显著协同" if syn else "显著拮抗/次可加"
            ws.cell(row=r, column=10, value=note).alignment = CENTER
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).border = THIN
            r += 1
        r += 1  # blank row between objects
    # legend
    ws.cell(row=r, column=1, value="绿=该指标最优臂  黄=INT 95%CI 不含0(显著)  "
            "INT=M(G1+A2)−M(A2)−M(G1)+M(A0), <0且指标↓越好=协同").font = Font(italic=True, size=9)
    ws.freeze_panes = "A1"
    widths = [22, 5, 5, 11, 11, 11, 13, 20, 22, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def sheet_gates(wb):
    ws = wb.create_sheet("Gate迁移(McNemar)")
    rows = read_tsv(GATES)
    if not rows:
        return
    cols = ["transition", "object_key", "gate", "n", "before_pass", "after_pass",
            "delta_pp", "p2f", "f2p", "mcnemar_p"]
    disp = ["迁移", "物体", "门", "n", "前通过", "后通过", "Δpp", "P→F", "F→P", "exact p"]
    for c, h in enumerate(disp, 1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, 1, len(disp))
    r = 2
    for row in rows:
        for c, k in enumerate(cols, 1):
            v = row.get(k, "")
            fv = fnum(v)
            cell = ws.cell(row=r, column=c, value=fv if fv is not None else v)
            cell.alignment = CENTER
            cell.border = THIN
        dpp = fnum(row.get("delta_pp"))
        p = fnum(row.get("mcnemar_p"))
        if dpp is not None and abs(dpp) >= 8:
            ws.cell(row=r, column=7).fill = BEST_FILL if dpp > 0 else PatternFill(
                "solid", fgColor="FFC7CE")
        if p is not None and p < 0.05:
            ws.cell(row=r, column=10).fill = SIG_FILL
        r += 1
    ws.freeze_panes = "A2"
    for i, w in enumerate([12, 9, 16, 5, 8, 8, 7, 6, 6, 9], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _truthy(v) -> bool:
    return str(v).strip().lower() in ("true", "1", "1.0", "yes", "pass")


def sheet_g1a2_by_case(wb):
    """Dedicated per-case sheet for the G1+A2 arm: metrics + 12-gate breakdown."""
    ws = wb.create_sheet("G1+A2逐例")
    metrics = {r["case_id"]: r for r in read_tsv(BY_CASE)}
    g1a2 = {r["case_id"]: r for r in read_tsv(ARM_CACHE) if r.get("arm") == "G1A2"}
    cases = sorted(metrics, key=lambda c: (OBJ_ORDER.get(c.split("_")[0], 9), c))

    id_cols = ["case_id", "物体", "variant"]
    ov_cols = ["12门通过", "通过门数"]
    metric_disp = [d for _, d, _ in METRICS]
    headers = id_cols + ov_cols + metric_disp + GATES_12
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, 1, len(headers))

    r = 2
    for cid in cases:
        m = metrics[cid]
        g = g1a2.get(cid, {})
        col = 1
        ws.cell(row=r, column=col, value=cid); col += 1
        ws.cell(row=r, column=col, value=cid.split("_")[0]).alignment = CENTER; col += 1
        ws.cell(row=r, column=col, value=m.get("retarget_variant_id", "")).alignment = CENTER
        col += 1
        # overall 12-gate pass
        overall = _truthy(g.get("numeric_release_pass_12gate"))
        oc = ws.cell(row=r, column=col, value="通过" if overall else "未过")
        oc.alignment = CENTER
        oc.fill = BEST_FILL if overall else PatternFill("solid", fgColor="FFC7CE")
        col += 1
        # gates passed count
        npass = sum(1 for gt in GATES_12 if _truthy(g.get(f"{gt}_gate_pass")))
        gc = ws.cell(row=r, column=col, value=f"{npass}/12")
        gc.alignment = CENTER
        if npass < 12:
            gc.fill = SIG_FILL
        col += 1
        # G1A2 metric values
        for key, _, _ in METRICS:
            v = fnum(m.get(f"{key}__G1A2"))
            cell = ws.cell(row=r, column=col, value=None if v is None else round(v, 3))
            cell.alignment = CENTER
            col += 1
        # individual gate P/F
        for gt in GATES_12:
            ok = _truthy(g.get(f"{gt}_gate_pass"))
            cell = ws.cell(row=r, column=col, value="P" if ok else "F")
            cell.alignment = CENTER
            if not ok:
                cell.fill = PatternFill("solid", fgColor="FFC7CE")
            col += 1
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).border = THIN
        r += 1

    ws.freeze_panes = "D2"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 8
    ws.column_dimensions["E"].width = 8
    for i in range(6, 6 + len(METRICS)):
        ws.column_dimensions[get_column_letter(i)].width = 13
    for i in range(6 + len(METRICS), len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 6


def sheet_box004_ex086(wb):
    """box004 four-arm comparison excluding the two 086 cases (n=4 sensitivity check)."""
    excl = {"box004_20231003_2_086_p1", "box004_20231003_2_086_p2"}
    keep = [r for r in read_tsv(BY_CASE)
            if r["object_key"] == "box004" and r["case_id"] not in excl]
    ws = wb.create_sheet("box004去086(n=4)")

    def mean(key, arm):
        vs = [fnum(r.get(f"{key}__{arm}")) for r in keep]
        vs = [v for v in vs if v is not None]
        return sum(vs) / len(vs) if vs else None

    ws.cell(row=1, column=1, value=f"box004 去除 086_p1/p2 后四臂对比  n={len(keep)}  "
            f"(保留: {', '.join(r['case_id'] for r in keep)})").font = Font(bold=True, size=11)
    for c in range(1, 8):
        ws.cell(row=1, column=c).fill = OBJ_FILL
    headers = ["指标", "方向", "G1+A2", "G1", "A2", "PRG", "INT (=G1A2−G1−A2+A0)"]
    for c, h in enumerate(headers, 1):
        ws.cell(row=2, column=c, value=h)
    style_header(ws, 2, len(headers))
    r = 3
    for key, disp, direction in METRICS:
        vals = {a: mean(key, a) for a, _ in ARM_COLS}
        present = [v for v in vals.values() if v is not None]
        best = (min if direction == "↓" else max)(present) if present else None
        intv = None
        if all(vals[a] is not None for a in ("G1A2", "G1", "A2", "A0")):
            intv = vals["G1A2"] - vals["G1"] - vals["A2"] + vals["A0"]
        ws.cell(row=r, column=1, value=disp)
        ws.cell(row=r, column=2, value=direction).alignment = CENTER
        for c, (a, _) in enumerate(ARM_COLS, 3):
            v = vals[a]
            cell = ws.cell(row=r, column=c, value=None if v is None else round(v, 3))
            cell.alignment = CENTER
            if v is not None and best is not None and abs(v - best) < 1e-9:
                cell.fill = BEST_FILL
        ic = ws.cell(row=r, column=7, value=None if intv is None else round(intv, 3))
        ic.alignment = CENTER
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).border = THIN
        r += 1
    ws.cell(row=r + 1, column=1, value="注：obj_ori/root_ori 上 A2 单用崩溃、G1+A2 救回，"
            "INT 比全集 n=6 更深（-3.51→-5.43）；见 log284 §5.5。").font = Font(italic=True, size=9)
    ws.freeze_panes = "A3"
    for i, w in enumerate([22, 5, 11, 11, 11, 11, 22], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def sheet_by_case(wb):
    ws = wb.create_sheet("逐例(by_case)")
    rows = read_tsv(BY_CASE)
    if not rows:
        return
    # keep case_id/object/variant + the 4-arm columns for the curated metrics
    keep = ["case_id", "object_key", "retarget_variant_id"]
    for key, _, _ in METRICS:
        for a in ("G1A2", "G1", "A2", "A0", "INT"):
            k = f"{key}__{a}"
            if k in rows[0]:
                keep.append(k)
    for c, k in enumerate(keep, 1):
        ws.cell(row=1, column=c, value=k)
    style_header(ws, 1, len(keep))
    for r, row in enumerate(rows, 2):
        for c, k in enumerate(keep, 1):
            fv = fnum(row.get(k))
            ws.cell(row=r, column=c,
                    value=round(fv, 3) if fv is not None else row.get(k, ""))
    ws.freeze_panes = "D2"
    ws.column_dimensions["A"].width = 22


def main() -> int:
    rows = read_tsv(BY_OBJECT)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet_by_object(wb, rows)
    sheet_g1a2_by_case(wb)
    sheet_box004_ex086(wb)
    sheet_gates(wb)
    sheet_by_case(wb)
    wb.save(OUT)
    print(f"[xlsx] wrote {C.rel(OUT)}  ({len(wb.sheetnames)} sheets: {wb.sheetnames})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
