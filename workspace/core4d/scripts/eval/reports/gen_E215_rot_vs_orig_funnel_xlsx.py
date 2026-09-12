#!/usr/bin/env python3
"""E215 rot-vs-orig comparison under the E201 14-gate wide/narrow funnel.

E215's own evaluator scores a 6-gate C4 verdict. This report re-scores the SAME
already-computed rollouts through the repo's authoritative 14-gate funnel
(4 hard + 10 banded, ``E201/funnel_config.py``), the caliber every arm-comparison
xlsx uses, so rot is judged on the identical ruler as E199/E200/E208.

No CEM/eval re-run: the funnel needs only fields the E215 evaluator already wrote
into ``e215_rot_rollout.tsv`` (102 rows = 34 orig + 34 rot0 + 34 rot1, all 34
cases paired). We classify each row inline (E215 never ran classify_funnel) and
emit the same detail/summary shape as gen_E200_arm_funnel_compare_xlsx.py, but
the two compared columns are **orig vs rot** (rot = rot0 ∪ rot1) instead of the
two reward arms.

Sheets:
  * Overview  — caliber provenance, counts, source sha256, registered exclusions.
  * Detail    — per case, orig/rot0/rot1 adjacent: 4 hard (value+PASS) + 10 banded
                (value + narrow PASS + wide PASS) + funnel layer.
  * Summary   — pooled + per arm_group + per object: layer distribution, overall
                wide/narrow accept, per-gate wide & narrow pass rate, key-metric
                means and Δ(rot − orig).

Usage:
    .venv/bin/python workspace/core4d/scripts/eval/reports/gen_E215_rot_vs_orig_funnel_xlsx.py
    ... --out <path>
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[5]
sys.path.insert(0, str(REPO / "workspace/core4d/scripts/experiments/E201"))
import funnel_config as FC  # noqa: E402

EVAL_DIR = REPO / "workspace/core4d/results/E215/s6_downstream/eval/aug"
ROLLOUT_TSV = EVAL_DIR / "e215_rot_rollout.tsv"
SUMMARY_JSON = EVAL_DIR / "e215_rot_eval_summary.json"
DEFAULT_OUT = EVAL_DIR / "E215_rot_vs_orig_funnel.xlsx"

HARD = FC.HARD_GATES
BANDED = FC.BANDED_GATES
VARIANT_ORDER = {"orig": 0, "rot0": 1, "rot1": 2}
LAYERS = ["L1_reject", "L2_review", "L3_auto", "L3_review"]

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
NAVY = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF")
LAYER_FILL = {
    "L1_reject": PatternFill("solid", fgColor="C00000"),
    "L2_review": PatternFill("solid", fgColor="FFD966"),
    "L3_auto": PatternFill("solid", fgColor="70AD47"),
    "L3_review": PatternFill("solid", fgColor="F4B183"),
}


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh, delimiter="\t"))


def finite(v: Any) -> float:
    try:
        f = float(v)
        return f if math.isfinite(f) else math.nan
    except (TypeError, ValueError):
        return math.nan


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# --------------------------------------------------------------------------
# Inline funnel classification (E215 never ran classify_funnel).
# --------------------------------------------------------------------------
def classify(row: dict[str, Any], is_orig: bool) -> dict[str, str]:
    """Add hard/narrow/wide pass + failed lists + 3-tier layer to a rollout row.

    Layer rule mirrors classify_funnel.py: any hard fail or banded-WIDE fail => L1;
    all WIDE pass but >=1 NARROW fail => L2; all NARROW pass => L3 (orig auto,
    aug review, same auto/review split E200 uses for its coloring)."""
    hard_ok, hard_failed = FC.hard_gate_result(row)
    n_ok, n_failed = FC.banded_gate_result(row, "narrow")
    w_ok, w_failed = FC.banded_gate_result(row, "wide")
    if not hard_ok or not w_ok:
        layer = "L1_reject"
    elif not n_ok:
        layer = "L2_review"
    else:
        layer = "L3_auto" if is_orig else "L3_review"
    return {
        "hard_pass": str(hard_ok), "hard_failed": ",".join(hard_failed),
        "narrow_pass": str(n_ok), "narrow_failed": ",".join(n_failed),
        "wide_pass": str(w_ok), "wide_failed": ",".join(w_failed),
        "layer": layer,
    }


def _truthy(v: Any) -> bool:
    return str(v).strip().lower() in ("true", "1")


def load_rows() -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for r in read_tsv(ROLLOUT_TSV):
        variant = (r.get("aug_variant") or "").strip()
        if variant not in VARIANT_ORDER:
            continue
        r.update(classify(r, is_orig=(variant == "orig")))
        rows.append(r)
    return rows


# --------------------------------------------------------------------------
# Cell helpers (verbatim discipline from gen_E200_arm_funnel_compare_xlsx).
# --------------------------------------------------------------------------
def _cell(ws, r, c, value, fill=None, font=None):
    cell = ws.cell(row=r, column=c, value=value)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    return cell


def _num(ws, r, c, value):
    v = finite(value)
    _cell(ws, r, c, round(v, 4) if math.isfinite(v) else "")


def write_detail(ws, rows: list[dict[str, Any]]) -> None:
    header = ["object_key", "case_id", "arm_group", "aug_variant", "layer", "base_variant"]
    for name, _f, _op, _thr in HARD:
        header += [name, f"{name}✓"]
    for name, _f, _op, narrow, wide in BANDED:
        header += [name, f"{name} N({narrow:g})", f"{name} W({wide:g})"]
    for c, name in enumerate(header, start=1):
        cell = _cell(ws, 1, c, name, NAVY, HEAD_FONT)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "G2"

    for r, row in enumerate(rows, start=2):
        _cell(ws, r, 1, row.get("object_key", ""))
        _cell(ws, r, 2, row.get("case_id", ""))
        _cell(ws, r, 3, row.get("arm_group", ""))
        _cell(ws, r, 4, row.get("aug_variant", ""))
        _cell(ws, r, 5, row.get("layer", ""), LAYER_FILL.get(row.get("layer", "")), HEAD_FONT)
        _cell(ws, r, 6, row.get("base_variant", ""))
        c = 7
        hard_failed = set((row.get("hard_failed") or "").split(","))
        for name, field, _op, _thr in HARD:
            if field == "fall_flag":
                _cell(ws, r, c, str(row.get("fall_flag", "")))
            else:
                _num(ws, r, c, row.get(field))
            c += 1
            ok = name not in hard_failed
            _cell(ws, r, c, "PASS" if ok else "FAIL", GREEN if ok else RED)
            c += 1
        wide_failed = set((row.get("wide_failed") or "").split(","))
        narrow_failed = set((row.get("narrow_failed") or "").split(","))
        for name, field, _op, _n, _w in BANDED:
            _num(ws, r, c, row.get(field)); c += 1
            n_ok = name not in narrow_failed
            _cell(ws, r, c, "PASS" if n_ok else "FAIL", GREEN if n_ok else RED); c += 1
            w_ok = name not in wide_failed
            _cell(ws, r, c, "PASS" if w_ok else "FAIL", GREEN if w_ok else RED); c += 1

    widths = [10, 34, 16, 10, 11, 12] + [12, 6] * len(HARD) + [12, 9, 9] * len(BANDED)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _rate(rows: list[dict[str, Any]], pred) -> tuple[int, int]:
    return sum(1 for r in rows if pred(r)), len(rows)


def write_summary(ws, orig: list[dict[str, Any]], rot: list[dict[str, Any]]) -> None:
    import numpy as np

    arm_groups = sorted({r["arm_group"] for r in orig + rot})
    objects = sorted({r["object_key"] for r in orig + rot})

    def gate_pass(row, name, caliber):
        col = "narrow_failed" if caliber == "narrow" else "wide_failed"
        return name not in set((row.get(col) or "").split(","))

    def hard_pass(row, name):
        return name not in set((row.get("hard_failed") or "").split(","))

    def block(title: str, o_rows, r_rows, top: int) -> int:
        _cell(ws, top, 1, f"{title}  (orig {len(o_rows)} vs rot {len(r_rows)})", font=Font(bold=True))
        heads = ["metric / gate", "orig", "orig rate", "rot", "rot rate", "Δ (rot−orig)"]
        for c, h in enumerate(heads, start=1):
            _cell(ws, top + 1, c, h, NAVY, HEAD_FONT)
        r = top + 2

        def line(label, pred):
            nonlocal r
            on, ot = _rate(o_rows, pred); rn, rt = _rate(r_rows, pred)
            orate = on / ot if ot else float("nan")
            rrate = rn / rt if rt else float("nan")
            _cell(ws, r, 1, label)
            _cell(ws, r, 2, f"{on}/{ot}"); _cell(ws, r, 3, round(orate, 3) if ot else "")
            _cell(ws, r, 4, f"{rn}/{rt}"); _cell(ws, r, 5, round(rrate, 3) if rt else "")
            _cell(ws, r, 6, round((rrate - orate) * 100, 1) if ot and rt else "")
            r += 1

        for lay in LAYERS:
            line(lay, lambda x, L=lay: x.get("layer") == L)
        line("accept L3 (auto+review)", lambda x: x.get("layer") in ("L3_auto", "L3_review"))
        _cell(ws, r, 1, "— overall accept —"); ws.cell(row=r, column=1).font = Font(italic=True); r += 1
        line("hard_all_pass", lambda x: _truthy(x.get("hard_pass")))
        line("WIDE_all_pass (not L1)", lambda x: _truthy(x.get("hard_pass")) and _truthy(x.get("wide_pass")))
        line("NARROW_all_pass", lambda x: _truthy(x.get("hard_pass")) and _truthy(x.get("narrow_pass")))
        _cell(ws, r, 1, "— hard gates —"); ws.cell(row=r, column=1).font = Font(italic=True); r += 1
        for name, _f, _op, _thr in HARD:
            line(f"hard:{name}", lambda x, N=name: hard_pass(x, N))
        _cell(ws, r, 1, "— banded gates (N=narrow, W=wide) —"); ws.cell(row=r, column=1).font = Font(italic=True); r += 1
        for name, _f, _op, narrow, wide in BANDED:
            line(f"{name} N({narrow:g})", lambda x, N=name: gate_pass(x, N, "narrow"))
            line(f"{name} W({wide:g})", lambda x, N=name: gate_pass(x, N, "wide"))
        _cell(ws, r, 1, "— key-metric mean —"); ws.cell(row=r, column=1).font = Font(italic=True); r += 1
        for name, field, _op, *_ in BANDED:
            def mean(rows, f=field):
                v = np.array([finite(x.get(f)) for x in rows]); v = v[np.isfinite(v)]
                return float(v.mean()) if v.size else float("nan")
            om, rm = mean(o_rows), mean(r_rows)
            _cell(ws, r, 1, f"mean:{name}")
            _cell(ws, r, 2, round(om, 4) if math.isfinite(om) else "")
            _cell(ws, r, 4, round(rm, 4) if math.isfinite(rm) else "")
            _cell(ws, r, 6, round(rm - om, 4) if math.isfinite(om) and math.isfinite(rm) else "")
            r += 1
        return r + 1

    row_ptr = block("OVERALL", orig, rot, 1)
    for grp in arm_groups:
        row_ptr = block(f"arm_group={grp}",
                        [r for r in orig if r["arm_group"] == grp],
                        [r for r in rot if r["arm_group"] == grp], row_ptr)
    for obj in objects:
        row_ptr = block(f"object={obj}",
                        [r for r in orig if r["object_key"] == obj],
                        [r for r in rot if r["object_key"] == obj], row_ptr)
    ws.column_dimensions["A"].width = 34
    for col in "BCDEF":
        ws.column_dimensions[col].width = 14


def write_overview(ws, rows: list[dict[str, Any]], summary: dict[str, Any]) -> None:
    orig = [r for r in rows if r["aug_variant"] == "orig"]
    rot = [r for r in rows if r["aug_variant"] != "orig"]
    n_cases = len({r["case_id"] for r in rows})
    excl = summary.get("failures", []) + [
        {"case_id": c, "aug_variant": "orig", "error": "orig baseline missing"}
        for c in summary.get("orig_baseline_missing", [])
    ]
    lines = [
        ("E215 rot-vs-orig · 14-gate 宽/窄漏斗", "", ""),
        ("caliber 来源", "E201/funnel_config.py", f"{len(HARD)} hard + {len(BANDED)} banded = 14 gates"),
        ("分层规则", "hard fail 或 banded 破 WIDE→L1;全过 WIDE 破 NARROW→L2;全过 NARROW→L3", "classify_funnel.py"),
        ("rot-ori 口径", "旋转 45° 的 obj-ori 抬升是固有代价,报告不判劣化", "见 detail obj_ori 列"),
        ("配对 case 数", n_cases, "每 case = orig + rot0 + rot1"),
        ("orig 行 / rot 行", f"{len(orig)} / {len(rot)}", "rot = rot0 ∪ rot1"),
        ("源 rollout TSV", str(ROLLOUT_TSV.relative_to(REPO)), sha256(ROLLOUT_TSV)),
        ("metric standard", summary.get("metric_standard_id", ""), "同 E215 evaluator"),
        ("生成时间", summary.get("generated_at", ""), "e215_rot_eval_summary.json"),
        ("已登记排除", f"{len(excl)} 项", "eval 打分失败 / 无 orig 基线,下方列出"),
    ]
    for row in lines:
        ws.append(row)
    ws.append([])
    ws.append(["— 排除项(case / variant / 原因)—"])
    for e in excl:
        ws.append([e.get("case_id", ""), e.get("aug_variant", ""), e.get("error", "")])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 70
    for cell in ws[1]:
        cell.font = Font(bold=True, size=12)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()

    FC.assert_monotonic()
    if not ROLLOUT_TSV.is_file():
        raise SystemExit(f"missing rollout TSV: {ROLLOUT_TSV}")
    summary = json.loads(SUMMARY_JSON.read_text(encoding="utf-8")) if SUMMARY_JSON.is_file() else {}

    rows = load_rows()
    by_case: dict[str, set[str]] = {}
    for r in rows:
        by_case.setdefault(r["case_id"], set()).add(r["aug_variant"])
    incomplete = {c: sorted(v) for c, v in by_case.items() if len(v) != 3}
    if incomplete:
        print(f"[warn] {len(incomplete)} cases not 3-variant complete: {incomplete}", file=sys.stderr)

    rows.sort(key=lambda r: (r["object_key"], r["case_id"], VARIANT_ORDER.get(r["aug_variant"], 9)))
    orig = [r for r in rows if r["aug_variant"] == "orig"]
    rot = [r for r in rows if r["aug_variant"] != "orig"]

    wb = Workbook()
    write_overview(wb.active, rows, summary)
    wb.active.title = "Overview"
    write_detail(wb.create_sheet("Detail"), rows)
    write_summary(wb.create_sheet("Summary"), orig, rot)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(json.dumps({
        "status": "PASS",
        "output": str(args.out),
        "sheets": wb.sheetnames,
        "rows": len(rows),
        "orig": len(orig), "rot": len(rot),
        "cases": len(by_case),
        "incomplete_cases": incomplete,
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
