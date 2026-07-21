#!/usr/bin/env python3
"""Generate the E170 paired-validation workbook from canonical TSV/JSON outputs."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_EVAL = Path("workspace/core4d/results/E170/s6_downstream/eval/full")
DEFAULT_KEYFRAMES = Path("workspace/core4d/results/E170/s6_downstream/evidence/visual_qc/keyframe_manifest.tsv")
SHEETS = (
    ("Case Metrics", "e170_case_metrics.tsv"),
    ("Paired Deltas", "e170_paired_deltas.tsv"),
    ("Group Summary", "e170_group_summary.tsv"),
    ("Worst Cases", "e170_worst_cases.tsv"),
    ("Manual Review", "user_manual_review_template.tsv"),
    ("Codex Verification", "codex_verification.tsv"),
    ("Not Ready", "e170_not_ready.tsv"),
    ("Eval Errors", "e170_evaluation_errors.tsv"),
)


def read_tsv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if not path.is_file() or path.stat().st_size == 0:
        return [], []
    with path.open("r", encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream, delimiter="\t")
        return list(reader.fieldnames or []), list(reader)


def value(raw: str) -> Any:
    text = str(raw)
    if text.lower() == "true":
        return True
    if text.lower() == "false":
        return False
    if text == "":
        return None
    try:
        return float(text) if any(char in text.lower() for char in (".", "e")) else int(text)
    except ValueError:
        return text


def style_sheet(sheet: Any) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        width = min(45, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[letter].width = width
        for cell in column:
            cell.font = Font(name="Arial", size=9, bold=cell.row == 1, color="FFFFFF" if cell.row == 1 else "000000")
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def add_tsv_sheet(workbook: Workbook, title: str, path: Path) -> tuple[Any, dict[str, int]]:
    fields, rows = read_tsv(path)
    sheet = workbook.create_sheet(title)
    if not fields:
        sheet.append(["status"])
        sheet.append(["not_generated"])
        style_sheet(sheet)
        return sheet, {"status": 1}
    sheet.append(fields)
    for row in rows:
        sheet.append([value(row.get(field, "")) for field in fields])
    style_sheet(sheet)
    return sheet, {field: index + 1 for index, field in enumerate(fields)}


def formula_range(sheet: str, column: int, last: int = 29) -> str:
    letter = get_column_letter(column)
    return f"'{sheet}'!{letter}2:{letter}{last}"


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--eval-dir", type=Path, default=DEFAULT_EVAL)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--keyframe-manifest", type=Path, default=DEFAULT_KEYFRAMES)
    parser.add_argument("--require-keyframes", action="store_true")
    args = parser.parse_args()
    eval_dir = args.eval_dir
    output = args.output or eval_dir / "E170_box021_prg_full_validation.xlsx"
    summary = json.loads((eval_dir / "summary.json").read_text(encoding="utf-8"))

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    overview.append(["E170 Box021 PRG Full Validation", "Value", "Source / rule"])
    overview.append(["Metric standard", summary["metric_standard_id"], "summary.json"])
    overview.append(["Evaluated rows", None, "COUNTA Case Metrics case_id"])
    overview.append(["Numeric release pass", None, "COUNTIF numeric_release_pass"])
    overview.append(["User reviewed", None, "COUNTIF user_manual_review_status=reviewed"])
    overview.append(["Manual operational USE", None, "COUNTIF manual_operational_use"])
    overview.append(["Strict release usable", None, "COUNTIF strict_release_usable"])
    overview.append(["Strict recovery (E168 DNU)", None, "COUNTIFS strict + E168 DO_NOT_USE"])
    overview.append(["Strict retention (E168 USE)", None, "COUNTIFS strict + E168 USE"])
    overview.append(["Gate-health pass", None, "COUNTIF leg_gate_health_pass"])
    overview.append(["Machine recommendation", None, "22/9/12 strong; 18/6/10 partial; requires 28 reviews"])
    overview.append(["Final promotion decision", "PENDING_USER_DECISION", "User authority"])
    overview.append(["Manifest SHA256", summary["manifest_sha256"], summary["manifest"]])
    overview.append(["Baseline SHA256", summary["baseline_sha256"], summary["baseline"]])

    mappings = {}
    for title, filename in SHEETS:
        _, mapping = add_tsv_sheet(workbook, title, eval_dir / filename)
        mappings[title] = mapping
    keyframe_fields, keyframe_rows = read_tsv(args.keyframe_manifest)
    if args.require_keyframes:
        expected_events = {
            "pre_contact",
            "max_lower_body_penetration",
            "max_hand_penetration",
            "max_object_speed",
            "final",
        }
        keyframe_cases = {row.get("case_id", "") for row in keyframe_rows}
        keyframe_events = {row.get("event", "") for row in keyframe_rows}
        events_by_case: dict[str, list[str]] = {}
        for row in keyframe_rows:
            events_by_case.setdefault(row.get("case_id", ""), []).append(row.get("event", ""))
        invalid_cases = {
            case_id: events
            for case_id, events in events_by_case.items()
            if not case_id or len(events) != 5 or set(events) != expected_events
        }
        if (
            len(keyframe_rows) != 140
            or len(keyframe_cases) != 28
            or keyframe_events != expected_events
            or invalid_cases
        ):
            raise RuntimeError(
                "strict keyframe workbook gate failed: "
                f"rows={len(keyframe_rows)} cases={len(keyframe_cases)} "
                f"events={sorted(keyframe_events)} invalid_cases={invalid_cases}"
            )
    _, mappings["Keyframe Evidence"] = add_tsv_sheet(
        workbook,
        "Keyframe Evidence",
        args.keyframe_manifest,
    )
    case = mappings["Case Metrics"]
    required = ("case_id", "numeric_release_pass", "user_manual_review_status", "manual_operational_use", "strict_release_usable", "e168_manual_use_decision", "leg_gate_health_pass")
    missing = [field for field in required if field not in case]
    if missing:
        raise RuntimeError(f"Case Metrics missing workbook fields: {missing}")
    overview["B3"] = f"=COUNTA({formula_range('Case Metrics', case['case_id'])})"
    overview["B4"] = f"=COUNTIF({formula_range('Case Metrics', case['numeric_release_pass'])},TRUE)"
    overview["B5"] = f'=COUNTIF({formula_range("Case Metrics", case["user_manual_review_status"])},"reviewed")'
    overview["B6"] = f"=COUNTIF({formula_range('Case Metrics', case['manual_operational_use'])},TRUE)"
    overview["B7"] = f"=COUNTIF({formula_range('Case Metrics', case['strict_release_usable'])},TRUE)"
    overview["B8"] = f'=COUNTIFS({formula_range("Case Metrics", case["strict_release_usable"])},TRUE,{formula_range("Case Metrics", case["e168_manual_use_decision"])},"DO_NOT_USE")'
    overview["B9"] = f'=COUNTIFS({formula_range("Case Metrics", case["strict_release_usable"])},TRUE,{formula_range("Case Metrics", case["e168_manual_use_decision"])},"USE")'
    overview["B10"] = f"=COUNTIF({formula_range('Case Metrics', case['leg_gate_health_pass'])},TRUE)"
    overview["B11"] = '=IF(B5<>28,"PENDING_USER_REVIEW",IF(AND(B7>=22,B8>=9,B9>=12),"STRONG",IF(AND(B7>=18,B8>=6,B9>=10),"PARTIAL","FAIL")))'
    overview.freeze_panes = "A2"
    overview.column_dimensions["A"].width = 31
    overview.column_dimensions["B"].width = 26
    overview.column_dimensions["C"].width = 65
    for cell in overview[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=12)
        cell.fill = PatternFill("solid", fgColor="17365D")
    for row in overview.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in (overview["B5"], overview["B11"], overview["B12"]):
        cell.fill = PatternFill("solid", fgColor="FFF2CC")

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(json.dumps({"output": str(output), "sheets": workbook.sheetnames, "case_rows": overview["B3"].value, "keyframe_rows": len(keyframe_rows), "keyframe_fields": len(keyframe_fields), "formulas": 9}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
