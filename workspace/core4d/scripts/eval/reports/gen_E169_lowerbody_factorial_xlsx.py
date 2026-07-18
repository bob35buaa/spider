#!/usr/bin/env python3
"""Build the formatted E169 factorial evaluation workbook."""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO = Path(__file__).resolve().parents[5]
DEFAULT_EVAL = REPO / "workspace/core4d/results/E169/eval/full"
NAVY, TEAL, GREEN, RED, AMBER = "1F4E78", "0F6B78", "548235", "C00000", "BF8F00"
WHITE, LIGHT_GREEN, LIGHT_RED = "FFFFFF", "E2F0D9", "FCE4D6"


def read_tsv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if not path.is_file() or path.stat().st_size == 0:
        return [], []
    with path.open("r", encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream, delimiter="\t")
        return list(reader.fieldnames or []), list(reader)


def value(raw: Any) -> Any:
    text = str(raw).strip() if raw is not None else ""
    if not text:
        return None
    if text.lower() in {"true", "false"}:
        return text.lower() == "true"
    if re.fullmatch(r"[-+]?\d+", text):
        return int(text)
    try:
        number = float(text)
        return number if math.isfinite(number) else None
    except ValueError:
        return text


def table(ws, fields: list[str], rows: list[dict[str, str]], fill: str = NAVY) -> None:
    for col, field in enumerate(fields, 1):
        cell = ws.cell(1, col, field)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_idx, row in enumerate(rows, 2):
        for col, field in enumerate(fields, 1):
            cell = ws.cell(row_idx, col, value(row.get(field)))
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(max(1, len(fields)))}{max(1, len(rows) + 1)}"
    ws.sheet_view.showGridLines = False
    for col, field in enumerate(fields, 1):
        width = max([len(field), *[len(str(row.get(field, ""))) for row in rows[:80]]]) + 2
        ws.column_dimensions[get_column_letter(col)].width = min(max(width, 10), 36)


def overview(ws, summary: dict[str, Any], case_fields: list[str], row_count: int) -> None:
    ws.merge_cells("A1:H1")
    ws["A1"] = "E169 Lower-body / Object 2x2x2 Factorial"
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="left")
    ws["A3"], ws["B3"] = "Metric standard", summary["metric_standard_id"]
    ws["A4"], ws["B4"] = "Generated", summary["generated_at"]
    labels = ("Evaluated", "Numeric pass", "E169 acceptance", "G rows", "G health pass")
    for col, label in enumerate(labels, 1):
        ws.cell(6, col, label)
        ws.cell(6, col).font = Font(name="Arial", bold=True, color=WHITE)
        ws.cell(6, col).fill = PatternFill("solid", fgColor=TEAL)
    if row_count:
        col = {field: get_column_letter(index + 1) for index, field in enumerate(case_fields)}
        ws["A7"] = f"=COUNTA('Complete Metrics'!${col['variant']}$2:${col['variant']}${row_count + 1})"
        ws["B7"] = f"=COUNTIF('Complete Metrics'!${col['numeric_release_pass']}$2:${col['numeric_release_pass']}${row_count + 1},TRUE)"
        ws["C7"] = f"=COUNTIF('Complete Metrics'!${col['e169_acceptance_pass']}$2:${col['e169_acceptance_pass']}${row_count + 1},TRUE)"
        ws["D7"] = f"=COUNTIF('Complete Metrics'!${col['g_enabled']}$2:${col['g_enabled']}${row_count + 1},TRUE)"
        ws["E7"] = f"=COUNTIF('Complete Metrics'!${col['leg_gate_health_pass']}$2:${col['leg_gate_health_pass']}${row_count + 1},TRUE)"
    else:
        for col in range(1, 6):
            ws.cell(7, col, "=0")
    ws["A10"] = "Frozen thresholds"
    ws["A10"].font = Font(name="Arial", bold=True, color=WHITE)
    ws["A10"].fill = PatternFill("solid", fgColor=GREEN)
    for index, (key, threshold) in enumerate(summary["thresholds"].items(), 11):
        ws.cell(index, 1, key)
        ws.cell(index, 2, threshold)
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and cell.row not in {1, 6, 10}:
                cell.font = Font(name="Arial", size=10)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 30
    ws.sheet_view.showGridLines = False


def boolean_format(ws, fields: list[str]) -> None:
    for field in ("numeric_release_pass", "e169_acceptance_pass", "leg_gate_health_pass"):
        if field not in fields or ws.max_row < 2:
            continue
        letter = get_column_letter(fields.index(field) + 1)
        region = f"{letter}2:{letter}{ws.max_row}"
        ws.conditional_formatting.add(region, CellIsRule(operator="equal", formula=["TRUE"], fill=PatternFill("solid", fgColor=LIGHT_GREEN)))
        ws.conditional_formatting.add(region, CellIsRule(operator="equal", formula=["FALSE"], fill=PatternFill("solid", fgColor=LIGHT_RED)))


def validate(path: Path, expected_rows: int, expected_metrics: int) -> None:
    workbook = load_workbook(path, data_only=False)
    expected = [
        "Overview",
        "E168 Key Metrics",
        "Complete Metrics",
        "Cell Summary",
        "Metric Summary",
        "Worst Cases",
        "Factorial Contrasts",
        "Manual Review",
        "Not Ready",
        "Errors",
        "Manifest Snapshot",
        "Notes",
    ]
    if workbook.sheetnames != expected:
        raise RuntimeError(f"unexpected sheets: {workbook.sheetnames}")
    if workbook["Complete Metrics"].max_row - 1 != expected_rows:
        raise RuntimeError("case metric row count mismatch")
    if workbook["Complete Metrics"].max_column != expected_metrics:
        raise RuntimeError("complete metric column count mismatch")
    if workbook["E168 Key Metrics"].max_row - 1 != expected_rows:
        raise RuntimeError("E168 key metric row count mismatch")
    errors = []
    formulas = 0
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
                if isinstance(cell.value, str) and any(token in cell.value for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")):
                    errors.append(f"{ws.title}!{cell.coordinate}")
    if not formulas or errors:
        raise RuntimeError(f"formula validation failed formulas={formulas} errors={errors[:10]}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--eval-dir", type=Path, default=DEFAULT_EVAL)
    parser.add_argument("--output", type=Path, default=None)
    args = parser.parse_args()
    root = args.eval_dir.resolve()
    output = args.output or root / "E169_lowerbody_object_factorial_metrics.xlsx"
    summary = json.loads((root / "summary.json").read_text(encoding="utf-8"))
    specs = [
        ("E168 Key Metrics", "e169_e168_key_metrics.tsv", TEAL),
        ("Complete Metrics", "e169_case_metrics.tsv", NAVY),
        ("Cell Summary", "e169_cell_summary.tsv", GREEN),
        ("Metric Summary", "e169_metric_summary.tsv", GREEN),
        ("Worst Cases", "e169_worst_case_rankings.tsv", RED),
        ("Factorial Contrasts", "e169_factorial_contrasts.tsv", TEAL),
        ("Manual Review", "manual_review_template.tsv", AMBER),
        ("Not Ready", "e169_not_ready.tsv", AMBER),
        ("Errors", "e169_evaluation_errors.tsv", RED),
        ("Manifest Snapshot", "evaluated_manifest_snapshot.tsv", NAVY),
    ]
    source = {name: read_tsv(root / filename) for name, filename, _ in specs}
    case_fields, case_rows = source["Complete Metrics"]
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    ws = workbook.create_sheet("Overview")
    overview(ws, summary, case_fields, len(case_rows))
    for name, _filename, color in specs:
        fields, rows = source[name]
        ws = workbook.create_sheet(name)
        table(ws, fields, rows, color)
        if name in {"E168 Key Metrics", "Complete Metrics"}:
            boolean_format(ws, fields)
        if name == "Factorial Contrasts" and rows and "improvement_effect" in fields:
            letter = get_column_letter(fields.index("improvement_effect") + 1)
            ws.conditional_formatting.add(
                f"{letter}2:{letter}{ws.max_row}",
                ColorScaleRule(start_type="min", start_color=LIGHT_RED, mid_type="percentile", mid_value=50, mid_color="FFFFFF", end_type="max", end_color=LIGHT_GREEN),
            )
    ws = workbook.create_sheet("Notes")
    notes = [
        ("Experiment", "E169: P=physical pairs, R=soft lower-body reward, G=CEM leg gate."),
        ("Factorial effect", "raw_effect_on_minus_off uses the standard balanced 2^3 contrast; improvement_effect is positive when quality improves."),
        ("Acceptance", "e169_acceptance_pass requires all E168-aligned numeric gates and, for G rows, the frozen leg-gate health contract."),
        ("E168 parity", "E168 Key Metrics contains the same 23 core metrics used by E168 metric summaries, plus E169 leg physics/gate diagnostics."),
        ("Complete metrics", "Complete Metrics contains every column from e169_case_metrics.tsv; Metric Summary and Worst Cases restore the E168 analysis views."),
        ("Manual review", "Review all 28 new rows for visible penetration, stepping/support on the box, object kick, collision jitter, and posture failure."),
        ("Source", summary["manifest"]),
    ]
    table(ws, ["Item", "Description"], [{"Item": key, "Description": text} for key, text in notes], TEAL)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    validate(output, len(case_rows), len(case_fields))
    print(f"wrote {output} rows={len(case_rows)} sheets={len(workbook.sheetnames)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
