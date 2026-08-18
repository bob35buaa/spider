#!/usr/bin/env python3
"""E200 three-arm MASTER workbook: noPRG / PRG / PRG+G1+A2, fusing three evidence
layers per physical case, under the E201 14-gate wide/narrow funnel.

Layers fused:
  1. Upstream gate/funnel  (funnel_config 14-gate wide/narrow + 3-tier layer)
       noPRG  aug+orig -> E200_noprg_funnel_rollout.tsv
       PRG    aug+orig -> E199_funnel_rollout.tsv          (E199 fullscale)
       G1A2   aug      -> E200_prg_g1a2_funnel_rollout.tsv  (if present)
       G1A2   orig     -> E198 arm_cache (arm=G1A2) scored through funnel_config
  2. Human manual review (USE / DO_NOT_USE)
       PRG aug + PRG orig#variant -> E199 fullscale user_manual_review_filled (case#variant)
       PRG orig                   -> E170 (box021) + E172 (box004) + E173 (rest)
       G1A2 orig                  -> E198 full_factorial user_manual_review_filled
       noPRG                      -> none
  3. Downstream RL binary success (orig / 38 canonical case)
       PRG   -> R018-12 By Case (SUGAR-W / Holo-W)
       noPRG -> R018-20 By Case
       rule (user 2026-08-18): RL FAIL iff  SUGAR-W == 0  OR  Holo-W <= 0.10
                               => rl_ok = 1 iff SUGAR-W>0 AND Holo-W>0.10

Sheets:
  by_rollout — every (arm, case, variant): funnel layer/wide/narrow + key metrics
               + manual decision + rl_ok (orig only).
  by_case    — one row per physical case; three arm blocks (orig layer/WIDE/NARROW,
               aug narrow x/3), manual (PRG/G1A2), RL binary (noPRG/PRG).
  summary    — per arm x {aug,orig}: wide/narrow pass rate + layer mix; manual USE
               rate; RL success rate; and consistency cross-tabs (NARROW-pass and
               manual-USE vs RL success on the 38 canonical cases).
  臂间对比(McNemar) — paired arm comparison on the aug set (three arms re-ran the
               SAME E199 aug trajectories, so aug rollouts are genuinely paired;
               orig baselines differ per pipeline and are excluded). Block A: per-
               gate arm means across the 249 common rollouts + Δ(PRG−noPRG)/Δ(G1A2−PRG).
               Block B: McNemar-exact gate migration for noPRG→PRG, PRG→G1A2,
               noPRG→G1A2 (14 gates at hard/NARROW口径 + 3 composite funnel rows).

Usage:
    .venv/bin/python workspace/core4d/scripts/eval/reports/gen_E200_three_arm_master_xlsx.py
"""

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[5]
sys.path.insert(0, str(REPO / "workspace/core4d/scripts/experiments/E201"))
import funnel_config as FC  # noqa: E402

RES = REPO / "workspace/core4d/results"
FUNNEL = RES / "E201/funnel"
ARMS = ["noPRG", "PRG", "PRG+G1+A2"]
FUNNEL_TSV = {
    "noPRG": FUNNEL / "E200_noprg_funnel_rollout.tsv",
    "PRG": FUNNEL / "E199_funnel_rollout.tsv",
    "PRG+G1+A2": FUNNEL / "E200_prg_g1a2_funnel_rollout.tsv",  # may be absent until G1A2 classify runs
}
CASE_METRICS = {  # numeric values join (for by_rollout key metrics)
    "noPRG": RES / "E200/s6_downstream/eval/noprg/e200_noprg_case_metrics.tsv",
    "PRG": RES / "E199/s6_downstream/eval/fullscale_augmentation/e199_fullscale_case_metrics.tsv",
    "PRG+G1+A2": RES / "E200/s6_downstream/eval/prg_g1a2/e200_prg_g1a2_case_metrics.tsv",
}
E198_ARM_CACHE = RES / "E198/s6_downstream/eval/full_factorial/e198_arm_cache.tsv"
R018_PRG = REPO / "tmp/R018-12_upstream_downstream_metrics.xlsx"
R018_NOPRG = REPO / "tmp/R018-20_upstream_downstream_metrics.xlsx"
MANUAL_PRG_AUG = RES / "E199/s6_downstream/eval/fullscale_augmentation/user_manual_review_filled.tsv"
MANUAL_PRG_ORIG = [RES / "E170/s6_downstream/eval/full/user_manual_review_template.tsv",
                   RES / "E172/s6_downstream/eval/full/user_manual_review_filled.tsv",
                   RES / "E173/s6_downstream/eval/full/user_manual_review_filled.tsv"]
MANUAL_G1A2 = RES / "E198/s6_downstream/eval/full_factorial/user_manual_review_filled.tsv"
OUT = RES / "E200/s6_downstream/eval/E200_three_arm_master.xlsx"

VARIANTS = ["orig", "trans0", "trans1", "trans2"]
# full 14-gate value fields (from funnel_config, same order/口径): 4 hard + 10 banded
HARD_VAL = [(n, f) for n, f, _op, _thr in FC.HARD_GATES]          # (name, field)
BANDED_VAL = [(n, f, nar, wide) for n, f, _op, nar, wide in FC.BANDED_GATES]  # (name, field, narrow, wide)
GATE_FIELDS_ALL = [f for _n, f in HARD_VAL] + [f for _n, f, _nr, _w in BANDED_VAL]
# body_z is not stored in case_metrics (recomputed by classify); read it from the funnel row.
_BODYZ = "body_z_err_p95_m"

GREEN = PatternFill("solid", fgColor="C6EFCE"); RED = PatternFill("solid", fgColor="FFC7CE")
GREY = PatternFill("solid", fgColor="D9D9D9"); NAVY = PatternFill("solid", fgColor="1F3864")
HEAD = Font(bold=True, color="FFFFFF")
LAYER_FILL = {"L1_reject": PatternFill("solid", fgColor="C00000"),
              "L2_review": PatternFill("solid", fgColor="FFD966"),
              "L3_auto": PatternFill("solid", fgColor="70AD47"),
              "L3_review": PatternFill("solid", fgColor="F4B183"),
              "L3_narrow": PatternFill("solid", fgColor="70AD47")}


def finite(v: Any) -> float:
    try:
        f = float(v); return f if math.isfinite(f) else math.nan
    except (TypeError, ValueError):
        return math.nan


def read_tsv(p: Path) -> list[dict[str, str]]:
    if not p.is_file():
        return []
    with p.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh, delimiter="\t"))


def norm(cid: str) -> str:
    """physical case key: strip #variant, _person -> _p."""
    cid = cid.split("#")[0]
    return re.sub(r"_person(\d)$", r"_p\1", cid)


def variant_of(row: dict[str, str]) -> str:
    v = row.get("aug_variant") or ""
    if v in ("trans0", "trans1", "trans2"):
        return v
    return "orig"


# ---- gate/funnel per rollout ------------------------------------------------
def load_funnel_arm(arm: str) -> dict[tuple[str, str], dict[str, Any]]:
    """(norm_case, variant) -> funnel record for an arm (aug+orig from its rollout tsv)."""
    out: dict[tuple[str, str], dict[str, Any]] = {}
    vals = {(r.get("case_id", ""), variant_of(r)): r for r in read_tsv(CASE_METRICS[arm])}
    for r in read_tsv(FUNNEL_TSV[arm]):
        var = variant_of(r)
        key = (norm(r["case_id"]), var)
        v = vals.get((r["case_id"], var), {})
        gatevals = {f: (r.get(_BODYZ) if f == _BODYZ else v.get(f)) for f in GATE_FIELDS_ALL}
        out[key] = {
            "arm": arm, "object_key": r.get("object_key", ""), "case_id": norm(r["case_id"]), "variant": var,
            "layer": r.get("layer", ""), "hard_pass": r.get("hard_pass", ""),
            "wide_pass": r.get("wide_pass", ""), "narrow_pass": r.get("narrow_pass", ""),
            "narrow_failed": r.get("narrow_failed", ""), "wide_failed": r.get("wide_failed", ""),
            "hard_failed": r.get("hard_failed", ""), "gatevals": gatevals,
        }
    return out


def load_g1a2_orig_from_armcache() -> dict[tuple[str, str], dict[str, Any]]:
    """G1A2 orig funnel from E198 arm_cache (has body_z + all fields); per-row layer only."""
    out: dict[tuple[str, str], dict[str, Any]] = {}
    for r in read_tsv(E198_ARM_CACHE):
        if r.get("arm") != "G1A2":
            continue
        hard_ok, hard_f = FC.hard_gate_result(r)
        wide_ok, wide_f = FC.banded_gate_result(r, "wide")
        narrow_ok, narrow_f = FC.banded_gate_result(r, "narrow")
        layer = "L1_reject" if (not hard_ok or not wide_ok) else ("L2_review" if not narrow_ok else "L3_narrow")
        out[(norm(r["case_id"]), "orig")] = {
            "arm": "PRG+G1+A2", "object_key": r.get("object_key", ""), "case_id": norm(r["case_id"]),
            "variant": "orig", "layer": layer, "hard_pass": hard_ok, "wide_pass": wide_ok, "narrow_pass": narrow_ok,
            "narrow_failed": ",".join(narrow_f), "wide_failed": ",".join(wide_f), "hard_failed": ",".join(hard_f),
            "gatevals": {f: r.get(f) for f in GATE_FIELDS_ALL},
        }
    return out


def _arbitrate_g1a2_orig(recs: dict[tuple[str, str], dict[str, Any]]) -> None:
    """G1A2 orig comes from arm_cache un-arbitrated (L3_narrow). Family-check it
    against the present aug siblings so its L3 label matches every other arm:
    L3_auto iff all other present variants of the case pass narrow, else L3_review."""
    for (nc, var), r in recs.items():
        if var != "orig" or r.get("layer") != "L3_narrow":
            continue
        others = [recs[(nc, v)] for v in ("trans0", "trans1", "trans2") if (nc, v) in recs]
        all_ok = all(_b(o["narrow_pass"]) for o in others)  # all([]) is True -> lone orig auto
        r["layer"] = "L3_auto" if all_ok else "L3_review"


# ---- manual review ----------------------------------------------------------
def load_manual() -> dict[tuple[str, str, str], dict[str, str]]:
    """(norm_case, variant, arm) -> {decision, label, note}."""
    out: dict[tuple[str, str, str], dict[str, str]] = {}

    def put(cid_raw: str, variant: str, arm: str, r: dict[str, str]) -> None:
        out[(norm(cid_raw), variant, arm)] = {
            "decision": r.get("manual_use_decision", ""), "label": r.get("manual_quality_label", ""),
            "note": r.get("manual_review_note", "")}

    # PRG aug+orig: E199 fullscale, case_id like "..._person1#trans0"
    for r in read_tsv(MANUAL_PRG_AUG):
        cid = r["case_id"]
        var = cid.split("#")[1] if "#" in cid else "orig"
        var = var if var in VARIANTS else "orig"
        put(cid, var, "PRG", r)
    # PRG orig: E170/E172/E173 (variant=orig)
    for src in MANUAL_PRG_ORIG:
        for r in read_tsv(src):
            out.setdefault((norm(r["case_id"]), "orig", "PRG"),
                           {"decision": r.get("manual_use_decision", ""),
                            "label": r.get("manual_quality_label", ""), "note": r.get("manual_review_note", "")})
    # G1A2 orig: E198
    for r in read_tsv(MANUAL_G1A2):
        put(r["case_id"], "orig", "PRG+G1+A2", r)
    return out


# ---- downstream RL binary ---------------------------------------------------
def _r018_by_case(path: Path) -> dict[str, tuple[float, float]]:
    """norm_case -> (SUGAR-W, Holo-W) success rates from R018 By Case (cols 17,19)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["By Case"]
    out: dict[str, tuple[float, float]] = {}
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r or not r[0]:
            continue
        out[norm(str(r[0]))] = (finite(r[17]), finite(r[19]))  # SUGAR-W col17, Holo-W col19
    return out


def rl_ok(sugar_w: float, holo_w: float) -> Any:
    if not math.isfinite(sugar_w) or not math.isfinite(holo_w):
        return None
    return 0 if (sugar_w == 0 or holo_w <= 0.10) else 1


# ---- assemble ---------------------------------------------------------------
def build():
    funnel: dict[str, dict[tuple[str, str], dict[str, Any]]] = {a: load_funnel_arm(a) for a in ARMS}
    funnel["PRG+G1+A2"].update(load_g1a2_orig_from_armcache())  # add G1A2 orig
    _arbitrate_g1a2_orig(funnel["PRG+G1+A2"])  # L3_narrow -> L3_auto/L3_review (same口径 as other arms)
    manual = load_manual()
    rl = {"PRG": _r018_by_case(R018_PRG), "noPRG": _r018_by_case(R018_NOPRG)}
    rl_arm = {"PRG": "PRG", "noPRG": "noPRG"}  # G1A2 has no RL

    rollouts: list[dict[str, Any]] = []
    for arm in ARMS:
        for (nc, var), rec in funnel[arm].items():
            row = dict(rec)
            m = manual.get((nc, var, arm))
            row["manual_decision"] = m["decision"] if m else ""
            row["manual_label"] = m["label"] if m else ""
            # RL only for orig
            sw = hw = math.nan
            if var == "orig" and arm in rl and nc in rl[arm]:
                sw, hw = rl[arm][nc]
            row["SUGAR_W"] = sw; row["Holo_W"] = hw; row["rl_ok"] = rl_ok(sw, hw)
            rollouts.append(row)
    return funnel, manual, rl, rollouts


# ---- xlsx writers -----------------------------------------------------------
def _c(ws, r, c, v, fill=None, font=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    return cell


def _n(ws, r, c, v):
    f = finite(v); _c(ws, r, c, round(f, 4) if math.isfinite(f) else "")


def _b(v) -> bool | None:
    s = str(v).strip().lower()
    if s in ("true", "1"):
        return True
    if s in ("false", "0"):
        return False
    return None


def write_by_rollout(ws, rollouts):
    # id block -> hard-gate overall -> 4 hard (value + ✓) -> 10 banded (value + N + W)
    # -> manual -> downstream RL. Every one of the 14 gates shows its numeric value.
    hdr = ["arm", "object", "case_id", "variant", "layer", "hard_all", "WIDE_all", "NARROW_all"]
    for name, _f in HARD_VAL:
        hdr += [name, f"{name}✓"]
    for name, _f, nar, wide in BANDED_VAL:
        hdr += [name, f"{name} N({nar:g})", f"{name} W({wide:g})"]
    hdr += ["manual", "manual_label", "SUGAR-W", "Holo-W", "rl_ok"]
    for c, h in enumerate(hdr, 1):
        _c(ws, 1, c, h, NAVY, HEAD)
        ws.cell(row=1, column=c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "E2"
    order = {"orig": 0, "trans0": 1, "trans1": 2, "trans2": 3}
    aorder = {"noPRG": 0, "PRG": 1, "PRG+G1+A2": 2}
    rollouts = sorted(rollouts, key=lambda r: (r["object_key"], r["case_id"], order.get(r["variant"], 9), aorder[r["arm"]]))
    for i, r in enumerate(rollouts, 2):
        gv = r.get("gatevals", {})
        hard_f = set((r.get("hard_failed") or "").split(","))
        narrow_f = set((r.get("narrow_failed") or "").split(","))
        wide_f = set((r.get("wide_failed") or "").split(","))
        _c(ws, i, 1, r["arm"]); _c(ws, i, 2, r["object_key"]); _c(ws, i, 3, r["case_id"]); _c(ws, i, 4, r["variant"])
        _c(ws, i, 5, r["layer"], LAYER_FILL.get(r["layer"]))
        for c, key in ((6, "hard_pass"), (7, "wide_pass"), (8, "narrow_pass")):
            b = _b(r[key]); _c(ws, i, c, "P" if b else "F" if b is False else "", GREEN if b else RED if b is False else None)
        col = 9
        for name, field in HARD_VAL:
            if field == "fall_flag":
                _c(ws, i, col, str(gv.get(field, "")))
            else:
                _n(ws, i, col, gv.get(field))
            col += 1
            ok = name not in hard_f
            _c(ws, i, col, "P" if ok else "F", GREEN if ok else RED); col += 1
        for name, field, _nr, _w in BANDED_VAL:
            _n(ws, i, col, gv.get(field)); col += 1
            n_ok = name not in narrow_f
            _c(ws, i, col, "P" if n_ok else "F", GREEN if n_ok else RED); col += 1
            w_ok = name not in wide_f
            _c(ws, i, col, "P" if w_ok else "F", GREEN if w_ok else RED); col += 1
        md = r["manual_decision"]
        _c(ws, i, col, md, GREEN if md == "USE" else RED if md == "DO_NOT_USE" else None); col += 1
        _c(ws, i, col, r["manual_label"]); col += 1
        _n(ws, i, col, r["SUGAR_W"]); col += 1
        _n(ws, i, col, r["Holo_W"]); col += 1
        ok = r["rl_ok"]
        _c(ws, i, col, ("ok" if ok == 1 else "FAIL" if ok == 0 else ""), GREEN if ok == 1 else RED if ok == 0 else None)
    widths = [10, 9, 30, 8, 11, 6, 6, 7] + [9, 5] * len(HARD_VAL) + [9, 8, 8] * len(BANDED_VAL) + [11, 14, 9, 9, 7]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_by_case(ws, funnel, manual, rl):
    cases = sorted({(rec["object_key"], nc) for a in ARMS for (nc, _v), rec in funnel[a].items()},
                   key=lambda x: (x[0], x[1]))
    # header: id + per-arm (orig layer/W/N, aug N x/3) + manual(PRG/G1A2) + RL(noPRG/PRG)
    hdr = ["object", "case_id"]
    for a in ARMS:
        hdr += [f"{a}:orig layer", f"{a}:orig W", f"{a}:orig N", f"{a}:aug N(x/3)"]
    hdr += ["PRG manual(orig)", "G1A2 manual(orig)", "noPRG RL", "PRG RL",
            "PRG SUGAR-W", "PRG Holo-W", "noPRG SUGAR-W", "noPRG Holo-W"]
    for c, h in enumerate(hdr, 1):
        _c(ws, 1, c, h, NAVY, HEAD); ws.cell(row=1, column=c).alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = "C2"
    for i, (obj, nc) in enumerate(cases, 2):
        _c(ws, i, 1, obj); _c(ws, i, 2, nc)
        col = 3
        for a in ARMS:
            o = funnel[a].get((nc, "orig"))
            _c(ws, i, col, o["layer"] if o else "", LAYER_FILL.get(o["layer"]) if o else None); col += 1
            for key in ("wide_pass", "narrow_pass"):
                b = _b(o[key]) if o else None
                _c(ws, i, col, "P" if b else "F" if b is False else "", GREEN if b else RED if b is False else None); col += 1
            naug = sum(1 for v in ("trans0", "trans1", "trans2")
                       if (nc, v) in funnel[a] and _b(funnel[a][(nc, v)]["narrow_pass"]))
            taug = sum(1 for v in ("trans0", "trans1", "trans2") if (nc, v) in funnel[a])
            _c(ws, i, col, f"{naug}/{taug}" if taug else ""); col += 1
        pm = manual.get((nc, "orig", "PRG")); gm = manual.get((nc, "orig", "PRG+G1+A2"))
        _c(ws, i, col, pm["decision"] if pm else "", GREEN if pm and pm["decision"] == "USE" else RED if pm and pm["decision"] == "DO_NOT_USE" else None); col += 1
        _c(ws, i, col, gm["decision"] if gm else "", GREEN if gm and gm["decision"] == "USE" else RED if gm and gm["decision"] == "DO_NOT_USE" else None); col += 1
        for arm in ("noPRG", "PRG"):
            v = rl[arm].get(nc)
            ok = rl_ok(*v) if v else None
            _c(ws, i, col, "ok" if ok == 1 else "FAIL" if ok == 0 else "", GREEN if ok == 1 else RED if ok == 0 else None); col += 1
        for arm in ("PRG", "noPRG"):
            v = rl[arm].get(nc)
            _n(ws, i, col, v[0] if v else math.nan); col += 1
            _n(ws, i, col, v[1] if v else math.nan); col += 1
    widths = [9, 30] + [11, 5, 5, 9] * len(ARMS) + [15, 15, 8, 8, 11, 11, 11, 11]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_summary(ws, rollouts, rl):
    r = 1

    def hdr(row, cols):
        for c, h in enumerate(cols, 1):
            _c(ws, row, c, h, NAVY, HEAD)

    _c(ws, r, 1, "A) 上游 funnel 通过率（逐 arm × orig/aug）", font=Font(bold=True)); r += 1
    hdr(r, ["arm", "group", "n", "WIDE pass", "NARROW pass", "L1", "L2", "L3_auto", "L3_review/narrow"]); r += 1
    for arm in ARMS:
        for grp, sel in (("orig", "orig"), ("aug", None)):
            rs = [x for x in rollouts if x["arm"] == arm and (x["variant"] == "orig" if grp == "orig" else x["variant"] != "orig")]
            if not rs:
                continue
            n = len(rs)
            wp = sum(1 for x in rs if _b(x["wide_pass"])); npn = sum(1 for x in rs if _b(x["narrow_pass"]))
            lc = {L: sum(1 for x in rs if x["layer"] == L) for L in ("L1_reject", "L2_review", "L3_auto")}
            l3rn = sum(1 for x in rs if x["layer"] in ("L3_review", "L3_narrow"))
            _c(ws, r, 1, arm); _c(ws, r, 2, grp); _c(ws, r, 3, n)
            _c(ws, r, 4, f"{wp}/{n} ({wp/n:.0%})"); _c(ws, r, 5, f"{npn}/{n} ({npn/n:.0%})")
            _c(ws, r, 6, lc["L1_reject"]); _c(ws, r, 7, lc["L2_review"]); _c(ws, r, 8, lc["L3_auto"]); _c(ws, r, 9, l3rn)
            r += 1
    r += 1
    _c(ws, r, 1, "B) 人工 USE 率（有标注的 arm/组）", font=Font(bold=True)); r += 1
    hdr(r, ["arm", "group", "n_reviewed", "USE", "DO_NOT_USE", "USE rate"]); r += 1
    for arm in ARMS:
        for grp in ("orig", "aug"):
            rs = [x for x in rollouts if x["arm"] == arm and x["manual_decision"] and
                  (x["variant"] == "orig" if grp == "orig" else x["variant"] != "orig")]
            if not rs:
                continue
            use = sum(1 for x in rs if x["manual_decision"] == "USE")
            _c(ws, r, 1, arm); _c(ws, r, 2, grp); _c(ws, r, 3, len(rs)); _c(ws, r, 4, use)
            _c(ws, r, 5, len(rs) - use); _c(ws, r, 6, round(use / len(rs), 3)); r += 1
    r += 1
    _c(ws, r, 1, "C) 下游 RL 二值成功率（orig，规则: FAIL iff SUGAR-W==0 or Holo-W<=0.10）", font=Font(bold=True)); r += 1
    hdr(r, ["arm", "n_case", "rl_ok", "rl_fail", "success rate"]); r += 1
    for arm in ("noPRG", "PRG"):
        oks = [rl_ok(*v) for v in rl[arm].values()]
        oks = [o for o in oks if o is not None]
        n = len(oks); s = sum(oks)
        _c(ws, r, 1, arm); _c(ws, r, 2, n); _c(ws, r, 3, s); _c(ws, r, 4, n - s)
        _c(ws, r, 5, round(s / n, 3) if n else ""); r += 1
    r += 1
    _c(ws, r, 1, "D) 一致性（38 canonical orig）: 上游门/人工 vs 下游 RL", font=Font(bold=True)); r += 1
    hdr(r, ["arm", "predictor", "pred+ & RL_ok", "pred+ & RL_fail", "pred- & RL_ok", "pred- & RL_fail", "note"]); r += 1
    for arm in ("noPRG", "PRG"):
        rs = [x for x in rollouts if x["arm"] == arm and x["variant"] == "orig" and x["rl_ok"] is not None]
        for pname, pred in (("NARROW_pass", lambda x: _b(x["narrow_pass"]) is True),
                            ("manual_USE", lambda x: x["manual_decision"] == "USE")):
            if pname == "manual_USE" and not any(x["manual_decision"] for x in rs):
                continue
            tp = sum(1 for x in rs if pred(x) and x["rl_ok"] == 1)
            fp = sum(1 for x in rs if pred(x) and x["rl_ok"] == 0)
            fn = sum(1 for x in rs if not pred(x) and x["rl_ok"] == 1)
            tn = sum(1 for x in rs if not pred(x) and x["rl_ok"] == 0)
            _c(ws, r, 1, arm); _c(ws, r, 2, pname)
            _c(ws, r, 3, tp); _c(ws, r, 4, fp); _c(ws, r, 5, fn); _c(ws, r, 6, tn)
            _c(ws, r, 7, f"n={len(rs)}"); r += 1
    ws.column_dimensions["A"].width = 12
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 16


# ---- arm×arm paired comparison (McNemar) ------------------------------------
# The three arms re-ran the SAME E199 aug trajectories, so aug rollouts are a
# genuine paired set (same trajectory, different downstream arm). orig rollouts
# are each arm's own baseline (different pipelines: noPRG=E190, PRG=E199 orig,
# G1A2=E198) and are NOT paired, so this analysis is restricted to aug.
DIR = {n: ("↑" if op == ">=" else "↓") for n, _f, op, *_ in FC.HARD_GATES}
DIR.update({n: ("↑" if op == ">=" else "↓") for n, _f, op, *_ in FC.BANDED_GATES})
GATE_META = ([(n, f, True) for n, f in HARD_VAL]              # (name, field, is_hard)
             + [(n, f, False) for n, f, _nr, _w in BANDED_VAL])
YELLOW = PatternFill("solid", fgColor="FFEB9C")


def mcnemar_exact(p2f: int, f2p: int) -> float:
    from math import comb
    n = p2f + f2p
    if n == 0:
        return 1.0
    k = min(p2f, f2p)
    tail = sum(comb(n, i) for i in range(0, k + 1)) / (2 ** n)
    return float(min(1.0, 2.0 * tail))


def _gate_pass(rec: dict[str, Any], name: str, is_hard: bool) -> bool:
    """per-gate pass on the primary口径 (hard for hard gates, NARROW for banded)."""
    field = "hard_failed" if is_hard else "narrow_failed"
    return name not in set((rec.get(field) or "").split(","))


def write_arm_compare(ws, funnel):
    aug = {a: {k: v for k, v in funnel[a].items() if k[1] != "orig"} for a in ARMS}
    common = set(aug[ARMS[0]]) & set(aug[ARMS[1]]) & set(aug[ARMS[2]])
    r = 1
    _c(ws, r, 1, "本表仅用 aug 配对：三 arm 复跑同一条 E199 增强轨迹，逐 rollout 配对可比；"
                 "orig 各 arm 基线不同管线，不做配对。", font=Font(bold=True, color="C00000")); r += 2

    # -- Block A: per-gate arm means on the set common to all three arms --
    _c(ws, r, 1, f"A) 三臂 × 14-gate 数值均值（aug，三臂共有 {len(common)} 条 rollout）", font=Font(bold=True)); r += 1
    hdr = ["门", "方向", "n"] + ARMS + ["Δ(PRG−noPRG)", "Δ(G1A2−PRG)"]
    for c, h in enumerate(hdr, 1):
        _c(ws, r, c, h, NAVY, HEAD)
    r += 1
    for name, field, is_hard in GATE_META:
        d = DIR[name]
        _c(ws, r, 1, name); _c(ws, r, 2, d); _c(ws, r, 3, len(common))
        means = {}
        for c, a in enumerate(ARMS, 4):
            if field == "fall_flag":
                vals = [1.0 if _b(aug[a][k]["gatevals"].get(field)) else 0.0 for k in common]
            else:
                vals = [finite(aug[a][k]["gatevals"].get(field)) for k in common]
                vals = [v for v in vals if math.isfinite(v)]
            m = sum(vals) / len(vals) if vals else math.nan
            means[a] = m
            _c(ws, r, c, round(m, 4) if math.isfinite(m) else "")
        for c, (hi, lo) in enumerate(((ARMS[1], ARMS[0]), (ARMS[2], ARMS[1])), 7):
            delta = means[hi] - means[lo]
            fill = None
            if math.isfinite(delta) and abs(delta) > 1e-9:
                improved = (delta > 0) if d == "↑" else (delta < 0)
                fill = GREEN if improved else RED
            _c(ws, r, c, round(delta, 4) if math.isfinite(delta) else "", fill)
        r += 1
    r += 1

    # -- Block B: McNemar gate migration across arm pairs --
    _c(ws, r, 1, "B) 臂间门迁移（McNemar exact，aug 配对；banded 用 NARROW 口径，p<0.05 高亮）",
       font=Font(bold=True)); r += 1
    hdr = ["迁移", "门", "口径", "n", "前通过", "后通过", "Δpp", "P→F", "F→P", "exact p"]
    for c, h in enumerate(hdr, 1):
        _c(ws, r, c, h, NAVY, HEAD)
    r += 1
    transitions = [("noPRG→PRG", "noPRG", "PRG"),
                   ("PRG→G1A2", "PRG", "PRG+G1+A2"),
                   ("noPRG→G1A2", "noPRG", "PRG+G1+A2")]
    comp = [("NARROW_all", "narrow_pass", "复合"), ("WIDE_all", "wide_pass", "复合"),
            ("hard_all", "hard_pass", "复合")]
    for tname, before, after in transitions:
        sel = sorted(set(aug[before]) & set(aug[after]))
        # per-gate (hard + banded narrow)
        for name, _field, is_hard in GATE_META:
            kind = "hard" if is_hard else "narrow"
            bp = sum(_gate_pass(aug[before][k], name, is_hard) for k in sel)
            ap_ = sum(_gate_pass(aug[after][k], name, is_hard) for k in sel)
            p2f = sum(_gate_pass(aug[before][k], name, is_hard)
                      and not _gate_pass(aug[after][k], name, is_hard) for k in sel)
            f2p = sum((not _gate_pass(aug[before][k], name, is_hard))
                      and _gate_pass(aug[after][k], name, is_hard) for k in sel)
            _write_mig_row(ws, r, tname, name, kind, len(sel), bp, ap_, p2f, f2p); r += 1
        # composite funnel decisions
        for cname, key, kind in comp:
            bp = sum(1 for k in sel if _b(aug[before][k][key]))
            ap_ = sum(1 for k in sel if _b(aug[after][k][key]))
            p2f = sum(1 for k in sel if _b(aug[before][k][key]) and not _b(aug[after][k][key]))
            f2p = sum(1 for k in sel if (not _b(aug[before][k][key])) and _b(aug[after][k][key]))
            _write_mig_row(ws, r, tname, cname, kind, len(sel), bp, ap_, p2f, f2p); r += 1
    widths = [12, 12, 8, 5, 8, 8, 8, 6, 6, 9]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_mig_row(ws, r, tname, gate, kind, n, bp, ap_, p2f, f2p):
    dpp = round((ap_ - bp) / n * 100, 2) if n else 0.0
    p = mcnemar_exact(p2f, f2p)
    sig = p < 0.05
    _c(ws, r, 1, tname); _c(ws, r, 2, gate); _c(ws, r, 3, kind); _c(ws, r, 4, n)
    _c(ws, r, 5, bp); _c(ws, r, 6, ap_)
    _c(ws, r, 7, dpp, GREEN if dpp > 0 else RED if dpp < 0 else None)
    _c(ws, r, 8, p2f); _c(ws, r, 9, f2p)
    _c(ws, r, 10, round(p, 6), YELLOW if sig else None)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=OUT)
    args = ap.parse_args()
    funnel, manual, rl, rollouts = build()
    for a in ARMS:
        n_aug = sum(1 for (nc, v) in funnel[a] if v != "orig")
        n_orig = sum(1 for (nc, v) in funnel[a] if v == "orig")
        print(f"[{a}] funnel rollouts: {n_aug} aug + {n_orig} orig")
    print(f"[manual] {len(manual)} entries | [RL] PRG {len(rl['PRG'])}, noPRG {len(rl['noPRG'])} cases")
    wb = Workbook()
    write_by_rollout(wb.active, rollouts); wb.active.title = "by_rollout"
    write_by_case(wb.create_sheet("by_case"), funnel, manual, rl)
    write_summary(wb.create_sheet("summary"), rollouts, rl)
    write_arm_compare(wb.create_sheet("臂间对比(McNemar)"), funnel)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"[done] wrote {args.out} ({len(rollouts)} rollout rows)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
