#!/usr/bin/env python3
"""Build the formula-driven E194 noPRG / PRG / G1 comparison workbook."""

from __future__ import annotations

import hashlib
import math
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "experiments/E194"))
import e194_g1_expansion_common as C  # noqa: E402

EVAL = C.RESULTS / "s6_downstream/eval/full_g1_expansion"
OUTPUT = EVAL / "E194_noPRG_PRG_G1_comparison.xlsx"
E173_BOX001_USE_AUTHORITY = (
    C.RESULTS.parent
    / "E173/s6_downstream/rl_export/box001_user_approved/box001_user_approved_source_rows.tsv"
)
E194_G1_REVIEW = EVAL / "user_manual_review_filled.tsv"
EXCLUDED_BOX001_CASE = "box001_20231023_110_p1"
PRIMARY = "track_obj_z_abs_err_cm_mean"
METRICS = (
    (PRIMARY, "Object z MAE (cm)", "lower"),
    ("track_obj_pos_err_cm_mean", "Object 3D pos MAE (cm)", "lower"),
    ("track_obj_ori_err_deg_mean", "Object orientation error (deg)", "lower"),
    ("body_z_err_p95_m", "Body z error p95 (m)", "lower"),
    ("track_pelvis_z_err_terminal_m", "Pelvis z terminal error (m)", "lower"),
    ("track_root_pos_err_cm_mean", "Body/root position error (cm)", "lower"),
    ("track_root_ori_err_deg_mean", "Body/root orientation error (deg)", "lower"),
    ("track_eef_pos_err_cm_mean", "Hand position error (cm)", "lower"),
    ("track_eef_ori_err_deg_mean", "Hand orientation error (deg)", "lower"),
    ("hand_object_physics_contact_3mm_in_mask_frac", "3mm contact frac", "higher"),
    ("hand_object_physics_contact_in_mask_frac", "Raw contact frac", "higher"),
    ("hand_object_release_false_contact_3mm_frac", "False release frac", "lower"),
    ("hand_object_physics_penetration_3mm_frame_frac", "Hand penetration frac", "lower"),
    ("leg_penetration_frac", "Leg penetration frac", "lower"),
)
GATES = ("fall", "body_z", "contact", "release", "hand_penetration", "lower_body",
         "root_pos", "root_ori", "hand_pos", "hand_ori", "object_pos", "object_ori")
G1_REVIEW_FIELDS = (
    "user_manual_review_status",
    "manual_use_decision",
    "manual_quality_label",
    "manual_failure_taxonomy",
    "manual_review_note",
    "manual_reviewer",
    "manual_reviewed_at",
    "paired_video",
)
NAVY = "17365D"
BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_RED = "FCE4D6"
LIGHT_GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7B7B7")


def number(value: Any) -> Any:
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (TypeError, ValueError):
        return value


def truth(value: Any) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "on"}


def title(ws, text: str, subtitle: str, columns: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws["A1"] = text
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="666666")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30


def header(ws, row: int, columns: int) -> None:
    for cell in ws[row][:columns]:
        cell.font = Font(name="Arial", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    ws.row_dimensions[row].height = 32


def finish_table(ws, start_row: int, end_row: int, end_col: int, name: str) -> None:
    table = Table(displayName=name, ref=f"A{start_row}:{get_column_letter(end_col)}{end_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False,
                                          showLastColumn=False, showColumnStripes=False)
    ws.add_table(table)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(end_col)}{end_row}"
    ws.freeze_panes = f"A{start_row + 1}"


def set_widths(ws, widths: dict[int, float], default: float = 13) -> None:
    for index in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(index, default)


def add_delta_quality_formatting(ws, cell_range: str, direction: str) -> None:
    """Apply a zero-centered gradient: green=better, red=worse, yellow=no change."""
    if direction not in {"higher", "lower"}:
        raise ValueError(f"unknown metric direction: {direction}")
    start_color, end_color = (("F8696B", "63BE7B") if direction == "higher" else ("63BE7B", "F8696B"))
    ws.conditional_formatting.add(
        cell_range,
        ColorScaleRule(
            start_type="min", start_color=start_color,
            mid_type="num", mid_value=0, mid_color="FFEB84",
            end_type="max", end_color=end_color,
        ),
    )


def add_readme(
    wb: Workbook,
    summary: dict[str, Any],
    authority_sha256: str,
    g1_review_sha256: str,
    corrected_overlay: bool = False,
    corrected_count: int = 0,
) -> None:
    ws = wb.active
    ws.title = "README"
    title(ws, "E194 noPRG / PRG / G1 Full comparison",
          "72 matched cases; each delta column marks whether higher or lower is better.", 4)
    g1_authority = (
        f"E194 Full G1 with E196 corrected overlay ({corrected_count} cases)"
        if corrected_overlay else "E194 Full expansion"
    )
    rows = [
        ("Status", summary.get("status"), "Expected: pass", ""),
        ("Cardinality", "216 arm-case rows", "144 paired comparison rows", "1,728 gate migrations"),
        ("Objects", "box001: 28", "box023: 16", "box021: 28"),
        ("noPRG authority", "box001: E189 Full", "box023: E179 Full", "box021: E168 production"),
        ("PRG authority", "E173 Full (box001/box023)", "E170 Full (box021)", "exact E194 source authority"),
        ("G1 authority", g1_authority, "kp_pos=500, kp_rot=50", "gravcomp=1.0"),
        ("Metric contract", summary.get("metric_standard_id", ""), "public core scorer", "12 frozen gates"),
        ("Formula convention", "Delta = after - before", "green = improvement; red = regression", "headers mark ↑/↓ better"),
        ("PRG human authority", "E173 box001_user_approved_source_rows.tsv", "box001: 13 USE / 15 DNU", "box023/box021: not applicable"),
        ("Authority rule", "membership = USE", "absent within E194 box001 universe = DNU", f"source SHA256: {authority_sha256}"),
        ("G1 box001 human review", "28/28 reviewed: 15 USE / 13 DNU", "primary 27: PRG 13 USE / G1 15 USE", f"source SHA256: {g1_review_sha256}"),
        ("Manual paired result", "USE→USE 8 / USE→DNU 5", "DNU→USE 7 / DNU→DNU 7", "McNemar p=0.774414; not comprehensive"),
        ("Workbook sheets", "Arm Case Metrics / Paired Comparison", "Box001 Human Review / Failure Modes", "12-Gate / Gate migrations / Visual Review"),
        ("Primary evidence", "E194 three-arm metrics + E196 corrected overlay" if corrected_overlay else "e194_three_arm_case_metrics.tsv", "formula-derived paired sheet", "formula-derived gate sheets"),
    ]
    for row_index, values in enumerate(rows, 4):
        for col_index, value in enumerate(values, 1):
            cell = ws.cell(row_index, col_index, value)
            cell.font = Font(name="Arial", bold=col_index == 1)
            cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE if row_index % 2 == 0 else WHITE)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
    set_widths(ws, {1: 24, 2: 34, 3: 34, 4: 34})
    ws.freeze_panes = "A4"


def add_arm_cases(wb: Workbook, rows: list[dict[str, str]]) -> None:
    ws = wb.create_sheet("Arm Case Metrics")
    fields = ["arm", "case_id", "object_key", "source_exp", "retarget_variant_id"]
    fields += [key for key, _, _ in METRICS]
    fields += ["fall_flag", "numeric_release_pass_12gate"] + [f"{gate}_gate_pass" for gate in GATES]
    fields += ["numeric_failure_modes", "result_sha256", "scene_sha256"]
    labels = ["Arm", "Case", "Object", "Source exp", "Retarget variant"] + [label for _, label, _ in METRICS]
    labels += ["Fall flag", "Overall 12-gate pass"] + [f"{gate} pass" for gate in GATES]
    labels += ["Numeric failure modes", "Result SHA256", "Scene SHA256"]
    title(ws, "Arm-level case metrics", "Observed public-core values; source rows are hardcoded evidence, not formulas.", len(fields))
    for col, label in enumerate(labels, 1):
        ws.cell(4, col, label)
    header(ws, 4, len(fields))
    rank = {key: index for index, key in enumerate(C.OBJECT_ORDER)}
    arm_rank = {"noPRG": 0, "PRG": 1, "G1": 2}
    for row_index, raw in enumerate(sorted(rows, key=lambda row: (rank[row["object_key"]], row["case_id"], arm_rank[row["arm"]])), 5):
        for col, field in enumerate(fields, 1):
            value: Any = raw.get(field, "")
            if field in {key for key, _, _ in METRICS} or field == "fall_flag":
                value = number(value)
            elif field == "numeric_release_pass_12gate" or field.endswith("_pass"):
                value = truth(value)
            ws.cell(row_index, col, value)
            ws.cell(row_index, col).font = Font(name="Arial", size=9)
        for col in range(6, 6 + len(METRICS)):
            ws.cell(row_index, col).number_format = "0.0000"
    finish_table(ws, 4, ws.max_row, len(fields), "ArmCaseMetrics")
    set_widths(ws, {1: 11, 2: 36, 3: 12, 4: 12, 5: 18, len(fields)-1: 18, len(fields): 18}, 14)


def add_paired(
    wb: Workbook,
    rows: list[dict[str, str]],
    box001_use_rows: dict[str, dict[str, str]],
    authority_sha256: str,
    g1_reviews: dict[str, dict[str, str]],
    g1_review_sha256: str,
) -> None:
    ws = wb.create_sheet("Paired Comparison")
    labels = ["Comparison", "Before", "After", "Case", "Object"]
    for _, label, direction in METRICS:
        arrow = "↑ better" if direction == "higher" else "↓ better"
        labels += [f"Before {label}", f"After {label}", f"Delta {label} ({arrow})"]
    labels += ["Before 12-gate pass", "After 12-gate pass",
               "Before numeric failure modes (blank=PASS)", "After numeric failure modes (blank=PASS)",
               "PRG authority scope", "PRG authority source", "PRG authority source SHA256",
               "PRG authority status", "PRG authoritative decision",
               "PRG in approved USE set", "PRG membership rule",
               "PRG source exp", "PRG authority row ref",
               "G1 review source", "G1 review source SHA256", "G1 review status",
               "G1 manual decision", "G1 quality label", "G1 failure taxonomy",
               "G1 review note", "G1 reviewer", "G1 reviewed at", "G1 paired video",
               "PRG→G1 manual migration", "Box001 analysis inclusion", "Overall migration"]
    title(ws, "Same-case paired comparisons",
          "Delta = after - before; green=improvement, red=regression. PRG uses E173 box001 RL-export membership authority; G1 uses the completed E194 28-case review. Primary analysis excludes box001_20231023_110_p1.", len(labels))
    for col, label in enumerate(labels, 1):
        ws.cell(4, col, label)
    header(ws, 4, len(labels))
    for row_index, raw in enumerate(rows, 5):
        authority_row = box001_use_rows.get(raw["case_id"])
        if raw["object_key"] == "box001":
            in_use_set = authority_row is not None
            prg_decision = "USE" if in_use_set else "DO_NOT_USE"
            g1_review = g1_reviews[raw["case_id"]]
            g1_decision = g1_review["manual_use_decision"]
            authority_values: list[Any] = [
                "E173_BOX001_FINAL_RL_EXPORT",
                "E173 box001_user_approved_source_rows.tsv",
                authority_sha256,
                "AUTHORITATIVE_FINAL" if in_use_set else "AUTHORITATIVE_COMPLEMENT",
                prg_decision,
                in_use_set,
                ("case_id present in approved source rows" if in_use_set else
                 "case_id absent from approved source rows within frozen E194 box001 28-case universe"),
                "E173",
                authority_row.get("authority_ref", "") if authority_row else "",
            ]
            g1_values: list[Any] = [
                "E194 full_g1_expansion/user_manual_review_filled.tsv",
                g1_review_sha256,
                *[g1_review.get(field, "") for field in G1_REVIEW_FIELDS],
                f"{prg_decision}_TO_{g1_decision}",
                "EXCLUDED_BY_USER" if raw["case_id"] == EXCLUDED_BOX001_CASE else "INCLUDED_PRIMARY_27",
            ]
        else:
            authority_values = [
                "NOT_APPLICABLE_BOX001_AUTHORITY",
                "E173 box001_user_approved_source_rows.tsv",
                authority_sha256,
                "NOT_APPLICABLE",
                "",
                "",
                "box001-only authority; no decision imputed for box023/box021",
                "",
                "",
            ]
            g1_values = [
                "E194 full_g1_expansion/user_manual_review_filled.tsv",
                g1_review_sha256,
                "NOT_APPLICABLE",
                "", "", "", "", "", "", "",
                "NOT_APPLICABLE",
                "NOT_APPLICABLE_BOX001_REVIEW",
            ]
        base_values: list[Any] = [raw["comparison"], raw["before_arm"], raw["after_arm"], raw["case_id"], raw["object_key"]]
        for metric, _, _ in METRICS:
            base_values += [number(raw[f"before_{metric}"]), number(raw[f"after_{metric}"]), None]
        base_values += [truth(raw["before_numeric_release_pass_12gate"]), truth(raw["after_numeric_release_pass_12gate"]),
                        raw.get("before_numeric_failure_modes", ""), raw.get("after_numeric_failure_modes", ""),
                        *authority_values,
                        *g1_values,
                        None]
        for col, value in enumerate(base_values, 1):
            ws.cell(row_index, col, value)
            ws.cell(row_index, col).font = Font(name="Arial", size=9)
        for metric_index in range(len(METRICS)):
            before_col = 6 + metric_index * 3
            after_col = before_col + 1
            delta_col = before_col + 2
            before_letter, after_letter = get_column_letter(before_col), get_column_letter(after_col)
            ws.cell(row_index, delta_col,
                    f'=IF(OR({before_letter}{row_index}="",{after_letter}{row_index}=""),"",{after_letter}{row_index}-{before_letter}{row_index})')
            for col in (before_col, after_col):
                ws.cell(row_index, col).number_format = "0.0000"
            ws.cell(row_index, delta_col).number_format = "0.0000;(0.0000);-"
        before_gate_col = 6 + len(METRICS) * 3
        after_gate_col, migration_col = before_gate_col + 1, len(labels)
        b, a = get_column_letter(before_gate_col), get_column_letter(after_gate_col)
        ws.cell(row_index, migration_col, f'=IF(AND({b}{row_index},{a}{row_index}),"PASS_TO_PASS",IF(AND({b}{row_index},NOT({a}{row_index})),"PASS_TO_FAIL",IF({a}{row_index},"FAIL_TO_PASS","FAIL_TO_FAIL")))')
        authority_scope_col = before_gate_col + 4
        authority_decision_col = authority_scope_col + 4
        g1_source_col = authority_scope_col + 9
        g1_decision_col = g1_source_col + 3
        manual_migration_col = g1_source_col + 10
        for col in range(authority_scope_col, migration_col):
            ws.cell(row_index, col).alignment = Alignment(vertical="top", wrap_text=True)
        decision = ws.cell(row_index, authority_decision_col).value
        if decision == "USE":
            ws.cell(row_index, authority_decision_col).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
        elif decision == "DO_NOT_USE":
            ws.cell(row_index, authority_decision_col).fill = PatternFill("solid", fgColor=LIGHT_RED)
        g1_decision = ws.cell(row_index, g1_decision_col).value
        if g1_decision == "USE":
            ws.cell(row_index, g1_decision_col).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
        elif g1_decision == "DO_NOT_USE":
            ws.cell(row_index, g1_decision_col).fill = PatternFill("solid", fgColor=LIGHT_RED)
        manual_migration = ws.cell(row_index, manual_migration_col).value
        migration_fill = {
            "DO_NOT_USE_TO_USE": LIGHT_GREEN,
            "USE_TO_DO_NOT_USE": LIGHT_RED,
            "USE_TO_USE": LIGHT_GREEN,
            "DO_NOT_USE_TO_DO_NOT_USE": LIGHT_GRAY,
        }.get(manual_migration)
        if migration_fill:
            ws.cell(row_index, manual_migration_col).fill = PatternFill("solid", fgColor=migration_fill)
    finish_table(ws, 4, ws.max_row, len(labels), "PairedComparison")
    for metric_index, (_, _, direction) in enumerate(METRICS):
        delta_letter = get_column_letter(8 + metric_index * 3)
        add_delta_quality_formatting(ws, f"{delta_letter}5:{delta_letter}{ws.max_row}", direction)
    before_gate_col = 6 + len(METRICS) * 3
    authority_scope_col = before_gate_col + 4
    g1_source_col = authority_scope_col + 9
    set_widths(ws, {1: 18, 2: 10, 3: 10, 4: 36, 5: 12,
                    before_gate_col + 2: 36, before_gate_col + 3: 36,
                    authority_scope_col: 32, authority_scope_col + 1: 38,
                    authority_scope_col + 2: 66, authority_scope_col + 3: 24,
                    authority_scope_col + 4: 24, authority_scope_col + 5: 22,
                    authority_scope_col + 6: 54, authority_scope_col + 7: 16,
                    authority_scope_col + 8: 52,
                    g1_source_col: 44, g1_source_col + 1: 66, g1_source_col + 2: 18,
                    g1_source_col + 3: 20, g1_source_col + 4: 22,
                    g1_source_col + 5: 24, g1_source_col + 6: 42,
                    g1_source_col + 7: 18, g1_source_col + 8: 26,
                    g1_source_col + 9: 42, g1_source_col + 10: 28,
                    g1_source_col + 11: 26, len(labels): 18}, 15)


def add_box001_human_review(
    wb: Workbook,
    paired_rows: list[dict[str, str]],
    box001_use_rows: dict[str, dict[str, str]],
    g1_reviews: dict[str, dict[str, str]],
    g1_review_sha256: str,
    corrected_overlay: bool = False,
) -> None:
    ws = wb.create_sheet("Box001 Human Review")
    title(
        ws,
        "E194 G1 box001 completed human review",
        ("Primary scope excludes box001_20231023_110_p1. PRG membership authority: E173 final box001 RL export. "
         "G1 authority: completed 28-case E194 review. Corrected E196 overlay rows are not re-labeled as human-reviewed."
         if corrected_overlay else
         "Primary scope excludes box001_20231023_110_p1. PRG membership authority: E173 final box001 RL export. G1 authority: completed 28-case E194 review."),
        21,
    )
    summary_labels = ["Primary 27-case summary", "Count / value", "Rate / secondary value", "Interpretation"]
    for col, label in enumerate(summary_labels, 1):
        ws.cell(4, col, label)
    header(ws, 4, len(summary_labels))
    detail_start, detail_end = 21, 48
    summary = [
        ("Included cases", f'=COUNTIF($B${detail_start}:$B${detail_end},"INCLUDED_PRIMARY_27")', "", f"Excludes {EXCLUDED_BOX001_CASE}"),
        ("PRG USE", f'=COUNTIFS($B${detail_start}:$B${detail_end},"INCLUDED_PRIMARY_27",$C${detail_start}:$C${detail_end},"USE")', '=B6/B5', "E173 authority membership"),
        ("G1 USE", f'=COUNTIFS($B${detail_start}:$B${detail_end},"INCLUDED_PRIMARY_27",$D${detail_start}:$D${detail_end},"USE")', '=B7/B5', "Completed E194 review"),
        ("Net G1−PRG USE", '=B7-B6', '=C7-C6', "Positive is favorable"),
        ("USE→USE", f'=COUNTIFS($B${detail_start}:$B${detail_end},"INCLUDED_PRIMARY_27",$E${detail_start}:$E${detail_end},"USE_TO_USE")', '=B9/B5', "Stable usable"),
        ("USE→DNU", f'=COUNTIFS($B${detail_start}:$B${detail_end},"INCLUDED_PRIMARY_27",$E${detail_start}:$E${detail_end},"USE_TO_DO_NOT_USE")', '=B10/B5', "Manual regression"),
        ("DNU→USE", f'=COUNTIFS($B${detail_start}:$B${detail_end},"INCLUDED_PRIMARY_27",$E${detail_start}:$E${detail_end},"DO_NOT_USE_TO_USE")', '=B11/B5', "Manual recovery"),
        ("DNU→DNU", f'=COUNTIFS($B${detail_start}:$B${detail_end},"INCLUDED_PRIMARY_27",$E${detail_start}:$E${detail_end},"DO_NOT_USE_TO_DO_NOT_USE")', '=B12/B5', "Stable unusable"),
        ("Agreement", '=B9+B12', '=B13/B5', "Same decision"),
        ("Churn", '=B10+B11', '=B14/B5', "Decision changed"),
        ("Exact McNemar p", '=MIN(1,2*BINOMDIST(MIN(B10,B11),B10+B11,0.5,TRUE))', "", "Two-sided exact paired test"),
        ("All-28 sensitivity", "PRG USE 13 / G1 USE 15", "USE→DNU 5 / DNU→USE 7", "Excluded case is DNU→DNU; conclusion unchanged"),
        ("Integrated judgment", "NOT_COMPREHENSIVE_IMPROVEMENT", "", "Strong z/3D/penetration gains, but manual churn and orientation/contact regressions remain"),
    ]
    for row_index, values in enumerate(summary, 5):
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
            ws.cell(row_index, col).font = Font(name="Arial", size=9, bold=col == 1)
            ws.cell(row_index, col).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row_index, col).fill = PatternFill("solid", fgColor=LIGHT_BLUE if row_index % 2 else WHITE)
    for row_index in (6, 7, 8, 9, 10, 11, 12, 13, 14):
        ws.cell(row_index, 3).number_format = "0.0%"
    ws.cell(15, 2).number_format = "0.000000"

    labels = [
        "Case", "Primary analysis inclusion", "PRG authoritative decision", "G1 manual decision",
        "Manual migration", "G1 quality", "G1 review status", "G1 reviewer", "G1 reviewed at",
        "PRG 12-gate pass", "G1 12-gate pass", "PRG numeric failure modes", "G1 numeric failure modes",
        "Delta object z MAE cm (↓ better)", "Delta object 3D pos cm (↓ better)",
        "Delta object orientation deg (↓ better)", "Delta 3mm contact frac (↑ better)",
        "Delta raw contact frac (↑ better)", "Delta hand penetration frac (↓ better)",
        "Delta leg penetration frac (↓ better)", "G1 review source SHA256",
    ]
    for col, label in enumerate(labels, 1):
        ws.cell(20, col, label)
    header(ws, 20, len(labels))
    selected = [row for row in paired_rows if row["comparison"] == "PRG_to_G1" and row["object_key"] == "box001"]
    if len(selected) != 28:
        raise ValueError(f"box001 PRG_to_G1 rows={len(selected)} expected=28")
    paired_excel_rows = {
        row["case_id"]: index for index, row in enumerate(paired_rows, 5) if row["comparison"] == "PRG_to_G1"
    }
    key_metrics = (
        "track_obj_z_abs_err_cm_mean",
        "track_obj_pos_err_cm_mean",
        "track_obj_ori_err_deg_mean",
        "hand_object_physics_contact_3mm_in_mask_frac",
        "hand_object_physics_contact_in_mask_frac",
        "hand_object_physics_penetration_3mm_frame_frac",
        "leg_penetration_frac",
    )
    metric_delta_cols = {key: 8 + index * 3 for index, (key, _, _) in enumerate(METRICS)}
    for row_index, raw in enumerate(sorted(selected, key=lambda row: row["case_id"]), detail_start):
        case_id = raw["case_id"]
        prg_decision = "USE" if case_id in box001_use_rows else "DO_NOT_USE"
        review = g1_reviews[case_id]
        g1_decision = review["manual_use_decision"]
        source_row = paired_excel_rows[case_id]
        values: list[Any] = [
            case_id,
            "EXCLUDED_BY_USER" if case_id == EXCLUDED_BOX001_CASE else "INCLUDED_PRIMARY_27",
            prg_decision,
            g1_decision,
            f"{prg_decision}_TO_{g1_decision}",
            review["manual_quality_label"],
            review["user_manual_review_status"],
            review["manual_reviewer"],
            review["manual_reviewed_at"],
            f"='Paired Comparison'!{get_column_letter(6 + len(METRICS) * 3)}{source_row}",
            f"='Paired Comparison'!{get_column_letter(7 + len(METRICS) * 3)}{source_row}",
            f"='Paired Comparison'!{get_column_letter(8 + len(METRICS) * 3)}{source_row}",
            f"='Paired Comparison'!{get_column_letter(9 + len(METRICS) * 3)}{source_row}",
        ]
        values += [
            f"='Paired Comparison'!{get_column_letter(metric_delta_cols[key])}{source_row}"
            for key in key_metrics
        ]
        values.append(g1_review_sha256)
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
            ws.cell(row_index, col).font = Font(name="Arial", size=9)
            ws.cell(row_index, col).alignment = Alignment(vertical="top", wrap_text=True)
        for col in range(14, 21):
            ws.cell(row_index, col).number_format = "0.0000;(0.0000);-"
        ws.cell(row_index, 3).fill = PatternFill("solid", fgColor=LIGHT_GREEN if prg_decision == "USE" else LIGHT_RED)
        ws.cell(row_index, 4).fill = PatternFill("solid", fgColor=LIGHT_GREEN if g1_decision == "USE" else LIGHT_RED)
        migration_fill = {
            "DO_NOT_USE_TO_USE": LIGHT_GREEN,
            "USE_TO_DO_NOT_USE": LIGHT_RED,
            "USE_TO_USE": LIGHT_GREEN,
            "DO_NOT_USE_TO_DO_NOT_USE": LIGHT_GRAY,
        }[f"{prg_decision}_TO_{g1_decision}"]
        ws.cell(row_index, 5).fill = PatternFill("solid", fgColor=migration_fill)
        if case_id == EXCLUDED_BOX001_CASE:
            ws.cell(row_index, 2).fill = PatternFill("solid", fgColor=LIGHT_GRAY)
    finish_table(ws, 20, detail_end, len(labels), "Box001HumanReview")
    for col, direction in zip(range(14, 21), ("lower", "lower", "lower", "higher", "higher", "lower", "lower")):
        add_delta_quality_formatting(ws, f"{get_column_letter(col)}{detail_start}:{get_column_letter(col)}{detail_end}", direction)
    set_widths(ws, {1: 36, 2: 24, 3: 24, 4: 20, 5: 28, 6: 20, 7: 18, 8: 16, 9: 26,
                    10: 18, 11: 18, 12: 40, 13: 40, 21: 66}, 18)


def add_prg_g1_failure_modes(wb: Workbook, rows: list[dict[str, str]]) -> None:
    ws = wb.create_sheet("PRG-G1 Failure Modes")
    labels = ["Case", "Object", "PRG 12-gate pass", "G1 12-gate pass",
              "PRG numeric failure modes", "G1 numeric failure modes",
              "PRG failure count", "G1 failure count", "Delta failure count (↓ better)", "Overall migration"]
    title(ws, "PRG / G1 numeric failure modes",
          "72 same-case comparisons. Modes are comma-separated failed numeric gates; PASS means no failure mode. Delta = G1 count - PRG count.", len(labels))
    for col, label in enumerate(labels, 1):
        ws.cell(4, col, label)
    header(ws, 4, len(labels))
    selected = [row for row in rows if row["comparison"] == "PRG_to_G1"]
    if len(selected) != C.N_CASES:
        raise ValueError(f"PRG_to_G1 rows={len(selected)} expected={C.N_CASES}")
    for row_index, raw in enumerate(selected, 5):
        prg_modes = raw.get("before_numeric_failure_modes", "") or "PASS"
        g1_modes = raw.get("after_numeric_failure_modes", "") or "PASS"
        values = [raw["case_id"], raw["object_key"], truth(raw["before_numeric_release_pass_12gate"]),
                  truth(raw["after_numeric_release_pass_12gate"]), prg_modes, g1_modes]
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
            ws.cell(row_index, col).font = Font(name="Arial", size=9)
            ws.cell(row_index, col).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row_index, 7, f'=IF(E{row_index}="PASS",0,LEN(E{row_index})-LEN(SUBSTITUTE(E{row_index},",",""))+1)')
        ws.cell(row_index, 8, f'=IF(F{row_index}="PASS",0,LEN(F{row_index})-LEN(SUBSTITUTE(F{row_index},",",""))+1)')
        ws.cell(row_index, 9, f'=H{row_index}-G{row_index}')
        ws.cell(row_index, 10,
                f'=IF(AND(C{row_index},D{row_index}),"PASS_TO_PASS",IF(AND(C{row_index},NOT(D{row_index})),"PASS_TO_FAIL",IF(D{row_index},"FAIL_TO_PASS","FAIL_TO_FAIL")))')
        for col in range(7, 11):
            ws.cell(row_index, col).font = Font(name="Arial", size=9)
        ws.cell(row_index, 9).number_format = "0;(0);-"
        ws.cell(row_index, 5).fill = PatternFill("solid", fgColor=LIGHT_GREEN if prg_modes == "PASS" else LIGHT_RED)
        ws.cell(row_index, 6).fill = PatternFill("solid", fgColor=LIGHT_GREEN if g1_modes == "PASS" else LIGHT_RED)
    finish_table(ws, 4, ws.max_row, len(labels), "PRGG1FailureModes")
    add_delta_quality_formatting(ws, f"I5:I{ws.max_row}", "lower")
    set_widths(ws, {1: 36, 2: 12, 3: 16, 4: 16, 5: 42, 6: 42, 7: 16, 8: 16, 9: 24, 10: 18})


def add_by_object(wb: Workbook) -> None:
    ws = wb.create_sheet("By Object")
    labels = ["Comparison", "Before", "After", "Object", "n"]
    for _, label, direction in METRICS:
        arrow = "↑ better" if direction == "higher" else "↓ better"
        labels += [f"Before mean {label}", f"After mean {label}", f"Delta mean {label} ({arrow})"]
    labels += ["Before overall pass rate", "After overall pass rate", "Delta pass rate"]
    title(ws, "By-object formula summary", "All summary values are formulas over the Paired Comparison sheet. ALL is a case-weighted aggregate.", len(labels))
    for col, label in enumerate(labels, 1):
        ws.cell(4, col, label)
    header(ws, 4, len(labels))
    paired_last = 4 + C.N_CASES * 2
    source = "'Paired Comparison'!"
    for row_index, (comparison, before, after, object_key) in enumerate(
            [(comparison, before, after, object_key) for comparison, before, after in
             (("noPRG_to_PRG", "noPRG", "PRG"), ("PRG_to_G1", "PRG", "G1"))
             for object_key in (*C.OBJECT_ORDER, "ALL")], 5):
        ws.cell(row_index, 1, comparison); ws.cell(row_index, 2, before); ws.cell(row_index, 3, after); ws.cell(row_index, 4, object_key)
        object_filter = f',{source}$E$5:$E${paired_last},$D{row_index}' if object_key != "ALL" else ""
        ws.cell(row_index, 5, f'=COUNTIFS({source}$A$5:$A${paired_last},$A{row_index}{object_filter})')
        for metric_index in range(len(METRICS)):
            source_before_col = get_column_letter(6 + metric_index * 3)
            source_after_col = get_column_letter(7 + metric_index * 3)
            out_before_col = 6 + metric_index * 3
            out_after_col = out_before_col + 1
            out_delta_col = out_before_col + 2
            criteria = f'{source}$A$5:$A${paired_last},$A{row_index}{object_filter}'
            ws.cell(row_index, out_before_col,
                    f'=IFERROR(AVERAGEIFS({source}${source_before_col}$5:${source_before_col}${paired_last},{criteria}),"")')
            ws.cell(row_index, out_after_col,
                    f'=IFERROR(AVERAGEIFS({source}${source_after_col}$5:${source_after_col}${paired_last},{criteria}),"")')
            out_before_letter, out_after_letter = get_column_letter(out_before_col), get_column_letter(out_after_col)
            ws.cell(row_index, out_delta_col,
                    f'=IF(OR({out_before_letter}{row_index}="",{out_after_letter}{row_index}=""),"",{out_after_letter}{row_index}-{out_before_letter}{row_index})')
            for col in (out_before_col, out_after_col):
                ws.cell(row_index, col).number_format = "0.0000"
            ws.cell(row_index, out_delta_col).number_format = "0.0000;(0.0000);-"
        before_gate_source = get_column_letter(6 + len(METRICS) * 3)
        after_gate_source = get_column_letter(7 + len(METRICS) * 3)
        before_col, after_col, delta_col = len(labels) - 2, len(labels) - 1, len(labels)
        gate_criteria = f'{source}$A$5:$A${paired_last},$A{row_index}{object_filter}'
        ws.cell(row_index, before_col, f'=IFERROR(COUNTIFS({gate_criteria},{source}${before_gate_source}$5:${before_gate_source}${paired_last},TRUE)/$E{row_index},"")')
        ws.cell(row_index, after_col, f'=IFERROR(COUNTIFS({gate_criteria},{source}${after_gate_source}$5:${after_gate_source}${paired_last},TRUE)/$E{row_index},"")')
        ws.cell(row_index, delta_col, f'={get_column_letter(after_col)}{row_index}-{get_column_letter(before_col)}{row_index}')
        for col in (before_col, after_col, delta_col):
            ws.cell(row_index, col).number_format = "0.0%;[Red](0.0%);-"
        for col in range(1, len(labels) + 1):
            ws.cell(row_index, col).font = Font(name="Arial", size=9, bold=object_key == "ALL")
            if object_key == "ALL":
                ws.cell(row_index, col).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    finish_table(ws, 4, ws.max_row, len(labels), "ByObjectSummary")
    for metric_index, (_, _, direction) in enumerate(METRICS):
        delta_letter = get_column_letter(8 + metric_index * 3)
        add_delta_quality_formatting(ws, f"{delta_letter}5:{delta_letter}{ws.max_row}", direction)
    gate_delta_letter = get_column_letter(len(labels))
    add_delta_quality_formatting(ws, f"{gate_delta_letter}5:{gate_delta_letter}{ws.max_row}", "higher")
    set_widths(ws, {1: 18, 2: 10, 3: 10, 4: 12, 5: 8}, 16)


def add_gate_summary(wb: Workbook, migrations: list[dict[str, str]]) -> None:
    detail = wb.create_sheet("Gate Migrations")
    labels = ["Comparison", "Case", "Object", "Gate", "Before pass", "After pass", "Migration"]
    title(detail, "12-gate case migrations", "Two comparisons × 72 cases × 12 gates = 1,728 observed rows.", len(labels))
    for col, label in enumerate(labels, 1): detail.cell(4, col, label)
    header(detail, 4, len(labels))
    for row_index, raw in enumerate(migrations, 5):
        values = [raw["comparison"], raw["case_id"], raw["object_key"], raw["gate"], truth(raw["before_pass"]),
                  truth(raw["after_pass"]), raw["migration"]]
        for col, value in enumerate(values, 1):
            detail.cell(row_index, col, value); detail.cell(row_index, col).font = Font(name="Arial", size=9)
    finish_table(detail, 4, detail.max_row, len(labels), "GateMigrations")
    set_widths(detail, {1: 18, 2: 36, 3: 12, 4: 18, 5: 13, 6: 13, 7: 18})

    ws = wb.create_sheet("Gate Summary")
    labels = ["Comparison", "Object", "Gate", "n", "Before pass", "After pass", "Before rate", "After rate", "Delta rate",
              "PASS→FAIL", "FAIL→PASS"]
    title(ws, "Gate formula summary", "Counts and rates are Excel formulas over Gate Migrations.", len(labels))
    for col, label in enumerate(labels, 1): ws.cell(4, col, label)
    header(ws, 4, len(labels))
    source = "'Gate Migrations'!"; last = 4 + C.N_CASES * 2 * len(GATES)
    combos = [(comparison, object_key, gate) for comparison in ("noPRG_to_PRG", "PRG_to_G1")
              for object_key in (*C.OBJECT_ORDER, "ALL") for gate in GATES]
    for row_index, (comparison, object_key, gate) in enumerate(combos, 5):
        ws.cell(row_index, 1, comparison); ws.cell(row_index, 2, object_key); ws.cell(row_index, 3, gate)
        obj_criteria = f',{source}$C$5:$C${last},$B{row_index}' if object_key != "ALL" else ""
        common = f'{source}$A$5:$A${last},$A{row_index},{source}$D$5:$D${last},$C{row_index}{obj_criteria}'
        ws.cell(row_index, 4, f'=COUNTIFS({common})')
        ws.cell(row_index, 5, f'=COUNTIFS({common},{source}$E$5:$E${last},TRUE)')
        ws.cell(row_index, 6, f'=COUNTIFS({common},{source}$F$5:$F${last},TRUE)')
        ws.cell(row_index, 7, f'=IFERROR(E{row_index}/D{row_index},"")')
        ws.cell(row_index, 8, f'=IFERROR(F{row_index}/D{row_index},"")')
        ws.cell(row_index, 9, f'=H{row_index}-G{row_index}')
        ws.cell(row_index, 10, f'=COUNTIFS({common},{source}$G$5:$G${last},"PASS_TO_FAIL")')
        ws.cell(row_index, 11, f'=COUNTIFS({common},{source}$G$5:$G${last},"FAIL_TO_PASS")')
        for col in (7, 8, 9): ws.cell(row_index, col).number_format = "0.0%;[Red](0.0%);-"
        for col in range(1, len(labels) + 1): ws.cell(row_index, col).font = Font(name="Arial", size=9, bold=object_key == "ALL")
    finish_table(ws, 4, ws.max_row, len(labels), "GateSummary")
    add_delta_quality_formatting(ws, f"I5:I{ws.max_row}", "higher")
    set_widths(ws, {1: 18, 2: 12, 3: 18}, 13)


def add_12gate_comparison(wb: Workbook) -> None:
    ws = wb.create_sheet("12-Gate Comparison")
    labels = ["Object", "Gate", "n", "noPRG pass", "noPRG rate", "PRG pass", "PRG rate", "G1 pass", "G1 rate",
              "PRG−noPRG (pp)", "G1−PRG (pp)", "G1−noPRG (pp)",
              "noPRG→PRG P→F", "noPRG→PRG F→P", "noPRG→PRG McNemar p",
              "PRG→G1 P→F", "PRG→G1 F→P", "PRG→G1 McNemar p"]
    title(ws, "12-gate noPRG / PRG / G1 comparison", "Pass counts/rates and paired flips are formulas over arm-level and migration evidence. pp = percentage points.", len(labels))
    for col, label in enumerate(labels, 1):
        ws.cell(4, col, label)
    header(ws, 4, len(labels))
    case_source, case_last = "'Arm Case Metrics'!", 4 + C.N_CASES * 3
    migration_source, migration_last = "'Gate Migrations'!", 4 + C.N_CASES * 2 * len(GATES)
    gate_start_col = 5 + len(METRICS) + 3
    gate_col = {gate: get_column_letter(gate_start_col + index) for index, gate in enumerate(GATES)}
    row_index = 5
    for object_key in (*C.OBJECT_ORDER, "ALL"):
        for gate in GATES:
            ws.cell(row_index, 1, object_key)
            ws.cell(row_index, 2, gate)
            object_case = f',{case_source}$C$5:$C${case_last},$A{row_index}' if object_key != "ALL" else ""
            ws.cell(row_index, 3, f'=COUNTIFS({case_source}$A$5:$A${case_last},"PRG"{object_case})')
            for arm, count_col, rate_col in (("noPRG", 4, 5), ("PRG", 6, 7), ("G1", 8, 9)):
                ws.cell(row_index, count_col,
                        f'=COUNTIFS({case_source}$A$5:$A${case_last},"{arm}",{case_source}${gate_col[gate]}$5:${gate_col[gate]}${case_last},TRUE{object_case})')
                ws.cell(row_index, rate_col, f'=IFERROR({get_column_letter(count_col)}{row_index}/$C{row_index},"")')
                ws.cell(row_index, rate_col).number_format = "0.0%;[Red](0.0%);-"
            ws.cell(row_index, 10, f'=(G{row_index}-E{row_index})*100')
            ws.cell(row_index, 11, f'=(I{row_index}-G{row_index})*100')
            ws.cell(row_index, 12, f'=(I{row_index}-E{row_index})*100')
            object_migration = f',{migration_source}$C$5:$C${migration_last},$A{row_index}' if object_key != "ALL" else ""
            for comparison, p2f_col, f2p_col, p_col in (("noPRG_to_PRG", 13, 14, 15), ("PRG_to_G1", 16, 17, 18)):
                common = (f'{migration_source}$A$5:$A${migration_last},"{comparison}",'
                          f'{migration_source}$D$5:$D${migration_last},$B{row_index}{object_migration}')
                ws.cell(row_index, p2f_col, f'=COUNTIFS({common},{migration_source}$G$5:$G${migration_last},"PASS_TO_FAIL")')
                ws.cell(row_index, f2p_col, f'=COUNTIFS({common},{migration_source}$G$5:$G${migration_last},"FAIL_TO_PASS")')
                a, b = get_column_letter(p2f_col), get_column_letter(f2p_col)
                # BINOMDIST is the Excel-2007-compatible spelling understood by both
                # Excel and LibreOffice; BINOM.DIST becomes #NAME? in the latter's
                # OOXML recalculation path.
                ws.cell(row_index, p_col, f'=MIN(1,2*BINOMDIST(MIN({a}{row_index},{b}{row_index}),{a}{row_index}+{b}{row_index},0.5,TRUE))')
                ws.cell(row_index, p_col).number_format = "0.000000"
            for col in (10, 11, 12):
                ws.cell(row_index, col).number_format = "0.0;[Red](0.0);-"
            for col in range(1, len(labels) + 1):
                ws.cell(row_index, col).font = Font(name="Arial", size=9, bold=object_key == "ALL")
                if object_key == "ALL":
                    ws.cell(row_index, col).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            row_index += 1
    finish_table(ws, 4, ws.max_row, len(labels), "TwelveGateComparison")
    for col in (10, 11, 12):
        letter = get_column_letter(col)
        add_delta_quality_formatting(ws, f"{letter}5:{letter}{ws.max_row}", "higher")
    set_widths(ws, {1: 12, 2: 18, 3: 8, 10: 17, 11: 15, 12: 17, 15: 20, 18: 18}, 13)


def add_12gate_overall(wb: Workbook) -> None:
    ws = wb.create_sheet("12-Gate Overall")
    labels = ["Object", "n", "Gate decisions", "noPRG gate pass", "noPRG gate rate", "noPRG strict pass", "noPRG strict rate",
              "PRG gate pass", "PRG gate rate", "PRG strict pass", "PRG strict rate",
              "G1 gate pass", "G1 gate rate", "G1 strict pass", "G1 strict rate",
              "PRG−noPRG gate (pp)", "G1−PRG gate (pp)", "G1−noPRG gate (pp)",
              "PRG−noPRG strict (pp)", "G1−PRG strict (pp)", "G1−noPRG strict (pp)"]
    title(ws, "12-gate aggregate and strict all-gate pass", "Gate rate pools n×12 decisions; strict rate requires a case to pass all 12 gates.", len(labels))
    for col, label in enumerate(labels, 1):
        ws.cell(4, col, label)
    header(ws, 4, len(labels))
    comparison_source = "'12-Gate Comparison'!"; comparison_last = 4 + 4 * len(GATES)
    strict_col = get_column_letter(5 + len(METRICS) + 2)
    case_source, case_last = "'Arm Case Metrics'!", 4 + C.N_CASES * 3
    for row_index, object_key in enumerate((*C.OBJECT_ORDER, "ALL"), 5):
        ws.cell(row_index, 1, object_key)
        object_case = f',{case_source}$C$5:$C${case_last},$A{row_index}' if object_key != "ALL" else ""
        ws.cell(row_index, 2, f'=COUNTIFS({case_source}$A$5:$A${case_last},"PRG"{object_case})')
        ws.cell(row_index, 3, f'=B{row_index}*12')
        object_comparison = f',{comparison_source}$A$5:$A${comparison_last},$A{row_index}'
        for arm, gate_count_col, gate_rate_col, strict_count_col, strict_rate_col, comparison_count_col in (
                ("noPRG", 4, 5, 6, 7, "D"), ("PRG", 8, 9, 10, 11, "F"), ("G1", 12, 13, 14, 15, "H")):
            ws.cell(row_index, gate_count_col,
                    f'=SUMIFS({comparison_source}${comparison_count_col}$5:${comparison_count_col}${comparison_last}{object_comparison})')
            ws.cell(row_index, gate_rate_col, f'={get_column_letter(gate_count_col)}{row_index}/$C{row_index}')
            ws.cell(row_index, strict_count_col,
                    f'=COUNTIFS({case_source}$A$5:$A${case_last},"{arm}",{case_source}${strict_col}$5:${strict_col}${case_last},TRUE{object_case})')
            ws.cell(row_index, strict_rate_col, f'={get_column_letter(strict_count_col)}{row_index}/$B{row_index}')
            for col in (gate_rate_col, strict_rate_col):
                ws.cell(row_index, col).number_format = "0.0%;[Red](0.0%);-"
        for col, formula in ((16, f'=(I{row_index}-E{row_index})*100'), (17, f'=(M{row_index}-I{row_index})*100'),
                             (18, f'=(M{row_index}-E{row_index})*100'), (19, f'=(K{row_index}-G{row_index})*100'),
                             (20, f'=(O{row_index}-K{row_index})*100'), (21, f'=(O{row_index}-G{row_index})*100')):
            ws.cell(row_index, col, formula); ws.cell(row_index, col).number_format = "0.0;[Red](0.0);-"
        for col in range(1, len(labels) + 1):
            ws.cell(row_index, col).font = Font(name="Arial", size=9, bold=object_key == "ALL")
            if object_key == "ALL":
                ws.cell(row_index, col).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    finish_table(ws, 4, ws.max_row, len(labels), "TwelveGateOverall")
    for col in range(16, 22):
        letter = get_column_letter(col)
        add_delta_quality_formatting(ws, f"{letter}5:{letter}{ws.max_row}", "higher")
    set_widths(ws, {1: 12, 2: 8, 3: 14}, 16)


def add_visual_review(wb: Workbook) -> None:
    ws = wb.create_sheet("Visual Review")
    title(ws, "Mandatory paired visual review", "Concrete grasp/lift/carry/place observations are sourced from the experiment log.", 8)
    labels = ["Case", "Object", "Selection reason", "noPRG observation", "PRG observation", "G1 observation", "Verdict", "Frame evidence"]
    for col, label in enumerate(labels, 1): ws.cell(4, col, label)
    header(ws, 4, len(labels))
    review_path = EVAL / "e194_three_arm_visual_review.tsv"
    rows = C.read_tsv(review_path) if review_path.is_file() else []
    for row_index, raw in enumerate(rows, 5):
        values = [raw.get("case_id", ""), raw.get("object_key", ""), raw.get("selection_reason", ""),
                  raw.get("noprg_observation", ""), raw.get("prg_observation", ""), raw.get("g1_observation", ""),
                  raw.get("verdict", ""), raw.get("frame_evidence", "")]
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value); ws.cell(row_index, col).font = Font(name="Arial", size=9)
            ws.cell(row_index, col).alignment = Alignment(vertical="top", wrap_text=True)
    if rows:
        finish_table(ws, 4, ws.max_row, len(labels), "VisualReview")
    ws.freeze_panes = "A5"
    set_widths(ws, {1: 36, 2: 12, 3: 32, 4: 44, 5: 44, 6: 44, 7: 20, 8: 48})


def main() -> int:
    summary = C.read_tsv(EVAL / "e194_three_arm_case_metrics.tsv")
    paired = C.read_tsv(EVAL / "e194_three_arm_paired_deltas.tsv")
    migrations = C.read_tsv(EVAL / "e194_three_arm_gate_migrations.tsv")
    gate_metrics = C.read_tsv(EVAL / "e194_three_arm_12gate_by_object.tsv")
    gate_overall = C.read_tsv(EVAL / "e194_three_arm_12gate_overall.tsv")
    authority_rows = C.read_tsv(E173_BOX001_USE_AUTHORITY)
    g1_review_rows = C.read_tsv(E194_G1_REVIEW)
    required_authority_fields = {
        "case_id", "object_key", "manual_use_decision", "source_exp_id", "authority_ref"
    }
    if not authority_rows or not required_authority_fields.issubset(authority_rows[0]):
        raise SystemExit(f"invalid E173 box001 USE authority: {E173_BOX001_USE_AUTHORITY}")
    box001_use_rows = {row["case_id"]: row for row in authority_rows}
    if len(box001_use_rows) != len(authority_rows):
        raise SystemExit("duplicate case_id in E173 box001 USE authority")
    if any(
        row["object_key"] != "box001"
        or row["manual_use_decision"] != "USE"
        or row["source_exp_id"] != "E173"
        for row in authority_rows
    ):
        raise SystemExit("E173 box001 USE authority contains an invalid object/decision/source row")
    required_g1_review_fields = {"case_id", *G1_REVIEW_FIELDS}
    if not g1_review_rows or not required_g1_review_fields.issubset(g1_review_rows[0]):
        raise SystemExit(f"invalid E194 G1 box001 review source: {E194_G1_REVIEW}")
    g1_reviews = {row["case_id"]: row for row in g1_review_rows}
    if len(g1_reviews) != len(g1_review_rows):
        raise SystemExit("duplicate case_id in E194 G1 box001 review source")
    import json
    payload = json.loads((EVAL / "e194_three_arm_summary.json").read_text(encoding="utf-8"))
    if len(summary) != 216 or len(paired) != 144 or len(migrations) != 1728 or len(gate_metrics) != 48 or len(gate_overall) != 4 or payload.get("status") != "pass":
        raise SystemExit("three-arm inputs are incomplete")
    paired_cases = {row["case_id"] for row in paired}
    box001_cases = {row["case_id"] for row in paired if row["object_key"] == "box001"}
    if len(paired_cases) != 72 or len(box001_cases) != 28:
        raise SystemExit("unexpected E194 paired-case universe")
    if not set(box001_use_rows).issubset(box001_cases):
        raise SystemExit("E173 box001 USE authority is not a subset of the E194 box001 universe")
    if len(box001_use_rows) != 13:
        raise SystemExit("unexpected E173 box001 approved USE authority cardinality")
    if set(g1_reviews) != box001_cases or len(g1_reviews) != 28:
        raise SystemExit("E194 G1 review source does not exactly cover the box001 28-case universe")
    if any(
        row["user_manual_review_status"] != "reviewed"
        or row["manual_use_decision"] not in {"USE", "DO_NOT_USE"}
        or row["manual_quality_label"] not in {"CLEAN", "MINOR_ACCEPTABLE", "UNUSABLE"}
        for row in g1_review_rows
    ):
        raise SystemExit("E194 G1 review source contains an incomplete or invalid decision")
    authority_sha256 = hashlib.sha256(E173_BOX001_USE_AUTHORITY.read_bytes()).hexdigest()
    g1_review_sha256 = hashlib.sha256(E194_G1_REVIEW.read_bytes()).hexdigest()
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    add_readme(wb, payload, authority_sha256, g1_review_sha256)
    add_arm_cases(wb, summary)
    add_paired(wb, paired, box001_use_rows, authority_sha256, g1_reviews, g1_review_sha256)
    add_box001_human_review(wb, paired, box001_use_rows, g1_reviews, g1_review_sha256)
    add_prg_g1_failure_modes(wb, paired)
    add_by_object(wb)
    add_gate_summary(wb, migrations)
    add_12gate_comparison(wb)
    add_12gate_overall(wb)
    add_visual_review(wb)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font.name != "Arial":
                    cell.font = Font(name="Arial", size=cell.font.sz or 10, bold=cell.font.bold,
                                     italic=cell.font.italic, color=cell.font.color)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(C.rel(OUTPUT))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
