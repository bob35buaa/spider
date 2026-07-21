#!/usr/bin/env python3
"""Create a formatted workbook for the current E168 available-case evaluation."""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


REPO = Path(__file__).resolve().parents[5]
DEFAULT_EVAL_DIR = (
    REPO
    / "workspace/core4d/results/E168/s6_downstream/cem/eval/e167a_aligned_available"
)
DEFAULT_OUTPUT_NAME = "E168_E167A_aligned_available_case_metrics.xlsx"

NAVY = "1F4E78"
TEAL = "0F6B78"
GREEN = "548235"
LIGHT_GREEN = "E2F0D9"
RED = "C00000"
LIGHT_RED = "FCE4D6"
AMBER = "BF8F00"
LIGHT_AMBER = "FFF2CC"
BLUE_GRAY = "D9E2F3"
LIGHT_BLUE = "DDEBF7"
GRAY = "666666"
LIGHT_GRAY = "E7E6E6"
WHITE = "FFFFFF"

PATH_FIELDS = {
    "qpos_path",
    "scene_xml",
    "z_reference_path",
    "result_npz",
    "outdir_npz",
    "config_act",
    "video",
    "trajectory",
    "contact_mask",
    "target_task",
    "target_scene",
    "scene_act",
    "override_path",
    "log",
}

OVERVIEW_COLUMNS = [
    ("case_id", "case"),
    ("object_key", "物体"),
    ("source_person", "person"),
    ("preferred_pool", "GPU池"),
    ("gpu_id", "GPU"),
    ("retarget_variant_id", "retarget"),
    ("release_status", "numeric状态"),
    ("failure_modes", "失败模式"),
    ("manual_use_decision", "人工决定"),
    ("manual_quality_label", "人工质量"),
    ("visual_gate_pass", "视觉通过"),
    ("manual_review_status", "核验状态"),
    ("manual_review_note", "人工备注"),
    ("numeric_release_pass", "numeric pass"),
    ("tracking_gate_pass", "tracking"),
    ("holosoma_z_gate_pass", "z-only"),
    ("contact_gate_pass", "contact"),
    ("release_gate_applicable", "release适用"),
    ("release_gate_pass", "release"),
    ("release_gate_status", "release状态"),
    ("penetration_gate_pass", "penetration"),
    ("lower_body_gate_pass", "lower body"),
    ("fall_flag", "fall"),
    ("track_pelvis_z_err_terminal_m", "pelvis终点误差(m)"),
    ("body_z_err_p95_m", "body-z p95(m)"),
    ("body_z_err_peak_m", "body-z peak(m)"),
    ("track_joint_err_deg_mean", "关节误差(°)"),
    ("track_eef_pos_err_cm_mean", "EEF误差(cm)"),
    ("track_obj_pos_err_cm_mean", "物体误差(cm)"),
    ("hand_object_physics_contact_in_mask_frac", "raw接触"),
    ("hand_object_release_false_contact_3mm_frac", "release误接触3mm"),
    ("hand_object_physics_penetration_3mm_frame_frac", "物理穿透3mm"),
    ("leg_penetration_frac", "下肢干涉"),
    ("qpos_jerk_l2_p95", "qpos jerk p95"),
    ("foot_slip_max_m", "foot slip max(m)"),
    ("video", "视频"),
]


def read_tsv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if not path.is_file() or path.stat().st_size == 0:
        return [], []
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        return list(reader.fieldnames or []), list(reader)


def excel_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (bool, int, float)):
        return value
    text = str(value).strip()
    if text == "":
        return None
    if text.lower() == "true":
        return True
    if text.lower() == "false":
        return False
    if re.fullmatch(r"[-+]?\d+", text):
        try:
            return int(text)
        except ValueError:
            return text
    if re.fullmatch(r"[-+]?(?:\d+\.\d*|\.\d+)(?:[eE][-+]?\d+)?", text):
        try:
            number = float(text)
        except ValueError:
            return text
        return number if math.isfinite(number) else None
    return text


def repo_path(raw: str | Path) -> Path:
    path = Path(raw)
    if path.exists():
        return path.resolve()
    text = str(raw)
    for marker in ("example_datasets/", "workspace/", "logs/"):
        if marker in text:
            return REPO / (marker + text.split(marker, 1)[1])
    return path if path.is_absolute() else REPO / path


def header_style(cell, fill: str = NAVY) -> None:
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Arial", size=10, color=WHITE, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(bottom=Side(style="thin", color=WHITE))


def body_style(cell) -> None:
    cell.font = Font(name="Arial", size=10, color="000000")
    cell.alignment = Alignment(vertical="center", wrap_text=False)
    cell.border = Border(bottom=Side(style="hair", color=LIGHT_GRAY))
    if isinstance(cell.value, float):
        cell.number_format = "0.000"


def add_path_hyperlinks(ws, fields: list[str], header_row: int) -> None:
    for col_idx, field in enumerate(fields, start=1):
        if field not in PATH_FIELDS:
            continue
        for row_idx in range(header_row + 1, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if not isinstance(cell.value, str) or not cell.value:
                continue
            target = repo_path(cell.value)
            if target.exists():
                cell.hyperlink = str(target)
                cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")


def set_widths(ws, fields: list[str], header_row: int, max_width: int = 28) -> None:
    for col_idx, field in enumerate(fields, start=1):
        values = [ws.cell(row=header_row, column=col_idx).value]
        values.extend(
            ws.cell(row=row_idx, column=col_idx).value
            for row_idx in range(header_row + 1, min(ws.max_row, header_row + 80) + 1)
        )
        width = max(9, min(max(len(str(value)) for value in values if value is not None) + 2, max_width))
        if field in {"case_id", "sequence_key"}:
            width = max(width, 30)
        elif field in {"release_status", "failure_modes", "release_gate_status"}:
            width = max(width, 28)
        elif field in PATH_FIELDS:
            width = max(width, 34)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def write_table(
    ws,
    fields: list[str],
    rows: list[dict[str, Any]],
    *,
    labels: dict[str, str] | None = None,
    header_row: int = 1,
    header_fill: str = NAVY,
    max_width: int = 28,
) -> None:
    labels = labels or {}
    for col_idx, field in enumerate(fields, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=labels.get(field, field))
        header_style(cell, header_fill)
    ws.row_dimensions[header_row].height = 42
    for row_idx, row in enumerate(rows, start=header_row + 1):
        for col_idx, field in enumerate(fields, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=excel_value(row.get(field)))
            body_style(cell)
        if row_idx % 2 == 1:
            for cell in ws[row_idx]:
                cell.fill = PatternFill("solid", fgColor="F7F9FC")
    ws.freeze_panes = f"A{header_row + 1}"
    if rows and fields:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(fields))}{header_row + len(rows)}"
    ws.sheet_view.showGridLines = False
    add_path_hyperlinks(ws, fields, header_row)
    set_widths(ws, fields, header_row, max_width=max_width)


def style_gate_cells(ws, header_row: int, fields: list[str]) -> None:
    gate_fields = {
        "numeric_release_pass",
        "tracking_gate_pass",
        "holosoma_z_gate_pass",
        "contact_gate_pass",
        "release_gate_pass",
        "penetration_gate_pass",
        "lower_body_gate_pass",
        "visual_gate_pass",
    }
    for field in gate_fields:
        if field not in fields:
            continue
        col = fields.index(field) + 1
        letter = get_column_letter(col)
        region = f"{letter}{header_row + 1}:{letter}{ws.max_row}"
        ws.conditional_formatting.add(
            region,
            CellIsRule(operator="equal", formula=["TRUE"], fill=PatternFill("solid", fgColor=LIGHT_GREEN)),
        )
        ws.conditional_formatting.add(
            region,
            CellIsRule(operator="equal", formula=["FALSE"], fill=PatternFill("solid", fgColor=LIGHT_RED)),
        )


def write_overview(ws, rows: list[dict[str, str]], summary: dict[str, Any]) -> None:
    fields = [field for field, _ in OVERVIEW_COLUMNS]
    labels = dict(OVERVIEW_COLUMNS)
    last_col = get_column_letter(len(fields))
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = summary.get("scope_label") or "E168 E167A-aligned 已有 case 全面评测"
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(f"A2:{last_col}2")
    counts = summary["counts"]
    ws["A2"] = (
        f"快照 {summary['generated_at']} | 完整产物 {counts['evaluated']}/{counts['manifest_rows']} | "
        f"numeric pass {counts['numeric_pass']} | 人工可用 {counts['manual_use']} | "
        f"人工拒绝 {counts['manual_do_not_use']} | errors {counts['evaluation_errors']} | "
        f"metric {summary['metric_standard_id']}"
    )
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color=GRAY)

    summary_headers = ["门控", "通过", "适用/总数", "失败或N/A", "阈值/说明"]
    for col_idx, label in enumerate(summary_headers, start=1):
        header_style(ws.cell(row=4, column=col_idx, value=label), TEAL)

    table_header = 17
    first_data = table_header + 1
    last_data = table_header + len(rows)
    col_map = {field: get_column_letter(idx + 1) for idx, field in enumerate(fields)}
    case_range = f"${col_map['case_id']}${first_data}:${col_map['case_id']}${last_data}"

    gate_rows = [
        ("全部 numeric gates", "numeric_release_pass", "<= all hard gates"),
        ("tracking", "tracking_gate_pass", "no fall + pelvis terminal <=0.08m"),
        ("fixed-reference z-only", "holosoma_z_gate_pass", "body-z p95 <=0.20m"),
        ("raw contact", "contact_gate_pass", "in-mask physics contact >=0.50"),
        ("release", "release_gate_pass", "false contact 3mm <=0.30"),
        (">3mm penetration", "penetration_gate_pass", "frame frac <=0.30"),
        ("lower body", "lower_body_gate_pass", "leg/object interference <=0.10"),
        ("no fall", "fall_flag", "fall_flag=false"),
    ]
    for row_idx, (label, field, note) in enumerate(gate_rows, start=5):
        ws.cell(row=row_idx, column=1, value=label)
        col = col_map[field]
        if field == "release_gate_pass":
            applicable_col = col_map["release_gate_applicable"]
            ws.cell(row=row_idx, column=2, value=f'=COUNTIF(${col}${first_data}:${col}${last_data},TRUE)')
            ws.cell(
                row=row_idx,
                column=3,
                value=f'=COUNTIF(${applicable_col}${first_data}:${applicable_col}${last_data},TRUE)',
            )
            ws.cell(
                row=row_idx,
                column=4,
                value=f'=COUNTA({case_range})-COUNTIF(${applicable_col}${first_data}:${applicable_col}${last_data},TRUE)',
            )
        elif field == "fall_flag":
            ws.cell(row=row_idx, column=2, value=f'=COUNTIF(${col}${first_data}:${col}${last_data},FALSE)')
            ws.cell(row=row_idx, column=3, value=f'=COUNTA({case_range})')
            ws.cell(row=row_idx, column=4, value=f'=COUNTIF(${col}${first_data}:${col}${last_data},TRUE)')
        else:
            ws.cell(row=row_idx, column=2, value=f'=COUNTIF(${col}${first_data}:${col}${last_data},TRUE)')
            ws.cell(row=row_idx, column=3, value=f'=COUNTA({case_range})')
            ws.cell(row=row_idx, column=4, value=f'=COUNTIF(${col}${first_data}:${col}${last_data},FALSE)')
        ws.cell(row=row_idx, column=5, value=note)
        for cell in ws[row_idx][:5]:
            body_style(cell)
        ws.cell(row=row_idx, column=2).font = Font(name="Arial", size=10, bold=True, color=GREEN)
        ws.cell(row=row_idx, column=4).font = Font(name="Arial", size=10, bold=True, color=RED)

    manual_rows = [
        (13, "人工已核验", "manual_review_status", "reviewed", "用户逐一视频核验"),
        (14, "人工可用", "manual_use_decision", "USE", "包含无问题及小瑕疵可接受"),
        (15, "人工不可用", "manual_use_decision", "DO_NOT_USE", "存在大问题，禁止使用"),
    ]
    for row_idx, label, field, criterion, note in manual_rows:
        col = col_map[field]
        ws.cell(row=row_idx, column=1, value=label)
        ws.cell(
            row=row_idx,
            column=2,
            value=f'=COUNTIF(${col}${first_data}:${col}${last_data},"{criterion}")',
        )
        ws.cell(row=row_idx, column=3, value=f'=COUNTA({case_range})')
        ws.cell(row=row_idx, column=4, value=f"={get_column_letter(3)}{row_idx}-{get_column_letter(2)}{row_idx}")
        ws.cell(row=row_idx, column=5, value=note)
        for cell in ws[row_idx][:5]:
            body_style(cell)
        ws.cell(row=row_idx, column=2).font = Font(name="Arial", size=10, bold=True, color=GREEN)

    write_table(
        ws,
        fields,
        rows,
        labels=labels,
        header_row=table_header,
        header_fill=NAVY,
        max_width=34,
    )
    style_gate_cells(ws, table_header, fields)
    ws.freeze_panes = f"A{first_data}"
    ws.auto_filter.ref = f"A{table_header}:{last_col}{last_data}"
    ws.column_dimensions["A"].width = 31
    ws.column_dimensions["G"].width = 34
    ws.column_dimensions["H"].width = 30


def write_manual_review(
    ws,
    fields: list[str],
    rows: list[dict[str, str]],
    scope_label: str = "",
) -> None:
    if not fields:
        fields = [
            "case_id",
            "manual_review_status",
            "manual_use_decision",
            "manual_quality_label",
            "visual_gate_pass",
            "manual_review_note",
            "reviewer",
            "reviewed_at",
            "source_workbook",
        ]
    last_col = get_column_letter(len(fields))
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = (
        f"{scope_label}：范围内人工核验"
        if scope_label
        else "E168 Box021 用户人工核验"
    )
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=TEAL)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    write_table(ws, fields, rows, header_row=6, header_fill=TEAL, max_width=50)
    first_data = 7
    last_data = 6 + len(rows)
    decision_col = get_column_letter(fields.index("manual_use_decision") + 1)
    status_col = get_column_letter(fields.index("manual_review_status") + 1)
    if rows:
        summary_rows = [
            ("已核验", f'=COUNTIF(${status_col}${first_data}:${status_col}${last_data},"reviewed")'),
            ("人工可用", f'=COUNTIF(${decision_col}${first_data}:${decision_col}${last_data},"USE")'),
            ("人工不可用", f'=COUNTIF(${decision_col}${first_data}:${decision_col}${last_data},"DO_NOT_USE")'),
        ]
    else:
        summary_rows = [("已核验", "=0"), ("人工可用", "=0"), ("人工不可用", "=0")]
    for col_idx, (label, formula) in enumerate(summary_rows, start=1):
        header_style(ws.cell(row=3, column=col_idx, value=label), TEAL)
        cell = ws.cell(row=4, column=col_idx, value=formula)
        body_style(cell)
        cell.font = Font(name="Arial", size=12, bold=True, color=GREEN if col_idx < 3 else RED)
        cell.alignment = Alignment(horizontal="center")

    if rows:
        for row_idx in range(first_data, last_data + 1):
            decision = ws.cell(row=row_idx, column=fields.index("manual_use_decision") + 1)
            quality = ws.cell(row=row_idx, column=fields.index("manual_quality_label") + 1)
            fill = LIGHT_GREEN if decision.value == "USE" else LIGHT_RED
            decision.fill = PatternFill("solid", fgColor=fill)
            quality.fill = PatternFill("solid", fgColor=fill)
    ws.freeze_panes = "A7"


def write_notes(ws, summary: dict[str, Any], eval_dir: Path, output: Path) -> None:
    counts = summary["counts"]
    rows = [
        {"项": "范围标签", "说明": summary.get("scope_label") or "E168 available-case evaluation"},
        {"项": "范围", "说明": f"当前可评测产物 {counts['evaluated']}/{counts['manifest_rows']}；未就绪 {counts['not_ready']}。"},
        {"项": "case IDs", "说明": ", ".join(summary.get("scope_case_ids", []))},
        {"项": "numeric pass", "说明": f"{counts['numeric_pass']}/{counts['evaluated']}；数值门槛与人工结论分别保留，不用数值结果覆盖用户核验。"},
        {"项": "人工核验", "说明": f"已核验 {counts['manual_reviewed']} 条；USE {counts['manual_use']} 条；DO_NOT_USE {counts['manual_do_not_use']} 条；其余 {counts['manual_pending']} 条待核验。"},
        {"项": "核心标准", "说明": summary["metric_standard_id"]},
        {"项": "z-only 参考", "说明": "四个腕踝 body 对 manifest 固定 kinematic trajectory；body_z_err_p95_m <=0.20m。peak 仅保留为诊断指标。"},
        {"项": "门槛", "说明": "raw contact >=0.50；release false contact <=0.30；>3mm penetration frame frac <=0.30；lower-body interference <=0.10。"},
        {"项": "legacy zgate caveat", "说明": "E167 legacy intra-tick reference 只作诊断；不能替代固定 kinematic reference。"},
        {"项": "relative gate", "说明": summary["relative_gate_status"]},
        {"项": "release N/A", "说明": "最后 reference-contact 帧之后没有 trailing frame 时，显式标为 NOT_APPLICABLE_NO_RELEASE_WINDOW；不填 0、不判失败。"},
        {"项": "visual", "说明": "人工决定来自用户逐一核验；未核验的新产物保持 PENDING_MANUAL_REVIEW。"},
        {"项": "完整指标", "说明": "“完整指标”sheet 原样收录 e168_case_metrics.tsv 的全部列，并转换为 Excel 数值/布尔类型。"},
        {"项": "源目录", "说明": str(eval_dir)},
        {"项": "工作簿", "说明": str(output)},
        {"项": "生成脚本", "说明": "workspace/core4d/scripts/eval/reports/gen_E168_available_metrics_xlsx.py"},
    ]
    write_table(ws, ["项", "说明"], rows, header_fill=TEAL, max_width=100)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(vertical="top", wrap_text=True)


def validate_structure(path: Path, expected_rows: int, expected_metrics: int) -> None:
    wb = load_workbook(path, data_only=False, read_only=False)
    expected_sheets = [
        "门控总览",
        "人工核验",
        "完整指标",
        "指标统计",
        "分组统计",
        "最差样本",
        "未就绪",
        "评测错误",
        "评测快照",
        "说明",
    ]
    if wb.sheetnames != expected_sheets:
        raise RuntimeError(f"unexpected sheets: {wb.sheetnames}")
    if wb["完整指标"].max_row != expected_rows + 1:
        raise RuntimeError("完整指标 row count mismatch")
    if wb["完整指标"].max_column != expected_metrics:
        raise RuntimeError("完整指标 column count mismatch")
    formula_count = 0
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
                if isinstance(cell.value, str) and cell.value in {
                    "#REF!",
                    "#DIV/0!",
                    "#VALUE!",
                    "#N/A",
                    "#NAME?",
                }:
                    bad.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    if formula_count == 0:
        raise RuntimeError("expected formulas in gate summary")
    if bad:
        raise RuntimeError("formula errors found: " + "; ".join(bad[:20]))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--eval-dir", type=Path, default=DEFAULT_EVAL_DIR)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    eval_dir = args.eval_dir.expanduser().resolve()
    output = (
        args.output.expanduser().resolve()
        if args.output
        else eval_dir / DEFAULT_OUTPUT_NAME
    )
    summary = json.loads((eval_dir / "summary.json").read_text(encoding="utf-8"))

    case_fields, case_rows = read_tsv(eval_dir / "e168_case_metrics.tsv")
    metric_fields, metric_rows = read_tsv(eval_dir / "e168_metric_summary.tsv")
    group_fields, group_rows = read_tsv(eval_dir / "e168_group_summary.tsv")
    ranking_fields, ranking_rows = read_tsv(eval_dir / "e168_worst_case_rankings.tsv")
    not_ready_fields, not_ready_rows = read_tsv(eval_dir / "e168_not_ready.tsv")
    error_fields, error_rows = read_tsv(eval_dir / "e168_evaluation_errors.tsv")
    snapshot_fields, snapshot_rows = read_tsv(eval_dir / "evaluated_manifest_snapshot.tsv")
    manual_fields, manual_rows = read_tsv(eval_dir / "e168_manual_review_snapshot.tsv")

    if len(case_rows) != int(summary["counts"]["evaluated"]):
        raise SystemExit("case metrics row count does not match summary.json")
    if not case_fields:
        raise SystemExit("e168_case_metrics.tsv has no fields")

    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    ws = wb.create_sheet("门控总览")
    ws.sheet_properties.tabColor = NAVY
    write_overview(ws, case_rows, summary)

    ws = wb.create_sheet("人工核验")
    ws.sheet_properties.tabColor = TEAL
    write_manual_review(
        ws,
        manual_fields,
        manual_rows,
        scope_label=summary.get("scope_label", ""),
    )

    ws = wb.create_sheet("完整指标")
    ws.sheet_properties.tabColor = TEAL
    write_table(ws, case_fields, case_rows, max_width=34)
    style_gate_cells(ws, 1, case_fields)

    ws = wb.create_sheet("指标统计")
    ws.sheet_properties.tabColor = GREEN
    write_table(ws, metric_fields, metric_rows, header_fill=GREEN)

    ws = wb.create_sheet("分组统计")
    ws.sheet_properties.tabColor = GREEN
    write_table(ws, group_fields, group_rows, header_fill=GREEN, max_width=36)

    ws = wb.create_sheet("最差样本")
    ws.sheet_properties.tabColor = RED
    write_table(ws, ranking_fields, ranking_rows, header_fill=RED, max_width=36)

    ws = wb.create_sheet("未就绪")
    ws.sheet_properties.tabColor = AMBER
    write_table(ws, not_ready_fields, not_ready_rows, header_fill=AMBER, max_width=36)

    ws = wb.create_sheet("评测错误")
    ws.sheet_properties.tabColor = RED
    if not error_fields:
        error_fields = ["status", "detail"]
        error_rows = [{"status": "NO_EVALUATION_ERRORS", "detail": "0 errors in this snapshot"}]
    write_table(ws, error_fields, error_rows, header_fill=RED, max_width=60)

    ws = wb.create_sheet("评测快照")
    ws.sheet_properties.tabColor = GRAY
    write_table(ws, snapshot_fields, snapshot_rows, header_fill=GRAY, max_width=36)

    ws = wb.create_sheet("说明")
    ws.sheet_properties.tabColor = TEAL
    write_notes(ws, summary, eval_dir, output)

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    validate_structure(output, len(case_rows), len(case_fields))
    print(output)
    print(
        f"case_rows={len(case_rows)} metric_columns={len(case_fields)} "
        f"sheets={len(wb.sheetnames)}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
