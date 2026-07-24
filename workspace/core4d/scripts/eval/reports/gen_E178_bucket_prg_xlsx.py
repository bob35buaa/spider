#!/usr/bin/env python3
"""Generate the E178 bucket PRG full-validation workbook."""

from __future__ import annotations

import argparse
import csv
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_EVAL = Path("workspace/core4d/results/E178/s6_downstream/eval/full")
DEFAULT_KEYFRAMES = Path(
    "workspace/core4d/results/E178/s6_downstream/render/full/"
    "keyframes_all_20260724_165645"
)
REVIEW_FIELDS = (
    "case_id",
    "user_manual_review_status",
    "manual_use_decision",
    "manual_quality_label",
    "manual_failure_taxonomy",
    "manual_review_note",
    "manual_reviewer",
    "manual_reviewed_at",
    "paired_video",
)
WORST_DIRECTIONS = {
    "body_z_err_p95_m": "high",
    "track_pelvis_z_err_terminal_m": "high",
    "hand_object_physics_contact_in_mask_frac": "low",
    "hand_object_physics_contact_3mm_in_mask_frac": "low",
    "hand_object_release_false_contact_3mm_frac": "high",
    "hand_object_physics_penetration_3mm_frame_frac": "high",
    "leg_penetration_frac": "high",
    "leg_near_2cm_frac": "high",
    "leg_object_physics_contact_frac": "high",
    "track_root_pos_err_cm_mean": "high",
    "track_root_ori_err_deg_mean": "high",
    "track_eef_pos_err_cm_mean": "high",
    "track_eef_ori_err_deg_mean": "high",
    "track_obj_pos_err_cm_mean": "high",
    "track_obj_ori_err_deg_mean": "high",
    "qpos_accel_l2_p95": "high",
    "qpos_jerk_l2_p95": "high",
    "trackbody_jerk_p95": "high",
    "ankle_jerk_p95": "high",
    "obj_speed_max": "high",
    "foot_slip_max_m": "high",
}


def read_tsv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if not path.is_file() or path.stat().st_size <= 1:
        return [], []
    with path.open("r", encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream, delimiter="\t")
        return list(reader.fieldnames or []), list(reader)


def value(raw: str) -> Any:
    text = str(raw)
    if text.lower() == "true":
        return True
    if text.lower() == "false":
        return False
    if text == "":
        return None
    try:
        return float(text) if any(char in text.lower() for char in (".", "e")) else int(text)
    except ValueError:
        return text


def finite(raw: str) -> float | None:
    try:
        result = float(raw)
    except (TypeError, ValueError):
        return None
    return result if result == result and abs(result) != float("inf") else None


def style_sheet(sheet: Any) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        width = min(45, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[letter].width = width
        for cell in column:
            cell.font = Font(
                name="Arial",
                size=9,
                bold=cell.row == 1,
                color="FFFFFF" if cell.row == 1 else "000000",
            )
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")


def add_rows_sheet(
    workbook: Workbook,
    title: str,
    fields: list[str],
    rows: list[dict[str, Any]],
) -> tuple[dict[str, int], int]:
    sheet = workbook.create_sheet(title)
    if not fields:
        sheet.append(["status"])
        sheet.append(["not_generated_or_empty"])
        style_sheet(sheet)
        return {"status": 1}, 0
    sheet.append(fields)
    for row in rows:
        sheet.append([value(row.get(field, "")) for field in fields])
    style_sheet(sheet)
    return {field: index + 1 for index, field in enumerate(fields)}, len(rows)


def col_range(sheet: str, column: int, last: int) -> str:
    letter = get_column_letter(column)
    return f"'{sheet}'!{letter}2:{letter}{last}"


def worst_rows(case_rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for metric, direction in WORST_DIRECTIONS.items():
        ranked = [
            (number, row)
            for row in case_rows
            if (number := finite(row.get(metric, ""))) is not None
        ]
        ranked.sort(key=lambda item: item[0], reverse=direction == "high")
        for rank, (number, row) in enumerate(ranked[:5], 1):
            output.append(
                {
                    "metric": metric,
                    "worse_direction": direction,
                    "rank": rank,
                    "case_id": row["case_id"],
                    "value": number,
                    "numeric_failure_modes": row.get("numeric_failure_modes", ""),
                    "video": row.get("video", ""),
                }
            )
    return output


def manual_rows(
    case_rows: list[dict[str, str]],
    filled: dict[str, dict[str, str]],
) -> list[dict[str, Any]]:
    output = []
    for case in case_rows:
        row = {field: "" for field in REVIEW_FIELDS}
        row.update(filled.get(case["case_id"], {}))
        row["case_id"] = case["case_id"]
        row["user_manual_review_status"] = row["user_manual_review_status"] or "pending"
        row["manual_use_decision"] = row["manual_use_decision"] or "PENDING"
        row["paired_video"] = row["paired_video"] or case.get("video", "")
        output.append(row)
    return output


def codex_rows(
    case_rows: list[dict[str, str]],
    baseline: dict[str, dict[str, str]],
    keyframes: Path,
) -> tuple[list[str], list[dict[str, Any]]]:
    fields = [
        "case_id",
        "codex_metric_verification_status",
        "codex_visual_spotcheck_status",
        "codex_visual_selection_reasons",
        "codex_visual_findings",
        "codex_visual_reviewer",
        "codex_visual_reviewed_at",
        "numeric_release_pass",
        "numeric_failure_modes",
        "e174_numeric_release_pass",
        "person",
        "retarget_variant_id",
        "object_key",
        "paired_video",
        "keyframe_sheet",
    ]
    reviewed_at = datetime.now().astimezone().isoformat(timespec="seconds")
    rows = []
    for case in case_rows:
        old = baseline.get(case["case_id"], {})
        rows.append(
            {
                "case_id": case["case_id"],
                "codex_metric_verification_status": "EVAL_COMPLETE",
                "codex_visual_spotcheck_status": "MIDPOINT_FRAME_QC_ONLY",
                "codex_visual_selection_reasons": f"object_group:{case['object_key']}",
                "codex_visual_findings": (
                    "Renderer/replay readable in grouped midpoint montage; "
                    "full-sequence user review pending"
                ),
                "codex_visual_reviewer": "codex",
                "codex_visual_reviewed_at": reviewed_at,
                "numeric_release_pass": case.get("numeric_release_pass", ""),
                "numeric_failure_modes": case.get("numeric_failure_modes", ""),
                "e174_numeric_release_pass": old.get("numeric_release_pass", ""),
                "person": case.get("person", ""),
                "retarget_variant_id": case.get("retarget_variant_id", ""),
                "object_key": case.get("object_key", ""),
                "paired_video": case.get("video", ""),
                "keyframe_sheet": str(
                    keyframes / f"{case['variant']}_mid.jpg"
                ),
            }
        )
    return fields, rows


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--eval-dir", type=Path, default=DEFAULT_EVAL)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--keyframes", type=Path, default=DEFAULT_KEYFRAMES)
    args = parser.parse_args()

    eval_dir = args.eval_dir
    output = args.output or eval_dir / "E178_buckets_prg_full_validation.xlsx"
    summary = json.loads((eval_dir / "summary.json").read_text(encoding="utf-8"))
    counts = summary["counts"]
    failures = summary["numeric_failure_counts"]

    case_fields, cases = read_tsv(eval_dir / "e178_case_metrics.tsv")
    paired_fields, paired = read_tsv(eval_dir / "e178_vs_e174_paired_deltas.tsv")
    group_fields, groups = read_tsv(eval_dir / "e178_group_summary.tsv")
    filled_fields, filled_rows = read_tsv(eval_dir / "user_manual_review_filled.tsv")
    filled = {row["case_id"]: row for row in filled_rows}
    baseline_fields, baseline_rows = read_tsv(
        Path("workspace/core4d/results/E174/s6_downstream/eval/full/e174_case_metrics.tsv")
    )
    baseline = {row["case_id"]: row for row in baseline_rows}
    not_ready_fields, not_ready = read_tsv(eval_dir / "e178_not_ready.tsv")
    error_fields, errors = read_tsv(eval_dir / "e178_evaluation_errors.tsv")

    if len(cases) != counts["evaluated"]:
        raise RuntimeError(
            f"Case Metrics row mismatch: {len(cases)} != {counts['evaluated']}"
        )
    if len(paired) != counts["paired_rows"]:
        raise RuntimeError(
            f"Paired Deltas row mismatch: {len(paired)} != {counts['paired_rows']}"
        )

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    mappings: dict[str, dict[str, int]] = {}
    row_counts: dict[str, int] = {}

    mappings["Case Metrics"], row_counts["Case Metrics"] = add_rows_sheet(
        workbook, "Case Metrics", case_fields, cases
    )
    mappings["Paired Deltas"], row_counts["Paired Deltas"] = add_rows_sheet(
        workbook, "Paired Deltas", paired_fields, paired
    )
    mappings["Group Summary"], row_counts["Group Summary"] = add_rows_sheet(
        workbook, "Group Summary", group_fields, groups
    )
    worst = worst_rows(cases)
    worst_fields = [
        "metric",
        "worse_direction",
        "rank",
        "case_id",
        "value",
        "numeric_failure_modes",
        "video",
    ]
    mappings["Worst Cases"], row_counts["Worst Cases"] = add_rows_sheet(
        workbook, "Worst Cases", worst_fields, worst
    )
    manual = manual_rows(cases, filled)
    mappings["Manual Review"], row_counts["Manual Review"] = add_rows_sheet(
        workbook, "Manual Review", list(REVIEW_FIELDS), manual
    )
    codex_fields, codex = codex_rows(cases, baseline, args.keyframes)
    mappings["Codex Verification"], row_counts["Codex Verification"] = add_rows_sheet(
        workbook, "Codex Verification", codex_fields, codex
    )
    mappings["Not Ready"], row_counts["Not Ready"] = add_rows_sheet(
        workbook, "Not Ready", not_ready_fields, not_ready
    )
    mappings["Eval Errors"], row_counts["Eval Errors"] = add_rows_sheet(
        workbook, "Eval Errors", error_fields, errors
    )

    required_case = (
        "case_id",
        "object_key",
        "numeric_release_pass",
        "leg_gate_health_pass",
    )
    required_pair = ("e174_numeric_release_pass", "e178_numeric_release_pass")
    missing = [field for field in required_case if field not in mappings["Case Metrics"]]
    missing += [field for field in required_pair if field not in mappings["Paired Deltas"]]
    if missing:
        raise RuntimeError(f"Workbook source fields missing: {missing}")

    case = mappings["Case Metrics"]
    paired_map = mappings["Paired Deltas"]
    review = mappings["Manual Review"]
    case_last = row_counts["Case Metrics"] + 1
    paired_last = row_counts["Paired Deltas"] + 1
    review_last = row_counts["Manual Review"] + 1
    cid = col_range("Case Metrics", case["case_id"], case_last)
    obj = col_range("Case Metrics", case["object_key"], case_last)
    npass = col_range("Case Metrics", case["numeric_release_pass"], case_last)
    gate = col_range("Case Metrics", case["leg_gate_health_pass"], case_last)
    old_pass = col_range(
        "Paired Deltas", paired_map["e174_numeric_release_pass"], paired_last
    )
    new_pass = col_range(
        "Paired Deltas", paired_map["e178_numeric_release_pass"], paired_last
    )
    reviewed = col_range(
        "Manual Review", review["user_manual_review_status"], review_last
    )
    use_decision = col_range(
        "Manual Review", review["manual_use_decision"], review_last
    )

    overview.append(
        [
            "E178 PRG Full Validation — bucket003 / bucket004 / bucket007",
            "Value",
            "Source / rule",
        ]
    )
    rows: list[tuple[str, Any, str]] = [
        ("Metric standard", summary["metric_standard_id"], "summary.json"),
        ("Evaluated rows", f"=COUNTA({cid})", "COUNTA Case Metrics case_id"),
        ("Numeric release pass", f"=COUNTIF({npass},TRUE)", "COUNTIF numeric_release_pass"),
        ("Numeric pass rate", "=B4/B3", "pass / evaluated"),
        (
            "— bucket003",
            f'=COUNTIFS({obj},"bucket003",{npass},TRUE)&" / "&COUNTIF({obj},"bucket003")',
            "per-object numeric pass",
        ),
        (
            "— bucket004",
            f'=COUNTIFS({obj},"bucket004",{npass},TRUE)&" / "&COUNTIF({obj},"bucket004")',
            "per-object numeric pass",
        ),
        (
            "— bucket007",
            f'=COUNTIFS({obj},"bucket007",{npass},TRUE)&" / "&COUNTIF({obj},"bucket007")',
            "per-object numeric pass",
        ),
        (
            "E174 fail → E178 pass",
            f"=COUNTIFS({old_pass},FALSE,{new_pass},TRUE)",
            "paired numeric transition",
        ),
        (
            "E174 pass → E178 fail",
            f"=COUNTIFS({old_pass},TRUE,{new_pass},FALSE)",
            "paired numeric transition",
        ),
        ("Fail: hand_penetration", failures.get("hand_penetration", 0), "summary.json"),
        ("Fail: lower_body", failures.get("lower_body", 0), "summary.json"),
        ("Fail: release", failures.get("release", 0), "summary.json"),
        ("Fail: contact", failures.get("contact", 0), "summary.json"),
        (
            "Fail: fall / body_z",
            f"{failures.get('fall', 0)} / {failures.get('body_z', 0)}",
            "summary.json",
        ),
        ("Fail: root_pos (>20cm)", failures.get("root_pos", 0), "summary.json"),
        ("Fail: root_ori (>20°)", failures.get("root_ori", 0), "summary.json"),
        ("Fail: hand_pos (>20cm)", failures.get("hand_pos", 0), "summary.json"),
        ("Fail: hand_ori (>20°)", failures.get("hand_ori", 0), "summary.json"),
        ("Fail: object_pos (>20cm)", failures.get("object_pos", 0), "summary.json"),
        ("Fail: object_ori (>10°)", failures.get("object_ori", 0), "summary.json"),
        ("Gate-health pass", f"=COUNTIF({gate},TRUE)", "COUNTIF leg_gate_health_pass"),
        (
            "User reviewed",
            f'=COUNTIF({reviewed},"reviewed")',
            "Manual Review authority",
        ),
        (
            "Manual USE",
            f'=COUNTIF({use_decision},"USE")',
            "Manual Review authority",
        ),
        (
            "Manual DO_NOT_USE",
            f'=COUNTIF({use_decision},"DO_NOT_USE")',
            "Manual Review authority",
        ),
        (
            "Machine recommendation",
            f'=IF(B23<{counts["evaluated"]},"PENDING_USER_REVIEW","USER_REVIEW_COMPLETE")',
            f"requires {counts['evaluated']} manual reviews",
        ),
        (
            "Final promotion decision",
            f'=IF(B23<{counts["evaluated"]},"PENDING_USER_REVIEW","USER_DECIDED")',
            "user authority",
        ),
        (
            "KEY FINDING",
            (
                "Tracking-gated numeric pass is "
                f"{counts['numeric_pass']}/{counts['evaluated']}"
            ),
            "includes root/hand/object pos+ori gates",
        ),
        ("Manifest SHA256", summary["manifest_sha256"], summary["manifest"]),
        ("Baseline SHA256", summary["baseline_sha256"], summary["baseline"]),
    ]
    for row in rows:
        overview.append(row)

    overview.freeze_panes = "A2"
    overview.column_dimensions["A"].width = 36
    overview.column_dimensions["B"].width = 34
    overview.column_dimensions["C"].width = 68
    overview.sheet_properties.pageSetUpPr.fitToPage = True
    overview.page_setup.orientation = "landscape"
    overview.page_setup.fitToWidth = 1
    overview.page_setup.fitToHeight = 1
    overview.print_area = f"A1:C{overview.max_row}"
    for cell in overview[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=12)
        cell.fill = PatternFill("solid", fgColor="17365D")
    for row in overview.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    overview["B5"].number_format = "0.0%"
    for cell in (overview["B5"], overview["B26"], overview["B27"], overview["B28"]):
        cell.fill = PatternFill("solid", fgColor="FFF2CC")

    workbook.move_sheet("Overview", -(len(workbook.sheetnames) - 1))
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(
        json.dumps(
            {
                "output": str(output),
                "sheets": workbook.sheetnames,
                "evaluated": counts["evaluated"],
                "numeric_pass": counts["numeric_pass"],
                "manual_review_rows": len(manual),
                "worst_case_rows": len(worst),
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
