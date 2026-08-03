#!/usr/bin/env python3
"""Generate the E187 versus E178 22-case paired evaluation workbook."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

E187_EVAL = Path("workspace/core4d/results/E187/s6_downstream/eval/full")
E187_MANIFEST = Path(
    "workspace/core4d/results/E187/s6_downstream/manifests/"
    "e187_full_evaluation_manifest.tsv"
)
E178_METRICS = Path(
    "workspace/core4d/results/E178/s6_downstream/eval/full/e178_case_metrics.tsv"
)
OUTPUT = E187_EVAL / "E187_vs_E178_paired_evaluation.xlsx"
PAIR_METRICS = {
    "body_z_err_p95_m": "lower",
    "track_pelvis_z_err_terminal_m": "lower",
    "hand_object_physics_contact_in_mask_frac": "higher",
    "hand_object_physics_contact_3mm_in_mask_frac": "higher",
    "hand_object_release_false_contact_3mm_frac": "lower",
    "hand_object_physics_penetration_3mm_frame_frac": "lower",
    "leg_penetration_frac": "lower",
    "leg_near_2cm_frac": "lower",
    "leg_object_physics_contact_frac": "lower",
    "track_root_pos_err_cm_mean": "lower",
    "track_root_ori_err_deg_mean": "lower",
    "track_eef_pos_err_cm_mean": "lower",
    "track_eef_ori_err_deg_mean": "lower",
    "track_obj_pos_err_cm_mean": "lower",
    "track_obj_ori_err_deg_mean": "lower",
    "qpos_accel_l2_p95": "lower",
    "qpos_jerk_l2_p95": "lower",
    "trackbody_jerk_p95": "lower",
    "ankle_jerk_p95": "lower",
    "obj_speed_max": "lower",
    "foot_slip_max_m": "lower",
}
GATES = (
    "fall",
    "body_z",
    "contact",
    "release",
    "hand_penetration",
    "lower_body",
    "root_pos",
    "root_ori",
    "hand_pos",
    "hand_ori",
    "object_pos",
    "object_ori",
)
GATE_THRESHOLDS = {
    "fall": "fall_flag = FALSE",
    "body_z": "body_z_err_p95_m <= 0.20 m",
    "contact": "contact_in_mask >= 0.50",
    "release": "release_false_3mm <= 0.30 or N/A",
    "hand_penetration": "hand_penetration_3mm <= 0.30",
    "lower_body": "leg_penetration <= 0.10",
    "root_pos": "root position mean <= 20 cm",
    "root_ori": "root orientation mean <= 20 deg",
    "hand_pos": "hand position mean <= 20 cm",
    "hand_ori": "hand orientation mean <= 20 deg",
    "object_pos": "object position mean <= 20 cm",
    "object_ori": "object orientation mean <= 10 deg",
}
OBJECT_METRICS = (
    "hand_object_physics_contact_in_mask_frac",
    "hand_object_physics_penetration_3mm_frame_frac",
    "leg_penetration_frac",
    "track_root_pos_err_cm_mean",
    "track_eef_pos_err_cm_mean",
    "track_obj_pos_err_cm_mean",
)

NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
GREEN = "E2F0D9"
RED = "FCE4D6"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"


def read_tsv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream, delimiter="\t")
        return list(reader.fieldnames or []), list(reader)


def value(raw: Any) -> Any:
    text = str(raw)
    if text.lower() == "true":
        return True
    if text.lower() == "false":
        return False
    if text == "":
        return None
    try:
        return float(text) if any(char in text.lower() for char in ".e") else int(text)
    except ValueError:
        return text


def style_sheet(sheet: Any, freeze: str = "A2") -> None:
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=9, color="000000")
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        sample = [str(cell.value or "") for cell in column[: min(len(column), 24)]]
        sheet.column_dimensions[letter].width = min(
            42, max(11, max(map(len, sample)) + 2)
        )


def add_data_sheet(
    workbook: Workbook,
    title: str,
    fields: list[str],
    rows: list[dict[str, Any]],
) -> tuple[dict[str, int], int]:
    sheet = workbook.create_sheet(title)
    sheet.append(fields)
    for row in rows:
        sheet.append([value(row.get(field, "")) for field in fields])
    style_sheet(sheet)
    return {field: index + 1 for index, field in enumerate(fields)}, len(rows) + 1


def col_range(sheet: str, column: int, last: int) -> str:
    letter = get_column_letter(column)
    return f"'{sheet}'!${letter}$2:${letter}${last}"


def lookup_formula(
    sheet: str,
    key_col: int,
    value_col: int,
    source_last: int,
    target_row: int,
) -> str:
    keys = col_range(sheet, key_col, source_last)
    values = col_range(sheet, value_col, source_last)
    return f"=INDEX({values},MATCH($A{target_row},{keys},0))"


def add_paired_sheet(
    workbook: Workbook,
    cases: list[dict[str, str]],
    e187_map: dict[str, int],
    e187_last: int,
    e178_map: dict[str, int],
    e178_last: int,
) -> tuple[dict[str, int], int]:
    sheet = workbook.create_sheet("Paired Comparison")
    fields = [
        "case_id",
        "object_key",
        "person",
        "E178 numeric pass",
        "E187 numeric pass",
        "Gate transition",
        "E178 failure modes",
        "E187 failure modes",
    ]
    for metric in PAIR_METRICS:
        fields.extend(
            (
                f"{metric} | E178",
                f"{metric} | E187",
                f"{metric} | Delta (E187-E178)",
                f"{metric} | Improvement",
            )
        )
    sheet.append(fields)
    mapping = {field: index + 1 for index, field in enumerate(fields)}
    for row_index, case in enumerate(cases, 2):
        sheet.cell(row_index, 1, case["case_id"])
        for target_col, field in ((2, "object_key"), (3, "person")):
            sheet.cell(
                row_index,
                target_col,
                lookup_formula(
                    "E187 Metrics",
                    e187_map["case_id"],
                    e187_map[field],
                    e187_last,
                    row_index,
                ),
            )
        for target_col, source, mapping_source, last, field in (
            (4, "E178 Baseline", e178_map, e178_last, "numeric_release_pass"),
            (5, "E187 Metrics", e187_map, e187_last, "numeric_release_pass"),
            (7, "E178 Baseline", e178_map, e178_last, "numeric_failure_modes"),
            (8, "E187 Metrics", e187_map, e187_last, "numeric_failure_modes"),
        ):
            sheet.cell(
                row_index,
                target_col,
                lookup_formula(
                    source,
                    mapping_source["case_id"],
                    mapping_source[field],
                    last,
                    row_index,
                ),
            )
        sheet.cell(
            row_index,
            6,
            f'=IF(D{row_index}=E{row_index},"UNCHANGED_"&IF(E{row_index},"PASS","FAIL"),'
            f'IF(E{row_index},"IMPROVED_TO_PASS","REGRESSED_TO_FAIL"))',
        )
        target_col = 9
        for metric, direction in PAIR_METRICS.items():
            sheet.cell(
                row_index,
                target_col,
                lookup_formula(
                    "E178 Baseline",
                    e178_map["case_id"],
                    e178_map[metric],
                    e178_last,
                    row_index,
                ),
            )
            sheet.cell(
                row_index,
                target_col + 1,
                lookup_formula(
                    "E187 Metrics",
                    e187_map["case_id"],
                    e187_map[metric],
                    e187_last,
                    row_index,
                ),
            )
            sheet.cell(
                row_index,
                target_col + 2,
                (
                    f"={get_column_letter(target_col + 1)}{row_index}"
                    f"-{get_column_letter(target_col)}{row_index}"
                ),
            )
            sign = "" if direction == "higher" else "-"
            sheet.cell(
                row_index,
                target_col + 3,
                f"={sign}{get_column_letter(target_col + 2)}{row_index}",
            )
            for column in range(target_col, target_col + 4):
                sheet.cell(row_index, column).number_format = "0.0000"
            target_col += 4
    style_sheet(sheet, "D2")
    for column in range(12, len(fields) + 1, 4):
        letter = get_column_letter(column)
        cell_range = f"{letter}2:{letter}{len(cases) + 1}"
        sheet.conditional_formatting.add(
            cell_range,
            CellIsRule(
                operator="greaterThan",
                formula=["0"],
                fill=PatternFill("solid", fgColor=GREEN),
            ),
        )
        sheet.conditional_formatting.add(
            cell_range,
            CellIsRule(
                operator="lessThan",
                formula=["0"],
                fill=PatternFill("solid", fgColor=RED),
            ),
        )
    sheet.conditional_formatting.add(
        f"F2:F{len(cases) + 1}",
        FormulaRule(
            formula=['ISNUMBER(SEARCH("IMPROVED",F2))'],
            fill=PatternFill("solid", fgColor=GREEN),
        ),
    )
    sheet.conditional_formatting.add(
        f"F2:F{len(cases) + 1}",
        FormulaRule(
            formula=['ISNUMBER(SEARCH("REGRESSED",F2))'],
            fill=PatternFill("solid", fgColor=RED),
        ),
    )
    return mapping, len(cases) + 1


def add_overview(
    workbook: Workbook,
    summary: dict[str, Any],
    paired_map: dict[str, int],
    paired_last: int,
) -> None:
    sheet = workbook.active
    sheet.title = "Overview"
    sheet.append(["E187 vs E178 — 22-case Paired Evaluation", "Value", "Source / rule"])
    paired_case = col_range("Paired Comparison", paired_map["case_id"], paired_last)
    old_pass = col_range(
        "Paired Comparison", paired_map["E178 numeric pass"], paired_last
    )
    new_pass = col_range(
        "Paired Comparison", paired_map["E187 numeric pass"], paired_last
    )
    transitions = col_range(
        "Paired Comparison", paired_map["Gate transition"], paired_last
    )
    rows = [
        ("Evaluation status", summary["status"].upper(), "summary.json"),
        ("C9 technical status", "FAIL", "frozen governance authority"),
        ("Progression authority", "USER_WAIVED", "user waiver; not technical PASS"),
        ("Paired rows", f"=COUNTA({paired_case})", "exact keep22 case_id match"),
        ("E178 numeric pass", f"=COUNTIF({old_pass},TRUE)", "same 12-gate contract"),
        ("E187 numeric pass", f"=COUNTIF({new_pass},TRUE)", "same 12-gate contract"),
        ("Pass-count delta", "=B7-B6", "E187 minus E178"),
        (
            "Improved to pass",
            f'=COUNTIF({transitions},"IMPROVED_TO_PASS")',
            "paired transition",
        ),
        (
            "Regressed to fail",
            f'=COUNTIF({transitions},"REGRESSED_TO_FAIL")',
            "paired transition",
        ),
        (
            "Unchanged pass",
            f'=COUNTIF({transitions},"UNCHANGED_PASS")',
            "paired transition",
        ),
        (
            "Unchanged fail",
            f'=COUNTIF({transitions},"UNCHANGED_FAIL")',
            "paired transition",
        ),
        (
            "Metric standard",
            summary["metric_standard_id"],
            "shared low-geometry evaluator",
        ),
        ("Delta definition", "E187 - E178", "Paired Comparison delta columns"),
        (
            "Improvement definition",
            "higher-is-better: delta; lower-is-better: -delta",
            "positive always means E187 better",
        ),
        ("Tracking gates", "root/hand/object pos+ori: 20/20/20/20/20/10", "cm / deg"),
        ("Manifest SHA256", summary["manifest_sha256"], summary["manifest"]),
        ("E178 metrics SHA256", summary["baseline_sha256"], summary["baseline"]),
        ("Generated at", summary["generated_at"], "summary.json"),
        (
            "KEY FINDING",
            '="E187 pass "&B7&"/"&B5&" vs E178 "&B6&"/"&B5',
            "paired numeric outcome",
        ),
    ]
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 72
    sheet.column_dimensions["C"].width = 58
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=12, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in (sheet["B3"], sheet["B4"]):
        cell.fill = PatternFill("solid", fgColor=YELLOW)
    for cell in (sheet["B8"], sheet["B9"], sheet["B10"], sheet["B20"]):
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)


def add_object_summary(
    workbook: Workbook,
    e187_map: dict[str, int],
    e187_last: int,
    e178_map: dict[str, int],
    e178_last: int,
) -> None:
    sheet = workbook.create_sheet("Object Summary")
    fields = ["object_key", "rows", "E178 pass", "E187 pass", "pass delta"]
    for metric in OBJECT_METRICS:
        fields.extend(
            (
                f"{metric} | E178 mean",
                f"{metric} | E187 mean",
                f"{metric} | improvement mean",
            )
        )
    sheet.append(fields)
    for row_index, object_key in enumerate(("bucket003", "bucket004", "bucket007"), 2):
        sheet.cell(row_index, 1, object_key)
        old_obj = col_range("E178 Baseline", e178_map["object_key"], e178_last)
        new_obj = col_range("E187 Metrics", e187_map["object_key"], e187_last)
        old_pass = col_range(
            "E178 Baseline", e178_map["numeric_release_pass"], e178_last
        )
        new_pass = col_range(
            "E187 Metrics", e187_map["numeric_release_pass"], e187_last
        )
        sheet.cell(row_index, 2, f"=COUNTIF({new_obj},A{row_index})")
        sheet.cell(row_index, 3, f"=COUNTIFS({old_obj},A{row_index},{old_pass},TRUE)")
        sheet.cell(row_index, 4, f"=COUNTIFS({new_obj},A{row_index},{new_pass},TRUE)")
        sheet.cell(row_index, 5, f"=D{row_index}-C{row_index}")
        target_col = 6
        for metric in OBJECT_METRICS:
            old_values = col_range("E178 Baseline", e178_map[metric], e178_last)
            new_values = col_range("E187 Metrics", e187_map[metric], e187_last)
            sheet.cell(
                row_index,
                target_col,
                f"=AVERAGEIF({old_obj},A{row_index},{old_values})",
            )
            sheet.cell(
                row_index,
                target_col + 1,
                f"=AVERAGEIF({new_obj},A{row_index},{new_values})",
            )
            sign = "" if PAIR_METRICS[metric] == "higher" else "-"
            old_letter = get_column_letter(target_col)
            new_letter = get_column_letter(target_col + 1)
            sheet.cell(
                row_index,
                target_col + 2,
                f"={sign}({new_letter}{row_index}-{old_letter}{row_index})",
            )
            for column in range(target_col, target_col + 3):
                sheet.cell(row_index, column).number_format = "0.0000"
            target_col += 3
    style_sheet(sheet, "B2")


def add_gate_transitions(
    workbook: Workbook,
    e187_map: dict[str, int],
    e187_last: int,
    e178_map: dict[str, int],
    e178_last: int,
) -> None:
    sheet = workbook.create_sheet("Gate Transitions")
    sheet.append(
        [
            "gate",
            "threshold",
            "E178 pass",
            "E187 pass",
            "improved",
            "regressed",
            "unchanged pass",
            "unchanged fail",
        ]
    )
    for row_index, gate in enumerate(GATES, 2):
        old = col_range("E178 Baseline", e178_map[f"{gate}_gate_pass"], e178_last)
        new = col_range("E187 Metrics", e187_map[f"{gate}_gate_pass"], e187_last)
        sheet.append(
            [
                gate,
                GATE_THRESHOLDS[gate],
                f"=COUNTIF({old},TRUE)",
                f"=COUNTIF({new},TRUE)",
                f"=COUNTIFS({old},FALSE,{new},TRUE)",
                f"=COUNTIFS({old},TRUE,{new},FALSE)",
                f"=COUNTIFS({old},TRUE,{new},TRUE)",
                f"=COUNTIFS({old},FALSE,{new},FALSE)",
            ]
        )
    style_sheet(sheet)


def add_ranked_sheet(
    workbook: Workbook,
    title: str,
    paired_map: dict[str, int],
    paired_last: int,
    best: bool,
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(
        ["metric", "direction", "rank", "case_id", "improvement", "E178", "E187"]
    )
    row_index = 2
    case_range = col_range("Paired Comparison", paired_map["case_id"], paired_last)
    for metric, direction in PAIR_METRICS.items():
        improvement_col = paired_map[f"{metric} | Improvement"]
        old_col = paired_map[f"{metric} | E178"]
        new_col = paired_map[f"{metric} | E187"]
        improvements = col_range("Paired Comparison", improvement_col, paired_last)
        for rank in range(1, 4):
            sheet.cell(row_index, 1, metric)
            sheet.cell(row_index, 2, direction)
            sheet.cell(row_index, 3, rank)
            fn = "LARGE" if best else "SMALL"
            sheet.cell(row_index, 5, f"={fn}({improvements},C{row_index})")
            sheet.cell(
                row_index,
                4,
                f"=INDEX({case_range},MATCH(E{row_index},{improvements},0))",
            )
            ranked_case = f"D{row_index}"
            for target_col, source_col in ((6, old_col), (7, new_col)):
                values = col_range("Paired Comparison", source_col, paired_last)
                sheet.cell(
                    row_index,
                    target_col,
                    f"=INDEX({values},MATCH({ranked_case},{case_range},0))",
                )
                sheet.cell(row_index, target_col).number_format = "0.0000"
            sheet.cell(row_index, 5).number_format = "0.0000"
            row_index += 1
    style_sheet(sheet)


def add_provenance(workbook: Workbook, manifest_rows: list[dict[str, str]]) -> None:
    fields = [
        "case_id",
        "execution_kind",
        "c9_technical_status",
        "c9_progression_authority",
        "row_manifest",
        "row_manifest_sha256",
        "result_npz",
        "result_sha256",
        "config_act",
        "config_sha256",
        "scene_act",
        "scene_sha256",
        "video",
        "video_sha256",
        "queue_manifest_sha256",
        "source_e178_variant",
        "source_e178_result_npz",
        "source_e178_manifest_sha256",
        "source_e178_metrics_sha256",
        "cem_samples",
        "cem_opt_steps",
        "cem_seed",
    ]
    add_data_sheet(workbook, "Artifact Provenance", fields, manifest_rows)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--eval-dir", type=Path, default=E187_EVAL)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    summary = json.loads((args.eval_dir / "summary.json").read_text(encoding="utf-8"))
    e187_fields, e187_rows = read_tsv(args.eval_dir / "e187_case_metrics.tsv")
    _, paired_rows = read_tsv(args.eval_dir / "e187_vs_e178_paired_deltas.tsv")
    manifest_fields, manifest_rows = read_tsv(E187_MANIFEST)
    e178_fields, e178_all = read_tsv(E178_METRICS)
    case_ids = [row["case_id"] for row in e187_rows]
    e178_by_case = {row["case_id"]: row for row in e178_all}
    e178_rows = [e178_by_case[case_id] for case_id in case_ids]
    if summary["status"] != "pass" or summary["counts"]["evaluated"] != 22:
        raise RuntimeError("E187 evaluation is not closed at 22/22 PASS")
    if len(e187_rows) != 22 or len(e178_rows) != 22 or len(paired_rows) != 22:
        raise RuntimeError("paired row count is not 22")
    if len(manifest_rows) != 22 or len(set(case_ids)) != 22:
        raise RuntimeError("manifest/case authority is not exact keep22")
    required = {
        "case_id",
        "object_key",
        "person",
        "numeric_release_pass",
        "numeric_failure_modes",
        *PAIR_METRICS,
        *(f"{gate}_gate_pass" for gate in GATES),
    }
    missing = sorted(required - set(e187_fields))
    missing += sorted(required - set(e178_fields))
    if missing:
        raise RuntimeError(f"source fields missing: {missing}")

    workbook = Workbook()
    e187_map, e187_last = add_data_sheet(
        workbook, "E187 Metrics", e187_fields, e187_rows
    )
    e178_map, e178_last = add_data_sheet(
        workbook, "E178 Baseline", e178_fields, e178_rows
    )
    paired_map, paired_last = add_paired_sheet(
        workbook,
        e187_rows,
        e187_map,
        e187_last,
        e178_map,
        e178_last,
    )
    add_overview(workbook, summary, paired_map, paired_last)
    add_object_summary(workbook, e187_map, e187_last, e178_map, e178_last)
    add_gate_transitions(workbook, e187_map, e187_last, e178_map, e178_last)
    add_ranked_sheet(workbook, "Best Improvements", paired_map, paired_last, True)
    add_ranked_sheet(workbook, "Worst Regressions", paired_map, paired_last, False)
    add_provenance(workbook, manifest_rows)
    workbook.move_sheet("Overview", offset=-(len(workbook.sheetnames) - 1))
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(
        json.dumps(
            {
                "status": "PASS",
                "output": str(args.output),
                "sheets": workbook.sheetnames,
                "paired_rows": len(e187_rows),
                "formulas": sum(
                    1
                    for sheet in workbook.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                    if isinstance(cell.value, str) and cell.value.startswith("=")
                ),
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
