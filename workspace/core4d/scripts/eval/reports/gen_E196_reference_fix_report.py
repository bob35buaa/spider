#!/usr/bin/env python3
"""Generate E196 by-object evidence, report, and formula-driven workbook."""

from __future__ import annotations

import argparse
import math
import statistics
import sys
from collections import Counter
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "experiments/E196"))
import e196_reference_fix_common as C  # noqa: E402


EVAL = C.RESULTS / "s6_downstream/eval/full_reference_fix"
CASE_PATH = EVAL / "e196_reference_fix_by_case.tsv"
METRICS_PATH = EVAL / "e196_reference_fix_case_metrics.tsv"
INTEGRITY_PATH = EVAL / "e196_reference_integrity_audit.tsv"
EVAL_SUMMARY_PATH = EVAL / "e196_reference_fix_eval_summary.json"
WORKBOOK = EVAL / "E196_reference_fix_comparison.xlsx"
REPORT = EVAL / "E196_reference_metadata_integrity_fix_report.md"

METRICS = (
    ("track_obj_z_abs_err_cm_mean", "Object z MAE (cm)", "lower"),
    ("track_obj_pos_err_cm_mean", "Object 3D position error (cm)", "lower"),
    ("track_obj_ori_err_deg_mean", "Object orientation error (deg)", "lower"),
    ("body_z_err_p95_m", "Body z error p95 (m)", "lower"),
    ("track_root_pos_err_cm_mean", "Body/root position error (cm)", "lower"),
    ("track_root_ori_err_deg_mean", "Body/root orientation error (deg)", "lower"),
    ("track_eef_pos_err_cm_mean", "Combined hand position error (cm)", "lower"),
    ("track_eef_ori_err_deg_mean", "Combined hand orientation error (deg)", "lower"),
    ("hand_object_physics_contact_3mm_in_mask_frac", "3 mm contact fraction", "higher"),
    ("hand_object_physics_contact_in_mask_frac", "Raw contact fraction", "higher"),
    ("hand_object_release_false_contact_3mm_frac", "False release fraction", "lower"),
    ("hand_object_physics_penetration_3mm_frame_frac", "Hand penetration fraction", "lower"),
    ("leg_penetration_frac", "Leg penetration fraction", "lower"),
    ("fall_flag", "Fall flag", "lower"),
)
COMPARISONS = (
    ("corrected_vs_prg", "prg", "corrected_g1", "PRG", "G1 corrected"),
    (
        "corrected_vs_contaminated_g1",
        "contaminated_g1",
        "corrected_g1",
        "G1 contaminated",
        "G1 corrected",
    ),
)
GATES = (
    "fall", "body_z", "contact", "release", "hand_penetration", "lower_body",
    "root_pos", "root_ori", "hand_pos", "hand_ori", "object_pos", "object_ori",
)
SUBSETS = (
    "ALL_AFFECTED_29", "BOX001_ALL_AFFECTED_21", "BOX001_PRIMARY_20", "BOX023_8"
)
NAVY, BLUE, WHITE = "17365D", "4472C4", "FFFFFF"
THIN = Side(style="thin", color="B7B7B7")


def number(value: Any) -> float:
    if isinstance(value, bool):
        return float(value)
    lowered = str(value).strip().lower()
    if lowered in {"true", "false"}:
        return float(lowered == "true")
    try:
        result = float(value)
    except (TypeError, ValueError):
        return math.nan
    return result if math.isfinite(result) else math.nan


def truth(value: Any) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "on"}


def subset_rows(rows: list[dict[str, str]], subset: str) -> list[dict[str, str]]:
    if subset == "ALL_AFFECTED_29":
        return rows
    if subset == "BOX001_ALL_AFFECTED_21":
        return [row for row in rows if row["object_key"] == "box001"]
    if subset == "BOX001_PRIMARY_20":
        return [
            row for row in rows
            if row["object_key"] == "box001" and truth(row["box001_primary"])
        ]
    if subset == "BOX023_8":
        return [row for row in rows if row["object_key"] == "box023"]
    raise ValueError(subset)


def bootstrap(values: list[float], seed: int) -> tuple[float, float]:
    array = np.asarray([value for value in values if math.isfinite(value)], dtype=np.float64)
    if not array.size:
        return math.nan, math.nan
    rng = np.random.default_rng(seed)
    indices = rng.integers(0, len(array), size=(10_000, len(array)))
    draws = array[indices].mean(axis=1)
    low, high = np.quantile(draws, [0.025, 0.975])
    return float(low), float(high)


def by_object(rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    seed = 0
    for subset in SUBSETS:
        group = subset_rows(rows, subset)
        for comparison, baseline, candidate, _, _ in COMPARISONS:
            for metric, label, direction in METRICS:
                before = [number(row[f"{baseline}_{metric}"]) for row in group]
                after = [number(row[f"{candidate}_{metric}"]) for row in group]
                raw = [right - left for left, right in zip(before, after)]
                improvement = [value if direction == "higher" else -value for value in raw]
                low, high = bootstrap(improvement, seed)
                seed += 1
                output.append(
                    {
                        "subset": subset,
                        "n_cases": len(group),
                        "comparison": comparison,
                        "metric": metric,
                        "metric_label": label,
                        "direction": direction,
                        "baseline_mean": statistics.fmean(before),
                        "candidate_mean": statistics.fmean(after),
                        "raw_delta_mean": statistics.fmean(raw),
                        "improvement_mean": statistics.fmean(improvement),
                        "improvement_median": statistics.median(improvement),
                        "improvement_min": min(improvement),
                        "improvement_max": max(improvement),
                        "improvement_bootstrap_ci95_low": low,
                        "improvement_bootstrap_ci95_high": high,
                        "improved_cases": sum(value > 0 for value in improvement),
                        "unchanged_cases": sum(abs(value) <= 1e-12 for value in improvement),
                        "worsened_cases": sum(value < 0 for value in improvement),
                    }
                )
    return output


def gate_rows(rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    output = []
    for row in rows:
        for comparison, baseline, candidate, baseline_label, candidate_label in COMPARISONS:
            for gate in GATES:
                before = truth(row[f"{baseline}_{gate}_gate_pass"])
                after = truth(row[f"{candidate}_{gate}_gate_pass"])
                migration = (
                    "PASS_TO_PASS" if before and after else "PASS_TO_FAIL" if before
                    else "FAIL_TO_PASS" if after else "FAIL_TO_FAIL"
                )
                output.append(
                    {
                        "comparison": comparison,
                        "baseline_arm": baseline_label,
                        "candidate_arm": candidate_label,
                        "case_id": row["case_id"],
                        "object_key": row["object_key"],
                        "gate": gate,
                        "baseline_pass": before,
                        "candidate_pass": after,
                        "migration": migration,
                    }
                )
    return output


def title(ws, text: str, subtitle: str, columns: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    ws["A1"] = text
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="666666")
    ws["A2"].alignment = Alignment(wrap_text=True)


def header(ws, row: int, columns: int) -> None:
    for cell in ws[row][:columns]:
        cell.font = Font(name="Arial", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)


def table(ws, start: int, end: int, columns: int, name: str) -> None:
    item = Table(displayName=name, ref=f"A{start}:{get_column_letter(columns)}{end}")
    item.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(item)
    ws.freeze_panes = f"A{start + 1}"
    ws.auto_filter.ref = item.ref


def widths(ws, values: dict[int, float], default: float = 15) -> None:
    for index in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(index)].width = values.get(index, default)


def add_readme(wb: Workbook, decision: str, integrity: list[dict[str, str]]) -> None:
    ws = wb.active
    ws.title = "README"
    title(ws, "E196 reference-fix comparison", "29 affected cases; formulas recalculate deltas and direction-aware improvements.", 4)
    rows = (
        ("Decision", decision, "Technical reference contract and corrected G1 performance are separate claims", ""),
        ("Scope", "29 mismatch cases", "box001=21; box023=8", "box001 primary excludes box001_20231023_110_p1"),
        ("Arms", "PRG", "G1 contaminated (E194)", "G1 corrected (E196)"),
        ("Delta", "Raw delta = candidate - baseline", "Improvement flips sign for lower-is-better metrics", "Positive improvement is always better"),
        ("Color", "Improvement >0 green", "Improvement =0 yellow", "Improvement <0 red"),
        ("Integrity", f"{sum(truth(row.get('integrity_pass')) for row in integrity)}/29 pass", "runtime/meta/XML parity", "target max <1e-4 deg"),
        ("CI", "10,000 paired bootstrap draws", "seed fixed per row", "CI cells are evidence values; workbook means/deltas are formulas"),
        ("Sources", "e196_reference_fix_by_case.tsv", "e196_reference_integrity_audit.tsv", "e196_reference_fix_by_object.tsv"),
    )
    for row_index, values in enumerate(rows, 4):
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_index, column, value)
            cell.font = Font(name="Arial", bold=column == 1)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN)
    widths(ws, {1: 22, 2: 32, 3: 36, 4: 42})


def add_paired(wb: Workbook, rows: list[dict[str, str]]) -> None:
    ws = wb.create_sheet("Paired Comparison")
    labels = (
        "Comparison", "Baseline arm", "Candidate arm", "Case", "Object", "Box001 primary",
        "Metric key", "Metric", "Direction", "Baseline value", "Candidate value", "Raw Delta", "Improvement",
    )
    title(ws, "Same-case paired comparison", "Raw Delta = candidate-baseline. Improvement is direction-aware; positive is always better.", len(labels))
    for column, label in enumerate(labels, 1):
        ws.cell(4, column, label)
    header(ws, 4, len(labels))
    row_index = 5
    for raw in rows:
        for comparison, baseline, candidate, baseline_label, candidate_label in COMPARISONS:
            for metric, label, direction in METRICS:
                values = (
                    comparison, baseline_label, candidate_label, raw["case_id"], raw["object_key"],
                    truth(raw["box001_primary"]), metric, label, direction,
                    number(raw[f"{baseline}_{metric}"]), number(raw[f"{candidate}_{metric}"]),
                )
                for column, value in enumerate(values, 1):
                    ws.cell(row_index, column, value)
                    ws.cell(row_index, column).font = Font(name="Arial", size=9)
                ws.cell(row_index, 12, f"=K{row_index}-J{row_index}")
                ws.cell(row_index, 13, f'=IF(I{row_index}="higher",L{row_index},-L{row_index})')
                for column in range(10, 14):
                    ws.cell(row_index, column).number_format = "0.0000"
                row_index += 1
    table(ws, 4, row_index - 1, len(labels), "E196PairedComparison")
    ws.conditional_formatting.add(
        f"M5:M{row_index - 1}",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0,
                       mid_color="FFEB84", end_type="max", end_color="63BE7B"),
    )
    widths(ws, {1: 31, 2: 18, 3: 18, 4: 38, 5: 12, 7: 42, 8: 34, 9: 12}, 16)


def criteria(subset: str, row: int) -> list[str]:
    source = "'Paired Comparison'!"
    output = [f"{source}$A:$A,$B{row}", f"{source}$G:$G,$C{row}"]
    if subset in {"BOX001_ALL_AFFECTED_21", "BOX001_PRIMARY_20"}:
        output.append(f'{source}$E:$E,"box001"')
    elif subset == "BOX023_8":
        output.append(f'{source}$E:$E,"box023"')
    if subset == "BOX001_PRIMARY_20":
        output.append(f"{source}$F:$F,TRUE")
    return output


def add_by_object(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("By Object")
    labels = (
        "Subset", "Comparison", "Metric key", "Metric", "Direction", "N", "Baseline mean",
        "Candidate mean", "Raw Delta", "Improvement", "Median improvement", "Bootstrap 95% CI low",
        "Bootstrap 95% CI high", "Improved cases", "Worsened cases",
    )
    title(ws, "By-object formula summary", "Means, deltas, improvements and counts are Excel formulas over Paired Comparison; bootstrap CI uses fixed-seed evidence.", len(labels))
    for column, label in enumerate(labels, 1):
        ws.cell(4, column, label)
    header(ws, 4, len(labels))
    for row_index, raw in enumerate(rows, 5):
        values = (raw["subset"], raw["comparison"], raw["metric"], raw["metric_label"], raw["direction"])
        for column, value in enumerate(values, 1):
            ws.cell(row_index, column, value)
        crit = ",".join(criteria(raw["subset"], row_index))
        ws.cell(row_index, 6, f"=COUNTIFS({crit})")
        ws.cell(row_index, 7, f"=AVERAGEIFS('Paired Comparison'!$J:$J,{crit})")
        ws.cell(row_index, 8, f"=AVERAGEIFS('Paired Comparison'!$K:$K,{crit})")
        ws.cell(row_index, 9, f"=H{row_index}-G{row_index}")
        ws.cell(row_index, 10, f'=IF(E{row_index}="higher",I{row_index},-I{row_index})')
        for column, key in (
            (11, "improvement_median"), (12, "improvement_bootstrap_ci95_low"),
            (13, "improvement_bootstrap_ci95_high"),
        ):
            ws.cell(row_index, column, raw[key])
            ws.cell(row_index, column).comment = Comment(
                "Source: e196_reference_fix_by_object.tsv; paired bootstrap 10,000 draws with frozen seed.",
                "Codex",
            )
        imp_crit = crit + "," + f"'Paired Comparison'!$M:$M,\">0\""
        bad_crit = crit + "," + f"'Paired Comparison'!$M:$M,\"<0\""
        ws.cell(row_index, 14, f"=COUNTIFS({imp_crit})")
        ws.cell(row_index, 15, f"=COUNTIFS({bad_crit})")
        for column in range(7, 14):
            ws.cell(row_index, column).number_format = "0.0000"
        for cell in ws[row_index]:
            cell.font = Font(name="Arial", size=9)
    table(ws, 4, ws.max_row, len(labels), "E196ByObject")
    ws.conditional_formatting.add(
        f"J5:J{ws.max_row}",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0,
                       mid_color="FFEB84", end_type="max", end_color="63BE7B"),
    )
    widths(ws, {1: 29, 2: 31, 3: 42, 4: 34, 5: 12}, 16)


def add_raw_sheet(wb: Workbook, name: str, title_text: str, rows: list[dict[str, Any]], table_name: str) -> None:
    ws = wb.create_sheet(name)
    fields = list(rows[0])
    title(ws, title_text, "Hardcoded source evidence; derived calculations are kept in formula-driven sheets.", len(fields))
    for column, field in enumerate(fields, 1):
        ws.cell(4, column, field)
    header(ws, 4, len(fields))
    for row_index, raw in enumerate(rows, 5):
        for column, field in enumerate(fields, 1):
            value: Any = raw.get(field, "")
            parsed = number(value)
            if math.isfinite(parsed):
                value = parsed
            elif str(value).lower() in {"true", "false"}:
                value = truth(value)
            ws.cell(row_index, column, value)
            ws.cell(row_index, column).font = Font(name="Arial", size=8)
    table(ws, 4, ws.max_row, len(fields), table_name)
    widths(ws, {1: 34, 2: 18, 3: 18}, 15)


def add_gates(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("12-Gate")
    labels = ("Comparison", "Baseline arm", "Candidate arm", "Case", "Object", "Gate", "Baseline pass", "Candidate pass", "Migration")
    title(ws, "Twelve-gate migrations", "Every case and gate is shown for corrected G1 against PRG and contaminated G1.", len(labels))
    for column, label in enumerate(labels, 1):
        ws.cell(4, column, label)
    header(ws, 4, len(labels))
    for row_index, raw in enumerate(rows, 5):
        for column, label in enumerate(labels[:-1], 1):
            field = ("comparison", "baseline_arm", "candidate_arm", "case_id", "object_key", "gate", "baseline_pass", "candidate_pass")[column - 1]
            ws.cell(row_index, column, raw[field])
        ws.cell(row_index, 9, f'=IF(AND(G{row_index},H{row_index}),"PASS_TO_PASS",IF(G{row_index},"PASS_TO_FAIL",IF(H{row_index},"FAIL_TO_PASS","FAIL_TO_FAIL")))')
        for cell in ws[row_index]:
            cell.font = Font(name="Arial", size=9)
    table(ws, 4, ws.max_row, len(labels), "E196TwelveGate")
    widths(ws, {1: 31, 2: 18, 3: 18, 4: 38, 5: 12, 6: 20, 9: 18}, 15)


def build_workbook(
    rows: list[dict[str, str]], by_rows: list[dict[str, Any]], integrity: list[dict[str, str]],
    gates: list[dict[str, Any]], metrics: list[dict[str, str]], decision: str,
) -> None:
    wb = Workbook()
    add_readme(wb, decision, integrity)
    add_paired(wb, rows)
    add_by_object(wb, by_rows)
    add_gates(wb, gates)
    add_raw_sheet(wb, "Integrity Audit", "Reference-integrity audit", integrity, "E196Integrity")
    add_raw_sheet(wb, "Case Metrics", "Three-arm public-core metrics", metrics, "E196CaseMetrics")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.font.name != "Arial":
                    cell.font = Font(name="Arial", size=cell.font.sz or 10, bold=cell.font.bold,
                                     italic=cell.font.italic, color=cell.font.color)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    wb.save(WORKBOOK)


def lookup(rows: list[dict[str, Any]], subset: str, comparison: str, metric: str) -> dict[str, Any]:
    return next(
        row for row in rows
        if row["subset"] == subset and row["comparison"] == comparison and row["metric"] == metric
    )


def gate_counts(metrics: list[dict[str, str]]) -> dict[str, int]:
    return dict(Counter(row["arm"] for row in metrics if truth(row.get("numeric_release_pass_12gate"))))


def write_report(
    rows: list[dict[str, str]], by_rows: list[dict[str, Any]], integrity: list[dict[str, str]],
    metrics: list[dict[str, str]], decision: str,
) -> None:
    gate_pass = gate_counts(metrics)
    key_rows = []
    for subset in ("ALL_AFFECTED_29", "BOX001_PRIMARY_20", "BOX023_8"):
        for metric in ("track_obj_ori_err_deg_mean", "track_obj_pos_err_cm_mean", "track_obj_z_abs_err_cm_mean"):
            key_rows.append(lookup(by_rows, subset, "corrected_vs_prg", metric))
    long_tails = [
        row for row in rows
        if number(row["raw_delta_corrected_vs_prg_track_obj_ori_err_deg_mean"]) > 5.0
    ]
    lines = [
        "# E196 reference metadata integrity fix report", "",
        "_Corrected Euler reference contract for the 29 E194 mismatch cases · 2026-08-12_", "", "---", "",
        "## 📋 Abstract", "",
        f"Decision: `{decision}`. E196 reran exactly 29 affected cases (box001=21, box023=8) with fail-closed metadata resolution. "
        f"Reference integrity passed `{sum(truth(row.get('integrity_pass')) for row in integrity)}/29`; performance is reported separately from the technical fix.", "",
        "## 🔬 Methodology", "",
        "PRG and contaminated G1 use the frozen E194 public-core authority. Corrected G1 is rescored with the same public core and 12-gate thresholds. "
        "All continuous comparisons are paired by case; positive `improvement` always means better.", "",
        "```mermaid", "flowchart LR",
        "    accTitle: E196 Evidence Closure",
        "    accDescr: Corrected metadata passes runtime parity before public-core scoring, paired comparison, and visual review produce the final E196 decision.", "",
        "    metadata[\"🔧 Resolve metadata\"] --> parity{\"🔍 Parity passed?\"}",
        "    parity -->|No| stop_run([\"❌ Stop experiment\"])",
        "    parity -->|Yes| score_cases[\"📊 Score 29 cases\"]",
        "    score_cases --> pair_arms[\"🔗 Pair three arms\"]",
        "    pair_arms --> review_video[\"🔍 Review videos\"]",
        "    review_video --> decision_node([\"✅ Record decision\"])", "",
        "    classDef action fill:#dbeafe,stroke:#2563eb,stroke-width:2px,color:#1e3a5f",
        "    classDef decision fill:#fef9c3,stroke:#ca8a04,stroke-width:2px,color:#713f12",
        "    classDef danger fill:#fee2e2,stroke:#dc2626,stroke-width:2px,color:#7f1d1d",
        "    classDef success fill:#dcfce7,stroke:#16a34a,stroke-width:2px,color:#14532d", "",
        "    class metadata,score_cases,pair_arms,review_video action",
        "    class parity decision", "    class stop_run danger", "    class decision_node success", "```", "",
        "## 📊 Findings", "", "### Reference integrity", "",
        "| Check | Result | Threshold |", "| --- | ---: | ---: |",
        f"| Runtime/meta/XML parity | {sum(truth(row.get('integrity_pass')) for row in integrity)}/29 | 29/29 |",
        f"| Runtime target max | {max(number(row['runtime_target_vs_raw_ori_err_deg_max']) for row in integrity):.8f}° | <0.0001° |",
        f"| Axis target max | {max(number(row['axis_target_vs_raw_ori_err_deg_max']) for row in integrity):.8f}° | <0.0001° |",
        f"| Public orientation reproduction max | {max(number(row['g1_vs_raw_ori_err_reproduction_abs_diff_deg']) for row in integrity):.10f}° | <0.000001° |", "",
        "### Corrected G1 versus PRG", "",
        "| Subset | Metric | PRG | Corrected G1 | Raw delta | Improvement | 95% CI |", "| --- | --- | ---: | ---: | ---: | ---: | ---: |",
    ]
    for row in key_rows:
        lines.append(
            f"| {row['subset']} | {row['metric_label']} | {row['baseline_mean']:.4f} | {row['candidate_mean']:.4f} | "
            f"{row['raw_delta_mean']:+.4f} | {row['improvement_mean']:+.4f} | "
            f"[{row['improvement_bootstrap_ci95_low']:+.4f}, {row['improvement_bootstrap_ci95_high']:+.4f}] |"
        )
    lines += ["", "### Twelve-gate result", "", "| Arm | Pass | Total |", "| --- | ---: | ---: |"]
    for arm in ("PRG", "G1_contaminated", "G1_corrected"):
        lines.append(f"| {arm} | {gate_pass.get(arm, 0)} | 29 |")
    lines += ["", "### New orientation long tails", ""]
    if long_tails:
        lines += ["| Case | Corrected − PRG |", "| --- | ---: |"]
        for row in sorted(long_tails, key=lambda item: number(item["raw_delta_corrected_vs_prg_track_obj_ori_err_deg_mean"]), reverse=True):
            lines.append(f"| `{row['case_id']}` | {number(row['raw_delta_corrected_vs_prg_track_obj_ori_err_deg_mean']):+.4f}° |")
    else:
        lines.append("No corrected G1 case has orientation regression `>5°` versus PRG.")
    lines += ["", "## 💡 Interpretation", "",
              "The reference contract claim is determined only by runtime/meta/XML/world-pose evidence. Performance claims use corrected G1 versus PRG; contaminated E194 G1 is retained only as historical evidence and is not used for gravcomp causality.", "",
              "## ⚠️ Limitations", "",
              "- The rerun set contains only the 29 convention-mismatch cases; the 43 convention-match cases were intentionally not rerun",
              "- Hand tracking is the public core's combined left/right EEF mean; no unsupported per-hand metric is imputed",
              "- Bootstrap intervals quantify paired effect uncertainty but do not replace case-level and visual failure review", "",
              "## 🔗 Artifacts", "",
              "- `e196_reference_fix_by_case.tsv`", "- `e196_reference_fix_by_object.tsv`",
              "- `e196_reference_integrity_audit.tsv`", "- `E196_reference_fix_comparison.xlsx`",
              "- `../../render/full_reference_fix/three_arm_video_manifest.tsv`", ""]
    REPORT.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--require-all", action="store_true")
    args = parser.parse_args()
    rows = C.read_tsv(CASE_PATH)
    metrics = C.read_tsv(METRICS_PATH)
    integrity = C.read_tsv(INTEGRITY_PATH)
    if args.require_all and (len(rows) != 29 or len(metrics) != 87 or len(integrity) != 29):
        raise SystemExit(f"E196 report inputs incomplete: cases={len(rows)} metrics={len(metrics)} integrity={len(integrity)}")
    technical = len(integrity) == 29 and all(truth(row.get("integrity_pass")) for row in integrity)
    by_rows = by_object(rows)
    gates = gate_rows(rows)
    C.write_tsv(EVAL / "e196_reference_fix_by_object.tsv", by_rows)
    C.write_tsv(EVAL / "e196_reference_fix_gate_migrations.tsv", gates)
    if not technical:
        decision = "INCOMPLETE"
    else:
        key = [
            lookup(by_rows, subset, "corrected_vs_prg", metric)["improvement_mean"]
            for subset in ("BOX001_PRIMARY_20", "BOX023_8")
            for metric in ("track_obj_ori_err_deg_mean", "track_obj_pos_err_cm_mean")
        ]
        decision = (
            "REFERENCE_FIX_VALIDATED_G1_IMPROVES" if all(value > 0 for value in key)
            else "REFERENCE_FIX_VALIDATED_G1_REGRESSION" if all(value < 0 for value in key)
            else "REFERENCE_FIX_VALIDATED_G1_MIXED"
        )
    build_workbook(rows, by_rows, integrity, gates, metrics, decision)
    write_report(rows, by_rows, integrity, metrics, decision)
    payload = {
        "created_at": C.now(), "decision": decision, "technical_reference_fix_validated": technical,
        "cases": len(rows), "case_set_sha256": C.case_set_sha256({row["case_id"] for row in rows}),
        "integrity_pass": sum(truth(row.get("integrity_pass")) for row in integrity),
        "strict12_pass": gate_counts(metrics),
        "sha256": {
            "by_case": C.sha256(CASE_PATH), "by_object": C.sha256(EVAL / "e196_reference_fix_by_object.tsv"),
            "integrity": C.sha256(INTEGRITY_PATH), "gate_migrations": C.sha256(EVAL / "e196_reference_fix_gate_migrations.tsv"),
            "workbook": C.sha256(WORKBOOK), "report": C.sha256(REPORT),
        },
    }
    C.write_json(EVAL / "e196_reference_fix_summary.json", payload)
    print(payload)
    return 1 if args.require_all and (not technical or len(rows) != 29) else 0


if __name__ == "__main__":
    raise SystemExit(main())
