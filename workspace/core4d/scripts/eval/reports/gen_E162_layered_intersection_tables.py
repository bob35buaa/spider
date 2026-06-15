#!/usr/bin/env python3
"""Generate layered, intersection-only E162 comparison workbooks.

This report intentionally does not overwrite the broad E162 workbook.  It
separates method-level comparisons by experiment stage and treats E147/E148 as
data sources for the same SPIDER+rubberhand method.
"""

from __future__ import annotations

import csv
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO = Path(__file__).resolve().parents[5]
OUT_DIR = REPO / "workspace/core4d/results/E162/post_e147_rl_safe_reeval/eval/full"

E162_METRICS = OUT_DIR / "e162_method_metrics.tsv"
E154_OMNI = REPO / "workspace/core4d/results/E154/omniretarget_eval/e154_omniretarget_method_metrics.tsv"
E156_METRICS = REPO / "workspace/core4d/results/E156/clean8_gate_decay/eval/full/e156_method_metrics.tsv"

CONTACT_METRIC = "hand_object_physics_contact_in_mask_frac"
CONTACT_DROP_FAIL_TH = -0.05


@dataclass(frozen=True)
class MethodDef:
    key: str
    label: str
    source: str
    method_id: str
    note: str


@dataclass(frozen=True)
class LayerSpec:
    title: str
    out_name: str
    baseline_key: str
    methods: list[MethodDef]


LAYER1 = LayerSpec(
    title="基础路线：OmniRetarget / rubberhand / gate / b1 / gate+b1 调参",
    out_name="E162_layer1_foundation_intersection.xlsx",
    baseline_key="rubberhand",
    methods=[
        MethodDef("omni", "OmniRetarget", "e154_omni", "OmniRetarget", "3-case SPIDER 输入"),
        MethodDef("rubberhand", "SPIDER+rubberhand", "e162", "E152_baseline", "同一方法；case 来源可为 E147 或 E148"),
        MethodDef("gateA", "+gateA", "e162", "E152_gateA", "E152 gate only"),
        MethodDef("b1", "+b1", "e162", "E152_b1", "E151/E152 b1 mesh"),
        MethodDef("gateA_b1_default", "+gateA+b1 default", "e162", "E152_gateA_b1", "min_sdf=-10mm, max_viol=5%"),
        MethodDef("gateA_b1_sdf005_v05", "+gateA+b1 -5mm/5%", "e162", "E153_gateA_b1_sdf005_v05", "E153 threshold sweep"),
        MethodDef("gateA_b1_sdf005_v10", "+gateA+b1 -5mm/10%", "e162", "E153_gateA_b1_sdf005_v10", "E153 threshold sweep"),
        MethodDef("gateA_b1_sdf010_v05", "+gateA+b1 -10mm/5%", "e162", "E153_gateA_b1_sdf010_v05", "E153 threshold sweep"),
        MethodDef("gateA_b1_sdf010_v10", "+gateA+b1 -10mm/10%", "e162", "E153_gateA_b1_sdf010_v10", "E153 threshold sweep"),
        MethodDef("gateA_b1_sdf015_v05", "+gateA+b1 -15mm/5%", "e162", "E153_gateA_b1_sdf015_v05", "E153 threshold sweep"),
        MethodDef("gateA_b1_sdf015_v10", "+gateA+b1 -15mm/10%", "e162", "E153_gateA_b1_sdf015_v10", "E153 threshold sweep"),
    ],
)

LAYER2 = LayerSpec(
    title="E156 之后：clean6 交集上的 surface/posture/release 系列",
    out_name="E162_layer2_post_E156_intersection.xlsx",
    baseline_key="rubberhand",
    methods=[
        MethodDef("omni", "OmniRetarget", "e156", "OmniRetarget", "E156 clean8 输入；本表取 clean6 交集"),
        MethodDef("rubberhand", "SPIDER+rubberhand", "e156", "spider-rubberhand", "E156 rubberhand baseline"),
        MethodDef("gateA", "+gateA", "e156", "+gateA", "E156 gateA"),
        MethodDef("e155_decay", "E155 decay", "e156", "E155_decay", "E156 中复用/补跑的 release smooth baseline"),
        MethodDef("surfaceA", "surfaceBand-A", "e162", "E158_gateA_surfaceBandA", "E158"),
        MethodDef("surfaceA2", "surfaceBand-A2", "e162", "E159_gateA_surfaceBandA2", "E159"),
        MethodDef(
            "posture",
            "postureRerankA",
            "e162",
            "E161_gateA_surfaceBandA2_postureRerankA",
            "同一方法；E161 是 E160 postureRerankA 的 clean8 扩展/复用",
        ),
        MethodDef(
            "release_decay",
            "surfaceBand releaseDecay",
            "e162",
            "E161_gateA_surfaceBandA2_postureRerankA_surfaceBandReleaseDecay",
            "E161 M1",
        ),
        MethodDef(
            "strict_mask",
            "surfaceBand strictMask",
            "e162",
            "E161_gateA_surfaceBandA2_postureRerankA_surfaceBandStrictMask",
            "E161 M2",
        ),
    ],
)


SUMMARY_METRICS = [
    ("几何接近5cm", "hand_geom_near_5cm_frac", "pct"),
    ("几何穿透2mm", "hand_geom_penetration_2mm_frac", "pct"),
    ("物理接触(raw)", CONTACT_METRIC, "pct"),
    ("物理穿透3mm", "hand_object_physics_penetration_3mm_frame_frac", "pct"),
    ("关节误差(°)", "track_joint_err_deg_mean", "num"),
    ("末端位置误差(cm)", "track_eef_pos_err_cm_mean", "num"),
    ("末端朝向误差(°)", "track_eef_ori_err_deg_mean", "num"),
    ("物体位置误差(cm)", "track_obj_pos_err_cm_mean", "num"),
    ("物体朝向误差(°)", "track_obj_ori_err_deg_mean", "num"),
]


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def finite(value: Any, default: float = math.nan) -> float:
    try:
        out = float(value)
    except (TypeError, ValueError):
        return default
    return out if math.isfinite(out) else default


def bool_text(value: Any) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "pass"}


def tracked_ok(row: dict[str, Any]) -> bool:
    value = row.get("success_tracked")
    if value not in {None, ""}:
        return bool_text(value)
    terminal_pz = finite(row.get("track_pelvis_z_err_terminal_m"))
    if math.isfinite(terminal_pz):
        return terminal_pz <= 0.08
    return True


def mean(values: list[Any]) -> float:
    vals = [finite(v) for v in values]
    vals = [v for v in vals if math.isfinite(v)]
    return sum(vals) / len(vals) if vals else math.nan


def case_id(row: dict[str, Any]) -> str:
    return str(row.get("short_case_id") or row.get("case_id") or "")


def normalize_row(row: dict[str, str], method: MethodDef) -> dict[str, Any]:
    out: dict[str, Any] = dict(row)
    out["layer_method_key"] = method.key
    out["layer_method_label"] = method.label
    out["layer_source"] = method.source
    out["layer_method_id"] = method.method_id
    out["layer_note"] = method.note
    out["short_case_id"] = case_id(row)
    return out


def load_indexes() -> dict[str, dict[str, dict[str, dict[str, Any]]]]:
    indexes: dict[str, dict[str, dict[str, dict[str, Any]]]] = {"e162": {}, "e154_omni": {}, "e156": {}}

    for row in read_tsv(E162_METRICS):
        method = row.get("canonical_method_id", "")
        indexes["e162"].setdefault(method, {})[case_id(row)] = row

    for row in read_tsv(E154_OMNI):
        indexes["e154_omni"].setdefault("OmniRetarget", {})[case_id(row)] = row

    for row in read_tsv(E156_METRICS):
        method = row.get("method", "")
        indexes["e156"].setdefault(method, {})[case_id(row)] = row

    return indexes


def layer_rows(
    spec: LayerSpec,
    indexes: dict[str, dict[str, dict[str, dict[str, Any]]]],
) -> tuple[list[str], dict[str, dict[str, dict[str, Any]]], list[dict[str, Any]]]:
    by_method: dict[str, dict[str, dict[str, Any]]] = {}
    audit_rows: list[dict[str, Any]] = []

    for method in spec.methods:
        rows = indexes.get(method.source, {}).get(method.method_id, {})
        normalized = {case: normalize_row(row, method) for case, row in rows.items()}
        by_method[method.key] = normalized

    case_sets = [set(by_method[method.key].keys()) for method in spec.methods]
    intersection = sorted(set.intersection(*case_sets)) if case_sets else []

    for method in spec.methods:
        cases = sorted(by_method[method.key].keys())
        excluded = [case for case in cases if case not in intersection]
        audit_rows.append(
            {
                "方法": method.label,
                "method_id": method.method_id,
                "source": method.source,
                "可用case数": len(cases),
                "纳入case数": len(intersection),
                "剔除case": ",".join(excluded),
                "说明": method.note,
            }
        )

    return intersection, by_method, audit_rows


def row_status(row: dict[str, Any], baseline: dict[str, Any]) -> tuple[str, float]:
    delta = finite(row.get(CONTACT_METRIC)) - finite(baseline.get(CONTACT_METRIC))
    if not math.isfinite(delta):
        return "接触指标缺失", math.nan
    if delta < CONTACT_DROP_FAIL_TH:
        return "接触退化", delta
    if bool_text(row.get("fall_flag")):
        return "摔倒", delta
    if not tracked_ok(row):
        return "tracking失败", delta
    return "通过", delta


def build_per_case(
    spec: LayerSpec,
    cases: list[str],
    by_method: dict[str, dict[str, dict[str, Any]]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    baseline_rows = by_method[spec.baseline_key]

    for case in cases:
        baseline = baseline_rows[case]
        for method in spec.methods:
            row = by_method[method.key][case]
            status, delta = row_status(row, baseline)
            item: dict[str, Any] = {
                "方法": method.label,
                "case": case,
                "状态": status,
                "物理接触Δ_vs_rubberhand": delta,
                "source": row.get("source_exp_id") or row.get("source_exp") or row.get("source_exp_id") or method.source,
                "method_id": method.method_id,
                "说明": method.note,
            }
            for _, key, _ in SUMMARY_METRICS:
                item[key] = finite(row.get(key))
            item["success_tracked"] = tracked_ok(row)
            item["fall_flag"] = bool_text(row.get("fall_flag"))
            rows.append(item)
    return rows


def build_summary(spec: LayerSpec, cases: list[str], per_case: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for method in spec.methods:
        rows = [row for row in per_case if row["方法"] == method.label]
        failures = [row for row in rows if row["状态"] != "通过"]
        contact_failures = [row for row in rows if row["状态"] == "接触退化"]
        item: dict[str, Any] = {
            "方法": method.label,
            "case数": len(cases),
            "通过case": sum(1 for row in rows if row["状态"] == "通过"),
            "失败case": ",".join(row["case"] for row in failures),
            "接触退化case": ",".join(row["case"] for row in contact_failures),
            "物理接触Δ最差": min(
                [finite(row.get("物理接触Δ_vs_rubberhand")) for row in rows if math.isfinite(finite(row.get("物理接触Δ_vs_rubberhand")))],
                default=math.nan,
            ),
            "source/method_id": method.method_id,
            "说明": method.note,
        }
        for _, key, _ in SUMMARY_METRICS:
            item[f"{key}_mean"] = mean([row.get(key) for row in rows])
        out.append(item)
    return out


def xlsx_value(value: Any) -> Any:
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float) and not math.isfinite(value):
        return ""
    return value


def append_table(
    wb: Workbook,
    title: str,
    rows: list[dict[str, Any]],
    columns: list[tuple[str, str | Callable[[dict[str, Any]], Any], str]],
    *,
    fill: str,
) -> Any:
    if len(wb.sheetnames) == 1 and wb.active.max_row == 1 and wb.active["A1"].value is None:
        ws = wb.active
        ws.delete_rows(1)
    else:
        ws = wb.create_sheet(title)
    ws.title = title
    ws.append([header for header, _, _ in columns])
    for row in rows:
        out = []
        for _, key, _ in columns:
            value = key(row) if callable(key) else row.get(key, "")
            out.append(xlsx_value(value))
        ws.append(out)

    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for col_idx, (_, _, fmt) in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        for cell in ws[letter][1:]:
            if fmt == "pct":
                cell.number_format = "0.0%"
            elif fmt == "pct_delta":
                cell.number_format = "0.0%;[Red]-0.0%"
            elif fmt == "num":
                cell.number_format = "0.00"
        width = min(max(len(str(cell.value or "")) for cell in ws[letter]) + 2, 38)
        ws.column_dimensions[letter].width = max(width, 10)
    return ws


def write_workbook(spec: LayerSpec, cases: list[str], summary: list[dict[str, Any]], per_case: list[dict[str, Any]], audit: list[dict[str, Any]]) -> Path:
    wb = Workbook()
    main_cols = [
        ("方法", "方法", "text"),
        ("case数", "case数", "int"),
        ("通过case", "通过case", "int"),
        ("失败case", "失败case", "text"),
        ("接触退化case", "接触退化case", "text"),
        ("物理接触Δ最差", "物理接触Δ最差", "pct_delta"),
        ("几何接近5cm", "hand_geom_near_5cm_frac_mean", "pct"),
        ("几何穿透2mm", "hand_geom_penetration_2mm_frac_mean", "pct"),
        ("物理接触(raw)", f"{CONTACT_METRIC}_mean", "pct"),
        ("物理穿透3mm", "hand_object_physics_penetration_3mm_frame_frac_mean", "pct"),
        ("关节误差(°)", "track_joint_err_deg_mean_mean", "num"),
        ("末端位置误差(cm)", "track_eef_pos_err_cm_mean_mean", "num"),
        ("末端朝向误差(°)", "track_eef_ori_err_deg_mean_mean", "num"),
        ("物体位置误差(cm)", "track_obj_pos_err_cm_mean_mean", "num"),
        ("物体朝向误差(°)", "track_obj_ori_err_deg_mean_mean", "num"),
        ("说明", "说明", "text"),
    ]
    append_table(wb, "主表", summary, main_cols, fill="1F4E79")

    per_case_cols = [
        ("方法", "方法", "text"),
        ("case", "case", "text"),
        ("状态", "状态", "text"),
        ("物理接触Δ_vs_rubberhand", "物理接触Δ_vs_rubberhand", "pct_delta"),
        ("几何接近5cm", "hand_geom_near_5cm_frac", "pct"),
        ("几何穿透2mm", "hand_geom_penetration_2mm_frac", "pct"),
        ("物理接触(raw)", CONTACT_METRIC, "pct"),
        ("物理穿透3mm", "hand_object_physics_penetration_3mm_frame_frac", "pct"),
        ("关节误差(°)", "track_joint_err_deg_mean", "num"),
        ("末端位置误差(cm)", "track_eef_pos_err_cm_mean", "num"),
        ("末端朝向误差(°)", "track_eef_ori_err_deg_mean", "num"),
        ("物体位置误差(cm)", "track_obj_pos_err_cm_mean", "num"),
        ("物体朝向误差(°)", "track_obj_ori_err_deg_mean", "num"),
        ("source", "source", "text"),
        ("method_id", "method_id", "text"),
    ]
    append_table(wb, "逐case", per_case, per_case_cols, fill="548235")

    audit_cols = [
        ("方法", "方法", "text"),
        ("method_id", "method_id", "text"),
        ("source", "source", "text"),
        ("可用case数", "可用case数", "int"),
        ("纳入case数", "纳入case数", "int"),
        ("剔除case", "剔除case", "text"),
        ("说明", "说明", "text"),
    ]
    append_table(wb, "交集审计", audit, audit_cols, fill="8064A2")

    info_rows = [
        {"项": "分层", "说明": spec.title},
        {"项": "case交集", "说明": ",".join(cases)},
        {"项": "baseline", "说明": "所有 delta/status 均相对同 case 的 SPIDER+rubberhand，不再相对 E147。"},
        {"项": "E147/E148", "说明": "E147 与 E148 是同一种 SPIDER+rubberhand 方法的不同结果批次；方法表合并，只在 source/method_id 审计中保留来源。"},
        {"项": "hard gate", "说明": f"{CONTACT_METRIC} 相对 rubberhand 下降 < {CONTACT_DROP_FAIL_TH} 判为接触退化。"},
        {"项": "tracking命名", "说明": "tracking 列按 SPIDER Table 4 风格：Joint Err., Pos. Err., Ori. Err., Obj. Pos. Err., Obj. Ori. Err."},
    ]
    append_table(wb, "说明", info_rows, [("项", "项", "text"), ("说明", "说明", "text")], fill="C65911")

    out_path = OUT_DIR / spec.out_name
    wb.save(out_path)
    return out_path


def main() -> None:
    indexes = load_indexes()
    for spec in [LAYER1, LAYER2]:
        cases, by_method, audit = layer_rows(spec, indexes)
        if not cases:
            raise RuntimeError(f"{spec.title}: empty case intersection")
        per_case = build_per_case(spec, cases, by_method)
        summary = build_summary(spec, cases, per_case)
        out_path = write_workbook(spec, cases, summary, per_case, audit)
        print(f"saved -> {out_path} (intersection_cases={len(cases)})")


if __name__ == "__main__":
    main()
