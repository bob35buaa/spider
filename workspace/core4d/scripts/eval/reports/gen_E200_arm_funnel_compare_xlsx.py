#!/usr/bin/env python3
"""E200 arm comparison under the E201 14-gate wide/narrow funnel: noPRG vs PRG.

Both arms score the SAME 249 translation-aug rollouts (E199 aug trajectories),
so they pair 1:1 on (case_id, aug_variant). This reads the classify_funnel.py
outputs (per-rollout hard/wide/narrow pass + failed + 3-tier layer, body_z
already recomputed there) and joins numeric gate values from case_metrics.

  * PRG   arm = E199 funnel rollout (group==aug)  -> E199_funnel_rollout.tsv
  * noPRG arm = E200_noprg funnel rollout          -> E200_noprg_funnel_rollout.tsv

Emits:
  * detail  — per rollout (both arms, adjacent): 4 hard (value+PASS) + 10 banded
              (value + narrow PASS + wide PASS) + layer + family_flag + arm.
  * summary — noPRG vs PRG, overall + per object: layer distribution, overall
              wide/narrow accept rate, per-gate wide & narrow pass rate, and
              key-metric means.

Caliber = E201 funnel_config (single source of truth). NOTE the noPRG arm has no
same-contract orig row (its baseline is E190), so its family arbitration uses only
the 3 trans arms — layer numbers are informational; the gate pass rates (same 249
aug, identical scoring) are the clean apples-to-apples comparison.

Usage:
    .venv/bin/python workspace/core4d/scripts/eval/reports/gen_E200_arm_funnel_compare_xlsx.py
"""

from __future__ import annotations

import argparse
import csv
import math
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[5]
sys.path.insert(0, str(REPO / "workspace/core4d/scripts/experiments/E201"))
import funnel_config as FC  # noqa: E402

E201_FUNNEL = REPO / "workspace/core4d/results/E201/funnel"
E200_EVAL = REPO / "workspace/core4d/results/E200/s6_downstream/eval"
PRG_FUNNEL = E201_FUNNEL / "E199_funnel_rollout.tsv"
NOPRG_FUNNEL = E201_FUNNEL / "E200_noprg_funnel_rollout.tsv"
PRG_CASE_METRICS = REPO / "workspace/core4d/results/E199/s6_downstream/eval/fullscale_augmentation/e199_fullscale_case_metrics.tsv"
NOPRG_CASE_METRICS = E200_EVAL / "noprg/e200_noprg_case_metrics.tsv"
DEFAULT_OUT = E200_EVAL / "E200_noprg_vs_prg_funnel.xlsx"

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
NAVY = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF")
LAYER_FILL = {
    "L1_reject": PatternFill("solid", fgColor="C00000"),
    "L2_review": PatternFill("solid", fgColor="FFD966"),
    "L3_auto": PatternFill("solid", fgColor="70AD47"),
    "L3_review": PatternFill("solid", fgColor="F4B183"),
}
LAYERS = ["L1_reject", "L2_review", "L3_auto", "L3_review"]
HARD = FC.HARD_GATES
BANDED = FC.BANDED_GATES
VARIANT_ORDER = {"trans0": 0, "trans1": 1, "trans2": 2}
ARM_ORDER = {"PRG": 0, "noPRG": 1}


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh, delimiter="\t"))


def finite(v: Any) -> float:
    try:
        f = float(v)
        return f if math.isfinite(f) else math.nan
    except (TypeError, ValueError):
        return math.nan


def _norm(cid: str) -> str:
    return cid.replace("_person", "_p")


def load_arm(funnel_tsv: Path, case_metrics: Path, arm: str, only_aug: bool) -> list[dict[str, Any]]:
    """Join per-rollout funnel layer/pass with numeric gate values from case_metrics."""
    values = {(m["case_id"], m.get("aug_variant", "")): m for m in read_tsv(case_metrics)}
    out: list[dict[str, Any]] = []
    for lr in read_tsv(funnel_tsv):
        if only_aug and lr.get("group") != "aug":
            continue
        row = dict(values.get((lr["case_id"], lr.get("aug_variant", "")), {}))
        row.update(lr)                 # funnel layer/pass/failed + recomputed body_z win
        row["arm"] = arm
        out.append(row)
    return out


def _truthy(v: Any) -> bool:
    return str(v).strip().lower() in ("true", "1")


def _cell(ws, r, c, value, fill=None, font=None):
    cell = ws.cell(row=r, column=c, value=value)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    return cell


def _num(ws, r, c, value):
    v = finite(value)
    _cell(ws, r, c, round(v, 4) if math.isfinite(v) else "")


def write_detail(ws, rows: list[dict[str, Any]]) -> None:
    header = ["object_key", "case_id", "arm", "aug_variant", "layer", "family_flag"]
    for name, _f, _op, _thr in HARD:
        header += [name, f"{name}✓"]
    for name, _f, _op, narrow, wide in BANDED:
        header += [name, f"{name} N({narrow:g})", f"{name} W({wide:g})"]
    for c, name in enumerate(header, start=1):
        cell = _cell(ws, 1, c, name, NAVY, HEAD_FONT)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "G2"

    for r, row in enumerate(rows, start=2):
        _cell(ws, r, 1, row.get("object_key", ""))
        _cell(ws, r, 2, row.get("case_id", ""))
        _cell(ws, r, 3, row.get("arm", ""))
        _cell(ws, r, 4, row.get("aug_variant", ""))
        _cell(ws, r, 5, row.get("layer", ""), LAYER_FILL.get(row.get("layer", "")), HEAD_FONT)
        _cell(ws, r, 6, row.get("family_flag", ""))
        c = 7
        hard_failed = set((row.get("hard_failed") or "").split(","))
        for name, field, _op, _thr in HARD:
            if field == "fall_flag":
                _cell(ws, r, c, str(row.get("fall_flag", "")))
            elif name == "body_z":
                _num(ws, r, c, row.get("body_z_err_p95_m"))
            else:
                _num(ws, r, c, row.get(field))
            c += 1
            ok = name not in hard_failed
            _cell(ws, r, c, "PASS" if ok else "FAIL", GREEN if ok else RED)
            c += 1
        wide_failed = set((row.get("wide_failed") or "").split(","))
        narrow_failed = set((row.get("narrow_failed") or "").split(","))
        for name, field, _op, _n, _w in BANDED:
            _num(ws, r, c, row.get(field)); c += 1
            n_ok = name not in narrow_failed
            _cell(ws, r, c, "PASS" if n_ok else "FAIL", GREEN if n_ok else RED); c += 1
            w_ok = name not in wide_failed
            _cell(ws, r, c, "PASS" if w_ok else "FAIL", GREEN if w_ok else RED); c += 1

    widths = [10, 34, 7, 10, 11, 20] + [11, 6] * len(HARD) + [11, 9, 9] * len(BANDED)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _rate(rows: list[dict[str, Any]], pred) -> tuple[int, int]:
    return sum(1 for r in rows if pred(r)), len(rows)


def write_summary(ws, noprg: list[dict[str, Any]], prg: list[dict[str, Any]]) -> None:
    objects = sorted({r["object_key"] for r in noprg + prg})

    def gate_pass(row, name, caliber):
        col = "narrow_failed" if caliber == "narrow" else "wide_failed"
        return name not in set((row.get(col) or "").split(","))

    def hard_pass(row, name):
        return name not in set((row.get("hard_failed") or "").split(","))

    def block(title: str, np_rows, prg_rows, top: int) -> int:
        _cell(ws, top, 1, f"{title}  (noPRG {len(np_rows)} vs PRG {len(prg_rows)})", font=Font(bold=True))
        heads = ["metric / gate", "noPRG", "noPRG rate", "PRG", "PRG rate", "Δrate pp"]
        for c, h in enumerate(heads, start=1):
            _cell(ws, top + 1, c, h, NAVY, HEAD_FONT)
        r = top + 2

        def line(label, np_pred, prg_pred=None):
            nonlocal r
            prg_pred = prg_pred or np_pred
            nn, nt = _rate(np_rows, np_pred); pn, pt = _rate(prg_rows, prg_pred)
            nr = nn / nt if nt else float("nan"); pr = pn / pt if pt else float("nan")
            _cell(ws, r, 1, label)
            _cell(ws, r, 2, f"{nn}/{nt}"); _cell(ws, r, 3, round(nr, 3) if nt else "")
            _cell(ws, r, 4, f"{pn}/{pt}"); _cell(ws, r, 5, round(pr, 3) if pt else "")
            _cell(ws, r, 6, round((nr - pr) * 100, 1) if nt and pt else "")
            r += 1

        # layer distribution
        for lay in LAYERS:
            line(lay, lambda x, L=lay: x.get("layer") == L)
        line("auto-decided (L1+L3auto)", lambda x: x.get("layer") in ("L1_reject", "L3_auto"))
        line("accept L3 (auto+review)", lambda x: x.get("layer") in ("L3_auto", "L3_review"))
        # overall calibers
        _cell(ws, r, 1, "— overall accept —"); ws.cell(row=r, column=1).font = Font(italic=True); r += 1
        line("hard_all_pass", lambda x: _truthy(x.get("hard_pass")))
        line("WIDE_all_pass (not L1)", lambda x: _truthy(x.get("hard_pass")) and _truthy(x.get("wide_pass")))
        line("NARROW_all_pass", lambda x: _truthy(x.get("hard_pass")) and _truthy(x.get("narrow_pass")))
        # per hard gate
        _cell(ws, r, 1, "— hard gates —"); ws.cell(row=r, column=1).font = Font(italic=True); r += 1
        for name, _f, _op, _thr in HARD:
            line(f"hard:{name}", lambda x, N=name: hard_pass(x, N))
        # per banded gate: narrow then wide
        _cell(ws, r, 1, "— banded gates (N=narrow, W=wide) —"); ws.cell(row=r, column=1).font = Font(italic=True); r += 1
        for name, _f, _op, narrow, wide in BANDED:
            line(f"{name} N({narrow:g})", lambda x, N=name: gate_pass(x, N, "narrow"))
            line(f"{name} W({wide:g})", lambda x, N=name: gate_pass(x, N, "wide"))
        # key-metric means
        _cell(ws, r, 1, "— key-metric mean —"); ws.cell(row=r, column=1).font = Font(italic=True); r += 1
        import numpy as np
        for name, field, _op, *_ in BANDED:
            def mean(rows):
                v = np.array([finite(x.get(field)) for x in rows]); v = v[np.isfinite(v)]
                return float(v.mean()) if v.size else float("nan")
            nm, pm = mean(np_rows), mean(prg_rows)
            _cell(ws, r, 1, f"mean:{name}")
            _cell(ws, r, 2, round(nm, 4)); _cell(ws, r, 4, round(pm, 4))
            _cell(ws, r, 6, round(nm - pm, 4) if math.isfinite(nm) and math.isfinite(pm) else "")
            r += 1
        return r + 1

    row_ptr = block("OVERALL", noprg, prg, 1)
    for obj in objects:
        row_ptr = block(obj, [r for r in noprg if r["object_key"] == obj],
                        [r for r in prg if r["object_key"] == obj], row_ptr)
    ws.column_dimensions["A"].width = 30
    for col in "BCDEF":
        ws.column_dimensions[col].width = 13


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()

    # both arms compared on aug only; orig rows (if present) only fed the family
    # arbitration inside classify_funnel and are excluded from the gate-rate compare.
    prg = load_arm(PRG_FUNNEL, PRG_CASE_METRICS, "PRG", only_aug=True)
    noprg = load_arm(NOPRG_FUNNEL, NOPRG_CASE_METRICS, "noPRG", only_aug=True)
    print(f"[load] noPRG {len(noprg)} aug, PRG {len(prg)} aug")

    combined = prg + noprg
    combined.sort(key=lambda r: (r["object_key"], _norm(r["case_id"]),
                                 VARIANT_ORDER.get(r["aug_variant"], 9), ARM_ORDER.get(r["arm"], 9)))
    wb = Workbook()
    write_detail(wb.active, combined); wb.active.title = "detail"
    write_summary(wb.create_sheet("summary"), noprg, prg)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"[done] wrote {args.out} (noPRG {len(noprg)} + PRG {len(prg)} = {len(combined)} rows)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
