#!/usr/bin/env python3
"""Generate the 11-sheet E188 versus E187 paired evaluation workbook."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[5]
ROOT = REPO / "workspace/core4d/results/E188"
EVAL = ROOT / "s6_downstream/eval/full"
MANIFEST = ROOT / "s6_downstream/manifests/e188_full_evaluation_manifest.tsv"
AUTHORITY = ROOT / "s0_environment/authority_manifest.tsv"
E187 = REPO / "workspace/core4d/results/E187/s6_downstream/eval/full/e187_case_metrics.tsv"
VIDEOS = ROOT / "s6_downstream/render/full/e188_videos/e188_video_manifest.tsv"
PAIRED_VIDEOS = ROOT / "s6_downstream/render/full/paired_e187_vs_e188/paired_video_manifest.tsv"
OUTPUT = EVAL / "E188_vs_E187_paired_evaluation.xlsx"
METRICS = {
    "hand_object_physics_contact_in_mask_frac": "higher",
    "hand_object_physics_penetration_3mm_frame_frac": "lower",
    "leg_penetration_frac": "lower",
    "track_root_pos_err_cm_mean": "lower",
    "track_root_ori_err_deg_mean": "lower",
    "track_eef_pos_err_cm_mean": "lower",
    "track_eef_ori_err_deg_mean": "lower",
    "track_obj_pos_err_cm_mean": "lower",
    "track_obj_ori_err_deg_mean": "lower",
}
NAVY = "17365D"; BLUE = "4472C4"; LIGHT = "D9EAF7"; GREEN = "C6EFCE"; RED = "FFC7CE"; WHITE = "FFFFFF"


def read_tsv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream, delimiter="\t")
        return list(reader.fieldnames or []), list(reader)


def value(raw: Any) -> Any:
    if raw is None or raw == "": return ""
    text = str(raw).strip()
    if text.lower() in {"true", "false"}: return text.lower() == "true"
    try:
        number = float(text)
        return int(number) if number.is_integer() and not any(char in text.lower() for char in (".", "e")) else number
    except ValueError:
        return raw


def style(sheet: Any, freeze: str = "A2") -> None:
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    for col in range(1, sheet.max_column + 1):
        width = max(len(str(sheet.cell(row=row, column=col).value or "")) for row in range(1, min(sheet.max_row, 50) + 1))
        sheet.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 11), 38)
    sheet.sheet_view.showGridLines = False


def data_sheet(wb: Workbook, name: str, fields: list[str], rows: list[dict[str, str]]) -> Any:
    sheet = wb.create_sheet(name)
    sheet.append(fields)
    for row in rows: sheet.append([value(row.get(field, "")) for field in fields])
    style(sheet)
    return sheet


def index_formula(sheet: str, case_cell: str, column: int, last_row: int) -> str:
    col = get_column_letter(column)
    return f'=INDEX(\'{sheet}\'!${col}$2:${col}${last_row},MATCH({case_cell},\'{sheet}\'!$A$2:$A${last_row},0))'


def paired_sheet(wb: Workbook, manifest_rows: list[dict[str, str]], e188_fields: list[str], e188_rows: list[dict[str, str]], e187_fields: list[str], e187_rows: list[dict[str, str]]) -> tuple[Any, dict[str, int]]:
    sheet = wb.create_sheet("Paired Comparison")
    headers = ["case_id", "object_key", "device_scope", "E187 numeric pass", "E188 numeric pass", "numeric transition", "E187 lower-body pass", "E188 lower-body pass", "E187 worker", "E188 worker", "E187 device", "E188 device", "old mass kg", "new mass kg"]
    for metric in METRICS: headers.extend([f"{metric} | E187", f"{metric} | E188", f"{metric} | improvement"])
    sheet.append(headers); mapping = {name: index + 1 for index, name in enumerate(headers)}
    e187_map = {field: index + 1 for index, field in enumerate(e187_fields)}; e188_map = {field: index + 1 for index, field in enumerate(e188_fields)}
    for excel_row, source in enumerate(manifest_rows, 2):
        sheet.cell(excel_row, 1, source["case_id"]); sheet.cell(excel_row, 2, source["object_key"]); sheet.cell(excel_row, 3, source["device_scope"])
        sheet.cell(excel_row, 4, index_formula("E187 Baseline", f"$A{excel_row}", e187_map["numeric_release_pass"], len(e187_rows) + 1))
        sheet.cell(excel_row, 5, index_formula("E188 Metrics", f"$A{excel_row}", e188_map["numeric_release_pass"], len(e188_rows) + 1))
        sheet.cell(excel_row, 6, f'=IF(D{excel_row},"E187_PASS_TO_E188_"&IF(E{excel_row},"PASS","FAIL"),"E187_FAIL_TO_E188_"&IF(E{excel_row},"PASS","FAIL"))')
        sheet.cell(excel_row, 7, index_formula("E187 Baseline", f"$A{excel_row}", e187_map["lower_body_gate_pass"], len(e187_rows) + 1))
        sheet.cell(excel_row, 8, index_formula("E188 Metrics", f"$A{excel_row}", e188_map["lower_body_gate_pass"], len(e188_rows) + 1))
        for col, field in enumerate(("e187_worker", "e188_worker", "e187_device", "e188_device", "old_mass_kg", "new_mass_kg"), 9): sheet.cell(excel_row, col, value(source[field]))
        col = 15
        for metric, direction in METRICS.items():
            sheet.cell(excel_row, col, index_formula("E187 Baseline", f"$A{excel_row}", e187_map[metric], len(e187_rows) + 1))
            sheet.cell(excel_row, col + 1, index_formula("E188 Metrics", f"$A{excel_row}", e188_map[metric], len(e188_rows) + 1))
            sheet.cell(excel_row, col + 2, f"={get_column_letter(col + 1)}{excel_row}-{get_column_letter(col)}{excel_row}" if direction == "higher" else f"={get_column_letter(col)}{excel_row}-{get_column_letter(col + 1)}{excel_row}")
            col += 3
    style(sheet)
    for col, name in enumerate(headers, 1):
        if name.endswith("| improvement"):
            rng = f"{get_column_letter(col)}2:{get_column_letter(col)}{sheet.max_row}"
            sheet.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    return sheet, mapping


def overview(wb: Workbook, paired_map: dict[str, int], summary: dict[str, Any], paired_rows: int) -> None:
    sheet = wb.create_sheet("Overview", 0)
    sheet.append(["E188 vs E187 — Bucket 2→5 kg Paired Evaluation", "Value", "Source / rule"])
    sheet.append(["Experiment status", "CEM 15/15; Eval 15/15; videos 15/15", "E188 closure + Phase 6/7 audits"])
    sheet.append(["Interpretation boundary", "same-device local-4 is strongest mass-only evidence", "cross-device11 confounds mass and device"])
    sheet.append(["Paired rows", f"=COUNTA('Paired Comparison'!A2:A{paired_rows + 1})", "exact E188 authority"])
    sheet.append(["E187 numeric pass", f"=COUNTIF('Paired Comparison'!D2:D{paired_rows + 1},TRUE)", "frozen 12 gates"])
    sheet.append(["E188 numeric pass", f"=COUNTIF('Paired Comparison'!E2:E{paired_rows + 1},TRUE)", "same 12 gates"])
    sheet.append(["PASS→FAIL", f'=COUNTIF(\'Paired Comparison\'!F2:F{paired_rows + 1},"E187_PASS_TO_E188_FAIL")', "C8 maximum 1"])
    sheet.append(["E187 lower-body pass", f"=COUNTIF('Paired Comparison'!G2:G{paired_rows + 1},TRUE)", "gate ≤0.10"])
    sheet.append(["E188 lower-body pass", f"=COUNTIF('Paired Comparison'!H2:H{paired_rows + 1},TRUE)", "C5 requires ≥7"])
    leg_col = get_column_letter(paired_map["leg_penetration_frac | improvement"]); contact_col = get_column_letter(paired_map["hand_object_physics_contact_in_mask_frac | improvement"]); hand_col = get_column_letter(paired_map["hand_object_physics_penetration_3mm_frame_frac | improvement"])
    sheet.append(["Mean leg improvement", f"=AVERAGE('Paired Comparison'!{leg_col}2:{leg_col}{paired_rows + 1})", "E187−E188; C5 requires ≥0.05"])
    sheet.append(["Leg nondegraded cases", f'=COUNTIF(\'Paired Comparison\'!{leg_col}2:{leg_col}{paired_rows + 1},">=0")', "C5 requires ≥10"])
    sheet.append(["Mean contact improvement", f"=AVERAGE('Paired Comparison'!{contact_col}2:{contact_col}{paired_rows + 1})", "E188−E187"])
    sheet.append(["Mean hand-penetration improvement", f"=AVERAGE('Paired Comparison'!{hand_col}2:{hand_col}{paired_rows + 1})", "E187−E188"])
    sheet.append(["C5 lower-body", "=AND(B9>=7,B10>=0.05,B11>=10)", "plan212"])
    sheet.append(["C6 contact/hand penetration", "=AND(B12>=-0.03,B13>=-0.03)", "plan212"])
    track_cols = [get_column_letter(paired_map[f"{metric} | improvement"]) for metric in METRICS if metric.startswith("track_")]
    track_checks = ",".join(f"AVERAGE('Paired Comparison'!{col}2:{col}{paired_rows + 1})>=-2" for col in track_cols)
    sheet.append(["C7 tracking", f"=AND({track_checks})", "regression per mean ≤2cm/2°"])
    sheet.append(["C8 total gates", "=AND(B6>=5,B7<=1)", "plan212"])
    sheet.append(["Bootstrap", f"{summary['bootstrap']['draws']} draws, seed {summary['bootstrap']['seed']}", "Device Stratified sheet"])
    style(sheet, "A2"); sheet.column_dimensions["A"].width = 34; sheet.column_dimensions["B"].width = 46; sheet.column_dimensions["C"].width = 48
    for row in range(14, 18):
        sheet.conditional_formatting.add(f"B{row}", CellIsRule(operator="equal", formula=["TRUE"], fill=PatternFill("solid", fgColor=GREEN)))
        sheet.conditional_formatting.add(f"B{row}", CellIsRule(operator="equal", formula=["FALSE"], fill=PatternFill("solid", fgColor=RED)))


def object_summary(wb: Workbook, paired_map: dict[str, int], n: int) -> None:
    sheet = wb.create_sheet("Object Summary")
    headers = ["object_key", "rows", "E187 pass", "E188 pass"]
    for metric in METRICS: headers.append(f"{metric} | improvement mean")
    sheet.append(headers)
    for row, obj in enumerate(("bucket003", "bucket007"), 2):
        sheet.cell(row, 1, obj); sheet.cell(row, 2, f'=COUNTIF(\'Paired Comparison\'!$B$2:$B${n + 1},$A{row})'); sheet.cell(row, 3, f'=COUNTIFS(\'Paired Comparison\'!$B$2:$B${n + 1},$A{row},\'Paired Comparison\'!$D$2:$D${n + 1},TRUE)'); sheet.cell(row, 4, f'=COUNTIFS(\'Paired Comparison\'!$B$2:$B${n + 1},$A{row},\'Paired Comparison\'!$E$2:$E${n + 1},TRUE)')
        for col, metric in enumerate(METRICS, 5):
            metric_col = get_column_letter(paired_map[f"{metric} | improvement"])
            sheet.cell(row, col, f'=AVERAGEIF(\'Paired Comparison\'!$B$2:$B${n + 1},$A{row},\'Paired Comparison\'!${metric_col}$2:${metric_col}${n + 1})')
    style(sheet)


def ranked_sheet(wb: Workbook, name: str, manifest_rows: list[dict[str, str]], paired_map: dict[str, int], reverse: bool) -> None:
    _, deltas = read_tsv(EVAL / "e188_vs_e187_paired_deltas.tsv")
    ordered = sorted(deltas, key=lambda row: float(row["improvement_leg_penetration_frac"]), reverse=reverse)
    sheet = wb.create_sheet(name); headers = ["rank", "case_id", "device_scope", "gate transition", *[f"{metric} improvement" for metric in METRICS]]; sheet.append(headers)
    for excel_row, row in enumerate(ordered, 2):
        source_row = next(index + 2 for index, source in enumerate(manifest_rows) if source["case_id"] == row["case_id"])
        sheet.cell(excel_row, 1, excel_row - 1); sheet.cell(excel_row, 2, f"='Paired Comparison'!A{source_row}"); sheet.cell(excel_row, 3, f"='Paired Comparison'!C{source_row}"); sheet.cell(excel_row, 4, f"='Paired Comparison'!F{source_row}")
        for col, metric in enumerate(METRICS, 5):
            source_col = get_column_letter(paired_map[f"{metric} | improvement"]); sheet.cell(excel_row, col, f"='Paired Comparison'!{source_col}{source_row}")
    style(sheet)


def main() -> int:
    authority_fields, authority_rows = read_tsv(AUTHORITY); manifest_fields, manifest_rows = read_tsv(MANIFEST); e188_fields, e188_rows = read_tsv(EVAL / "e188_case_metrics.tsv"); e187_fields, e187_all = read_tsv(E187); case_ids = {row["case_id"] for row in manifest_rows}; e187_rows = [row for row in e187_all if row["case_id"] in case_ids]
    device_fields, device_rows = read_tsv(EVAL / "e188_device_stratified_summary.tsv"); gate_fields, gate_rows = read_tsv(EVAL / "e188_gate_transitions.tsv"); video_fields, video_rows = read_tsv(VIDEOS); paired_video_fields, paired_video_rows = read_tsv(PAIRED_VIDEOS); summary = json.loads((EVAL / "summary.json").read_text(encoding="utf-8"))
    if not (len(authority_rows) == len(manifest_rows) == len(e188_rows) == len(e187_rows) == len(video_rows) == len(paired_video_rows) == 15): raise RuntimeError("E188 workbook inputs are not closed at 15 rows")
    wb = Workbook(); wb.remove(wb.active)
    data_sheet(wb, "Case Mass Audit", authority_fields, authority_rows); data_sheet(wb, "E188 Metrics", e188_fields, e188_rows); data_sheet(wb, "E187 Baseline", e187_fields, e187_rows)
    paired, paired_map = paired_sheet(wb, manifest_rows, e188_fields, e188_rows, e187_fields, e187_rows)
    data_sheet(wb, "Device Stratified", device_fields, device_rows); data_sheet(wb, "Gate Transitions", gate_fields, gate_rows); object_summary(wb, paired_map, len(manifest_rows)); ranked_sheet(wb, "Best Improvements", manifest_rows, paired_map, True); ranked_sheet(wb, "Worst Regressions", manifest_rows, paired_map, False)
    video_by_case = {row["case_id"]: row for row in video_rows}; paired_by_case = {row["case_id"]: row for row in paired_video_rows}; provenance_fields = ["case_id", "row_manifest", "row_manifest_sha256", "result_npz", "result_sha256", "scene_act", "scene_sha256", "trajectory", "trajectory_sha256", "contact_mask", "contact_mask_sha256", "video", "video_sha256", "paired_video", "paired_video_sha256", "render_mode", "e188_worker", "e188_device"]
    provenance_rows = [{**row, "video": video_by_case[row["case_id"]]["video"], "video_sha256": video_by_case[row["case_id"]]["video_sha256"], "paired_video": paired_by_case[row["case_id"]]["paired_video"], "paired_video_sha256": paired_by_case[row["case_id"]]["paired_video_sha256"]} for row in manifest_rows]
    data_sheet(wb, "Artifact Provenance", provenance_fields, provenance_rows); overview(wb, paired_map, summary, len(manifest_rows))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, float): cell.number_format = "0.0000;-0.0000;-"
    wb.calculation.fullCalcOnLoad = True; wb.calculation.forceFullCalc = True; OUTPUT.parent.mkdir(parents=True, exist_ok=True); wb.save(OUTPUT)
    print(json.dumps({"status": "PASS", "output": str(OUTPUT.relative_to(REPO)), "sheets": wb.sheetnames, "rows": len(manifest_rows)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
