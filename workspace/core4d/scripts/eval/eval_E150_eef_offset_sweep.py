#!/usr/bin/env python3
"""Evaluate E150 contact-anchor eef_offset sweep."""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import math
import statistics
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

REPO = Path(__file__).resolve().parents[4]
VARIANTS_TSV = REPO / "workspace/core4d/scripts/E150/variants.tsv"
E147_EVAL = REPO / "workspace/core4d/scripts/eval/eval_E147_rubber_hand_collision.py"
RESULT_ROOT = REPO / "workspace/core4d/results/E150/contact_anchor_eef_offset_sweep"

ANCHORS = ["off05", "off08", "off11"]

SUMMARY_METRICS = [
    "eef_near_5cm_frac",
    "eef_near_10cm_frac",
    "hand_geom_near_5cm_frac",
    "hand_geom_near_10cm_frac",
    "hand_geom_penetration_frac",
    "hand_geom_deep_penetration_2cm_frac",
    "hand_object_physics_contact_frac",
    "leg_penetration_frac",
    "object_floor_contact_frac",
    "pelvis_min_m",
    "obj_err_mean_m",
]

DELTA_METRICS = [
    "eef_near_5cm_frac",
    "hand_geom_near_5cm_frac",
    "hand_geom_near_10cm_frac",
    "hand_geom_penetration_frac",
    "hand_geom_deep_penetration_2cm_frac",
    "hand_object_physics_contact_frac",
    "leg_penetration_frac",
    "object_floor_contact_frac",
    "obj_err_mean_m",
]


def rel(path: Path | str) -> str:
    if not path:
        return ""
    p = Path(path)
    try:
        return str(p.resolve().relative_to(REPO.resolve()))
    except Exception:
        return str(path)


def repo_path(text: str | Path) -> Path:
    p = Path(text)
    return p if p.is_absolute() else REPO / p


def fmt(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return f"{value:.8g}" if math.isfinite(value) else ""
    return str(value)


def finite(value: Any) -> float | None:
    if value in ("", None):
        return None
    try:
        out = float(value)
    except Exception:
        return None
    return out if math.isfinite(out) else None


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def write_tsv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: fmt(row.get(field, "")) for field in fields})


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def load_e147_eval():
    spec = importlib.util.spec_from_file_location("eval_E147_for_E150", E147_EVAL)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load {E147_EVAL}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def eval_one(eval_mod: Any, row: dict[str, str]) -> dict[str, Any]:
    offset_x = float(row["eef_offset_x"])
    eval_mod.EEF_OFFSET = np.asarray([offset_x, 0.0, 0.0], dtype=np.float64)
    qpos_path = repo_path(row["outdir_npz"])
    scene_xml = repo_path(row["rubber_scene_act"])
    item = eval_mod.evaluate_sequence(
        row=row,
        method=f"rubber_hull {row['anchor_variant']}",
        hand_collision_variant_id=row["hand_collision_variant_id"],
        qpos_path=qpos_path,
        scene_xml=scene_xml,
    )
    item.update(
        {
            "short_case_id": row["short_case_id"],
            "e148_variant": row["e148_variant"],
            "anchor_variant": row["anchor_variant"],
            "eef_offset_x": offset_x,
            "run_status": row["run_status"],
            "source_exp": row["source_exp"],
            "baseline_variant": row["baseline_variant"],
            "result_npz": row["result_npz"],
            "video": row["video"],
        }
    )
    return item


def build_metric_rows(rows: list[dict[str, str]], *, allow_missing: bool) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[str]]:
    eval_mod = load_e147_eval()
    metric_rows: list[dict[str, Any]] = []
    missing: list[dict[str, Any]] = []
    fields = [
        "short_case_id",
        "e148_variant",
        "anchor_variant",
        "eef_offset_x",
        "run_status",
        "source_exp",
        "baseline_variant",
        "result_npz",
        "video",
        *eval_mod.METRIC_FIELDS,
    ]
    for row in rows:
        outdir = repo_path(row["outdir_npz"])
        if not outdir.is_file():
            missing.append({"variant": row["variant"], "anchor_variant": row["anchor_variant"], "missing": rel(outdir)})
            if allow_missing:
                continue
            raise FileNotFoundError(f"missing E150 qpos for {row['variant']}: {outdir}")
        metric_rows.append(eval_one(eval_mod, row))
    return metric_rows, missing, fields


def avg(values: list[float]) -> float:
    vals = [v for v in values if math.isfinite(v)]
    return sum(vals) / len(vals) if vals else math.nan


def stdev(values: list[float]) -> float:
    vals = [v for v in values if math.isfinite(v)]
    return statistics.pstdev(vals) if len(vals) > 1 else 0.0


def summary_rows(metric_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for anchor in ANCHORS:
        rows = [r for r in metric_rows if r["anchor_variant"] == anchor]
        if not rows:
            continue
        item: dict[str, Any] = {
            "anchor_variant": anchor,
            "eef_offset_x": rows[0]["eef_offset_x"],
            "case_count": len(rows),
            "fall_count": sum(1 for r in rows if r.get("fall_flag")),
        }
        for metric in SUMMARY_METRICS:
            vals = [float(r[metric]) for r in rows if finite(r.get(metric)) is not None]
            item[f"{metric}_mean"] = avg(vals)
            item[f"{metric}_std"] = stdev(vals)
            item[f"{metric}_worst"] = max(vals) if metric != "pelvis_min_m" else min(vals)
        out.append(item)
    return out


def object_group_rows(metric_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    object_keys = sorted({r["object_key"] for r in metric_rows})
    for object_key in object_keys:
        for anchor in ANCHORS:
            rows = [r for r in metric_rows if r["object_key"] == object_key and r["anchor_variant"] == anchor]
            if not rows:
                continue
            item: dict[str, Any] = {"object_key": object_key, "anchor_variant": anchor, "case_count": len(rows)}
            for metric in ("hand_geom_near_5cm_frac", "hand_geom_penetration_frac", "hand_object_physics_contact_frac", "leg_penetration_frac"):
                item[metric] = avg([float(r[metric]) for r in rows if finite(r.get(metric)) is not None])
            out.append(item)
    return out


def delta_rows(metric_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_case_anchor = {(r["short_case_id"], r["anchor_variant"]): r for r in metric_rows}
    out: list[dict[str, Any]] = []
    for short_case_id in sorted({r["short_case_id"] for r in metric_rows}):
        base = by_case_anchor.get((short_case_id, "off05"))
        if not base:
            continue
        for anchor in ("off08", "off11"):
            row = by_case_anchor.get((short_case_id, anchor))
            if not row:
                continue
            item: dict[str, Any] = {
                "short_case_id": short_case_id,
                "case_id": row["case_id"],
                "object_key": row["object_key"],
                "anchor_variant": anchor,
                "eef_offset_x": row["eef_offset_x"],
                "baseline_anchor": "off05",
                "baseline_source_exp": base["source_exp"],
                "run_source_exp": row["source_exp"],
            }
            for metric in DELTA_METRICS:
                item[f"{metric}_baseline"] = base[metric]
                item[f"{metric}_run"] = row[metric]
                item[f"{metric}_delta"] = float(row[metric]) - float(base[metric])
            item["fall_flag"] = row["fall_flag"]
            item["success_near5_plus3pp_pen_not_up"] = (
                item["hand_geom_near_5cm_frac_delta"] >= 0.03
                and item["hand_geom_penetration_frac_delta"] <= 0.0
                and not bool(row["fall_flag"])
            )
            out.append(item)
    return out


def delta_summary_rows(deltas: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for anchor in ("off08", "off11"):
        rows = [r for r in deltas if r["anchor_variant"] == anchor]
        if not rows:
            continue
        item: dict[str, Any] = {
            "anchor_variant": anchor,
            "case_count": len(rows),
            "success_cases": sum(1 for r in rows if r["success_near5_plus3pp_pen_not_up"]),
        }
        for metric in DELTA_METRICS:
            vals = [float(r[f"{metric}_delta"]) for r in rows if finite(r.get(f"{metric}_delta")) is not None]
            item[f"{metric}_delta_mean"] = avg(vals)
            item[f"{metric}_delta_std"] = stdev(vals)
            item[f"{metric}_delta_worst"] = max(vals) if metric not in {"pelvis_min_m"} else min(vals)
        out.append(item)
    return out


def append_sheet(ws, rows: list[dict[str, Any]], fields: list[str]) -> None:
    ws.append(fields)
    for row in rows:
        ws.append([row.get(field, "") for field in fields])
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    for col in ws.columns:
        letter = col[0].column_letter
        width = min(max(len(str(cell.value or "")) for cell in col) + 2, 48)
        ws.column_dimensions[letter].width = width


def write_workbook(path: Path, sheets: dict[str, tuple[list[dict[str, Any]], list[str]]]) -> None:
    wb = Workbook()
    first = True
    for name, (rows, fields) in sheets.items():
        ws = wb.active if first else wb.create_sheet(name)
        ws.title = name
        append_sheet(ws, rows, fields)
        first = False
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_summary_md(path: Path, summaries: list[dict[str, Any]], delta_summaries: list[dict[str, Any]], deltas: list[dict[str, Any]]) -> None:
    lines = [
        "# E150 contact anchor eef_offset sweep",
        "",
        "Benchmark: E149 relaxed8 valid-like. Baseline off05 reuses E148/E147 rubber trajectories; off08/off11 are E150 CEM runs.",
        "",
        "## Offset averages",
        "",
        "| anchor | cases | fall | hand 5cm | EEF 5cm | hand pen | deep2 | physics contact | leg pen | obj err |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for row in summaries:
        lines.append(
            f"| `{row['anchor_variant']}` | {row['case_count']} | {row['fall_count']} | "
            f"{row['hand_geom_near_5cm_frac_mean']:.4f} | {row['eef_near_5cm_frac_mean']:.4f} | "
            f"{row['hand_geom_penetration_frac_mean']:.4f} | {row['hand_geom_deep_penetration_2cm_frac_mean']:.4f} | "
            f"{row['hand_object_physics_contact_frac_mean']:.4f} | {row['leg_penetration_frac_mean']:.4f} | "
            f"{row['obj_err_mean_m_mean']:.4f} |"
        )
    lines += [
        "",
        "## Mean deltas vs off05",
        "",
        "| anchor | success cases | hand 5cm delta | EEF 5cm delta | hand pen delta | deep2 delta | physics contact delta | leg pen delta | obj err delta |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for row in delta_summaries:
        lines.append(
            f"| `{row['anchor_variant']}` | {row['success_cases']}/{row['case_count']} | "
            f"{row['hand_geom_near_5cm_frac_delta_mean']:+.4f} | {row['eef_near_5cm_frac_delta_mean']:+.4f} | "
            f"{row['hand_geom_penetration_frac_delta_mean']:+.4f} | {row['hand_geom_deep_penetration_2cm_frac_delta_mean']:+.4f} | "
            f"{row['hand_object_physics_contact_frac_delta_mean']:+.4f} | {row['leg_penetration_frac_delta_mean']:+.4f} | "
            f"{row['obj_err_mean_m_delta_mean']:+.4f} |"
        )
    winners = [row for row in delta_summaries if row["hand_geom_near_5cm_frac_delta_mean"] >= 0.03 and row["hand_geom_penetration_frac_delta_mean"] <= 0.0]
    lines += [
        "",
        "## Criterion",
        "",
        "Predefined pass: mean hand-geom 5cm contact delta >= +0.03, hand penetration not higher than off05, no fall, and no obvious object/leg degradation.",
        "",
        f"- aggregate near+penetration criterion met by: `{', '.join(row['anchor_variant'] for row in winners) if winners else 'none'}`",
        f"- per-case near+penetration successes: `{sum(1 for row in deltas if row['success_near5_plus3pp_pen_not_up'])}/{len(deltas)}`",
        "",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--stage", default="full", choices=["smoke", "full"])
    parser.add_argument("--allow-missing", action="store_true")
    args = parser.parse_args()

    rows = read_tsv(VARIANTS_TSV)
    out_dir = RESULT_ROOT / "eval" / args.stage
    metric_rows, missing, metric_fields = build_metric_rows(rows, allow_missing=args.allow_missing)
    summaries = summary_rows(metric_rows)
    deltas = delta_rows(metric_rows)
    delta_summaries = delta_summary_rows(deltas)
    objects = object_group_rows(metric_rows)

    write_tsv(out_dir / "e150_method_metrics.tsv", metric_rows, metric_fields)
    write_tsv(out_dir / "e150_offset_summary.tsv", summaries, list(summaries[0].keys()) if summaries else [])
    write_tsv(out_dir / "e150_offset_delta.tsv", deltas, list(deltas[0].keys()) if deltas else [])
    write_tsv(out_dir / "e150_offset_delta_summary.tsv", delta_summaries, list(delta_summaries[0].keys()) if delta_summaries else [])
    write_tsv(out_dir / "e150_object_group_summary.tsv", objects, list(objects[0].keys()) if objects else [])
    write_tsv(out_dir / "e150_missing.tsv", missing, ["variant", "anchor_variant", "missing"])
    write_json(
        out_dir / "e150_eval_summary.json",
        {
            "stage": args.stage,
            "manifest_rows": len(rows),
            "method_metric_rows": len(metric_rows),
            "missing_rows": len(missing),
            "delta_rows": len(deltas),
            "offset_summary_rows": len(summaries),
            "xlsx": rel(RESULT_ROOT / "comparison/E150_contact_anchor_eef_offset_sweep.xlsx"),
        },
    )
    write_summary_md(out_dir / "e150_summary.md", summaries, delta_summaries, deltas)
    write_workbook(
        RESULT_ROOT / "comparison/E150_contact_anchor_eef_offset_sweep.xlsx",
        {
            "offset平均": (summaries, list(summaries[0].keys()) if summaries else []),
            "offset_delta": (deltas, list(deltas[0].keys()) if deltas else []),
            "delta平均": (delta_summaries, list(delta_summaries[0].keys()) if delta_summaries else []),
            "逐case_metrics": (metric_rows, metric_fields),
            "object分组": (objects, list(objects[0].keys()) if objects else []),
        },
    )
    print(f"method_rows={len(metric_rows)} delta_rows={len(deltas)} missing={len(missing)}")


if __name__ == "__main__":
    main()
