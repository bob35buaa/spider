#!/usr/bin/env python3
"""Evaluate E166 foot/smooth CEM arms against E163 baseline."""

from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
import sys
from pathlib import Path
from typing import Any

import numpy as np
import mujoco
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    import yaml
except Exception:  # pragma: no cover
    yaml = None

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from eval.core.core_metrics import EvalConfig, evaluate_sequence, mj_id, npz_qpos, person_idx_from_case  # noqa: E402
from eval_E156_clean8_gate_decay import (  # noqa: E402
    add_success_flags,
    contact_mask_for_case,
    kin_ref_for_scene,
    repo_path,
)


REPO = Path(__file__).resolve().parents[5]
VARIANTS = REPO / "workspace/core4d/scripts/experiments/E166/variants.tsv"
RESULT_ROOT = REPO / "workspace/core4d/results/E166/foot_smooth_retarget"
EVAL_ROOT = RESULT_ROOT / "eval"
ARM_ORDER = ["baseline", "B1", "B2", "A", "A_B2_postSmooth", "AplusB"]
CASE_ORDER = ["box021_035_p2", "box004_082_p1", "box004_083_p2"]
CONTACT_DROP_FAIL_TH = -0.05
PENETRATION_RISE_FAIL_TH = 0.05

METRIC_KEYS = [
    "success_tracked",
    "fall_flag",
    "hand_object_physics_contact_in_mask_frac",
    "hand_object_physics_contact_in_rl_mask_frac",
    "hand_object_physics_contact_3mm_in_mask_frac",
    "hand_object_physics_contact_5mm_in_mask_frac",
    "hand_object_physics_penetration_3mm_frame_frac",
    "hand_object_physics_penetration_5mm_frame_frac",
    "hand_geom_penetration_2mm_frac",
    "hand_geom_penetration_5mm_frac",
    "hand_object_release_false_contact_3mm_frac",
    "track_joint_err_deg_mean",
    "track_eef_pos_err_cm_mean",
    "track_eef_ori_err_deg_mean",
    "track_obj_pos_err_cm_mean",
    "track_obj_ori_err_deg_mean",
    "track_root_pos_err_cm_mean",
    "track_root_ori_err_deg_mean",
    "qpos_speed_l2_p95",
    "qpos_accel_l2_p95",
    "qpos_jerk_l2_p95",
    "trackbody_speed_max",
    "ankle_speed_max",
    "wrist_speed_max",
    "trackbody_acc_max",
    "ankle_acc_max",
    "trackbody_jerk_p95",
    "ankle_jerk_p95",
    "obj_speed_max",
    "foot_slip_max_m",
    "foot_ground_dev_max_m",
    "foot_grounded_frame_frac",
]

HEALTH_AGGS = {
    "sample_smooth_accel_p95_mean": "mean",
    "sample_smooth_accel_p95_max": "max",
    "sample_smooth_jerk_p95_mean": "mean",
    "sample_smooth_jerk_p95_max": "max",
    "sample_smooth_penalty_mean": "mean",
    "sample_smooth_penalty_max": "max",
    "sample_foot_slip_speed_mean_mean": "mean",
    "sample_foot_slip_speed_peak_mean": "mean",
    "sample_foot_slip_speed_peak_max": "max",
    "sample_foot_ground_dev_mean_mean": "mean",
    "sample_foot_ground_dev_peak_mean": "mean",
    "sample_foot_ground_dev_peak_max": "max",
    "sample_foot_penalty_mean": "mean",
    "sample_foot_penalty_max": "max",
}

TRACK_BODY_NAMES = [
    "left_ankle_roll_link",
    "right_ankle_roll_link",
    "left_wrist_yaw_link",
    "right_wrist_yaw_link",
]
ANKLE_BODY_NAMES = ["left_ankle_roll_link", "right_ankle_roll_link"]
WRIST_BODY_NAMES = ["left_wrist_yaw_link", "right_wrist_yaw_link"]


def rel(path: str | Path) -> str:
    p = Path(path)
    try:
        return str(p.resolve().relative_to(REPO.resolve()))
    except Exception:
        return str(path)


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def write_tsv(path: Path, rows: list[dict[str, Any]], fields: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fields is None:
        seen: list[str] = []
        for row in rows:
            for key in row:
                if key not in seen:
                    seen.append(key)
        fields = seen
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: format_tsv_value(row.get(field, "")) for field in fields})


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def format_tsv_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, float) and not math.isfinite(value):
        return ""
    return value


def finite(value: Any, default: float = math.nan) -> float:
    try:
        val = float(value)
    except Exception:
        return default
    return val if math.isfinite(val) else default


def mean(values: list[Any]) -> float:
    vals = [finite(v) for v in values]
    vals = [v for v in vals if math.isfinite(v)]
    return statistics.fmean(vals) if vals else math.nan


def bool_text(value: Any) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "pass"}


def load_yaml(path: Path) -> dict[str, Any]:
    if yaml is None or not path.is_file():
        return {}
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    return data if isinstance(data, dict) else {}


def close_enough(value: Any, target: float, tol: float = 1e-7) -> bool:
    try:
        return abs(float(value) - target) <= tol
    except Exception:
        return False


def qpos_path_for_row(row: dict[str, str]) -> Path:
    if row["arm_kind"] == "postprocess":
        return repo_path(row["postprocess_output_npz"])
    outdir = repo_path(row["outdir_npz"])
    root = repo_path(row["result_npz"])
    return outdir if outdir.is_file() else root


def video_path_for_row(row: dict[str, str]) -> Path:
    return repo_path(row["video"])


def config_path_for_row(row: dict[str, str]) -> Path:
    if row["arm_kind"] != "cem":
        return Path("")
    outdir = Path(row["outdir_npz"])
    if outdir.name == "trajectory_mjwp_act.npz":
        return repo_path(outdir.parent / "config_act.yaml")
    return Path("")


def artifact_check(row: dict[str, str]) -> dict[str, Any]:
    qpos_path = qpos_path_for_row(row)
    video_path = video_path_for_row(row)
    config_path = config_path_for_row(row)
    smooth_report_path = qpos_path.with_suffix(qpos_path.suffix + ".smooth_report.json")
    cfg = load_yaml(config_path)
    arm = row["arm"]
    checks: dict[str, Any] = {
        "variant": row["variant"],
        "short_case_id": row["short_case_id"],
        "arm": arm,
        "arm_kind": row["arm_kind"],
        "qpos_exists": qpos_path.is_file(),
        "video_exists": row["arm_kind"] == "postprocess"
        or (arm in {"baseline", "B1", "A", "AplusB"} and video_path.is_file()),
        "config_exists": row["arm_kind"] != "cem" or config_path.is_file(),
        "qpos_path": rel(qpos_path),
        "video": rel(video_path) if video_path else "",
        "config_act": rel(config_path) if config_path else "",
        "smooth_report": rel(smooth_report_path) if row["arm_kind"] == "postprocess" else "",
    }
    if row["arm_kind"] == "postprocess":
        checks["smooth_report_exists"] = smooth_report_path.is_file()
    if arm in {"B1", "AplusB"}:
        checks["config_smooth_enabled_ok"] = cfg.get("cem_smooth_enabled") is True
        checks["config_smooth_accel_weight_ok"] = close_enough(cfg.get("cem_smooth_accel_weight"), 0.0005)
        checks["config_smooth_jerk_weight_ok"] = close_enough(cfg.get("cem_smooth_jerk_weight"), 0.00002)
    elif row["arm_kind"] == "cem":
        checks["config_smooth_enabled_ok"] = cfg.get("cem_smooth_enabled") in {False, None}
        checks["config_smooth_accel_weight_ok"] = close_enough(cfg.get("cem_smooth_accel_weight", 0.0), 0.0)
        checks["config_smooth_jerk_weight_ok"] = close_enough(cfg.get("cem_smooth_jerk_weight", 0.0), 0.0)
    if arm in {"A", "AplusB"}:
        checks["config_ankle_weight_ok"] = close_enough(cfg.get("local_frame_ankle_weight"), 2.0)
        checks["config_foot_slip_enabled_ok"] = cfg.get("foot_slip_enabled") is True
        checks["config_foot_ground_enabled_ok"] = cfg.get("foot_ground_enabled") is True
    elif row["arm_kind"] == "cem":
        checks["config_ankle_weight_ok"] = close_enough(cfg.get("local_frame_ankle_weight", 1.0), 1.0)
        checks["config_foot_slip_enabled_ok"] = cfg.get("foot_slip_enabled") in {False, None}
        checks["config_foot_ground_enabled_ok"] = cfg.get("foot_ground_enabled") in {False, None}
    if row["arm_kind"] == "cem":
        checks["config_peak_margin_disabled_ok"] = cfg.get("cem_peak_margin_enabled") in {False, None}
    bool_keys = [key for key in checks if key.endswith("_exists") or key.endswith("_ok")]
    checks["artifact_ok"] = all(bool(checks[key]) for key in bool_keys)
    return checks


def reduce_array(arr: np.ndarray, agg: str) -> float:
    vals = np.asarray(arr, dtype=np.float64)
    vals = vals[np.isfinite(vals)]
    if vals.size == 0:
        return math.nan
    if agg == "max":
        return float(vals.max())
    return float(vals.mean())


def qpos_kinematic_health(qpos_path: Path) -> dict[str, Any]:
    out = {
        "qpos_speed_l2_p95": math.nan,
        "qpos_accel_l2_p95": math.nan,
        "qpos_jerk_l2_p95": math.nan,
        "qpos_frames": 0,
    }
    if not qpos_path.is_file():
        return out
    data = np.load(qpos_path, allow_pickle=True)
    if "qpos" not in data.files:
        return out
    qpos = np.asarray(data["qpos"], dtype=np.float64)
    if qpos.ndim < 2 or qpos.shape[0] < 4:
        return out
    flat = qpos.reshape(qpos.shape[0], -1)
    fps = 50.0
    if "time" in data.files:
        time = np.asarray(data["time"], dtype=np.float64).reshape(qpos.shape[0], -1)[:, 0]
        dt = np.diff(time)
        dt = dt[np.isfinite(dt) & (dt > 0)]
        if dt.size:
            fps = float(1.0 / np.median(dt))
    speed = np.linalg.norm(np.diff(flat, axis=0), axis=1) * fps
    accel = np.linalg.norm(np.diff(flat, n=2, axis=0), axis=1) * (fps**2)
    jerk = np.linalg.norm(np.diff(flat, n=3, axis=0), axis=1) * (fps**3)
    out.update(
        {
            "qpos_speed_l2_p95": float(np.percentile(speed, 95)) if speed.size else math.nan,
            "qpos_accel_l2_p95": float(np.percentile(accel, 95)) if accel.size else math.nan,
            "qpos_jerk_l2_p95": float(np.percentile(jerk, 95)) if jerk.size else math.nan,
            "qpos_frames": int(qpos.shape[0]),
        }
    )
    return out


def fps_from_npz(qpos_path: Path, default: float) -> float:
    if not qpos_path.is_file():
        return default
    data = np.load(qpos_path, allow_pickle=True)
    if "time" not in data.files:
        return default
    time = np.asarray(data["time"], dtype=np.float64).reshape(data["time"].shape[0], -1)[:, 0]
    dt = np.diff(time)
    dt = dt[np.isfinite(dt) & (dt > 0)]
    return float(1.0 / np.median(dt)) if dt.size else default


def contiguous_segments(mask: np.ndarray) -> list[tuple[int, int]]:
    idx = np.flatnonzero(mask)
    if idx.size == 0:
        return []
    breaks = np.where(np.diff(idx) > 1)[0]
    starts = np.r_[idx[0], idx[breaks + 1]]
    ends = np.r_[idx[breaks], idx[-1]]
    return [(int(s), int(e) + 1) for s, e in zip(starts, ends)]


def foot_motion_metrics(ankles: np.ndarray) -> dict[str, float]:
    ground_z = np.percentile(ankles[:, :, 2], 5, axis=0)
    grounded = ankles[:, :, 2] <= (ground_z[None, :] + 0.05)
    slip_values: list[float] = []
    ground_dev_values: list[float] = []
    for foot_idx in range(ankles.shape[1]):
        for start, end in contiguous_segments(grounded[:, foot_idx]):
            if end - start < 2:
                continue
            xy = ankles[start:end, foot_idx, :2]
            z = ankles[start:end, foot_idx, 2]
            slip_values.append(float(np.linalg.norm(xy - xy[0], axis=-1).max()))
            ground_dev_values.append(float(np.abs(z - ground_z[foot_idx]).max()))
    return {
        "foot_slip_max_m": max(slip_values) if slip_values else math.nan,
        "foot_ground_dev_max_m": max(ground_dev_values) if ground_dev_values else math.nan,
        "foot_grounded_frame_frac": float(np.mean(np.any(grounded, axis=1))),
    }


def body_motion_health(qpos_path: Path, scene_xml: Path, cfg: EvalConfig) -> dict[str, Any]:
    out = {
        "trackbody_speed_max": math.nan,
        "ankle_speed_max": math.nan,
        "wrist_speed_max": math.nan,
        "trackbody_acc_max": math.nan,
        "ankle_acc_max": math.nan,
        "trackbody_jerk_p95": math.nan,
        "ankle_jerk_p95": math.nan,
        "obj_speed_max": math.nan,
        "foot_slip_max_m": math.nan,
        "foot_ground_dev_max_m": math.nan,
        "foot_grounded_frame_frac": math.nan,
    }
    if not qpos_path.is_file() or not scene_xml.is_file():
        return out
    qpos, _ = npz_qpos(qpos_path)
    if qpos.shape[0] < 4:
        return out
    model = mujoco.MjModel.from_xml_path(str(scene_xml))
    data = mujoco.MjData(model)
    track_ids = [
        bid
        for name in TRACK_BODY_NAMES
        if (bid := mj_id(model, mujoco.mjtObj.mjOBJ_BODY, name)) >= 0
    ]
    ankle_ids = [
        bid
        for name in ANKLE_BODY_NAMES
        if (bid := mj_id(model, mujoco.mjtObj.mjOBJ_BODY, name)) >= 0
    ]
    wrist_ids = [
        bid
        for name in WRIST_BODY_NAMES
        if (bid := mj_id(model, mujoco.mjtObj.mjOBJ_BODY, name)) >= 0
    ]
    object_id = mj_id(model, mujoco.mjtObj.mjOBJ_BODY, "object")
    if not track_ids or not ankle_ids or object_id < 0:
        return out

    body_pos = []
    ankle_pos = []
    wrist_pos = []
    obj_pos = []
    for q in qpos:
        data.qpos[:] = q
        data.qvel[:] = 0.0
        mujoco.mj_forward(model, data)
        body_pos.append(data.xpos[track_ids].copy())
        ankle_pos.append(data.xpos[ankle_ids].copy())
        if wrist_ids:
            wrist_pos.append(data.xpos[wrist_ids].copy())
        obj_pos.append(data.xpos[object_id].copy())

    fps = fps_from_npz(qpos_path, cfg.fps)
    tracked = np.asarray(body_pos, dtype=np.float64)
    ankles = np.asarray(ankle_pos, dtype=np.float64)
    wrists = np.asarray(wrist_pos, dtype=np.float64) if wrist_pos else np.empty((len(qpos), 0, 3))
    obj = np.asarray(obj_pos, dtype=np.float64)
    track_speed = np.linalg.norm(np.diff(tracked, axis=0), axis=-1) * fps
    ankle_speed = np.linalg.norm(np.diff(ankles, axis=0), axis=-1) * fps
    track_acc = np.linalg.norm(np.diff(tracked, n=2, axis=0), axis=-1) * (fps**2)
    ankle_acc = np.linalg.norm(np.diff(ankles, n=2, axis=0), axis=-1) * (fps**2)
    track_jerk = np.linalg.norm(np.diff(tracked, n=3, axis=0), axis=-1) * (fps**3)
    ankle_jerk = np.linalg.norm(np.diff(ankles, n=3, axis=0), axis=-1) * (fps**3)
    obj_speed = np.linalg.norm(np.diff(obj, axis=0), axis=-1) * fps

    out.update(
        {
            "trackbody_speed_max": float(np.max(track_speed)) if track_speed.size else math.nan,
            "ankle_speed_max": float(np.max(ankle_speed)) if ankle_speed.size else math.nan,
            "wrist_speed_max": float(np.max(np.linalg.norm(np.diff(wrists, axis=0), axis=-1) * fps))
            if wrists.size
            else math.nan,
            "trackbody_acc_max": float(np.max(track_acc)) if track_acc.size else math.nan,
            "ankle_acc_max": float(np.max(ankle_acc)) if ankle_acc.size else math.nan,
            "trackbody_jerk_p95": float(np.percentile(track_jerk, 95)) if track_jerk.size else math.nan,
            "ankle_jerk_p95": float(np.percentile(ankle_jerk, 95)) if ankle_jerk.size else math.nan,
            "obj_speed_max": float(np.max(obj_speed)) if obj_speed.size else math.nan,
        }
    )
    out.update(foot_motion_metrics(ankles))
    return out


def run_health(qpos_path: Path, scene_xml: Path, cfg: EvalConfig) -> dict[str, Any]:
    out: dict[str, Any] = {key: math.nan for key in HEALTH_AGGS}
    out.update(qpos_kinematic_health(qpos_path))
    out.update(body_motion_health(qpos_path, scene_xml, cfg))
    if not qpos_path.is_file():
        return out
    data = np.load(qpos_path, allow_pickle=True)
    for key, agg in HEALTH_AGGS.items():
        if key in data.files:
            out[key] = reduce_array(data[key], agg)
    out["has_smooth_health"] = all(
        key in data.files for key in ("sample_smooth_accel_p95_mean", "sample_smooth_jerk_p95_mean")
    )
    out["has_foot_health"] = all(
        key in data.files for key in ("sample_foot_slip_speed_peak_mean", "sample_foot_ground_dev_peak_mean")
    )
    return out


def evaluate_one(row: dict[str, str], cfg: EvalConfig) -> dict[str, Any] | None:
    qpos_path = qpos_path_for_row(row)
    scene = repo_path(row["rubber_scene_act"])
    if not qpos_path.is_file() or not scene.is_file():
        return None
    item = evaluate_sequence(
        row=row,
        method=row["method"],
        hand_collision_variant_id="rubber_hull",
        qpos_path=qpos_path,
        scene_xml=scene,
        config=cfg,
        kin_ref_path=kin_ref_for_scene(scene),
        contact_mask_path=contact_mask_for_case(row["short_case_id"]),
        person_idx=person_idx_from_case(row["short_case_id"]),
    )
    add_success_flags(item, cfg)
    item.update(
        {
            "short_case_id": row["short_case_id"],
            "variant": row["variant"],
            "arm": row["arm"],
            "arm_kind": row["arm_kind"],
            "method": row["method"],
            "source_exp": row["source_exp"],
            "split": row["split"],
            "result_npz": rel(qpos_path),
            "video": rel(video_path_for_row(row)) if video_path_for_row(row).is_file() else "",
        }
    )
    item.update(run_health(qpos_path, scene, cfg))
    return item


def add_baseline_deltas(rows: list[dict[str, Any]]) -> None:
    baseline_by_case = {
        row["short_case_id"]: row for row in rows if row.get("arm") == "baseline"
    }
    delta_keys = [key for key in METRIC_KEYS if key not in {"success_tracked", "fall_flag"}]
    for row in rows:
        base = baseline_by_case.get(row["short_case_id"])
        for key in delta_keys:
            row[f"{key}_delta_vs_baseline"] = (
                finite(row.get(key)) - finite(base.get(key)) if base else math.nan
            )
        if row.get("arm") == "baseline":
            row["contact_no_regression"] = True
            row["penetration_no_regression"] = True
            row["spider_gate_pass"] = bool(row.get("success_tracked")) and not bool(row.get("fall_flag"))
            row["status"] = "baseline"
            continue
        raw_delta = finite(row.get("hand_object_physics_contact_in_mask_frac_delta_vs_baseline"))
        clean3_delta = finite(row.get("hand_object_physics_contact_3mm_in_mask_frac_delta_vs_baseline"))
        pen3_delta = finite(row.get("hand_object_physics_penetration_3mm_frame_frac_delta_vs_baseline"))
        row["contact_no_regression"] = raw_delta >= CONTACT_DROP_FAIL_TH and clean3_delta >= CONTACT_DROP_FAIL_TH
        row["penetration_no_regression"] = pen3_delta <= PENETRATION_RISE_FAIL_TH
        row["spider_gate_pass"] = bool(row.get("success_tracked")) and not bool(row.get("fall_flag")) and bool(
            row["contact_no_regression"]
        ) and bool(row["penetration_no_regression"])
        if not bool(row.get("success_tracked")):
            row["status"] = "tracking_fail"
        elif bool(row.get("fall_flag")):
            row["status"] = "fall"
        elif not bool(row["contact_no_regression"]):
            row["status"] = "contact_regression"
        elif not bool(row["penetration_no_regression"]):
            row["status"] = "penetration_regression"
        else:
            row["status"] = "pass"


def summarize(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for arm in ARM_ORDER:
        arm_rows = [row for row in rows if row.get("arm") == arm]
        if not arm_rows:
            continue
        item: dict[str, Any] = {
            "arm": arm,
            "n_cases": len(arm_rows),
            "pass_cases": sum(1 for row in arm_rows if row.get("spider_gate_pass")),
            "tracked_cases": sum(1 for row in arm_rows if row.get("success_tracked")),
            "fall_cases": sum(1 for row in arm_rows if row.get("fall_flag")),
            "failed_cases": ",".join(str(row.get("short_case_id")) for row in arm_rows if not row.get("spider_gate_pass")),
        }
        for key in METRIC_KEYS:
            if key in {"success_tracked", "fall_flag"}:
                item[f"{key}_mean"] = mean([1.0 if row.get(key) else 0.0 for row in arm_rows])
            else:
                item[f"{key}_mean"] = mean([row.get(key) for row in arm_rows])
                item[f"{key}_delta_vs_baseline_mean"] = mean(
                    [row.get(f"{key}_delta_vs_baseline") for row in arm_rows if row.get("arm") != "baseline"]
                )
        out.append(item)
    return out


def write_summary_md(path: Path, summary_rows: list[dict[str, Any]], metric_rows: list[dict[str, Any]], missing: list[dict[str, Any]]) -> None:
    lines = [
        "# E166 foot/smooth retarget eval",
        "",
        "## Summary",
        "",
        "| arm | n | pass | tracked | fall | raw contact | clean3 contact | pen3 | qpos jerk p95 | track jerk p95 | ankle acc max | foot slip | failed |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---|",
    ]
    for row in summary_rows:
        lines.append(
            "| {arm} | {n_cases} | {pass_cases} | {tracked_cases} | {fall_cases} | {raw:.3f} | {clean:.3f} | {pen:.3f} | {qjerk:.1f} | {tjerk:.1f} | {aacc:.1f} | {slip:.3f} | {failed} |".format(
                arm=row.get("arm", ""),
                n_cases=row.get("n_cases", ""),
                pass_cases=row.get("pass_cases", ""),
                tracked_cases=row.get("tracked_cases", ""),
                fall_cases=row.get("fall_cases", ""),
                raw=finite(row.get("hand_object_physics_contact_in_mask_frac_mean")),
                clean=finite(row.get("hand_object_physics_contact_3mm_in_mask_frac_mean")),
                pen=finite(row.get("hand_object_physics_penetration_3mm_frame_frac_mean")),
                qjerk=finite(row.get("qpos_jerk_l2_p95_mean")),
                tjerk=finite(row.get("trackbody_jerk_p95_mean")),
                aacc=finite(row.get("ankle_acc_max_mean")),
                slip=finite(row.get("foot_slip_max_m_mean")),
                failed=row.get("failed_cases", ""),
            )
        )
    lines.extend(["", "## Per Case", ""])
    lines.extend(
        [
            "| case | arm | status | raw Δ | clean3 Δ | pen3 Δ | qpos jerk Δ | artifact |",
            "|---|---|---|---:|---:|---:|---:|---|",
        ]
    )
    for row in sorted(metric_rows, key=lambda r: (CASE_ORDER.index(r["short_case_id"]), ARM_ORDER.index(r["arm"]))):
        lines.append(
            "| {case} | {arm} | {status} | {raw:.3f} | {clean:.3f} | {pen:.3f} | {jerk:.1f} | {artifact} |".format(
                case=row.get("short_case_id", ""),
                arm=row.get("arm", ""),
                status=row.get("status", ""),
                raw=finite(row.get("hand_object_physics_contact_in_mask_frac_delta_vs_baseline")),
                clean=finite(row.get("hand_object_physics_contact_3mm_in_mask_frac_delta_vs_baseline")),
                pen=finite(row.get("hand_object_physics_penetration_3mm_frame_frac_delta_vs_baseline")),
                jerk=finite(row.get("qpos_jerk_l2_p95_delta_vs_baseline")),
                artifact=row.get("artifact_ok", ""),
            )
        )
    if missing:
        lines.extend(["", "## Missing", ""])
        for row in missing:
            lines.append(f"- {row['short_case_id']} {row['arm']}: {row['reason']}")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def xlsx_value(value: Any) -> Any:
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float) and not math.isfinite(value):
        return ""
    return value


def append_table(
    wb: Workbook,
    title: str,
    rows: list[dict[str, Any]],
    columns: list[tuple[str, str, str]],
    fill: str,
    *,
    hidden: bool = False,
) -> None:
    if len(wb.sheetnames) == 1 and wb.active.max_row == 1 and wb.active["A1"].value is None:
        ws = wb.active
        ws.delete_rows(1)
    else:
        ws = wb.create_sheet(title)
    ws.title = title
    ws.append([header for header, _, _ in columns])
    for row in rows:
        ws.append([xlsx_value(row.get(key, "")) for _, key, _ in columns])

    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, (_, _, fmt) in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        for cell in ws[letter][1:]:
            if fmt == "pct":
                cell.number_format = "0.0%"
            elif fmt == "pct_delta":
                cell.number_format = "0.0%;[Red]-0.0%"
            elif fmt == "num":
                cell.number_format = "0.00"
            elif fmt == "int":
                cell.number_format = "0"
            elif fmt == "m":
                cell.number_format = "0.000"
        width = min(max(len(str(cell.value or "")) for cell in ws[letter]) + 2, 50)
        ws.column_dimensions[letter].width = max(width, 10)
    if hidden:
        ws.sheet_state = "hidden"


def workbook_main_rows(summary_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    notes = {
        "baseline": "E163 narrowSurfaceBand clean8 产物复用；E166 对比基线。",
        "B1": "CEM 内 smooth accel/jerk penalty；接触 gate 通过，FK jerk 下降但 qpos jerk 未降。",
        "B2": "handoff 后处理平滑；降抖最强，但 box021_035_p2 raw contact 回归。",
        "A": "ankle weight + foot slip/ground penalty；本轮 CEM-only 最稳候选。",
        "A_B2_postSmooth": "先跑 A，再做 B2 同款 CPU 后处理平滑；用于验证 A+B2 是否比 A/AplusB 更好。",
        "AplusB": "A 与 B1 叠加；clean3/penetration 最好，但 tracking/ankle 代价更高。",
    }
    out: list[dict[str, Any]] = []
    for row in sorted(summary_rows, key=lambda r: ARM_ORDER.index(str(r.get("arm")))):
        item = dict(row)
        arm = str(item.get("arm", ""))
        item["方法"] = arm
        item["说明"] = notes.get(arm, "")
        item["失败case"] = item.get("failed_cases", "")
        item["通过case"] = item.get("pass_cases", "")
        item["case数"] = item.get("n_cases", "")
        out.append(item)
    return out


def workbook_case_rows(metric_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for row in metric_rows:
        item = dict(row)
        item["case"] = row.get("short_case_id", "")
        item["方法"] = row.get("arm", "")
        if row.get("status") == "baseline":
            item["状态"] = "baseline"
        elif row.get("status") == "pass":
            item["状态"] = "通过"
        elif row.get("status") == "contact_regression":
            item["状态"] = "接触退化"
        elif row.get("status") == "penetration_regression":
            item["状态"] = "穿透退化"
        elif row.get("status") == "fall":
            item["状态"] = "摔倒"
        elif row.get("status") == "tracking_fail":
            item["状态"] = "tracking失败"
        else:
            item["状态"] = row.get("status", "")
        item["Table4完整"] = all(
            math.isfinite(finite(row.get(key)))
            for key in (
                "track_joint_err_deg_mean",
                "track_eef_pos_err_cm_mean",
                "track_eef_ori_err_deg_mean",
                "track_obj_pos_err_cm_mean",
                "track_obj_ori_err_deg_mean",
            )
        )
        out.append(item)
    return sorted(out, key=lambda r: (CASE_ORDER.index(str(r.get("case"))), ARM_ORDER.index(str(r.get("方法")))))


def workbook_health_rows(metric_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return sorted(
        [dict(row, case=row.get("short_case_id", ""), 方法=row.get("arm", "")) for row in metric_rows],
        key=lambda r: (CASE_ORDER.index(str(r.get("case"))), ARM_ORDER.index(str(r.get("方法")))),
    )


def write_workbook(
    eval_dir: Path,
    summary_rows: list[dict[str, Any]],
    metric_rows: list[dict[str, Any]],
    artifact_rows: list[dict[str, Any]],
) -> Path:
    wb = Workbook()
    main_cols = [
        ("方法", "方法", "text"),
        ("case数", "case数", "int"),
        ("通过case", "通过case", "int"),
        ("失败case", "失败case", "text"),
        ("tracked", "tracked_cases", "int"),
        ("fall", "fall_cases", "int"),
        ("物理接触(raw)", "hand_object_physics_contact_in_mask_frac_mean", "pct"),
        ("raw Δ vs baseline", "hand_object_physics_contact_in_mask_frac_delta_vs_baseline_mean", "pct_delta"),
        ("RL mask接触", "hand_object_physics_contact_in_rl_mask_frac_mean", "pct"),
        ("clean3接触", "hand_object_physics_contact_3mm_in_mask_frac_mean", "pct"),
        ("clean3 Δ", "hand_object_physics_contact_3mm_in_mask_frac_delta_vs_baseline_mean", "pct_delta"),
        ("clean5接触", "hand_object_physics_contact_5mm_in_mask_frac_mean", "pct"),
        ("物理穿透3mm", "hand_object_physics_penetration_3mm_frame_frac_mean", "pct"),
        ("物理穿透3mm Δ", "hand_object_physics_penetration_3mm_frame_frac_delta_vs_baseline_mean", "pct_delta"),
        ("几何穿透2mm", "hand_geom_penetration_2mm_frac_mean", "pct"),
        ("release误接触3mm", "hand_object_release_false_contact_3mm_frac_mean", "pct"),
        ("关节误差(°)", "track_joint_err_deg_mean_mean", "num"),
        ("EEF误差(cm)", "track_eef_pos_err_cm_mean_mean", "num"),
        ("物体位置误差(cm)", "track_obj_pos_err_cm_mean_mean", "num"),
        ("root误差(cm)", "track_root_pos_err_cm_mean_mean", "num"),
        ("qpos jerk p95", "qpos_jerk_l2_p95_mean", "num"),
        ("qpos jerk Δ", "qpos_jerk_l2_p95_delta_vs_baseline_mean", "num"),
        ("trackbody jerk p95", "trackbody_jerk_p95_mean", "num"),
        ("trackbody jerk Δ", "trackbody_jerk_p95_delta_vs_baseline_mean", "num"),
        ("ankle acc max", "ankle_acc_max_mean", "num"),
        ("ankle acc Δ", "ankle_acc_max_delta_vs_baseline_mean", "num"),
        ("obj speed max", "obj_speed_max_mean", "num"),
        ("foot slip max(m)", "foot_slip_max_m_mean", "m"),
        ("foot slip Δ", "foot_slip_max_m_delta_vs_baseline_mean", "m"),
        ("grounded frame frac", "foot_grounded_frame_frac_mean", "pct"),
        ("说明", "说明", "text"),
    ]
    append_table(wb, "主表", workbook_main_rows(summary_rows), main_cols, "1F4E79")

    case_cols = [
        ("方法", "方法", "text"),
        ("case", "case", "text"),
        ("状态", "状态", "text"),
        ("SPIDER gate", "spider_gate_pass", "text"),
        ("contact no-reg", "contact_no_regression", "text"),
        ("penetration no-reg", "penetration_no_regression", "text"),
        ("raw接触", "hand_object_physics_contact_in_mask_frac", "pct"),
        ("raw Δ", "hand_object_physics_contact_in_mask_frac_delta_vs_baseline", "pct_delta"),
        ("RL mask接触", "hand_object_physics_contact_in_rl_mask_frac", "pct"),
        ("clean3接触", "hand_object_physics_contact_3mm_in_mask_frac", "pct"),
        ("clean3 Δ", "hand_object_physics_contact_3mm_in_mask_frac_delta_vs_baseline", "pct_delta"),
        ("clean5接触", "hand_object_physics_contact_5mm_in_mask_frac", "pct"),
        ("物理穿透3mm", "hand_object_physics_penetration_3mm_frame_frac", "pct"),
        ("物理穿透3mm Δ", "hand_object_physics_penetration_3mm_frame_frac_delta_vs_baseline", "pct_delta"),
        ("物理穿透5mm", "hand_object_physics_penetration_5mm_frame_frac", "pct"),
        ("几何穿透2mm", "hand_geom_penetration_2mm_frac", "pct"),
        ("release误接触3mm", "hand_object_release_false_contact_3mm_frac", "pct"),
        ("tracked", "success_tracked", "text"),
        ("fall", "fall_flag", "text"),
        ("Table4完整", "Table4完整", "text"),
        ("关节误差(°)", "track_joint_err_deg_mean", "num"),
        ("EEF误差(cm)", "track_eef_pos_err_cm_mean", "num"),
        ("EEF朝向误差(°)", "track_eef_ori_err_deg_mean", "num"),
        ("物体位置误差(cm)", "track_obj_pos_err_cm_mean", "num"),
        ("物体朝向误差(°)", "track_obj_ori_err_deg_mean", "num"),
        ("root误差(cm)", "track_root_pos_err_cm_mean", "num"),
        ("root朝向误差(°)", "track_root_ori_err_deg_mean", "num"),
        ("qpos jerk p95", "qpos_jerk_l2_p95", "num"),
        ("qpos jerk Δ", "qpos_jerk_l2_p95_delta_vs_baseline", "num"),
        ("trackbody jerk p95", "trackbody_jerk_p95", "num"),
        ("trackbody jerk Δ", "trackbody_jerk_p95_delta_vs_baseline", "num"),
        ("ankle acc max", "ankle_acc_max", "num"),
        ("ankle acc Δ", "ankle_acc_max_delta_vs_baseline", "num"),
        ("obj speed max", "obj_speed_max", "num"),
        ("obj speed Δ", "obj_speed_max_delta_vs_baseline", "num"),
        ("foot slip max(m)", "foot_slip_max_m", "m"),
        ("foot slip Δ", "foot_slip_max_m_delta_vs_baseline", "m"),
        ("ground dev max(m)", "foot_ground_dev_max_m", "m"),
        ("grounded frame frac", "foot_grounded_frame_frac", "pct"),
        ("qpos_path", "qpos_path", "text"),
        ("video", "video", "text"),
    ]
    append_table(wb, "逐case", workbook_case_rows(metric_rows), case_cols, "548235")

    health_cols = [
        ("方法", "方法", "text"),
        ("case", "case", "text"),
        ("has smooth health", "has_smooth_health", "text"),
        ("has foot health", "has_foot_health", "text"),
        ("smooth accel p95 mean", "sample_smooth_accel_p95_mean", "num"),
        ("smooth accel p95 max", "sample_smooth_accel_p95_max", "num"),
        ("smooth jerk p95 mean", "sample_smooth_jerk_p95_mean", "num"),
        ("smooth jerk p95 max", "sample_smooth_jerk_p95_max", "num"),
        ("smooth penalty mean", "sample_smooth_penalty_mean", "num"),
        ("smooth penalty max", "sample_smooth_penalty_max", "num"),
        ("foot slip speed mean", "sample_foot_slip_speed_mean_mean", "num"),
        ("foot slip speed peak mean", "sample_foot_slip_speed_peak_mean", "num"),
        ("foot slip speed peak max", "sample_foot_slip_speed_peak_max", "num"),
        ("foot ground dev mean", "sample_foot_ground_dev_mean_mean", "m"),
        ("foot ground dev peak mean", "sample_foot_ground_dev_peak_mean", "m"),
        ("foot ground dev peak max", "sample_foot_ground_dev_peak_max", "m"),
        ("foot penalty mean", "sample_foot_penalty_mean", "num"),
        ("foot penalty max", "sample_foot_penalty_max", "num"),
        ("FK foot slip max(m)", "foot_slip_max_m", "m"),
        ("FK foot ground dev max(m)", "foot_ground_dev_max_m", "m"),
        ("FK grounded frame frac", "foot_grounded_frame_frac", "pct"),
    ]
    append_table(wb, "SmoothFoot健康度", workbook_health_rows(metric_rows), health_cols, "7030A0")

    artifact_cols = [
        ("variant", "variant", "text"),
        ("case", "short_case_id", "text"),
        ("arm", "arm", "text"),
        ("kind", "arm_kind", "text"),
        ("artifact_ok", "artifact_ok", "text"),
        ("qpos", "qpos_exists", "text"),
        ("video", "video_exists", "text"),
        ("config", "config_exists", "text"),
        ("smooth report", "smooth_report_exists", "text"),
        ("smooth enabled", "config_smooth_enabled_ok", "text"),
        ("smooth accel weight", "config_smooth_accel_weight_ok", "text"),
        ("smooth jerk weight", "config_smooth_jerk_weight_ok", "text"),
        ("ankle weight", "config_ankle_weight_ok", "text"),
        ("foot slip enabled", "config_foot_slip_enabled_ok", "text"),
        ("foot ground enabled", "config_foot_ground_enabled_ok", "text"),
        ("peak margin disabled", "config_peak_margin_disabled_ok", "text"),
        ("qpos_path", "qpos_path", "text"),
        ("config_act", "config_act", "text"),
    ]
    append_table(wb, "E166产物检查", artifact_rows, artifact_cols, "8064A2")

    info_rows = [
        {"项": "范围", "说明": "E166 CEM 侧三 case：box021_035_p2、box004_082_p1、box004_083_p2；排除 box023 pathology，不补 box026 label。"},
        {"项": "baseline", "说明": "复用 E163 narrowSurfaceBand clean8 对应三 case，作为 E166 CEM arms 的对比基线。"},
        {"项": "B1", "说明": "CEM sample-level smooth accel/jerk penalty，默认不做 hard gate。"},
        {"项": "B2", "说明": "CPU handoff 后处理平滑；不重跑 CEM，因此无 config_act。"},
        {"项": "A", "说明": "local-frame ankle weight=2.0，并启用 foot slip / foot ground penalty。"},
        {"项": "A_B2_postSmooth", "说明": "先用 A 的 CEM 结果，再执行 B2 同款 CPU 后处理平滑；不是 A+B1。"},
        {"项": "AplusB", "说明": "A 与 B1 叠加。"},
        {"项": "SPIDER gate", "说明": "tracked、no fall、raw/clean3 contact delta >= -0.05、3mm penetration delta <= +0.05。"},
        {"项": "判读", "说明": "A 是 CEM-only 最强候选；A_B2_postSmooth 用来检验 A 后处理平滑；AplusB 是 A+B1。"},
    ]
    append_table(wb, "说明", info_rows, [("项", "项", "text"), ("说明", "说明", "text")], "7F6000")

    raw_cols = [(key, key, "text") for key in sorted({key for row in metric_rows for key in row})]
    append_table(wb, "原始metrics", metric_rows, raw_cols, "666666", hidden=True)

    out_path = eval_dir / "E166_foot_smooth_vs_E163_three_case_eval.xlsx"
    wb.save(out_path)
    return out_path


def evaluate(stage: str, allow_missing: bool, variants_path: Path = VARIANTS, eval_dir: Path | None = None) -> dict[str, Any]:
    global CASE_ORDER
    rows = read_tsv(variants_path)
    dynamic_cases: list[str] = []
    for row in rows:
        case = row["short_case_id"]
        if case not in dynamic_cases:
            dynamic_cases.append(case)
    CASE_ORDER = [case for case in CASE_ORDER if case in dynamic_cases] + [
        case for case in dynamic_cases if case not in CASE_ORDER
    ]
    if eval_dir is None:
        eval_dir = EVAL_ROOT / stage
    cfg = EvalConfig()
    metric_rows: list[dict[str, Any]] = []
    artifact_rows: list[dict[str, Any]] = []
    missing_rows: list[dict[str, Any]] = []

    for row in rows:
        artifact = artifact_check(row)
        artifact_rows.append(artifact)
        item = evaluate_one(row, cfg)
        if item is None:
            missing_rows.append(
                {
                    "short_case_id": row["short_case_id"],
                    "arm": row["arm"],
                    "variant": row["variant"],
                    "reason": "missing_qpos_or_scene",
                    "qpos_path": rel(qpos_path_for_row(row)),
                }
            )
            continue
        item.update(artifact)
        metric_rows.append(item)

    add_baseline_deltas(metric_rows)
    summary_rows = summarize(metric_rows)

    metric_fields = [
        "short_case_id",
        "arm",
        "variant",
        "method",
        "status",
        "spider_gate_pass",
        "success_tracked",
        "fall_flag",
        "contact_no_regression",
        "penetration_no_regression",
        "artifact_ok",
        *METRIC_KEYS,
        *[f"{key}_delta_vs_baseline" for key in METRIC_KEYS if key not in {"success_tracked", "fall_flag"}],
        *HEALTH_AGGS.keys(),
        "has_smooth_health",
        "has_foot_health",
        "qpos_path",
        "config_act",
        "video",
    ]
    write_tsv(eval_dir / "e166_arm_metrics.tsv", metric_rows, metric_fields)
    write_tsv(eval_dir / "e166_arm_summary.tsv", summary_rows)
    write_tsv(eval_dir / "e166_artifact_check.tsv", artifact_rows)
    write_tsv(eval_dir / "e166_missing.tsv", missing_rows)
    write_summary_md(eval_dir / "summary.md", summary_rows, metric_rows, missing_rows)
    xlsx = write_workbook(eval_dir, summary_rows, metric_rows, artifact_rows)
    payload = {
        "stage": stage,
        "rows_expected": len(rows),
        "rows_evaluated": len(metric_rows),
        "missing_rows": len(missing_rows),
        "allow_missing": allow_missing,
        "all_artifacts_ok_for_evaluated": all(bool(row.get("artifact_ok")) for row in metric_rows),
        "summary": summary_rows,
        "xlsx": rel(xlsx),
    }
    write_json(eval_dir / "summary.json", payload)
    if missing_rows and not allow_missing:
        raise SystemExit(f"E166 eval missing {len(missing_rows)} rows; see {rel(eval_dir / 'e166_missing.tsv')}")
    print(
        "E166 eval: "
        f"stage={stage} expected={len(rows)} evaluated={len(metric_rows)} "
        f"missing={len(missing_rows)} out={rel(eval_dir)} xlsx={rel(xlsx)}"
    )
    return payload


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("stage", nargs="?", default="full")
    parser.add_argument("--allow-missing", action="store_true")
    parser.add_argument("--variants", type=Path, default=VARIANTS)
    parser.add_argument("--eval-dir", type=Path, default=None)
    args = parser.parse_args()
    evaluate(args.stage, args.allow_missing, args.variants, args.eval_dir)


if __name__ == "__main__":
    main()
