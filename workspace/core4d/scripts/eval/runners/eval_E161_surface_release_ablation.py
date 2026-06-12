#!/usr/bin/env python3
"""Evaluate E161 surfaceBand release ablation clean8 benchmark."""

from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
import sys
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from eval.core.core_metrics import (  # noqa: E402
    EVAL_METRIC_STANDARD_ID,
    EvalConfig,
    STANDARD_DELTA_METRICS,
    STANDARD_MASK_DELTA_METRICS,
    STANDARD_TRACK_DIAG,
)
from eval_E156_clean8_gate_decay import (  # noqa: E402
    evaluate_omniretarget_row,
    evaluate_row,
    finite,
    fmt,
    read_tsv,
    repo_path,
    write_json,
    write_tsv,
)


REPO = Path(__file__).resolve().parents[5]
E156_VARIANTS = REPO / "workspace/core4d/scripts/experiments/E156/variants.tsv"
E159_VARIANTS = REPO / "workspace/core4d/scripts/experiments/E159/variants.tsv"
E161_VARIANTS = REPO / "workspace/core4d/scripts/experiments/E161/variants.tsv"
RESULT_ROOT = REPO / "workspace/core4d/results/E161/surface_release_ablation"

CLEAN8_CASES = [
    "box021_035_p1",
    "box021_035_p2",
    "box021_029_p2",
    "box004_083_p1",
    "box004_083_p2",
    "box023_person2",
    "box004_082_p1",
    "box026_139_p1",
]
E161_METHODS = [
    "gateA+surfaceBand-A2+postureRerankA",
    "gateA+surfaceBand-A2+postureRerankA+surfaceBandReleaseDecay",
    "gateA+surfaceBand-A2+postureRerankA+surfaceBandStrictMask",
]
METHOD_ORDER = [
    "OmniRetarget",
    "spider-rubberhand",
    "+gateA",
    "E155_decay",
    "gateA+surfaceBand-A2",
] + E161_METHODS

TRACK_DIAG = list(STANDARD_TRACK_DIAG)
DELTA_METRICS = list(STANDARD_DELTA_METRICS)
MASK_DELTA = list(STANDARD_MASK_DELTA_METRICS)

GATE_HEALTH_KEYS = {
    "cem_hand_gate_valid_frac": ("mean", "hand_gate_valid_frac"),
    "cem_gate_fallback_used": ("mean", "gate_fallback_used"),
    "cem_hand_gate_min_sdf_min": ("min", "hand_gate_min_sdf_min_m"),
    "sample_hand_gate_violation_pct_mean": ("mean", "hand_gate_violation_pct"),
}
SURFACE_HEALTH_KEYS = {
    "surface_band_rew_mean": ("mean", "surface_band_rew_mean"),
    "surface_band_penalty_mean": ("mean", "surface_band_penalty_mean"),
    "surface_band_gate_mean": ("mean", "surface_band_gate_mean"),
    "surface_band_decay_factor_mean": ("mean", "surface_band_decay_factor_mean"),
    "surface_band_sdf_mean": ("mean", "surface_band_sdf_mean_m"),
    "surface_band_score_mean": ("mean", "surface_band_score_mean"),
    "surface_band_penetration_mean": ("mean", "surface_band_penetration_mean_m"),
}
POSTURE_HEALTH_KEYS = {
    "cem_posture_gate_valid_frac": ("mean", "posture_gate_valid_frac"),
    "cem_posture_gate_selected_valid_frac": ("mean", "posture_gate_selected_valid_frac"),
    "cem_posture_gate_fallback_used": ("mean", "posture_gate_fallback_used"),
    "sample_posture_mean_z_err_mean": ("mean", "posture_mean_z_err_m"),
    "sample_posture_terminal_z_err_mean": ("mean", "posture_terminal_z_err_m"),
    "sample_posture_max_z_drop_mean": ("mean", "posture_max_z_drop_m"),
    "sample_posture_violation_mean": ("mean", "posture_violation"),
}
RELEASE_HEALTH_KEYS = [
    "surface_band_release_gate_mean",
    "surface_band_release_rew_mean",
    "surface_band_release_active_frac",
    "surface_band_release_decay_factor_mean",
]
HEALTH_DESTS = (
    [dst for _, dst in GATE_HEALTH_KEYS.values()]
    + [dst for _, dst in SURFACE_HEALTH_KEYS.values()]
    + [dst for _, dst in POSTURE_HEALTH_KEYS.values()]
    + RELEASE_HEALTH_KEYS
)

SUMMARY_KEYS = [
    "success_tracked",
    "track_pelvis_z_err_terminal_m",
    "hand_object_release_false_contact_3mm_frac",
    "hand_object_release_false_contact_5mm_frac",
    "hand_object_physics_contact_3mm_in_mask_frac",
    "hand_object_physics_contact_5mm_in_mask_frac",
    "hand_object_physics_penetration_3mm_frame_frac",
    "hand_object_physics_penetration_5mm_frame_frac",
    "hand_geom_penetration_2mm_frac",
    "hand_geom_penetration_5mm_frac",
    "leg_penetration_frac",
    "obj_err_mean_m",
] + HEALTH_DESTS

EXTRA_DELTA_METRICS = [
    "hand_object_release_false_contact_3mm_frac",
    "hand_object_release_false_contact_5mm_frac",
]

XLSX_METRICS = [
    ("tracked", "success_tracked_cases", +1),
    ("relF3", "hand_object_release_false_contact_3mm_frac_mean", -1),
    ("inmaskC3", "hand_object_physics_contact_3mm_in_mask_frac_mean", +1),
    ("physPen3", "hand_object_physics_penetration_3mm_frame_frac_mean", -1),
    ("pen2", "hand_geom_penetration_2mm_frac_mean", -1),
    ("releaseGate", "surface_band_release_gate_mean_mean", -1),
    ("releaseRew", "surface_band_release_rew_mean_mean", -1),
    ("pzTerm", "track_pelvis_z_err_terminal_m_mean", -1),
]


def mean(values: list[Any]) -> float:
    vals = [finite(value) for value in values]
    vals = [value for value in vals if math.isfinite(value)]
    return statistics.fmean(vals) if vals else math.nan


def worst(values: list[Any], high_is_bad: bool = True) -> float:
    vals = [finite(value) for value in values]
    vals = [value for value in vals if math.isfinite(value)]
    if not vals:
        return math.nan
    return max(vals) if high_is_bad else min(vals)


def stage_qpos(row: dict[str, str], stage: str) -> Path:
    if stage == "full" or row.get("run_status") == "reuse_e160":
        return repo_path(row["outdir_npz"])
    return RESULT_ROOT / "cem" / stage / f"{row['variant']}_outdir_{stage}" / "trajectory_mjwp_act.npz"


def stage_video(row: dict[str, str], stage: str) -> Path:
    if stage == "full" or row.get("run_status") == "reuse_e160":
        return repo_path(row["video"])
    return RESULT_ROOT / "cem" / stage / f"{row['variant']}_{stage}.mp4"


def resize_contact_mask(mask: np.ndarray, target_len: int) -> np.ndarray:
    if mask.shape[0] == target_len:
        return mask.astype(bool)
    if mask.shape[0] <= 0:
        return np.zeros((target_len, mask.shape[1]), dtype=bool)
    idx = np.round(np.linspace(0, mask.shape[0] - 1, target_len)).astype(np.int64)
    return mask[idx].astype(bool)


def aligned_contact_mask(row: dict[str, str], frames: int) -> np.ndarray | None:
    mask_path = repo_path(row.get("mask_path", ""))
    if not mask_path.is_file():
        return None
    data = np.load(mask_path, allow_pickle=True)
    person_idx = int(row.get("person_idx") or 0)
    key = ""
    if "spider_contact_mask_3cm" in data.files and data["spider_contact_mask_3cm"].shape[0] == frames:
        key = "spider_contact_mask_3cm"
    elif "eval_contact_mask_3cm" in data.files and data["eval_contact_mask_3cm"].shape[0] == frames:
        key = "eval_contact_mask_3cm"
    elif "eval_contact_mask_3cm" in data.files:
        key = "eval_contact_mask_3cm"
    elif "spider_contact_mask_3cm" in data.files:
        key = "spider_contact_mask_3cm"
    else:
        return None

    sm = np.asarray(data[key])
    if sm.ndim != 3 or person_idx >= sm.shape[1] or sm.shape[2] < 2:
        return None
    return resize_contact_mask(sm[:, person_idx, :2], frames)


def release_mask(row: dict[str, str], frames: int) -> np.ndarray:
    out = np.zeros(frames, dtype=bool)
    mask = aligned_contact_mask(row, frames)
    if mask is None:
        return out
    m = mask[:, 0] | mask[:, 1]
    if not m.any():
        return out
    last_c = int(len(m) - 1 - np.argmax(m[::-1]))
    out[last_c + 1 : len(m)] = True
    return out


def health_from_npz(qpos_path: Path, row: dict[str, str]) -> dict[str, Any]:
    out: dict[str, Any] = {key: "" for key in HEALTH_DESTS}
    if not qpos_path.is_file():
        return out
    data = np.load(qpos_path, allow_pickle=True)
    for src, (agg, dst) in {**GATE_HEALTH_KEYS, **SURFACE_HEALTH_KEYS, **POSTURE_HEALTH_KEYS}.items():
        if src not in data.files:
            continue
        arr = np.asarray(data[src], dtype=np.float64)
        arr = arr[np.isfinite(arr)]
        if arr.size:
            out[dst] = float(arr.min()) if agg == "min" else float(arr.mean())

    if "surface_band_rew_mean" in data.files:
        frames = np.asarray(data["surface_band_rew_mean"]).shape[0]
        release = release_mask(row, frames)
        if release.any():
            for src, dst in [
                ("surface_band_gate_mean", "surface_band_release_gate_mean"),
                ("surface_band_rew_mean", "surface_band_release_rew_mean"),
                ("surface_band_decay_factor_mean", "surface_band_release_decay_factor_mean"),
            ]:
                if src in data.files:
                    arr = np.asarray(data[src], dtype=np.float64)
                    out[dst] = float(np.nanmean(arr[release]))
            if "surface_band_gate_mean" in data.files:
                arr = np.asarray(data["surface_band_gate_mean"], dtype=np.float64)
                out["surface_band_release_active_frac"] = float(np.nanmean(arr[release] > 0.5))
            if row.get("method_group") == "strictMask":
                mask = aligned_contact_mask(row, frames)
                if mask is not None:
                    current_gate = (mask[:, 0] | mask[:, 1]).astype(np.float64)
                    out["surface_band_release_gate_mean"] = float(np.nanmean(current_gate[release]))
                    out["surface_band_release_active_frac"] = float(np.nanmean(current_gate[release] > 0.5))
    return out


def evaluate_with_health(row: dict[str, str], qpos_path: Path, cfg: EvalConfig, stage: str) -> dict[str, Any] | None:
    metrics = evaluate_row(row, qpos_path, cfg)
    if metrics is None:
        return None
    metrics.update(health_from_npz(qpos_path, row))
    metrics["video"] = str(stage_video(row, stage))
    return metrics


def summarize_method(method: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
    out: dict[str, Any] = {
        "method": method,
        "n_cases": len(rows),
        "success_tracked_cases": sum(1 for row in rows if row.get("success_tracked")),
        "fall_cases": sum(1 for row in rows if row.get("fall_flag")),
    }
    high_good = {
        "success_tracked",
        "hand_object_physics_contact_3mm_in_mask_frac",
        "hand_object_physics_contact_5mm_in_mask_frac",
        "hand_gate_valid_frac",
        "surface_band_rew_mean",
        "surface_band_score_mean",
        "posture_gate_valid_frac",
        "posture_gate_selected_valid_frac",
    }
    for key in SUMMARY_KEYS:
        if key == "success_tracked":
            vals = [1.0 if row.get(key) else 0.0 for row in rows]
            out[f"{key}_mean"] = mean(vals)
            out[f"{key}_worst"] = worst(vals, high_is_bad=False)
        else:
            out[f"{key}_mean"] = mean([row.get(key) for row in rows])
            out[f"{key}_worst"] = worst([row.get(key) for row in rows], high_is_bad=key not in high_good)
    return out


def delta_row(case: str, run: dict[str, Any], ref: dict[str, Any], ref_method: str) -> dict[str, Any]:
    row: dict[str, Any] = {
        "short_case_id": case,
        "variant": run["variant"],
        "method": run["method"],
        "reference_method": ref_method,
        "reference_variant": ref["variant"],
        "success_tracked": run.get("success_tracked", False),
        "fall_flag": run.get("fall_flag", False),
    }
    for key in TRACK_DIAG:
        row[key] = run.get(key, math.nan)
    for key in MASK_DELTA + DELTA_METRICS + EXTRA_DELTA_METRICS + RELEASE_HEALTH_KEYS:
        if f"{key}_delta" in row:
            continue
        row[f"{key}_ref"] = ref.get(key, math.nan)
        row[f"{key}_run"] = run.get(key, math.nan)
        row[f"{key}_delta"] = finite(run.get(key)) - finite(ref.get(key))
    return row


def build_delta_rows(metric_rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    by_case_method = {(row["short_case_id"], row["method"]): row for row in metric_rows}
    vs_omni: list[dict[str, Any]] = []
    vs_baseline: list[dict[str, Any]] = []
    vs_gate: list[dict[str, Any]] = []
    vs_m0: list[dict[str, Any]] = []
    for case in CLEAN8_CASES:
        omni = by_case_method.get((case, "OmniRetarget"))
        base = by_case_method.get((case, "spider-rubberhand"))
        gate = by_case_method.get((case, "+gateA"))
        m0 = by_case_method.get((case, E161_METHODS[0]))
        for method in METHOD_ORDER:
            run = by_case_method.get((case, method))
            if not run:
                continue
            if omni and method != "OmniRetarget":
                vs_omni.append(delta_row(case, run, omni, "OmniRetarget"))
            if base and method not in {"OmniRetarget", "spider-rubberhand"}:
                vs_baseline.append(delta_row(case, run, base, "spider-rubberhand"))
            if gate and method in {"E155_decay", "gateA+surfaceBand-A2", *E161_METHODS}:
                vs_gate.append(delta_row(case, run, gate, "+gateA"))
            if m0 and method in E161_METHODS[1:]:
                vs_m0.append(delta_row(case, run, m0, E161_METHODS[0]))
    return vs_omni, vs_baseline, vs_gate, vs_m0


def claims_summary(metric_rows: list[dict[str, Any]], delta_vs_m0: list[dict[str, Any]]) -> dict[str, Any]:
    by_method = {method: [row for row in metric_rows if row["method"] == method] for method in E161_METHODS}
    m0 = by_method[E161_METHODS[0]]
    out: dict[str, Any] = {"checks": {}, "per_method": {}}
    m0_success = sum(1 for row in m0 if row.get("success_tracked"))
    m0_fall = sum(1 for row in m0 if row.get("fall_flag"))
    out["checks"]["C1_M0_missing0_full8"] = len(m0) == 8
    out["checks"]["C1_M0_success_ge_7"] = m0_success >= 7
    out["checks"]["C1_M0_fall_le_1"] = m0_fall <= 1

    for method in E161_METHODS[1:]:
        rows = by_method[method]
        deltas = [row for row in delta_vs_m0 if row["method"] == method]
        box021 = next((row for row in rows if row["short_case_id"] == "box021_029_p2"), {})
        rel_delta = mean([row.get("hand_object_release_false_contact_3mm_frac_delta") for row in deltas])
        inmask_delta = mean([row.get("hand_object_physics_contact_3mm_in_mask_frac_delta") for row in deltas])
        phys_pen3_delta = mean([row.get("hand_object_physics_penetration_3mm_frame_frac_delta") for row in deltas])
        pen2_delta = mean([row.get("hand_geom_penetration_2mm_frac_delta") for row in deltas])
        pz_delta = mean([row.get("track_pelvis_z_err_terminal_m_delta") for row in deltas])
        success = sum(1 for row in rows if row.get("success_tracked"))
        fall = sum(1 for row in rows if row.get("fall_flag"))
        strict_gate_ok = True
        if method.endswith("surfaceBandStrictMask"):
            strict_gate_ok = (
                mean([row.get("surface_band_release_gate_mean") for row in rows]) <= 0.05
                and mean([row.get("surface_band_release_rew_mean") for row in rows]) <= 0.05
            )
        checks = {
            "full8": len(rows) == 8,
            "release_mean_delta_le_neg_0_10": rel_delta <= -0.10,
            "box021_releaseF3_le_0_25": finite(box021.get("hand_object_release_false_contact_3mm_frac")) <= 0.25,
            "inmaskC3_drop_le_0_08": inmask_delta >= -0.08,
            "physPen3_delta_le_0_03": phys_pen3_delta <= 0.03,
            "pen2_delta_le_0_03": pen2_delta <= 0.03,
            "success_not_worse_by_more_than_1": success >= m0_success - 1,
            "fall_not_worse_by_more_than_1": fall <= m0_fall + 1,
            "pz_delta_le_0_03": pz_delta <= 0.03,
            "strict_release_gate_ok": strict_gate_ok,
        }
        out["per_method"][method] = {
            "checks": checks,
            "pass": all(checks.values()),
            "releaseF3_delta_vs_M0": rel_delta,
            "inmaskC3_delta_vs_M0": inmask_delta,
            "physPen3_delta_vs_M0": phys_pen3_delta,
            "pen2_delta_vs_M0": pen2_delta,
            "pz_delta_vs_M0": pz_delta,
            "success_tracked_cases": success,
            "fall_cases": fall,
            "box021_releaseF3": finite(box021.get("hand_object_release_false_contact_3mm_frac")),
        }
    out["promote_candidates"] = [method for method, item in out["per_method"].items() if item["pass"]]
    return out


def rank_styles(values: list[float], direction: int) -> tuple[int | None, int | None]:
    indexed = [(idx, val) for idx, val in enumerate(values) if math.isfinite(val)]
    if not indexed:
        return None, None
    indexed.sort(key=lambda item: item[1], reverse=direction > 0)
    best = indexed[0][0]
    second = indexed[1][0] if len(indexed) > 1 else None
    return best, second


def write_xlsx(
    path: Path,
    method_summary: list[dict[str, Any]],
    metric_rows: list[dict[str, Any]],
    delta_vs_omni: list[dict[str, Any]],
    delta_vs_baseline: list[dict[str, Any]],
    delta_vs_gate: list[dict[str, Any]],
    delta_vs_m0: list[dict[str, Any]],
    claims: dict[str, Any],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "method_summary"
    headers = ["method", "n_cases"] + [label for label, _, _ in XLSX_METRICS]
    rows = sorted(method_summary, key=lambda row: METHOD_ORDER.index(row["method"]) if row["method"] in METHOD_ORDER else 99)
    ws.append(headers)
    for row in rows:
        ws.append([row.get("method"), row.get("n_cases")] + [row.get(key) for _, key, _ in XLSX_METRICS])
    for col in range(1, len(headers) + 1):
        ws.cell(1, col).font = Font(bold=True)
        ws.cell(1, col).fill = PatternFill("solid", fgColor="D9EAF7")
    for metric_idx, (_, key, direction) in enumerate(XLSX_METRICS, start=3):
        vals = [finite(row.get(key)) for row in rows]
        best, second = rank_styles(vals, direction)
        if best is not None:
            ws.cell(best + 2, metric_idx).font = Font(bold=True, color="000000")
        if second is not None:
            ws.cell(second + 2, metric_idx).font = Font(underline="single")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"

    for title, rows2 in (
        ("per_case_metrics", metric_rows),
        ("delta_vs_omniretarget", delta_vs_omni),
        ("delta_vs_baseline", delta_vs_baseline),
        ("delta_vs_gateA", delta_vs_gate),
        ("delta_vs_M0", delta_vs_m0),
    ):
        sheet = wb.create_sheet(title)
        fields = sorted({key for row in rows2 for key in row.keys()})
        front = ["short_case_id", "variant", "method", "reference_method", "success_tracked"]
        fields = [field for field in front if field in fields] + [field for field in fields if field not in front]
        sheet.append(fields)
        for row in rows2:
            sheet.append([row.get(field, "") for field in fields])
        for col in range(1, len(fields) + 1):
            sheet.cell(1, col).font = Font(bold=True)
            sheet.cell(1, col).fill = PatternFill("solid", fgColor="D9EAD3")
        sheet.freeze_panes = "A2"

    meta = wb.create_sheet("summary")
    meta.append(["metric_standard_id", EVAL_METRIC_STANDARD_ID])
    for key, val in claims.get("checks", {}).items():
        meta.append([key, val])
    for method, item in claims.get("per_method", {}).items():
        meta.append([method, item.get("pass")])
        for key, val in item.items():
            if key != "checks":
                meta.append([f"{method}.{key}", val])
        for key, val in item.get("checks", {}).items():
            meta.append([f"{method}.checks.{key}", val])
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("stage", nargs="?", default="full", choices=["smoke", "full"])
    parser.add_argument("--allow-missing", action="store_true")
    args = parser.parse_args()

    cfg = EvalConfig()
    eval_dir = RESULT_ROOT / "eval" / args.stage
    e156_rows = [row for row in read_tsv(E156_VARIANTS) if row["short_case_id"] in CLEAN8_CASES]
    e159_rows = [row for row in read_tsv(E159_VARIANTS) if row["short_case_id"] in CLEAN8_CASES]
    e161_rows = read_tsv(E161_VARIANTS)
    rows = e156_rows + e159_rows + e161_rows
    metric_rows: list[dict[str, Any]] = []
    missing: list[str] = []

    baseline_rows = [row for row in e156_rows if row["method"] == "spider-rubberhand"]
    for row in baseline_rows:
        metrics = evaluate_omniretarget_row(row, eval_dir, cfg)
        if metrics is None:
            missing.append(f"E161_{row['short_case_id']}_omniretarget")
            continue
        metrics.update({key: "" for key in HEALTH_DESTS})
        metric_rows.append(metrics)

    for row in rows:
        qpos = stage_qpos(row, args.stage) if row.get("source_exp") == "E161" else repo_path(row["outdir_npz"])
        metrics = evaluate_with_health(row, qpos, cfg, args.stage)
        if metrics is None:
            missing.append(row["variant"])
            continue
        metric_rows.append(metrics)

    method_summary = [summarize_method(method, [row for row in metric_rows if row["method"] == method]) for method in METHOD_ORDER]
    delta_vs_omni, delta_vs_baseline, delta_vs_gate, delta_vs_m0 = build_delta_rows(metric_rows)
    claims = claims_summary(metric_rows, delta_vs_m0)

    metric_fields = (
        [
            "short_case_id",
            "variant",
            "method",
            "method_group",
            "run_status",
            "source_exp",
            "split",
            "qpos_frames",
            "success_tracked",
            "pelvis_min_m",
            "fall_flag",
            "hand_geom_near_5cm_frac",
            "hand_geom_near_10cm_frac",
            "hand_geom_penetration_2mm_frac",
            "hand_geom_penetration_5mm_frac",
            "hand_object_physics_contact_3mm_frac",
            "hand_object_physics_contact_5mm_frac",
            "hand_object_physics_contact_3mm_in_mask_frac",
            "hand_object_physics_contact_5mm_in_mask_frac",
            "hand_object_physics_penetration_3mm_frame_frac",
            "hand_object_physics_penetration_5mm_frame_frac",
            "hand_object_release_false_contact_3mm_frac",
            "hand_object_release_false_contact_5mm_frac",
            "leg_penetration_frac",
            "obj_err_mean_m",
        ]
        + TRACK_DIAG
        + HEALTH_DESTS
        + ["result_npz", "video"]
    )
    delta_fields = (
        [
            "short_case_id",
            "variant",
            "method",
            "reference_method",
            "reference_variant",
            "success_tracked",
            "fall_flag",
        ]
        + TRACK_DIAG
        + [f"{key}_{suffix}" for key in MASK_DELTA for suffix in ("ref", "run", "delta")]
        + [f"{key}_{suffix}" for key in DELTA_METRICS for suffix in ("ref", "run", "delta")]
        + [f"{key}_{suffix}" for key in EXTRA_DELTA_METRICS for suffix in ("ref", "run", "delta")]
        + [f"{key}_{suffix}" for key in RELEASE_HEALTH_KEYS for suffix in ("ref", "run", "delta")]
    )
    summary_fields = sorted({key for row in method_summary for key in row.keys()})
    summary_fields = ["method", "n_cases", "success_tracked_cases", "fall_cases"] + [
        field for field in summary_fields if field not in {"method", "n_cases", "success_tracked_cases", "fall_cases"}
    ]

    write_tsv(eval_dir / "e161_method_metrics.tsv", metric_rows, metric_fields)
    write_tsv(eval_dir / "e161_method_summary.tsv", method_summary, summary_fields)
    write_tsv(eval_dir / "e161_delta_vs_omniretarget.tsv", delta_vs_omni, delta_fields)
    write_tsv(eval_dir / "e161_delta_vs_spider_rubberhand.tsv", delta_vs_baseline, delta_fields)
    write_tsv(eval_dir / "e161_delta_vs_gateA.tsv", delta_vs_gate, delta_fields)
    write_tsv(eval_dir / "e161_delta_vs_M0.tsv", delta_vs_m0, delta_fields)
    write_json(
        eval_dir / "e161_eval_summary.json",
        {
            "metric_standard_id": EVAL_METRIC_STANDARD_ID,
            "stage": args.stage,
            "metric_rows": len(metric_rows),
            "method_rows": len(method_summary),
            "missing": missing,
            "allow_missing": bool(args.allow_missing),
            "claims": claims,
        },
    )
    write_xlsx(
        eval_dir / "E161_surface_release_ablation_clean8.xlsx",
        method_summary,
        metric_rows,
        delta_vs_omni,
        delta_vs_baseline,
        delta_vs_gate,
        delta_vs_m0,
        claims,
    )
    print(
        f"E161 eval: metric_rows={len(metric_rows)} methods={len(method_summary)} "
        f"delta_vs_M0={len(delta_vs_m0)} missing={len(missing)}"
    )
    if missing:
        print("MISSING:", missing)
        if not args.allow_missing:
            raise SystemExit(1)


if __name__ == "__main__":
    main()
