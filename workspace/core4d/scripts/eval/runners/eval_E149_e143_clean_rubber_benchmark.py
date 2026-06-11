#!/usr/bin/env python3
"""Build an E143 clean-benchmark view of the E148 rubber hand comparison."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

REPO = Path(__file__).resolve().parents[5]
E148_CASE_TSV = (
    REPO
    / "workspace/core4d/results/E148/e143_24case_rubber_hand_collision/eval/full/e148_case_comparison.tsv"
)
E143_ANALYSIS_TSV = REPO / "workspace/core4d/results/E143/spider_contact_failure_analysis/data/case_analysis.tsv"
E143_ANALYSIS_JSON = REPO / "workspace/core4d/results/E143/spider_contact_failure_analysis/data/analysis_summary.json"
OUT_DIR = REPO / "workspace/core4d/results/E149/e143_clean_rubber_benchmark"

CLEAN6 = [
    "box021_035_p1",
    "box021_035_p2",
    "box021_029_p2",
    "box004_083_p1",
    "box004_083_p2",
    "box023_person2",
]

RELAXED8 = [
    "box021_035_p2",
    "box023_person2",
    "box021_029_p2",
    "box004_082_p1",
    "box004_083_p2",
    "box021_035_p1",
    "box026_139_p1",
    "box004_083_p1",
]

BENCHMARKS = {
    "clean6_primary": CLEAN6,
    "relaxed8_valid_like": RELAXED8,
}

METHODS = [
    ("OmniRetarget", "OmniRetarget"),
    ("sphere Spider", "sphere_spider"),
    ("rubber hand Spider", "rubber_hand_spider"),
]

METRICS = [
    "手物接触",
    "5cm",
    "10cm",
    "手物穿透",
    "深穿透2cm",
    "腿穿透",
    "body穿透",
    "物体触地",
    "pelvis_min_m",
]

DIFF_METRICS = ["手物接触", "5cm", "10cm", "手物穿透", "腿穿透"]

ANNOTATION_FIELDS = [
    "is_box_hold",
    "note",
    "rl_success",
    "color_label",
    "category",
    "category_reason",
]


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def write_tsv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: fmt(row.get(field, "")) for field in fields})


def fmt(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.8g}"
    return str(value)


def as_float(row: dict[str, str], key: str) -> float:
    text = row.get(key, "")
    if text == "":
        raise ValueError(f"missing numeric field {key} in {row.get('case_id')}")
    return float(text)


def as_bool(value: str) -> bool:
    return str(value or "").strip().lower() in {"1", "true", "yes", "pass"}


def short_case_id(row: dict[str, str]) -> str:
    variant = row["e143_variant"]
    if not variant.startswith("E143_") or not variant.endswith("_raw_mask_ref_fk"):
        raise ValueError(f"unexpected E143 variant: {variant}")
    return variant.removeprefix("E143_").removesuffix("_raw_mask_ref_fk")


def mean(values: list[float]) -> float:
    return sum(values) / len(values) if values else float("nan")


def load_inputs() -> tuple[dict[str, dict[str, str]], dict[str, dict[str, str]], dict[str, Any]]:
    e148_rows = read_tsv(E148_CASE_TSV)
    by_short: dict[str, dict[str, str]] = {}
    for row in e148_rows:
        short = short_case_id(row)
        row["short_case_id"] = short
        by_short[short] = row

    annotation_rows = read_tsv(E143_ANALYSIS_TSV)
    annotations = {row["case_id"]: row for row in annotation_rows}
    analysis = json.loads(E143_ANALYSIS_JSON.read_text(encoding="utf-8"))
    return by_short, annotations, analysis


def method_summary_rows(by_short: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for bench, cases in BENCHMARKS.items():
        rows = [by_short[c] for c in cases]
        for method_name, method_key in METHODS:
            item: dict[str, Any] = {
                "benchmark": bench,
                "method": method_name,
                "case_count": len(rows),
                "fall_count": sum(int(as_bool(row.get(f"fall_{method_key}", ""))) for row in rows),
            }
            for metric in METRICS:
                item[metric] = mean([as_float(row, f"{metric}_{method_key}") for row in rows])
            out.append(item)
    return out


def diff_summary_rows(summary_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_bench_method = {(row["benchmark"], row["method"]): row for row in summary_rows}
    out: list[dict[str, Any]] = []
    for bench in BENCHMARKS:
        omni = by_bench_method[(bench, "OmniRetarget")]
        sphere = by_bench_method[(bench, "sphere Spider")]
        rubber = by_bench_method[(bench, "rubber hand Spider")]
        for metric in DIFF_METRICS:
            out.append(
                {
                    "benchmark": bench,
                    "metric": metric,
                    "rubber-sphere": rubber[metric] - sphere[metric],
                    "OmniRetarget-rubber": omni[metric] - rubber[metric],
                    "rubber-OmniRetarget": rubber[metric] - omni[metric],
                }
            )
    return out


def case_rows(by_short: dict[str, dict[str, str]], annotations: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for bench, cases in BENCHMARKS.items():
        for short in cases:
            src = by_short[short]
            ann = annotations.get(short, {})
            row: dict[str, Any] = {
                "benchmark": bench,
                "short_case_id": short,
                "e148_case_id": src["case_id"],
                "object_key": src["object_key"],
                "run_status": src["run_status"],
                "rubber_source_exp": src["rubber_source_exp"],
                "e143_variant": src["e143_variant"],
                "e147_variant": src["e147_variant"],
            }
            for field in ANNOTATION_FIELDS:
                row[field] = ann.get(field, "")
            for metric in DIFF_METRICS:
                for method_name, method_key in METHODS:
                    row[f"{metric}_{method_name}"] = as_float(src, f"{metric}_{method_key}")
                row[f"{metric}_rubber-sphere"] = as_float(src, f"{metric}_rubber_hand_spider") - as_float(
                    src, f"{metric}_sphere_spider"
                )
                row[f"{metric}_OmniRetarget-rubber"] = as_float(src, f"{metric}_OmniRetarget") - as_float(
                    src, f"{metric}_rubber_hand_spider"
                )
            rows.append(row)
    return rows


def annotation_source_rows(annotations: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    out = []
    wanted = set(RELAXED8)
    for short in RELAXED8:
        ann = annotations.get(short, {})
        out.append(
            {
                "short_case_id": short,
                "in_clean6": "yes" if short in CLEAN6 else "no",
                "in_relaxed8": "yes" if short in wanted else "no",
                **{field: ann.get(field, "") for field in ANNOTATION_FIELDS},
            }
        )
    return out


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def write_markdown(
    path: Path,
    summary_rows: list[dict[str, Any]],
    diff_rows: list[dict[str, Any]],
    analysis: dict[str, Any],
) -> None:
    by_bench = {}
    for row in summary_rows:
        by_bench.setdefault(row["benchmark"], []).append(row)
    diffs_by_bench = {}
    for row in diff_rows:
        diffs_by_bench.setdefault(row["benchmark"], []).append(row)

    lines = [
        "# E149 E143 clean rubber benchmark",
        "",
        "This is an eval-only clean benchmark view of E148. No CEM/RL/remote job was launched.",
        "",
        f"- clean6 cases: `{len(CLEAN6)}`",
        f"- relaxed8 valid-like cases: `{len(RELAXED8)}`",
        f"- E143 annotated valid-like mean raw-Omni contact delta: `{analysis.get('valid_like_mean_contact_delta', '')}`",
        "",
    ]
    for bench in BENCHMARKS:
        lines += [
            f"## {bench}",
            "",
            "| method | cases | fall | 手物接触 | 5cm | 10cm | 手物穿透 | 深穿透2cm | 腿穿透 |",
            "|---|---:|---:|---:|---:|---:|---:|---:|---:|",
        ]
        for row in by_bench[bench]:
            lines.append(
                f"| `{row['method']}` | {row['case_count']} | {row['fall_count']} | "
                f"{row['手物接触']:.4f} | {row['5cm']:.4f} | {row['10cm']:.4f} | "
                f"{row['手物穿透']:.4f} | {row['深穿透2cm']:.4f} | {row['腿穿透']:.4f} |"
            )
        lines += [
            "",
            "| metric | rubber-sphere | OmniRetarget-rubber |",
            "|---|---:|---:|",
        ]
        for row in diffs_by_bench[bench]:
            lines.append(
                f"| {row['metric']} | {row['rubber-sphere']:+.4f} | "
                f"{row['OmniRetarget-rubber']:+.4f} |"
            )
        lines.append("")
    lines += [
        "## Interpretation",
        "",
        "- On clean6/relaxed8, rubber hand Spider still gives small positive 5cm/10cm deltas over sphere Spider.",
        "- Rubber reduces hand-object penetration substantially versus sphere Spider.",
        "- Rubber lowers physics-contact fraction and raises leg penetration on these cleaner cases, so the conclusion remains a geometry tradeoff, not a complete contact-quality win.",
        "",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def append_sheet(ws, rows: list[dict[str, Any]], fields: list[str]) -> None:
    ws.append(fields)
    for row in rows:
        ws.append([row.get(field, "") for field in fields])
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    for col in ws.columns:
        letter = col[0].column_letter
        width = min(max(len(str(cell.value or "")) for cell in col) + 2, 42)
        ws.column_dimensions[letter].width = width


def write_workbook(
    path: Path,
    summary_rows: list[dict[str, Any]],
    diff_rows: list[dict[str, Any]],
    case_rows_all: list[dict[str, Any]],
    annotation_rows: list[dict[str, Any]],
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "benchmark平均"
    summary_fields = ["benchmark", "method", "case_count", "fall_count", *METRICS]
    append_sheet(ws, summary_rows, summary_fields)

    ws = wb.create_sheet("benchmark_diff")
    append_sheet(ws, diff_rows, ["benchmark", "metric", "rubber-sphere", "OmniRetarget-rubber", "rubber-OmniRetarget"])

    case_fields = [
        "benchmark",
        "short_case_id",
        "e148_case_id",
        "object_key",
        "run_status",
        "rubber_source_exp",
        *ANNOTATION_FIELDS,
    ]
    for metric in DIFF_METRICS:
        case_fields.extend(
            [
                f"{metric}_OmniRetarget",
                f"{metric}_sphere Spider",
                f"{metric}_rubber hand Spider",
                f"{metric}_rubber-sphere",
                f"{metric}_OmniRetarget-rubber",
            ]
        )

    ws = wb.create_sheet("clean6逐case")
    append_sheet(ws, [row for row in case_rows_all if row["benchmark"] == "clean6_primary"], case_fields)
    ws = wb.create_sheet("relaxed8逐case")
    append_sheet(ws, [row for row in case_rows_all if row["benchmark"] == "relaxed8_valid_like"], case_fields)

    ws = wb.create_sheet("case_source_annotations")
    append_sheet(ws, annotation_rows, ["short_case_id", "in_clean6", "in_relaxed8", *ANNOTATION_FIELDS])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    by_short, annotations, analysis = load_inputs()
    missing = {bench: [case for case in cases if case not in by_short] for bench, cases in BENCHMARKS.items()}
    bad = {bench: cases for bench, cases in missing.items() if cases}
    if bad:
        raise SystemExit(f"missing benchmark rows in E148 comparison: {bad}")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    summary = method_summary_rows(by_short)
    diffs = diff_summary_rows(summary)
    cases = case_rows(by_short, annotations)
    annotation_rows = annotation_source_rows(annotations)

    write_tsv(OUT_DIR / "e149_method_summary.tsv", summary, ["benchmark", "method", "case_count", "fall_count", *METRICS])
    write_tsv(
        OUT_DIR / "e149_diff_summary.tsv",
        diffs,
        ["benchmark", "metric", "rubber-sphere", "OmniRetarget-rubber", "rubber-OmniRetarget"],
    )
    case_fields = list(cases[0].keys())
    write_tsv(OUT_DIR / "e149_case_comparison.tsv", cases, case_fields)
    write_tsv(OUT_DIR / "e149_case_source_annotations.tsv", annotation_rows, list(annotation_rows[0].keys()))
    write_json(
        OUT_DIR / "e149_summary.json",
        {
            "benchmarks": BENCHMARKS,
            "method_summary_rows": len(summary),
            "diff_summary_rows": len(diffs),
            "case_rows": len(cases),
            "xlsx": str((OUT_DIR / "E149_e143_clean_rubber_benchmark.xlsx").relative_to(REPO)),
        },
    )
    write_markdown(OUT_DIR / "e149_summary.md", summary, diffs, analysis)
    write_workbook(OUT_DIR / "E149_e143_clean_rubber_benchmark.xlsx", summary, diffs, cases, annotation_rows)
    print(f"wrote {OUT_DIR.relative_to(REPO)}")
    print(f"summary_rows={len(summary)} diff_rows={len(diffs)} case_rows={len(cases)}")


if __name__ == "__main__":
    main()
