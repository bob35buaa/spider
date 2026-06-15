#!/usr/bin/env python3
"""E162 post-E147 RL-safe unified re-evaluation."""

from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from eval.core.core_metrics import (  # noqa: E402
    EVAL_METRIC_STANDARD_ID,
    METRIC_FIELDS,
    EvalConfig,
    STANDARD_DELTA_METRICS,
    STANDARD_MASK_DELTA_METRICS,
    STANDARD_TRACK_DIAG,
    evaluate_sequence,
)


REPO = Path(__file__).resolve().parents[5]
VARIANTS_TSV = REPO / "workspace/core4d/scripts/experiments/E162/variants.tsv"
RESULT_ROOT = REPO / "workspace/core4d/results/E162/post_e147_rl_safe_reeval"
CONTACT_GATE_METRIC = "hand_object_physics_contact_in_mask_frac"
CONTACT_CLEAN3_METRIC = "hand_object_physics_contact_3mm_in_mask_frac"
CONTACT_DROP_FAIL_TH = -0.05

TABLE4_FIELDS = [
    "track_joint_err_deg_mean",
    "track_eef_pos_err_cm_mean",
    "track_eef_ori_err_deg_mean",
    "track_root_pos_err_cm_mean",
    "track_root_ori_err_deg_mean",
    "track_obj_pos_err_cm_mean",
    "track_obj_ori_err_deg_mean",
]

SUMMARY_METRICS = [
    "success_tracked",
    "hand_geom_near_5cm_frac",
    "hand_geom_penetration_2mm_frac",
    "hand_geom_penetration_5mm_frac",
    CONTACT_GATE_METRIC,
    CONTACT_CLEAN3_METRIC,
    "hand_object_physics_contact_5mm_in_mask_frac",
    "hand_object_physics_penetration_3mm_frame_frac",
    "hand_object_physics_penetration_5mm_frame_frac",
    "hand_object_release_false_contact_3mm_frac",
    "hand_object_release_false_contact_5mm_frac",
    "leg_penetration_frac",
    "obj_err_mean_m",
    *TABLE4_FIELDS,
]

DELTA_METRICS = list(dict.fromkeys([
    CONTACT_GATE_METRIC,
    "hand_geom_near_5cm_frac",
    *STANDARD_MASK_DELTA_METRICS,
    *STANDARD_DELTA_METRICS,
    "hand_object_physics_contact_5mm_in_mask_frac",
    "hand_object_release_false_contact_3mm_frac",
    "hand_object_release_false_contact_5mm_frac",
    *TABLE4_FIELDS,
]))

METHOD_LABELS_ZH = {
    "E147_spider_rubberhand": "E147基线 rubberhand",
    "E148_spiderrubberhand": "E148扩展 rubberhand",
    "E150_off05": "E150 offset 5cm",
    "E150_off08": "E150 offset 8cm",
    "E150_off11": "E150 offset 11cm",
    "E151_baseline": "E151 baseline",
    "E151_b2_sup": "E151 B2支撑",
    "E151_b2_tip": "E151 B2指尖",
    "E151_b1_mesh": "E151 B1 mesh",
    "E152_baseline": "E152 baseline",
    "E152_gateA": "E152 gateA",
    "E152_b1": "E152 B1",
    "E152_gateA_b1": "E152 gateA+B1",
    "E153_baseline": "E153 baseline",
    "E153_b1": "E153 B1",
    "E153_gateA_b1_sdf005_v05": "E153 sdf005/v05",
    "E153_gateA_b1_sdf005_v10": "E153 sdf005/v10",
    "E153_gateA_b1_sdf010_v05": "E153 sdf010/v05",
    "E153_gateA_b1_sdf010_v10": "E153 sdf010/v10",
    "E153_gateA_b1_sdf015_v05": "E153 sdf015/v05",
    "E153_gateA_b1_sdf015_v10": "E153 sdf015/v10",
    "E155_ramp5": "E155 ramp5",
    "E155_ramp10": "E155 ramp10",
    "E155_decay": "E155 decay",
    "E155_neutral": "E155 neutral",
    "E156_spiderrubberhand": "E156 rubberhand",
    "E156_gateA": "E156 gateA",
    "E156_E155_decay": "E156 decay",
    "E158_gateA_surfaceBandA": "E158 surfaceBand-A",
    "E159_gateA_surfaceBandA2": "E159 surfaceBand-A2",
    "E160_gateA_surfaceBandA2_postureRerankA": "E160 postureRerankA",
    "E161_gateA_surfaceBandA2_postureRerankA": "E161 M0 posture",
    "E161_gateA_surfaceBandA2_postureRerankA_surfaceBandReleaseDecay": "E161 M1 releaseDecay",
    "E161_gateA_surfaceBandA2_postureRerankA_surfaceBandStrictMask": "E161 M2 strictMask",
}

CASE_STATUS_ZH = {
    "pass": "通过",
    "contact_regression_fail": "接触退化",
    "fall_fail": "摔倒",
    "tracking_fail": "tracking失败",
    "contact_metric_missing": "接触指标缺失",
    "eval_missing": "评测缺失",
    "needs_baseline_or_manual_review": "缺E147基线/需审查",
}

BASELINE_STATUS_ZH = {
    "ok": "有E147基线",
    "missing_e147_case": "缺E147 case",
    "missing_e147_contact_mask": "缺E147 mask",
    "missing_e147_qpos": "缺E147 qpos",
    "missing_e147_scene": "缺E147 scene",
    "missing_e147_kin_ref": "缺E147 GT",
}


def repo_path(path: str | Path) -> Path:
    p = Path(path)
    return p if p.is_absolute() else REPO / p


def rel(path: Path | str) -> str:
    if not path:
        return ""
    p = Path(path)
    try:
        return str(p.resolve().relative_to(REPO.resolve()))
    except Exception:
        return str(path)


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def fmt(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return f"{value:.8g}" if math.isfinite(value) else ""
    return str(value)


def write_tsv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: fmt(row.get(field, "")) for field in fields})


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def finite(value: Any, default: float = math.nan) -> float:
    try:
        out = float(value)
    except (TypeError, ValueError):
        return default
    return out if math.isfinite(out) else default


def mean(values: list[Any]) -> float:
    vals = [finite(v) for v in values]
    vals = [v for v in vals if math.isfinite(v)]
    return statistics.fmean(vals) if vals else math.nan


def worst(values: list[Any], high_is_bad: bool = True) -> float:
    vals = [finite(v) for v in values]
    vals = [v for v in vals if math.isfinite(v)]
    if not vals:
        return math.nan
    return max(vals) if high_is_bad else min(vals)


def bool_text(value: Any) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "pass"}


def missing_blocking_eval(row: dict[str, str]) -> list[str]:
    missing = []
    for label, key in [("qpos", "qpos_path"), ("scene", "scene_xml")]:
        if not row.get(key) or not repo_path(row[key]).is_file():
            missing.append(f"missing_{label}")
    return missing


def add_success_flags(row: dict[str, Any], cfg: EvalConfig) -> None:
    pz_term = finite(row.get("track_pelvis_z_err_terminal_m"), math.inf)
    row["success_tracked"] = bool(
        not bool_text(row.get("fall_flag"))
        and math.isfinite(pz_term)
        and pz_term <= cfg.track_pelvis_terminal_th_m
    )


def evaluate_manifest_row(row: dict[str, str], cfg: EvalConfig) -> dict[str, Any] | None:
    blocking = missing_blocking_eval(row)
    if blocking:
        return None
    eval_row = {
        "case_id": row["case_id"],
        "variant": row["variant"],
        "object_key": row.get("object_key", ""),
        "object_category": row.get("object_category", ""),
        "expected_quality": row.get("expected_quality", ""),
    }
    contact_mask = repo_path(row["contact_mask_path"]) if row.get("contact_mask_path") else None
    kin_ref = repo_path(row["kin_ref_path"]) if row.get("kin_ref_path") else None
    metrics = evaluate_sequence(
        row=eval_row,
        method=row["canonical_method_id"],
        hand_collision_variant_id=row.get("hand_collision_variant_id") or "rubber_hull",
        qpos_path=repo_path(row["qpos_path"]),
        scene_xml=repo_path(row["scene_xml"]),
        config=cfg,
        kin_ref_path=kin_ref,
        contact_mask_path=contact_mask,
        person_idx=int(row["person_idx"]) if row.get("person_idx") not in {"", None} else None,
    )
    add_success_flags(metrics, cfg)
    metrics.update(
        {
            "row_id": row["row_id"],
            "manifest_exp_id": row["manifest_exp_id"],
            "source_exp_id": row["source_exp_id"],
            "source_method_id": row["source_method_id"],
            "canonical_method_id": row["canonical_method_id"],
            "method_display": row["method_display"],
            "short_case_id": row["short_case_id"],
            "video_path": row.get("video_path", ""),
            "contact_mask_path": row.get("contact_mask_path", ""),
            "kin_ref_path": row.get("kin_ref_path", ""),
            "baseline_status": row.get("baseline_status", ""),
            "baseline_method_id": row.get("baseline_method_id", ""),
            "baseline_source_exp_id": row.get("baseline_source_exp_id", ""),
            "baseline_variant": row.get("baseline_variant", ""),
            "baseline_qpos_path": row.get("baseline_qpos_path", ""),
            "baseline_scene_xml": row.get("baseline_scene_xml", ""),
            "row_status": row.get("row_status", ""),
            "eval_status": "ok",
        }
    )
    return metrics


def evaluate_synthetic_baseline(row: dict[str, str], cfg: EvalConfig) -> dict[str, Any] | None:
    if not row.get("baseline_qpos_path") or not row.get("baseline_scene_xml"):
        return None
    if not repo_path(row["baseline_qpos_path"]).is_file() or not repo_path(row["baseline_scene_xml"]).is_file():
        return None
    base_row = {
        "case_id": row["short_case_id"],
        "variant": row.get("baseline_variant") or f"E147_{row['short_case_id']}_synthetic_baseline",
        "object_key": row.get("object_key", ""),
        "object_category": row.get("object_category", ""),
        "expected_quality": row.get("expected_quality", ""),
    }
    metrics = evaluate_sequence(
        row=base_row,
        method="E147_spider_rubberhand",
        hand_collision_variant_id="rubber_hull",
        qpos_path=repo_path(row["baseline_qpos_path"]),
        scene_xml=repo_path(row["baseline_scene_xml"]),
        config=cfg,
        kin_ref_path=repo_path(row["baseline_kin_ref_path"]) if row.get("baseline_kin_ref_path") else None,
        contact_mask_path=repo_path(row["baseline_contact_mask_path"]) if row.get("baseline_contact_mask_path") else None,
        person_idx=int(row["person_idx"]) if row.get("person_idx") not in {"", None} else None,
    )
    add_success_flags(metrics, cfg)
    metrics.update(
        {
            "row_id": f"synthetic_E147_{row['short_case_id']}",
            "manifest_exp_id": "E147",
            "source_exp_id": "E147",
            "source_method_id": "spider-rubberhand",
            "canonical_method_id": "E147_spider_rubberhand",
            "method_display": "spider-rubberhand",
            "short_case_id": row["short_case_id"],
            "video_path": "",
            "contact_mask_path": row.get("baseline_contact_mask_path", ""),
            "kin_ref_path": row.get("baseline_kin_ref_path", ""),
            "baseline_status": row.get("baseline_status", ""),
            "baseline_method_id": "E147_spider_rubberhand",
            "baseline_source_exp_id": "E147",
            "baseline_variant": row.get("baseline_variant", ""),
            "baseline_qpos_path": row.get("baseline_qpos_path", ""),
            "baseline_scene_xml": row.get("baseline_scene_xml", ""),
            "row_status": "synthetic_baseline",
            "eval_status": "ok",
        }
    )
    return metrics


def build_baseline_by_case(metric_rows: list[dict[str, Any]], manifest_rows: list[dict[str, str]], cfg: EvalConfig) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for row in metric_rows:
        if row.get("canonical_method_id") == "E147_spider_rubberhand":
            out.setdefault(row["short_case_id"], row)
    for row in manifest_rows:
        if row.get("baseline_status") != "ok" or row["short_case_id"] in out:
            continue
        baseline = evaluate_synthetic_baseline(row, cfg)
        if baseline is not None:
            out[row["short_case_id"]] = baseline
    return out


def delta_row(run: dict[str, Any], baseline: dict[str, Any] | None) -> dict[str, Any]:
    row: dict[str, Any] = {
        "row_id": run["row_id"],
        "manifest_exp_id": run["manifest_exp_id"],
        "source_exp_id": run["source_exp_id"],
        "source_method_id": run["source_method_id"],
        "canonical_method_id": run["canonical_method_id"],
        "method_display": run["method_display"],
        "short_case_id": run["short_case_id"],
        "case_id": run["case_id"],
        "variant": run["variant"],
        "baseline_status": run.get("baseline_status", ""),
        "baseline_variant": run.get("baseline_variant", ""),
        "success_tracked": run.get("success_tracked", False),
        "fall_flag": run.get("fall_flag", False),
        "eval_status": run.get("eval_status", ""),
    }
    for key in STANDARD_TRACK_DIAG:
        row[key] = run.get(key, math.nan)
    if baseline is not None:
        for key in DELTA_METRICS:
            row[f"{key}_baseline"] = baseline.get(key, math.nan)
            row[f"{key}_run"] = run.get(key, math.nan)
            row[f"{key}_delta_vs_E147"] = finite(run.get(key)) - finite(baseline.get(key))
    contact_delta = finite(row.get(f"{CONTACT_GATE_METRIC}_delta_vs_E147"))
    row["contact3_inmask_delta_vs_E147"] = contact_delta
    row["contact_inmask_delta_vs_E147"] = contact_delta
    row["contact_gate_delta_vs_E147"] = contact_delta
    row["clean3mm_contact_inmask_delta_vs_E147"] = finite(
        row.get(f"{CONTACT_CLEAN3_METRIC}_delta_vs_E147")
    )
    row["case_status"] = case_status(row)
    row["rl_safe_case_pass"] = row["case_status"] == "pass"
    return row


def case_status(row: dict[str, Any]) -> str:
    baseline_status = str(row.get("baseline_status") or "")
    if row.get("eval_status") != "ok":
        return "eval_missing"
    if baseline_status != "ok":
        return "needs_baseline_or_manual_review"
    contact_delta = finite(row.get("contact3_inmask_delta_vs_E147"))
    if not math.isfinite(contact_delta):
        return "contact_metric_missing"
    if contact_delta < CONTACT_DROP_FAIL_TH:
        return "contact_regression_fail"
    if bool_text(row.get("fall_flag")):
        return "fall_fail"
    if not bool_text(row.get("success_tracked")):
        return "tracking_fail"
    return "pass"


def missing_delta_rows(missing_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for row in missing_rows:
        item = {
            "row_id": row.get("row_id", ""),
            "manifest_exp_id": row.get("manifest_exp_id", ""),
            "source_exp_id": row.get("source_exp_id", ""),
            "source_method_id": row.get("source_method_id", ""),
            "canonical_method_id": row.get("canonical_method_id", ""),
            "method_display": row.get("method_display", ""),
            "short_case_id": row.get("short_case_id", ""),
            "case_id": row.get("case_id", ""),
            "variant": row.get("variant", ""),
            "baseline_status": row.get("baseline_status", ""),
            "baseline_variant": row.get("baseline_variant", ""),
            "success_tracked": False,
            "fall_flag": "",
            "eval_status": row.get("eval_status", "eval_missing"),
            "contact3_inmask_delta_vs_E147": math.nan,
            "contact_inmask_delta_vs_E147": math.nan,
            "contact_gate_delta_vs_E147": math.nan,
            "clean3mm_contact_inmask_delta_vs_E147": math.nan,
            "case_status": "eval_missing",
            "rl_safe_case_pass": False,
            "missing_reason": row.get("missing_reason", ""),
        }
        out.append(item)
    return out


def summarize_method(method: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
    applicable = [r for r in rows if r.get("baseline_status") == "ok" and r.get("eval_status") == "ok"]
    failures = [r for r in rows if r.get("case_status") != "pass"]
    contact_fail = [r for r in rows if r.get("case_status") == "contact_regression_fail"]
    tracking_fail = [r for r in rows if r.get("case_status") == "tracking_fail"]
    fall_fail = [r for r in rows if r.get("case_status") == "fall_fail"]
    baseline_missing = [r for r in rows if r.get("baseline_status") != "ok"]
    eval_missing = [r for r in rows if r.get("eval_status") != "ok"]
    out: dict[str, Any] = {
        "canonical_method_id": method,
        "method_display": rows[0].get("method_display", method) if rows else method,
        "source_exp_ids": ",".join(sorted({str(r.get("source_exp_id", "")) for r in rows if r.get("source_exp_id")})),
        "n_rows": len(rows),
        "n_applicable_cases": len(applicable),
        "rl_safe_case_pass_count": sum(1 for r in applicable if r.get("rl_safe_case_pass")),
        "rl_safe_method_pass": bool(applicable) and not failures,
        "failed_cases": ",".join(r["short_case_id"] for r in failures),
        "contact_regression_failed_cases": ",".join(r["short_case_id"] for r in contact_fail),
        "tracking_failed_cases": ",".join(r["short_case_id"] for r in tracking_fail),
        "fall_failed_cases": ",".join(r["short_case_id"] for r in fall_fail),
        "baseline_missing_cases": ",".join(r["short_case_id"] for r in baseline_missing),
        "eval_missing_cases": ",".join(r["short_case_id"] for r in eval_missing),
    }
    high_good = {
        "success_tracked",
        CONTACT_GATE_METRIC,
        "hand_object_physics_contact_5mm_in_mask_frac",
    }
    for key in SUMMARY_METRICS:
        if key == "success_tracked":
            vals = [1.0 if bool_text(r.get(key)) else 0.0 for r in applicable]
        else:
            vals = [r.get(f"{key}_run", r.get(key)) for r in applicable]
        out[f"{key}_mean"] = mean(vals)
        out[f"{key}_worst"] = worst(vals, high_is_bad=key not in high_good)
    out["contact3_inmask_delta_vs_E147_mean"] = mean([r.get("contact3_inmask_delta_vs_E147") for r in applicable])
    out["contact3_inmask_delta_vs_E147_worst"] = worst(
        [r.get("contact3_inmask_delta_vs_E147") for r in applicable],
        high_is_bad=False,
    )
    out["contact_inmask_delta_vs_E147_mean"] = mean([r.get("contact_inmask_delta_vs_E147") for r in applicable])
    out["contact_inmask_delta_vs_E147_worst"] = worst(
        [r.get("contact_inmask_delta_vs_E147") for r in applicable],
        high_is_bad=False,
    )
    out["clean3mm_contact_inmask_delta_vs_E147_mean"] = mean(
        [r.get("clean3mm_contact_inmask_delta_vs_E147") for r in applicable]
    )
    out["clean3mm_contact_inmask_delta_vs_E147_worst"] = worst(
        [r.get("clean3mm_contact_inmask_delta_vs_E147") for r in applicable],
        high_is_bad=False,
    )
    return out


def method_label_zh(method_id: str) -> str:
    return METHOD_LABELS_ZH.get(method_id, method_id)


def method_sort_key(row: dict[str, Any]) -> tuple[int, str]:
    method = str(row.get("canonical_method_id", ""))
    m = re_match_exp(method)
    if m is not None:
        return m, method
    return 999, method


def re_match_exp(text: str) -> int | None:
    match = None
    for idx, ch in enumerate(text):
        if ch == "E" and idx + 3 < len(text) and text[idx + 1 : idx + 4].isdigit():
            match = int(text[idx + 1 : idx + 4])
            break
    return match


def xlsx_value(row: dict[str, Any], key: str | Any) -> Any:
    if callable(key):
        value = key(row)
    else:
        value = row.get(key, "")
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float) and not math.isfinite(value):
        return ""
    return value


def bool_zh(value: Any) -> str:
    return "是" if bool_text(value) else "否"


def case_status_zh(value: Any) -> str:
    return CASE_STATUS_ZH.get(str(value), str(value))


def baseline_status_zh(value: Any) -> str:
    text = str(value)
    if text in BASELINE_STATUS_ZH:
        return BASELINE_STATUS_ZH[text]
    if text.startswith("missing_e147_"):
        return text.replace("missing_e147_", "缺E147 ")
    return text


def method_rl_safe_label(row: dict[str, Any]) -> str:
    applicable = int(finite(row.get("n_applicable_cases"), 0.0))
    passed = int(finite(row.get("rl_safe_case_pass_count"), 0.0))
    if applicable <= 0:
        return "未判定"
    if passed < applicable:
        return "否"
    if row.get("baseline_missing_cases"):
        return "可比通过/缺基线"
    if row.get("eval_missing_cases"):
        return "可比通过/缺评测"
    return "是"


def method_failure_cases(row: dict[str, Any]) -> str:
    cases: list[str] = []
    for key in [
        "contact_regression_failed_cases",
        "tracking_failed_cases",
        "fall_failed_cases",
        "eval_missing_cases",
    ]:
        text = str(row.get(key) or "")
        for case in [part for part in text.split(",") if part]:
            if case not in cases:
                cases.append(case)
    return ",".join(cases)


def append_sheet(
    wb: Workbook,
    title: str,
    rows: list[dict[str, Any]],
    columns: list[tuple[str, str | Any]],
    *,
    active: bool = False,
    fill: str = "D9EAF7",
) -> Any:
    ws = wb.active if active else wb.create_sheet(title)
    ws.title = title
    ws.append([header for header, _ in columns])
    for row in rows:
        ws.append([xlsx_value(row, key) for _, key in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, col_cells in enumerate(ws.columns, start=1):
        width = min(max(len(str(cell.value or "")) for cell in col_cells) + 2, 34)
        ws.column_dimensions[get_column_letter(idx)].width = max(width, 10)
    return ws


def detail_fields(rows: list[dict[str, Any]], front: list[str]) -> list[str]:
    fields = sorted({key for row in rows for key in row.keys()})
    return [field for field in front if field in fields] + [field for field in fields if field not in front]


def append_detail_sheet(wb: Workbook, title: str, rows: list[dict[str, Any]], front: list[str], *, hidden: bool = False) -> None:
    fields = detail_fields(rows, front)
    ws = append_sheet(wb, title, rows, [(field, field) for field in fields], fill="EADCF8")
    if hidden:
        ws.sheet_state = "hidden"


def xlsx_info_rows(meta: dict[str, Any]) -> list[dict[str, Any]]:
    rows = [
        {
            "项": "E147 vs E148",
            "说明": "E147 是统一 baseline：原始 10-case spider-rubberhand。E148 是 E143 24-case rubberhand 扩展；重叠 case 复用 E147，新 case 由 E148 跑出。E148 不是本轮 baseline。",
        },
        {
            "项": "RL hard gate",
            "说明": "使用 raw hand_object_physics_contact_in_mask_frac，即真实 3cm contact mask 内是否有物理接触。相对 E147 下降 < -0.05 直接判接触退化。",
        },
        {
            "项": "clean 3mm contact",
            "说明": "clean 3mm 是 raw contact 的子集，只保留 contact.dist >= -0.003m 的接触；它用于诊断接触质量，不能抵消 raw contact 下降。",
        },
        {
            "项": "baseline missing",
            "说明": "缺 E147 同 case baseline 或缺 E147 contact mask 的 rows 只报告绝对指标，不判 RL-safe pass。",
        },
    ]
    rows.extend({"项": key, "说明": value} for key, value in meta.items())
    return rows


def write_xlsx(
    path: Path,
    method_summary: list[dict[str, Any]],
    metric_rows: list[dict[str, Any]],
    delta_rows: list[dict[str, Any]],
    case_failures: list[dict[str, Any]],
    baseline_missing: list[dict[str, Any]],
    meta: dict[str, Any],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    method_rows = sorted(method_summary, key=method_sort_key)
    main_columns: list[tuple[str, str | Any]] = [
        ("方法", lambda r: method_label_zh(str(r.get("canonical_method_id", "")))),
        ("E147可比case", "n_applicable_cases"),
        ("RL通过case", "rl_safe_case_pass_count"),
        ("RL安全", method_rl_safe_label),
        ("失败case", method_failure_cases),
        ("接触退化case", "contact_regression_failed_cases"),
        ("缺E147基线", "baseline_missing_cases"),
        ("几何接近5cm", "hand_geom_near_5cm_frac_mean"),
        ("几何穿透2mm", "hand_geom_penetration_2mm_frac_mean"),
        ("物理接触", f"{CONTACT_GATE_METRIC}_mean"),
        ("物理接触Δ最差", "contact_inmask_delta_vs_E147_worst"),
        ("物理穿透3mm", "hand_object_physics_penetration_3mm_frame_frac_mean"),
        ("关节误差(°)", "track_joint_err_deg_mean_mean"),
        ("末端位置误差(cm)", "track_eef_pos_err_cm_mean_mean"),
        ("末端朝向误差(°)", "track_eef_ori_err_deg_mean_mean"),
        ("物体位置误差(cm)", "track_obj_pos_err_cm_mean_mean"),
        ("物体朝向误差(°)", "track_obj_ori_err_deg_mean_mean"),
    ]
    append_sheet(wb, "主表", method_rows, main_columns, active=True, fill="D9EAF7")

    case_rows = sorted(delta_rows, key=lambda r: (str(r.get("short_case_id", "")), str(r.get("canonical_method_id", ""))))
    case_columns: list[tuple[str, str | Any]] = [
        ("方法", lambda r: method_label_zh(str(r.get("canonical_method_id", "")))),
        ("case", "short_case_id"),
        ("状态", lambda r: case_status_zh(r.get("case_status", ""))),
        ("RL通过", "rl_safe_case_pass"),
        ("raw接触Δ", "contact_inmask_delta_vs_E147"),
        ("E147 raw接触", f"{CONTACT_GATE_METRIC}_baseline"),
        ("本方法 raw接触", f"{CONTACT_GATE_METRIC}_run"),
        ("clean3接触Δ", "clean3mm_contact_inmask_delta_vs_E147"),
        ("E147 clean3", f"{CONTACT_CLEAN3_METRIC}_baseline"),
        ("本方法 clean3", f"{CONTACT_CLEAN3_METRIC}_run"),
        ("几何接近5cm", "hand_geom_near_5cm_frac_run"),
        ("几何穿透2mm", "hand_geom_penetration_2mm_frac_run"),
        ("物理穿透3mm", "hand_object_physics_penetration_3mm_frame_frac_run"),
        ("关节误差(°)", "track_joint_err_deg_mean"),
        ("末端位置误差(cm)", "track_eef_pos_err_cm_mean"),
        ("末端朝向误差(°)", "track_eef_ori_err_deg_mean"),
        ("物体位置误差(cm)", "track_obj_pos_err_cm_mean"),
        ("物体朝向误差(°)", "track_obj_ori_err_deg_mean"),
        ("baseline状态", lambda r: baseline_status_zh(r.get("baseline_status", ""))),
    ]
    append_sheet(wb, "逐case", case_rows, case_columns, fill="D9EAD3")

    tracking_rows = sorted(metric_rows, key=lambda r: (str(r.get("canonical_method_id", "")), str(r.get("short_case_id", ""))))
    tracking_columns: list[tuple[str, str | Any]] = [
        ("方法", lambda r: method_label_zh(str(r.get("canonical_method_id", "")))),
        ("case", "short_case_id"),
        ("关节误差(°)", "track_joint_err_deg_mean"),
        ("末端位置误差(cm)", "track_eef_pos_err_cm_mean"),
        ("末端朝向误差(°)", "track_eef_ori_err_deg_mean"),
        ("Root位置误差(cm)", "track_root_pos_err_cm_mean"),
        ("Root朝向误差(°)", "track_root_ori_err_deg_mean"),
        ("物体位置误差(cm)", "track_obj_pos_err_cm_mean"),
        ("物体朝向误差(°)", "track_obj_ori_err_deg_mean"),
    ]
    append_sheet(wb, "Tracking指标", tracking_rows, tracking_columns, fill="FFF2CC")
    append_sheet(wb, "说明", xlsx_info_rows(meta), [("项", "项"), ("说明", "说明")], fill="FCE4D6")

    detail_front = [
        "canonical_method_id",
        "short_case_id",
        "case_id",
        "variant",
        "case_status",
        "rl_safe_method_pass",
        "rl_safe_case_pass",
        "contact_inmask_delta_vs_E147",
        "clean3mm_contact_inmask_delta_vs_E147",
    ]
    append_detail_sheet(wb, "method_summary_full", method_summary, detail_front, hidden=True)
    append_detail_sheet(wb, "per_case_full", metric_rows, detail_front, hidden=True)
    append_detail_sheet(wb, "delta_full", delta_rows, detail_front, hidden=True)
    append_detail_sheet(wb, "failures_full", case_failures, detail_front, hidden=True)
    append_detail_sheet(wb, "baseline_missing_full", baseline_missing, detail_front, hidden=True)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("stage", nargs="?", default="full", choices=["full"])
    parser.add_argument("--allow-missing", action="store_true")
    args = parser.parse_args()

    cfg = EvalConfig()
    eval_dir = RESULT_ROOT / "eval" / args.stage
    manifest_rows = read_tsv(VARIANTS_TSV)
    metric_rows: list[dict[str, Any]] = []
    missing_rows: list[dict[str, Any]] = []

    for row in manifest_rows:
        blocking = missing_blocking_eval(row)
        if blocking:
            item = dict(row)
            item["eval_status"] = "eval_missing"
            item["missing_reason"] = ";".join(blocking)
            missing_rows.append(item)
            continue
        try:
            metrics = evaluate_manifest_row(row, cfg)
        except Exception as exc:
            item = dict(row)
            item["eval_status"] = "eval_error"
            item["missing_reason"] = f"{type(exc).__name__}: {exc}"
            missing_rows.append(item)
            if not args.allow_missing:
                raise
            continue
        if metrics is None:
            item = dict(row)
            item["eval_status"] = "eval_missing"
            item["missing_reason"] = "unknown"
            missing_rows.append(item)
            continue
        metric_rows.append(metrics)

    baseline_by_case = build_baseline_by_case(metric_rows, manifest_rows, cfg)
    delta_rows = [delta_row(row, baseline_by_case.get(row["short_case_id"])) for row in metric_rows]
    delta_rows.extend(missing_delta_rows(missing_rows))
    case_failures = [row for row in delta_rows if row.get("case_status") != "pass"]
    baseline_missing = [row for row in delta_rows if row.get("baseline_status") != "ok"]
    methods = sorted({row["canonical_method_id"] for row in delta_rows})
    method_summary = [summarize_method(method, [row for row in delta_rows if row["canonical_method_id"] == method]) for method in methods]

    metric_front = [
        "row_id",
        "manifest_exp_id",
        "source_exp_id",
        "source_method_id",
        "canonical_method_id",
        "method_display",
        "short_case_id",
        "success_tracked",
        "eval_status",
        "baseline_status",
    ]
    metric_fields = metric_front + [field for field in METRIC_FIELDS if field not in metric_front] + [
        "video_path",
        "contact_mask_path",
        "kin_ref_path",
        "baseline_variant",
        "baseline_qpos_path",
        "baseline_scene_xml",
    ]
    delta_fields = sorted({key for row in delta_rows for key in row.keys()})
    delta_front = [
        "canonical_method_id",
        "method_display",
        "short_case_id",
        "case_id",
        "variant",
        "case_status",
        "rl_safe_case_pass",
        "contact3_inmask_delta_vs_E147",
        "contact_inmask_delta_vs_E147",
        "clean3mm_contact_inmask_delta_vs_E147",
        "success_tracked",
        "fall_flag",
        "baseline_status",
        "eval_status",
        "missing_reason",
    ]
    delta_fields = [field for field in delta_front if field in delta_fields] + [
        field for field in delta_fields if field not in delta_front
    ]
    summary_fields = sorted({key for row in method_summary for key in row.keys()})
    summary_front = [
        "canonical_method_id",
        "method_display",
        "source_exp_ids",
        "n_rows",
        "n_applicable_cases",
        "rl_safe_case_pass_count",
        "rl_safe_method_pass",
        "failed_cases",
        "contact_regression_failed_cases",
        "tracking_failed_cases",
        "fall_failed_cases",
        "baseline_missing_cases",
        "eval_missing_cases",
        "contact3_inmask_delta_vs_E147_mean",
        "contact3_inmask_delta_vs_E147_worst",
        "contact_inmask_delta_vs_E147_mean",
        "contact_inmask_delta_vs_E147_worst",
        "clean3mm_contact_inmask_delta_vs_E147_mean",
        "clean3mm_contact_inmask_delta_vs_E147_worst",
    ]
    summary_fields = [field for field in summary_front if field in summary_fields] + [
        field for field in summary_fields if field not in summary_front
    ]

    meta = {
        "metric_standard_id": EVAL_METRIC_STANDARD_ID,
        "stage": args.stage,
        "manifest_rows": len(manifest_rows),
        "metric_rows": len(metric_rows),
        "delta_rows": len(delta_rows),
        "case_failure_rows": len(case_failures),
        "baseline_missing_rows": len(baseline_missing),
        "missing_rows": len(missing_rows),
        "contact_gate_metric": CONTACT_GATE_METRIC,
        "contact_drop_fail_threshold": CONTACT_DROP_FAIL_TH,
    }

    write_tsv(eval_dir / "e162_method_metrics.tsv", metric_rows, metric_fields)
    write_tsv(eval_dir / "e162_delta_vs_E147.tsv", delta_rows, delta_fields)
    write_tsv(eval_dir / "e162_case_failures.tsv", case_failures, delta_fields)
    write_tsv(eval_dir / "e162_method_summary.tsv", method_summary, summary_fields)
    write_tsv(eval_dir / "e162_baseline_missing.tsv", baseline_missing, delta_fields)
    write_json(eval_dir / "e162_eval_summary.json", meta)
    write_xlsx(
        eval_dir / "E162_post_E147_RL_safe_reeval.xlsx",
        method_summary,
        metric_rows,
        delta_rows,
        case_failures,
        baseline_missing,
        meta,
    )
    print(
        f"E162 eval: manifest_rows={len(manifest_rows)} metric_rows={len(metric_rows)} "
        f"failures={len(case_failures)} baseline_missing={len(baseline_missing)} missing={len(missing_rows)}"
    )
    if missing_rows and not args.allow_missing:
        print("MISSING:", [f"{row.get('variant')}:{row.get('missing_reason')}" for row in missing_rows[:20]])
        raise SystemExit(1)


if __name__ == "__main__":
    main()
