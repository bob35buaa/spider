#!/usr/bin/env python3
"""Evaluate E163 narrow symmetric surfaceBand probe and clean8 extension."""

from __future__ import annotations

import argparse
import csv
import json
import math
import sys
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    import yaml
except Exception:  # pragma: no cover - PyYAML is present in the experiment env.
    yaml = None

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from eval.core.core_metrics import EvalConfig  # noqa: E402
from eval_E156_clean8_gate_decay import (  # noqa: E402
    evaluate_row,
    finite,
    read_tsv,
    repo_path,
    write_json,
    write_tsv,
)


REPO = Path(__file__).resolve().parents[5]
E163_VARIANTS = REPO / "workspace/core4d/scripts/experiments/E163/variants.tsv"
E156_METRICS = REPO / "workspace/core4d/results/E156/clean8_gate_decay/eval/full/e156_method_metrics.tsv"
E162_METRICS = REPO / "workspace/core4d/results/E162/post_e147_rl_safe_reeval/eval/full/e162_method_metrics.tsv"
RESULT_ROOT = REPO / "workspace/core4d/results/E163/narrow_surface_band"

THREE_CASES = ["box023_person2", "box021_029_p2", "box004_083_p2"]
CLEAN8_CASES = [
    "box023_person2",
    "box021_029_p2",
    "box004_083_p2",
    "box021_035_p1",
    "box021_035_p2",
    "box004_083_p1",
    "box004_082_p1",
    "box026_139_p1",
]
CONTACT_METRIC = "hand_object_physics_contact_in_mask_frac"
CONTACT_DROP_FAIL_TH = -0.05
RAW_CONTACT_BASELINES = {
    "box023_person2": 0.9077,
    "box021_029_p2": 0.3818,
    "box004_083_p2": 0.5323,
}
RAW_CONTACT_PASS_MIN = {
    "box023_person2": 0.8577,
    "box021_029_p2": 0.3318,
    "box004_083_p2": 0.4823,
}
BOX023_E161_RELEASE_DECAY_RAW = 0.7538

TABLE4_TRACKING = [
    "track_joint_err_deg_mean",
    "track_eef_pos_err_cm_mean",
    "track_eef_ori_err_deg_mean",
    "track_root_pos_err_cm_mean",
    "track_root_ori_err_deg_mean",
    "track_obj_pos_err_cm_mean",
    "track_obj_ori_err_deg_mean",
]

CORE_METRICS = [
    CONTACT_METRIC,
    "hand_object_physics_contact_3mm_in_mask_frac",
    "hand_object_physics_contact_5mm_in_mask_frac",
    "hand_object_physics_penetration_3mm_frame_frac",
    "hand_object_physics_penetration_5mm_frame_frac",
    "hand_geom_penetration_2mm_frac",
    "hand_geom_penetration_5mm_frac",
    "hand_object_release_false_contact_3mm_frac",
    "hand_object_release_false_contact_5mm_frac",
    *TABLE4_TRACKING,
]

METHOD_SPECS = [
    ("omni", "OmniRetarget", "e156", "OmniRetarget", "SPIDER 输入参考"),
    ("rubberhand", "SPIDER+rubberhand", "e156", "spider-rubberhand", "统一 baseline"),
    ("gateA", "+gateA", "e156", "+gateA", "E156 gateA"),
    ("e155_decay", "E155 decay", "e156", "E155_decay", "release smooth transition 参考"),
    ("surfaceA", "surfaceBand-A", "e162", "E158_gateA_surfaceBandA", "E158 宽 band + penalty"),
    ("surfaceA2", "surfaceBand-A2", "e162", "E159_gateA_surfaceBandA2", "E159 宽 band no-penalty"),
    ("posture", "postureRerankA", "e162", "E161_gateA_surfaceBandA2_postureRerankA", "E160/E161 posture rerank"),
    (
        "release_decay",
        "surfaceBand releaseDecay",
        "e162",
        "E161_gateA_surfaceBandA2_postureRerankA_surfaceBandReleaseDecay",
        "E161 direct base",
    ),
    ("e163", "E163 narrowSurfaceBand", "e163", "E163", "本轮新方法"),
]


def target_cases_for_stage(stage: str) -> list[str]:
    return CLEAN8_CASES if stage == "clean8" else THREE_CASES


def artifact_stage_for_eval(stage: str) -> str:
    return "full" if stage == "clean8" else stage


def rel(path: str | Path) -> str:
    p = Path(path)
    try:
        return str(p.resolve().relative_to(REPO.resolve()))
    except Exception:
        return str(path)


def bool_text(value: Any) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "pass"}


def finite_or_blank(value: Any) -> Any:
    val = finite(value)
    return val if math.isfinite(val) else ""


def mean(values: list[Any]) -> float:
    vals = [finite(v) for v in values]
    vals = [v for v in vals if math.isfinite(v)]
    return sum(vals) / len(vals) if vals else math.nan


def stage_paths(row: dict[str, str], stage: str) -> tuple[Path, Path, Path, Path]:
    variant = row["variant"]
    root = RESULT_ROOT / "cem" / stage
    outdir = root / f"{variant}_outdir_{stage}"
    return (
        root / f"{variant}.npz",
        outdir / "trajectory_mjwp_act.npz",
        outdir / "config_act.yaml",
        root / f"{variant}_{stage}.mp4",
    )


def load_yaml(path: Path) -> dict[str, Any]:
    if yaml is None or not path.is_file():
        return {}
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    return data if isinstance(data, dict) else {}


def close_enough(value: Any, target: float, tol: float = 1e-7) -> bool:
    try:
        return abs(float(value) - target) <= tol
    except Exception:
        return False


def artifact_check(row: dict[str, str], stage: str) -> dict[str, Any]:
    root_npz, outdir_npz, config_act, video = stage_paths(row, artifact_stage_for_eval(stage))
    cfg = load_yaml(config_act)
    checks = {
        "variant": row["variant"],
        "short_case_id": row["short_case_id"],
        "root_npz_exists": root_npz.is_file(),
        "trajectory_mjwp_act_exists": outdir_npz.is_file(),
        "config_act_exists": config_act.is_file(),
        "full_mp4_exists": video.is_file(),
        "config_surface_band_min_sdf_ok": close_enough(cfg.get("surface_band_min_sdf_m"), -0.001),
        "config_surface_band_width_ok": close_enough(cfg.get("surface_band_width_m"), 0.003),
        "config_surface_band_sigma_ok": close_enough(cfg.get("surface_band_sigma"), 0.0015),
        "config_surface_band_score_mode_ok": cfg.get("surface_band_score_mode") == "symmetric_abs",
        "config_contact_mask_source_ok": cfg.get("contact_hdmi_mask_source") == "core4d_3cm",
        "config_path": rel(config_act),
        "root_npz": rel(root_npz),
        "trajectory_mjwp_act": rel(outdir_npz),
        "video": rel(video),
    }
    bool_keys = [key for key in checks if key.endswith("_exists") or key.endswith("_ok")]
    checks["artifact_ok"] = all(bool(checks[key]) for key in bool_keys)
    return checks


def run_health(qpos_path: Path) -> dict[str, Any]:
    keys = {
        "cem_hand_gate_valid_frac": ("mean", "hand_gate_valid_frac"),
        "cem_gate_fallback_used": ("mean", "gate_fallback_used"),
        "cem_hand_gate_min_sdf_min": ("min", "hand_gate_min_sdf_min_m"),
        "sample_hand_gate_violation_pct_mean": ("mean", "hand_gate_violation_pct"),
        "surface_band_rew_mean": ("mean", "surface_band_rew_mean"),
        "surface_band_penalty_mean": ("mean", "surface_band_penalty_mean"),
        "surface_band_gate_mean": ("mean", "surface_band_gate_mean"),
        "surface_band_decay_factor_mean": ("mean", "surface_band_decay_factor_mean"),
        "surface_band_sdf_mean": ("mean", "surface_band_sdf_mean_m"),
        "surface_band_score_mean": ("mean", "surface_band_score_mean"),
        "surface_band_penetration_mean": ("mean", "surface_band_penetration_mean_m"),
        "cem_posture_gate_valid_frac": ("mean", "posture_gate_valid_frac"),
        "cem_posture_gate_selected_valid_frac": ("mean", "posture_gate_selected_valid_frac"),
        "cem_posture_gate_fallback_used": ("mean", "posture_gate_fallback_used"),
    }
    out = {dst: "" for _, dst in keys.values()}
    if not qpos_path.is_file():
        return out
    data = np.load(qpos_path, allow_pickle=True)
    for src, (agg, dst) in keys.items():
        if src not in data.files:
            continue
        arr = np.asarray(data[src], dtype=np.float64)
        arr = arr[np.isfinite(arr)]
        if arr.size:
            out[dst] = float(arr.min()) if agg == "min" else float(arr.mean())
    return out


def evaluate_e163(
    stage: str,
    cfg: EvalConfig,
    target_cases: list[str],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    target_set = set(target_cases)
    rows = [row for row in read_tsv(E163_VARIANTS) if row.get("short_case_id") in target_set]
    metric_rows: list[dict[str, Any]] = []
    artifact_rows: list[dict[str, Any]] = []
    missing_rows: list[dict[str, Any]] = []

    for row in rows:
        root_npz, outdir_npz, _, _ = stage_paths(row, artifact_stage_for_eval(stage))
        artifact = artifact_check(row, stage)
        artifact_rows.append(artifact)
        qpos_path = outdir_npz if outdir_npz.is_file() else root_npz
        if not qpos_path.is_file():
            missing_rows.append({"variant": row["variant"], "short_case_id": row["short_case_id"], "reason": "missing_qpos"})
            continue
        eval_row = dict(row)
        eval_row["method"] = "E163 narrowSurfaceBand"
        metrics = evaluate_row(eval_row, qpos_path, cfg)
        if metrics is None:
            missing_rows.append({"variant": row["variant"], "short_case_id": row["short_case_id"], "reason": "evaluate_row_failed"})
            continue
        metrics.update(run_health(qpos_path))
        metrics["method_label"] = "E163 narrowSurfaceBand"
        metrics["source"] = "E163"
        metrics["artifact_ok"] = artifact["artifact_ok"]
        metrics["config_act"] = artifact["config_path"]
        metric_rows.append(metrics)
    return metric_rows, artifact_rows, missing_rows


def load_reference_rows(
    e163_rows: list[dict[str, Any]],
    target_cases: list[str],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    indexes: dict[str, dict[str, dict[str, dict[str, Any]]]] = {"e156": {}, "e162": {}, "e163": {}}
    for row in read_tsv(E156_METRICS):
        indexes["e156"].setdefault(row.get("method", ""), {})[row.get("short_case_id", "")] = row
    for row in read_tsv(E162_METRICS):
        indexes["e162"].setdefault(row.get("canonical_method_id", ""), {})[row.get("short_case_id", "")] = row
    for row in e163_rows:
        indexes["e163"].setdefault("E163", {})[row.get("short_case_id", "")] = row

    rows: list[dict[str, Any]] = []
    missing: list[dict[str, Any]] = []
    for key, label, source, method_id, note in METHOD_SPECS:
        for case in target_cases:
            row = indexes.get(source, {}).get(method_id, {}).get(case)
            if row is None:
                missing.append({"method": label, "source": source, "method_id": method_id, "short_case_id": case})
                continue
            item = dict(row)
            item["method_key"] = key
            item["method_label"] = label
            item["method_note"] = note
            item["short_case_id"] = case
            rows.append(item)
    return rows, missing


def tracked_ok(row: dict[str, Any]) -> bool:
    value = row.get("success_tracked")
    if value not in {None, ""}:
        return bool_text(value)
    terminal_pz = finite(row.get("track_pelvis_z_err_terminal_m"))
    return math.isfinite(terminal_pz) and terminal_pz <= 0.08


def table4_complete(row: dict[str, Any]) -> bool:
    return all(math.isfinite(finite(row.get(key))) for key in TABLE4_TRACKING)


def case_status(row: dict[str, Any], baseline: dict[str, Any]) -> tuple[str, float]:
    if row.get("method_label") == "E163 narrowSurfaceBand" and not bool(row.get("artifact_ok", False)):
        return "产物/配置缺失", math.nan
    raw = finite(row.get(CONTACT_METRIC))
    base_raw = finite(baseline.get(CONTACT_METRIC))
    delta = raw - base_raw
    if not math.isfinite(delta):
        return "接触指标缺失", math.nan
    if delta < CONTACT_DROP_FAIL_TH:
        return "接触退化", delta
    if bool_text(row.get("fall_flag")):
        return "摔倒", delta
    if not tracked_ok(row):
        return "tracking失败", delta
    if not table4_complete(row):
        return "tracking指标缺失", delta
    return "通过", delta


def build_per_case(rows: list[dict[str, Any]], target_cases: list[str]) -> list[dict[str, Any]]:
    by_case_method = {(row["short_case_id"], row["method_label"]): row for row in rows}
    out: list[dict[str, Any]] = []
    for case in target_cases:
        baseline = by_case_method[(case, "SPIDER+rubberhand")]
        for _, label, _, method_id, note in METHOD_SPECS:
            row = by_case_method.get((case, label))
            if row is None:
                continue
            base_raw = finite(baseline.get(CONTACT_METRIC))
            raw_min = RAW_CONTACT_PASS_MIN.get(case, base_raw + CONTACT_DROP_FAIL_TH)
            status, delta = case_status(row, baseline)
            item: dict[str, Any] = {
                "方法": label,
                "case": case,
                "状态": status,
                "物理接触Δ_vs_rubberhand": delta,
                "raw接触下限": raw_min,
                "source/method_id": method_id,
                "说明": note,
                "success_tracked": tracked_ok(row),
                "fall_flag": bool_text(row.get("fall_flag")),
                "Table4完整": table4_complete(row),
            }
            for key in CORE_METRICS:
                item[key] = finite_or_blank(row.get(key))
            out.append(item)
    return out


def build_summary(per_case: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for _, label, _, method_id, note in METHOD_SPECS:
        rows = [row for row in per_case if row["方法"] == label]
        failures = [row for row in rows if row["状态"] != "通过"]
        contact_failures = [row for row in rows if row["状态"] == "接触退化"]
        item: dict[str, Any] = {
            "方法": label,
            "case数": len(rows),
            "通过case": sum(1 for row in rows if row["状态"] == "通过"),
            "失败case": ",".join(row["case"] for row in failures),
            "接触退化case": ",".join(row["case"] for row in contact_failures),
            "物理接触Δ最差": min(
                [finite(row.get("物理接触Δ_vs_rubberhand")) for row in rows if math.isfinite(finite(row.get("物理接触Δ_vs_rubberhand")))],
                default=math.nan,
            ),
            "source/method_id": method_id,
            "说明": note,
        }
        for key in CORE_METRICS:
            item[f"{key}_mean"] = mean([row.get(key) for row in rows])
        out.append(item)
    return out


def xlsx_value(value: Any) -> Any:
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float) and not math.isfinite(value):
        return ""
    return value


def append_table(
    wb: Workbook,
    title: str,
    rows: list[dict[str, Any]],
    columns: list[tuple[str, str, str]],
    fill: str,
) -> None:
    if len(wb.sheetnames) == 1 and wb.active.max_row == 1 and wb.active["A1"].value is None:
        ws = wb.active
        ws.delete_rows(1)
    else:
        ws = wb.create_sheet(title)
    ws.title = title
    ws.append([header for header, _, _ in columns])
    for row in rows:
        ws.append([xlsx_value(row.get(key, "")) for _, key, _ in columns])

    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, (_, _, fmt) in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        for cell in ws[letter][1:]:
            if fmt == "pct":
                cell.number_format = "0.0%"
            elif fmt == "pct_delta":
                cell.number_format = "0.0%;[Red]-0.0%"
            elif fmt == "num":
                cell.number_format = "0.00"
        width = min(max(len(str(cell.value or "")) for cell in ws[letter]) + 2, 42)
        ws.column_dimensions[letter].width = max(width, 10)


def write_workbook(
    out_dir: Path,
    stage: str,
    summary: list[dict[str, Any]],
    per_case: list[dict[str, Any]],
    artifacts: list[dict[str, Any]],
) -> Path:
    wb = Workbook()
    main_cols = [
        ("方法", "方法", "text"),
        ("case数", "case数", "int"),
        ("通过case", "通过case", "int"),
        ("失败case", "失败case", "text"),
        ("接触退化case", "接触退化case", "text"),
        ("物理接触Δ最差", "物理接触Δ最差", "pct_delta"),
        ("物理接触(raw)", f"{CONTACT_METRIC}_mean", "pct"),
        ("clean3接触", "hand_object_physics_contact_3mm_in_mask_frac_mean", "pct"),
        ("clean5接触", "hand_object_physics_contact_5mm_in_mask_frac_mean", "pct"),
        ("物理穿透3mm", "hand_object_physics_penetration_3mm_frame_frac_mean", "pct"),
        ("几何穿透2mm", "hand_geom_penetration_2mm_frac_mean", "pct"),
        ("release误接触3mm", "hand_object_release_false_contact_3mm_frac_mean", "pct"),
        ("关节误差(°)", "track_joint_err_deg_mean_mean", "num"),
        ("末端位置误差(cm)", "track_eef_pos_err_cm_mean_mean", "num"),
        ("末端朝向误差(°)", "track_eef_ori_err_deg_mean_mean", "num"),
        ("物体位置误差(cm)", "track_obj_pos_err_cm_mean_mean", "num"),
        ("物体朝向误差(°)", "track_obj_ori_err_deg_mean_mean", "num"),
        ("说明", "说明", "text"),
    ]
    append_table(wb, "主表", summary, main_cols, "1F4E79")

    per_case_cols = [
        ("方法", "方法", "text"),
        ("case", "case", "text"),
        ("状态", "状态", "text"),
        ("物理接触Δ_vs_rubberhand", "物理接触Δ_vs_rubberhand", "pct_delta"),
        ("raw接触下限", "raw接触下限", "pct"),
        ("物理接触(raw)", CONTACT_METRIC, "pct"),
        ("clean3接触", "hand_object_physics_contact_3mm_in_mask_frac", "pct"),
        ("clean5接触", "hand_object_physics_contact_5mm_in_mask_frac", "pct"),
        ("物理穿透3mm", "hand_object_physics_penetration_3mm_frame_frac", "pct"),
        ("物理穿透5mm", "hand_object_physics_penetration_5mm_frame_frac", "pct"),
        ("几何穿透2mm", "hand_geom_penetration_2mm_frac", "pct"),
        ("几何穿透5mm", "hand_geom_penetration_5mm_frac", "pct"),
        ("release误接触3mm", "hand_object_release_false_contact_3mm_frac", "pct"),
        ("tracked", "success_tracked", "text"),
        ("fall", "fall_flag", "text"),
        ("Table4完整", "Table4完整", "text"),
        ("关节误差(°)", "track_joint_err_deg_mean", "num"),
        ("末端位置误差(cm)", "track_eef_pos_err_cm_mean", "num"),
        ("末端朝向误差(°)", "track_eef_ori_err_deg_mean", "num"),
        ("物体位置误差(cm)", "track_obj_pos_err_cm_mean", "num"),
        ("物体朝向误差(°)", "track_obj_ori_err_deg_mean", "num"),
        ("source/method_id", "source/method_id", "text"),
    ]
    append_table(wb, "逐case", per_case, per_case_cols, "548235")

    artifact_cols = [
        ("variant", "variant", "text"),
        ("case", "short_case_id", "text"),
        ("artifact_ok", "artifact_ok", "text"),
        ("root_npz", "root_npz_exists", "text"),
        ("trajectory", "trajectory_mjwp_act_exists", "text"),
        ("config", "config_act_exists", "text"),
        ("mp4", "full_mp4_exists", "text"),
        ("min_sdf_ok", "config_surface_band_min_sdf_ok", "text"),
        ("width_ok", "config_surface_band_width_ok", "text"),
        ("sigma_ok", "config_surface_band_sigma_ok", "text"),
        ("score_mode_ok", "config_surface_band_score_mode_ok", "text"),
        ("mask_source_ok", "config_contact_mask_source_ok", "text"),
        ("config_path", "config_path", "text"),
    ]
    append_table(wb, "E163产物检查", artifacts, artifact_cols, "8064A2")

    info_rows = [
        {"项": "baseline", "说明": "所有 delta/status 均相对同 case 的 SPIDER+rubberhand。"},
        {"项": "raw contact hard gate", "说明": "raw hand_object_physics_contact_in_mask_frac 相对 rubberhand 下降超过 0.05 判接触退化。"},
        {"项": "box023 sanity", "说明": f"E163 必须超过 E161 releaseDecay raw contact {BOX023_E161_RELEASE_DECAY_RAW:.4f}；真正 pass 为 >=0.8577。"},
        {"项": "E163 config", "说明": "surface_band_min_sdf=-1mm, width=+3mm, sigma=1.5mm, score_mode=symmetric_abs, mask_source=core4d_3cm。"},
        {"项": "secondary", "说明": "clean3/5、physPen、geomPen、release false 只辅助解释，不能抵消 raw contact fail。"},
        {"项": "stage", "说明": f"当前评测阶段为 {stage}；clean8 使用 cem/full 产物并写到 eval/clean8，不覆盖三 case eval/full。"},
    ]
    append_table(wb, "说明", info_rows, [("项", "项", "text"), ("说明", "说明", "text")], "C65911")

    workbook_name = "E163_narrow_surface_band_clean8_eval.xlsx" if stage == "clean8" else "E163_narrow_surface_band_three_case_eval.xlsx"
    out_path = out_dir / workbook_name
    wb.save(out_path)
    return out_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("stage", nargs="?", default="full", choices=["smoke", "full", "clean8"])
    parser.add_argument("--allow-missing", action="store_true")
    args = parser.parse_args()

    target_cases = target_cases_for_stage(args.stage)
    out_dir = RESULT_ROOT / "eval" / args.stage
    out_dir.mkdir(parents=True, exist_ok=True)

    cfg = EvalConfig()
    e163_rows, artifact_rows, e163_missing = evaluate_e163(args.stage, cfg, target_cases)
    all_rows, ref_missing = load_reference_rows(e163_rows, target_cases)
    per_case = build_per_case(all_rows, target_cases)
    summary = build_summary(per_case)

    write_tsv(out_dir / "e163_method_metrics.tsv", all_rows, sorted({key for row in all_rows for key in row.keys()}))
    write_tsv(out_dir / "e163_case_status.tsv", per_case, sorted({key for row in per_case for key in row.keys()}))
    write_tsv(out_dir / "e163_method_summary.tsv", summary, sorted({key for row in summary for key in row.keys()}))
    write_tsv(out_dir / "e163_artifact_check.tsv", artifact_rows, sorted({key for row in artifact_rows for key in row.keys()}))
    write_tsv(out_dir / "e163_missing.tsv", e163_missing, sorted({key for row in e163_missing for key in row.keys()}) if e163_missing else ["reason"])
    write_tsv(out_dir / "e163_reference_missing.tsv", ref_missing, sorted({key for row in ref_missing for key in row.keys()}) if ref_missing else ["reason"])
    xlsx = write_workbook(out_dir, args.stage, summary, per_case, artifact_rows)

    e163_case_rows = [row for row in per_case if row["方法"] == "E163 narrowSurfaceBand"]
    eval_summary = {
        "stage": args.stage,
        "target_cases": target_cases,
        "e163_metric_rows": len(e163_rows),
        "all_metric_rows": len(all_rows),
        "missing_rows": len(e163_missing),
        "reference_missing_rows": len(ref_missing),
        "e163_pass_cases": sum(1 for row in e163_case_rows if row["状态"] == "通过"),
        "e163_failed_cases": [row["case"] for row in e163_case_rows if row["状态"] != "通过"],
        "e163_contact_failed_cases": [row["case"] for row in e163_case_rows if row["状态"] == "接触退化"],
        "e163_box023_raw_contact": next(
            (row.get(CONTACT_METRIC) for row in e163_case_rows if row["case"] == "box023_person2"),
            math.nan,
        ),
        "xlsx": rel(xlsx),
    }
    write_json(out_dir / "e163_eval_summary.json", eval_summary)

    print(
        "E163 eval: "
        f"stage={args.stage} e163_rows={len(e163_rows)} all_rows={len(all_rows)} "
        f"missing={len(e163_missing)} ref_missing={len(ref_missing)} "
        f"e163_pass={eval_summary['e163_pass_cases']}/{len(target_cases)} xlsx={rel(xlsx)}"
    )
    if e163_missing and not args.allow_missing:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
