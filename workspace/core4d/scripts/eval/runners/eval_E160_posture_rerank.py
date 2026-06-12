#!/usr/bin/env python3
"""Evaluate E160 gateA + surfaceBand-A2 + postureRerankA benchmark."""

from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
import sys
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from eval.core.core_metrics import (  # noqa: E402
    EVAL_METRIC_STANDARD_ID,
    EvalConfig,
    STANDARD_DELTA_METRICS,
    STANDARD_MASK_DELTA_METRICS,
    STANDARD_TRACK_DIAG,
)
from eval_E156_clean8_gate_decay import (  # noqa: E402
    evaluate_omniretarget_row,
    evaluate_row,
    finite,
    fmt,
    read_tsv,
    repo_path,
    write_json,
    write_tsv,
)


REPO = Path(__file__).resolve().parents[5]
E156_VARIANTS = REPO / "workspace/core4d/scripts/experiments/E156/variants.tsv"
E159_VARIANTS = REPO / "workspace/core4d/scripts/experiments/E159/variants.tsv"
E160_VARIANTS = REPO / "workspace/core4d/scripts/experiments/E160/variants.tsv"
RESULT_ROOT = REPO / "workspace/core4d/results/E160/posture_rerank"

E160_CASES = [
    "box021_029_p2",
    "box004_083_p2",
    "box023_person2",
]
METHOD_ORDER = [
    "OmniRetarget",
    "spider-rubberhand",
    "+gateA",
    "E155_decay",
    "gateA+surfaceBand-A2",
    "gateA+surfaceBand-A2+postureRerankA",
]
TRACK_DIAG = list(STANDARD_TRACK_DIAG)
DELTA_METRICS = list(STANDARD_DELTA_METRICS)
MASK_DELTA = list(STANDARD_MASK_DELTA_METRICS)
SURFACE_HEALTH_KEYS = {
    "surface_band_rew_mean": "surface_band_rew_mean",
    "surface_band_penalty_mean": "surface_band_penalty_mean",
    "surface_band_sdf_mean": "surface_band_sdf_mean_m",
    "surface_band_score_mean": "surface_band_score_mean",
    "surface_band_penetration_mean": "surface_band_penetration_mean_m",
}
GATE_HEALTH_KEYS = {
    "hand_gate_valid_frac": "hand_gate_valid_frac",
    "gate_fallback_used": "gate_fallback_used",
    "hand_gate_min_sdf_min_m": "hand_gate_min_sdf_min_m",
    "hand_gate_violation_pct": "hand_gate_violation_pct",
}
POSTURE_HEALTH_KEYS = {
    "cem_posture_gate_valid_frac": "posture_gate_valid_frac",
    "cem_posture_gate_selected_valid_frac": "posture_gate_selected_valid_frac",
    "cem_posture_gate_fallback_used": "posture_gate_fallback_used",
    "sample_posture_mean_z_err_mean": "posture_mean_z_err_m",
    "sample_posture_terminal_z_err_mean": "posture_terminal_z_err_m",
    "sample_posture_max_z_drop_mean": "posture_max_z_drop_m",
    "sample_posture_violation_mean": "posture_violation",
}

SUMMARY_KEYS = [
    "success_tracked",
    "track_pelvis_z_err_terminal_m",
    "hand_object_release_false_contact_3mm_frac",
    "hand_object_release_false_contact_5mm_frac",
    "hand_object_physics_contact_3mm_in_mask_frac",
    "hand_object_physics_contact_5mm_in_mask_frac",
    "hand_object_physics_penetration_3mm_frame_frac",
    "hand_object_physics_penetration_5mm_frame_frac",
    "hand_geom_penetration_2mm_frac",
    "hand_geom_penetration_5mm_frac",
    "leg_penetration_frac",
    "obj_err_mean_m",
] + list(GATE_HEALTH_KEYS.values()) + list(SURFACE_HEALTH_KEYS.values()) + list(POSTURE_HEALTH_KEYS.values())

EXTRA_DELTA_METRICS = [
    "hand_object_release_false_contact_3mm_frac",
    "hand_object_release_false_contact_5mm_frac",
]

XLSX_METRICS = [
    ("tracked", "success_tracked_cases", +1),
    ("relF3", "hand_object_release_false_contact_3mm_frac_mean", -1),
    ("inmaskC3", "hand_object_physics_contact_3mm_in_mask_frac_mean", +1),
    ("physPen3", "hand_object_physics_penetration_3mm_frame_frac_mean", -1),
    ("pen2", "hand_geom_penetration_2mm_frac_mean", -1),
    ("legPen", "leg_penetration_frac_mean", -1),
    ("objErr", "obj_err_mean_m_mean", -1),
]


def mean(values: list[Any]) -> float:
    vals = [finite(v) for v in values]
    vals = [v for v in vals if math.isfinite(v)]
    return statistics.fmean(vals) if vals else math.nan


def worst(values: list[Any], high_is_bad: bool = True) -> float:
    vals = [finite(v) for v in values]
    vals = [v for v in vals if math.isfinite(v)]
    if not vals:
        return math.nan
    return max(vals) if high_is_bad else min(vals)


def run_health(qpos_path: Path) -> dict[str, Any]:
    out = {
        key: ""
        for key in (
            list(GATE_HEALTH_KEYS.values())
            + list(SURFACE_HEALTH_KEYS.values())
            + list(POSTURE_HEALTH_KEYS.values())
        )
    }
    if not qpos_path.is_file():
        return out
    data = np.load(qpos_path, allow_pickle=True)
    source_map = {
        "cem_hand_gate_valid_frac": "hand_gate_valid_frac",
        "cem_gate_fallback_used": "gate_fallback_used",
        "cem_hand_gate_min_sdf_min": "hand_gate_min_sdf_min_m",
        "sample_hand_gate_violation_pct_mean": "hand_gate_violation_pct",
    }
    for src, dst in source_map.items():
        if src in data.files:
            arr = np.asarray(data[src], dtype=np.float64)
            arr = arr[np.isfinite(arr)]
            if arr.size:
                out[dst] = float(arr.min()) if "min_sdf_min" in src else float(arr.mean())
    for src, dst in SURFACE_HEALTH_KEYS.items():
        if src in data.files:
            arr = np.asarray(data[src], dtype=np.float64)
            arr = arr[np.isfinite(arr)]
            if arr.size:
                out[dst] = float(arr.mean())
    for src, dst in POSTURE_HEALTH_KEYS.items():
        if src in data.files:
            arr = np.asarray(data[src], dtype=np.float64)
            arr = arr[np.isfinite(arr)]
            if arr.size:
                out[dst] = float(arr.mean())
    return out


def summarize_method(method: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
    out: dict[str, Any] = {
        "method": method,
        "n_cases": len(rows),
        "success_tracked_cases": sum(1 for row in rows if row.get("success_tracked")),
        "fall_cases": sum(1 for row in rows if row.get("fall_flag")),
    }
    high_good = {
        "success_tracked",
        "hand_object_physics_contact_3mm_in_mask_frac",
        "hand_object_physics_contact_5mm_in_mask_frac",
        "hand_gate_valid_frac",
        "surface_band_rew_mean",
        "surface_band_score_mean",
        "posture_gate_valid_frac",
        "posture_gate_selected_valid_frac",
    }
    for key in SUMMARY_KEYS:
        if key == "success_tracked":
            vals = [1.0 if row.get(key) else 0.0 for row in rows]
            out[f"{key}_mean"] = mean(vals)
            out[f"{key}_worst"] = worst(vals, high_is_bad=False)
        else:
            out[f"{key}_mean"] = mean([row.get(key) for row in rows])
            out[f"{key}_worst"] = worst([row.get(key) for row in rows], high_is_bad=key not in high_good)
    return out


def delta_row(case: str, run: dict[str, Any], ref: dict[str, Any], ref_method: str) -> dict[str, Any]:
    row: dict[str, Any] = {
        "short_case_id": case,
        "variant": run["variant"],
        "method": run["method"],
        "reference_method": ref_method,
        "reference_variant": ref["variant"],
        "success_tracked": run.get("success_tracked", False),
        "fall_flag": run.get("fall_flag", False),
    }
    for key in TRACK_DIAG:
        row[key] = run.get(key, math.nan)
    for key in MASK_DELTA:
        row[f"{key}_ref"] = ref.get(key, math.nan)
        row[f"{key}_run"] = run.get(key, math.nan)
        row[f"{key}_delta"] = finite(run.get(key)) - finite(ref.get(key))
    for key in DELTA_METRICS:
        row[f"{key}_ref"] = ref.get(key, math.nan)
        row[f"{key}_run"] = run.get(key, math.nan)
        row[f"{key}_delta"] = finite(run.get(key)) - finite(ref.get(key))
    for key in EXTRA_DELTA_METRICS:
        if f"{key}_delta" in row:
            continue
        row[f"{key}_ref"] = ref.get(key, math.nan)
        row[f"{key}_run"] = run.get(key, math.nan)
        row[f"{key}_delta"] = finite(run.get(key)) - finite(ref.get(key))
    return row


def build_delta_rows(metric_rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    by_case_method = {(row["short_case_id"], row["method"]): row for row in metric_rows}
    vs_omni: list[dict[str, Any]] = []
    vs_baseline: list[dict[str, Any]] = []
    vs_gate: list[dict[str, Any]] = []
    for case in E160_CASES:
        omni = by_case_method.get((case, "OmniRetarget"))
        base = by_case_method.get((case, "spider-rubberhand"))
        gate = by_case_method.get((case, "+gateA"))
        for method in [
            "spider-rubberhand",
            "+gateA",
            "E155_decay",
            "gateA+surfaceBand-A2",
            "gateA+surfaceBand-A2+postureRerankA",
        ]:
            run = by_case_method.get((case, method))
            if omni and run:
                vs_omni.append(delta_row(case, run, omni, "OmniRetarget"))
            if base and run and method != "spider-rubberhand":
                vs_baseline.append(delta_row(case, run, base, "spider-rubberhand"))
            if gate and run and method in {"E155_decay", "gateA+surfaceBand-A2", "gateA+surfaceBand-A2+postureRerankA"}:
                vs_gate.append(delta_row(case, run, gate, "+gateA"))
    return vs_omni, vs_baseline, vs_gate


def success_summary(method_summary: list[dict[str, Any]], delta_vs_gate: list[dict[str, Any]]) -> dict[str, Any]:
    rerank = next((row for row in method_summary if row["method"] == "gateA+surfaceBand-A2+postureRerankA"), None)
    a2 = next((row for row in method_summary if row["method"] == "gateA+surfaceBand-A2"), None)
    if not rerank or not a2:
        return {"promote_to_clean6_probe": False, "reason": "missing E159 A2 or E160 summary"}
    rows_by_method_case = {
        (row["method"], row["short_case_id"]): row
        for row in delta_vs_gate
        if row["method"] in {"gateA+surfaceBand-A2", "gateA+surfaceBand-A2+postureRerankA"}
    }
    rerank_rows = [row for row in delta_vs_gate if row["method"] == "gateA+surfaceBand-A2+postureRerankA"]
    rel3_delta = mean([row.get("hand_object_release_false_contact_3mm_frac_delta") for row in rerank_rows])
    inmask_delta = mean([row.get("hand_object_physics_contact_3mm_in_mask_frac_delta") for row in rerank_rows])
    phys_pen3_delta = mean([row.get("hand_object_physics_penetration_3mm_frame_frac_delta") for row in rerank_rows])
    pen2_delta = mean([row.get("hand_geom_penetration_2mm_frac_delta") for row in rerank_rows])

    def mean_run_delta_vs_a2(metric: str) -> float:
        vals = []
        for case in E160_CASES:
            run = rows_by_method_case.get(("gateA+surfaceBand-A2+postureRerankA", case))
            ref = rows_by_method_case.get(("gateA+surfaceBand-A2", case))
            if run and ref:
                vals.append(finite(run.get(f"{metric}_run")) - finite(ref.get(f"{metric}_run")))
        return mean(vals)

    delta_vs_a2 = {
        "inmaskC3": mean_run_delta_vs_a2("hand_object_physics_contact_3mm_in_mask_frac"),
        "physPen3": mean_run_delta_vs_a2("hand_object_physics_penetration_3mm_frame_frac"),
        "pen2": mean_run_delta_vs_a2("hand_geom_penetration_2mm_frac"),
        "releaseF3": mean_run_delta_vs_a2("hand_object_release_false_contact_3mm_frac"),
    }
    box021 = rows_by_method_case.get(("gateA+surfaceBand-A2+postureRerankA", "box021_029_p2"), {})
    box004_run = rows_by_method_case.get(("gateA+surfaceBand-A2+postureRerankA", "box004_083_p2"), {})
    box004_a2 = rows_by_method_case.get(("gateA+surfaceBand-A2", "box004_083_p2"), {})
    box023 = rows_by_method_case.get(("gateA+surfaceBand-A2+postureRerankA", "box023_person2"), {})
    box004_rel_delta = finite(box004_run.get("hand_object_release_false_contact_3mm_frac_run")) - finite(
        box004_a2.get("hand_object_release_false_contact_3mm_frac_run")
    )
    checks = {
        "C1_box021_tracked": bool(box021.get("success_tracked")),
        "C1_box021_pz_term_le_0_08": finite(box021.get("track_pelvis_z_err_terminal_m")) <= 0.08,
        "C2_mean_inmaskC3_vs_A2_drop_le_0_10": delta_vs_a2["inmaskC3"] >= -0.10,
        "C3_mean_physPen3_vs_A2_le_0_03": delta_vs_a2["physPen3"] <= 0.03,
        "C3_mean_pen2_vs_A2_le_0_03": delta_vs_a2["pen2"] <= 0.03,
        "C4_box004_releaseF3_not_worse_A2": box004_rel_delta <= 0.0,
        "C5_box023_tracked": bool(box023.get("success_tracked")),
    }
    return {
        "promote_to_clean6_probe": all(checks.values()),
        "checks": checks,
        "mean_delta_vs_gateA": {
            "release_false3": rel3_delta,
            "inmaskC3": inmask_delta,
            "physPen3": phys_pen3_delta,
            "pen2": pen2_delta,
        },
        "mean_delta_vs_E159_A2": delta_vs_a2,
        "box004_releaseF3_delta_vs_E159_A2": box004_rel_delta,
    }


def rank_styles(values: list[float], direction: int) -> tuple[int | None, int | None]:
    indexed = [(idx, val) for idx, val in enumerate(values) if math.isfinite(val)]
    if not indexed:
        return None, None
    indexed.sort(key=lambda item: item[1], reverse=direction > 0)
    best = indexed[0][0]
    second = indexed[1][0] if len(indexed) > 1 else None
    return best, second


def write_xlsx(
    path: Path,
    method_summary: list[dict[str, Any]],
    metric_rows: list[dict[str, Any]],
    delta_vs_omni: list[dict[str, Any]],
    delta_vs_baseline: list[dict[str, Any]],
    delta_vs_gate: list[dict[str, Any]],
    success: dict[str, Any],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "method_summary"
    headers = ["method", "n_cases"] + [label for label, _, _ in XLSX_METRICS]
    rows = sorted(method_summary, key=lambda row: METHOD_ORDER.index(row["method"]) if row["method"] in METHOD_ORDER else 99)
    ws.append(headers)
    for row in rows:
        ws.append([row.get("method"), row.get("n_cases")] + [row.get(key) for _, key, _ in XLSX_METRICS])
    for col in range(1, len(headers) + 1):
        ws.cell(1, col).font = Font(bold=True)
        ws.cell(1, col).fill = PatternFill("solid", fgColor="D9EAF7")
    for metric_idx, (_, key, direction) in enumerate(XLSX_METRICS, start=3):
        vals = [finite(row.get(key)) for row in rows]
        best, second = rank_styles(vals, direction)
        if best is not None:
            ws.cell(best + 2, metric_idx).font = Font(bold=True, color="000000")
        if second is not None:
            ws.cell(second + 2, metric_idx).font = Font(underline="single")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"

    for title, rows2 in (
        ("per_case_metrics", metric_rows),
        ("delta_vs_omniretarget", delta_vs_omni),
        ("delta_vs_baseline", delta_vs_baseline),
        ("delta_vs_gateA", delta_vs_gate),
    ):
        sheet = wb.create_sheet(title)
        fields = sorted({key for row in rows2 for key in row.keys()})
        front = ["short_case_id", "variant", "method", "reference_method", "success_tracked"]
        fields = [field for field in front if field in fields] + [field for field in fields if field not in front]
        sheet.append(fields)
        for row in rows2:
            sheet.append([row.get(field, "") for field in fields])
        for col in range(1, len(fields) + 1):
            sheet.cell(1, col).font = Font(bold=True)
            sheet.cell(1, col).fill = PatternFill("solid", fgColor="D9EAD3")
        sheet.freeze_panes = "A2"

    meta = wb.create_sheet("summary")
    meta.append(["metric_standard_id", EVAL_METRIC_STANDARD_ID])
    meta.append(["promote_to_clean6_probe", success.get("promote_to_clean6_probe")])
    for key, val in success.get("checks", {}).items():
        meta.append([key, val])
    for key, val in success.get("mean_delta_vs_gateA", {}).items():
        meta.append([f"mean_delta_vs_gateA.{key}", val])
    for key, val in success.get("mean_delta_vs_E159_A2", {}).items():
        meta.append([f"mean_delta_vs_E159_A2.{key}", val])
    meta.append(
        [
            "box004_releaseF3_delta_vs_E159_A2",
            success.get("box004_releaseF3_delta_vs_E159_A2"),
        ]
    )
    wb.save(path)


def e156_case_rows() -> list[dict[str, str]]:
    rows = read_tsv(E156_VARIANTS)
    return [row for row in rows if row["short_case_id"] in E160_CASES]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("stage", nargs="?", default="full", choices=["smoke", "full"])
    parser.add_argument("--allow-missing", action="store_true")
    args = parser.parse_args()

    cfg = EvalConfig()
    eval_dir = RESULT_ROOT / "eval" / args.stage
    e156_rows = e156_case_rows()
    e159_rows = [row for row in read_tsv(E159_VARIANTS) if row["short_case_id"] in E160_CASES]
    e160_rows = read_tsv(E160_VARIANTS)
    rows = e156_rows + e159_rows + e160_rows
    metric_rows: list[dict[str, Any]] = []
    missing: list[str] = []

    baseline_rows = [row for row in e156_rows if row["method"] == "spider-rubberhand"]
    for row in baseline_rows:
        metrics = evaluate_omniretarget_row(row, eval_dir, cfg)
        if metrics is None:
            missing.append(f"E160_{row['short_case_id']}_omniretarget")
            continue
        metric_rows.append(metrics)

    for row in rows:
        qpos = repo_path(row["outdir_npz"])
        metrics = evaluate_row(row, qpos, cfg)
        if metrics is None:
            missing.append(row["variant"])
            continue
        metrics.update(run_health(qpos))
        metric_rows.append(metrics)

    method_summary = [summarize_method(method, [row for row in metric_rows if row["method"] == method]) for method in METHOD_ORDER]
    delta_vs_omni, delta_vs_baseline, delta_vs_gate = build_delta_rows(metric_rows)
    success = success_summary(method_summary, delta_vs_gate)

    metric_fields = (
        [
            "short_case_id",
            "variant",
            "method",
            "method_group",
            "run_status",
            "source_exp",
            "split",
            "qpos_frames",
            "success_tracked",
            "pelvis_min_m",
            "fall_flag",
            "hand_geom_near_5cm_frac",
            "hand_geom_near_10cm_frac",
            "hand_geom_penetration_2mm_frac",
            "hand_geom_penetration_5mm_frac",
            "hand_object_physics_contact_3mm_frac",
            "hand_object_physics_contact_5mm_frac",
            "hand_object_physics_contact_3mm_in_mask_frac",
            "hand_object_physics_contact_5mm_in_mask_frac",
            "hand_object_physics_penetration_3mm_frame_frac",
            "hand_object_physics_penetration_5mm_frame_frac",
            "hand_object_release_false_contact_3mm_frac",
            "hand_object_release_false_contact_5mm_frac",
            "leg_penetration_frac",
            "obj_err_mean_m",
        ]
        + TRACK_DIAG
        + list(GATE_HEALTH_KEYS.values())
        + list(SURFACE_HEALTH_KEYS.values())
        + list(POSTURE_HEALTH_KEYS.values())
        + ["result_npz", "video"]
    )
    delta_fields = (
        [
            "short_case_id",
            "variant",
            "method",
            "reference_method",
            "reference_variant",
            "success_tracked",
            "fall_flag",
        ]
        + TRACK_DIAG
        + [f"{key}_{suffix}" for key in MASK_DELTA for suffix in ("ref", "run", "delta")]
        + [f"{key}_{suffix}" for key in DELTA_METRICS for suffix in ("ref", "run", "delta")]
        + [f"{key}_{suffix}" for key in EXTRA_DELTA_METRICS for suffix in ("ref", "run", "delta")]
    )
    summary_fields = sorted({key for row in method_summary for key in row.keys()})
    summary_fields = ["method", "n_cases", "success_tracked_cases", "fall_cases"] + [
        field for field in summary_fields if field not in {"method", "n_cases", "success_tracked_cases", "fall_cases"}
    ]

    write_tsv(eval_dir / "e160_method_metrics.tsv", metric_rows, metric_fields)
    write_tsv(eval_dir / "e160_method_summary.tsv", method_summary, summary_fields)
    write_tsv(eval_dir / "e160_delta_vs_omniretarget.tsv", delta_vs_omni, delta_fields)
    write_tsv(eval_dir / "e160_delta_vs_spider_rubberhand.tsv", delta_vs_baseline, delta_fields)
    write_tsv(eval_dir / "e160_delta_vs_gateA.tsv", delta_vs_gate, delta_fields)
    write_json(
        eval_dir / "e160_eval_summary.json",
        {
            "metric_standard_id": EVAL_METRIC_STANDARD_ID,
            "stage": args.stage,
            "metric_rows": len(metric_rows),
            "method_rows": len(method_summary),
            "missing": missing,
            "allow_missing": bool(args.allow_missing),
            "success": success,
        },
    )
    write_xlsx(
        eval_dir / "E160_posture_rerank_3case.xlsx",
        method_summary,
        metric_rows,
        delta_vs_omni,
        delta_vs_baseline,
        delta_vs_gate,
        success,
    )
    print(
        f"E160 eval: metric_rows={len(metric_rows)} methods={len(method_summary)} "
        f"delta_vs_gate={len(delta_vs_gate)} missing={len(missing)}"
    )
    if missing:
        print("MISSING:", missing)
        if not args.allow_missing:
            raise SystemExit(1)


if __name__ == "__main__":
    main()
