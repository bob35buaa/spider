#!/usr/bin/env python3
"""Evaluate E156 clean8 spider-rubberhand / +gateA / E155_decay benchmark."""

from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
import sys
from pathlib import Path
from typing import Any

import mujoco
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.spatial.transform import Rotation as R

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from eval.core.core_metrics import (  # noqa: E402
    EVAL_METRIC_STANDARD_ID,
    EvalConfig,
    STANDARD_DELTA_METRICS,
    STANDARD_MASK_DELTA_METRICS,
    STANDARD_TRACK_DIAG,
    contact_mask_for_case,
    evaluate_sequence,
    kin_ref_for_scene,
    person_idx_from_case,
)


REPO = Path(__file__).resolve().parents[5]
VARIANTS_TSV = REPO / "workspace/core4d/scripts/experiments/E156/variants.tsv"
RESULT_ROOT = REPO / "workspace/core4d/results/E156/clean8_gate_decay"
GATE_HEALTH_KEYS = {
    "cem_hand_gate_valid_frac": ("mean", "hand_gate_valid_frac"),
    "cem_gate_fallback_used": ("mean", "gate_fallback_used"),
    "cem_hand_gate_min_sdf_min": ("min", "hand_gate_min_sdf_min_m"),
    "sample_hand_gate_violation_pct_mean": ("mean", "hand_gate_violation_pct"),
}

METHOD_ORDER = ["OmniRetarget", "spider-rubberhand", "+gateA", "E155_decay"]
TRACK_DIAG = list(STANDARD_TRACK_DIAG)
DELTA_METRICS = list(STANDARD_DELTA_METRICS)
MASK_DELTA = list(STANDARD_MASK_DELTA_METRICS)

SUMMARY_KEYS = [
    "success_tracked",
    "track_pelvis_z_err_terminal_m",
    "hand_object_release_false_contact_3mm_frac",
    "hand_object_release_false_contact_5mm_frac",
    "hand_object_physics_contact_3mm_in_mask_frac",
    "hand_object_physics_contact_5mm_in_mask_frac",
    "hand_object_physics_penetration_3mm_frame_frac",
    "hand_object_physics_penetration_5mm_frame_frac",
    "hand_geom_penetration_2mm_frac",
    "hand_geom_penetration_5mm_frac",
    "leg_penetration_frac",
    "obj_err_mean_m",
    "hand_gate_valid_frac",
    "gate_fallback_used",
]

XLSX_METRICS = [
    ("tracked", "success_tracked_cases", +1),
    ("relF3", "hand_object_release_false_contact_3mm_frac_mean", -1),
    ("relF5", "hand_object_release_false_contact_5mm_frac_mean", -1),
    ("inmaskC3", "hand_object_physics_contact_3mm_in_mask_frac_mean", +1),
    ("inmaskC5", "hand_object_physics_contact_5mm_in_mask_frac_mean", +1),
    ("physPen3", "hand_object_physics_penetration_3mm_frame_frac_mean", -1),
    ("physPen5", "hand_object_physics_penetration_5mm_frame_frac_mean", -1),
    ("pen2", "hand_geom_penetration_2mm_frac_mean", -1),
    ("pen5", "hand_geom_penetration_5mm_frac_mean", -1),
    ("legPen", "leg_penetration_frac_mean", -1),
    ("objErr", "obj_err_mean_m_mean", -1),
]


def repo_path(path: str | Path) -> Path:
    p = Path(path)
    return p if p.is_absolute() else REPO / p


def rel(path: Path | str) -> str:
    p = Path(path)
    try:
        return str(p.resolve().relative_to(REPO.resolve()))
    except Exception:
        return str(path)


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def fmt(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return f"{value:.8g}" if math.isfinite(value) else ""
    return str(value)


def finite(value: Any, default: float = math.nan) -> float:
    try:
        out = float(value)
    except (TypeError, ValueError):
        return default
    return out if math.isfinite(out) else default


def write_tsv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: fmt(row.get(field, "")) for field in fields})


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def stage_qpos(row: dict[str, str], stage: str) -> Path:
    if stage == "full":
        return repo_path(row["outdir_npz"])
    return RESULT_ROOT / "cem" / stage / f"{row['variant']}_outdir_{stage}" / "trajectory_mjwp_act.npz"


def stage_video(row: dict[str, str], stage: str) -> Path:
    if stage == "full":
        return repo_path(row["video"])
    return RESULT_ROOT / "cem" / stage / f"{row['variant']}_{stage}.mp4"


def scene_euler_convention(scene_act: Path) -> str:
    meta = scene_act.parent / "scene_act_meta.json"
    if meta.is_file():
        return str(json.loads(meta.read_text(encoding="utf-8")).get("euler_convention", "XYZ"))
    return "XYZ"


def convert_freejoint_to_scene_act(qpos: np.ndarray, scene_act: Path) -> np.ndarray:
    model = mujoco.MjModel.from_xml_path(str(scene_act))
    if qpos.shape[1] == model.nq:
        return qpos.astype(np.float64, copy=True)
    nq_robot = model.nq - 6
    if qpos.shape[1] < nq_robot + 7:
        raise ValueError(f"cannot convert qpos shape={qpos.shape} for scene nq={model.nq}: {scene_act}")

    obj_body = mujoco.mj_name2id(model, mujoco.mjtObj.mjOBJ_BODY, "object")
    if obj_body < 0:
        raise ValueError(f"scene has no object body: {scene_act}")

    obj_pos_world = qpos[:, nq_robot : nq_robot + 3]
    obj_quat_wxyz = qpos[:, nq_robot + 3 : nq_robot + 7]
    body_pos = model.body_pos[obj_body]
    body_quat_wxyz = model.body_quat[obj_body]
    body_quat_xyzw = [body_quat_wxyz[1], body_quat_wxyz[2], body_quat_wxyz[3], body_quat_wxyz[0]]
    r_body = R.from_quat(body_quat_xyzw)

    obj_slide = r_body.inv().apply(obj_pos_world - body_pos[np.newaxis, :])
    obj_quat_xyzw = np.column_stack(
        [obj_quat_wxyz[:, 1], obj_quat_wxyz[:, 2], obj_quat_wxyz[:, 3], obj_quat_wxyz[:, 0]]
    )
    obj_euler = (r_body.inv() * R.from_quat(obj_quat_xyzw)).as_euler(scene_euler_convention(scene_act))

    out = np.zeros((qpos.shape[0], model.nq), dtype=np.float64)
    out[:, :nq_robot] = qpos[:, :nq_robot]
    out[:, nq_robot : nq_robot + 3] = obj_slide
    out[:, nq_robot + 3 : nq_robot + 6] = obj_euler
    return out


def omniretarget_qpos(row: dict[str, str], eval_dir: Path) -> Path:
    trajectory = repo_path(row["trajectory"])
    scene_act = repo_path(row["base_scene_act"])
    data = np.load(trajectory, allow_pickle=True)
    qpos = np.asarray(data["qpos"], dtype=np.float64)
    if qpos.ndim == 3:
        qpos = qpos[:, 0, :]
    if qpos.ndim != 2:
        raise ValueError(f"unsupported OmniRetarget qpos shape={qpos.shape}: {trajectory}")

    qpos_dir = eval_dir / "omni_converted_qpos"
    qpos_dir.mkdir(parents=True, exist_ok=True)
    qpos_path = qpos_dir / f"E156_{row['short_case_id']}_omniretarget_scene_act_qpos.npz"
    np.savez_compressed(qpos_path, qpos=convert_freejoint_to_scene_act(qpos, scene_act))
    return qpos_path


def gate_health(qpos_path: Path) -> dict[str, Any]:
    out: dict[str, Any] = {dst: "" for _, (_, dst) in GATE_HEALTH_KEYS.items()}
    if not qpos_path.is_file():
        return out
    data = np.load(qpos_path, allow_pickle=True)
    for key, (agg, dst) in GATE_HEALTH_KEYS.items():
        if key not in data.files:
            continue
        arr = np.asarray(data[key], dtype=np.float64)
        arr = arr[np.isfinite(arr)]
        if not arr.size:
            continue
        out[dst] = float(arr.mean()) if agg == "mean" else float(arr.min())
    return out


def add_success_flags(row: dict[str, Any], cfg: EvalConfig) -> None:
    pz_term = finite(row.get("track_pelvis_z_err_terminal_m"), math.inf)
    row["success_tracked"] = bool(
        not bool(row.get("fall_flag"))
        and math.isfinite(pz_term)
        and pz_term <= cfg.track_pelvis_terminal_th_m
    )


def evaluate_row(row: dict[str, str], qpos_path: Path, cfg: EvalConfig) -> dict[str, Any] | None:
    scene = repo_path(row["rubber_scene_act"])
    if not qpos_path.is_file() or not scene.is_file():
        return None
    item = evaluate_sequence(
        row=row,
        method=row["method"],
        hand_collision_variant_id="rubber_hull",
        qpos_path=qpos_path,
        scene_xml=scene,
        config=cfg,
        kin_ref_path=kin_ref_for_scene(scene),
        contact_mask_path=contact_mask_for_case(row["short_case_id"]),
        person_idx=person_idx_from_case(row["short_case_id"]),
    )
    add_success_flags(item, cfg)
    return {
        **item,
        "short_case_id": row["short_case_id"],
        "variant": row["variant"],
        "method": row["method"],
        "method_group": row["method_group"],
        "run_status": row["run_status"],
        "source_exp": row["source_exp"],
        "split": row["split"],
        "result_npz": rel(qpos_path),
        "video": rel(stage_video(row, "full")),
        **gate_health(qpos_path),
    }


def evaluate_omniretarget_row(row: dict[str, str], eval_dir: Path, cfg: EvalConfig) -> dict[str, Any] | None:
    scene = repo_path(row["base_scene_act"])
    trajectory = repo_path(row["trajectory"])
    if not scene.is_file() or not trajectory.is_file():
        return None
    qpos_path = omniretarget_qpos(row, eval_dir)
    eval_row = dict(row)
    eval_row["variant"] = f"E156_{row['short_case_id']}_omniretarget"
    item = evaluate_sequence(
        row=eval_row,
        method="OmniRetarget",
        hand_collision_variant_id="OmniRetarget",
        qpos_path=qpos_path,
        scene_xml=scene,
        config=cfg,
        kin_ref_path=trajectory,
        contact_mask_path=contact_mask_for_case(row["short_case_id"]),
        person_idx=person_idx_from_case(row["short_case_id"]),
    )
    add_success_flags(item, cfg)
    return {
        **item,
        "short_case_id": row["short_case_id"],
        "variant": eval_row["variant"],
        "method": "OmniRetarget",
        "method_group": "omniretarget",
        "run_status": "recomputed_reference",
        "source_exp": "E156",
        "split": "reference",
        "result_npz": rel(qpos_path),
        "video": "",
        **{dst: "" for _, (_, dst) in GATE_HEALTH_KEYS.items()},
    }


def mean(values: list[Any]) -> float:
    vals = [finite(v) for v in values]
    vals = [v for v in vals if math.isfinite(v)]
    return statistics.fmean(vals) if vals else math.nan


def worst(values: list[Any], high_is_bad: bool = True) -> float:
    vals = [finite(v) for v in values]
    vals = [v for v in vals if math.isfinite(v)]
    if not vals:
        return math.nan
    return max(vals) if high_is_bad else min(vals)


def summarize_method(method: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
    out: dict[str, Any] = {
        "method": method,
        "n_cases": len(rows),
        "success_tracked_cases": sum(1 for r in rows if r.get("success_tracked")),
        "fall_cases": sum(1 for r in rows if r.get("fall_flag")),
    }
    high_good = {
        "success_tracked",
        "hand_object_physics_contact_3mm_in_mask_frac",
        "hand_object_physics_contact_5mm_in_mask_frac",
        "hand_gate_valid_frac",
    }
    for key in SUMMARY_KEYS:
        if key == "success_tracked":
            out[f"{key}_mean"] = mean([1.0 if r.get(key) else 0.0 for r in rows])
            out[f"{key}_worst"] = worst([1.0 if r.get(key) else 0.0 for r in rows], high_is_bad=False)
            continue
        out[f"{key}_mean"] = mean([r.get(key) for r in rows])
        out[f"{key}_worst"] = worst([r.get(key) for r in rows], high_is_bad=key not in high_good)
    return out


def build_delta_rows(
    metric_rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    by_case_method = {(r["short_case_id"], r["method"]): r for r in metric_rows}
    vs_omni: list[dict[str, Any]] = []
    vs_baseline: list[dict[str, Any]] = []
    vs_gate: list[dict[str, Any]] = []
    for case in sorted({r["short_case_id"] for r in metric_rows}):
        omni = by_case_method.get((case, "OmniRetarget"))
        base = by_case_method.get((case, "spider-rubberhand"))
        gate = by_case_method.get((case, "+gateA"))
        for method in ("spider-rubberhand", "+gateA", "E155_decay"):
            run = by_case_method.get((case, method))
            if omni and run:
                vs_omni.append(delta_row(case, run, omni, "OmniRetarget"))
            if base and run and method != "spider-rubberhand":
                vs_baseline.append(delta_row(case, run, base, "spider-rubberhand"))
            if gate and run and method == "E155_decay":
                vs_gate.append(delta_row(case, run, gate, "+gateA"))
    return vs_omni, vs_baseline, vs_gate


def delta_row(case: str, run: dict[str, Any], ref: dict[str, Any], ref_method: str) -> dict[str, Any]:
    row: dict[str, Any] = {
        "short_case_id": case,
        "variant": run["variant"],
        "method": run["method"],
        "reference_method": ref_method,
        "reference_variant": ref["variant"],
        "success_tracked": run.get("success_tracked", False),
        "fall_flag": run.get("fall_flag", False),
    }
    for key in TRACK_DIAG:
        row[key] = run.get(key, math.nan)
    for key in MASK_DELTA:
        row[f"{key}_ref"] = ref.get(key, math.nan)
        row[f"{key}_run"] = run.get(key, math.nan)
        row[f"{key}_delta"] = finite(run.get(key)) - finite(ref.get(key))
    for key in DELTA_METRICS:
        row[f"{key}_ref"] = ref.get(key, math.nan)
        row[f"{key}_run"] = run.get(key, math.nan)
        row[f"{key}_delta"] = finite(run.get(key)) - finite(ref.get(key))
    return row


def promotion_summary(method_summary: list[dict[str, Any]], delta_vs_gate: list[dict[str, Any]]) -> dict[str, Any]:
    decay = next((r for r in method_summary if r["method"] == "E155_decay"), None)
    gate = next((r for r in method_summary if r["method"] == "+gateA"), None)
    if not decay:
        return {"promote_decay": False, "reason": "missing E155_decay summary"}
    if not gate:
        return {"promote_decay": False, "reason": "missing +gateA summary"}
    rel3_delta = finite(decay.get("hand_object_release_false_contact_3mm_frac_mean")) - finite(
        gate.get("hand_object_release_false_contact_3mm_frac_mean")
    )
    inmask_delta = mean([r.get("hand_object_physics_contact_3mm_in_mask_frac_delta") for r in delta_vs_gate])
    pen3_delta = mean([r.get("hand_object_physics_penetration_3mm_frame_frac_delta") for r in delta_vs_gate])
    checks = {
        "success_tracked_ge_7": decay["success_tracked_cases"] >= 7,
        "release_false_3mm_delta_le_neg_0_05": rel3_delta <= -0.05,
        "inmaskC3_delta_ge_neg_0_10": inmask_delta >= -0.10,
        "phys_pen3_delta_le_0_03": pen3_delta <= 0.03,
    }
    return {
        "promote_decay": all(checks.values()),
        "checks": checks,
        "mean_delta_vs_gateA": {
            "release_false_3mm": rel3_delta,
            "inmaskC3": inmask_delta,
            "phys_pen3": pen3_delta,
        },
    }


def rank_styles(values: list[float], direction: int) -> tuple[int | None, int | None]:
    indexed = [(i, v) for i, v in enumerate(values) if math.isfinite(v)]
    if not indexed:
        return None, None
    indexed.sort(key=lambda x: x[1], reverse=direction > 0)
    best = indexed[0][0]
    second = indexed[1][0] if len(indexed) > 1 else None
    return best, second


def write_xlsx(
    path: Path,
    method_summary: list[dict[str, Any]],
    metric_rows: list[dict[str, Any]],
    delta_vs_omni: list[dict[str, Any]],
    delta_vs_baseline: list[dict[str, Any]],
    delta_vs_gate: list[dict[str, Any]],
    promotion: dict[str, Any],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "method_summary"
    headers = ["method", "n_cases"] + [label for label, _, _ in XLSX_METRICS]
    ws.append(headers)
    rows = sorted(method_summary, key=lambda r: METHOD_ORDER.index(r["method"]) if r["method"] in METHOD_ORDER else 99)
    for row in rows:
        ws.append([row.get("method"), row.get("n_cases")] + [row.get(key) for _, key, _ in XLSX_METRICS])
    for col in range(1, len(headers) + 1):
        ws.cell(1, col).font = Font(bold=True)
        ws.cell(1, col).fill = PatternFill("solid", fgColor="D9EAF7")
    for metric_idx, (_, key, direction) in enumerate(XLSX_METRICS, start=3):
        vals = [finite(r.get(key)) for r in rows]
        best, second = rank_styles(vals, direction)
        if best is not None:
            ws.cell(best + 2, metric_idx).font = Font(bold=True, color="000000")
        if second is not None:
            ws.cell(second + 2, metric_idx).font = Font(underline="single")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"

    for title, rows2 in (
        ("per_case_metrics", metric_rows),
        ("delta_vs_omniretarget", delta_vs_omni),
        ("delta_vs_baseline", delta_vs_baseline),
        ("delta_vs_gateA", delta_vs_gate),
    ):
        sheet = wb.create_sheet(title)
        fields = sorted({k for r in rows2 for k in r.keys()})
        front = ["short_case_id", "variant", "method", "reference_method", "success_tracked"]
        fields = [f for f in front if f in fields] + [f for f in fields if f not in front]
        sheet.append(fields)
        for row in rows2:
            sheet.append([row.get(f, "") for f in fields])
        for col in range(1, len(fields) + 1):
            sheet.cell(1, col).font = Font(bold=True)
            sheet.cell(1, col).fill = PatternFill("solid", fgColor="D9EAD3")
        sheet.freeze_panes = "A2"

    ws_meta = wb.create_sheet("summary")
    ws_meta.append(["metric_standard_id", EVAL_METRIC_STANDARD_ID])
    ws_meta.append(["promote_decay", promotion.get("promote_decay")])
    for key, val in promotion.get("checks", {}).items():
        ws_meta.append([key, val])
    for key, val in promotion.get("mean_delta_vs_gateA", {}).items():
        ws_meta.append([f"mean_delta_vs_gateA.{key}", val])
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("stage", nargs="?", default="full", choices=["smoke", "full"])
    parser.add_argument("--allow-missing", action="store_true")
    args = parser.parse_args()

    cfg = EvalConfig()
    eval_dir = RESULT_ROOT / "eval" / args.stage
    rows = read_tsv(VARIANTS_TSV)
    metric_rows: list[dict[str, Any]] = []
    missing: list[str] = []
    baseline_rows = [row for row in rows if row.get("method") == "spider-rubberhand"]

    for row in baseline_rows:
        metrics = evaluate_omniretarget_row(row, eval_dir, cfg)
        if metrics is None:
            missing.append(f"E156_{row['short_case_id']}_omniretarget")
            continue
        metric_rows.append(metrics)

    for row in rows:
        qpos = stage_qpos(row, args.stage)
        metrics = evaluate_row(row, qpos, cfg)
        if metrics is None:
            missing.append(row["variant"])
            continue
        metric_rows.append(metrics)

    method_summary = [summarize_method(method, [r for r in metric_rows if r["method"] == method]) for method in METHOD_ORDER]
    delta_vs_omni, delta_vs_baseline, delta_vs_gate = build_delta_rows(metric_rows)
    promotion = promotion_summary(method_summary, delta_vs_gate)

    metric_fields = (
        [
            "short_case_id",
            "variant",
            "method",
            "method_group",
            "run_status",
            "source_exp",
            "split",
            "qpos_frames",
            "success_tracked",
            "pelvis_min_m",
            "fall_flag",
            "hand_geom_near_5cm_frac",
            "hand_geom_near_10cm_frac",
            "hand_geom_penetration_2mm_frac",
            "hand_geom_penetration_5mm_frac",
            "hand_object_physics_contact_3mm_frac",
            "hand_object_physics_contact_5mm_frac",
            "hand_object_physics_contact_3mm_in_mask_frac",
            "hand_object_physics_contact_5mm_in_mask_frac",
            "hand_object_physics_penetration_3mm_frame_frac",
            "hand_object_physics_penetration_5mm_frame_frac",
            "hand_object_release_false_contact_3mm_frac",
            "hand_object_release_false_contact_5mm_frac",
            "leg_penetration_frac",
            "obj_err_mean_m",
        ]
        + TRACK_DIAG
        + [dst for _, (_, dst) in GATE_HEALTH_KEYS.items()]
        + ["result_npz", "video"]
    )
    delta_fields = (
        [
            "short_case_id",
            "variant",
            "method",
            "reference_method",
            "reference_variant",
            "success_tracked",
            "fall_flag",
        ]
        + TRACK_DIAG
        + [f"{k}_{s}" for k in MASK_DELTA for s in ("ref", "run", "delta")]
        + [f"{k}_{s}" for k in DELTA_METRICS for s in ("ref", "run", "delta")]
    )
    summary_fields = sorted({k for row in method_summary for k in row.keys()})
    summary_fields = ["method", "n_cases", "success_tracked_cases", "fall_cases"] + [
        f for f in summary_fields if f not in {"method", "n_cases", "success_tracked_cases", "fall_cases"}
    ]

    write_tsv(eval_dir / "e156_method_metrics.tsv", metric_rows, metric_fields)
    write_tsv(eval_dir / "e156_method_summary.tsv", method_summary, summary_fields)
    write_tsv(eval_dir / "e156_delta_vs_omniretarget.tsv", delta_vs_omni, delta_fields)
    write_tsv(eval_dir / "e156_delta_vs_spider_rubberhand.tsv", delta_vs_baseline, delta_fields)
    write_tsv(eval_dir / "e156_delta_vs_gateA.tsv", delta_vs_gate, delta_fields)
    write_json(
        eval_dir / "e156_eval_summary.json",
        {
            "metric_standard_id": EVAL_METRIC_STANDARD_ID,
            "stage": args.stage,
            "metric_rows": len(metric_rows),
            "method_rows": len(method_summary),
            "delta_vs_omniretarget_rows": len(delta_vs_omni),
            "delta_vs_baseline_rows": len(delta_vs_baseline),
            "delta_vs_gate_rows": len(delta_vs_gate),
            "missing": missing,
            "allow_missing": bool(args.allow_missing),
            "promotion": promotion,
        },
    )
    write_xlsx(
        eval_dir / "E156_clean8_gate_decay_benchmark.xlsx",
        method_summary,
        metric_rows,
        delta_vs_omni,
        delta_vs_baseline,
        delta_vs_gate,
        promotion,
    )
    print(
        f"E156 eval: metric_rows={len(metric_rows)} methods={len(method_summary)} "
        f"delta_vs_omni={len(delta_vs_omni)} delta_vs_baseline={len(delta_vs_baseline)} "
        f"delta_vs_gate={len(delta_vs_gate)} missing={len(missing)}"
    )
    if missing:
        print("MISSING:", missing)
        if not args.allow_missing:
            raise SystemExit(1)


if __name__ == "__main__":
    main()
