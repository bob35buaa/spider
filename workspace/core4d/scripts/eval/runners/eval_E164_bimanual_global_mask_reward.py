#!/usr/bin/env python3
"""Evaluate E164 with E164 global masks for all compared methods."""

from __future__ import annotations

import argparse
import csv
import json
import math
import sys
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    import yaml
except Exception:  # pragma: no cover
    yaml = None

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from eval.core.core_metrics import (  # noqa: E402
    EvalConfig,
    evaluate_sequence,
    kin_ref_for_scene,
)


REPO = Path(__file__).resolve().parents[5]
E156_VARIANTS = REPO / "workspace/core4d/scripts/experiments/E156/variants.tsv"
E161_VARIANTS = REPO / "workspace/core4d/scripts/experiments/E161/variants.tsv"
E163_VARIANTS = REPO / "workspace/core4d/scripts/experiments/E163/variants.tsv"
E164_VARIANTS = REPO / "workspace/core4d/scripts/experiments/E164/variants.tsv"
RESULT_ROOT = REPO / "workspace/core4d/results/E164/bimanual_global_mask_reward"

TARGET_CASES = ["box023_person2", "box004_082_p1", "box026_139_p1"]
CONTACT_METRIC = "hand_object_physics_contact_in_mask_frac"
CONTACT_DROP_FAIL_TH = -0.05

TABLE4_TRACKING = [
    "track_joint_err_deg_mean",
    "track_eef_pos_err_cm_mean",
    "track_eef_ori_err_deg_mean",
    "track_root_pos_err_cm_mean",
    "track_root_ori_err_deg_mean",
    "track_obj_pos_err_cm_mean",
    "track_obj_ori_err_deg_mean",
]

CORE_METRICS = [
    CONTACT_METRIC,
    "hand_object_physics_contact_in_rl_mask_frac",
    "rl_object_contact_ref_frac",
    "rl_object_contact_filled_frame_count",
    "hand_object_physics_contact_3mm_in_mask_frac",
    "hand_object_physics_contact_5mm_in_mask_frac",
    "hand_object_physics_penetration_3mm_frame_frac",
    "hand_object_physics_penetration_5mm_frame_frac",
    "hand_geom_penetration_2mm_frac",
    "hand_geom_penetration_5mm_frac",
    "hand_geom_penetration_2mm_in_mask_frac",
    "hand_geom_penetration_5mm_in_mask_frac",
    "hand_object_release_false_contact_3mm_frac",
    "hand_object_release_false_contact_5mm_frac",
    *TABLE4_TRACKING,
]

METHOD_ORDER = [
    ("rubberhand", "SPIDER+rubberhand", "E156", "spider-rubberhand"),
    ("e161_release", "E161 releaseDecay", "E161", "releaseDecay"),
    ("e163_narrow", "E163 narrowSurfaceBand", "E163", "narrowSurfaceBand"),
    ("e164", "E164 bimanualGlobalMaskReward", "E164", "bimanualGlobalMaskReward"),
]

HEALTH_KEYS = {
    "cem_hand_gate_valid_frac": ("mean", "hand_gate_valid_frac"),
    "cem_gate_fallback_used": ("mean", "gate_fallback_used"),
    "cem_hand_gate_min_sdf_min": ("min", "hand_gate_min_sdf_min_m"),
    "sample_hand_gate_violation_pct_mean": ("mean", "hand_gate_violation_pct"),
    "surface_band_rew_mean": ("mean", "surface_band_rew_mean"),
    "surface_band_penalty_mean": ("mean", "surface_band_penalty_mean"),
    "surface_band_gate_mean": ("mean", "surface_band_gate_mean"),
    "surface_band_decay_factor_mean": ("mean", "surface_band_decay_factor_mean"),
    "surface_band_sdf_mean": ("mean", "surface_band_sdf_mean_m"),
    "surface_band_left_sdf_mean": ("mean", "surface_band_left_sdf_mean_m"),
    "surface_band_right_sdf_mean": ("mean", "surface_band_right_sdf_mean_m"),
    "surface_band_bimanual_gate_mean": ("mean", "surface_band_bimanual_gate_mean"),
    "surface_band_bimanual_score_mean": ("mean", "surface_band_bimanual_score_mean"),
    "contact_hdmi_bimanual_gate_mean": ("mean", "contact_hdmi_bimanual_gate_mean"),
    "contact_hdmi_bimanual_score_mean": ("mean", "contact_hdmi_bimanual_score_mean"),
    "cem_posture_gate_valid_frac": ("mean", "posture_gate_valid_frac"),
    "cem_posture_gate_fallback_used": ("mean", "posture_gate_fallback_used"),
}


def rel(path: str | Path) -> str:
    p = Path(path)
    try:
        return str(p.resolve().relative_to(REPO.resolve()))
    except Exception:
        return str(path)


def repo_path(path: str | Path) -> Path:
    p = Path(path)
    return p if p.is_absolute() else REPO / p


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def fmt(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return f"{value:.8g}" if math.isfinite(value) else ""
    return str(value)


def finite(value: Any, default: float = math.nan) -> float:
    try:
        out = float(value)
    except (TypeError, ValueError):
        return default
    return out if math.isfinite(out) else default


def write_tsv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: fmt(row.get(field, "")) for field in fields})


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def bool_text(value: Any) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "pass"}


def mean(values: list[Any]) -> float:
    vals = [finite(v) for v in values]
    vals = [v for v in vals if math.isfinite(v)]
    return sum(vals) / len(vals) if vals else math.nan


def load_yaml(path: Path) -> dict[str, Any]:
    if yaml is None or not path.is_file():
        return {}
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    return data if isinstance(data, dict) else {}


def close_enough(value: Any, target: float, tol: float = 1e-7) -> bool:
    try:
        return abs(float(value) - target) <= tol
    except Exception:
        return False


def e164_stage_paths(row: dict[str, str], stage: str) -> tuple[Path, Path, Path, Path]:
    variant = row["variant"]
    root = RESULT_ROOT / "cem" / stage
    outdir = root / f"{variant}_outdir_{stage}"
    return (
        root / f"{variant}.npz",
        outdir / "trajectory_mjwp_act.npz",
        outdir / "config_act.yaml",
        root / f"{variant}_{stage}.mp4",
    )


def qpos_path_for(row: dict[str, str], source: str, stage: str) -> Path:
    if source == "E164":
        root_npz, outdir_npz, _, _ = e164_stage_paths(row, stage)
        return outdir_npz if outdir_npz.is_file() else root_npz
    outdir_npz = repo_path(row.get("outdir_npz", ""))
    root_npz = repo_path(row.get("result_npz", ""))
    return outdir_npz if outdir_npz.is_file() else root_npz


def video_path_for(row: dict[str, str], source: str, stage: str) -> Path:
    if source == "E164":
        return e164_stage_paths(row, stage)[3]
    return repo_path(row.get("video", ""))


def run_health(qpos_path: Path) -> dict[str, Any]:
    out: dict[str, Any] = {dst: "" for _, (_, dst) in HEALTH_KEYS.items()}
    if not qpos_path.is_file():
        return out
    data = np.load(qpos_path, allow_pickle=True)
    for key, (agg, dst) in HEALTH_KEYS.items():
        if key not in data.files:
            continue
        arr = np.asarray(data[key], dtype=np.float64)
        arr = arr[np.isfinite(arr)]
        if arr.size:
            out[dst] = float(arr.min()) if agg == "min" else float(arr.mean())
    return out


def add_success_flags(row: dict[str, Any], cfg: EvalConfig) -> None:
    pz_term = finite(row.get("track_pelvis_z_err_terminal_m"), math.inf)
    row["success_tracked"] = bool(
        not bool(row.get("fall_flag"))
        and math.isfinite(pz_term)
        and pz_term <= cfg.track_pelvis_terminal_th_m
    )


def artifact_check(row: dict[str, str], stage: str) -> dict[str, Any]:
    root_npz, outdir_npz, config_act, video = e164_stage_paths(row, stage)
    cfg = load_yaml(config_act)
    qpos_path = outdir_npz if outdir_npz.is_file() else root_npz
    keys: set[str] = set()
    if qpos_path.is_file():
        with np.load(qpos_path, allow_pickle=True) as data:
            keys = set(data.files)
    checks = {
        "variant": row["variant"],
        "short_case_id": row["short_case_id"],
        "root_npz_exists": root_npz.is_file(),
        "trajectory_mjwp_act_exists": outdir_npz.is_file(),
        "config_act_exists": config_act.is_file(),
        "full_mp4_exists": video.is_file(),
        "e164_mask_exists": repo_path(row["mask_path"]).is_file(),
        "config_surface_band_min_sdf_ok": close_enough(cfg.get("surface_band_min_sdf_m"), -0.001),
        "config_surface_band_width_ok": close_enough(cfg.get("surface_band_width_m"), 0.003),
        "config_surface_band_sigma_ok": close_enough(cfg.get("surface_band_sigma"), 0.0015),
        "config_surface_band_score_mode_ok": cfg.get("surface_band_score_mode") == "symmetric_abs",
        "config_surface_bimanual_ok": bool(cfg.get("surface_band_bimanual_required")) is True,
        "config_contact_bimanual_ok": bool(cfg.get("contact_hdmi_bimanual_required")) is True,
        "config_contact_mask_axis_ok": cfg.get("contact_hdmi_mask_time_axis") == "spider",
        "diag_surface_bimanual_exists": "surface_band_bimanual_gate_mean" in keys,
        "diag_contact_bimanual_exists": "contact_hdmi_bimanual_gate_mean" in keys,
        "config_path": rel(config_act),
        "root_npz": rel(root_npz),
        "trajectory_mjwp_act": rel(outdir_npz),
        "video": rel(video),
        "mask_path": row["mask_path"],
    }
    bool_keys = [key for key in checks if key.endswith("_exists") or key.endswith("_ok")]
    checks["artifact_ok"] = all(bool(checks[key]) for key in bool_keys)
    return checks


def source_indexes() -> dict[str, dict[tuple[str, str], dict[str, str]]]:
    e156 = {
        (row["short_case_id"], row["method"]): row
        for row in read_tsv(E156_VARIANTS)
        if row.get("short_case_id") in TARGET_CASES
    }
    e161 = {
        (row["short_case_id"], row["method_group"]): row
        for row in read_tsv(E161_VARIANTS)
        if row.get("short_case_id") in TARGET_CASES
    }
    e163 = {
        (row["short_case_id"], row["method_group"]): row
        for row in read_tsv(E163_VARIANTS)
        if row.get("short_case_id") in TARGET_CASES
    }
    e164 = {
        (row["short_case_id"], row["method_group"]): row
        for row in read_tsv(E164_VARIANTS)
        if row.get("short_case_id") in TARGET_CASES
    }
    return {"E156": e156, "E161": e161, "E163": e163, "E164": e164}


def eval_one(
    row: dict[str, str],
    label: str,
    source: str,
    method_key: str,
    stage: str,
    mask_path: Path,
    person_idx: int,
    cfg: EvalConfig,
) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    qpos_path = qpos_path_for(row, source, stage)
    scene = repo_path(row["rubber_scene_act"])
    if not qpos_path.is_file():
        return None, {"method_label": label, "short_case_id": row["short_case_id"], "reason": "missing_qpos", "qpos_path": rel(qpos_path)}
    if not scene.is_file():
        return None, {"method_label": label, "short_case_id": row["short_case_id"], "reason": "missing_scene", "scene": rel(scene)}
    item = evaluate_sequence(
        row=row,
        method=label,
        hand_collision_variant_id="rubber_hull",
        qpos_path=qpos_path,
        scene_xml=scene,
        config=cfg,
        kin_ref_path=kin_ref_for_scene(scene),
        contact_mask_path=mask_path,
        person_idx=person_idx,
    )
    add_success_flags(item, cfg)
    item.update(run_health(qpos_path))
    item.update(
        {
            "short_case_id": row["short_case_id"],
            "variant": row["variant"],
            "method": label,
            "method_label": label,
            "method_key": method_key,
            "source_exp": source,
            "source_method_group": row.get("method_group", row.get("method", "")),
            "split": row.get("split", ""),
            "result_npz": rel(qpos_path),
            "video": rel(video_path_for(row, source, stage)),
            "e164_mask_path": rel(mask_path),
            "e164_mask_person_idx": person_idx,
        }
    )
    return item, None


def evaluate(stage: str, cfg: EvalConfig) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    indexes = source_indexes()
    e164_rows = [indexes["E164"][(case, "bimanualGlobalMaskReward")] for case in TARGET_CASES]
    e164_by_case = {row["short_case_id"]: row for row in e164_rows}
    artifact_rows = [artifact_check(row, stage) for row in e164_rows]
    metric_rows: list[dict[str, Any]] = []
    missing_rows: list[dict[str, Any]] = []

    for case in TARGET_CASES:
        e164_row = e164_by_case[case]
        mask_path = repo_path(e164_row["mask_path"])
        person_idx = int(e164_row["person_idx"])
        specs = [
            ("SPIDER+rubberhand", "rubberhand", "E156", indexes["E156"].get((case, "spider-rubberhand"))),
            ("E161 releaseDecay", "e161_release", "E161", indexes["E161"].get((case, "releaseDecay"))),
            ("E163 narrowSurfaceBand", "e163_narrow", "E163", indexes["E163"].get((case, "narrowSurfaceBand"))),
            ("E164 bimanualGlobalMaskReward", "e164", "E164", e164_row),
        ]
        for label, key, source, row in specs:
            if row is None:
                missing_rows.append({"method_label": label, "short_case_id": case, "reason": "missing_manifest_row"})
                continue
            metrics, missing = eval_one(row, label, source, key, stage, mask_path, person_idx, cfg)
            if metrics is not None:
                if label == "E164 bimanualGlobalMaskReward":
                    artifact = next(a for a in artifact_rows if a["short_case_id"] == case)
                    metrics["artifact_ok"] = artifact["artifact_ok"]
                    metrics["config_act"] = artifact["config_path"]
                else:
                    metrics["artifact_ok"] = True
                metric_rows.append(metrics)
            if missing is not None:
                missing_rows.append(missing)
    return metric_rows, artifact_rows, missing_rows


def tracked_ok(row: dict[str, Any]) -> bool:
    value = row.get("success_tracked")
    if value not in {None, ""}:
        return bool_text(value)
    terminal_pz = finite(row.get("track_pelvis_z_err_terminal_m"))
    return math.isfinite(terminal_pz) and terminal_pz <= 0.08


def table4_complete(row: dict[str, Any]) -> bool:
    return all(math.isfinite(finite(row.get(key))) for key in TABLE4_TRACKING)


def case_status(row: dict[str, Any], baseline: dict[str, Any]) -> tuple[str, float]:
    if row.get("method_label") == "E164 bimanualGlobalMaskReward" and not bool(row.get("artifact_ok", False)):
        return "产物/配置缺失", math.nan
    raw = finite(row.get(CONTACT_METRIC))
    base_raw = finite(baseline.get(CONTACT_METRIC))
    delta = raw - base_raw
    if not math.isfinite(delta):
        return "接触指标缺失", math.nan
    if delta < CONTACT_DROP_FAIL_TH:
        return "接触退化", delta
    if bool_text(row.get("fall_flag")):
        return "摔倒", delta
    if not tracked_ok(row):
        return "tracking失败", delta
    if not table4_complete(row):
        return "tracking指标缺失", delta
    return "通过", delta


def build_per_case(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_case_method = {(row["short_case_id"], row["method_label"]): row for row in rows}
    out: list[dict[str, Any]] = []
    labels = [label for _, label, _, _ in METHOD_ORDER]
    for case in TARGET_CASES:
        baseline = by_case_method.get((case, "SPIDER+rubberhand"))
        if baseline is None:
            continue
        base_raw = finite(baseline.get(CONTACT_METRIC))
        raw_min = base_raw + CONTACT_DROP_FAIL_TH if math.isfinite(base_raw) else math.nan
        for label in labels:
            row = by_case_method.get((case, label))
            if row is None:
                continue
            status, delta = case_status(row, baseline)
            item: dict[str, Any] = {
                "方法": label,
                "case": case,
                "状态": status,
                "物理接触Δ_vs_rubberhand": delta,
                "RL接触Δ_vs_rubberhand": (
                    finite(row.get("hand_object_physics_contact_in_rl_mask_frac"))
                    - finite(baseline.get("hand_object_physics_contact_in_rl_mask_frac"))
                ),
                "raw接触下限": raw_min,
                "success_tracked": tracked_ok(row),
                "fall_flag": bool_text(row.get("fall_flag")),
                "Table4完整": table4_complete(row),
                "source_exp": row.get("source_exp", ""),
                "variant": row.get("variant", ""),
                "result_npz": row.get("result_npz", ""),
                "e164_mask_path": row.get("e164_mask_path", ""),
            }
            for key in CORE_METRICS:
                item[key] = row.get(key, "")
            for key in HEALTH_KEYS.values():
                item[key[1]] = row.get(key[1], "")
            out.append(item)
    return out


def build_summary(per_case: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for _, label, source, method_group in METHOD_ORDER:
        rows = [row for row in per_case if row["方法"] == label]
        failures = [row for row in rows if row["状态"] != "通过"]
        contact_failures = [row for row in rows if row["状态"] == "接触退化"]
        item: dict[str, Any] = {
            "方法": label,
            "case数": len(rows),
            "通过case": sum(1 for row in rows if row["状态"] == "通过"),
            "失败case": ",".join(row["case"] for row in failures),
            "接触退化case": ",".join(row["case"] for row in contact_failures),
            "物理接触Δ最差": min(
                [finite(row.get("物理接触Δ_vs_rubberhand")) for row in rows if math.isfinite(finite(row.get("物理接触Δ_vs_rubberhand")))],
                default=math.nan,
            ),
            "source/method_group": f"{source}:{method_group}",
        }
        for key in CORE_METRICS:
            item[f"{key}_mean"] = mean([row.get(key) for row in rows])
        out.append(item)
    return out


def xlsx_value(value: Any) -> Any:
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float) and not math.isfinite(value):
        return ""
    return value


def append_table(
    wb: Workbook,
    title: str,
    rows: list[dict[str, Any]],
    columns: list[tuple[str, str, str]],
    fill: str,
) -> None:
    if len(wb.sheetnames) == 1 and wb.active.max_row == 1 and wb.active["A1"].value is None:
        ws = wb.active
        ws.delete_rows(1)
    else:
        ws = wb.create_sheet(title)
    ws.title = title
    ws.append([header for header, _, _ in columns])
    for row in rows:
        ws.append([xlsx_value(row.get(key, "")) for _, key, _ in columns])

    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, (_, _, fmt_kind) in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        for cell in ws[letter][1:]:
            if fmt_kind == "pct":
                cell.number_format = "0.0%"
            elif fmt_kind == "pct_delta":
                cell.number_format = "0.0%;[Red]-0.0%"
            elif fmt_kind == "num":
                cell.number_format = "0.00"
        width = min(max(len(str(cell.value or "")) for cell in ws[letter]) + 2, 46)
        ws.column_dimensions[letter].width = max(width, 10)


def write_workbook(out_dir: Path, stage: str, summary: list[dict[str, Any]], per_case: list[dict[str, Any]], artifacts: list[dict[str, Any]]) -> Path:
    wb = Workbook()
    main_cols = [
        ("方法", "方法", "text"),
        ("case数", "case数", "int"),
        ("通过case", "通过case", "int"),
        ("失败case", "失败case", "text"),
        ("接触退化case", "接触退化case", "text"),
        ("物理接触Δ最差", "物理接触Δ最差", "pct_delta"),
        ("物理接触(raw)", f"{CONTACT_METRIC}_mean", "pct"),
        ("RL接触(raw)", "hand_object_physics_contact_in_rl_mask_frac_mean", "pct"),
        ("RL mask占比", "rl_object_contact_ref_frac_mean", "pct"),
        ("clean3接触", "hand_object_physics_contact_3mm_in_mask_frac_mean", "pct"),
        ("clean5接触", "hand_object_physics_contact_5mm_in_mask_frac_mean", "pct"),
        ("物理穿透3mm", "hand_object_physics_penetration_3mm_frame_frac_mean", "pct"),
        ("几何穿透2mm", "hand_geom_penetration_2mm_frac_mean", "pct"),
        ("关节误差(°)", "track_joint_err_deg_mean_mean", "num"),
        ("末端位置误差(cm)", "track_eef_pos_err_cm_mean_mean", "num"),
        ("物体位置误差(cm)", "track_obj_pos_err_cm_mean_mean", "num"),
        ("bimanual gate", "surface_band_bimanual_gate_mean_mean", "pct"),
        ("source", "source/method_group", "text"),
    ]
    append_table(wb, "主表", summary, main_cols, "1F4E79")

    per_case_cols = [
        ("方法", "方法", "text"),
        ("case", "case", "text"),
        ("状态", "状态", "text"),
        ("物理接触Δ_vs_rubberhand", "物理接触Δ_vs_rubberhand", "pct_delta"),
        ("RL接触Δ_vs_rubberhand", "RL接触Δ_vs_rubberhand", "pct_delta"),
        ("raw接触下限", "raw接触下限", "pct"),
        ("物理接触(raw)", CONTACT_METRIC, "pct"),
        ("RL接触(raw)", "hand_object_physics_contact_in_rl_mask_frac", "pct"),
        ("RL mask占比", "rl_object_contact_ref_frac", "pct"),
        ("RL补洞帧数", "rl_object_contact_filled_frame_count", "int"),
        ("clean3接触", "hand_object_physics_contact_3mm_in_mask_frac", "pct"),
        ("clean5接触", "hand_object_physics_contact_5mm_in_mask_frac", "pct"),
        ("物理穿透3mm", "hand_object_physics_penetration_3mm_frame_frac", "pct"),
        ("几何穿透2mm", "hand_geom_penetration_2mm_frac", "pct"),
        ("tracked", "success_tracked", "text"),
        ("fall", "fall_flag", "text"),
        ("Table4完整", "Table4完整", "text"),
        ("关节误差(°)", "track_joint_err_deg_mean", "num"),
        ("末端位置误差(cm)", "track_eef_pos_err_cm_mean", "num"),
        ("物体位置误差(cm)", "track_obj_pos_err_cm_mean", "num"),
        ("surface bimanual gate", "surface_band_bimanual_gate_mean", "pct"),
        ("contact bimanual gate", "contact_hdmi_bimanual_gate_mean", "pct"),
        ("mask", "e164_mask_path", "text"),
    ]
    append_table(wb, "逐case", per_case, per_case_cols, "548235")

    artifact_cols = [
        ("variant", "variant", "text"),
        ("case", "short_case_id", "text"),
        ("artifact_ok", "artifact_ok", "text"),
        ("root_npz", "root_npz_exists", "text"),
        ("trajectory", "trajectory_mjwp_act_exists", "text"),
        ("config", "config_act_exists", "text"),
        ("mp4", "full_mp4_exists", "text"),
        ("mask", "e164_mask_exists", "text"),
        ("surface_bimanual", "config_surface_bimanual_ok", "text"),
        ("contact_bimanual", "config_contact_bimanual_ok", "text"),
        ("axis_spider", "config_contact_mask_axis_ok", "text"),
        ("diag_surface", "diag_surface_bimanual_exists", "text"),
        ("diag_contact", "diag_contact_bimanual_exists", "text"),
        ("config_path", "config_path", "text"),
    ]
    append_table(wb, "E164产物检查", artifacts, artifact_cols, "8064A2")

    info_rows = [
        {"项": "mask", "说明": "所有方法的 masked contact/penetration 均用 E164 global spider mask 重算。"},
        {"项": "reward", "说明": "E164 CEM reward 要求 surfaceBand 和 contact_hdmi 都走双手 AND/瓶颈手 min。"},
        {"项": "raw contact gate", "说明": "raw contact in mask 相对 SPIDER+rubberhand 下降超过 0.05 判接触退化。"},
        {"项": "RL contact", "说明": "按 Holosoma downstream exporter 口径：spider_contact_mask_3cm 取 max(L,R)，填补 <=5 帧内部断口，统计该 RL object_contact 窗口内 raw contact 比例。"},
        {"项": "stage", "说明": f"当前评测阶段: {stage}。"},
    ]
    append_table(wb, "说明", info_rows, [("项", "项", "text"), ("说明", "说明", "text")], "C65911")

    out_path = out_dir / f"E164_bimanual_global_mask_reward_{stage}_eval.xlsx"
    wb.save(out_path)
    return out_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("stage", nargs="?", default="full", choices=["smoke", "full"])
    parser.add_argument("--allow-missing", action="store_true")
    args = parser.parse_args()

    out_dir = RESULT_ROOT / "eval" / args.stage
    out_dir.mkdir(parents=True, exist_ok=True)

    cfg = EvalConfig()
    metric_rows, artifact_rows, missing_rows = evaluate(args.stage, cfg)
    per_case = build_per_case(metric_rows)
    summary = build_summary(per_case)

    write_tsv(out_dir / "e164_method_metrics.tsv", metric_rows, sorted({key for row in metric_rows for key in row.keys()}))
    write_tsv(out_dir / "e164_case_status.tsv", per_case, sorted({key for row in per_case for key in row.keys()}))
    write_tsv(out_dir / "e164_method_summary.tsv", summary, sorted({key for row in summary for key in row.keys()}))
    write_tsv(out_dir / "e164_artifact_check.tsv", artifact_rows, sorted({key for row in artifact_rows for key in row.keys()}))
    write_tsv(out_dir / "e164_missing.tsv", missing_rows, sorted({key for row in missing_rows for key in row.keys()}) if missing_rows else ["reason"])
    xlsx = write_workbook(out_dir, args.stage, summary, per_case, artifact_rows)

    e164_case_rows = [row for row in per_case if row["方法"] == "E164 bimanualGlobalMaskReward"]
    eval_summary = {
        "stage": args.stage,
        "target_cases": TARGET_CASES,
        "method_metric_rows": len(metric_rows),
        "missing_rows": len(missing_rows),
        "e164_pass_cases": sum(1 for row in e164_case_rows if row["状态"] == "通过"),
        "e164_failed_cases": [row["case"] for row in e164_case_rows if row["状态"] != "通过"],
        "e164_contact_failed_cases": [row["case"] for row in e164_case_rows if row["状态"] == "接触退化"],
        "xlsx": rel(xlsx),
    }
    write_json(out_dir / "e164_eval_summary.json", eval_summary)
    print(
        "E164 eval: "
        f"stage={args.stage} rows={len(metric_rows)} missing={len(missing_rows)} "
        f"e164_pass={eval_summary['e164_pass_cases']}/{len(TARGET_CASES)} xlsx={rel(xlsx)}"
    )
    if missing_rows and not args.allow_missing:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
