#!/usr/bin/env python3
"""Evaluate E148 E143-24 OmniRetarget vs sphere Spider vs rubber hand Spider."""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import math
from pathlib import Path
from typing import Any

import mujoco
import numpy as np
from scipy.spatial.transform import Rotation as R

REPO = Path(__file__).resolve().parents[4]
VARIANTS_TSV = REPO / "workspace/core4d/scripts/E148/variants.tsv"
E147_EVAL = REPO / "workspace/core4d/scripts/eval/eval_E147_rubber_hand_collision.py"
RESULT_ROOT = REPO / "workspace/core4d/results/E148/e143_24case_rubber_hand_collision"
FILTERED_EXCLUDE_CASES = {"bucket004_20231003_1_012_p1"}

METRICS = [
    ("手物接触", "hand_object_physics_contact_frac", "0.0%"),
    ("5cm", "hand_geom_near_5cm_frac", "0.0%"),
    ("10cm", "hand_geom_near_10cm_frac", "0.0%"),
    ("手物穿透", "hand_geom_penetration_frac", "0.0%"),
    ("深穿透2cm", "hand_geom_deep_penetration_2cm_frac", "0.0%"),
    ("腿穿透", "leg_penetration_frac", "0.0%"),
    ("body穿透", "body_penetration_frac", "0.0%"),
    ("物体触地", "object_floor_contact_frac", "0.0%"),
    ("pelvis_min_m", "pelvis_min_m", "0.000"),
]

METHOD_FIELDS_PREFIX = ["method_key", "source_exp", "run_status"]


def rel(path: Path | str) -> str:
    if not path:
        return ""
    p = Path(path)
    try:
        return str(p.resolve().relative_to(REPO.resolve()))
    except Exception:
        return str(path)


def repo_path(text: str | Path) -> Path:
    p = Path(text)
    return p if p.is_absolute() else REPO / p


def finite(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        out = float(value)
    except Exception:
        return None
    return out if math.isfinite(out) else None


def parse_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"true", "1", "yes", "pass"}


def fmt(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return f"{value:.8g}" if math.isfinite(value) else ""
    return str(value)


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def write_tsv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: fmt(row.get(field, "")) for field in fields})


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def load_e147_eval():
    spec = importlib.util.spec_from_file_location("eval_E147_for_E148", E147_EVAL)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load {E147_EVAL}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def scene_euler_convention(scene_act: Path) -> str:
    meta = scene_act.parent / "scene_act_meta.json"
    if meta.is_file():
        return str(json.loads(meta.read_text(encoding="utf-8")).get("euler_convention", "XYZ"))
    return "XYZ"


def convert_freejoint_to_scene_act(qpos: np.ndarray, scene_act: Path) -> np.ndarray:
    model = mujoco.MjModel.from_xml_path(str(scene_act))
    if qpos.shape[1] == model.nq:
        return qpos.astype(np.float64, copy=True)
    nq_robot = model.nq - 6
    if qpos.shape[1] < nq_robot + 7:
        raise ValueError(f"cannot convert qpos shape={qpos.shape} for scene nq={model.nq}: {scene_act}")

    obj_body = mujoco.mj_name2id(model, mujoco.mjtObj.mjOBJ_BODY, "object")
    if obj_body < 0:
        raise ValueError(f"scene has no object body: {scene_act}")

    obj_pos_world = qpos[:, nq_robot : nq_robot + 3]
    obj_quat_wxyz = qpos[:, nq_robot + 3 : nq_robot + 7]
    body_pos = model.body_pos[obj_body]
    body_quat_wxyz = model.body_quat[obj_body]
    body_quat_xyzw = [body_quat_wxyz[1], body_quat_wxyz[2], body_quat_wxyz[3], body_quat_wxyz[0]]
    r_body = R.from_quat(body_quat_xyzw)

    obj_slide = r_body.inv().apply(obj_pos_world - body_pos[np.newaxis, :])
    obj_quat_xyzw = np.column_stack(
        [obj_quat_wxyz[:, 1], obj_quat_wxyz[:, 2], obj_quat_wxyz[:, 3], obj_quat_wxyz[:, 0]]
    )
    obj_euler = (r_body.inv() * R.from_quat(obj_quat_xyzw)).as_euler(scene_euler_convention(scene_act))

    out = np.zeros((qpos.shape[0], model.nq), dtype=np.float64)
    out[:, :nq_robot] = qpos[:, :nq_robot]
    out[:, nq_robot : nq_robot + 3] = obj_slide
    out[:, nq_robot + 3 : nq_robot + 6] = obj_euler
    return out


def converted_omni_qpos(row: dict[str, str], scene_act: Path, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    qpos_path = out_dir / f"{row['variant']}_omni_scene_act_qpos.npz"
    data = np.load(repo_path(row["omni_qpos_path"]), allow_pickle=True)
    qpos = np.asarray(data["qpos"], dtype=np.float64)
    if qpos.ndim == 3:
        qpos = qpos[:, 0, :]
    if qpos.ndim != 2:
        raise ValueError(f"unsupported OmniRetarget qpos shape={qpos.shape}: {row['omni_qpos_path']}")
    np.savez(qpos_path, qpos=convert_freejoint_to_scene_act(qpos, scene_act))
    return qpos_path


def eval_sequence(eval_mod: Any, row: dict[str, str], *, method: str, method_key: str, source_exp: str, run_status: str, qpos_path: Path, scene_xml: Path) -> dict[str, Any]:
    item = eval_mod.evaluate_sequence(
        row=row,
        method=method,
        hand_collision_variant_id=method_key,
        qpos_path=qpos_path,
        scene_xml=scene_xml,
    )
    item["method_key"] = method_key
    item["source_exp"] = source_exp
    item["run_status"] = run_status
    return item


def build_method_rows(rows: list[dict[str, str]], stage: str, results_dir: Path, allow_missing_rubber: bool) -> tuple[list[dict[str, Any]], list[dict[str, str]], list[str]]:
    eval_mod = load_e147_eval()
    metric_fields = list(eval_mod.METRIC_FIELDS)
    metrics: list[dict[str, Any]] = []
    missing: list[dict[str, str]] = []
    omni_dir = RESULT_ROOT / "comparison/omni_converted_qpos"
    for row in rows:
        case_id = row["case_id"]
        base_scene = repo_path(row["base_scene_act"])
        omni_qpos = converted_omni_qpos(row, base_scene, omni_dir)
        metrics.append(
            eval_sequence(
                eval_mod,
                row,
                method="OmniRetarget kinematic replay",
                method_key="OmniRetarget",
                source_exp="E143",
                run_status="recomputed_omni",
                qpos_path=omni_qpos,
                scene_xml=base_scene,
            )
        )

        sphere_path = repo_path(row["sphere_outdir_npz"])
        metrics.append(
            eval_sequence(
                eval_mod,
                row,
                method="sphere Spider raw_mask_ref_fk",
                method_key="sphere_spider",
                source_exp="E143",
                run_status="reuse_e143",
                qpos_path=sphere_path,
                scene_xml=base_scene,
            )
        )

        rubber_path = repo_path(row["rubber_outdir_npz"])
        rubber_scene = repo_path(row["rubber_scene_act"])
        if not rubber_path.is_file():
            missing.append({"case_id": case_id, "variant": row["variant"], "missing": rel(rubber_path)})
            continue
        metrics.append(
            eval_sequence(
                eval_mod,
                row,
                method="rubber hand Spider",
                method_key="rubber_hand_spider",
                source_exp=row.get("reuse_source_exp") or "E148",
                run_status=row.get("run_status", ""),
                qpos_path=rubber_path,
                scene_xml=rubber_scene,
            )
        )
    if missing and not allow_missing_rubber:
        raise FileNotFoundError("missing rubber outputs:\n" + "\n".join(f"{m['variant']}: {m['missing']}" for m in missing))
    return metrics, missing, METHOD_FIELDS_PREFIX + metric_fields


def metric_lookup(method_rows: list[dict[str, Any]]) -> dict[tuple[str, str], dict[str, Any]]:
    out = {}
    for row in method_rows:
        out[(str(row["case_id"]), str(row["method_key"]))] = row
    return out


def build_case_rows(rows: list[dict[str, str]], method_rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[str]]:
    lookup = metric_lookup(method_rows)
    fields = [
        "case_id",
        "object_key",
        "run_status",
        "rubber_source_exp",
        "e143_variant",
        "e147_variant",
        "expected_quality",
    ]
    for label, _metric, _fmt in METRICS:
        fields.extend(
            [
                f"{label}_OmniRetarget",
                f"{label}_sphere_spider",
                f"{label}_rubber_hand_spider",
                f"{label}_rubber-sphere",
                f"{label}_rubber-Omni",
                f"{label}_Omni-rubber",
            ]
        )
    fields.extend(
        [
            "fall_OmniRetarget",
            "fall_sphere_spider",
            "fall_rubber_hand_spider",
            "sphere_npz",
            "rubber_npz",
            "rubber_video",
        ]
    )
    out = []
    for row in rows:
        case_id = row["case_id"]
        if (case_id, "rubber_hand_spider") not in lookup:
            continue
        omni = lookup[(case_id, "OmniRetarget")]
        sphere = lookup[(case_id, "sphere_spider")]
        rubber = lookup[(case_id, "rubber_hand_spider")]
        item: dict[str, Any] = {
            "case_id": case_id,
            "object_key": row["object_key"],
            "run_status": row["run_status"],
            "rubber_source_exp": row.get("reuse_source_exp") or "E148",
            "e143_variant": row["e143_variant"],
            "e147_variant": row.get("e147_variant", ""),
            "expected_quality": row.get("expected_quality", ""),
            "fall_OmniRetarget": parse_bool(omni.get("fall_flag")),
            "fall_sphere_spider": parse_bool(sphere.get("fall_flag")),
            "fall_rubber_hand_spider": parse_bool(rubber.get("fall_flag")),
            "sphere_npz": row.get("sphere_outdir_npz", ""),
            "rubber_npz": row.get("rubber_outdir_npz", ""),
            "rubber_video": row.get("rubber_video", ""),
        }
        for label, metric, _fmt in METRICS:
            item[f"{label}_OmniRetarget"] = finite(omni.get(metric))
            item[f"{label}_sphere_spider"] = finite(sphere.get(metric))
            item[f"{label}_rubber_hand_spider"] = finite(rubber.get(metric))
            item[f"{label}_rubber-sphere"] = None
            item[f"{label}_rubber-Omni"] = None
            item[f"{label}_Omni-rubber"] = None
        out.append(item)
    return out, fields


def mean(vals: list[Any]) -> float:
    nums = [float(v) for v in vals if finite(v) is not None]
    return sum(nums) / len(nums) if nums else math.nan


def summary_for(case_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    methods = [
        ("OmniRetarget", "OmniRetarget"),
        ("sphere Spider", "sphere_spider"),
        ("rubber hand Spider", "rubber_hand_spider"),
    ]
    for label, key in methods:
        item: dict[str, Any] = {"method": label, "case_count": len(case_rows)}
        item["fall_count"] = sum(bool(row.get(f"fall_{key}")) for row in case_rows)
        for metric_label, _metric, _fmt in METRICS:
            item[metric_label] = mean([row.get(f"{metric_label}_{key}") for row in case_rows])
        out.append(item)
    return out


def write_xlsx(path: Path, case_rows: list[dict[str, Any]], case_fields: list[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter, quote_sheetname

    wb = Workbook()
    ws_avg = wb.active
    ws_avg.title = "24case平均"
    ws_case = wb.create_sheet("逐case对比")
    ws_favg = wb.create_sheet("filtered平均")
    ws_fcase = wb.create_sheet("filtered逐case")

    full_rows = case_rows
    filtered_rows = [row for row in case_rows if row["case_id"] not in FILTERED_EXCLUDE_CASES]

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    delta_fill = PatternFill("solid", fgColor="FCE4D6")
    path_fill = PatternFill("solid", fgColor="E2F0D9")

    def write_case_sheet(ws, rows):
        ws.append(case_fields)
        for row in rows:
            ws.append([row.get(field) for field in case_fields])
        field_to_col = {field: idx + 1 for idx, field in enumerate(case_fields)}
        for r in range(2, len(rows) + 2):
            for label, _metric, _fmt in METRICS:
                rub = get_column_letter(field_to_col[f"{label}_rubber_hand_spider"])
                sph = get_column_letter(field_to_col[f"{label}_sphere_spider"])
                omni = get_column_letter(field_to_col[f"{label}_OmniRetarget"])
                ws.cell(r, field_to_col[f"{label}_rubber-sphere"]).value = f"={rub}{r}-{sph}{r}"
                ws.cell(r, field_to_col[f"{label}_rubber-Omni"]).value = f"={rub}{r}-{omni}{r}"
                ws.cell(r, field_to_col[f"{label}_Omni-rubber"]).value = f"={omni}{r}-{rub}{r}"

    def write_avg_sheet(ws, case_ws, rows):
        avg_fields = ["method", "case_count", "fall_count"] + [label for label, _metric, _fmt in METRICS]
        ws.append(avg_fields)
        field_to_col = {field: idx + 1 for idx, field in enumerate(case_fields)}
        method_col_map = {
            "OmniRetarget": "OmniRetarget",
            "sphere Spider": "sphere_spider",
            "rubber hand Spider": "rubber_hand_spider",
        }
        case_sheet = quote_sheetname(case_ws.title)
        row_start = 2
        row_end = len(rows) + 1
        for row_idx, method in enumerate(["OmniRetarget", "sphere Spider", "rubber hand Spider"], start=2):
            key = method_col_map[method]
            ws.cell(row_idx, 1).value = method
            ws.cell(row_idx, 2).value = f"=COUNTA({case_sheet}!$A${row_start}:$A${row_end})"
            fall_col = get_column_letter(field_to_col[f"fall_{key}"])
            ws.cell(row_idx, 3).value = f"=COUNTIF({case_sheet}!{fall_col}${row_start}:{fall_col}${row_end},TRUE)"
            for col_idx, (label, _metric, _fmt) in enumerate(METRICS, start=4):
                metric_col = get_column_letter(field_to_col[f"{label}_{key}"])
                ws.cell(row_idx, col_idx).value = f"=AVERAGE({case_sheet}!{metric_col}${row_start}:{metric_col}${row_end})"

    write_case_sheet(ws_case, full_rows)
    write_case_sheet(ws_fcase, filtered_rows)
    write_avg_sheet(ws_avg, ws_case, full_rows)
    write_avg_sheet(ws_favg, ws_fcase, filtered_rows)

    for ws in (ws_avg, ws_case, ws_favg, ws_fcase):
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            header = str(ws.cell(1, col_idx).value or "")
            max_len = max(len(header), *(len(str(ws.cell(r, col_idx).value or "")) for r in range(2, min(ws.max_row, 30) + 1)))
            ws.column_dimensions[letter].width = max(10, min(45, max_len + 2))
            if "rubber-" in header or "Omni-rubber" in header:
                ws.cell(1, col_idx).fill = delta_fill
            if header.endswith("_npz") or header.endswith("_video"):
                ws.cell(1, col_idx).fill = path_fill
            for row_idx in range(2, ws.max_row + 1):
                cell = ws.cell(row_idx, col_idx)
                cell.font = Font(name="Arial")
                if isinstance(cell.value, float) or (isinstance(cell.value, str) and cell.value.startswith("=AVERAGE")):
                    cell.number_format = "0.000" if "pelvis_min_m" in header else "0.0%"
                cell.alignment = Alignment(vertical="top", wrap_text=header.endswith("_npz") or header.endswith("_video"))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_summary(path: Path, case_rows: list[dict[str, Any]]) -> None:
    full = summary_for(case_rows)
    filtered = summary_for([row for row in case_rows if row["case_id"] not in FILTERED_EXCLUDE_CASES])

    def table(rows: list[dict[str, Any]]) -> list[str]:
        lines = [
            "| method | cases | fall | 5cm | 10cm | hand pen | deep2 | leg pen | object floor | pelvis min |",
            "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
        ]
        for row in rows:
            lines.append(
                f"| `{row['method']}` | {row['case_count']} | {row['fall_count']} | "
                f"{row['5cm']:.4f} | {row['10cm']:.4f} | {row['手物穿透']:.4f} | "
                f"{row['深穿透2cm']:.4f} | {row['腿穿透']:.4f} | {row['物体触地']:.4f} | {row['pelvis_min_m']:.4f} |"
            )
        return lines

    def diff_lines(rows: list[dict[str, Any]], title: str) -> list[str]:
        by_method = {row["method"]: row for row in rows}
        omni = by_method["OmniRetarget"]
        sphere = by_method["sphere Spider"]
        rubber = by_method["rubber hand Spider"]
        lines = [
            f"## {title} diffs",
            "",
            "| metric | OmniRetarget-rubber | rubber-sphere |",
            "|---|---:|---:|",
        ]
        for label in ("5cm", "10cm", "手物穿透", "腿穿透"):
            lines.append(f"| {label} | {omni[label] - rubber[label]:.4f} | {rubber[label] - sphere[label]:.4f} |")
        return lines

    lines = [
        "# E148 E143-24 rubber hand collision comparison",
        "",
        f"- full cases: `{len(case_rows)}`",
        f"- filtered excludes: `{', '.join(sorted(FILTERED_EXCLUDE_CASES))}`",
        "",
        "## 24-case average",
        "",
        *table(full),
        "",
        "## Filtered average",
        "",
        *table(filtered),
        "",
        *diff_lines(full, "24-case"),
        "",
        *diff_lines(filtered, "Filtered"),
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--stage", default="full", choices=("smoke", "full"))
    parser.add_argument("--variants", type=Path, default=VARIANTS_TSV)
    parser.add_argument("--results-dir", type=Path, default=None)
    parser.add_argument("--out-dir", type=Path, default=None)
    parser.add_argument("--allow-missing-rubber", action="store_true")
    args = parser.parse_args()

    _results_dir = args.results_dir or (RESULT_ROOT / "cem" / args.stage)
    out_dir = args.out_dir or (RESULT_ROOT / "eval" / args.stage)
    comparison_dir = RESULT_ROOT / "comparison"
    rows = read_tsv(args.variants)
    method_rows, missing, method_fields = build_method_rows(rows, args.stage, _results_dir, args.allow_missing_rubber)
    case_rows, case_fields = build_case_rows(rows, method_rows)

    out_dir.mkdir(parents=True, exist_ok=True)
    write_tsv(out_dir / "e148_method_metrics.tsv", method_rows, method_fields)
    write_tsv(out_dir / "e148_case_comparison.tsv", case_rows, case_fields)
    write_tsv(out_dir / "e148_filtered_case_comparison.tsv", [r for r in case_rows if r["case_id"] not in FILTERED_EXCLUDE_CASES], case_fields)
    if missing:
        write_tsv(out_dir / "e148_missing_rubber.tsv", missing, ["case_id", "variant", "missing"])
    summary = {
        "stage": args.stage,
        "manifest_rows": len(rows),
        "method_metric_rows": len(method_rows),
        "case_rows": len(case_rows),
        "filtered_case_rows": sum(row["case_id"] not in FILTERED_EXCLUDE_CASES for row in case_rows),
        "missing_rubber_rows": len(missing),
        "reuse_e147_rows": sum(row.get("run_status") == "reuse_e147" for row in rows),
        "to_run_rows": sum(row.get("run_status") == "to_run" for row in rows),
        "xlsx": rel(comparison_dir / "E148_e143_omni_sphere_rubber_hand_comparison.xlsx"),
    }
    write_json(out_dir / "e148_eval_summary.json", summary)
    write_summary(out_dir / "e148_summary.md", case_rows)
    write_xlsx(comparison_dir / "E148_e143_omni_sphere_rubber_hand_comparison.xlsx", case_rows, case_fields)
    print(f"[E148 eval] wrote method_rows={len(method_rows)} case_rows={len(case_rows)} missing_rubber={len(missing)} to {out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
