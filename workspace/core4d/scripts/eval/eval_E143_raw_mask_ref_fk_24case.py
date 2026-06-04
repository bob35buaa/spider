#!/usr/bin/env python3
"""Evaluate E143 raw_mask_ref_fk 24-case sweep against OmniRetarget."""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import math
import os
from pathlib import Path
from typing import Any


THIS = Path(__file__).resolve()
REPO = THIS.parents[4]
TASK_ROOT = REPO / "example_datasets/processed/core4d/unitree_g1/humanoid_object"
VARIANTS_TSV = REPO / "workspace/core4d/scripts/E143/variants.tsv"
E109_METHODS = REPO / "workspace/core4d/results/E109/expanded_24_work_cases/expanded_24_method_metrics.tsv"
E109_FAIR_METRICS = REPO / "workspace/core4d/results/E109/fair_eval/method_case_metrics.tsv"
E090_EVAL = REPO / "workspace/core4d/scripts/eval/eval_E090.py"
E105_EVAL = REPO / "workspace/core4d/scripts/eval/eval_E105_box026_clean_cem.py"
OUT_ROOT = REPO / "workspace/core4d/results/E143/raw_mask_ref_fk_24case_omniretarget_comparison"


FIELDS = [
    "ordinal",
    "variant",
    "e109_case_id",
    "case_id",
    "object_key",
    "target_variant_id",
    "source_task",
    "derived_task",
    "person_idx",
    "split",
    "ablation",
    "mask_path",
    "mask_kind",
    "source_mask_path",
    "baseline_npz_path",
    "baseline_run_id",
    "omni_qpos_path",
    "omni_scene_xml",
    "spider_scene_xml",
    "override",
    "run_status",
    "reuse_source_exp",
    "reuse_npz_path",
    "reuse_video_path",
    "reuse_outdir_npz_path",
]

CORE_METRICS = [
    "hand_object_contact",
    "hand_near_5cm",
    "hand_near_10cm",
    "hand_object_penetration",
    "leg_penetration",
]


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(REPO))
    except ValueError:
        return str(path)


def repo_path(path: str) -> Path:
    p = Path(path)
    return p if p.is_absolute() else REPO / p


def repo_path_from_env(name: str, default: Path) -> Path:
    raw = os.environ.get(name)
    path = Path(raw) if raw else default
    return path if path.is_absolute() else REPO / path


def finite(value: Any, default: float | None = None) -> float | None:
    if value is None:
        return default
    text = str(value).strip()
    if text == "":
        return default
    try:
        out = float(text)
    except ValueError:
        return default
    if math.isnan(out):
        return default
    return out


def fmt(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return f"{value:.6f}"
    return str(value)


def read_variants() -> list[dict[str, str]]:
    with VARIANTS_TSV.open("r", encoding="utf-8", newline="") as f:
        lines = (line for line in f if line.strip() and not line.startswith("#"))
        return list(csv.DictReader(lines, fieldnames=FIELDS, delimiter="\t"))


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def write_tsv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: fmt(row.get(field)) for field in fields})


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Could not load {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def e109_metrics() -> dict[tuple[str, str], dict[str, Any]]:
    out: dict[tuple[str, str], dict[str, Any]] = {}
    for row in read_tsv(E109_METHODS):
        method = "OmniRetarget" if row["method"] == "OmniRetarget" else "ref_fk"
        out[(row["case_id"], method)] = {
            "e109_case_id": row["case_id"],
            "case_id": row["case_id"],
            "object_key": row["object_key"],
            "method": method,
            "hand_object_contact": finite(row.get("hand_object_physics_contact_frac")),
            "hand_near_5cm": finite(row.get("hand_geom_near_5cm_frac")),
            "hand_near_10cm": finite(row.get("hand_geom_near_10cm_frac")),
            "hand_object_penetration": finite(row.get("hand_geom_penetration_frac")),
            "hand_object_deep_penetration_2cm": finite(row.get("hand_geom_deep_penetration_2cm_frac")),
            "leg_penetration": finite(row.get("leg_penetration_frac")),
            "pelvis_min_m": finite(row.get("pelvis_min_m")),
            "qpos_path": row.get("qpos_path", ""),
            "scene_xml": row.get("scene_xml", ""),
            "run_id": row.get("run_id", ""),
        }
    return out


def spider_status_by_case(variants: list[dict[str, str]]) -> dict[str, str]:
    status_by_run = {}
    for row in read_tsv(E109_FAIR_METRICS):
        if row.get("method_family") == "spider":
            status_by_run[row["variant"]] = row.get("cem_status", "")
    return {row["e109_case_id"]: status_by_run.get(row["baseline_run_id"], "") for row in variants}


def is_spider_success(status: str) -> bool:
    text = status.strip()
    return text.upper() == "WORK" or text.lower() == "pass"


def expected_paths(stage: str, row: dict[str, str], results_dir: Path) -> dict[str, Path]:
    if row["run_status"] == "already_done":
        return {
            "root": repo_path(row["reuse_npz_path"]),
            "video": repo_path(row["reuse_video_path"]),
            "outdir": repo_path(row["reuse_outdir_npz_path"]),
        }
    variant = row["variant"]
    return {
        "root": results_dir / f"{variant}.npz",
        "video": results_dir / f"{variant}_{stage}.mp4",
        "outdir": results_dir / f"{variant}_outdir_{stage}/trajectory_mjwp_act.npz",
    }


def missing_outputs(stage: str, rows: list[dict[str, str]], results_dir: Path, allow_missing: bool) -> list[dict[str, str]]:
    missing: list[dict[str, str]] = []
    for row in rows:
        paths = expected_paths(stage, row, results_dir)
        absent = [rel(path) for path in paths.values() if not path.is_file()]
        if absent:
            missing.append({"variant": row["variant"], "case_id": row["case_id"], "missing": ";".join(absent)})
    if missing and not allow_missing:
        print("E143 eval is deferred until all raw-mask outputs exist. Missing:")
        for row in missing:
            print(f"- {row['variant']}: {row['missing']}")
        raise SystemExit(1)
    return missing


def compute_raw_metrics(modules: tuple[Any, Any], row: dict[str, str], stage: str, results_dir: Path) -> dict[str, Any] | None:
    paths = expected_paths(stage, row, results_dir)
    npz = paths["outdir"]
    if not npz.is_file():
        return None
    eval_e090, eval_e105 = modules
    scene = TASK_ROOT / row["derived_task"] / "scene_act.xml"
    metrics = eval_e090.compute_sim_metrics(npz, scene)
    metrics.update(eval_e105.replay_metrics(npz, scene))
    metrics.update(eval_e105.leg_object_metrics(row["variant"], npz, scene, results_dir))

    ts_path = repo_path(str(metrics["legobj_timeseries_csv"]))
    with ts_path.open("r", encoding="utf-8", newline="") as f:
        hand_sdf = [float(item["hand_box_sdf_min_m"]) for item in csv.DictReader(f)]
    hand_near_5 = sum(value <= 0.05 for value in hand_sdf) / len(hand_sdf)
    hand_near_10 = sum(value <= 0.10 for value in hand_sdf) / len(hand_sdf)
    hand_pen = sum(value < 0.0 for value in hand_sdf) / len(hand_sdf)
    hand_deep = sum(value < -0.02 for value in hand_sdf) / len(hand_sdf)
    work_status, stage_pass = eval_e105.status(stage, metrics)
    lowerbody_pass = bool(metrics["lowerbody_strict_pass"])
    strict_status = "WORK" if work_status == "WORK" and lowerbody_pass else "FAIL"
    return {
        **row,
        "method": "raw_mask_ref_fk",
        "stage": stage,
        "hand_object_contact": finite(metrics.get("hand_object_contact_physics_frac")),
        "hand_near_5cm": hand_near_5,
        "hand_near_10cm": hand_near_10,
        "hand_object_penetration": hand_pen,
        "hand_object_deep_penetration_2cm": hand_deep,
        "leg_penetration": finite(metrics.get("leg_box_interference_frac")),
        "pelvis_min_m": finite(metrics.get("pelvis_min_m")),
        "obj_err_mean_m": finite(metrics.get("obj_err_mean_m")),
        "obj_err_max_m": finite(metrics.get("obj_err_max_m")),
        "work_status": work_status,
        "work_status_lowerbody_strict": strict_status,
        "stage_pass": stage_pass,
        "npz_path": rel(npz),
        "root_npz_path": rel(paths["root"]),
        "video_path": rel(paths["video"]),
        "scene_xml": rel(scene),
        "legobj_timeseries_csv": metrics.get("legobj_timeseries_csv", ""),
    }


def aggregate(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for method in ["OmniRetarget", "ref_fk", "raw_mask_ref_fk"]:
        subset = [row for row in rows if row["method"] == method]
        item: dict[str, Any] = {"method": method, "case_count": len(subset)}
        for metric in CORE_METRICS:
            vals = [row.get(metric) for row in subset if row.get(metric) is not None]
            item[metric] = sum(float(v) for v in vals) / len(vals) if vals else None
        out.append(item)
    return out


def joined_case_rows(all_rows: list[dict[str, Any]], spider_status: dict[str, str]) -> list[dict[str, Any]]:
    by_case: dict[str, dict[str, dict[str, Any]]] = {}
    for row in all_rows:
        by_case.setdefault(row["e109_case_id"], {})[row["method"]] = row
    out = []
    for case_id in sorted(by_case):
        methods = by_case[case_id]
        omni = methods.get("OmniRetarget", {})
        ref = methods.get("ref_fk", {})
        raw = methods.get("raw_mask_ref_fk", {})
        row: dict[str, Any] = {
            "e109_case_id": case_id,
            "case_id": raw.get("case_id") or ref.get("case_id") or omni.get("case_id") or case_id,
            "object_key": raw.get("object_key") or ref.get("object_key") or omni.get("object_key", ""),
            "spider_cem_status": spider_status.get(case_id, ""),
            "spider_success_case": is_spider_success(spider_status.get(case_id, "")),
            "raw_status": "evaluated" if raw else "missing_raw_mask",
            "raw_video_path": raw.get("video_path", ""),
            "raw_npz_path": raw.get("root_npz_path", ""),
        }
        for metric in CORE_METRICS:
            row[f"OmniRetarget_{metric}"] = omni.get(metric)
            row[f"ref_fk_{metric}"] = ref.get(metric)
            row[f"raw_mask_ref_fk_{metric}"] = raw.get(metric)
            row[f"delta_raw_vs_spider_{metric}"] = (
                float(raw[metric]) - float(ref[metric])
                if raw.get(metric) is not None and ref.get(metric) is not None
                else None
            )
            row[f"delta_raw_vs_omni_{metric}"] = (
                float(raw[metric]) - float(omni[metric])
                if raw.get(metric) is not None and omni.get(metric) is not None
                else None
            )
        row["raw_exceeds_spider_contact"] = bool(
            row.get("delta_raw_vs_spider_hand_object_contact") is not None
            and row["delta_raw_vs_spider_hand_object_contact"] > 0
        )
        row["raw_exceeds_omni_contact"] = bool(
            row.get("delta_raw_vs_omni_hand_object_contact") is not None
            and row["delta_raw_vs_omni_hand_object_contact"] > 0
        )
        out.append(row)
    return out


def worst_case_rows(case_rows: list[dict[str, Any]], limit: int = 8) -> list[dict[str, Any]]:
    scored = [
        row
        for row in case_rows
        if row.get("spider_success_case") and row.get("delta_raw_vs_spider_hand_object_contact") is not None
    ]
    scored.sort(key=lambda row: float(row["delta_raw_vs_spider_hand_object_contact"]))
    fields = [
        "case_id",
        "e109_case_id",
        "object_key",
        "spider_cem_status",
        "ref_fk_hand_object_contact",
        "raw_mask_ref_fk_hand_object_contact",
        "delta_raw_vs_spider_hand_object_contact",
        "OmniRetarget_hand_object_contact",
        "delta_raw_vs_omni_hand_object_contact",
        "ref_fk_hand_near_5cm",
        "raw_mask_ref_fk_hand_near_5cm",
        "delta_raw_vs_spider_hand_near_5cm",
        "ref_fk_hand_near_10cm",
        "raw_mask_ref_fk_hand_near_10cm",
        "delta_raw_vs_spider_hand_near_10cm",
        "ref_fk_leg_penetration",
        "raw_mask_ref_fk_leg_penetration",
        "delta_raw_vs_spider_leg_penetration",
        "raw_video_path",
        "raw_npz_path",
    ]
    return [{field: row.get(field) for field in fields} for row in scored[:limit]]


def concise_case_rows(case_rows: list[dict[str, Any]], include_paths: bool = False) -> list[dict[str, Any]]:
    metric_labels = [
        ("手物接触", "hand_object_contact"),
        ("5cm", "hand_near_5cm"),
        ("10cm", "hand_near_10cm"),
        ("手物穿透", "hand_object_penetration"),
        ("腿穿透", "leg_penetration"),
    ]
    rows = []
    for row in case_rows:
        out: dict[str, Any] = {
            "case_id": row.get("case_id"),
            "object_key": row.get("object_key"),
            "Spider状态": row.get("spider_cem_status"),
        }
        for label, metric in metric_labels:
            out[f"{label}_raw_mask_ref_fk"] = row.get(f"raw_mask_ref_fk_{metric}")
            out[f"{label}_OmniRetarget"] = row.get(f"OmniRetarget_{metric}")
            out[f"{label}_raw-Omni差值"] = row.get(f"delta_raw_vs_omni_{metric}")
        if include_paths:
            out["npz/视频位置"] = f"npz: {row.get('raw_npz_path', '')}\nvideo: {row.get('raw_video_path', '')}"
        rows.append(out)
    rows.sort(
        key=lambda row: (
            row.get("手物接触_raw-Omni差值") is None,
            -float(row.get("手物接触_raw-Omni差值") or 0.0),
        )
    )
    return rows


def write_xlsx(
    path: Path,
    summary_rows: list[dict[str, Any]],
    concise_case_rows_all: list[dict[str, Any]],
    spider_success_summary_rows: list[dict[str, Any]],
    concise_spider_success_case_rows: list[dict[str, Any]],
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "方法平均"
    sheets = [
        (ws, summary_rows),
        (wb.create_sheet("逐case对比"), concise_case_rows_all),
        (wb.create_sheet("Spider成功方法平均"), spider_success_summary_rows),
        (wb.create_sheet("Spider成功逐case"), concise_spider_success_case_rows),
    ]
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet, rows in sheets:
        fields = list(rows[0].keys()) if rows else []
        sheet.append(fields)
        for row in rows:
            sheet.append([row.get(field) for field in fields])
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for col_idx, field in enumerate(fields, start=1):
            width = max(12, min(48, max(len(str(field)), *(len(str(r.get(field, ""))) for r in rows)) + 2))
            sheet.column_dimensions[get_column_letter(col_idx)].width = width
            for cell in sheet[get_column_letter(col_idx)]:
                cell.font = Font(name="Arial")
                if isinstance(cell.value, float):
                    cell.number_format = "0.0%"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--stage", default="full", choices=["smoke", "full"])
    parser.add_argument("--allow-missing", action="store_true")
    parser.add_argument("--results-dir", type=Path, default=None)
    parser.add_argument("--out-dir", type=Path, default=OUT_ROOT)
    args = parser.parse_args()

    results_dir = repo_path_from_env("RESULTS", args.results_dir or REPO / f"workspace/core4d/results/E143/cem/{args.stage}")
    out_dir = repo_path(args.out_dir.as_posix())
    out_dir.mkdir(parents=True, exist_ok=True)
    variants = read_variants()
    spider_status = spider_status_by_case(variants)
    missing = missing_outputs(args.stage, variants, results_dir, args.allow_missing)
    modules = (load_module("eval_E090_for_E143", E090_EVAL), load_module("eval_E105_for_E143", E105_EVAL))

    baseline_rows = list(e109_metrics().values())
    raw_rows = []
    for row in variants:
        metrics = compute_raw_metrics(modules, row, args.stage, results_dir)
        if metrics is not None:
            raw_rows.append(metrics)
    all_rows = baseline_rows + raw_rows
    summary_rows = aggregate(all_rows)
    case_rows = joined_case_rows(all_rows, spider_status)
    spider_success_case_ids = {row["e109_case_id"] for row in case_rows if row.get("spider_success_case")}
    spider_success_rows = [row for row in all_rows if row["e109_case_id"] in spider_success_case_ids]
    spider_success_summary_rows = aggregate(spider_success_rows)
    spider_success_case_rows = [row for row in case_rows if row.get("spider_success_case")]
    spider_success_worst_rows = worst_case_rows(case_rows)
    concise_all_case_rows = concise_case_rows(case_rows)
    concise_spider_success_case_rows = concise_case_rows(spider_success_case_rows, include_paths=True)

    raw_fields = sorted({key for row in raw_rows for key in row}) if raw_rows else FIELDS
    case_fields = list(case_rows[0].keys()) if case_rows else []
    summary_fields = list(summary_rows[0].keys()) if summary_rows else []
    spider_success_case_fields = list(spider_success_case_rows[0].keys()) if spider_success_case_rows else case_fields
    spider_success_worst_fields = list(spider_success_worst_rows[0].keys()) if spider_success_worst_rows else []
    write_tsv(out_dir / "e143_method_summary.tsv", summary_rows, summary_fields)
    write_tsv(out_dir / "e143_case_by_case.tsv", case_rows, case_fields)
    write_tsv(out_dir / "e143_spider_success_method_summary.tsv", spider_success_summary_rows, summary_fields)
    write_tsv(out_dir / "e143_spider_success_case_by_case.tsv", spider_success_case_rows, spider_success_case_fields)
    write_tsv(out_dir / "e143_spider_success_worst_cases.tsv", spider_success_worst_rows, spider_success_worst_fields)
    write_tsv(out_dir / "e143_raw_mask_rows.tsv", raw_rows, raw_fields)
    missing_path = out_dir / "e143_missing_outputs.tsv"
    if missing:
        write_tsv(missing_path, missing, ["variant", "case_id", "missing"])
    elif missing_path.exists():
        missing_path.unlink()

    summary = {
        "stage": args.stage,
        "manifest_rows": len(variants),
        "raw_evaluated_rows": len(raw_rows),
        "missing_rows": len(missing),
        "case_rows": len(case_rows),
        "raw_exceeds_omni_contact_cases": sum(1 for row in case_rows if row.get("raw_exceeds_omni_contact")),
        "spider_success_case_rows": len(spider_success_case_rows),
        "spider_success_raw_exceeds_spider_contact_cases": sum(
            1
            for row in spider_success_case_rows
            if row.get("raw_mask_ref_fk_hand_object_contact") is not None
            and row.get("ref_fk_hand_object_contact") is not None
            and float(row["raw_mask_ref_fk_hand_object_contact"]) > float(row["ref_fk_hand_object_contact"])
        ),
        "spider_success_raw_exceeds_omni_contact_cases": sum(
            1
            for row in spider_success_case_rows
            if row.get("raw_mask_ref_fk_hand_object_contact") is not None
            and row.get("OmniRetarget_hand_object_contact") is not None
            and float(row["raw_mask_ref_fk_hand_object_contact"]) > float(row["OmniRetarget_hand_object_contact"])
        ),
        "method_summary_tsv": rel(out_dir / "e143_method_summary.tsv"),
        "case_by_case_tsv": rel(out_dir / "e143_case_by_case.tsv"),
        "spider_success_method_summary_tsv": rel(out_dir / "e143_spider_success_method_summary.tsv"),
        "spider_success_case_by_case_tsv": rel(out_dir / "e143_spider_success_case_by_case.tsv"),
        "spider_success_worst_cases_tsv": rel(out_dir / "e143_spider_success_worst_cases.tsv"),
        "raw_rows_tsv": rel(out_dir / "e143_raw_mask_rows.tsv"),
        "xlsx": rel(out_dir / "E143_raw_mask_ref_fk_24case_omniretarget_comparison.xlsx"),
    }
    (out_dir / "e143_summary.json").write_text(json.dumps(summary, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    lines = [
        "# E143 Raw-Mask Ref-FK 24-Case OmniRetarget Comparison",
        "",
        f"- stage: `{args.stage}`",
        f"- manifest rows: `{len(variants)}`",
        f"- raw evaluated rows: `{len(raw_rows)}`",
        f"- missing rows: `{len(missing)}`",
        f"- raw cases exceeding OmniRetarget hand-object contact: `{summary['raw_exceeds_omni_contact_cases']}`",
        f"- Spider success rows (`WORK/pass`): `{len(spider_success_case_rows)}`",
        f"- Spider success rows where raw exceeds Spider contact: `{summary['spider_success_raw_exceeds_spider_contact_cases']}`",
        f"- Spider success rows where raw exceeds OmniRetarget contact: `{summary['spider_success_raw_exceeds_omni_contact_cases']}`",
        "- xlsx case sheets compare only `raw_mask_ref_fk` and `OmniRetarget`, sorted by hand-object contact raw-Omni delta",
        "",
        "| method | cases | contact | 5cm | 10cm | hand penetration | leg penetration |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for row in summary_rows:
        lines.append(
            f"| `{row['method']}` | {row['case_count']} | "
            f"{float(row['hand_object_contact']) * 100:.1f}% | "
            f"{float(row['hand_near_5cm']) * 100:.1f}% | "
            f"{float(row['hand_near_10cm']) * 100:.1f}% | "
            f"{float(row['hand_object_penetration']) * 100:.1f}% | "
            f"{float(row['leg_penetration']) * 100:.1f}% |"
        )
    lines.extend(
        [
            "",
            "## Spider Success Cases Only",
            "",
            "| method | cases | contact | 5cm | 10cm | hand penetration | leg penetration |",
            "|---|---:|---:|---:|---:|---:|---:|",
        ]
    )
    for row in spider_success_summary_rows:
        lines.append(
            f"| `{row['method']}` | {row['case_count']} | "
            f"{float(row['hand_object_contact']) * 100:.1f}% | "
            f"{float(row['hand_near_5cm']) * 100:.1f}% | "
            f"{float(row['hand_near_10cm']) * 100:.1f}% | "
            f"{float(row['hand_object_penetration']) * 100:.1f}% | "
            f"{float(row['leg_penetration']) * 100:.1f}% |"
        )
    lines.extend(
        [
            "",
            "xlsx 只保留精简逐 case 表。`Spider成功逐case` 只比较 `raw_mask_ref_fk` 和 `OmniRetarget`，"
            "并按手物接触 `raw_mask_ref_fk - OmniRetarget` 差值从高到低排序。",
        ]
    )
    if missing:
        lines.extend(["", "## Missing Outputs", ""])
        for row in missing:
            lines.append(f"- `{row['variant']}`: {row['missing']}")
    (out_dir / "e143_summary.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    write_xlsx(
        out_dir / "E143_raw_mask_ref_fk_24case_omniretarget_comparison.xlsx",
        summary_rows,
        concise_all_case_rows,
        spider_success_summary_rows,
        concise_spider_success_case_rows,
    )
    print(f"wrote {rel(out_dir / 'e143_summary.md')}")
    print(f"wrote {rel(out_dir / 'E143_raw_mask_ref_fk_24case_omniretarget_comparison.xlsx')}")


if __name__ == "__main__":
    main()
