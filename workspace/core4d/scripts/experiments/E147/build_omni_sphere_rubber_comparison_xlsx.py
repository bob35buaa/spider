#!/usr/bin/env python3
"""Build E147 OmniRetarget vs sphere Spider vs rubber-hand Spider comparison xlsx."""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import math
from pathlib import Path
from typing import Any

import mujoco
import numpy as np
from scipy.spatial.transform import Rotation as R


REPO = Path(__file__).resolve().parents[4]
VARIANTS_TSV = REPO / "workspace/core4d/scripts/E147/variants.tsv"
E147_EVAL = REPO / "workspace/core4d/scripts/eval/eval_E147_rubber_hand_collision.py"
EVAL_ROOT = REPO / "workspace/core4d/results/E147/rubber_hand_collision/eval/full"
METHOD_METRICS_TSV = EVAL_ROOT / "e147_method_metrics.tsv"
PAIR_DELTA_TSV = EVAL_ROOT / "e147_pair_delta.tsv"
DOWNSTREAM_TSV = EVAL_ROOT / "e147_downstream_evidence_input.tsv"
OUT_DIR = REPO / "workspace/core4d/results/E147/rubber_hand_collision/comparison"

METHODS = [
    ("OmniRetarget", "OmniRetarget"),
    ("sphere5cm", "sphere Spider"),
    ("rubber_hull", "rubber hand Spider"),
]

METRICS = [
    ("手物接触", "hand_object_physics_contact_frac", "0.0%"),
    ("5cm", "hand_geom_near_5cm_frac", "0.0%"),
    ("10cm", "hand_geom_near_10cm_frac", "0.0%"),
    ("手物穿透", "hand_geom_penetration_frac", "0.0%"),
    ("深穿透2cm", "hand_geom_deep_penetration_2cm_frac", "0.0%"),
    ("腿穿透", "leg_penetration_frac", "0.0%"),
    ("body穿透", "body_penetration_frac", "0.0%"),
    ("物体触地", "object_floor_contact_frac", "0.0%"),
    ("pelvis_min_m", "pelvis_min_m", "0.000"),
]


def repo_path(text: str | Path) -> Path:
    p = Path(text)
    return p if p.is_absolute() else REPO / p


def rel(path: Path | str) -> str:
    p = Path(path)
    try:
        return str(p.resolve().relative_to(REPO.resolve()))
    except Exception:
        return str(path)


def finite(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        out = float(value)
    except Exception:
        return None
    if not math.isfinite(out):
        return None
    return out


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def write_tsv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: "" if row.get(field, "") is None else row.get(field, "") for field in fields})


def parse_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"true", "1", "yes", "pass"}


def load_eval_module():
    spec = importlib.util.spec_from_file_location("eval_E147_for_xlsx", E147_EVAL)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load {E147_EVAL}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def scene_euler_convention(scene_act: Path) -> str:
    meta = scene_act.parent / "scene_act_meta.json"
    if meta.is_file():
        return str(json.loads(meta.read_text(encoding="utf-8")).get("euler_convention", "XYZ"))
    return "XYZ"


def convert_freejoint_to_scene_act(qpos: np.ndarray, scene_act: Path) -> np.ndarray:
    model = mujoco.MjModel.from_xml_path(str(scene_act))
    if qpos.shape[1] == model.nq:
        return qpos.astype(np.float64, copy=True)
    nq_robot = model.nq - 6
    if qpos.shape[1] < nq_robot + 7:
        raise ValueError(f"cannot convert qpos shape={qpos.shape} for scene nq={model.nq}: {scene_act}")

    obj_body = mujoco.mj_name2id(model, mujoco.mjtObj.mjOBJ_BODY, "object")
    if obj_body < 0:
        raise ValueError(f"scene has no object body: {scene_act}")

    obj_pos_world = qpos[:, nq_robot : nq_robot + 3]
    obj_quat_wxyz = qpos[:, nq_robot + 3 : nq_robot + 7]
    body_pos = model.body_pos[obj_body]
    body_quat_wxyz = model.body_quat[obj_body]
    body_quat_xyzw = [body_quat_wxyz[1], body_quat_wxyz[2], body_quat_wxyz[3], body_quat_wxyz[0]]
    r_body = R.from_quat(body_quat_xyzw)

    obj_slide = r_body.inv().apply(obj_pos_world - body_pos[np.newaxis, :])
    obj_quat_xyzw = np.column_stack(
        [obj_quat_wxyz[:, 1], obj_quat_wxyz[:, 2], obj_quat_wxyz[:, 3], obj_quat_wxyz[:, 0]]
    )
    obj_euler = (r_body.inv() * R.from_quat(obj_quat_xyzw)).as_euler(scene_euler_convention(scene_act))

    out = np.zeros((qpos.shape[0], model.nq), dtype=np.float64)
    out[:, :nq_robot] = qpos[:, :nq_robot]
    out[:, nq_robot : nq_robot + 3] = obj_slide
    out[:, nq_robot + 3 : nq_robot + 6] = obj_euler
    return out


def ensure_omni_metrics(rows: list[dict[str, str]], out_dir: Path) -> list[dict[str, Any]]:
    eval_mod = load_eval_module()
    qpos_dir = out_dir / "omni_converted_qpos"
    qpos_dir.mkdir(parents=True, exist_ok=True)
    metrics: list[dict[str, Any]] = []
    for row in rows:
        trajectory = repo_path(row["trajectory"])
        scene_act = repo_path(row["base_scene_act"])
        data = np.load(trajectory, allow_pickle=True)
        qpos = np.asarray(data["qpos"], dtype=np.float64)
        qpos_converted = convert_freejoint_to_scene_act(qpos, scene_act)
        qpos_path = qpos_dir / f"{row['variant']}_omni_scene_act_qpos.npz"
        np.savez(qpos_path, qpos=qpos_converted)
        metric = eval_mod.evaluate_sequence(
            row=row,
            method="OmniRetarget kinematic replay",
            hand_collision_variant_id="OmniRetarget",
            qpos_path=qpos_path,
            scene_xml=scene_act,
        )
        metric["source_trajectory"] = rel(trajectory)
        metric["qpos_path"] = rel(qpos_path)
        metrics.append(metric)
    return metrics


def metric_lookup(method_rows: list[dict[str, Any]]) -> dict[tuple[str, str], dict[str, Any]]:
    out: dict[tuple[str, str], dict[str, Any]] = {}
    for row in method_rows:
        out[(str(row["case_id"]), str(row["hand_collision_variant_id"]))] = row
    return out


def build_case_rows(
    variants: list[dict[str, str]],
    method_rows: list[dict[str, Any]],
    pair_rows: list[dict[str, str]],
    downstream_rows: list[dict[str, str]],
) -> tuple[list[dict[str, Any]], list[str]]:
    metrics = metric_lookup(method_rows)
    pairs = {row["case_id"]: row for row in pair_rows}
    downstream = {row["case_id"]: row for row in downstream_rows}
    fields = [
        "case_id",
        "object_key",
        "expected_quality",
        "sphere_historical_cem_status",
        "rubber_cem_status",
        "rubber_failure_mode",
        "rubber_ab_status",
    ]
    for label, _metric, _fmt in METRICS:
        fields.extend(
            [
                f"{label}_OmniRetarget",
                f"{label}_sphere_spider",
                f"{label}_rubber_hand_spider",
                f"{label}_rubber-sphere",
                f"{label}_rubber-Omni",
            ]
        )
    fields.extend(
        [
            "fall_OmniRetarget",
            "fall_sphere_spider",
            "fall_rubber_hand_spider",
            "omni_qpos_path",
            "sphere_npz",
            "rubber_npz",
            "rubber_video",
        ]
    )

    out = []
    for row in variants:
        case_id = row["case_id"]
        omni = metrics[(case_id, "OmniRetarget")]
        sphere = metrics[(case_id, "sphere5cm")]
        rubber = metrics[(case_id, "rubber_hull")]
        pair = pairs.get(case_id, {})
        ev = downstream.get(case_id, {})
        item: dict[str, Any] = {
            "case_id": case_id,
            "object_key": row["object_key"],
            "expected_quality": row["expected_quality"],
            "sphere_historical_cem_status": row["historical_cem_status"],
            "rubber_cem_status": ev.get("cem_status", ""),
            "rubber_failure_mode": ev.get("downstream_failure_mode", ""),
            "rubber_ab_status": pair.get("ab_status", ""),
            "fall_OmniRetarget": parse_bool(omni.get("fall_flag")),
            "fall_sphere_spider": parse_bool(sphere.get("fall_flag")),
            "fall_rubber_hand_spider": parse_bool(rubber.get("fall_flag")),
            "omni_qpos_path": omni.get("qpos_path", ""),
            "sphere_npz": sphere.get("qpos_path", ""),
            "rubber_npz": rubber.get("qpos_path", ""),
            "rubber_video": ev.get("cem_video", ""),
        }
        for label, metric, _fmt in METRICS:
            item[f"{label}_OmniRetarget"] = finite(omni.get(metric))
            item[f"{label}_sphere_spider"] = finite(sphere.get(metric))
            item[f"{label}_rubber_hand_spider"] = finite(rubber.get(metric))
            item[f"{label}_rubber-sphere"] = None
            item[f"{label}_rubber-Omni"] = None
        out.append(item)
    return out, fields


def write_xlsx(path: Path, case_rows: list[dict[str, Any]], case_fields: list[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter, quote_sheetname

    wb = Workbook()
    ws_avg = wb.active
    ws_avg.title = "10case平均"
    ws_case = wb.create_sheet("逐case对比")

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    delta_fill = PatternFill("solid", fgColor="FCE4D6")
    path_fill = PatternFill("solid", fgColor="E2F0D9")

    ws_case.append(case_fields)
    for row in case_rows:
        ws_case.append([row.get(field) for field in case_fields])

    field_to_col = {field: idx + 1 for idx, field in enumerate(case_fields)}
    row_start = 2
    row_end = len(case_rows) + 1
    case_sheet = quote_sheetname(ws_case.title)

    for r in range(row_start, row_end + 1):
        for label, _metric, _fmt in METRICS:
            rub = get_column_letter(field_to_col[f"{label}_rubber_hand_spider"])
            sph = get_column_letter(field_to_col[f"{label}_sphere_spider"])
            omni = get_column_letter(field_to_col[f"{label}_OmniRetarget"])
            ws_case.cell(r, field_to_col[f"{label}_rubber-sphere"]).value = f"={rub}{r}-{sph}{r}"
            ws_case.cell(r, field_to_col[f"{label}_rubber-Omni"]).value = f"={rub}{r}-{omni}{r}"

    avg_fields = [
        "method",
        "case_count",
        "cem_pass_count",
        "fall_count",
        "hand_object_contact",
        "hand_near_5cm",
        "hand_near_10cm",
        "hand_object_penetration",
        "hand_deep_penetration_2cm",
        "leg_penetration",
        "body_penetration",
        "object_floor_contact",
        "pelvis_min_m",
        "notes",
    ]
    ws_avg.append(avg_fields)
    method_col_map = {
        "OmniRetarget": "OmniRetarget",
        "sphere Spider": "sphere_spider",
        "rubber hand Spider": "rubber_hand_spider",
    }
    metric_avg_map = {
        "hand_object_contact": "手物接触",
        "hand_near_5cm": "5cm",
        "hand_near_10cm": "10cm",
        "hand_object_penetration": "手物穿透",
        "hand_deep_penetration_2cm": "深穿透2cm",
        "leg_penetration": "腿穿透",
        "body_penetration": "body穿透",
        "object_floor_contact": "物体触地",
        "pelvis_min_m": "pelvis_min_m",
    }
    notes = {
        "OmniRetarget": "kinematic replay converted from trajectory_kinematic.npz",
        "sphere Spider": "historical CEM with sphere5cm hand collision",
        "rubber hand Spider": "E147 full CEM with rubber_hull hand collision",
    }
    for row_idx, method in enumerate(["OmniRetarget", "sphere Spider", "rubber hand Spider"], start=2):
        key = method_col_map[method]
        ws_avg.cell(row_idx, 1).value = method
        ws_avg.cell(row_idx, 2).value = f"=COUNTA({case_sheet}!$A${row_start}:$A${row_end})"
        if method == "OmniRetarget":
            ws_avg.cell(row_idx, 3).value = ""
        elif method == "sphere Spider":
            col = get_column_letter(field_to_col["sphere_historical_cem_status"])
            ws_avg.cell(row_idx, 3).value = f'=COUNTIF({case_sheet}!{col}${row_start}:{col}${row_end},"pass")+COUNTIF({case_sheet}!{col}${row_start}:{col}${row_end},"WORK")'
        else:
            col = get_column_letter(field_to_col["rubber_cem_status"])
            ws_avg.cell(row_idx, 3).value = f'=COUNTIF({case_sheet}!{col}${row_start}:{col}${row_end},"pass")'
        fall_col = get_column_letter(field_to_col[f"fall_{key}"])
        ws_avg.cell(row_idx, 4).value = f"=COUNTIF({case_sheet}!{fall_col}${row_start}:{fall_col}${row_end},TRUE)"
        for col_idx, avg_field in enumerate(avg_fields[4:-1], start=5):
            label = metric_avg_map[avg_field]
            metric_col = get_column_letter(field_to_col[f"{label}_{key}"])
            ws_avg.cell(row_idx, col_idx).value = f"=AVERAGE({case_sheet}!{metric_col}${row_start}:{metric_col}${row_end})"
        ws_avg.cell(row_idx, len(avg_fields)).value = notes[method]

    for ws in (ws_avg, ws_case):
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            header = str(ws.cell(1, col_idx).value or "")
            max_len = max(len(header), *(len(str(ws.cell(r, col_idx).value or "")) for r in range(2, min(ws.max_row, 30) + 1)))
            ws.column_dimensions[letter].width = max(10, min(42, max_len + 2))
            if "rubber-" in header:
                ws.cell(1, col_idx).fill = delta_fill
            if header.endswith("_path") or header.endswith("_npz") or header.endswith("_video"):
                ws.cell(1, col_idx).fill = path_fill
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row, col_idx)
                cell.font = Font(name="Arial")
                if isinstance(cell.value, float) or (isinstance(cell.value, str) and cell.value.startswith("=AVERAGE")):
                    if "pelvis_min_m" in header:
                        cell.number_format = "0.000"
                    else:
                        cell.number_format = "0.0%"
                cell.alignment = Alignment(vertical="top", wrap_text=header.endswith("_path") or header.endswith("_npz") or header.endswith("_video"))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out-dir", type=Path, default=OUT_DIR)
    args = parser.parse_args()
    out_dir = repo_path(args.out_dir)
    variants = read_tsv(VARIANTS_TSV)
    existing_metrics = read_tsv(METHOD_METRICS_TSV)
    omni_metrics = ensure_omni_metrics(variants, out_dir)
    method_rows: list[dict[str, Any]] = [*existing_metrics, *omni_metrics]
    pair_rows = read_tsv(PAIR_DELTA_TSV)
    downstream_rows = read_tsv(DOWNSTREAM_TSV)
    case_rows, case_fields = build_case_rows(variants, method_rows, pair_rows, downstream_rows)

    xlsx_path = out_dir / "E147_omni_sphere_rubber_hand_comparison.xlsx"
    write_xlsx(xlsx_path, case_rows, case_fields)
    write_tsv(out_dir / "e147_omni_sphere_rubber_case_comparison.tsv", case_rows, case_fields)
    print(rel(xlsx_path))


if __name__ == "__main__":
    main()
