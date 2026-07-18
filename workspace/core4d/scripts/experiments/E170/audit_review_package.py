#!/usr/bin/env python3
"""Audit E170 automated, Codex-review, and final-user package readiness.

This audit intentionally stops short of result-log/tracker completion.  It
proves plan items 1--6 for the review package and keeps three authorities
separate:

* ``postprocess``: automated artifacts/evaluation/video/xlsx are complete;
* ``pre_user``: Codex metric verification and selected visual review are done;
* ``final``: the user has supplied fresh 28/28 labels.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import yaml


REPO = Path(__file__).resolve().parents[5]
ROOT = REPO / "workspace/core4d/results/E170"
MANIFEST = ROOT / "s6_downstream/manifests/analysis_manifest.tsv"
EVAL = ROOT / "s6_downstream/eval/full"
KEYFRAMES = ROOT / "s6_downstream/evidence/visual_qc"
PAIR_ROOT = ROOT / "s6_downstream/render/full/paired"
AUDIT_ROOT = ROOT / "s6_downstream/evidence/completion"
DIAGNOSTICS = {
    "cem_leg_gate_valid_frac",
    "cem_leg_gate_selected_valid_frac",
    "cem_leg_gate_fallback_used",
    "sample_leg_gate_min_sdf_min",
    "sample_leg_gate_violation_pct_mean",
    "leg_object_penalty_mean",
}
PAIR_METRICS = {
    "body_z_err_p95_m",
    "track_pelvis_z_err_terminal_m",
    "hand_object_physics_contact_in_mask_frac",
    "hand_object_physics_contact_3mm_in_mask_frac",
    "hand_object_release_false_contact_3mm_frac",
    "hand_object_physics_penetration_3mm_frame_frac",
    "leg_penetration_frac",
    "leg_near_2cm_frac",
    "leg_object_physics_contact_frac",
    "track_root_pos_err_cm_mean",
    "track_root_ori_err_deg_mean",
    "track_eef_pos_err_cm_mean",
    "track_eef_ori_err_deg_mean",
    "track_obj_pos_err_cm_mean",
    "track_obj_ori_err_deg_mean",
    "qpos_accel_l2_p95",
    "qpos_jerk_l2_p95",
    "trackbody_jerk_p95",
    "ankle_jerk_p95",
    "obj_speed_max",
    "foot_slip_max_m",
}
KEYFRAME_EVENTS = {
    "pre_contact",
    "max_lower_body_penetration",
    "max_hand_penetration",
    "max_object_speed",
    "final",
}


def read_tsv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if not path.is_file() or path.stat().st_size == 0:
        return [], []
    with path.open("r", encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream, delimiter="\t")
        return list(reader.fieldnames or []), list(reader)


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8")) if path.is_file() else {}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def truthy(value: Any) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "on"}


class Audit:
    def __init__(self, stage: str) -> None:
        self.stage = stage
        self.checks: list[dict[str, Any]] = []

    def check(self, check_id: str, passed: bool, detail: str, evidence: str | Path) -> None:
        self.checks.append(
            {
                "check_id": check_id,
                "status": "pass" if passed else "fail",
                "detail": detail,
                "evidence": str(evidence),
            }
        )

    @property
    def passed(self) -> bool:
        return all(row["status"] == "pass" for row in self.checks)

    def write(self) -> Path:
        AUDIT_ROOT.mkdir(parents=True, exist_ok=True)
        path = AUDIT_ROOT / f"review_package_audit_{self.stage}.json"
        payload = {
            "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
            "stage": self.stage,
            "status": "pass" if self.passed else "fail",
            "passed": sum(row["status"] == "pass" for row in self.checks),
            "failed": sum(row["status"] == "fail" for row in self.checks),
            "checks": self.checks,
        }
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        print(json.dumps({key: payload[key] for key in ("stage", "status", "passed", "failed")}, sort_keys=True))
        return path


def audit_npz_contract(path: Path) -> list[str]:
    failures: list[str] = []
    if not path.is_file():
        return ["missing:outdir_npz"]
    try:
        with np.load(path, allow_pickle=True) as data:
            if "qpos" not in data.files or not np.isfinite(np.asarray(data["qpos"], dtype=float)).all():
                failures.append("qpos_missing_or_nonfinite")
            absent = DIAGNOSTICS - set(data.files)
            failures.extend(f"diag_missing:{key}" for key in sorted(absent))
            for key in sorted(DIAGNOSTICS - absent):
                try:
                    finite = np.isfinite(np.asarray(data[key], dtype=float)).all()
                except (TypeError, ValueError):
                    finite = False
                if not finite:
                    failures.append(f"diag_nonfinite:{key}")
    except Exception as exc:
        failures.append(f"load:{type(exc).__name__}:{exc}")
    return failures


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--stage", choices=("postprocess", "pre_user", "final"), default="postprocess")
    args = parser.parse_args()
    audit = Audit(args.stage)

    _, rows = read_tsv(MANIFEST)
    cases = [row.get("case_id", "") for row in rows]
    sources = Counter(row.get("execution_source", "") for row in rows)
    routes = Counter(row.get("retarget_variant_id", "") for row in rows)
    manifest_ok = (
        len(rows) == 28
        and len(set(cases)) == 28
        and sources == Counter({"E170": 24, "E169": 4})
        and routes == Counter({"omnirt_v1": 25, "omnirt_v2": 3})
        and all(truthy(row.get(key)) for row in rows for key in ("p_enabled", "r_enabled", "g_enabled"))
    )
    audit.check(
        "manifest_authority_28_4_24",
        manifest_ok,
        f"rows={len(rows)} unique={len(set(cases))} sources={dict(sources)} routes={dict(routes)}",
        MANIFEST,
    )

    for stage, expected in (("canary", 2), ("recovery_smoke", 2), ("full", 24), ("recovery_full", 2)):
        path = ROOT / f"s6_downstream/artifacts/{stage}/artifact_summary.json"
        summary = read_json(path)
        passed = (
            summary.get("status") == "pass"
            and summary.get("manifest_rows") == expected
            and summary.get("complete_rows") == expected
            and summary.get("artifact_files") == expected * 3
            and not summary.get("incomplete")
        )
        audit.check(
            f"artifact_summary_{stage}",
            passed,
            f"status={summary.get('status')} complete={summary.get('complete_rows')}/{summary.get('manifest_rows')} files={summary.get('artifact_files')}",
            path,
        )

    row_contract_failures: dict[str, list[str]] = {}
    for row in rows:
        failures: list[str] = []
        root_npz = REPO / row["result_npz"]
        outdir_npz = REPO / row["outdir_npz"]
        config_path = REPO / row["config_act"]
        scene_path = REPO / row["scene_act"]
        failures.extend(audit_npz_contract(outdir_npz))
        if not root_npz.is_file():
            failures.append("missing:result_npz")
        elif outdir_npz.is_file():
            try:
                with np.load(root_npz, allow_pickle=True) as root_data, np.load(outdir_npz, allow_pickle=True) as out_data:
                    if "qpos" not in root_data.files or "qpos" not in out_data.files or not np.array_equal(root_data["qpos"], out_data["qpos"]):
                        failures.append("root_outdir_qpos_mismatch")
            except Exception as exc:
                failures.append(f"root_outdir_load:{type(exc).__name__}:{exc}")
        if not config_path.is_file():
            failures.append("missing:config_act")
        else:
            try:
                config = yaml.safe_load(config_path.read_text(encoding="utf-8"))
                if (
                    config.get("scene_name") != row.get("scene_name")
                    or config.get("leg_object_penalty_scale") != 2.0
                    or config.get("cem_leg_gate_enabled") is not True
                ):
                    failures.append("effective_config")
            except Exception as exc:
                failures.append(f"config_load:{type(exc).__name__}:{exc}")
        if not scene_path.is_file():
            failures.append("missing:scene_act")
        elif sha256(scene_path) != row.get("effective_scene_sha256", ""):
            failures.append("sha:scene_act")
        for path_key, sha_key in (
            ("trajectory", "trajectory_sha256"),
            ("contact_mask", "contact_mask_sha256"),
            ("override_path", "override_sha256"),
        ):
            path = REPO / row[path_key]
            if not path.is_file():
                failures.append(f"missing:{path_key}")
            elif sha256(path) != row.get(sha_key, ""):
                failures.append(f"sha:{path_key}")
        if failures:
            row_contract_failures[row["case_id"]] = failures
    audit.check(
        "all_28_runtime_and_input_contracts",
        len(rows) == 28 and not row_contract_failures,
        f"rows={len(rows)} failures={row_contract_failures}",
        MANIFEST,
    )

    reuse_failures: dict[str, list[str]] = {}
    for row in rows:
        if row.get("execution_source") != "E169":
            continue
        failures: list[str] = []
        checks = (
            ("result_npz", "source_result_sha256"),
            ("outdir_npz", "source_outdir_sha256"),
            ("config_act", "source_config_sha256"),
            ("video", "source_video_sha256"),
        )
        for path_key, sha_key in checks:
            path = REPO / row[path_key]
            if not path.is_file():
                failures.append(f"missing:{path_key}")
            elif sha256(path) != row.get(sha_key, ""):
                failures.append(f"sha:{path_key}")
        failures.extend(audit_npz_contract(REPO / row["outdir_npz"]))
        if failures:
            reuse_failures[row["case_id"]] = failures
    audit.check(
        "reuse_4_sha_and_npz_contract",
        sources.get("E169") == 4 and not reuse_failures,
        f"reuse_rows={sources.get('E169', 0)} failures={reuse_failures}",
        MANIFEST,
    )

    eval_summary_path = EVAL / "summary.json"
    eval_summary = read_json(eval_summary_path)
    counts = eval_summary.get("counts", {})
    eval_counts_ok = counts == {
        "manifest_rows": 28,
        "evaluated": 28,
        "not_ready": 0,
        "errors": 0,
        "numeric_pass": counts.get("numeric_pass"),
    }
    audit.check("unified_evaluation_28", eval_counts_ok, f"counts={counts}", eval_summary_path)

    metric_fields, metrics = read_tsv(EVAL / "e170_case_metrics.tsv")
    paired_fields, paired = read_tsv(EVAL / "e170_paired_deltas.tsv")
    _, groups = read_tsv(EVAL / "e170_group_summary.tsv")
    _, worst = read_tsv(EVAL / "e170_worst_cases.tsv")
    _, not_ready = read_tsv(EVAL / "e170_not_ready.tsv")
    _, errors = read_tsv(EVAL / "e170_evaluation_errors.tsv")
    gate_fields = {
        "leg_gate_health_pass",
        "cem_leg_gate_valid_frac_mean",
        "cem_leg_gate_valid_frac_last_iter_mean",
        "cem_leg_gate_fallback_used_mean",
        "cem_leg_gate_selected_all_valid_mean",
    }
    tables_ok = (
        len(metrics) == 28
        and len({row.get("case_id", "") for row in metrics}) == 28
        and len(paired) == 28
        and not not_ready
        and not errors
        and bool(groups)
        and any(row.get("group_type") == "overall" for row in groups)
        and {row.get("metric", "") for row in worst} == PAIR_METRICS
        and PAIR_METRICS <= set(metric_fields)
        and {f"e168_{key}" for key in PAIR_METRICS} <= set(metric_fields)
        and gate_fields <= set(metric_fields)
        and {f"prg_{key}" for key in PAIR_METRICS} <= set(paired_fields)
        and {f"e168_{key}" for key in PAIR_METRICS} <= set(paired_fields)
        and {f"delta_{key}" for key in PAIR_METRICS} <= set(paired_fields)
    )
    audit.check(
        "evaluation_tables_complete",
        tables_ok,
        f"metrics={len(metrics)} paired={len(paired)} groups={len(groups)} worst_metrics={len({row.get('metric', '') for row in worst})} not_ready={len(not_ready)} errors={len(errors)}",
        EVAL,
    )

    new_videos = [REPO / row["video"] for row in rows if row.get("execution_source") == "E170"]
    reuse_videos = [REPO / row["video"] for row in rows if row.get("execution_source") == "E169"]
    paired_videos = [PAIR_ROOT / f"{case_id}_E168_vs_E170_PRG.mp4" for case_id in cases]
    videos_ok = (
        len(new_videos) == 24
        and len(reuse_videos) == 4
        and len(paired_videos) == 28
        and all(path.is_file() and path.stat().st_size > 0 for path in [*new_videos, *reuse_videos, *paired_videos])
    )
    audit.check(
        "videos_24_new_4_reuse_28_paired",
        videos_ok,
        f"new={sum(path.is_file() for path in new_videos)}/24 reuse={sum(path.is_file() for path in reuse_videos)}/4 paired={sum(path.is_file() for path in paired_videos)}/28",
        ROOT / "s6_downstream/render/full",
    )

    _, keyframes = read_tsv(KEYFRAMES / "keyframe_manifest.tsv")
    keyframe_summary = read_json(KEYFRAMES / "keyframe_summary.json")
    events_by_case: dict[str, list[str]] = {}
    keyframe_files_ok = True
    for row in keyframes:
        events_by_case.setdefault(row.get("case_id", ""), []).append(row.get("event", ""))
        for key in ("image_path", "sheet_path"):
            path = REPO / row.get(key, "")
            keyframe_files_ok &= path.is_file() and path.stat().st_size > 0
    keyframes_ok = (
        keyframe_summary.get("status") == "pass"
        and keyframe_summary.get("selected_cases") == 28
        and keyframe_summary.get("evidence_rows") == 140
        and len(keyframes) == 140
        and len(events_by_case) == 28
        and all(len(events) == 5 and set(events) == KEYFRAME_EVENTS for events in events_by_case.values())
        and all(row.get("mapping_status") == "exact_flattened_qpos_to_video_frame" for row in keyframes)
        and keyframe_files_ok
    )
    audit.check(
        "keyframes_28_x_5",
        keyframes_ok,
        f"summary={keyframe_summary.get('status')} cases={len(events_by_case)} rows={len(keyframes)} files_ok={keyframe_files_ok}",
        KEYFRAMES,
    )

    workbook = EVAL / "E170_box021_prg_full_validation.xlsx"
    recalc_path = EVAL / "xlsx_recalc_validation.json"
    recalc = read_json(recalc_path)
    workbook_shape: dict[str, int] = {}
    workbook_error = ""
    if workbook.is_file():
        try:
            from openpyxl import load_workbook

            book = load_workbook(workbook, read_only=True, data_only=False)
            for sheet_name in ("Case Metrics", "Manual Review", "Codex Verification", "Keyframe Evidence"):
                workbook_shape[sheet_name] = book[sheet_name].max_row - 1 if sheet_name in book.sheetnames else -1
            book.close()
        except Exception as exc:
            workbook_error = f"{type(exc).__name__}:{exc}"
    workbook_ok = (
        workbook.is_file()
        and workbook.stat().st_size > 0
        and recalc.get("status") == "success"
        and recalc.get("total_errors") == 0
        and int(recalc.get("total_formulas", 0)) > 0
        and workbook_shape
        == {
            "Case Metrics": 28,
            "Manual Review": 28,
            "Codex Verification": 28,
            "Keyframe Evidence": 140,
        }
        and not workbook_error
    )
    audit.check(
        "xlsx_recalculated_zero_errors",
        workbook_ok,
        f"xlsx={workbook.is_file()} recalc_status={recalc.get('status')} formula_errors={recalc.get('total_errors')} shape={workbook_shape} error={workbook_error}",
        recalc_path,
    )

    _, manual = read_tsv(EVAL / "user_manual_review_template.tsv")
    _, codex = read_tsv(EVAL / "codex_verification.tsv")
    manual_pending = (
        len(manual) == 28
        and all(row.get("user_manual_review_status") == "pending" for row in manual)
        and all(row.get("manual_use_decision") == "PENDING" for row in manual)
    )
    if args.stage in {"postprocess", "pre_user"}:
        audit.check(
            "user_authority_preserved_pending",
            manual_pending,
            f"rows={len(manual)} pending={sum(row.get('user_manual_review_status') == 'pending' for row in manual)}",
            EVAL / "user_manual_review_template.tsv",
        )
    else:
        valid_quality = {"NO_ISSUE", "MINOR_ACCEPTABLE", "MAJOR_ISSUE"}
        manual_final = (
            len(manual) == 28
            and all(row.get("user_manual_review_status") == "reviewed" for row in manual)
            and all(row.get("manual_use_decision") in {"USE", "DO_NOT_USE"} for row in manual)
            and all(row.get("manual_quality_label") in valid_quality for row in manual)
            and all(row.get("manual_reviewer") and row.get("manual_reviewed_at") for row in manual)
        )
        audit.check(
            "user_manual_review_28_final",
            manual_final,
            f"rows={len(manual)} reviewed={sum(row.get('user_manual_review_status') == 'reviewed' for row in manual)}",
            EVAL / "user_manual_review_template.tsv",
        )

    if args.stage in {"pre_user", "final"}:
        selected = [row for row in codex if row.get("codex_visual_selection_reasons")]
        codex_ok = (
            len(codex) == 28
            and all(not row.get("codex_metric_verification_status", "").startswith("PENDING") for row in codex)
            and bool(selected)
            and all(row.get("codex_visual_spotcheck_status") not in {"", "SELECTED_PENDING_REVIEW", "PENDING_FULL_EVALUATION"} for row in selected)
            and all(row.get("codex_visual_findings") and row.get("codex_visual_reviewer") and row.get("codex_visual_reviewed_at") for row in selected)
        )
        audit.check(
            "codex_metric_and_visual_review",
            codex_ok,
            f"rows={len(codex)} selected={len(selected)} reviewed_selected={sum(row.get('codex_visual_spotcheck_status') not in {'', 'SELECTED_PENDING_REVIEW', 'PENDING_FULL_EVALUATION'} for row in selected)}",
            EVAL / "codex_verification.tsv",
        )

    audit_path = audit.write()
    print(f"audit={audit_path}")
    return 0 if audit.passed else 2


if __name__ == "__main__":
    raise SystemExit(main())
