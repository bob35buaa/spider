#!/usr/bin/env python3
"""E200 RL-export case selection from the reviewed three-arm master workbook.

Reads the ``by_rollout`` sheet of ``E200_three_arm_master-review.xlsx`` and emits
three selection TSVs under ``results/E200/s6_downstream/rl_export/<version>/``:

  1. box001_only          -- box001 PRG+G1+A2, manual==USE                 (47)
  2. 5object_all          -- per-object arm, manual==USE                   (109)
  3. 5object_box001clean  -- box001 (USE & CLEAN) + 2 exemptions; rest USE (97)

Every version enforces hard numeric gates (total + orig/aug sub-counts). Any
mismatch aborts -- no approximation, no cherry-picking. This is selection only;
it does not resolve artifacts or partners (that is build_rl_export.py).
"""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

# Per-object export arm (the reviewer口径). box001/024/004 use PRG+G1+A2,
# box023 uses noPRG, box021 uses plain PRG.
ARM_BY_OBJECT = {
    "box001": "PRG+G1+A2",
    "box024": "PRG+G1+A2",
    "box004": "PRG+G1+A2",
    "box023": "noPRG",
    "box021": "PRG",
}

# v3 box001 exemptions: USE but manual_label != CLEAN, explicitly whitelisted.
EXEMPT_V3 = {
    ("box001_20231003_1_040_p2", "orig"),
    ("box001_20231003_1_040_p2", "trans1"),
}

# Hard numeric gates. (total, {object: (orig, aug)}) — sub-counts checked too.
EXPECTED = {
    "box001_only": {
        "total": 47,
        "by_object": {"box001": (10, 37)},
    },
    "5object_all": {
        "total": 109,
        "by_object": {
            "box001": (10, 37),
            "box024": (4, 12),
            "box004": (3, 0),
            "box023": (7, 12),
            "box021": (9, 15),
        },
    },
    "5object_box001clean": {
        "total": 97,
        "by_object": {
            "box001": (None, None),  # 35 total (33 clean + 2 exempt), split not fixed
            "box024": (4, 12),
            "box004": (3, 0),
            "box023": (7, 12),
            "box021": (9, 15),
        },
        "object_totals": {"box001": 35},
    },
}

SELECT_FIELDS = [
    "object",
    "arm",
    "case_id",
    "variant",
    "orig_or_aug",
    "manual",
    "manual_label",
]


def read_by_rollout(xlsx_path: Path) -> list[dict[str, str]]:
    """Load the ``by_rollout`` sheet as a list of dict rows (str-normalized)."""
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    if "by_rollout" not in workbook.sheetnames:
        raise SystemExit(f"sheet 'by_rollout' not found in {xlsx_path}")
    sheet = workbook["by_rollout"]
    it = sheet.iter_rows(values_only=True)
    header = [str(cell).strip() if cell is not None else "" for cell in next(it)]
    rows: list[dict[str, str]] = []
    for raw in it:
        cells = ["" if cell is None else str(cell).strip() for cell in raw]
        cells += [""] * (len(header) - len(cells))
        rows.append(dict(zip(header, cells, strict=False)))
    workbook.close()
    for col in ("arm", "object", "case_id", "variant", "manual", "manual_label"):
        if col not in header:
            raise SystemExit(f"required column '{col}' missing from by_rollout header")
    return rows


def orig_or_aug(variant: str) -> str:
    """Map the aug variant label to the orig/aug bucket."""
    return "orig" if variant == "orig" else "aug"


def select(rows: list[dict[str, str]], version: str) -> list[dict[str, str]]:
    """Apply the arm-per-object + manual rules for one export version."""
    objects = (
        ["box001"]
        if version == "box001_only"
        else ["box001", "box024", "box004", "box023", "box021"]
    )
    out: list[dict[str, str]] = []
    for row in rows:
        obj = row["object"]
        if obj not in objects:
            continue
        if row["arm"] != ARM_BY_OBJECT[obj]:
            continue
        case_id, variant, manual, label = (
            row["case_id"],
            row["variant"],
            row["manual"],
            row["manual_label"],
        )
        keep = False
        if version == "5object_box001clean" and obj == "box001":
            keep = (manual == "USE" and label == "CLEAN") or (
                (case_id, variant) in EXEMPT_V3
            )
        else:
            keep = manual == "USE"
        if not keep:
            continue
        out.append(
            {
                "object": obj,
                "arm": row["arm"],
                "case_id": case_id,
                "variant": variant,
                "orig_or_aug": orig_or_aug(variant),
                "manual": manual,
                "manual_label": label,
            }
        )
    return out


def check_gates(version: str, selected: list[dict[str, str]]) -> None:
    """Enforce the hard total + orig/aug sub-count gates for one version."""
    spec = EXPECTED[version]
    total = len(selected)
    if total != spec["total"]:
        raise SystemExit(
            f"[{version}] total gate failed: expected {spec['total']}, got {total}"
        )
    # exemptions must all be present for the clean version
    if version == "5object_box001clean":
        present = {(r["case_id"], r["variant"]) for r in selected}
        missing = EXEMPT_V3 - present
        if missing:
            raise SystemExit(f"[{version}] missing exemption cases: {sorted(missing)}")
    for obj, (n_orig, n_aug) in spec["by_object"].items():
        rows = [r for r in selected if r["object"] == obj]
        got_orig = sum(1 for r in rows if r["orig_or_aug"] == "orig")
        got_aug = sum(1 for r in rows if r["orig_or_aug"] == "aug")
        obj_total_expected = spec.get("object_totals", {}).get(obj)
        if obj_total_expected is not None and len(rows) != obj_total_expected:
            raise SystemExit(
                f"[{version}] {obj} total gate failed: "
                f"expected {obj_total_expected}, got {len(rows)}"
            )
        if n_orig is not None and got_orig != n_orig:
            raise SystemExit(
                f"[{version}] {obj} orig gate failed: expected {n_orig}, got {got_orig}"
            )
        if n_aug is not None and got_aug != n_aug:
            raise SystemExit(
                f"[{version}] {obj} aug gate failed: expected {n_aug}, got {got_aug}"
            )


def write_selection(out_dir: Path, rows: list[dict[str, str]]) -> Path:
    """Write one selection.tsv (sorted for determinism)."""
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "selection.tsv"
    ordered = sorted(rows, key=lambda r: (r["object"], r["case_id"], r["variant"]))
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(
            stream, fieldnames=SELECT_FIELDS, delimiter="\t", lineterminator="\n"
        )
        writer.writeheader()
        writer.writerows(ordered)
    return path


def main() -> int:
    """Select, gate, and write the three RL-export versions."""
    parser = argparse.ArgumentParser()
    parser.add_argument("--review-xlsx", type=Path, required=True)
    parser.add_argument("--out-root", type=Path, required=True)
    args = parser.parse_args()

    rows = read_by_rollout(args.review_xlsx.expanduser())
    for version in ("box001_only", "5object_all", "5object_box001clean"):
        selected = select(rows, version)
        check_gates(version, selected)
        path = write_selection(args.out_root.expanduser() / version, selected)
        n_orig = sum(1 for r in selected if r["orig_or_aug"] == "orig")
        n_aug = len(selected) - n_orig
        print(f"[{version}] wrote {len(selected)} rows ({n_orig} orig / {n_aug} aug) -> {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
