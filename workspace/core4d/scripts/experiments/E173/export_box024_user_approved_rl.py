#!/usr/bin/env python3
"""Export the E173 Box024 rows approved by the final user review."""

from __future__ import annotations

import json
import subprocess
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from export_box001_user_approved_rl import (
    AUDIT_FIELDS,
    EVIDENCE_FIELDS,
    REVIEW_FIELDS,
    local_path,
    read_tsv,
    rel,
    require_file,
    sha256,
    unique,
    write_json,
    write_tsv,
)


REPO = Path(__file__).resolve().parents[5]
RUN_ROOT = REPO / "workspace/core4d/results/E173"
AUTHORITY = (
    RUN_ROOT
    / "s6_downstream/eval/full/E173_boxes_prg_full_validation_modified_new.xlsx"
)
AUTHORITY_SHEET = "Case Metrics"
MANUAL_TSV = RUN_ROOT / "s6_downstream/eval/full/user_manual_review_filled.tsv"
METRICS = RUN_ROOT / "s6_downstream/eval/full/e173_case_metrics.tsv"
CEM_MANIFEST = RUN_ROOT / "s6_downstream/manifests/cem_full_manifest.tsv"
HANDOFF = RUN_ROOT / "s5_handoff/handoff_manifest.tsv"
RL_DIR = RUN_ROOT / "s6_downstream/rl_export/box024_user_approved"
EVIDENCE_DIR = RUN_ROOT / "s6_downstream/evidence/box024_user_approved"
S6_SCRIPTS = (
    REPO / "workspace/core4d/scripts/data_construction_v3/stages/s6_downstream"
)
STAGE2B_MANIFESTS = {
    variant: RUN_ROOT
    / f"s3_retarget/{variant}/ref_fk/stage2b_manifest_{variant}_ref_fk.tsv"
    for variant in ("omnirt_v1", "omnirt_v2")
}


def read_authority() -> tuple[list[dict[str, Any]], set[str], set[str]]:
    workbook = load_workbook(AUTHORITY, data_only=True, read_only=True)
    sheet = workbook[AUTHORITY_SHEET]
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows))
    required = {
        "case_id",
        "object_key",
        "user_manual_review_status",
        "manual_use_decision",
        "manual_quality_label",
        "manual_failure_taxonomy",
        "manual_review_note",
        "manual_reviewer",
        "manual_reviewed_at",
    }
    missing = sorted(required - set(headers))
    if missing:
        raise SystemExit(f"authority fields missing: {missing}")
    selected = [
        dict(zip(headers, values))
        for values in rows
        if values[headers.index("case_id")]
        and str(values[headers.index("case_id")]).startswith("box024_")
    ]
    invalid = [
        str(row["case_id"])
        for row in selected
        if row.get("object_key") != "box024"
        or row.get("user_manual_review_status") != "reviewed"
        or row.get("manual_use_decision") not in {"USE", "DO_NOT_USE"}
    ]
    if len(selected) != 9 or len({str(row["case_id"]) for row in selected}) != 9:
        raise SystemExit("authority must contain exactly 9 unique Box024 rows")
    if invalid:
        raise SystemExit(f"invalid Box024 final reviews: {invalid}")
    approved = {
        str(row["case_id"])
        for row in selected
        if row["manual_use_decision"] == "USE"
    }
    rejected = {
        str(row["case_id"])
        for row in selected
        if row["manual_use_decision"] == "DO_NOT_USE"
    }
    if len(approved) != 3 or len(rejected) != 6:
        raise SystemExit(
            f"review drift: expected USE=3/DNU=6, got {len(approved)}/{len(rejected)}"
        )
    snapshot = [
        {
            **{field: row.get(field, "") or "" for field in REVIEW_FIELDS[:9]},
            "authority_workbook": rel(AUTHORITY),
            "authority_sheet": AUTHORITY_SHEET,
        }
        for row in selected
    ]
    return snapshot, approved, rejected


def restore_trajectory(
    case_id: str,
    cem: dict[str, str],
    metric: dict[str, str],
    stage2b: dict[str, str],
) -> Path:
    trajectory = local_path(metric["trajectory"])
    expected_sha = cem.get("trajectory_sha256", "")
    if trajectory.is_file() and trajectory.stat().st_size > 0:
        if expected_sha and sha256(trajectory) != expected_sha:
            raise SystemExit(f"trajectory SHA mismatch: {case_id}")
        return trajectory

    trimmed = require_file(stage2b["trimmed_npz"], f"{case_id} trimmed_npz")
    task = cem["target_task"]
    person = cem["person"]
    source_scene = require_file(
        REPO
        / "example_datasets/processed/core4d/unitree_g1/humanoid_object"
        / f"box024_{person}/scene.xml",
        f"{case_id} source scene",
    )
    subprocess.run(
        [
            str(REPO / ".venv/bin/python"),
            str(
                REPO
                / "workspace/core4d/data_preprocess/create_spider_scene_from_template.py"
            ),
            "--source-scene",
            str(source_scene),
            "--task",
            task,
            "--qpos",
            str(trimmed),
            "--data-id",
            "0",
            "--date",
            cem["date"],
            "--seq",
            cem["seq"],
            "--person",
            person,
            "--object-name",
            "box024",
            "--object-model-rel",
            "box/box024_m.obj",
        ],
        cwd=REPO,
        check=True,
    )
    subprocess.run(
        [
            str(REPO / ".venv/bin/python"),
            str(REPO / "spider/process_datasets/core4d.py"),
            "--source-npz",
            str(trimmed),
            "--task",
            task,
            "--data-id",
            "0",
            "--no-show-viewer",
            "--no-save-video",
        ],
        cwd=REPO,
        check=True,
    )
    trajectory = require_file(trajectory, f"{case_id} restored trajectory")
    if expected_sha and sha256(trajectory) != expected_sha:
        raise SystemExit(f"restored trajectory SHA mismatch: {case_id}")
    return trajectory


def main() -> int:
    snapshot, approved, rejected = read_authority()
    review_by_case = unique(
        [{key: str(value) for key, value in row.items()} for row in snapshot],
        "authority snapshot",
    )
    _, metrics_rows = read_tsv(METRICS)
    metric_by_case = unique(
        [row for row in metrics_rows if row.get("object_key") == "box024"],
        "E173 Box024 metrics",
    )
    handoff_fields, handoff_rows = read_tsv(HANDOFF)
    _, cem_rows = read_tsv(CEM_MANIFEST)
    cem_by_case = unique(
        [row for row in cem_rows if row.get("object_key") == "box024"],
        "E173 Box024 CEM manifest",
    )
    stage2b_by_variant: dict[str, dict[str, dict[str, str]]] = {}
    for variant, manifest in STAGE2B_MANIFESTS.items():
        _, rows = read_tsv(manifest)
        stage2b_by_variant[variant] = unique(rows, f"E173 {variant} Stage2b")

    authority_cases = approved | rejected
    if set(metric_by_case) != authority_cases or set(cem_by_case) != authority_cases:
        raise SystemExit("Box024 authority/metrics/CEM sets do not match")

    RL_DIR.mkdir(parents=True, exist_ok=True)
    EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)
    review_snapshot = RL_DIR / "box024_manual_review_snapshot.tsv"
    write_tsv(review_snapshot, snapshot, REVIEW_FIELDS)

    scoped_handoff: list[dict[str, Any]] = []
    evidence_input: list[dict[str, Any]] = []
    source_audit: list[dict[str, Any]] = []
    for case_id in sorted(approved):
        review = review_by_case[case_id]
        metric = metric_by_case[case_id]
        cem = cem_by_case[case_id]
        variant = cem["retarget_variant_id"]
        stage2b = stage2b_by_variant[variant].get(case_id)
        if not stage2b or stage2b.get("stage2b_status") != "pass":
            raise SystemExit(f"selected Stage2b row is not pass: {case_id}")
        matches = [
            row
            for row in handoff_rows
            if row.get("case_id") == case_id
            and row.get("retarget_variant_id") == variant
            and row.get("target_variant_id") == cem["target_variant_id"]
            and row.get("handoff_decision") == "HANDOFF_READY"
            and row.get("target_gate_status") == "pass"
            and row.get("visual_qc_status") == "pass"
        ]
        if len(matches) != 1:
            raise SystemExit(f"expected one ready S5 handoff row for {case_id}")
        base = matches[0]
        trajectory = restore_trajectory(
            case_id,
            cem,
            metric,
            stage2b,
        )
        paths = {
            "scene_act": require_file(cem["scene_act"], f"{case_id} scene_act"),
            "trajectory": trajectory,
            "contact_mask": require_file(
                metric["contact_mask"], f"{case_id} contact_mask"
            ),
            "cem_result_npz": require_file(
                metric["result_npz"], f"{case_id} CEM result"
            ),
            "cem_video": require_file(metric["video"], f"{case_id} CEM video"),
            "config_act": require_file(
                metric["config_act"], f"{case_id} config_act"
            ),
        }
        if (
            cem.get("effective_scene_sha256")
            and sha256(paths["scene_act"]) != cem["effective_scene_sha256"]
        ):
            raise SystemExit(f"scene_act SHA mismatch: {case_id}")
        if (
            cem.get("contact_mask_sha256")
            and sha256(paths["contact_mask"]) != cem["contact_mask_sha256"]
        ):
            raise SystemExit(f"contact_mask SHA mismatch: {case_id}")

        notes = (
            f"E173 final user review={review['manual_use_decision']}/"
            f"{review['manual_quality_label']}; numeric_pass="
            f"{metric.get('numeric_release_pass', '')}; numeric_warnings="
            f"{metric.get('numeric_failure_modes', '') or 'none'}; "
            "manual operational approval does not overwrite numeric or gate-health facts"
        )
        task_dir = paths["trajectory"].parents[1]
        scoped_handoff.append(
            {
                **base,
                "short_case_id": case_id,
                "hand_collision_variant_id": cem["hand_collision_variant_id"],
                "source_exp_id": "E173",
                "spider_method_id": cem["spider_method_id"],
                "source_exp": "E173",
                "source_variant": cem["variant"],
                "source_metrics_ref": rel(METRICS),
                "source_review_ref": rel(AUTHORITY),
                "target_scene": rel(task_dir / "scene.xml"),
                "trajectory": rel(paths["trajectory"]),
                "scene_act": rel(paths["scene_act"]),
                "contact_mask": rel(paths["contact_mask"]),
                "stage2b_target_task": cem["target_task"],
                "stage2b_manifest_ref": rel(STAGE2B_MANIFESTS[variant]),
                "raw_contact_threshold_label": cem["contact_mask_label"],
                "notes": notes,
            }
        )
        evidence_input.append(
            {
                "case_id": case_id,
                "object_key": "box024",
                "object_name": base["object_name"],
                "date": base["date"],
                "seq": base["seq"],
                "person": base["person"],
                "person_idx": base["person_idx"],
                "retarget_variant_id": variant,
                "target_variant_id": cem["target_variant_id"],
                "hand_collision_variant_id": cem["hand_collision_variant_id"],
                "source_exp_id": "E173",
                "spider_method_id": cem["spider_method_id"],
                "cem_status": "pass",
                "rl_status": "not_run",
                "downstream_failure_mode": "",
                "downstream_notes": notes,
                "cem_run_id": cem["variant"],
                "cem_result_npz": rel(paths["cem_result_npz"]),
                "cem_video": rel(paths["cem_video"]),
                "cem_metrics_ref": rel(METRICS),
            }
        )
        source_audit.append(
            {
                "case_id": case_id,
                "object_key": "box024",
                "manual_use_decision": review["manual_use_decision"],
                "manual_quality_label": review["manual_quality_label"],
                "manual_failure_taxonomy": review["manual_failure_taxonomy"],
                "manual_review_note": review["manual_review_note"],
                "numeric_release_pass": metric["numeric_release_pass"],
                "numeric_failure_modes": metric["numeric_failure_modes"],
                "retarget_variant_id": variant,
                "target_variant_id": cem["target_variant_id"],
                "hand_collision_variant_id": cem["hand_collision_variant_id"],
                "source_exp_id": "E173",
                "spider_method_id": cem["spider_method_id"],
                **{field: rel(path) for field, path in paths.items()},
                "scene_act_sha256": sha256(paths["scene_act"]),
                "trajectory_sha256": sha256(paths["trajectory"]),
                "contact_mask_sha256": sha256(paths["contact_mask"]),
                "cem_result_sha256": sha256(paths["cem_result_npz"]),
                "cem_video_sha256": sha256(paths["cem_video"]),
                "config_act_sha256": sha256(paths["config_act"]),
                "metrics_ref": rel(METRICS),
                "authority_ref": rel(AUTHORITY),
                "authority_sha256": sha256(AUTHORITY),
            }
        )

    extra_handoff_fields = [
        "short_case_id",
        "source_exp_id",
        "spider_method_id",
        "source_exp",
        "source_variant",
        "source_metrics_ref",
        "source_review_ref",
    ]
    scoped_handoff_path = RL_DIR / "box024_user_approved_handoff_manifest.tsv"
    evidence_input_path = EVIDENCE_DIR / "downstream_evidence_input.tsv"
    source_audit_path = RL_DIR / "box024_user_approved_source_rows.tsv"
    write_tsv(
        scoped_handoff_path,
        scoped_handoff,
        [
            *handoff_fields,
            *[
                field
                for field in extra_handoff_fields
                if field not in handoff_fields
            ],
        ],
    )
    write_tsv(evidence_input_path, evidence_input, EVIDENCE_FIELDS)
    write_tsv(source_audit_path, source_audit, AUDIT_FIELDS)

    subprocess.run(
        [
            sys.executable,
            str(S6_SCRIPTS / "record_downstream_evidence.py"),
            "--handoff-manifest-tsv",
            str(scoped_handoff_path),
            "--evidence-tsv",
            str(evidence_input_path),
            "--out-dir",
            str(EVIDENCE_DIR),
            "--evidence-root",
            str(RUN_ROOT),
            "--source-ref",
            "E173_box024_user_review_20260724",
        ],
        cwd=REPO,
        check=True,
    )
    evidence_manifest = EVIDENCE_DIR / "downstream_evidence_manifest.tsv"
    subprocess.run(
        [
            sys.executable,
            str(S6_SCRIPTS / "export_rl_inputs.py"),
            "--handoff-manifest-tsv",
            str(scoped_handoff_path),
            "--cem-evidence-tsv",
            str(evidence_manifest),
            "--out-dir",
            str(RL_DIR),
            "--spider-repo",
            str(REPO),
        ],
        cwd=REPO,
        check=True,
    )

    _, rl_rows = read_tsv(RL_DIR / "rl_export_input.tsv")
    rl_cases = {row["case_id"] for row in rl_rows}
    if len(rl_rows) != 3 or rl_cases != approved:
        raise SystemExit("RL export set is not exactly the 3 approved Box024 cases")
    if any(row["object_key"] != "box024" for row in rl_rows):
        raise SystemExit("non-Box024 row entered RL export")
    if any(row["rl_export_decision"] != "RL_EXPORT_READY" for row in rl_rows):
        raise SystemExit("not all approved Box024 rows are RL_EXPORT_READY")
    if rl_cases & rejected:
        raise SystemExit("DO_NOT_USE row entered RL export")
    for row in rl_rows:
        for field in ("scene_act", "trajectory", "contact_mask", "cem_result_npz"):
            require_file(row[field], f"{row['case_id']} RL {field}")

    audit = {
        "experiment": "E173",
        "scope": "Box024 only",
        "status": "pass",
        "authority_workbook": rel(AUTHORITY),
        "authority_sheet": AUTHORITY_SHEET,
        "authority_sha256": sha256(AUTHORITY),
        "manual_review_tsv": rel(MANUAL_TSV),
        "manual_review_tsv_sha256": sha256(MANUAL_TSV),
        "authority_box024_rows": len(snapshot),
        "approved_rows": len(approved),
        "rejected_rows": len(rejected),
        "rl_ready_rows": len(rl_rows),
        "source_set_matches_approved": rl_cases == approved,
        "rejected_rows_in_export": sorted(rl_cases & rejected),
        "non_box024_rows": sorted(
            row["case_id"] for row in rl_rows if row["object_key"] != "box024"
        ),
        "numeric_release_counts": dict(
            Counter(row["numeric_release_pass"] for row in source_audit)
        ),
        "retarget_variant_counts": dict(
            Counter(row["retarget_variant_id"] for row in source_audit)
        ),
        "hand_collision_variant_counts": dict(
            Counter(row["hand_collision_variant_id"] for row in source_audit)
        ),
        "source_rows_sha256": sha256(source_audit_path),
        "rl_export_input_sha256": sha256(RL_DIR / "rl_export_input.tsv"),
    }
    write_json(RL_DIR / "box024_rl_export_audit.json", audit)
    summary = [
        "# E173 Box024 user-approved RL export",
        "",
        f"- authority: `{rel(AUTHORITY)}` / `{AUTHORITY_SHEET}.manual_use_decision`",
        f"- manual TSV: `{rel(MANUAL_TSV)}`",
        f"- reviewed Box024 rows: `{len(snapshot)}`",
        f"- approved / rejected: `{len(approved)} / {len(rejected)}`",
        f"- RL_EXPORT_READY: `{len(rl_rows)}/{len(approved)}`",
        "- object scope: `box024` only",
        f"- variants: `{audit['retarget_variant_counts']}`",
        f"- numeric release among approved: `{audit['numeric_release_counts']}`",
        "",
        "说明：人工 USE 是 operational allowlist；numeric/gate-health 事实保留在 source audit，不被人工标签覆盖。",
        "",
    ]
    (RL_DIR / "box024_rl_export_summary.md").write_text(
        "\n".join(summary), encoding="utf-8"
    )
    print(
        json.dumps(
            {
                "output": rel(RL_DIR / "rl_export_input.tsv"),
                "box024_reviewed": len(snapshot),
                "approved": len(approved),
                "rejected": len(rejected),
                "rl_ready": len(rl_rows),
                "status": "pass",
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
