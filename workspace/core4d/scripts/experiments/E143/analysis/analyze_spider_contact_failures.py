#!/usr/bin/env python3
"""Analyze annotated E143 OmniRetarget vs SPIDER contact failures."""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import cv2
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook


REPO = Path(__file__).resolve().parents[5]
DEFAULT_XLSX = REPO / "workspace/core4d/results/E143/raw_mask_ref_fk_24case_omniretarget_comparison/E143_raw_mask_ref_fk_24case_omniretarget_comparison-anno.xlsx"
DEFAULT_OUT = REPO / "workspace/core4d/results/E143/spider_contact_failure_analysis"
MOCAP_CMP_ROOT = REPO / "workspace/core4d/results/E143/mocap_omni_spider_cmp"


COLOR_LABELS = {
    "FFfdeada": "yellow_object_not_gt",
    "FFe6e0ec": "light_purple_box_rotation",
    "FFb3a2c7": "deep_purple_walk_up",
    "FFebf1de": "green_qualified",
    "FFd7e4bd": "green_qualified",
    "FFd9d9d9": "gray_not_expected",
}


CORE_METRICS = [
    "手物接触_raw_mask_ref_fk",
    "手物接触_OmniRetarget",
    "手物接触_raw-Omni差值",
    "5cm_raw_mask_ref_fk",
    "5cm_OmniRetarget",
    "5cm_raw-Omni差值",
    "10cm_raw_mask_ref_fk",
    "10cm_OmniRetarget",
    "10cm_raw-Omni差值",
    "手物穿透_raw_mask_ref_fk",
    "手物穿透_OmniRetarget",
    "手物穿透_raw-Omni差值",
    "腿穿透_raw_mask_ref_fk",
    "腿穿透_OmniRetarget",
    "腿穿透_raw-Omni差值",
]


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(REPO))
    except ValueError:
        return str(path)


def cell_rgb(cell: Any) -> str:
    color = cell.fill.fgColor
    rgb = color.rgb
    return str(rgb) if rgb else ""


def classify_from_note(note: str, color_label: str) -> tuple[str, str]:
    note = note or ""
    if color_label == "gray_not_expected" or "先不考虑" in note or "传递" in note:
        return "data_not_target_sequence", "非目标序列/传递动作，不应进入主比较集"
    if "优质case" in note or "还行" in note:
        return "spider_algorithm_gap_on_valid_case", "动捕与任务基本合格，差距主要看 SPIDER 接触保持"
    if "箱子轨迹不是gt" in note or "物体动捕抖动" in note or "箱子的轨迹落后" in note or "箱子抬起的高度不足" in note:
        return "mocap_or_object_quality", "物体轨迹/高度/抖动不合规，Spider 被错误目标牵引"
    if "箱子旋转" in note or "旋转180" in note or "旋转接近180" in note:
        return "mocap_or_object_quality", "物体旋转异常，接触点语义与可控轨迹不稳定"
    if "走上前" in note or "初始位置解算错误" in note or "接触点和动捕不一致" in note or "接触点很近" in note:
        return "omniretarget_or_reference_issue", "OmniRetarget/ref 接触点或时序语义已经偏离动捕"
    if "sim" in note or "摔倒" in note or "没有搬动" in note or "没有抬" in note:
        return "spider_algorithm_gap", "SPIDER 动力学执行没有跟上参考物体/接触状态"
    if "抖动" in note:
        return "mocap_or_object_quality", "动捕抖动影响接触判断"
    return "mixed_or_uncertain", "需要结合视频逐帧判断"


def parse_npz_video(text: str) -> tuple[str, str]:
    text = (text or "").replace("&#10;", "\n")
    npz = ""
    video = ""
    for line in text.splitlines():
        if line.startswith("npz:"):
            npz = line.split("npz:", 1)[1].strip()
        elif line.startswith("video:"):
            video = line.split("video:", 1)[1].strip()
    return npz, video


def load_annotations(xlsx_path: Path) -> tuple[list[dict[str, Any]], dict[str, str]]:
    wb = load_workbook(xlsx_path, data_only=False)
    ws = wb["Spider成功逐case"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows: list[dict[str, Any]] = []
    legend: dict[str, str] = {}
    for r in range(2, ws.max_row + 1):
        values = {headers[c - 1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
        case_id = values.get("case_id")
        note = values.get("备注") or ""
        first_color = cell_rgb(ws.cell(r, 1))
        note_color = cell_rgb(ws.cell(r, 3))
        color = note_color if note_color in COLOR_LABELS else first_color
        color_label = COLOR_LABELS.get(color, "")
        if case_id in ("黄色", "浅紫色", "深紫色", "绿色", "灰色"):
            legend[str(case_id)] = str(note)
            continue
        if not case_id:
            continue
        npz_path, video_path = parse_npz_video(str(values.get("npz/视频位置") or ""))
        category, category_reason = classify_from_note(str(note), color_label)
        row = {
            "case_id": str(case_id),
            "object_key": values.get("object_key") or "",
            "is_box_hold": values.get("是否抱箱子") or "",
            "note": str(note),
            "rl_success": values.get("RL是否成功") or "",
            "color_rgb": color,
            "color_label": color_label,
            "category": category,
            "category_reason": category_reason,
            "spider_npz": npz_path,
            "ref_sim_video": video_path,
            "mocap_ref_sim_video": rel(MOCAP_CMP_ROOT / f"{case_id}_mocap_omni_spider_cmp.mp4"),
        }
        for metric in CORE_METRICS:
            value = values.get(metric)
            row[metric] = float(value) if isinstance(value, (int, float)) else value
        rows.append(row)
    return rows, legend


def add_npz_diagnostics(rows: list[dict[str, Any]]) -> None:
    for row in rows:
        path = REPO / row["spider_npz"] if row["spider_npz"] else None
        if not path or not path.is_file():
            row.update({"npz_exists": False})
            continue
        data = np.load(path, allow_pickle=True)
        row["npz_exists"] = True
        for key in [
            "qpos",
            "trace_ref",
            "hold_contact_rew_mean",
            "hand_approach_rew_mean",
            "contact_hdmi_rew_mean",
            "task_obj_rew_mean",
            "leg_object_penalty_mean",
            "object_lift_rew_mean",
            "object_clearance_m_mean",
            "terminal_carry_gate_obj_rot_err_mean",
            "terminal_carry_gate_hand_near_frac_mean",
        ]:
            if key not in data.files:
                continue
            arr = np.asarray(data[key])
            row[f"{key}_shape"] = "x".join(str(x) for x in arr.shape)
            if arr.size and np.issubdtype(arr.dtype, np.number):
                row[f"{key}_mean"] = float(np.nanmean(arr))
                row[f"{key}_min"] = float(np.nanmin(arr))
                row[f"{key}_max"] = float(np.nanmax(arr))


def frame_triplet(video_path: Path, frame_idx: int, width: int = 420) -> np.ndarray | None:
    cap = cv2.VideoCapture(str(video_path))
    if not cap.isOpened():
        return None
    total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
    if total <= 0:
        cap.release()
        return None
    idx = max(0, min(total - 1, frame_idx))
    cap.set(cv2.CAP_PROP_POS_FRAMES, idx)
    ok, frame = cap.read()
    cap.release()
    if not ok:
        return None
    h, w = frame.shape[:2]
    target_h = int(round(h * (width / w)))
    return cv2.resize(frame, (width, target_h), interpolation=cv2.INTER_AREA)


def extract_case_frames(rows: list[dict[str, Any]], out_dir: Path) -> list[dict[str, str]]:
    fig_dir = out_dir / "figures" / "case_keyframes"
    fig_dir.mkdir(parents=True, exist_ok=True)
    exported = []
    representative = [
        "box021_035_p1",
        "box021_035_p2",
        "box023_person2",
        "box004_083_p1",
        "box026_134_p1",
        "box026_039_p2",
        "box026_138_p2",
        "box026_133_p1",
        "box026_137_p1",
        "box026_141_p1",
    ]
    by_case = {row["case_id"]: row for row in rows}
    for case_id in representative:
        row = by_case.get(case_id)
        if not row:
            continue
        video_path = REPO / row["mocap_ref_sim_video"]
        cap = cv2.VideoCapture(str(video_path))
        total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
        cap.release()
        if total <= 0:
            continue
        frames = []
        for frac in (0.15, 0.50, 0.85):
            img = frame_triplet(video_path, int(round((total - 1) * frac)), width=900)
            if img is not None:
                frames.append(img)
        if not frames:
            continue
        sheet = np.concatenate(frames, axis=0)
        safe = re.sub(r"[^A-Za-z0-9_.-]", "_", case_id)
        out_path = fig_dir / f"{safe}_keyframes.jpg"
        cv2.imwrite(str(out_path), sheet)
        row["keyframe_sheet"] = rel(out_path)
        exported.append({"case_id": case_id, "keyframe_sheet": rel(out_path)})
    return exported


def plot_summary_figures(rows: list[dict[str, Any]], out_dir: Path) -> list[dict[str, str]]:
    fig_dir = out_dir / "figures"
    fig_dir.mkdir(parents=True, exist_ok=True)
    outputs = []

    category_names = {
        "spider_algorithm_gap_on_valid_case": "SPIDER gap\n(valid)",
        "spider_algorithm_gap": "SPIDER gap\n(exec)",
        "mocap_or_object_quality": "mocap/object\nquality",
        "omniretarget_or_reference_issue": "Omni/ref\nissue",
        "data_not_target_sequence": "not target\nsequence",
        "mixed_or_uncertain": "mixed",
    }
    counts = Counter(row["category"] for row in rows)
    keys = list(counts.keys())
    plt.figure(figsize=(8, 4.5))
    plt.bar([category_names.get(k, k) for k in keys], [counts[k] for k in keys], color="#6b8fb3")
    plt.ylabel("case count")
    plt.title("Annotated root-cause groups")
    plt.tight_layout()
    out = fig_dir / "category_counts.png"
    plt.savefig(out, dpi=180)
    plt.close()
    outputs.append({"name": "category_counts", "path": rel(out)})

    sorted_rows = sorted(rows, key=lambda r: float(r["手物接触_raw-Omni差值"]))
    colors = {
        "spider_algorithm_gap_on_valid_case": "#4f9d69",
        "spider_algorithm_gap": "#d97941",
        "mocap_or_object_quality": "#9473b5",
        "omniretarget_or_reference_issue": "#5c6bc0",
        "data_not_target_sequence": "#9e9e9e",
        "mixed_or_uncertain": "#607d8b",
    }
    plt.figure(figsize=(11, 5.4))
    y = np.arange(len(sorted_rows))
    plt.barh(
        y,
        [float(r["手物接触_raw-Omni差值"]) for r in sorted_rows],
        color=[colors.get(r["category"], "#607d8b") for r in sorted_rows],
    )
    plt.yticks(y, [r["case_id"] for r in sorted_rows], fontsize=8)
    plt.axvline(0.0, color="black", linewidth=1)
    plt.xlabel("hand-object contact: SPIDER - OmniRetarget")
    plt.title("Per-case contact delta sorted by SPIDER deficit")
    plt.tight_layout()
    out = fig_dir / "contact_delta_by_case.png"
    plt.savefig(out, dpi=180)
    plt.close()
    outputs.append({"name": "contact_delta_by_case", "path": rel(out)})
    return outputs


def aggregate(rows: list[dict[str, Any]]) -> dict[str, Any]:
    counts = Counter(row["category"] for row in rows)
    colors = Counter(row["color_label"] or "none" for row in rows)
    by_category = {}
    for category, group in defaultdict(list, {k: [] for k in counts}).items():
        pass
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[row["category"]].append(row)
    for category, group in grouped.items():
        by_category[category] = {
            "n": len(group),
            "mean_contact_delta": float(np.mean([g["手物接触_raw-Omni差值"] for g in group])),
            "mean_near5_delta": float(np.mean([g["5cm_raw-Omni差值"] for g in group])),
            "mean_near10_delta": float(np.mean([g["10cm_raw-Omni差值"] for g in group])),
            "mean_leg_pen_delta": float(np.mean([g["腿穿透_raw-Omni差值"] for g in group])),
        }
    valid = [
        row
        for row in rows
        if row["category"] == "spider_algorithm_gap_on_valid_case"
        or row["color_label"] == "green_qualified"
        or "优质case" in row["note"]
        or row["note"] == "还行"
    ]
    return {
        "n_cases": len(rows),
        "category_counts": dict(counts),
        "color_counts": dict(colors),
        "by_category": by_category,
        "valid_like_cases": [row["case_id"] for row in valid],
        "valid_like_mean_contact_delta": float(np.mean([r["手物接触_raw-Omni差值"] for r in valid])) if valid else None,
        "valid_like_mean_near5_delta": float(np.mean([r["5cm_raw-Omni差值"] for r in valid])) if valid else None,
        "valid_like_mean_leg_pen_delta": float(np.mean([r["腿穿透_raw-Omni差值"] for r in valid])) if valid else None,
        "rl_success_cases": [row["case_id"] for row in rows if str(row["rl_success"]).strip() == "是"],
    }


def write_outputs(
    rows: list[dict[str, Any]],
    legend: dict[str, str],
    keyframes: list[dict[str, str]],
    summary_figures: list[dict[str, str]],
    out_dir: Path,
) -> None:
    data_dir = out_dir / "data"
    data_dir.mkdir(parents=True, exist_ok=True)
    fields = []
    seen = set()
    for row in rows:
        for key in row.keys():
            if key not in seen:
                seen.add(key)
                fields.append(key)
    with (data_dir / "case_analysis.tsv").open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    summary = aggregate(rows)
    summary["legend"] = legend
    summary["keyframes"] = keyframes
    summary["summary_figures"] = summary_figures
    (data_dir / "analysis_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    rows, legend = load_annotations(args.xlsx)
    add_npz_diagnostics(rows)
    keyframes = extract_case_frames(rows, args.out_dir)
    summary_figures = plot_summary_figures(rows, args.out_dir)
    write_outputs(rows, legend, keyframes, summary_figures, args.out_dir)
    summary = aggregate(rows)
    summary["summary_figures"] = summary_figures
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
